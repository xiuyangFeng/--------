"""Multi-case summary table (C14): one row per selected job, fixed columns, blanks where a value is absent.

Numbers are copied from ``summary.json`` exactly as stored (no unit conversion, no recomputation).
"""
from __future__ import annotations

import csv
import datetime as dt
import io
from typing import Any, Mapping

BRANCHES = ("主动脉", "左髂总", "左髂外", "左髂内", "右髂总", "右髂外", "右髂内")
BASE_COLUMNS = ("job_id", "case_id", "patient_id", "scan_label", "scan_date", "tags", "release_id", "family", "status",
                "review_status", "reviewer", "created_at", "input_sha256_12", "run_identity_12", "compute_total_s")
MORPHOLOGY_COLUMNS = ("max_diameter_mm", "max_diameter_s_mm", "sac_present", "sac_length_mm", "sac_volume_ml",
                      "neck_diameter_mm", "neck_length_mm", "lumen_volume_ml")
WALL_COLUMNS = ("wss_p99_pa", "wss_max_pa", "wss_mean_pa", "area_frac_low", "area_low_mm2", "area_frac_high", "area_high_mm2",
                "area_frac_very_high", "area_very_high_mm2", *(f"branch_p99_{b}" for b in BRANCHES), "population_percentile", "quality_grade")
VOLUME_COLUMNS = ("speed_p99_m_s", "speed_max_m_s", "pressure_min_pa", "pressure_max_pa", *(f"dp_{b}" for b in BRANCHES), "low_speed_regions")
TAIL_COLUMNS = ("findings_attention", "findings_note")
COLUMNS = (*BASE_COLUMNS, *MORPHOLOGY_COLUMNS, *WALL_COLUMNS, *VOLUME_COLUMNS, *TAIL_COLUMNS)


def _map(value: Any) -> Mapping[str, Any]:
    return value if isinstance(value, Mapping) else {}


def _num(value: Any):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return None if value != value else value


def family_of(job: Mapping[str, Any], summary: Mapping[str, Any]) -> str | None:
    fields = _map(summary.get("fields"))
    if fields:
        return "volume" if ("velocity" in fields or "pressure" in fields) else "wall"
    return job.get("family")


def morphology_columns(summary: Mapping[str, Any]) -> dict[str, Any]:
    """The §17.1 morphology columns; every value is copied from ``summary["morphology"]`` as stored."""
    morphology = _map(summary.get("morphology"))
    aorta = _map(morphology.get("aorta"))
    largest, sac, neck = _map(aorta.get("max")), _map(aorta.get("sac")), _map(aorta.get("neck"))
    present = sac.get("present")
    return {"max_diameter_mm": _num(largest.get("max_diameter_mm")),
            "max_diameter_s_mm": _num(largest.get("s_from_root_mm")),
            "sac_present": (1 if present else 0) if morphology else None,
            "sac_length_mm": _num(sac.get("length_mm")) if present else None,
            "sac_volume_ml": _num(sac.get("volume_ml")) if present else None,
            "neck_diameter_mm": _num(neck.get("diameter_mean_mm")) if neck.get("present") else None,
            "neck_length_mm": _num(neck.get("length_mm")) if neck.get("present") else None,
            "lumen_volume_ml": _num(morphology.get("lumen_volume_ml"))}


def build_row(job: Mapping[str, Any], summary: Mapping[str, Any] | None) -> dict[str, Any]:
    """One table row; ``summary`` is the parsed summary.json of a finished job (or None)."""
    summary = _map(summary)
    release = _map(job.get("model_release"))
    review = _map(job.get("review")) or _map(summary.get("review"))
    timing = _map(job.get("timing"))
    row: dict[str, Any] = {column: None for column in COLUMNS}
    row.update({
        "job_id": job.get("id"), "case_id": job.get("case_id") or summary.get("case_id"), "patient_id": job.get("patient_id"),
        "scan_label": job.get("scan_label"), "scan_date": job.get("scan_date"),
        "tags": ", ".join(str(t) for t in (job.get("tags") or [])) or None,
        "release_id": release.get("id") or summary.get("release"), "family": family_of(job, summary), "status": job.get("status"),
        "review_status": review.get("status") or "unreviewed", "reviewer": review.get("by") or None, "created_at": job.get("created_at"),
        "input_sha256_12": (str(job.get("input_sha256") or summary.get("input_sha256") or "")[:12] or None),
        "run_identity_12": (str(job.get("run_identity") or summary.get("run_identity") or "")[:12] or None),
        "compute_total_s": _num(timing.get("compute_s")) if timing.get("compute_s") is not None else _num(_map(summary.get("timing_s")).get("total")),
    })
    if not summary:
        return row
    row.update(morphology_columns(summary))
    peak, field = _map(summary.get("peak")), _map(summary.get("wss_field_pa"))
    if row["family"] == "wall":
        row.update({"wss_p99_pa": _num(peak.get("p99_pa")), "wss_max_pa": _num(peak.get("max_pa")), "wss_mean_pa": _num(field.get("mean")),
                    "area_frac_low": _num(field.get("area_frac_low")), "area_low_mm2": _num(field.get("area_low_mm2")),
                    "area_frac_high": _num(field.get("area_frac_high")), "area_high_mm2": _num(field.get("area_high_mm2")),
                    "area_frac_very_high": _num(field.get("area_frac_very_high")), "area_very_high_mm2": _num(field.get("area_very_high_mm2"))})
        for name, branch in _map(summary.get("per_branch")).items():
            if name in BRANCHES:
                row[f"branch_p99_{name}"] = _num(_map(branch).get("wss_p99_pa"))
        population = _map(_map(summary.get("reference_assessment")).get("population"))
        row["population_percentile"] = _num(population.get("percentile"))
        quality = _map(summary.get("quality"))
        row["quality_grade"] = quality.get("label") or quality.get("level") or None
    elif row["family"] == "volume":
        stats = _map(summary.get("volume_statistics"))
        speed, pressure = _map(stats.get("speed_m_s")), _map(stats.get("pressure_interior_pa"))
        row.update({"speed_p99_m_s": _num(speed.get("p99")), "speed_max_m_s": _num(speed.get("max")),
                    "pressure_min_pa": _num(pressure.get("min")), "pressure_max_pa": _num(pressure.get("max"))})
    items = [item for item in (_map(summary.get("findings")).get("items") or []) if isinstance(item, Mapping)]
    low_speed = 0
    for item in items:
        if item.get("kind") == "pressure_drop" and item.get("branch") in BRANCHES and row[f"dp_{item['branch']}"] is None:
            row[f"dp_{item['branch']}"] = _num(item.get("value"))
        if item.get("kind") == "low_speed_region":
            low_speed += 1
    if row["family"] == "volume":
        row["low_speed_regions"] = low_speed
    row["findings_attention"] = sum(1 for item in items if item.get("severity") == "attention")
    row["findings_note"] = sum(1 for item in items if item.get("severity") == "note")
    return row


def build_rows(records: list[tuple[Mapping[str, Any], Mapping[str, Any] | None]]) -> list[dict[str, Any]]:
    return [build_row(job, summary) for job, summary in records]


def to_csv(rows: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(COLUMNS)
    for row in rows:
        writer.writerow(["" if row.get(column) is None else row[column] for column in COLUMNS])
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def to_xlsx(rows: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(list(COLUMNS))
    for row in rows:
        sheet.append([row.get(column) for column in COLUMNS])
    sheet.freeze_panes = "A2"
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, min(40, len(column) + 4))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def population_block(release_dir: Any) -> dict[str, Any] | None:
    """The release's declared out-of-fold p99 population, or None when it declares none.

    Only the numbers a cohort histogram needs are returned; the full profile (protocol, caveats,
    validation) stays in ``reference.json`` and in each result's ``reference_assessment``.
    """
    import json
    from pathlib import Path
    sidecar = Path(release_dir) / "reference.json"
    if not sidecar.is_file():
        return None
    try:
        profile = json.loads(sidecar.read_text(encoding="utf-8")).get("population_reference")
    except (OSError, ValueError, TypeError):
        return None
    if not isinstance(profile, Mapping):
        return None
    values = [float(value) for value in (profile.get("values_pa") or [])
              if isinstance(value, (int, float)) and not isinstance(value, bool) and value == value]
    if not values:
        return None
    thresholds = [float(value) for value in (profile.get("thresholds_pa") or [])
                  if isinstance(value, (int, float)) and not isinstance(value, bool)]
    return {"metric": profile.get("metric") or "p99_pa", "values_pa": values, "case_count": len(values),
            "reference_set": profile.get("reference_set") or "", "thresholds_pa": thresholds or [0.4, 4.0, 7.0]}


def population_blocks(registry: Any, release_ids) -> dict[str, dict[str, Any]]:
    """Population references for the releases appearing in a cohort table; releases without one are omitted."""
    if registry is None:
        return {}
    out: dict[str, dict[str, Any]] = {}
    for release_id in dict.fromkeys(str(rid) for rid in release_ids if rid):
        if release_id in out:
            continue
        try:
            record = registry.describe(release_id)
        except Exception:
            continue
        block = population_block(getattr(record, "path", None) or "")
        if block:
            out[release_id] = block
    return out


def export_filename(fmt: str, when: dt.datetime | None = None) -> str:
    when = when or dt.datetime.now()
    return f"wss_summary_{when.strftime('%Y%m%d_%H%M')}.{fmt}"


__all__ = ["BRANCHES", "COLUMNS", "MORPHOLOGY_COLUMNS", "build_row", "build_rows", "export_filename", "family_of",
           "morphology_columns",
           "population_block", "population_blocks", "to_csv", "to_xlsx"]
