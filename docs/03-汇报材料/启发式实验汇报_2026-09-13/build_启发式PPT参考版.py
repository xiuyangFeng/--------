from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR
from PIL import Image

ROOT = Path('/public/newhome/cy/Digital_twin/GNN')
OUT = ROOT / 'docs/03-汇报材料/启发式实验汇报_2026-09-13/启发式实验汇报_参考版_2026-09-13.pptx'
ASSET = ROOT / 'docs/03-汇报材料/启发式实验汇报_2026-09-13/_启发式PPT参考版_assets'
ASSET.mkdir(exist_ok=True)

NAVY = RGBColor(22, 39, 67)
BLUE = RGBColor(39, 112, 181)
TEAL = RGBColor(39, 150, 146)
ORANGE = RGBColor(225, 126, 47)
RED = RGBColor(181, 70, 70)
INK = RGBColor(34, 40, 49)
GRAY = RGBColor(100, 112, 126)
LIGHT = RGBColor(241, 245, 249)
PALE_BLUE = RGBColor(226, 239, 250)
PALE_ORANGE = RGBColor(253, 239, 224)
PALE_RED = RGBColor(252, 232, 232)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def textbox(slide, text, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT,
            font='Noto Sans CJK SC', fill=None, line=None, margin=0.08, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid(); shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return shape

def title(slide, text, kicker=None):
    if kicker:
        textbox(slide, kicker.upper(), 0.55, 0.25, 12.2, 0.25, 9, BLUE, True)
    textbox(slide, text, 0.55, 0.52, 12.2, 0.52, 26, NAVY, True)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.12), Inches(12.25), Inches(0.025)).fill.solid()
    line = slide.shapes[-1]; line.fill.fore_color.rgb = BLUE; line.line.fill.background()

def footer(slide, text='参考稿｜正式汇报前请替换待补图并核对 run / checkpoint'):
    textbox(slide, text, 0.58, 7.18, 12.1, 0.18, 7.5, GRAY)

def add_img(slide, path, x, y, w=None, h=None, crop=False):
    path = str(path)
    if not Path(path).exists():
        return None
    try:
        im = Image.open(path)
        iw, ih = im.size
        if w and not h: h = w * ih / iw
        if h and not w: w = h * iw / ih
        return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    except Exception:
        return None

def rounded(slide, x, y, w, h, fill, line=None, radius=True):
    if line is None:
        line = fill
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    return sh

def arrow(slide, x1, y1, x2, y2, color=BLUE, width=2.2):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(width); ln.line.end_arrowhead = True
    return ln

def note_box(slide, text, x, y, w, h, fill=PALE_BLUE, color=NAVY):
    rounded(slide, x, y, w, h, fill, fill)
    textbox(slide, text, x+0.08, y+0.04, w-0.16, h-0.08, 13, color, True, valign=MSO_ANCHOR.MIDDLE)

def bullet_box(slide, heading, bullets, x, y, w, h, fill=LIGHT, accent=BLUE):
    rounded(slide, x, y, w, h, fill, fill)
    textbox(slide, heading, x+0.15, y+0.12, w-0.3, 0.28, 16, accent, True)
    txt = '\n'.join('• ' + b for b in bullets)
    textbox(slide, txt, x+0.15, y+0.48, w-0.3, h-0.58, 12.5, INK)

# 1 cover
s = prs.slides.add_slide(blank)
s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
textbox(s, 'WSS 预测实验：从“结果罗列”到“启发式证据链”', 0.75, 1.10, 11.9, 1.1, 30, RGBColor(255,255,255), True)
textbox(s, '三条路线｜模块作用｜失败原因｜下一步如何验证', 0.8, 2.45, 11.4, 0.55, 19, RGBColor(203,222,245))
for i,(lab,col,sub) in enumerate([('A 直接 WSS',TEAL,'wall → WSS'),('B 体场数据驱动',BLUE,'volume → u/p → WSS'),('C 体场 + 物理',ORANGE,'volume → u/p + physics → WSS')]):
    x = 0.95 + i*4.15
    rounded(s,x,4.1,3.55,1.35,col,col)
    textbox(s,lab,x+0.18,4.34,3.18,0.32,17,RGBColor(255,255,255),True,align=PP_ALIGN.CENTER)
    textbox(s,sub,x+0.18,4.78,3.18,0.26,13,RGBColor(255,255,255),False,align=PP_ALIGN.CENTER)
textbox(s,'参考版｜数据与图件截至 2026-09-13',0.82,6.85,8,0.3,11,RGBColor(180,195,215))

# 2 question / route map
s=prs.slides.add_slide(blank); title(s,'先把同一个问题拆成三条可比较的误差链','01 / 研究问题')
textbox(s,'目标不是列出所有实验，而是定位：哪个模块改变了什么？误差在哪一步被放大？',0.7,1.36,11.8,0.42,17,NAVY,True)
routes=[('A','直接壁面预测',TEAL,'壁面点 + 局部几何','网络直接输出 WSS','当前最稳定的 WSS 路线'),('B','数据驱动体场',BLUE,'体点 + 几何','输出速度/压力，再经冻结算子求 WSS','速度高分不代表 WSS 高分'),('C','体场 + 物理约束',ORANGE,'体点 + 几何','数据 loss + BC/PDE loss','需要先解释 loss 竞争')]
for i,(tag,name,col,inp,out,claim) in enumerate(routes):
    y=2.1+i*1.42
    rounded(s,0.8,y,0.62,0.62,col,col); textbox(s,tag,0.8,y+0.12,0.62,0.3,18,RGBColor(255,255,255),True,PP_ALIGN.CENTER)
    textbox(s,name,1.58,y+0.02,2.05,0.28,16,NAVY,True)
    note_box(s,inp,3.8,y-0.01,2.05,0.68,PALE_BLUE if i==1 else (PALE_ORANGE if i==2 else RGBColor(224,245,242)),NAVY)
    arrow(s,5.95,y+0.32,6.65,y+0.32,col)
    note_box(s,out,6.82,y-0.01,3.0,0.68,PALE_BLUE if i==1 else (PALE_ORANGE if i==2 else RGBColor(224,245,242)),NAVY)
    textbox(s,claim,10.15,y+0.08,2.4,0.48,13,col,True)
footer(s)

# 3 unified contract
s=prs.slides.add_slide(blank); title(s,'统一条件先说清楚，旧数据才不会和新结果混在一起','02 / 比较合同')
rows=[('数据母库','V5：172例；train138 / test34','历史 V1–V4：另列机制证据，不直接排冠军'),('时间帧','V5 峰值 step 1162','历史 V3 可能是 1146，图注单独写'),('主指标','R²_cb、逐病例 P10/中位数、Pa/m/s MAE','拟合 R² 只作校准诊断'),('选模','沿用各 run 冻结的 best/last 规则','不按 test 结果重新挑 checkpoint'),('最新代表','X5/X5X11；PF6/VF6','最新四组尚无同名后处理图')]
for i,(a,b,c) in enumerate(rows):
    y=1.55+i*0.78
    rounded(s,0.75,y,2.0,0.55,NAVY,NAVY); textbox(s,a,0.86,y+0.13,1.78,0.22,13,RGBColor(255,255,255),True,PP_ALIGN.CENTER)
    rounded(s,2.95,y,4.35,0.55,PALE_BLUE,PALE_BLUE); textbox(s,b,3.1,y+0.1,4.05,0.3,12.5,INK,True)
    rounded(s,7.55,y,5.0,0.55,LIGHT,LIGHT); textbox(s,c,7.7,y+0.1,4.7,0.3,12,GRAY)
textbox(s,'读图规则：正式指标来自原始同点预测；插值云图只用于展示。',0.88,5.77,11.5,0.42,17,RED,True,fill=PALE_RED,margin=0.18,valign=MSO_ANCHOR.MIDDLE)
footer(s)

# 4 current scoreboard
s=prs.slides.add_slide(blank); title(s,'现在能确认什么？三条路线的“现象”先分开','03 / 当前结果')
cards=[('A 直接 WSS','X5：R²_cb 0.7173 ± 0.004\nX5X11：0.7261（三 seed）','Murray 分支流量先验有配对增益；X5X11更像尾部候选',TEAL),('B 体场 data-only','PF6 压力：0.7894 ± 0.005\nVF6 速度：0.7964 ± 0.004','PF6 / VF6 已成为当前体场底座；尚未完成 VF6→WSS 闭环',BLUE),('C 物理约束','历史 V4：加入 BC/PDE 后\n数据拟合与 WSS 没有同步变好','已有 loss 与梯度证据；不能直接说方程错',ORANGE)]
for i,(head,val,desc,col) in enumerate(cards):
    x=0.78+i*4.15
    rounded(s,x,1.55,3.72,4.6,RGBColor(248,250,252),RGBColor(220,226,234))
    rounded(s,x,1.55,3.72,0.62,col,col); textbox(s,head,x+0.12,1.72,3.48,0.24,16,RGBColor(255,255,255),True,PP_ALIGN.CENTER)
    textbox(s,val,x+0.23,2.42,3.25,0.86,19,NAVY,True,PP_ALIGN.CENTER)
    textbox(s,desc,x+0.25,3.68,3.2,1.02,13,INK,False,PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
textbox(s,'这些数字说明“做到了什么”，下一页再用图解释“为什么”。',0.95,6.45,11.4,0.36,17,NAVY,True,PP_ALIGN.CENTER)
footer(s,'最新三 seed 指标来自 wave1b / wave2 / wave3；test34 曾参与开发筛选，不能称独立泛化验证')

# 5 network structure
s=prs.slides.add_slide(blank); title(s,'模块作用要绑定到网络结构，而不是只报模型名字','04 / A路线：直接 WSS')
net=ROOT/'docs/03-汇报材料/启发式实验汇报_2026-09-13/WSS_V5_R4_模型结构_20260909/01_整体网络结构.png'
add_img(s,net,0.62,1.42,7.45,4.2)
bullet_box(s,'结构 → 可能作用',[
    'LocalGeoPE：在局部邻域聚合时补充几何残差，帮助识别局部壁面形状',
    'SA2 local Transformer：扩大局部邻域内的信息交互',
    'SA3 / bottleneck context：提供长程上下文，但增益较小',
    'Murray F6：提供部署可得的分支流量先验，修正整体幅值与分支分配'
],8.35,1.52,4.2,3.45,fill=LIGHT,accent=TEAL)
note_box(s,'不要说“Transformer 有用”\n要说“它改善了哪一种误差”',8.35,5.25,4.2,0.86,RGBColor(224,245,242),TEAL)
footer(s,'网络图为 V5-R4 结构图；X5/X5X11 是后续实验，尚无同名结构后处理包')

# 6 direct WSS evidence
s=prs.slides.add_slide(blank); title(s,'A路线的证据：F6先验带来稳定增益，但尾部仍低估','05 / A路线：模块证据')
bullet_box(s,'已证实',[
    'X5 相对 X0：三 seed 配对 ΔR²_cb ≈ +0.069',
    'X5 三 seed：0.7173 ± 0.004，波动小于早期 C1 单 seed',
    'X5X11：整体约 0.7261；high-WSS / p99 同步改善'
],0.72,1.48,4.0,2.35,fill=RGBColor(224,245,242),accent=TEAL)
bullet_box(s,'仍未解释',[
    'X5X11 的热点 IoU 并未同步提升',
    'X5X11 p99 预测/真值约 0.733 ± 0.011，极值仍系统低估',
    '需要最好 / 中位 / 最差病例云图，确认是幅值、位置还是局部结构问题'
],0.72,4.02,4.0,2.35,fill=PALE_ORANGE,accent=ORANGE)
img=ROOT/'docs/03-汇报材料/_archive/figures/WSS最小路线_20260712/06_几何特征增益典型病例_预测云图对比.png'
add_img(s,img,5.05,1.52,7.55,4.88)
textbox(s,'可复用的历史几何特征增益图｜正式稿应替换为 X5/X5X11 的最新病例三联图',5.18,6.47,7.25,0.25,10,GRAY,False,PP_ALIGN.CENTER)
footer(s)

# 7 latest postview placeholder
s=prs.slides.add_slide(blank); title(s,'A路线还差一张关键图：不要只放最好病例','06 / A路线：待补后处理')
for i,(lab,col,desc) in enumerate([('最好病例',TEAL,'预测 R² 最高'),('中位病例',BLUE,'接近病例分布中位数'),('最差病例',RED,'预测 R² 最低')]):
    x=0.8+i*4.18
    rounded(s,x,1.65,3.65,3.95,RGBColor(248,250,252),RGBColor(210,218,228))
    rounded(s,x,1.65,3.65,0.55,col,col); textbox(s,lab,x+0.1,1.81,3.45,0.22,15,RGBColor(255,255,255),True,PP_ALIGN.CENTER)
    textbox(s,'CFD  |  Pred  |  Pred−CFD',x+0.2,2.55,3.25,0.28,15,NAVY,True,PP_ALIGN.CENTER)
    textbox(s,'待补：X5/X5X11\n真实 Pa 三联图\n共用视角与色标',x+0.35,3.25,2.95,1.15,18,GRAY,True,PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    textbox(s,desc,x+0.2,5.1,3.25,0.22,11,GRAY,False,PP_ALIGN.CENTER)
textbox(s,'病例必须按正式预测 R² 选择，不能按拟合 R² 或“看起来最好”选择。',1.05,6.18,11.2,0.38,17,RED,True,PP_ALIGN.CENTER)
footer(s,'当前尚未发现 X5/X5X11 对应 PNG/VTP 包；此页是你后续需要补做的核心图')

# 8 volume field existing
s=prs.slides.add_slide(blank); title(s,'B路线：体场整体指标较好，但“整体”不是“局部”','07 / B路线：体场 data-only')
img=ROOT/'docs/03-汇报材料/启发式实验汇报_2026-09-13/WSS_V5_体场_最好最差与横向R2_20260909/field_best_worst_overview.png'
add_img(s,img,0.58,1.45,7.45,3.84)
bullet_box(s,'现有证据（R5 历史单 seed）',[
    '压力、速度整体 R² 高于直接 WSS',
    '可视化显示最好/最差病例的空间差异',
    '该包是 R5P/R5V，不是最新 PF6/VF6，只作展示方法参考'
],8.35,1.48,4.25,2.28,fill=LIGHT,accent=BLUE)
note_box(s,'最新 PF6 / VF6 三 seed 已有指标\n但还没有对应后处理图包',8.35,4.12,4.25,1.08,PALE_ORANGE,ORANGE)
textbox(s,'需要补：PF6/VF6 的 best / median / worst，CFD、Pred、signed error 共用色标。',8.35,5.48,4.25,0.65,14,NAVY,True,fill=PALE_BLUE,margin=0.16,valign=MSO_ANCHOR.MIDDLE)
footer(s,'现有图件：R5 体场最好/最差包；连续截面需真实体单元 VTU，点云/slab 不是连续 Slice')

# 9 speed to wss diagnostic
s=prs.slides.add_slide(blank); title(s,'为什么速度 R² 高，WSS 仍可能明显变差？','08 / B路线：误差链')
for i,(lab,col,sub) in enumerate([('预测速度',BLUE,'整体速度可较好'),('近壁结构',ORANGE,'法向/切向局部误差'),('梯度算子',RED,'误差被放大'),('WSS',NAVY,'最终指标下降')]):
    x=0.72+i*3.05
    rounded(s,x,1.55,2.45,1.0,col,col); textbox(s,lab,x+0.1,1.78,2.25,0.25,17,RGBColor(255,255,255),True,PP_ALIGN.CENTER); textbox(s,sub,x+0.1,2.13,2.25,0.2,10,RGBColor(255,255,255),False,PP_ALIGN.CENTER)
    if i<3: arrow(s,x+2.47,2.05,x+2.9,2.05,GRAY,1.8)
bullet_box(s,'已有旧 R5V 审计（机制证据，不是 VF6）',[
    '预测速度 → 冻结 Profile-Secant WSS：R²_cb = 0.491',
    'CFD 真值速度 → 同一算子：R²_cb = 0.965',
    '预测速度 → 未校准物理核：R²_cb = 0.468',
    '支持“速度误差 / 梯度放大 / 算子链”需拆分排查，不能单凭此图断言只有近壁层问题'
],0.8,3.05,5.4,2.45,fill=PALE_BLUE,accent=BLUE)
img=ROOT/'outputs/wss_pinn/audits/cfd_velocity_wss_explore/01_pred_vs_truth/LIU_JIN_LIANG/analysis_panels.png'
add_img(s,img,6.62,3.05,5.85,3.15)
footer(s,'右图为旧 exploratory velocity→WSS 图；正式下一步要在 VF6 checkpoint 上重做同样闭环')

# 10 PINN existing evidence
s=prs.slides.add_slide(blank); title(s,'C路线：物理约束的问题，先看 loss 和梯度，不先看总分','09 / C路线：PINN诊断')
img1=ROOT/'docs/03-汇报材料/V4汇报/V4_稳态梯度与损失_2026-09-12/fig1_gradient_conflict_steady.png'
img2=ROOT/'docs/03-汇报材料/V4汇报/V4_稳态梯度与损失_2026-09-12/fig2_training_losses_steady.png'
add_img(s,img1,0.62,1.48,6.0,3.63)
add_img(s,img2,6.78,1.48,5.95,3.63)
note_box(s,'历史 V4 证据：data 与 no-slip 梯度长期反向；EMA 权重约 99.76% 训练时刻贴近上限。PDE 残差下降，但数据拟合更差。',1.02,5.45,11.25,0.86,PALE_RED,RED)
footer(s,'图件来自历史 V4；不能写成最新 PF6/VF6 的 physics 结果')

# 11 network/module matrix
s=prs.slides.add_slide(blank); title(s,'模块作用矩阵：把“有用”写成可检验的具体命题','10 / 机制总结')
cols=[('模块',1.65),('它改变了什么',4.2),('当前证据',3.0),('还缺什么',3.2)]
x=0.72
for h,w in cols:
    rounded(s,x,1.45,w,0.5,NAVY,NAVY); textbox(s,h,x+0.08,1.59,w-0.16,0.2,13,RGBColor(255,255,255),True,PP_ALIGN.CENTER); x+=w+0.08
data=[('局部几何 / LocalGeoPE','局部邻域表达','R4 早期有增益','X5 最新局部云图'),('Murray F6','分支流量/幅值先验','X5、PF6、VF6 配对增益','分支误差与最差病例对应'),('section token X11','截面上下文','整体/高值 R² 小幅改善','热点 IoU 是否改善'),('Transformer / 全局上下文','长程交互','增益小，不能解决局部误差','跨病例稳定性'),('BC / PDE / EMA','物理残差与边界','梯度冲突、权重饱和','修正尺度后再做小矩阵')]
for r,(a,b,c,d) in enumerate(data):
    y=2.04+r*0.82
    vals=[a,b,c,d]; x=0.72
    for j,(v,w) in enumerate(zip(vals,[1.65,4.2,3.0,3.2])):
        fill=RGBColor(248,250,252) if r%2==0 else LIGHT
        rounded(s,x,y,w,0.66,fill,RGBColor(225,230,237)); textbox(s,v,x+0.09,y+0.12,w-0.18,0.42,11.5,NAVY if j==0 else INK,j==0,PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE); x+=w+0.08
footer(s)

# 12 close / requests
s=prs.slides.add_slide(blank); title(s,'下一步不是继续堆实验，而是补齐三张能区分机制的图','11 / 讨论')
bullet_box(s,'优先级 1｜直接 WSS',[
    '补 X5/X5X11 最好、中位、最差三联图',
    '检查极值低估来自幅值、热点位置还是局部形态'
],0.72,1.5,3.85,2.18,fill=RGBColor(224,245,242),accent=TEAL)
bullet_box(s,'优先级 2｜体场 → WSS',[
    '在 VF6 上重做 CFD velocity oracle / predicted velocity / WSS 三列',
    '拆分预测速度误差与 WSS 算子误差'
],4.75,1.5,3.85,2.18,fill=PALE_BLUE,accent=BLUE)
bullet_box(s,'优先级 3｜PINN',[
    '保留 raw/weighted loss、gradient cosine、EMA λ',
    '先让权重控制器进入工作区，再讨论物理项是否有效'
],8.78,1.5,3.85,2.18,fill=PALE_ORANGE,accent=ORANGE)
textbox(s,'请老师帮助判断的三个问题',0.85,4.18,4.0,0.35,20,NAVY,True)
textbox(s,'1  先补局部高 WSS / 分叉图，还是先做更大规模稳定性？\n2  VF6 先查近壁采样，还是先查预测速度的径向/周向结构？\n3  PINN 是否先按梯度冲突与 λ 控制器 Gate 继续，而不是按当前 WSS 数值直接 No-Go？',0.92,4.75,11.6,1.35,16,INK,False,fill=LIGHT,margin=0.22)
footer(s,'参考稿结束｜正式版建议保留 10–12 页，附录再放完整实验矩阵')

prs.save(OUT)
print(OUT)
