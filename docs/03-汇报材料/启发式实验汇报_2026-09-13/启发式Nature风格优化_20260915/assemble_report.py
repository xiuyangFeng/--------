"""Assemble verified figure exports into slide decks and a vector PDF booklet."""
from pathlib import Path
import json, hashlib
import pymupdf
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation
from pptx.util import Inches

OUT=Path(__file__).resolve().parent
STEMS=['01_routes','02_protocol','03_backbone','04_x5','05_x11','06_volume_models',
       '07_murray_paired_effect','08_x11_tradeoff','09_current_case_distributions',
       '10_x5x11_cases','11_vf6_cases','12_pf6_cases','17_velocity_wss_gap','13_diagnostic_map',
       '14_data_loss_terms','15_physics_loss_terms','18_physics_conflicts',
       '16_historical_case_distributions','19_discussion',
       '10_x5x11_cases_selfmax','11_vf6_cases_selfmax','12_pf6_cases_selfmax']

NOTES={
 '01_routes':'先指出A直接输出WSS，B输出速度再经算子求WSS，PF6压力独立；C是历史物理诊断。',
 '02_protocol':'说明当前A/B使用V5；历史C不能混成同条件排名。病例R²均值与R²_cb不同。',
 '03_backbone':'主干负责几何编码与特征恢复。结构用途不等于独立消融净增益。',
 '04_x5':'沿箭头讲两处残差：support特征相加；log_z基础输出加patch修正。Murray是输入先验。',
 '05_x11':'突出截面上下文进入FiLM条件，不直接加到WSS；有收益也有代价。',
 '06_volume_models':'PF6/VF6独立训练；BT位置与局部分支不同，QAD不是PDE约束。',
 '07_murray_paired_effect':'三个目标按相同seed配对。支持先验收益同向，不声称统计显著或机制已经证明。',
 '08_x11_tradeoff':'整体R²与高值幅度改善，MAE和热点IoU退步。请老师讨论如何权衡。',
 '09_current_case_distributions':'每点是一病例三seed均值，横须为seed样本SD；展示全部34例，Median不是Most。',
 '10_x5x11_cases':'从最好到中位到最差看空间误差。场值和正式R²是s1234，选例依据三seed均值。',
 '11_vf6_cases':'完整体内点云投影，无壁面插值；速度大小好不代表方向和近壁梯度好。',
 '12_pf6_cases':'展示相对壁面压力；R²针对原始壁面与体内全域，不能当作仅壁面分数。',
 '17_velocity_wss_gap':'旧R5V与CFD速度通过同算子是诊断证据；数值不能贴给VF6。',
 '13_diagnostic_map':'不画未经测量的假想剖面。用三项待验证问题定位误差来源。',
 '14_data_loss_terms':'逐项核查u/v/w/p的数据拟合；所有epoch保留，粗线只是移动均值。',
 '15_physics_loss_terms':'六项都是raw损失，尺度不同；不能从大小直接判断加权贡献或梯度强弱。',
 '18_physics_conflicts':'左为PN梯度诊断，右为PN/PNPP共同数据损失，接近重合不是漏画；不据此否定PINN。',
 '16_historical_case_distributions':'V4历史test35、单seed、未收敛状态；全部负R²保留，只作机制诊断。',
 '19_discussion':'把讨论落到区分幅度/位置、场/算子、边界/优化的下一轮证据上。',
}

def validate(stems):
    records=[]
    for stem in stems:
        for ext in ('pdf','svg','png'):assert (OUT/'figures'/f'{stem}.{ext}').is_file(),(stem,ext)
        p=OUT/'qa'/f'{stem}.metadata.json';meta=json.loads(p.read_text())
        assert all(r['exit_code']==0 for r in meta['audits'].values()),(stem,meta['audits'])
        report=json.loads((OUT/'qa'/f'{stem}.collision.json').read_text())
        assert report['summary']['fail']==0,(stem,report['summary'])
        assert report['summary']['warn']==0,(stem,'review required',report['summary'])
        with Image.open(OUT/'figures'/f'{stem}.png') as im:assert im.size==(4800,2700),(stem,im.size)
        doc=pymupdf.open(OUT/'figures'/f'{stem}.pdf');assert len(doc)==1
        assert len(doc[0].get_text().strip())>20,(stem,'no editable PDF text')
        assert abs(doc[0].rect.width-1152)<.1 and abs(doc[0].rect.height-648)<.1
        doc.close();records.append(meta)
    return records

def deck(stems,name):
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    doc=pymupdf.open()
    for i,stem in enumerate(stems,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_picture(str(OUT/'figures'/f'{stem}.png'),0,0,width=prs.slide_width,height=prs.slide_height)
        meta=json.loads((OUT/'qa'/f'{stem}.metadata.json').read_text())
        base=stem.replace('_selfmax','')
        note=NOTES.get(base,'')
        if stem.endswith('_selfmax'):note+=' 附录：各自最大值归一化仅比较分布，幅值差已移除；分母固定取原始同点域。'
        slide.notes_slide.notes_text_frame.text=f'{i}. {meta["claim"]}\n\n讲解提示：{note}\n\n图件：figures/{stem}.png\n核验：qa/{stem}.metadata.json'
        with pymupdf.open(OUT/'figures'/f'{stem}.pdf') as pdf:doc.insert_pdf(pdf)
    prs.save(OUT/f'{name}.pptx');doc.save(OUT/f'{name}.pdf',garbage=4,deflate=True);doc.close()
    assert len(Presentation(OUT/f'{name}.pptx').slides)==len(stems)
    with pymupdf.open(OUT/f'{name}.pdf') as pdf:assert len(pdf)==len(stems)

def contact(stems,name,columns=4):
    width,cellw,cellh=2400,600,370
    rows=(len(stems)+columns-1)//columns
    sheet=Image.new('RGB',(columns*cellw,rows*cellh),'#E9EEF1');draw=ImageDraw.Draw(sheet)
    font=ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',18)
    for i,stem in enumerate(stems):
        x=(i%columns)*cellw;y=(i//columns)*cellh
        draw.text((x+10,y+5),f'{i+1:02d}  {stem}',font=font,fill='#203140')
        with Image.open(OUT/'figures'/f'{stem}.png') as im:
            im=im.convert('RGB');im.thumbnail((590,332),Image.Resampling.LANCZOS);sheet.paste(im,(x+5,y+32))
    sheet.save(OUT/name)

def main():
    records=validate(STEMS)
    deck(STEMS,'启发式实验汇报_Nature优化版_20260915')
    modules=['03_backbone','04_x5','05_x11','06_volume_models']
    deck(modules,'网络模块_Nature优化版_20260915')
    contact(STEMS,'全部图件总览.png');contact(modules,'网络模块总览.png',columns=2)
    pages=[]
    for i,(stem,m) in enumerate(zip(STEMS,records),1):pages.append({'page':i,'stem':stem,'claim':m['claim'],'appendix':i>19,'files':m['exports'],'note':NOTES.get(stem.replace('_selfmax',''),'')})
    result={'status':'passed','figure_count':len(STEMS),'main_pages':19,'appendix_pages':3,'network_pages':4,
      'formats':['pptx','pdf','svg','png'],'png_resolution':[4800,2700],
      'pptx_content':'high-resolution figure images with speaker notes; edit scientific diagrams via SVG or Python source',
      'checks':{'all_pdf_collision_failures':0,'all_pdf_collision_warnings':0,'all_pdf_text_audits_passed':True,'all_figure_exports_exist':True,'all_pdfs_have_selectable_text':True,'pptx_pdf_page_counts_match':True},
      'pages':pages}
    (OUT/'delivery_manifest.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in result.items() if k!='pages'},ensure_ascii=False))

if __name__=='__main__':main()
