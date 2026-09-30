"""R4 report figures, traced against the experiment config and strict-loaded checkpoint.
Run with /public/newhome/cy/.conda/envs/GNN/bin/python make_figures.py.
"""
from pathlib import Path
import json
import csv
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Circle
from matplotlib.font_manager import FontProperties, fontManager
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent
ROOT = next(p for p in OUT.parents if (p / 'training_wss_min').is_dir())
CFG = ROOT / 'training_wss_min/configs/v5_rerun_20260906/r4_lsa2_h2_logradius_v5feat_s1234.json'
cfg = json.loads(CFG.read_text())
RUN = ROOT / 'training_wss_min/runs' / cfg['name']
fontfile = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
fontManager.addfont(fontfile)
plt.rcParams.update({'font.family': ['DejaVu Sans', FontProperties(fname=fontfile).get_name()],
                     'axes.unicode_minus': False, 'pdf.fonttype': 42, 'svg.fonttype': 'none'})
INK='#162C46'; MUTED='#586B80'; BLUE='#2272B6'; TEAL='#098D8A'; PURPLE='#8551B8'; ORANGE='#C87725'
LIGHT={BLUE:'#EAF3FC',TEAL:'#E8F6F3',PURPLE:'#F3ECFA',ORANGE:'#FFF3E5',INK:'#EEF2F6'}
FIGS=[]
TEXTS=[]

def text(ax,x,y,s,size=12,color=INK,weight='normal',ha='left',va='center',**kw):
    t=ax.text(x,y,s,fontsize=size,color=color,fontweight=weight,ha=ha,va=va,linespacing=1.55,**kw)
    TEXTS.append(t)
    return t

def page(no,title,subtitle):
    fig,ax=plt.subplots(figsize=(20,11.25));fig.subplots_adjust(0,0,1,1)
    fig.patch.set_facecolor('#FAFCFF');ax.set_xlim(0,160);ax.set_ylim(0,90);ax.axis('off')
    ax.add_patch(FancyBboxPatch((3,80.5),1,6,boxstyle='round,pad=0',fc=BLUE,ec='none'))
    text(ax,6,85,title,25,weight='bold');text(ax,6,80.6,subtitle,11.5,MUTED)
    text(ax,154,85,f'{no:02d} / 03',13,MUTED,ha='right')
    text(ax,3,1.6,'V5 · R4 · seed 1234  |  依据实际配置、模型代码与 best checkpoint  |  2026-09-09',9,MUTED)
    text(ax,157,1.6,'尺寸记法：N × C（每例点数 × 通道数；省略 batch）',9,MUTED,ha='right')
    return fig,ax

def box(ax,x,y,w,h,title,lines=(),color=BLUE,fs=12,ts=14):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.35,rounding_size=1.1',lw=1.4,ec=color,fc=LIGHT[color]))
    text(ax,x+1.6,y+h-3,title,ts,color,weight='bold')
    if lines:
        gap=min(3.35,(h-7)/max(len(lines)-1,1))
        for i,l in enumerate(lines):text(ax,x+1.6,y+h-7-i*gap,l,fs)

def arrow(ax,points,color=INK,dashed=False,label=None,label_pos=None):
    for a,b in zip(points[:-2],points[1:-1]):ax.plot([a[0],b[0]],[a[1],b[1]],color=color,lw=1.65,ls='--' if dashed else '-')
    ax.add_patch(FancyArrowPatch(points[-2],points[-1],arrowstyle='-|>',mutation_scale=13,color=color,lw=1.65,linestyle='--' if dashed else '-'))
    if label:text(ax,*label_pos,label,10.5,color,ha='center',bbox={'facecolor':'#FAFCFF','edgecolor':'none','pad':1.5})

def add(ax,x,y,color=INK):
    ax.add_patch(Circle((x,y),1.4,fc='white',ec=color,lw=1.6));text(ax,x,y,'+',16,color,ha='center')

def save(fig,stem):
    fig.savefig(OUT/f'{stem}.png',dpi=220,facecolor=fig.get_facecolor())
    FIGS.append(fig)

# 1. End-to-end path. Dashed vertical arrows are real encoder skip tensors.
fig,ax=page(1,'V5·R4  |  从血管几何到逐点 WSS','PointNeXt-R + LocalGeoPE + L-SA2  ·  17 维几何输入  ·  参数量 573,351')
box(ax,3,51,25,24,'壁面 support 点云',['坐标 P：5000 × 3','特征 X：5000 × 17','原有 7D + V5 新增 10D','每 epoch 随机重采样','几何邻域由 P 构建'],BLUE,11.7)
box(ax,33,51,24,24,'Stem 逐点嵌入',['MLP：17 → 32 → 32','Linear + BN + ReLU','输出 5000 × 32','保留细分辨率特征'],BLUE,11.5)
box(ax,62,51,26,24,'SA1  +  残差块 ×1',['FPS：5000 → 125','KNN-64 + 覆盖补边','LocalGeoPE + max','InvRes：64 → 128 → 64','输出 125 × 64'],BLUE,11.4)
box(ax,94,51,28,24,'SA2  +  残差块 ×1',['FPS：125 → 125','Ball：r=0.10，K≤16','GeoPE → L-SA2 → max','InvRes：128 → 256 → 128','输出 125 × 128'],PURPLE,11.2)
box(ax,128,51,29,24,'SA3  最粗尺度',['FPS：125 → 32','Ball：r=0.20，K≤16','LocalGeoPE + max','本级残差块数：0','输出 32 × 256'],BLUE,11.7)
for a,b in [(28,33),(57,62),(88,94),(122,128)]:arrow(ax,[(a,63),(b,63)])
box(ax,94,24,28,18,'FP3  恢复到 SA2',['3-NN：32 → 125 点','拼接：256 + 128 = 384','MLP：384 → 128 → 128','输出 125 × 128'],TEAL,11.5)
box(ax,62,24,26,18,'FP2  恢复到 SA1',['3-NN：125 → 125 点','拼接：128 + 64 = 192','MLP：192 → 64 → 64','输出 125 × 64'],TEAL,11.3)
box(ax,33,24,24,18,'FP1  恢复到 support',['3-NN：125 → 5000 点','拼接：64 + 32 = 96','MLP：96 → 32 → 32','输出 5000 × 32'],TEAL,10.8,13)
box(ax,3,24,25,18,'Query 特征读取',['训练 SAME：直接读取','全壁面推理：3-NN 插值','5000 × 32 → Nq × 32','固定 support，只编码一次'],TEAL,10.8,13)
arrow(ax,[(143,51),(143,33),(122,33)])
for a,b in [(94,88),(62,57),(33,28)]:arrow(ax,[(a,33),(b,33)],TEAL)
for x,l in [(45,'skip 32D'),(75,'skip 64D'),(108,'skip 128D')]:arrow(ax,[(x,51),(x,42)],BLUE,True,l,(x,46.6))
text(ax,130,24.5,'FP = 特征传播\n3-NN 权重 ∝ 1 / d²\n插值后再拼接 skip',11,TEAL,va='bottom')
box(ax,3,5,35,12,'逐点回归头',['Linear 32 → 64 → 1；ReLU','预测 ẑ：Nq × 1（log_z 空间）'],ORANGE,11.5)
arrow(ax,[(15.5,24),(15.5,17)],ORANGE)
box(ax,44,5,47,12,'恢复物理量：峰值帧 1162',['WSS(Pa) = exp(1.299948·ẑ + 0.640838) − 10⁻⁶','每个 query 点输出一个 WSS 标量'],ORANGE,10.9)
arrow(ax,[(38,11),(44,11)],ORANGE)
box(ax,97,5,60,12,'训练监督：同一个标量回归输出',['L = MSE(ẑ, z) + 0.20 × Pinball₀.₉(ẑ, z)','z = [ln(WSS + 10⁻⁶) − μtrain] / σtrain'],INK,11.4)
arrow(ax,[(38,15),(41,15),(41,20),(127,20),(127,17)],ORANGE,True)
save(fig,'01_整体网络结构')

# 2. LocalGeoPE and exactly located local transformer.
fig,ax=page(2,'SA 编码细节  |  LocalGeoPE 与 L-SA2','三层 SA 均加入 LocalGeoPE；只有 SA2 在邻域 token 上运行局部 Transformer，然后进行 max pooling')
box(ax,3,52,28,24,'采样与邻域分组',['输入：P、特征 F、属性 a','FPS 选中心 i，邻居 j','token：[Δpᵢⱼ, Fⱼ]','a = [s̃, R̃, κ̃]','a 从标准化输入索引 3/4/5 取'],BLUE,11.3)
box(ax,39,64,47,12,'特征分支：逐邻居 MLP',['(Cin + 3) → Cout → Cout','两层均为 Linear + BN + ReLU'],BLUE,11.7)
box(ax,39,42,47,17,'几何分支：LocalGeoPE',['gᵢⱼ = [Δp/r, ‖Δp/r‖, Δa]，共 7D','MLP：7 → Cout → Cout','末层无 BN / 激活；权重、偏置零初始化'],TEAL,11.6)
arrow(ax,[(31,69),(39,69)]);arrow(ax,[(31,57),(35,57),(35,50),(39,50)],TEAL)
add(ax,94,69);arrow(ax,[(86,69),(92.6,69)]);arrow(ax,[(86,50),(94,50),(94,67.6)],TEAL)
box(ax,102,54,55,22,'聚合顺序',['SA1：token 相加 → max → InvRes ×1','SA2：token 相加 → L-SA2 → max → InvRes ×1','SA3：token 相加 → max','SA 边 MLP：35→64→64 / 67→128→128','                         / 131→256→256'],PURPLE,11.5)
arrow(ax,[(95.4,69),(102,69)],PURPLE)
text(ax,3,36,'展开 SA2 的单个邻域  |  每个中心独立处理 K≤16 个 128D token；共 125 个中心',15,PURPLE,weight='bold')
box(ax,3,16,24,14,'输入 T',['K × 128','已含相对位置与 GeoPE'],PURPLE,11)
box(ax,34,16,28,14,'Pre-LN + 4-head MHA',['Q、K、V：128D → 4 × 32D','邻域内 softmax(QKᵀ/√32)'],PURPLE,10.9)
add(ax,70,23,PURPLE)
box(ax,78,16,38,14,'Pre-LN + FFN',['Linear 128 → 256 → 128','中间 GELU；dropout = 0'],PURPLE,11.5)
add(ax,124,23,PURPLE)
box(ax,132,16,25,14,'Masked max',['沿 K 个 token 聚合','每个中心输出 128D'],PURPLE,11)
arrow(ax,[(27,23),(34,23)],PURPLE);arrow(ax,[(62,23),(68.6,23)],PURPLE,label='× γattn',label_pos=(65,19))
arrow(ax,[(71.4,23),(78,23)],PURPLE);arrow(ax,[(116,23),(122.6,23)],PURPLE,label='× γffn',label_pos=(119,19))
arrow(ax,[(125.4,23),(132,23)],PURPLE)
arrow(ax,[(29,23),(29,32.5),(70,32.5),(70,24.4)],PURPLE,True)
arrow(ax,[(74,23),(74,32.5),(124,32.5),(124,24.4)],PURPLE,True)
text(ax,3,9.5,'T₁ = T + γattn·MHA(LN(T))      T₂ = T₁ + γffn·FFN(LN(T₁))      γattn、γffn 均为可学习标量，初值 0.001',12,PURPLE)
text(ax,3,5.4,'SA1 的 knn_cover 可在 64 邻居之外补边；r=0.05 用于该级 GeoPE 缩放与后续残差 ball 邻域，未限制 SA1 的 KNN 分组。',10.7,MUTED)
save(fig,'02_SA与Transformer展开')

# 3. Residual update plus explicit interpolation and training objective.
fig,ax=page(3,'残差、解码与回归  |  模块展开','PointNeXt-R 同分辨率双残差更新；FP 通过插值和跳连逐级融合；最终监督发生在 log_z 空间')
text(ax,3,75,'A   LocalInvResBlock：两条可学习缩放残差分支',15,BLUE,weight='bold')
box(ax,3,51,21,18,'输入 F',['SA1：125 × 64','SA2：125 × 128','点数保持不变'],BLUE,11.5)
box(ax,31,51,44,18,'局部几何聚合',['同层 ball 邻域：[Δp, Fⱼ]','MLP：(C+3) → C → C；max','首层 BN + ReLU；末层线性','SA1 r=.05 / K≤64；SA2 r=.10 / K≤16'],BLUE,11)
add(ax,83,60,BLUE)
box(ax,92,51,39,18,'倒置瓶颈 Pointwise MLP',['C → 2C → C','首层 Linear + BN + ReLU','末层 Linear','64→128→64 / 128→256→128'],BLUE,11.3)
add(ax,140,60,BLUE)
arrow(ax,[(24,60),(31,60)],BLUE);arrow(ax,[(75,60),(81.6,60)],BLUE,label='× γlocal × D',label_pos=(78,55))
arrow(ax,[(84.4,60),(92,60)],BLUE);arrow(ax,[(131,60),(138.6,60)],BLUE,label='× γpw × D',label_pos=(135,55))
arrow(ax,[(141.4,60),(156,60)],BLUE,label='Fout',label_pos=(151,63))
arrow(ax,[(27,60),(27,71),(83,71),(83,61.4)],BLUE,True)
arrow(ax,[(87,60),(87,71),(140,71),(140,61.4)],BLUE,True)
text(ax,3,45,'h = F + D·γlocal·Local(F)； Fout = h + D·γpw·MLP(h)。γ 初值均为 0.001；相加后无激活。',12,BLUE)
text(ax,3,40.5,'D：按病例共享的 DropPath 缩放；SA1 丢弃率 0，SA2 丢弃率 0.10。同一残差块两分支共用 D；推理时 D=1。',11.5,MUTED)
text(ax,3,33.5,'B   Feature Propagation 与 query 插值',15,TEAL,weight='bold')
box(ax,3,11,31,18,'粗层 → 细层坐标',['选最近 3 个粗层点','wⱼ ∝ 1 / max(dⱼ², 10⁻¹⁶)','对 w 归一化后加权求和'],TEAL,11.3)
box(ax,41,11,38,18,'拼接同层 encoder skip',['[插值特征, skip] → MLP','FP3：384 → 128 → 128','FP2：192 → 64 → 64','FP1：96 → 32 → 32'],TEAL,11.4)
arrow(ax,[(34,20),(41,20)],TEAL)
box(ax,86,11,32,18,'Query → 回归头',['SAME 读取 / 全壁面 3-NN','32 → 64 → 1','隐藏层 ReLU，末层线性','每点独立预测 log_z 标量'],ORANGE,11.1)
arrow(ax,[(79,20),(86,20)],TEAL)
box(ax,125,11,32,18,'MSE + q90 pinball',['e = z − ẑ','ρ₀.₉(e) = max(0.9e, −0.1e)','L = mean(e²) + .2·mean(ρ)','同一输出参与两项损失'],ORANGE,10.8)
arrow(ax,[(118,20),(125,20)],ORANGE)
text(ax,3,6,'FP 各层两次 Linear 均接 BN + ReLU；query 插值在 32D 特征上完成，再送入回归头。q90 表示 pinball 的分位参数。',11,MUTED)
save(fig,'03_残差解码与回归展开')

# Actual configured order: 0-based indices match LocalGeoPE configuration.
rows=[
('0–2','原有7D','x, y, z',3,'对齐后的壁面三维坐标','无量纲，[-1,1]','分叉为原点；+Z朝入口、+X朝左髂总；除以逐例 max-abs 尺度；不做 z-score'),
('3','原有7D','abscissa_norm',1,'从入口根节点沿中心线树到映射点的累计弧长 / 本例壁面映射弧长最大值','无量纲，[0,1]','train138 z-score；同时进入 LocalGeoPE 的差分属性'),
('4','原有7D','local_radius',1,'映射中心线位置的平滑局部半径 R','mm','train138 z-score；同时进入 LocalGeoPE 的差分属性'),
('5','原有7D','curvature',1,'映射中心线位置的曲率 κ','mm⁻¹','signed_log1p → 训练集绝对值 P99 裁剪 → train138 z-score；同时用于 LocalGeoPE'),
('6','原有7D','log_local_radius',1,'局部半径的自然对数 ln(max(R, 10⁻⁶))，R 取 mm 数值','对数特征','先取 ln，再用 train138 z-score'),
('7','V5新增10D','rho',1,'到中心线的横向径向距离 r⊥ / 局部半径 R','无量纲','train138 z-score；不强制限制在 [0,1]'),
('8','V5新增10D','theta_sin',1,'中心线局部法平面内周向角的 sinθ','无量纲，[-1,1]','从 atlas 局部坐标架计算 θ，再取 sin，最后 train138 z-score'),
('9','V5新增10D','theta_cos',1,'中心线局部法平面内周向角的 cosθ','无量纲，[-1,1]','从 atlas 局部坐标架计算 θ，再取 cos，最后 train138 z-score'),
('10','V5新增10D','dr_ds',1,'局部半径沿中心线弧长的变化率 dR/ds','mm/mm，无量纲','train138 z-score'),
('11','V5新增10D','dist_to_junction_mm',1,'映射中心线点沿中心线树到最近分叉节点的距离','mm','train138 z-score；沿树距离'),
('12','V5新增10D','dist_to_endpoint_mm',1,'映射中心线点沿中心线树到最近开口端点的距离（入口或出口）','mm','train138 z-score；沿树距离'),
('13','V5新增10D','end_zone',1,'中心线端区特征保持/修正区域标记；含开口端点与子分支起始端','原始值 0/1','该实验仍做 train138 z-score；不是单纯的入口/出口类别'),
('14–16','V5新增10D','nx_aligned, ny_aligned, nz_aligned',3,'壁面点云 PCA 估计的外向法向，旋转到与 xyz 相同的解剖坐标架','原始单位法向的三个分量','三个分量分别使用 train138 z-score；标准化后不再是单位向量'),
]
headers=['输入索引（0起）','分组','配置字段名','维数','几何含义','标准化前单位/范围','进入网络前处理']
assert sum(r[3] for r in rows)==17
assert [s.strip() for r in rows for s in r[2].split(',')]==cfg['data']['input_features']
for ext,delimiter in [('csv',','),('tsv','\t')]:
    with (OUT/f'V5_R4_输入特征表.{ext}').open('w',encoding='utf-8-sig',newline='') as f:
        wr=csv.writer(f,delimiter=delimiter);wr.writerow(headers);wr.writerows(rows)
table='| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'
table+=''.join('| '+' | '.join(map(str,r))+' |\n' for r in rows)
stats=json.loads((RUN/'feature_stats.json').read_text())
wb=Workbook();ws=wb.active;ws.title='17维输入特征';ws.append(headers)
for r in rows:ws.append(r)
ws2=wb.create_sheet('实际训练统计');ws2.append(['输入索引','特征名','均值 μ','标准差 σ','clip','transform'])
for i,name in enumerate(cfg['data']['input_features']):
    s=stats.get(name,{})
    ws2.append([i,name,s.get('mean'),s.get('std'),s.get('clip'),s.get('transform','逐例归一化坐标，不做z-score')])
stage_rows=[
['输入','5000 × 17','17','pos 同时单独作为几何坐标'],
['Stem','5000 × 32','17 → 32 → 32','逐点 MLP'],
['SA1','125 × 64','边 MLP 35→64→64；GeoPE 7→64→64','FPS125；knn_cover64；max；InvRes×1'],
['SA1 InvRes','125 × 64','局部 67→64→64；逐点 64→128→64','ball .05 / ≤64；DropPath 0'],
['SA2','125 × 128','边 MLP 67→128→128；GeoPE 7→128→128','FPS125；ball .10 / ≤16；L-SA2→max；InvRes×1'],
['SA2 Transformer','每邻域 K × 128 → 128','4头×32；FFN 128→256→128','Pre-LN；两处缩放残差；max后为125×128'],
['SA2 InvRes','125 × 128','局部 131→128→128；逐点 128→256→128','ball .10 / ≤16；DropPath .10'],
['SA3','32 × 256','边 MLP 131→256→256；GeoPE 7→256→256','FPS32；ball .20 / ≤16；max；残差块0'],
['FP3','125 × 128','256+128=384 → 128 → 128','3-NN 特征插值 + SA2 skip'],
['FP2','125 × 64','128+64=192 → 64 → 64','3-NN 特征插值 + SA1 skip'],
['FP1','5000 × 32','64+32=96 → 32 → 32','3-NN 特征插值 + Stem skip'],
['Query','Nq × 32','32 → 32','SAME直接读取；独立query固定3-NN插值'],
['回归头','Nq × 1','32 → 64 → 1','ReLU；最后线性输出log_z'],
['物理输出','Nq × 1','1 → 1','exp(σtrain ẑ + μtrain) − eps，单位Pa'],
]
ws3=wb.create_sheet('通道与模块');ws3.append(['阶段','输出尺寸（单病例）','通道变换','模块/说明'])
for r in stage_rows:ws3.append(r)
for sheet,widths in [(ws,[19,16,43,8,67,27,84]),(ws2,[14,36,23,23,24,44]),(ws3,[24,28,61,67])]:
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for j,w in enumerate(widths,1):sheet.column_dimensions[get_column_letter(j)].width=w
    for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='162C46');cell.font=Font(name='微软雅黑',bold=True,color='FFFFFF',size=11)
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height=48
        for cell in row:
            cell.font=Font(name='微软雅黑',size=11);cell.alignment=Alignment(vertical='center',wrap_text=True)
            if row[0].row%2==0:cell.fill=PatternFill('solid',fgColor='EAF3FC')
wb.save(OUT/'V5_R4_输入特征与通道表.xlsx')

notes='''# V5·R4 模型结构与输入表

对应工作簿「实验矩阵总览」207 行、「教师汇报视图」182 行的 **★ V5·R4 R1 + V5 十维几何特征 s1234**。

## 图件

- `01_整体网络结构.png`：适合汇报主页面的完整结构。
- `02_SA与Transformer展开.png`：SA、LocalGeoPE、四头 Transformer 与实际模块顺序。
- `03_残差解码与回归展开.png`：PointNeXt-R 双残差、DropPath、FP 和监督目标。
- PNG 为 4400×2475。
- `V5_R4_输入特征与通道表.xlsx`：输入定义、真实训练统计、逐层通道三张表。原实验工作簿未修改。
- `V5_R4_输入特征表.tsv/.csv`：用于复制到 Excel / Word；本文表格也可直接复制。

## 可直接用于汇报的说明

模型以 5000 个壁面 support 点的 17 维几何特征为输入，经 Stem 映射到 32 维。三级编码器的点数依次为 125、125、32，通道依次为 64、128、256。每级局部聚合均加入 LocalGeoPE；第二级在每个邻域内加入四头 Transformer。第一、二级聚合后各接一个 PointNeXt-R 倒置瓶颈残差块。解码器通过 3-NN 插值、编码器跳连拼接和 MLP，逐级恢复到 5000×32 的 support 特征，再读取或插值到 query 点，经 32→64→1 回归头预测峰值 WSS 的 log_z 值，最后还原为 Pa。

原有 7D = xyz 3D + 归一化弧长 1D + 半径 1D + 曲率 1D + 对数半径 1D。
新增 10D = 管道坐标 3D + 半径坡度 1D + 距离与端区 3D + 法向 3D。

## 输入表（严格按配置顺序）

'''+table+'''

xyz 另作为 `pos` 传入 FPS、邻域搜索与插值；其已包含在 17 维特征中，不另计为 20 维。其余 14 个通道均用 train138 统计进行 z-score，包括角度编码、end_zone 与法向分量。统计来自全部训练病例的壁面顶点；推理使用冻结统计。原始范围是标准化前的范围。

ρ 中的 r⊥ 为点到最近 atlas 样点切线的横向距离；θ 在该样点的局部法平面架内定义。两项距离继承该映射样点的沿中心线树距离。法向由壁面点云 PCA 估计并定向，不是中心线局部坐标架的法向。

## 逐层通道

| 阶段 | 输出尺寸（每例） | 通道变换 | 模块/说明 |
| --- | --- | --- | --- |
'''+''.join('| '+' | '.join(r)+' |\n' for r in stage_rows)+'''

## 配置解释与实现要点

1. 配置字段为 `model.name=pointnetpp`，实际类是 `PointNetPlusPlusRegressor`，启用了 PointNeXt-R 残差扩展。汇报称 **PointNeXt-R + LocalGeoPE + L-SA2**，不能用独立 `pointnext.py` 中的默认模型代替。
2. `sa_center_counts=[125,125,32]` 优先于 ratios，因此点数为 5000→125→125→32。SA2 为同点数重采样与邻域特征提升。FP2 同样保持 125 点数。
3. SA1 是 KNN-64 加未覆盖点最近中心补边，邻域可能超过 64；SA2/3 分别是半径 0.10/0.20 的 ball-query，最多 16 邻居。所有 radius 均在逐例归一化坐标空间，不是 mm。SA1 的 0.05 仍用于 GeoPE 缩放与该级 InvRes 的 ball 邻域。
4. GeoPE 的 7 维边输入为 `[Δp/r (3), ||Δp/r|| (1), Δs̃, ΔR̃, Δκ̃ (3)]`。最后一层零初始化，可学习几何增量加到 SA 边特征；新增 10D 经 Stem 主通路输入，未额外并入 GeoPE 属性索引。
5. Transformer 位于 SA2 的逐邻居 token MLP + GeoPE 之后、max pooling 之前；它有自己的 attention/FFN 缩放残差。其后再接独立 `LocalInvResBlock`，两种残差模块应区分。注意力只在同一个中心的邻域中计算，不跨病例、也不是 32 点瓶颈全局注意力。
6. InvRes 采用局部聚合更新和 C→2C→C 逐点更新，两个 gamma 初始 0.001。两处相加后不额外激活。DropPath 在两个块间线性分配为 0 和 0.10，按病例抽取并在该块两分支共享。Transformer 的 gamma 初始也为 0.001，但其分支不使用此 DropPath。
7. FP3/2/1 的拼接通道分别为 384、192、96。3-NN 使用 PyG 的归一化逆平方距离权重；对 query 插值的是 32D 特征，之后才回归标量。SAME 点直接读取编码特征；独立 query 只使用坐标定位插值，本配置不将 query 自身的 17D 特征额外送入回归头。
8. 单输出头；H2 是实验损失配方标签，不能解释为双输出头。损失为 `mean((z-ẑ)^2) + .2 mean(max(.9(z-ẑ), -.1(z-ẑ)))`，在所有监督点上计算，q90 不表示只训练真值最高 10% 的点。
9. 标签是峰值帧 1162 的 CFD 壁面 WSS 标量。`z=[ln(WSS+1e-6)-0.6408381995]/1.2999484463`；输出恢复为 `exp(1.2999484463*ẑ+0.6408381995)-1e-6 Pa`。原实验没有 PDE/PINN 损失、速度/压力联合输出或额外 q90 输出头。
10. R4 中 bottleneck Transformer、coarse global attention、EdgeConv、SEP/QAD/local-attn query decoder、局部壁面分支与病例幅值头未启用。图中只展示实际启用模块。

## 训练与验证

train138/test34；seed1234；400 epoch；batch8；AdamW，lr=1e-3，warmup10 后 cosine 至1e-5；weight decay=1e-4；grad clip=1；AMP。每 epoch 随机5000 support=query（SAME），无旋转增强。按训练损失选 best；推理固定5000 support（seed1234）并对全壁面 query 分块（16384）预测。

已将 `ckpt_best.pt` 严格加载到配置构建的模型，全部键匹配；参数量 **573,351**。真实病例抽取5000点完成 CPU 前向核对，并检查 SAME/独立query 输出有限。源文件哈希、实际输出尺寸与检查结果见 `verification.json`。这些检查仅验证本次图件对应结构，不重新训练或改变实验指标。

## 依据

- `training_wss_min/configs/v5_rerun_20260906/r4_lsa2_h2_logradius_v5feat_s1234.json`
- `training_wss_min/baseline_models.py`：`PointNetPlusPlusRegressor`、`PointNetSetAbstraction`、`LocalNeighborhoodTransformer`、`LocalInvResBlock`、`FeaturePropagation`
- `training_wss_min/dataset.py`：特征组装、train统计、归一化/反归一化
- `training_wss_min/objectives.py`：MSE + pinball
- `wss_v5/views/wss_min_view.py`、`wss_v5/centerline_features.py`、`wss_pinn/v4/centerline_atlas.py`：几何定义
- 该 run 的 `config.json`、`feature_stats.json`、`wss_global_stats.json` 与 `ckpt_best.pt`
- `docs/02-推进与变更/00-V5设计与历史跟踪/WSS_V5_训练实验跟踪_历史卷_2026-09-06至09-20.md` §4
- 用户提供工作簿中的 R4 s1234 对应行
'''
(OUT/'README.md').write_text(notes,encoding='utf-8')

# Presentation check: flag texts outside their canvas.
outside=[]
for fig in FIGS:
    fig.canvas.draw();renderer=fig.canvas.get_renderer()
    for t in fig.axes[0].texts:
        bb=t.get_window_extent(renderer)
        if bb.x0 < 0 or bb.y0 < 0 or bb.x1 > fig.bbox.width or bb.y1 > fig.bbox.height:outside.append(t.get_text())
assert not outside, outside
print('Created 3 figures in PNG, XLSX/CSV/TSV, and README. No text outside canvas.')
