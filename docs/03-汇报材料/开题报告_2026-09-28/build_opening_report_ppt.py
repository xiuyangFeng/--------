"""Build editable opening-report deck; only existing data and document readouts.
Run: python3 build_opening_report_ppt.py
No training, inference, new evaluation, or PDF export.
"""
from pathlib import Path
import json, re, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[2]
A=OUT/'assets'/'final'; A.mkdir(parents=True,exist_ok=True)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
C={'ink':'203140','muted':'617584','blue':'3D7294','teal':'2D8B82','orange':'BC8150','red':'B85E57','line':'DDE5E8','pale':'F1F6F7','white':'FFFFFF','navy':'183A4D'}
FONT='Noto Sans CJK SC'
notes_map={}
md=OUT/'开题报告_故事与逐页讲稿.md'
if md.exists():
    for k,t in re.findall(r'### P(\d{2})[^\n]*\n(.*?)(?=\n### P\d{2}|\n## |\Z)',md.read_text(),re.S): notes_map[int(k)]=t

def color(v): return RGBColor.from_string(C.get(v,v))
def box(s,x,y,w,h,fill='pale',line=None):
    q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=color(fill);q.line.fill.background() if not line else None
    if line:q.line.color.rgb=color(line)
    return q

def txt(s,x,y,w,h,t,size=20,col='ink',bold=False,align=PP_ALIGN.LEFT):
    q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_after=Pt(4);p.line_spacing=1.13
        r=p.add_run();r.text=line;r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=color(col)
        for key in ('a:ea','a:cs'):
            el=OxmlElement(key);el.set('typeface',FONT);r._r.get_or_add_rPr().append(el)
    return q

def line(s,x1,y1,x2,y2,col='line',width=1.2,arrow=False):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));q.line.color.rgb=color(col);q.line.width=Pt(width)
    if arrow:
        e=OxmlElement('a:tailEnd');e.set('type','triangle');q._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(e)
    return q

def dot(s,x,y,r=.08,col='teal'):
    q=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-r),Inches(y-r),Inches(2*r),Inches(2*r));q.fill.solid();q.fill.fore_color.rgb=color(col);q.line.fill.background();return q

def pic(s,path,x,y,w,h):
    if not path.exists(): return False
    iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale
    s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh));return True

def native_vessel(s,x,y,w,h):
    # Conceptual aorto-iliac tree only. Broad lines carry no measured WSS.
    line(s,x+.48*w,y,x+.48*w,y+.53*h,'blue',28)
    for sign in [-1,1]:
        line(s,x+.48*w,y+.5*h,x+(.48+sign*.21)*w,y+.78*h,'blue',18)
        line(s,x+(.48+sign*.21)*w,y+.77*h,x+(.48+sign*.30)*w,y+h,'blue',12)
        line(s,x+(.48+sign*.21)*w,y+.77*h,x+(.48+sign*.12)*w,y+.97*h,'blue',9)
    txt(s,x,y+h+.16,w,.4,'解剖示意 · 无定量数据',10,'muted',align=PP_ALIGN.CENTER)

def snew(title,kicker,source='',note=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color('white');n=len(prs.slides)
    txt(s,.65,.3,11.9,.24,kicker,10,'teal',True);txt(s,.65,.82,12.05,.67,title,27,'ink',True)
    line(s,.65,1.56,12.68,1.56)
    txt(s,.65,7.08,11.4,.25,source or '腹主动脉—髂动脉 WSS 快速预测｜开题报告｜修订 2026-09-29',8.5,'muted')
    txt(s,12.05,7.05,.6,.3,f'{n:02d}',10,'muted',align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text=notes_map.get(n,note or title)+'\n\n【来源】'+source
    return s

def takeaway(s,t):
    box(s,.65,6.28,12.03,.53,'pale');txt(s,.85,6.38,11.63,.4,t,17,'teal',True)

def point(s,x,y,num,title,body,w=3.6):
    txt(s,x,y,w,.48,num,29,'teal',True);txt(s,x,y+.75,w,.65,title,22,'ink',True);txt(s,x,y+1.62,w,1.3,body,17,'muted')

def chart(s,x,y,w,h,categories,series,kind=XL_CHART_TYPE.COLUMN_CLUSTERED,low=0,high=1,fmt='0.000'):
    data=CategoryChartData();data.categories=categories
    for name,vs in series:data.add_series(name,vs)
    ch=s.shapes.add_chart(kind,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    ch.has_title=False;ch.has_legend=len(series)>1
    if ch.has_legend:ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.font.size=Pt(11)
    ch.value_axis.minimum_scale=low;ch.value_axis.maximum_scale=high
    ch.value_axis.has_title=True;ch.value_axis.axis_title.text_frame.text='R²'
    for par in ch.value_axis.axis_title.text_frame.paragraphs:
        for run in par.runs:run.font.name=FONT;run.font.size=Pt(13)
    ch.value_axis.tick_labels.number_format=fmt;ch.value_axis.tick_labels.font.size=Pt(11);ch.value_axis.tick_labels.font.name=FONT
    ch.category_axis.tick_labels.font.size=Pt(12);ch.category_axis.tick_labels.font.name=FONT
    ch.value_axis.has_major_gridlines=True;ch.value_axis.major_gridlines.format.line.color.rgb=color('line')
    for k,ser in enumerate(ch.series):
        col=['blue','teal','orange'][k%3];ser.format.fill.solid();ser.format.fill.fore_color.rgb=color(col);ser.format.line.color.rgb=color(col)
        if kind==XL_CHART_TYPE.LINE_MARKERS:ser.format.line.width=Pt(2.5)
    if kind==XL_CHART_TYPE.COLUMN_CLUSTERED and len(series)==1:
        for j,pt in enumerate(ch.series[0].points):
            pt.format.fill.solid();pt.format.fill.fore_color.rgb=color(['blue','teal'][j%2]);pt.format.line.fill.background()
    p=ch.plots[0];p.has_data_labels=True;p.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END if kind!=XL_CHART_TYPE.LINE_MARKERS else XL_LABEL_POSITION.ABOVE;p.data_labels.number_format=fmt;p.data_labels.font.size=Pt(13);p.data_labels.font.color.rgb=color('ink')
    return ch

# 01
s=snew('面向腹主动脉—髂动脉的\n壁面剪切应力快速预测与应用研究','开题报告｜建议题目','研究主线：X5D｜数据与结果口径更新至 2026-09-29', '我们希望把影像中的血管形态，转成可解释的力学信息。先解决这种信息如何快速、稳定获得，再研究它能否为随访与术后比较提供增量价值。')
# extend cover title height deliberately for two lines
s.shapes[1].height=Inches(1.3);s.shapes[1].width=Inches(7.2);s.shapes[2].top=Inches(2.14);s.shapes[2].width=Inches(6.8)
txt(s,.65,2.65,6.4,1.1,'让血管影像进一步回答：\n血流如何作用于管壁？',27,'teal',True)
txt(s,.65,4.15,6.2,1.1,'真实解剖数据  →  快速 WSS 预测\n→  可追溯三维报告',21,'muted')
if not pic(s,A/'cover_wss_with_scale.png',8.3,.82,3.3,5.18):native_vessel(s,8.2,2.6,3.4,2.6)
takeaway(s,'研究价值：降低力学信息的获取成本，为后续疾病关联研究提供基础。')
txt(s,7.2,6.03,5.3,.2,'既有 v5.1 病例预测示例，非当前底座结果',8.5,'muted',align=PP_ALIGN.CENTER)
# 02
s=snew('影像呈现形态，力学信息帮助理解形态背后的差异','01｜为什么值得做','背景：ESVS 2024 指南；Dua & Dalman 2010；Mutlu et al. 2023。WSS 的临床增量价值仍需验证。')
txt(s,.75,1.95,5.2,.55,'临床已经关注',20,'muted');txt(s,.75,2.65,5.2,1.55,'直径与增长速度\n症状与解剖结构\n术前、术后的形态变化',24,'ink',True)
line(s,6.3,2,6.3,5.65)
txt(s,7.0,1.95,5.2,.55,'本研究希望补充',20,'teal');txt(s,7.0,2.65,5.2,1.55,'哪些部位剪切偏低？\n哪些分支出现高剪切？\n变化发生在什么位置？',24,'ink',True)
txt(s,.75,5.3,11.8,.55,'把可见的形态与可定位的力学信息联系起来，支持可重复的病例比较。',19,'muted')
takeaway(s,'先建立可靠的测算工具，再检验它对生长、破裂和术后闭塞研究的价值。')
# 03
s=snew('WSS 描述血流沿管壁的摩擦作用','02｜这项指标为什么有价值','WSS = wall shear stress，壁面剪切应力，单位 Pa；机制联系不等同于因果或诊断。')
box(s,.8,4.67,5.05,.4,'blue');line(s,1.05,4.52,5.2,4.52,'orange',4,True)
for yy,xx in [(2.6,4.9),(3.2,4.2),(3.8,3.3)]:line(s,1.05,yy,xx,yy,'teal',2.5,True)
txt(s,.8,1.98,5.1,.5,'近壁血流示意',21,'ink',True);txt(s,1.1,5.25,4.5,.65,'切向作用 ≠ 管壁拉伸应力',18,'blue',True)
txt(s,6.75,2.0,5.5,.45,'它提供一张空间分布图',23,'ink',True)
txt(s,6.75,2.85,5.45,2.5,'瘤腔：关注低剪切与流动分离\n\n狭窄与分叉：关注局部高剪切\n\n病程：需结合周期指标与临床随访',19,'muted')
takeaway(s,'当前主线预测收缩期峰值帧 WSS 幅值；周期振荡与破裂风险需要额外证据。')
# 04
s=snew('研究难点之一，是让力学信息能够常规、批量获得','03｜为什么需要快速预测','CFD 用作本项目的计算参照；耗时依赖网格、边界条件、硬件和求解设置。')
for y,tag,st,co in [(2.4,'逐例 CFD',['影像分割','体网格','边界设定','数值求解','后处理'],'blue'),(4.55,'拟采用路线',['影像分割','表面模型','质量检查','快速预测','三维报告'],'teal')]:
    txt(s,.75,y-.65,11,.4,tag,22,co,True)
    for i,t in enumerate(st):
        x=.75+i*2.48;box(s,x,y,2.1,.72,'pale');txt(s,x,y+.18,2.1,.4,t,18,'ink',True,PP_ALIGN.CENTER)
        if i<4:line(s,x+2.13,y+.36,x+2.42,y+.36,co,1.6,True)
takeaway(s,'训练阶段用 CFD 提供参照；应用阶段从血管表面快速获得标准化力学描述。')
# 05
s=snew('国内外研究已证明加速潜力，可靠应用仍需充分验证','04｜国内外现状','代表来源：Kang et al. 2024；Li et al. 2025；Suk et al. 2024；Rygiel et al. 2025。¹预印本；完整文献见备份页。')
rows=[('国内参与研究','Kang 2024 · AAA 术后血流','物理约束点云预测时空血流'),('国内参与研究','Li 2025 · 颅内动脉瘤 WSS','结合几何与边界条件快速估计'),('国外研究','Suk 2024 · 几何深度学习','研究表面几何与 WSS 的映射'),('国外研究','Rygiel 2025 · 真实 AAA¹','加入脉动 WSS 与外部数据验证')]
for i,(a,b,c) in enumerate(rows):
    yy=2.04+i*.88;txt(s,.75,yy,2.15,.45,a,16,'teal',True);txt(s,2.85,yy,4.3,.45,b,18,'ink',True);txt(s,7.35,yy,5.1,.5,c,17,'muted');line(s,.75,yy+.64,12.5,yy+.64)
takeaway(s,'本课题接着回答：真实主—髂动脉中，哪些结果可信，如何接入实际处理流程？')
# 06
s=snew('本课题聚焦真实解剖、局部误差与实际输入三处缺口','05｜研究定位')
point(s,.8,2.1,'01','复杂解剖下的泛化','三类临床队列\n患者与相关单元分组评估')
point(s,4.95,2.1,'02','关键区域的可靠性','关注高剪切区低估\n同时检查低剪切瘤腔')
point(s,9.05,2.1,'03','从输入到报告的闭环','真实分割面与输入质量\n端到端耗时和失败处理',3.4)
takeaway(s,'贡献以“可验证的预测能力 + 清楚的边界 + 可复用的工具”共同体现。')
# 07
s=snew('研究目标：建立能预测、能解释、能使用的计算工具','06｜研究目标与内容')
for i,(a,b,c) in enumerate([('目标一','建立可靠的预测模型','在真实数据上还原 WSS 的空间分布'),('目标二','明确误差来源与适用边界','解释数据量、关键区域和输入变化的影响'),('目标三','形成可追溯的科研原型','将质量检查、预测和三维报告串成完整流程')]):
    y=2.05+i*1.15;txt(s,.8,y,1.8,.5,a,18,'teal',True);txt(s,2.65,y,4.2,.5,b,23,'ink',True);txt(s,7.25,y,5.1,.6,c,18,'muted')
takeaway(s,'评价同时覆盖整体误差、区域误差、跨患者表现和处理流程的稳定性。')
# 08
s=snew('让模型从真实计算样本中学习，再用于新几何','07｜总体研究路线')
# two-row workflow
for x,t,b in [(.8,'真实解剖 + CFD','统一协议 · 数据审查'),(5.0,'几何与局部表征','形态、朝向、分支尺度'),(9.2,'训练与分组验证','逐病例与关键区域评估')]:
    box(s,x,2.05,3.25,1.2,'pale');txt(s,x+.18,2.25,2.9,.45,t,19,'ink',True);txt(s,x+.18,2.87,2.9,.3,b,14,'muted')
line(s,4.1,2.67,4.8,2.67,'teal',2,True);line(s,8.3,2.67,9,2.67,'teal',2,True)
line(s,6.62,3.3,6.62,4.15,'teal',2,True)
txt(s,6.9,3.65,3.0,.3,'冻结学习到的映射',12,'muted')
for x,t in [(.8,'新的血管表面'),(5.,'冻结模型快速预测'),(9.2,'WSS 场与三维报告')]:
    box(s,x,4.42,3.25,.85,'pale');txt(s,x+.18,4.64,2.9,.4,t,19,'teal',True)
line(s,4.1,4.85,4.8,4.85,'teal',2,True);line(s,8.3,4.85,9,4.85,'teal',2,True)
takeaway(s,'推理输入来自几何；预测含义由训练时的 CFD 生理条件共同限定。')
# 09
s=snew('已建立覆盖三类临床队列的真实解剖数据基础','08｜目前完成到哪里','数据：v5.2 split 与 cv5_summary；261 个影像/解剖单元，不等同于 261 位独立患者。')
txt(s,.8,2.05,5.0,1.0,'261',56,'teal',True);txt(s,.8,3.18,5.1,.8,'影像 / 解剖单元\n170 在库 + 91 回收',21,'muted')
for i,(co,n,label) in enumerate([('AG',85,'动脉瘤生长'),('AAA',60,'破裂 / 未破裂'),('ILO',116,'髂支闭塞相关')]):
    y=2.05+i*1.08;txt(s,6.4,y,1.2,.4,co,19,'blue',True);txt(s,7.65,y,1.1,.5,str(n),26,'teal',True);txt(s,9.0,y,3.3,.5,label,19,'ink');line(s,6.4,y+.74,12.5,y+.74)
txt(s,.8,5.15,11.7,.63,'患者 / 相关单元分组交叉验证，避免同一患者的信息跨训练与留出折泄漏。',20,'ink',True)
takeaway(s,'数据审查与分组设计，是后续算法结果能够被解释的前提。')
# 10
s=snew('患者分组验证支持峰值 WSS 预测的可行性','09｜主要结果','CV5：261 单元、5 折、3 seed；Pa 空间病例等权 R²_cb。来源：主线 §35、§36.3–36.5。')
chart(s,.8,2.08,7.2,3.85,['原基线','当前训练底座'],[('R²_cb',[.7608,.7714])],low=0,high=1)
txt(s,8.55,2.18,3.8,.5,'0.771',45,'teal',True);txt(s,8.55,3.24,3.8,1.1,'三次训练的\n合并折外 R² 均值',20,'ink',True)
txt(s,8.55,4.72,3.8,.9,'当前配方已确认\n部署权重另行更新',18,'muted')
takeaway(s,'已建立计算预测能力；临床增量价值将通过后续独立数据与终点研究检验。')
# 11
s=snew('后续改进方向已有证据：扩大真实覆盖，减少高值低估','10｜从结果中学到了什么','左：v5.1 学习曲线（三 seed 均值）；右：v5.2 CV5 最高10% WSS 点的低估比例约值。')
txt(s,.8,1.88,6.7,.5,'真实数据增加，性能仍在提升',21,'ink',True)
chart(s,.8,2.53,6.5,3.3,['23','45','68','91'],[('R²_cb',[.6468,.6814,.6983,.7129])],XL_CHART_TYPE.LINE_MARKERS,.60,.76)
txt(s,1.6,5.84,5.5,.3,'每折训练单元数（约值）',12,'muted',align=PP_ALIGN.CENTER)
txt(s,8.0,1.88,4.5,.5,'关键区域仍需特别关注',21,'ink',True)
txt(s,8.,2.83,4.5,.82,'71% → 64%',32,'teal',True);txt(s,8.,3.88,4.5,1.35,'提高低估惩罚后\n高剪切区低估减少\n热点位置仍需继续改善',19,'muted')
takeaway(s,'59 例形变合成数据未带来明确增益；后续优先补覆盖差异，而非只增加样本数量。')
# 12
s=snew('已形成从血管表面到三维报告的科研原型','11｜部署可行性','历史验收：X5D_v51，34 例，阶段计时中位 31.7 s / P90 74.8 s；非 CTA 到报告总时长。')
if not pic(s,A/'deployment_pair.png',.8,1.93,7.25,3.95):
    native_vessel(s,1.5,2.15,2.3,2.5);native_vessel(s,5.,2.15,2.3,2.5)
txt(s,8.55,2.1,3.8,.72,'约 32 秒',38,'teal',True);txt(s,8.55,3.03,3.8,.72,'标准化表面输入后的\n历史阶段计时中位数',17,'muted')
txt(s,8.55,4.18,3.8,1.25,'三维 WSS 分布\n分支统计与位置查询\n结果导出与模型版本',19,'ink')
takeaway(s,'保留输入质量检查和人工分支确认；先服务科研病例分析与标准化比较。')
# 13
s=snew('预期贡献落在真实解剖、关键区域与完整流程上','12｜创新与贡献')
for i,(a,b) in enumerate([('方法','将几何与分支尺度信息用于局部 WSS 预测，检验关键区域改进。'),('证据','用分组验证、学习曲线和失效分析，说明方法何时有效。'),('工具','形成从血管表面到三维报告的流程，支持重复使用和追溯。')]):
    y=2.0+i*1.2;txt(s,.8,y,1.4,.6,a,26,'teal',True);txt(s,2.7,y,9.45,.75,b,21,'ink');line(s,.8,y+.92,12.5,y+.92)
takeaway(s,'同协议外部基线与真实分割面验证，将决定这些贡献能够支持多强的结论。')
# 14
s=snew('六个月计划以验证结果和可交付成果作为里程碑','13｜后续怎么做','相对开题后的建议排期，随数据获取与临床合作条件调整。')
for i,(time,title,tasks,accept) in enumerate([('0–2 月','补齐计算验证','同协议外部基线\n冻结未触碰确认集','交付：公平比较表、确认集方案'),('3–4 月','检验真实输入与生理变化','真实分割面与切口差异\n流量 / 分流敏感性','交付：误差边界与失败处理规则'),('5–6 月','完成原型验证与论文','不确定度校准、稳定性\n回顾性病例分析方案','交付：版本化软件、验证报告')]):
    x=.8+4.15*i;txt(s,x,2.05,3.55,.5,time,28,'teal',True);line(s,x,2.8,x+3.4,2.8,'teal',2);txt(s,x,3.12,3.5,.65,title,21,'ink',True);txt(s,x,4.02,3.45,1.1,tasks,18,'muted');txt(s,x,5.4,3.45,.65,accept,14,'blue',True)
takeaway(s,'每阶段都回答一个问题：是否更准、换输入是否可靠、别人能否重复使用。')
# 15
s=snew('部署从科研助手起步，逐步走向动态数字孪生','14｜应用前景')
for i,(t,b) in enumerate([('当前','血管表面 → WSS\n科研三维报告'),('下一步','真实影像验证\n院内批量病例分析'),('验证后','外部回顾性研究\n比较随访与术后变化'),('远期','个体生理 + 纵向更新\n数字孪生研究')]):
    x=.8+i*3.1;dot(s,x+.18,2.38,.11,'teal' if i<2 else 'orange');txt(s,x,2.83,2.65,.5,t,21,'ink',True);txt(s,x,3.65,2.65,1.15,b,18,'muted')
    if i<3:line(s,x+.36,2.38,x+2.95,2.38,'teal',2,True)
txt(s,.8,5.25,11.7,.52,'逐步补入真实生理测量、纵向影像与临床终点，检验其对决策的增量价值。',20,'blue',True)
takeaway(s,'当前是固定协议的计算代理原型；动态数字孪生是需要持续验证的远期方向。')
# 16
s=snew('让可获得的血管形态，成为可使用的力学信息','15｜开题结论')
for yy,num,t in [(2.1,'为什么做','力学信息有研究价值，逐例获取成本限制了使用。'),(3.38,'已有基础','真实解剖数据、分组验证结果、三维报告原型已建立。'),(4.66,'继续做什么','补齐独立验证、真实输入与生理边界，检验实际价值。')]:
    txt(s,.8,yy,2.5,.6,num,22,'teal',True);txt(s,3.3,yy,9.1,.75,t,22,'ink')
takeaway(s,'希望为腹主动脉—髂动脉疾病研究提供一套更易获取、可解释、可追溯的工具。')
# 17 metrics backup
s=snew('备份：模型版本、评估口径与不能混用的数字','备份 A｜答辩时按需使用','来源：X5D主线 §35–§36.5；部署34例验收。R²_cb 不等于分类准确率。')
rows=[('当前训练底座','v5.2 · X5Dcap_asym2','CV5 三 seed OOF 均值 0.7714'),('上一训练基线','v5.2 · X5Dcap','同 CV5 口径 0.7608'),('T2 配对证据','5 折 × 3 seed','折差均值 +0.0093；13/15 为正'),('回收单元留出','IND · 五 seed 集成','0.7758 vs 0.7644；非独立患者集'),('当前部署包','v5.1 · X5D 五 seed','34例历史全链路成功；中位31.7 s'),('形变合成59例','CV5 / IND 双协议','未过增益门；已搁置')]
for i,row in enumerate(rows):
    yy=1.96+i*.65
    for x,w,t in zip([.8,3.6,7.35],[2.6,3.55,5.1],row):txt(s,x,yy,w,.5,t,15,'ink')
    line(s,.8,yy+.5,12.5,yy+.5)
txt(s,.8,6.3,11.7,.5,'训练底座已更新；全量新权重和部署切换尚待单独确认。',18,'teal',True)
# 18 case backup
s=snew('备份：已有病例展示了可定位的 WSS 空间分布','备份 B｜病例展示','匿名示例 A；v5.1 单 seed 既有病例，同一峰值时刻、同一 Pa 色标；不作为 v5.2 群体精度证据。')
pic(s,A/'case_cfd_vs_prediction.png',.9,1.8,8.0,4.98)
txt(s,9.15,2.5,3.25,2.8,'观察空间分布\n\n局部预测仍有差异\n\n整体性能看分组验证',20,'muted')
# 19 refs backup
s=snew('备份：代表性参考文献与证据来源','备份 C｜研究背景与现状','文献核查见同目录 文献核查.md；本页不以跨论文指标进行排名。')
refs=[
'[1] Wanhainen et al. ESVS 2024 aorto-iliac aneurysm guidelines. DOI: 10.1016/j.ejvs.2023.11.002',
'[2] Dua & Dalman. Hemodynamic influences on abdominal aortic aneurysm disease. 2010. DOI: 10.1016/j.vph.2010.03.004',
'[3] Mutlu et al. WSS-derived parameters and AAA tissue mechanics. 2023. DOI: 10.1016/j.compbiomed.2023.106609',
'[4] Kang et al. Four-dimensional hemodynamic prediction of AAA following EVAR. 2024. DOI: 10.1063/5.0220173',
'[5] Li et al. A deep learning-based rapid strategy for computing WSS in intracranial aneurysms. 2025. DOI: 10.1063/5.0267607',
'[6] Suk et al. Mesh neural networks for SE(3)-equivariant hemodynamics estimation on the artery wall. 2024. DOI: 10.1016/j.compbiomed.2024.108328',
'[7] Rygiel et al. WSS Estimation in AAA: Towards Generalisable Neural Surrogate Models. 2025. arXiv:2507.22817（预印本）',
'[8] Griffo et al. Geometric deep learning-based coronary WSS estimation from real-world patients. 2026. DOI: 10.1016/j.compbiomed.2026.111583'
]
for i,t in enumerate(refs):txt(s,.8,1.86+i*.59,11.7,.54,t,11.5,'ink')

path=OUT/'开题报告PPT_几何驱动WSS与部署路线.pptx';prs.save(path)
summary={'slides':len(prs.slides),'notes':sum(bool(s.notes_slide.notes_text_frame.text) for s in prs.slides),'native_charts':sum(sh.has_chart for s in prs.slides for sh in s.shapes),'date':'2026-09-29','main_slides':16,'appendix_slides':3}
(OUT/'build_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2));print(json.dumps(summary,ensure_ascii=False));print(path)
