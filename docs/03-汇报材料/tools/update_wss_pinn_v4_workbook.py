#!/usr/bin/env python3
"""Fill V4 seed1234 rows 0-15 into the existing V2/V3 workbook.

Field columns reuse the official test35 evaluations.  Downstream WSS uses the
frozen Profile-Secant V3 calculator on peak-time full-volume velocity, matching
the V2/V3 test35 x 1200 protocol.  Checkpoint sensitivity compares
last_converged (identical weights to last) with the pre-registered milestone
epoch_07500.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import math
import os
import shutil
import sys
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import numpy as np


ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))

from wss_pinn.utils import atomic_write_json, sha256_file, utc_now
from wss_pinn.v4.config import load_config
from wss_pinn.v4.data import (
    V4Dataset,
    _bc_transform,
    _geometry_transform,
    _seed,
    _take,
)
from wss_pinn.v4.evaluate import matrix_config_path
from wss_pinn.v4.models import build_model
from wss_pinn.v4.physics import denormalize, strong_form_residuals

XLSX = (
    ROOT
    / "docs/03-汇报材料/"
    "WSS_PINN_V1_V2_V3_field-v4实验矩阵与指标汇总_2026-08-06.xlsx"
)
BACKUP = XLSX.with_name(XLSX.stem + "_备份_补V4_0-14前.xlsx")
BACKUP_15 = XLSX.with_name(XLSX.stem + "_备份_补V4_15前.xlsx")
OUTPUT = ROOT / "outputs/wss_pinn/audits/v4_workbook_0_14_20260823"
FIELD_METRICS = OUTPUT / "field_metrics_v4_0_15.json"
ROI_DIR = (
    ROOT
    / "outputs/wss_pinn/audits/v2_v3_linear_regression_fullpoints_20260807/anatomical_roi"
)
BASE_TOOL_PATH = ROOT / "docs/03-汇报材料/tools/update_wss_pinn_v2_v3_workbook.py"
MODE_LABELS = {
    "data": "DATA",
    "data_bc": "DATA+BC",
    "bc_pde_fixed": "BC+PDE-FIXED",
    "bc_pde_ema": "BC+PDE-EMA",
}
BACKBONE_LABELS = {"pointnet": "PointNet", "pointnetpp": "PointNet++"}
TIME_LABELS = {"steady_peak": "steady peak", "transient_81": "transient 81"}


def _load_base():
    spec = importlib.util.spec_from_file_location("v2v3_workbook_tool", BASE_TOOL_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError("cannot import V2/V3 workbook tool")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _mean_std(values: list[float]) -> dict[str, float]:
    array = np.asarray(values, dtype=np.float64)
    array = array[np.isfinite(array)]
    if not len(array):
        return {"mean": float("nan"), "std_population": float("nan"), "n": 0}
    return {
        "mean": float(np.mean(array)),
        "std_population": float(np.std(array)),
        "n": int(len(array)),
    }


def _format_mean_std(summary: dict[str, float], decimals: int = 4) -> str:
    if not summary or summary.get("n", 0) <= 0 or not np.isfinite(summary.get("mean", float("nan"))):
        return ""
    return f"{summary['mean']:.{decimals}f} ± {summary['std_population']:.{decimals}f}"


def _format_pct(summary: dict[str, float]) -> str:
    if not summary or summary.get("n", 0) <= 0 or not np.isfinite(summary.get("mean", float("nan"))):
        return ""
    return f"{100.0 * summary['mean']:.2f}% ± {100.0 * summary['std_population']:.2f}%"


def _pooled_r2(cases: list[dict[str, Any]], metric: str) -> float:
    count = 0
    true_sum = 0.0
    true_sq = 0.0
    error_sq = 0.0
    for case in cases:
        item = case["metrics"][metric]
        if item["count"] <= 0:
            continue
        count += int(item["count"])
        n = float(item["count"])
        mean = float(item["truth_mean"])
        var = float(item["truth_variance"])
        true_sum += mean * n
        true_sq += (var * n) + (mean * mean * n)
        error_sq += (float(item["rmse"]) ** 2) * n
    if count <= 0:
        return float("nan")
    total = true_sq - true_sum**2 / count
    return 1.0 - error_sq / max(total, 1e-12)


def linear_fit(truth: np.ndarray, prediction: np.ndarray) -> dict[str, float]:
    x = np.asarray(truth, dtype=np.float64).reshape(-1)
    y = np.asarray(prediction, dtype=np.float64).reshape(-1)
    finite = np.isfinite(x) & np.isfinite(y)
    x = x[finite]
    y = y[finite]
    x_mean = float(np.mean(x))
    y_mean = float(np.mean(y))
    x_c = x - x_mean
    y_c = y - y_mean
    x_ss = float(np.dot(x_c, x_c))
    y_ss = float(np.dot(y_c, y_c))
    if x_ss <= 1e-30 or y_ss <= 1e-30:
        return {"slope_a": float("nan"), "intercept_b": float("nan"), "regression_r2": float("nan")}
    slope = float(np.dot(x_c, y_c) / x_ss)
    intercept = y_mean - slope * x_mean
    fitted = y - (slope * x + intercept)
    return {
        "slope_a": slope,
        "intercept_b": intercept,
        "regression_r2": 1.0 - float(np.dot(fitted, fitted)) / y_ss,
        "count": int(len(x)),
    }


@dataclass(frozen=True)
class V4Arm:
    index: int
    experiment_id: str
    temporal_mode: str
    backbone: str
    training_mode: str
    seed: int
    config_path: Path
    run_dir: Path

    @property
    def key(self) -> str:
        return self.experiment_id

    @property
    def checkpoint_path(self) -> Path:
        return self.run_dir / "checkpoints" / "last_converged.pt"

    @property
    def official_eval_path(self) -> Path:
        return self.run_dir / "evaluation_official_last_converged_full.json"

    @property
    def milestone_eval_path(self) -> Path:
        return self.run_dir / "evaluation_official_epoch_07500_full.json"


def load_arm(index: int) -> V4Arm:
    config_path = matrix_config_path(index)
    config = load_config(config_path, require_assets=True)
    experiment = config["experiment"]
    return V4Arm(
        index=index,
        experiment_id=str(experiment["id"]),
        temporal_mode=str(experiment["temporal_mode"]),
        backbone=str(experiment["backbone"]),
        training_mode=str(experiment["training_mode"]),
        seed=int(experiment["seed"]),
        config_path=config_path,
        run_dir=config.run_dir,
    )


def _volume_case(v4_case: dict[str, Any]) -> SimpleNamespace:
    volume = json.loads(
        Path(v4_case["files"]["steady_volume_manifest"]["path"]).read_text(encoding="utf-8")
    )
    coords = np.load(volume["files"]["interior_coords"]["path"], mmap_mode="r")
    velocity = np.load(volume["files"]["velocity_m_s"]["path"], mmap_mode="r")
    is_wall = np.load(volume["files"]["interior_is_wall"]["path"], mmap_mode="r")
    return SimpleNamespace(
        case_id=str(v4_case["canonical_id"]),
        manifest=volume,
        coords=coords,
        velocity=velocity,
        is_wall=is_wall,
        peak_step=int(volume["peak_step"]),
    )


def _peak_time_s(v4_case: dict[str, Any], peak_step: int) -> float:
    transient = v4_case["files"]["transient"]
    steps = np.asarray(np.load(transient["steps"]["path"]))
    times = np.asarray(np.load(transient["times_s"]["path"]))
    hits = np.flatnonzero(steps == peak_step)
    if len(hits) != 1:
        raise ValueError(f"peak step {peak_step} not unique in {v4_case['canonical_id']}")
    return float(times[int(hits[0])])


def _support_features(dataset: V4Dataset, arrays, support_index: np.ndarray) -> np.ndarray:
    coords = np.asarray(arrays.coords[support_index], dtype=np.float32)
    geometry = np.asarray(arrays.geometry_raw[support_index], dtype=np.float32)
    if arrays.geometry_is_pretransformed:
        geometry = (
            geometry - np.asarray(dataset.stats["geometry"]["mean"], dtype=np.float32)
        ) / np.asarray(dataset.stats["geometry"]["std"], dtype=np.float32)
    else:
        geometry = _geometry_transform(geometry, dataset.stats["geometry"])
    return np.concatenate([coords, geometry], axis=1).astype(np.float32)


def official_query_truth_ranges() -> dict[str, Any]:
    """Global u/v/w/p ranges on the V4 official evaluation universe."""
    from wss_pinn.v4.evaluate import _CaseArrays

    config = load_config(matrix_config_path(0), require_assets=True)
    dataset = V4Dataset(config, roles=("test",))
    vel_min = np.full(3, np.inf)
    vel_max = np.full(3, -np.inf)
    p_min = np.inf
    p_max = -np.inf
    # Use a transient dataset for the 81-frame universe as well.
    transient_config = load_config(matrix_config_path(8), require_assets=True)
    datasets = {
        "steady_peak": dataset,
        "transient_81": V4Dataset(transient_config, roles=("test",)),
    }
    ranges: dict[str, Any] = {}
    for mode, ds in datasets.items():
        vel_min = np.full(3, np.inf)
        vel_max = np.full(3, -np.inf)
        p_min = np.inf
        p_max = -np.inf
        for index in range(len(ds)):
            arrays = _CaseArrays(ds, index)
            query = np.asarray(arrays.population, dtype=np.int64)
            if arrays.times_s is None:
                frames = [None]
            else:
                frames = list(range(len(arrays.times_s)))
            for frame in frames:
                velocity, pressure = arrays.query_truth(query, frame)
                vel_min = np.minimum(vel_min, np.min(velocity, axis=0))
                vel_max = np.maximum(vel_max, np.max(velocity, axis=0))
                p_min = min(p_min, float(np.min(pressure)))
                p_max = max(p_max, float(np.max(pressure)))
        ranges[mode] = {
            "velocity_component_range_m_s": (vel_max - vel_min).astype(np.float64),
            "pressure_range_pa": float(p_max - p_min),
        }
    return ranges


def summarize_official(report: dict[str, Any], ranges: dict[str, Any]) -> dict[str, Any]:
    cases = report["cases"]
    component_range = np.asarray(ranges["velocity_component_range_m_s"], dtype=np.float64)
    pressure_range = float(ranges["pressure_range_pa"])
    r2 = {}
    for name in ("u", "v", "w", "speed", "pressure", "near_wall_speed", "core_speed"):
        values = [float(case["metrics"][name]["r2"]) for case in cases if case["metrics"][name]["count"]]
        r2[name] = _mean_std(values)
    velocity_nmae = []
    pressure_nmae = []
    speed_mae = []
    speed_rmse = []
    for case in cases:
        component = np.asarray(
            [case["metrics"][name]["mae"] for name in ("u", "v", "w")], dtype=np.float64
        ) / np.maximum(component_range, 1e-30)
        velocity_nmae.append(float(np.mean(component)))
        pressure_nmae.append(float(case["metrics"]["pressure"]["mae"]) / max(pressure_range, 1e-30))
        speed_mae.append(float(case["metrics"]["speed"]["mae"]))
        speed_rmse.append(float(case["metrics"]["speed"]["rmse"]))
    return {
        "experiment": report["experiment"],
        "n_cases": len(cases),
        "e_rel_l2": _mean_std([float(case["e_rel_l2"]) for case in cases]),
        "r2": r2,
        "velocity_nmae": _mean_std(velocity_nmae),
        "pressure_nmae": _mean_std(pressure_nmae),
        "speed_mae": float(np.mean(speed_mae)),
        "speed_rmse": float(np.mean(speed_rmse)),
        "pooled_speed_r2": _pooled_r2(cases, "speed"),
        "pooled_pressure_r2": _pooled_r2(cases, "pressure"),
        "protocol": report.get("protocol"),
        "checkpoint": report.get("checkpoint", {}),
    }


def build_field_metrics() -> Path:
    ranges = official_query_truth_ranges()
    payload = {
        "schema_version": 1,
        "created_at": utc_now(),
        "ranges": {
            mode: {
                "velocity_component_range_m_s": value["velocity_component_range_m_s"].tolist(),
                "pressure_range_pa": value["pressure_range_pa"],
            }
            for mode, value in ranges.items()
        },
        "arms": {},
        "milestones": {},
    }
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for index in range(16):
        arm = load_arm(index)
        report = _json(arm.official_eval_path)
        payload["arms"][arm.key] = summarize_official(report, ranges[arm.temporal_mode])
        if arm.milestone_eval_path.is_file():
            payload["milestones"][arm.key] = summarize_official(
                _json(arm.milestone_eval_path), ranges[arm.temporal_mode]
            )
    atomic_write_json(FIELD_METRICS, payload)
    return FIELD_METRICS


def _predict_volume(
    model,
    encoded,
    coords: np.ndarray,
    query_time_s: float | None,
    device,
    field_stats,
    chunk_size: int,
) -> tuple[np.ndarray, np.ndarray]:
    import torch

    velocities = []
    pressures = []
    for start in range(0, len(coords), chunk_size):
        chunk = np.asarray(coords[start : start + chunk_size], dtype=np.float32)
        query = torch.as_tensor(np.array(chunk, copy=True), device=device)
        batch = torch.zeros(len(query), dtype=torch.long, device=device)
        time_tensor = None
        if query_time_s is not None:
            time_tensor = torch.full((len(query), 1), float(query_time_s), device=device)
        prediction = model.decode_query(encoded, query, batch, time_tensor)
        vel, pres = denormalize(prediction, field_stats)
        velocities.append(vel.detach().cpu().numpy().astype(np.float32))
        pressures.append(pres.detach().cpu().numpy().astype(np.float32))
    return np.concatenate(velocities), np.concatenate(pressures)


def run_arm(
    index: int,
    device_name: str,
    sample_count: int = 1200,
    chunk_size: int = 16384,
    case_limit: int | None = None,
    output_dir: Path | None = None,
) -> Path:
    import torch

    from wss_pinn.v4.evaluate import _CaseArrays

    base = _load_base()
    arm = load_arm(index)
    if not arm.checkpoint_path.is_file():
        raise FileNotFoundError(arm.checkpoint_path)
    config = load_config(arm.config_path, require_assets=True)
    dataset = V4Dataset(config, roles=("test",))
    device = torch.device(device_name)
    model = build_model(config).to(device=device, dtype=torch.float32)
    payload = torch.load(arm.checkpoint_path, map_location=device, weights_only=False)
    model.load_state_dict(payload["model"], strict=True)
    model.eval()

    out_root = Path(output_dir) if output_dir is not None else OUTPUT
    arm_dir = out_root / "arms" / arm.key
    cache_dir = arm_dir / f"point_cache_s{sample_count}"
    cache_dir.mkdir(parents=True, exist_ok=True)
    wss_configs = base._load_wss_configs()
    physics_rows = []
    fit_rows = []
    n_cases = len(dataset) if case_limit is None else min(int(case_limit), len(dataset))
    for case_index in range(n_cases):
        v4_case = dataset._case(case_index)
        arrays = _CaseArrays(dataset, case_index)
        volume = _volume_case(v4_case)
        query_time = (
            None
            if arm.temporal_mode == "steady_peak"
            else _peak_time_s(v4_case, volume.peak_step)
        )
        support_index = arrays.sample_indices(
            "eval_support", int(config["sampling"]["support_points"])
        )
        support_coords = torch.as_tensor(
            np.asarray(arrays.coords[support_index], dtype=np.float32), device=device
        )
        support_features = torch.as_tensor(
            _support_features(dataset, arrays, support_index), device=device
        )
        support_batch = torch.zeros(len(support_index), dtype=torch.long, device=device)
        bc_vector = torch.as_tensor(arrays.bc_vector, dtype=torch.float32, device=device).unsqueeze(0)
        encoded = model.encode_support(
            support_coords,
            support_features,
            support_batch,
            bc_vector,
            [arrays.case_id],
            int(config["train"]["seed"]),
        )
        pred_u, pred_p = _predict_volume(
            model,
            encoded,
            np.asarray(volume.coords),
            query_time,
            device,
            dataset.field_stats,
            chunk_size,
        )

        physics_count = min(int(config["sampling"]["physics_points"]), int(len(arrays.population)))
        rng = np.random.default_rng(_seed(int(config["train"]["seed"]), 0, arrays.case_id, "eval_physics"))
        phys_index = _take(rng, arrays.population, physics_count)
        # Residuals on the temporal evaluation population at peak / steady.
        phys_coords = torch.as_tensor(
            np.asarray(arrays.coords[phys_index], dtype=np.float32),
            device=device,
        ).requires_grad_(True)
        phys_batch = torch.zeros(len(phys_index), dtype=torch.long, device=device)
        phys_time = None
        if query_time is not None:
            phys_time = torch.full(
                (len(phys_index), 1), float(query_time), device=device
            ).requires_grad_(True)
        phys_pred = model.decode_query(encoded, phys_coords, phys_batch, phys_time)
        residuals = strong_form_residuals(
            phys_pred,
            phys_coords,
            phys_time,
            phys_batch,
            {
                "length_m": torch.as_tensor([float(v4_case["physics_scales"]["coordinate_length_m"])], device=device),
                "U_c_m_s": torch.as_tensor([float(v4_case["physics_scales"]["U_c_m_s"])], device=device),
                "L_c_m": torch.as_tensor([float(v4_case["physics_scales"]["L_c_m"])], device=device),
            },
            field_stats=dataset.field_stats,
            physics_config=config["physics"],
        )
        wall_path = Path(v4_case["files"]["boundary_faces"]["path"])
        with np.load(wall_path) as boundary:
            wall_coords = np.asarray(boundary["wall_coords"], dtype=np.float32)
        wall_u, _ = _predict_volume(
            model, encoded, wall_coords, query_time, device, dataset.field_stats, chunk_size
        )
        physics_rows.append(
            {
                "case_id": arrays.case_id,
                "continuity_rms_dimensionless": float(
                    torch.sqrt(torch.mean(torch.square(residuals["continuity"]))).detach().cpu()
                ),
                "momentum_rms_pa_per_m": float(
                    torch.sqrt(torch.mean(torch.square(residuals["momentum_phys_pa_per_m"]))).detach().cpu()
                ),
                "wall_speed_rms_m_s": float(np.sqrt(np.mean(np.square(np.linalg.norm(wall_u, axis=1))))),
            }
        )

        cache_path = cache_dir / (arrays.case_id.replace("/", "__") + ".npz")
        if not base._cache_valid(cache_path, arm, sample_count):
            wss_arrays = base._build_wss_cache_arrays(
                volume,
                base._bundle_arrays(volume),
                pred_u,
                arm,
                sample_count,
                *wss_configs[:4],
            )
            base._write_npz(cache_path, wss_arrays)

        roi_path = ROI_DIR / f"{arrays.case_id.replace('/', '__')}.npz"
        with np.load(roi_path, allow_pickle=False) as roi:
            anatomical = np.asarray(roi["anatomical_indices"], dtype=np.int64)
        truth_speed = np.linalg.norm(np.asarray(volume.velocity[anatomical], dtype=np.float64), axis=1)
        pred_speed = np.linalg.norm(np.asarray(pred_u[anatomical], dtype=np.float64), axis=1)
        fit_rows.append({"case_id": arrays.case_id, "speed": linear_fit(truth_speed, pred_speed)})
        print(
            json.dumps(
                {
                    "event": "v4_wss_case_completed",
                    "arm": arm.key,
                    "case": arrays.case_id,
                    "index": case_index,
                },
                ensure_ascii=False,
            ),
            flush=True,
        )

    wss_report, arrays_by_case = base._apply_calibrator(cache_dir)
    wss_fits = [
        linear_fit(arrays["truth_mag"], arrays["pred_mag"])
        for arrays in arrays_by_case.values()
    ]

    result = {
        "schema_version": 1,
        "created_at": utc_now(),
        "arm": arm.key,
        "checkpoint": str(arm.checkpoint_path),
        "checkpoint_sha256": sha256_file(arm.checkpoint_path),
        "wss": wss_report,
        "physics": {
            "continuity_rms_dimensionless": _mean_std(
                [row["continuity_rms_dimensionless"] for row in physics_rows]
            ),
            "momentum_rms_pa_per_m": _mean_std(
                [row["momentum_rms_pa_per_m"] for row in physics_rows]
            ),
            "wall_speed_rms_m_s": _mean_std([row["wall_speed_rms_m_s"] for row in physics_rows]),
            "cases": physics_rows,
        },
        "speed_fit": _mean_std_fit(fit_rows, "speed"),
        "wss_fit": {
            "slope_a": _mean_std([row["slope_a"] for row in wss_fits]),
            "intercept_b": _mean_std([row["intercept_b"] for row in wss_fits]),
            "regression_r2": _mean_std([row["regression_r2"] for row in wss_fits]),
            "protocol": "OLS y=a*x+b on the same test35 x 1200 WSS points",
        },
        "query_time": "steady peak volume" if arm.temporal_mode == "steady_peak" else "transient peak frame on full V1 volume",
    }
    path = arm_dir / "downstream_metrics.json"
    atomic_write_json(path, result)
    atomic_write_json(arm_dir / "wss_metrics.json", wss_report)
    return path


def _mean_std_fit(rows: list[dict[str, Any]], key: str) -> dict[str, Any]:
    return {
        "slope_a": _mean_std([row[key]["slope_a"] for row in rows]),
        "intercept_b": _mean_std([row[key]["intercept_b"] for row in rows]),
        "regression_r2": _mean_std([row[key]["regression_r2"] for row in rows]),
        "protocol": "OLS y=a*x+b on anatomical-ROI peak-volume speed",
    }


def update_workbook() -> Path:
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if not BACKUP.is_file():
        shutil.copy2(XLSX, BACKUP)
    if not BACKUP_15.is_file():
        shutil.copy2(XLSX, BACKUP_15)
    field = _json(FIELD_METRICS if FIELD_METRICS.is_file() else OUTPUT / "field_metrics_v4_0_14.json")
    wb = load_workbook(XLSX)
    ws = wb["实验矩阵汇总"]
    headers = [ws.cell(4, column).value for column in range(1, ws.max_column + 1)]
    header_index = {name: i + 1 for i, name in enumerate(headers) if name}
    extra = [
        "E_rel_l2_cb (mean±std)",
    ]
    for name in extra:
        if name not in header_index:
            column = ws.max_column + 1
            ws.cell(4, column, name)
            header_index[name] = column
            src = ws.cell(4, 16)
            dst = ws.cell(4, column)
            if src.has_style:
                dst.font = copy(src.font)
                dst.fill = copy(src.fill)
                dst.alignment = copy(src.alignment)
                dst.border = copy(src.border)
    ws.cell(1, 1, "WSS_PINN V2 / V3 / V4 实验矩阵与核心指标汇总（2026-08-28 更新）")
    note = ws.cell(3, 1)
    note.value = (
        "V4 行是 seed1234 official test35 screen：场指标来自全部评估宇宙（瞬态=eligible×81 帧）；"
        "E_rel_l2 为设计 primary。WSS 为冻结 Profile-Secant V3、峰值全场速度、test35×1200。"
        "index 0–15（seed1234 16 臂）已填。单 seed，禁止与 V2 SAME5K / V3 val-selected 裸比。"
    )

    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fills = ["FFF3E8", "FCE4D6"]
    start_row = 19
    for offset in range(16):
        row = start_row + offset
        arm = load_arm(offset)
        field_row = field["arms"][arm.key]
        down_path = OUTPUT / "arms" / arm.key / "downstream_metrics.json"
        down = _json(down_path) if down_path.is_file() else {}
        wss = down.get("wss", {})
        physics = down.get("physics", {})
        speed_fit = down.get("speed_fit", {})
        wss_fit = down.get("wss_fit", {})
        ckpt = field_row.get("checkpoint", {})
        epoch = ckpt.get("epoch")
        values = {
            "路线": "V4",
            "阶段": f"volume_uvwp_bc_rcr_v4 / {TIME_LABELS[arm.temporal_mode]}",
            "编号/臂": str(arm.index),
            "实验ID": arm.experiment_id,
            "架构/解码器": BACKBONE_LABELS[arm.backbone],
            "输入": "xyz+geom + A/Q/RCR",
            "训练模式/旋钮": MODE_LABELS[arm.training_mode],
            "Split": "train138 / test35",
            "Seed": 1234,
            "Epoch/预算": epoch if epoch is not None else "",
            "主checkpoint": "last_converged",
            "场评估集/协议": (
                "test35 official: full non-wall volume"
                if arm.temporal_mode == "steady_peak"
                else "test35 official: eligible cache x 81 frames"
            ),
            "u R²_cb (mean±std)": _format_mean_std(field_row["r2"]["u"]),
            "v R²_cb (mean±std)": _format_mean_std(field_row["r2"]["v"]),
            "w R²_cb (mean±std)": _format_mean_std(field_row["r2"]["w"]),
            "speed R²_cb (mean±std)": _format_mean_std(field_row["r2"]["speed"]),
            "speed R²_pooled (单值)": field_row["pooled_speed_r2"],
            "speed MAE(m/s)": field_row["speed_mae"],
            "speed RMSE(m/s)": field_row["speed_rmse"],
            "velocity NMAE (mean±std)": _format_pct(field_row["velocity_nmae"]),
            "p R²_cb (mean±std)": _format_mean_std(field_row["r2"]["pressure"]),
            "p R²_pooled (单值)": field_row["pooled_pressure_r2"],
            "pressure NMAE (mean±std)": _format_pct(field_row["pressure_nmae"]),
            "near-wall speed R²_cb (mean±std)": _format_mean_std(field_row["r2"]["near_wall_speed"]),
            "core speed R²_cb (mean±std)": _format_mean_std(field_row["r2"]["core_speed"]),
            "WSS R²_cb (mean±std)": _format_mean_std(wss.get("wss_r2", {})),
            "WSS scaled R²_cb (mean±std)": _format_mean_std(wss.get("wss_scaled_r2", {})),
            "WSS NMAE (mean±std)": _format_pct(wss.get("wss_nmae_range", {})),
            "WSS下游协议": "Profile-Secant V3; test35×1200 wall; peak full-volume predicted velocity",
            "continuity RMS": physics.get("continuity_rms_dimensionless", {}).get("mean"),
            "momentum RMS(Pa/m)": physics.get("momentum_rms_pa_per_m", {}).get("mean"),
            "wall RMS(m/s)": physics.get("wall_speed_rms_m_s", {}).get("mean"),
            "状态/判定": "seed1234 screen",
            "结论/备注": (
                f"primary E_rel_l2={_format_mean_std(field_row['e_rel_l2'])}; "
                "单 seed，非正式 Holm 赢家。"
            ),
            "speed fit a_cb (mean±std)": _format_mean_std(speed_fit.get("slope_a", {})),
            "speed fit b_cb (m/s, mean±std)": _format_mean_std(speed_fit.get("intercept_b", {})),
            "speed fit R²_cb (mean±std)": _format_mean_std(speed_fit.get("regression_r2", {})),
            "WSS fit a_cb (mean±std)": _format_mean_std(wss_fit.get("slope_a", {})),
            "WSS fit b_cb (Pa, mean±std)": _format_mean_std(wss_fit.get("intercept_b", {})),
            "WSS fit R²_cb (mean±std)": _format_mean_std(wss_fit.get("regression_r2", {})),
            "E_rel_l2_cb (mean±std)": _format_mean_std(field_row["e_rel_l2"]),
        }
        for name, value in values.items():
            if name not in header_index:
                continue
            cell = ws.cell(row, header_index[name], value)
            cell.border = border
            cell.fill = PatternFill("solid", fgColor=fills[offset % 2])
            cell.font = Font(name="Microsoft YaHei", size=9, color="222222")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if name in {
                "speed R²_pooled (单值)",
                "speed MAE(m/s)",
                "speed RMSE(m/s)",
                "p R²_pooled (单值)",
                "continuity RMS",
                "momentum RMS(Pa/m)",
                "wall RMS(m/s)",
            }:
                cell.number_format = "0.0000"
        ws.row_dimensions[row].height = 38

    _write_sensitivity_sheet(wb, field)
    wb.save(XLSX)
    return XLSX


def _write_sensitivity_sheet(wb, field: dict[str, Any]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    name = "V4 Checkpoint敏感性"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    navy = "17365D"
    blue = "2F75B5"
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = [
        "编号", "实验ID", "时间", "架构", "模式", "checkpoint",
        "E_rel_l2_cb (mean±std)", "speed R²_cb (mean±std)", "speed R²_pooled",
        "p R²_cb (mean±std)", "p R²_pooled",
        "near-wall speed R²_cb", "core speed R²_cb", "说明",
    ]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title = ws.cell(1, 1, "V4 checkpoint 敏感性（seed1234 index 0–15；主读数 last_converged）")
    title.font = Font(name="Microsoft YaHei", size=14, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=navy)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    note = ws.cell(
        2,
        1,
        "V4 没有 val/best_test。last 与 last_converged 的模型权值 SHA 相同，故 last 行复用 last_converged 场指标。"
        "epoch_07500 是预注册 milestone，只比较场指标，不按 test35 反选 checkpoint，也不重算 WSS。",
    )
    note.fill = PatternFill("solid", fgColor="D9EAF7")
    note.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 36
    for column, header in enumerate(headers, 1):
        cell = ws.cell(3, column, header)
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.border = border
    row = 4
    for index in range(16):
        arm = load_arm(index)
        main = field["arms"][arm.key]
        milestone = field.get("milestones", {}).get(arm.key)
        blocks = [
            ("last_converged", main, "主checkpoint"),
            ("last", main, "与 last_converged 权值相同"),
            ("epoch_07500", milestone, "固定预算 milestone"),
        ]
        for ckpt, payload, note_text in blocks:
            if payload is None:
                values = [arm.index, arm.experiment_id, arm.temporal_mode, arm.backbone, arm.training_mode, ckpt, "", "", "", "", "", "", "", "评估文件未齐"]
            else:
                values = [
                    arm.index,
                    arm.experiment_id,
                    arm.temporal_mode,
                    arm.backbone,
                    arm.training_mode,
                    ckpt,
                    _format_mean_std(payload["e_rel_l2"]),
                    _format_mean_std(payload["r2"]["speed"]),
                    payload["pooled_speed_r2"],
                    _format_mean_std(payload["r2"]["pressure"]),
                    payload["pooled_pressure_r2"],
                    _format_mean_std(payload["r2"]["near_wall_speed"]),
                    _format_mean_std(payload["r2"]["core_speed"]),
                    note_text,
                ]
            for column, value in enumerate(values, 1):
                cell = ws.cell(row, column, value)
                cell.border = border
                cell.font = Font(name="Microsoft YaHei", size=8)
                if ckpt == "last_converged":
                    cell.fill = PatternFill("solid", fgColor="E2F0D9")
            row += 1
    for column, width in enumerate([8, 36, 16, 12, 16, 16, 22, 22, 16, 20, 14, 20, 18, 28], 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B4"
    ws.sheet_view.zoomScale = 80


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=("field-metrics", "arm", "workbook"))
    parser.add_argument("--matrix-index", type=int)
    parser.add_argument("--device", default="cuda")
    parser.add_argument("--sample-count", type=int, default=1200)
    parser.add_argument("--limit-cases", type=int)
    parser.add_argument("--output-dir")
    args = parser.parse_args()
    if args.command == "field-metrics":
        path = build_field_metrics()
    elif args.command == "arm":
        if args.matrix_index is None:
            raise ValueError("--matrix-index is required")
        path = run_arm(
            args.matrix_index,
            args.device,
            sample_count=args.sample_count,
            case_limit=args.limit_cases,
            output_dir=Path(args.output_dir) if args.output_dir else None,
        )
    else:
        path = update_workbook()
    print(json.dumps({"event": "completed", "command": args.command, "path": str(path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
