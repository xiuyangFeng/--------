#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build an editable PPT deck for the WSS PointNet / PointNet++ architecture.

The deck is deliberately made from native PowerPoint shapes so every box, label,
and connector remains editable.  Architecture facts are reconciled against:

* training_wss_min/baseline_models.py
* training_wss_min/config.py
* docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx

Output:
    docs/03-汇报材料/PointNet_PointNet++网络框架与实验旋钮_20260722.pptx
"""

from __future__ import annotations

import math
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
DOCS = ROOT / "docs" / "03-汇报材料"
XLSX = DOCS / "WSS_PointNet实验矩阵与结果汇总last.xlsx"
OUT = DOCS / "PointNet_PointNet++网络框架与实验旋钮_20260722.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Noto Sans CJK SC"


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


INK = rgb("172033")
TEXT = rgb("384152")
MUTED = rgb("6B7280")
LINE = rgb("D8DEE8")
BG = rgb("F7F9FC")
WHITE = rgb("FFFFFF")
BLUE = rgb("2F78D1")
BLUE_DARK = rgb("1F5FAE")
BLUE_LIGHT = rgb("DFECFA")
GREEN = rgb("18A978")
GREEN_DARK = rgb("0D7654")
GREEN_MID = rgb("118A63")
GREEN_LIGHT = rgb("DDF4EB")
ORANGE = rgb("F0A000")
ORANGE_DARK = rgb("A86D00")
ORANGE_LIGHT = rgb("FFF1D2")
RED = rgb("E64A4A")
RED_LIGHT = rgb("FDE3E3")
PURPLE = rgb("5543B5")
PURPLE_LIGHT = rgb("E9E5FA")
GRAY_FILL = rgb("EEF1F5")
GRAY = rgb("A9B1BD")


def set_run_font(run, size: float, color=INK, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        elem = rpr.find(qn(tag))
        if elem is None:
            elem = rpr.makeelement(qn(tag), {})
            rpr.append(elem)
        elem.set("typeface", FONT)


def add_text(slide, x, y, w, h, text, *, size=12, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03,
             line_spacing=1.0, italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    set_run_font(r, size, color=color, bold=bold, italic=italic)
    return shape


def add_rich_text(slide, x, y, w, h, lines, *, margin=0.05,
                  valign=MSO_ANCHOR.TOP):
    """lines = [{runs:[(text, options)], align, space_after, line_spacing}, ...]."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, spec in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        p.line_spacing = spec.get("line_spacing", 1.0)
        if spec.get("space_before") is not None:
            p.space_before = Pt(spec["space_before"])
        if spec.get("space_after") is not None:
            p.space_after = Pt(spec["space_after"])
        for part, opts in spec.get("runs", []):
            r = p.add_run()
            r.text = part
            set_run_font(r, opts.get("size", 12), color=opts.get("color", INK),
                         bold=opts.get("bold", False), italic=opts.get("italic", False))
    return shape


def add_box(slide, x, y, w, h, text, *, fill=WHITE, line=LINE, text_color=INK,
            size=11, bold=False, radius=True, line_width=1.0,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.05):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.0
    r = p.add_run()
    r.text = text
    set_run_font(r, size, color=text_color, bold=bold)
    return shape


def add_arrow(slide, x1, y1, x2, y2, *, color=MUTED, width=1.6,
              dash=None, begin=False, end=True):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if dash is not None:
        conn.line.dash_style = dash
    if begin:
        conn.line.begin_arrowhead = True
    if end:
        conn.line.end_arrowhead = True
    return conn


def add_elbow(slide, x1, y1, xm, ym, x2, y2, *, color=MUTED, width=1.5,
              dash=None, end=True):
    add_arrow(slide, x1, y1, xm, ym, color=color, width=width, dash=dash, end=False)
    add_arrow(slide, xm, ym, x2, y2, color=color, width=width, dash=dash, end=end)


def add_title(slide, title, subtitle=None, *, accent=BLUE, title_size=22):
    add_text(slide, 0.48, 0.18, 12.35, 0.50, title, size=title_size, bold=True,
             valign=MSO_ANCHOR.MIDDLE)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(0.74),
                                  Inches(1.15), Inches(0.035))
    rule.fill.solid(); rule.fill.fore_color.rgb = accent
    rule.line.fill.background()
    if subtitle:
        add_text(slide, 0.48, 0.82, 12.35, 0.34, subtitle, size=11.4, color=MUTED)


def add_footer(slide, page: int, note="架构依据：baseline_models.py / config.py；实验取值依据：WSS_PointNet实验矩阵与结果汇总last.xlsx"):
    add_text(slide, 0.48, 7.18, 11.8, 0.2, note, size=7.5, color=MUTED,
             valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 12.35, 7.18, 0.48, 0.2, f"{page:02d}", size=8, color=MUTED,
             align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)


def add_pill(slide, x, y, w, text, *, fill=ORANGE_LIGHT, line=ORANGE,
             text_color=ORANGE_DARK, size=9, bold=True):
    return add_box(slide, x, y, w, 0.36, text, fill=fill, line=line,
                   text_color=text_color, size=size, bold=bold, line_width=1.0)


def add_callout(slide, x, y, w, h, title, body, *, target=None):
    add_box(slide, x, y, w, h, "", fill=ORANGE_LIGHT, line=ORANGE,
            text_color=INK, size=9, line_width=1.2)
    add_rich_text(slide, x + 0.10, y + 0.08, w - 0.20, h - 0.12, [
        {"runs": [(title, {"size": 10.1, "bold": True, "color": ORANGE_DARK})],
         "space_after": 2},
        {"runs": [(body, {"size": 8.1, "color": TEXT})], "line_spacing": 0.95},
    ])
    if target:
        tx, ty = target
        sx = min(max(tx, x + 0.18), x + w - 0.18)
        sy = y + h if ty > y + h / 2 else y
        add_arrow(slide, sx, sy, tx, ty, color=ORANGE, width=1.3)


def add_badge(slide, x, y, text, *, fill=GRAY_FILL, color=MUTED, w=None):
    w = w or max(0.55, 0.09 * len(text) + 0.30)
    return add_box(slide, x, y, w, 0.29, text, fill=fill, line=fill,
                   text_color=color, size=7.8, bold=True, line_width=0.4)


def validate_sources():
    if not XLSX.exists():
        raise FileNotFoundError(XLSX)
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    if "实验矩阵总览" not in wb.sheetnames:
        raise RuntimeError("Excel 缺少『实验矩阵总览』sheet")
    ws = wb["实验矩阵总览"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    ids = {str(row[0]) for row in rows if row and row[0]}
    required = {
        "Point容量6-256-512",
        "Q1V：Point++SA3顶点采样(SAME)",
        "Q1V raw KNN-8",
        "Q1V adaptive_cover",
        "D5-A stem 6→32→64（w64 KNN-8-cover）",
    }
    missing = sorted(required - ids)
    if missing:
        raise RuntimeError(f"Excel 关键实验行缺失：{missing}")
    wb.close()


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
    return slide


def slide_overview(prs):
    slide = blank_slide(prs)
    add_title(slide, "PointNet vs PointNet++：两条 WSS 点云回归主干",
              "蓝色＝PointNet 全局聚合；绿色＝PointNet++ 局部分层；橙色＝本项目实际改过的实验旋钮")

    # PointNet panel.
    add_box(slide, 0.48, 1.30, 6.05, 4.95, "", fill=BG, line=BLUE_LIGHT, line_width=1.2)
    add_text(slide, 0.72, 1.48, 5.6, 0.34, "PointNet｜全局上下文", size=17, bold=True, color=BLUE_DARK)
    add_text(slide, 0.72, 1.83, 5.45, 0.32, "每个点先独立编码，再用一次 global max 汇总整例信息", size=9.5, color=MUTED)
    xs = [0.75, 1.92, 3.18, 4.45, 5.55]
    boxes = [
        ("输入\nN×6", BLUE_LIGHT, BLUE_DARK, 0.86),
        ("共享 MLP\n逐点编码", BLUE, WHITE, 1.02),
        ("global\nmax-pool", ORANGE, WHITE, 1.02),
        ("local⊕global\n拼接", PURPLE, WHITE, 0.94),
        ("逐点解码\n→ WSS", RED, WHITE, 0.76),
    ]
    for i, (label, fill, tc, w) in enumerate(boxes):
        add_box(slide, xs[i], 2.50, w, 1.10, label, fill=fill, line=fill,
                text_color=tc, size=9.2, bold=True)
        if i < len(boxes) - 1:
            add_arrow(slide, xs[i] + w, 3.05, xs[i + 1] - 0.06, 3.05, color=BLUE, width=1.8)
    add_pill(slide, 0.78, 4.05, 1.65, "width / depth")
    add_pill(slide, 2.56, 4.05, 1.62, "support sampling")
    add_pill(slide, 4.30, 4.05, 1.53, "SAME / SEP")
    add_box(slide, 0.78, 4.62, 5.30, 0.92,
            "关键区别：PointNet 不建立局部邻域，因此没有 SA、ball query 或 KNN 层。",
            fill=WHITE, line=BLUE_LIGHT, text_color=BLUE_DARK, size=10.0, bold=True)

    # PointNet++ panel.
    add_box(slide, 6.80, 1.30, 6.05, 4.95, "", fill=BG, line=GREEN_LIGHT, line_width=1.2)
    add_text(slide, 7.04, 1.48, 5.6, 0.34, "PointNet++｜局部分层上下文", size=17, bold=True, color=GREEN_DARK)
    add_text(slide, 7.04, 1.83, 5.45, 0.32, "逐层选中心、组邻域、聚合，再通过 FP 插值回到逐点输出", size=9.5, color=MUTED)
    labels = ["support\nsample", "SA1\n500 centers", "SA2\n125", "SA3\n32", "FP×3\n3-NN", "WSS"]
    widths = [0.76, 0.92, 0.72, 0.72, 0.82, 0.58]
    x = 7.05
    for i, (label, w) in enumerate(zip(labels, widths)):
        fill = GREEN_LIGHT if i == 0 else (RED if i == 5 else GREEN)
        tc = INK if i == 0 else WHITE
        add_box(slide, x, 2.50, w, 1.10, label, fill=fill, line=GREEN if i == 0 else fill,
                text_color=tc, size=7.5, bold=True)
        if i < len(labels) - 1:
            add_arrow(slide, x + w, 3.05, x + w + 0.12, 3.05, color=GREEN, width=1.7)
        x += w + 0.15
    add_pill(slide, 7.08, 4.05, 1.25, "radius")
    add_pill(slide, 8.44, 4.05, 1.34, "nsample")
    add_pill(slide, 9.89, 4.05, 1.26, "ball / KNN")
    add_pill(slide, 11.26, 4.05, 1.18, "width")
    add_box(slide, 7.08, 4.62, 5.30, 0.92,
            "本项目的 KNN / cover / adaptive_cover 主要替换 SA1；SA2、SA3 仍保留 ball。",
            fill=WHITE, line=GREEN_LIGHT, text_color=GREEN_DARK, size=10.0, bold=True)

    add_box(slide, 0.48, 6.48, 12.37, 0.50,
            "汇报口径：主干网络画“层”；橙色标注画“你改过的东西”；灰色标注画“固定没有扫的东西”。",
            fill=ORANGE_LIGHT, line=ORANGE, text_color=ORANGE_DARK, size=11.0, bold=True)
    add_footer(slide, 1)


def slide_pointnet(prs):
    slide = blank_slide(prs)
    add_title(slide, "PointNet 网络：逐点编码 → 全局 max-pool → 逐点解码",
              "主展示配置 E2：输入 6 维，local 6→256→512，拼接 1024，decoder 1024→512→256→1",
              accent=BLUE, title_size=21)

    # Main architecture.
    y = 2.45
    add_box(slide, 0.62, y, 1.16, 1.05, "输入点云\nN×6\nxyzgeom", fill=BLUE_LIGHT,
            line=BLUE, text_color=INK, size=10, bold=True)
    add_arrow(slide, 1.78, y + 0.52, 2.04, y + 0.52, color=BLUE, width=2.0)
    add_box(slide, 2.04, y, 1.48, 1.05, "共享 MLP\n6 → 256", fill=BLUE, line=BLUE,
            text_color=WHITE, size=11, bold=True)
    add_arrow(slide, 3.52, y + 0.52, 3.76, y + 0.52, color=BLUE, width=2.0)
    add_box(slide, 3.76, y, 1.48, 1.05, "共享 MLP\n256 → 512", fill=BLUE_DARK,
            line=BLUE_DARK, text_color=WHITE, size=11, bold=True)
    # Local branch and global branch.
    add_arrow(slide, 5.24, y + 0.52, 5.62, y + 0.52, color=BLUE, width=2.0)
    add_box(slide, 5.62, 1.95, 1.43, 0.88, "global max\n逐病例 → 512", fill=ORANGE,
            line=ORANGE, text_color=WHITE, size=9.5, bold=True)
    add_box(slide, 5.62, 3.16, 1.43, 0.72, "逐点 local\n512", fill=BLUE_LIGHT,
            line=BLUE, text_color=BLUE_DARK, size=9.5, bold=True)
    add_arrow(slide, 5.24, y + 0.52, 5.62, 3.52, color=BLUE, width=1.5)
    add_box(slide, 7.42, y, 1.16, 1.05, "拼接\n512 + 512\n= 1024", fill=PURPLE_LIGHT,
            line=PURPLE, text_color=PURPLE, size=9.5, bold=True)
    add_arrow(slide, 7.05, 2.39, 7.42, 2.72, color=ORANGE, width=1.6)
    add_arrow(slide, 7.05, 3.52, 7.42, 3.23, color=BLUE, width=1.6)
    add_arrow(slide, 8.58, y + 0.52, 8.84, y + 0.52, color=PURPLE, width=2.0)
    add_box(slide, 8.84, y, 1.28, 1.05, "MLP\n1024→512", fill=PURPLE,
            line=PURPLE, text_color=WHITE, size=10.2, bold=True)
    add_arrow(slide, 10.12, y + 0.52, 10.36, y + 0.52, color=PURPLE, width=2.0)
    add_box(slide, 10.36, y, 1.12, 1.05, "MLP\n512→256", fill=PURPLE,
            line=PURPLE, text_color=WHITE, size=10.2, bold=True)
    add_arrow(slide, 11.48, y + 0.52, 11.72, y + 0.52, color=RED, width=2.0)
    add_box(slide, 11.72, y, 0.92, 1.05, "Linear\n256→1\nWSS", fill=RED,
            line=RED, text_color=WHITE, size=9.5, bold=True)

    # Callouts above and below the exact layers.
    add_callout(slide, 0.58, 1.28, 2.35, 0.88, "① 输入 / 采样",
                "xyz ↔ xyzgeom；FPS2000 ↔ random5000；support/query SAME 或 SEP", target=(1.18, 2.42))
    add_callout(slide, 3.02, 1.20, 2.42, 0.96, "② width / depth",
                "原始 32→64→128；E2 256→512；E4 增深 64→128→256→512", target=(4.35, 2.42))
    add_callout(slide, 8.55, 1.25, 2.70, 0.90, "③ decoder 容量",
                "与 local 宽度一起变化：256→128→64→1 / 512→256→1 / 深层 decoder", target=(9.50, 2.42))
    add_callout(slide, 0.58, 4.36, 2.70, 1.04, "④ support 点数",
                "2000 / 5000；random5000 每 epoch 重采，FPS 为固定几何覆盖", target=(1.18, 3.53))
    add_callout(slide, 3.47, 4.36, 2.72, 1.04, "⑤ global 聚合（固定）",
                "始终是 per-case max；没有局部邻域、没有 SA、没有 ball / KNN", target=(6.32, 2.86))
    add_callout(slide, 6.40, 4.36, 2.72, 1.04, "⑥ query 合同",
                "SAME：监督点=输入点；SEP：独立 query；评估统一 full-cloud query", target=(7.96, 3.54))
    add_callout(slide, 9.34, 4.36, 3.30, 1.04, "⑦ 训练目标（不属于网络层）",
                "global log-z / case-max；point/case/cohort 等权；MSE / raw-Huber", target=(12.16, 3.54))

    add_box(slide, 0.60, 5.72, 12.04, 1.14, "", fill=BG, line=LINE, line_width=1.0)
    add_rich_text(slide, 0.78, 5.84, 11.70, 0.92, [
        {"runs": [("可以在 PPT 中这样讲：", {"size": 11.5, "bold": True, "color": BLUE_DARK}),
                  ("PointNet 的实验主要改了“看多少点”和“每点特征通道有多宽/多深”；", {"size": 10.4, "color": TEXT})]},
        {"runs": [("它没有邻域搜索，所以 ball、KNN、radius、nsample 都不是 PointNet 的层。", {"size": 10.4, "bold": True, "color": INK})],
         "space_before": 2},
    ])
    add_footer(slide, 2)


def slide_pointnetpp(prs):
    slide = blank_slide(prs)
    add_title(slide, "PointNet++ SA3：FPS 选中心 + 局部分组 + 分层聚合 + 3-NN 上采样",
              "橙色标签直接标在被实验修改的层上；基准 Q1V：random5000 / SAME / FPS centers / ball16 / width32",
              accent=GREEN)

    # Encoder row.
    enc_y = 2.10
    nodes = [
        (0.55, 1.15, "support\n5000×6", GREEN_LIGHT, GREEN_DARK, GREEN),
        (1.95, 1.15, "stem\n6→32→32", GREEN_MID, WHITE, GREEN_MID),
        (3.37, 2.02, "SA1\n500 centers\nr=.05｜group≤16\nMLP 35→64→64", GREEN, WHITE, GREEN),
        (5.67, 1.88, "SA2\n125 centers\nr=.10｜ball≤16\nMLP 67→128→128", GREEN_MID, WHITE, GREEN_MID),
        (7.82, 1.88, "SA3\n32 centers\nr=.20｜ball≤16\nMLP 131→256→256", GREEN_DARK, WHITE, GREEN_DARK),
    ]
    for i, (x, w, text, fill, tc, linec) in enumerate(nodes):
        add_box(slide, x, enc_y, w, 1.17, text, fill=fill, line=linec,
                text_color=tc, size=8.6 if i >= 2 else 9.4, bold=True)
        if i < len(nodes) - 1:
            nx = nodes[i + 1][0]
            add_arrow(slide, x + w, enc_y + 0.585, nx - 0.08, enc_y + 0.585,
                      color=GREEN, width=1.8)
    add_text(slide, 0.57, 1.78, 8.9, 0.23,
             "编码器：中心采样 → 邻域分组 → [相对坐标 Δxyz + 点特征] → shared MLP → scatter-max",
             size=8.7, color=GREEN_DARK, bold=True)

    # Decoder row.
    dec_y = 4.22
    dec = [
        (7.82, 1.63, "FP3\n384→128→128", rgb("74C7AA"), INK),
        (5.78, 1.63, "FP2\n192→64→64", rgb("9BD9C4"), INK),
        (3.74, 1.63, "FP1\n96→32→32", rgb("C9ECE0"), INK),
        (1.95, 1.45, "head\n32→64→1", RED, WHITE),
        (0.55, 1.15, "query\n逐点 WSS", RED_LIGHT, RED),
    ]
    # SA3 -> FP3 and decoder arrows flow right to left.
    add_arrow(slide, 8.76, enc_y + 1.17, 8.64, dec_y - 0.08, color=GREEN, width=1.8)
    for i, (x, w, text, fill, tc) in enumerate(dec):
        add_box(slide, x, dec_y, w, 0.84, text, fill=fill, line=fill if tc == WHITE else GREEN,
                text_color=tc, size=9.0, bold=True)
        if i < len(dec) - 1:
            nx, nw, *_ = dec[i + 1]
            add_arrow(slide, x, dec_y + 0.42, nx + nw + 0.08, dec_y + 0.42,
                      color=GREEN if i < 2 else RED, width=1.8)
    add_text(slide, 3.75, 5.10, 5.85, 0.25,
             "Feature Propagation：3-NN interpolate（k=3，固定） + encoder skip 拼接",
             size=8.7, color=GREEN_DARK, bold=True, align=PP_ALIGN.CENTER)

    # Skip connections, dashed and behind callouts visually.
    add_elbow(slide, 4.38, 3.27, 4.38, 3.66, 4.56, 4.20, color=GRAY, width=1.0,
              dash=MSO_LINE_DASH_STYLE.DASH, end=True)
    add_elbow(slide, 6.61, 3.27, 6.61, 3.66, 6.59, 4.20, color=GRAY, width=1.0,
              dash=MSO_LINE_DASH_STYLE.DASH, end=True)
    add_elbow(slide, 8.78, 3.27, 8.78, 3.72, 8.64, 4.20, color=GRAY, width=1.0,
              dash=MSO_LINE_DASH_STYLE.DASH, end=True)

    # Compact knob strip aligned to exact stages.
    add_pill(slide, 0.57, 1.33, 1.15, "采样 / 点数", size=8.0)
    add_pill(slide, 1.95, 1.33, 1.15, "stem / width", size=8.0)
    add_pill(slide, 3.37, 1.33, 0.82, "center", size=7.8)
    add_pill(slide, 4.25, 1.33, 0.88, "grouping", size=7.8)
    add_pill(slide, 5.69, 1.33, 0.80, "radius", size=7.8)
    add_pill(slide, 6.55, 1.33, 0.84, "nsample", size=7.8)
    add_pill(slide, 7.82, 1.33, 1.10, "width × 4", size=7.8)
    add_pill(slide, 9.00, 1.33, 0.92, "decoder", size=7.8)

    # Detailed experiment values.
    add_box(slide, 9.98, 1.30, 2.86, 4.76, "", fill=BG, line=LINE, line_width=1.0)
    add_text(slide, 10.18, 1.50, 2.45, 0.28, "本项目实际扫过", size=13.0, bold=True, color=GREEN_DARK)
    lines = [
        ("support", "2000 / 5000 / 10000 / 全点"),
        ("sampling", "FPS / random / FPS-multistart"),
        ("centers", "500-125-32；250-125-32；125-125-32；比例中心"),
        ("center sampler", "FPS ↔ Random"),
        ("SA1 grouping", "ball / KNN / KNN-cover / adaptive-cover"),
        ("nsample / k", "8 / 10 / 16 / 32 / 64 / 128 / 256"),
        ("radius scale", "0.6× / 0.8× / 1.0× / 1.2× / 1.5×"),
        ("width / stem", "w32 ↔ w64；stem 6→32→64"),
        ("query decoder", "interpolate ↔ QAD-Lite"),
    ]
    yy = 1.90
    for name, values in lines:
        add_text(slide, 10.18, yy, 1.05, 0.22, name, size=8.0, bold=True, color=ORANGE_DARK)
        add_text(slide, 11.16, yy, 1.47, 0.39, values, size=7.5, color=TEXT)
        yy += 0.43
    add_box(slide, 10.17, 5.78, 2.48, 0.50, "SA1 常被替换；SA2 / SA3 通常固定 ball",
            fill=ORANGE_LIGHT, line=ORANGE, text_color=ORANGE_DARK, size=8.2, bold=True)

    add_box(slide, 0.55, 5.58, 9.05, 0.93,
            "层内含义：radius 决定“多远算邻居”；nsample / k 决定“每个中心最多/固定取多少邻居”；width 决定通道数；center count 决定下采样尺度。",
            fill=GREEN_LIGHT, line=GREEN, text_color=GREEN_DARK, size=10.0, bold=True)
    add_footer(slide, 3)


def draw_group_panel(slide, x, y, w, h, title, subtitle, mode, color):
    add_box(slide, x, y, w, h, "", fill=BG, line=LINE, line_width=1.0)
    add_text(slide, x + 0.12, y + 0.12, w - 0.24, 0.32, title, size=12.2, bold=True,
             color=color, align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.15, y + h - 0.58, w - 0.30, 0.42, subtitle, size=8.2,
             color=MUTED, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

    # Deterministic editable point field.
    pts = [
        (0.15, .18), (.25, .35), (.18, .65), (.38, .15), (.42, .46), (.33, .74),
        (.52, .27), (.56, .58), (.65, .12), (.70, .38), (.76, .68), (.86, .22),
        (.88, .52), (.60, .78), (.48, .84), (.20, .86), (.80, .88), (.35, .58),
    ]
    centers = [(0.30, 0.38), (0.68, 0.61)]
    px0, py0 = x + 0.20, y + 0.55
    pw, ph = w - 0.40, h - 1.25
    for px, py in pts:
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(px0 + px * pw - 0.035),
                                   Inches(py0 + py * ph - 0.035), Inches(0.07), Inches(0.07))
        c.fill.solid(); c.fill.fore_color.rgb = GRAY
        c.line.fill.background()

    center_abs = []
    for cx, cy in centers:
        center_abs.append((px0 + cx * pw, py0 + cy * ph))

    def dist(a, b):
        return math.hypot(a[0] - b[0], a[1] - b[1])

    # Edges first so center stars remain on top.
    abs_pts = [(px0 + px * pw, py0 + py * ph) for px, py in pts]
    if mode == "ball":
        radius = 0.80
        for center in center_abs:
            ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center[0] - radius),
                                           Inches(center[1] - radius), Inches(2 * radius), Inches(2 * radius))
            ring.fill.background()
            ring.line.color.rgb = color; ring.line.width = Pt(1.2)
            for p in abs_pts:
                if dist(p, center) <= radius:
                    add_arrow(slide, center[0], center[1], p[0], p[1], color=color, width=0.65, end=False)
    elif mode == "knn":
        for center in center_abs:
            for p in sorted(abs_pts, key=lambda q: dist(q, center))[:5]:
                add_arrow(slide, center[0], center[1], p[0], p[1], color=color, width=0.7, end=False)
    elif mode == "cover":
        covered = set()
        for ci, center in enumerate(center_abs):
            for p in sorted(abs_pts, key=lambda q: dist(q, center))[:5]:
                covered.add(p)
                add_arrow(slide, center[0], center[1], p[0], p[1], color=PURPLE, width=0.65, end=False)
        for p in abs_pts:
            if p not in covered:
                center = min(center_abs, key=lambda c: dist(p, c))
                add_arrow(slide, center[0], center[1], p[0], p[1], color=RED, width=0.8, end=False)
    else:  # adaptive
        for p in abs_pts:
            center = min(center_abs, key=lambda c: dist(p, c))
            add_arrow(slide, center[0], center[1], p[0], p[1], color=color, width=0.65, end=False)

    for cx, cy in center_abs:
        star = slide.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(cx - 0.08), Inches(cy - 0.08),
                                      Inches(0.16), Inches(0.16))
        star.fill.solid(); star.fill.fore_color.rgb = ORANGE
        star.line.color.rgb = ORANGE_DARK; star.line.width = Pt(0.8)


def slide_grouping(prs):
    slide = blank_slide(prs)
    add_title(slide, "SA1 邻域分组：ball、KNN、KNN-cover、adaptive-cover 到底改了什么",
              "★＝SA center；灰点＝上一层 source points；连线＝该点被分到哪个中心。实验矩阵中主要只替换 SA1。",
              accent=GREEN)
    draw_group_panel(slide, 0.48, 1.32, 3.00, 4.82, "ball query（基准）",
                     "半径 r 内取 ≤ nsample；\n有局部尺度约束，但可能漏覆盖", "ball", GREEN)
    draw_group_panel(slide, 3.60, 1.32, 3.00, 4.82, "raw KNN-k",
                     "不看半径，每个中心固定取 k 近邻；\n局部尺度可能跨得很远", "knn", PURPLE)
    draw_group_panel(slide, 6.72, 1.32, 3.00, 4.82, "KNN + coverage 修复",
                     "先 KNN，再把未覆盖点补给最近中心；\n覆盖 100%，但组间重叠可偏高", "cover", BLUE)
    draw_group_panel(slide, 9.84, 1.32, 3.00, 4.82, "adaptive_cover",
                     "先最近中心全覆盖，再按预算给第二归属；\n控制 overlap cap ≤ 1/3", "adaptive", ORANGE)

    add_box(slide, 0.48, 6.36, 12.36, 0.55,
            "一句话：ball 调“距离阈值 r + 上限 nsample”；KNN 调“固定邻居数 k”；cover 类再额外调“覆盖/重叠规则”。",
            fill=ORANGE_LIGHT, line=ORANGE, text_color=ORANGE_DARK, size=10.6, bold=True)
    add_footer(slide, 4)


def add_matrix_row(slide, y, category, knobs, values, location, takeaway, *, fill=WHITE):
    xs = [0.48, 1.95, 4.43, 8.08, 10.08]
    ws = [1.47, 2.48, 3.65, 2.00, 2.76]
    contents = [category, knobs, values, location, takeaway]
    colors = [INK, ORANGE_DARK, TEXT, GREEN_DARK, TEXT]
    sizes = [9.0, 8.7, 8.25, 8.5, 8.15]
    for x, w, text, c, sz in zip(xs, ws, contents, colors, sizes):
        add_box(slide, x, y, w, 0.85, text, fill=fill, line=LINE, text_color=c,
                size=sz, bold=(x in (xs[0], xs[1], xs[3])), radius=False,
                line_width=0.65, align=PP_ALIGN.LEFT, margin=0.08)


def slide_knob_map(prs):
    slide = blank_slide(prs)
    add_title(slide, "实验矩阵总览：到底改了哪些东西、改在网络哪里",
              "橙色＝做过对照的旋钮；“位置”列告诉听众这个改动属于输入、网络层还是训练目标。",
              accent=ORANGE)
    headers = ["类别", "实验旋钮", "取值 / 对照范围", "在图里的位置", "当前汇报结论"]
    xs = [0.48, 1.95, 4.43, 8.08, 10.08]
    ws = [1.47, 2.48, 3.65, 2.00, 2.76]
    for x, w, text in zip(xs, ws, headers):
        add_box(slide, x, 1.28, w, 0.52, text, fill=INK, line=INK, text_color=WHITE,
                size=9.2, bold=True, radius=False, align=PP_ALIGN.LEFT, margin=0.08)

    rows = [
        ("输入与数据", "输入特征 / 数据池", "xyz ↔ xyzgeom；cohort one-hot；AG / AAA / ILO 扩容", "网络输入之前", "数据与特征改变信息来源，不是网络层"),
        ("采样合同", "support sampling / size", "FPS2000；random5000；FPS-multistart5000；10000；全点", "Point / Point++ 输入", "random5000 每 epoch 重采是主要增益来源之一"),
        ("监督合同", "query SAME / SEP", "SAME；独立 query（SEP / INDEPENDENT）；评估 full-cloud query", "编码后到输出", "PointNet 与 Point++ 的响应不完全一致"),
        ("PointNet", "width / depth", "原始 32-64-128；E2 256-512；E4 加深；decoder 随之变化", "shared MLP + decoder", "宽通道带来一次性增益；继续加深无稳定收益"),
        ("PointNet++", "stem / width", "w32 ↔ w64；stem 6→32→32 / 6→32→64", "stem + SA/FP 通道", "w64 参数约 4×，未形成稳定收益"),
        ("PointNet++", "SA center 数 / 采样", "500-125-32；250-125-32；125-125-32；比例中心；FPS ↔ Random", "每个 SA 的下采样", "改变空间分辨率；随机中心整体回退"),
        ("PointNet++", "grouping / radius / nsample", "ball / KNN / KNN-cover / adaptive；r×0.6–1.5；k/n=8–256", "主要在 SA1", "覆盖率提高不等于 WSS 精度提高；ball16 仍是锚点"),
        ("解码与目标", "decoder / normalization / loss", "3-NN interpolate ↔ QAD-Lite；global log-z / case-max；MSE / raw-Huber；case/cohort balancing", "FP/输出后 + loss", "QAD 无稳定增益；目标权重属于训练协议"),
    ]
    y = 1.80
    for i, row in enumerate(rows):
        add_matrix_row(slide, y, *row, fill=WHITE if i % 2 == 0 else BG)
        y += 0.63

    add_box(slide, 0.48, 6.94, 12.36, 0.23,
            "固定但要画出来：PointNet++ FP 插值 k=3；SA2 / SA3 通常保持 ball；所有模型最终逐点输出 1 个 WSS 标量。",
            fill=GRAY_FILL, line=GRAY_FILL, text_color=MUTED, size=8.2, bold=True, radius=False)
    add_footer(slide, 5)


def build():
    validate_sources()
    prs = new_prs()
    slide_overview(prs)
    slide_pointnet(prs)
    slide_pointnetpp(prs)
    slide_grouping(prs)
    slide_knob_map(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
