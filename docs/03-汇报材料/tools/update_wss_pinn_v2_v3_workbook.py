#!/usr/bin/env python3
"""Rebuild the V2/V3 workbook metrics and evaluate downstream WSS.

The script is intentionally resumable.  Its subcommands are:

``field-metrics``
    Re-express the frozen V2/V3 evaluations as case mean +/- population std,
    including paper-aligned velocity/pressure range NMAE.
``arm``
    Predict the complete peak velocity field for one frozen arm, run the
    frozen Profile-Secant V3 downstream WSS calculator on the deterministic
    test35 x 1200 wall protocol, and write per-case caches/metrics.
``export-vtp``
    For V2 E8 or V3 E6, select the best/worst sampled-WSS-NMAE cases, rerun
    full-wall WSS, and export visualization-ready VTP point clouds.
``workbook``
    Rebuild the first workbook page with only V2/V3 and update all
    case-balanced R2 cells in the checkpoint pages to mean +/- std strings,
    including downstream WSS R2 on the test35 x 1200 same-point protocol.
``verify``
    Audit the backup, workbook, 14 WSS reports, and four requested VTP files.

No checkpoint selection or training is performed.  V2 uses its frozen
last@7500 support realization; V3 uses best_validation_data selected on val15.
"""

from __future__ import annotations

import argparse
from copy import copy
from dataclasses import dataclass
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import sys
import time
import types
from typing import Any

import joblib
import numpy as np
from scipy.spatial import cKDTree


ROOT = Path(__file__).resolve().parents[3]
XLSX = (
    ROOT
    / "docs/03-汇报材料/"
    "WSS_PINN_V1_V2_V3_field-v4实验矩阵与指标汇总_2026-08-06.xlsx"
)
BACKUP = XLSX.with_name(XLSX.stem + "_备份_修改前.xlsx")
WSS_R2_BACKUP = XLSX.with_name(XLSX.stem + "_备份_补WSS_R2前.xlsx")
WSS_SCALED_R2_BACKUP = XLSX.with_name(
    XLSX.stem + "_备份_补WSS_scaled_R2前.xlsx"
)
OUTPUT = ROOT / "outputs/wss_pinn/audits/v2_v3_profile_secant_wss_20260806"
CALCULATOR_SRC = ROOT / "wss_mri_calculator/src"
CALCULATOR_EXPERIMENT = (
    ROOT / "wss_mri_calculator/experiments/pointcloud_surface_mls_v4"
)
CALIBRATOR = (
    CALCULATOR_EXPERIMENT
    / "calibrator_profile_secant_high_tail_anchor10_v3.joblib"
)
V1_CONFIG = (
    ROOT
    / "wss_mri_calculator/experiments/pointcloud_adaptive_v1/config_frozen_v1.json"
)
V3_WSS_CONFIG = (
    ROOT
    / "wss_mri_calculator/experiments/pointcloud_normal_multiscale_v3/"
    "config_frozen_v3.json"
)
V4_CONFIG = CALCULATOR_EXPERIMENT / "config_test35_cache.json"
ALGORITHM_NAME = "Profile-Secant V3 (frozen anchor10 high-tail calibrator)"
ALGORITHM_PROTOCOL = "test35 x 1200 deterministic wall points; full predicted interior field"

sys.path.insert(0, str(ROOT))


def _carreau_yasuda_numpy(
    shear_rate_s_inv: np.ndarray | float,
    *,
    mu_inf_pa_s: float = 0.0035,
    mu_zero_pa_s: float = 0.16,
    lambda_s: float = 8.2,
    a: float = 0.64,
    n: float = 0.2128,
) -> np.ndarray:
    gamma = np.asarray(shear_rate_s_inv, dtype=np.float64)
    if np.any(gamma < 0):
        raise ValueError("shear rate must be non-negative")
    return mu_inf_pa_s + (mu_zero_pa_s - mu_inf_pa_s) * np.power(
        1.0 + np.power(lambda_s * gamma, a), (n - 1.0) / a
    )


# The active WSS-PINN route exposes only the Torch rheology function, while the
# frozen downstream calculator imports the historical NumPy symbol.
from wss_pinn.physics import rheology as _active_rheology

if not hasattr(_active_rheology, "carreau_yasuda_numpy"):
    _active_rheology.carreau_yasuda_numpy = _carreau_yasuda_numpy

# The active route also archived the old STL helper.  The present workflow
# uses PCA normals only, but the frozen calculator imports the symbol at module
# load time, so provide an explicit guard instead of editing either route.
_wall_shear_compat = types.ModuleType("wss_pinn.physics.wall_shear")


def _stl_face_geometry_unavailable(*_args, **_kwargs):
    raise RuntimeError("STL normals are disabled; this workflow is frozen to PCA normals")


_wall_shear_compat.stl_face_geometry = _stl_face_geometry_unavailable
sys.modules.setdefault("wss_pinn.physics.wall_shear", _wall_shear_compat)

sys.path.insert(0, str(CALCULATOR_SRC))

from apply_v4_high_tail_calibrator import validate_package
from calculate_wss_cfd import (
    evaluate as evaluate_wss,
    fit_wall_gradient,
    pca_wall_normals,
    wss_from_gradient,
)
import data_loader_cfd as dl
from calibrate_v4_high_tail_oof import current_safe_ratio, tail_ratio
from calibrate_v4_oof import build_case_scale_features, load_cache
from v4_profile_features import base_v4_view, build_profile_features
from wss_normal_multiscale_v3 import (
    NormalMultiscaleV3Config,
    fit_wall_gradient_normal_multiscale,
)
from wss_surface_mls_v4 import SurfaceMLSV4Config, fit_wall_gradient_surface_mls

from wss_pinn.config import ExperimentConfig
from wss_pinn.data.dataset import VolumeCase, _take_strict_volume
from wss_pinn.evaluate import _case_seed
from wss_pinn.models import build_model
from wss_pinn.physics.residuals import denormalize_fields
from wss_pinn.utils import sha256_file, utc_now


@dataclass(frozen=True)
class Arm:
    key: str
    route_label: str
    arm: str
    run_id: str
    architecture: str
    input_label: str
    mode_label: str
    stage_label: str
    epochs: int
    checkpoint: str
    evaluation_name: str
    evaluation_label: str

    @property
    def run_dir(self) -> Path:
        route = (
            "volume_uvwp_peak_same5k_e7500_v2"
            if self.route_label == "V2"
            else "volume_uvwp_peak_qs_smooth_v3"
        )
        return ROOT / "outputs/wss_pinn" / route / "runs" / self.run_id

    @property
    def config_path(self) -> Path:
        return self.run_dir / "resolved_config.json"

    @property
    def checkpoint_path(self) -> Path:
        return self.run_dir / "checkpoints" / f"{self.checkpoint}.pt"

    @property
    def evaluation_path(self) -> Path:
        return self.run_dir / self.evaluation_name


ARMS = [
    Arm("V2-E1", "V2", "E1", "VF-PN-XYZ-DATA-SAME5K-E7500-s1234-v2", "PointNet", "xyz", "data-only", "SAME5K-E7500", 7500, "last", "evaluation_same5k_last.json", "test35 SAME5K (same 5000)"),
    Arm("V2-E2", "V2", "E2", "VF-PN-XYZ-PINN-SAME5K-E7500-s1234-v2", "PointNet", "xyz", "PINN", "SAME5K-E7500", 7500, "last", "evaluation_same5k_last.json", "test35 SAME5K (same 5000)"),
    Arm("V2-E3", "V2", "E3", "VF-PN-XYZG-DATA-SAME5K-E7500-s1234-v2", "PointNet", "xyz+geom", "data-only", "SAME5K-E7500", 7500, "last", "evaluation_same5k_last.json", "test35 SAME5K (same 5000)"),
    Arm("V2-E4", "V2", "E4", "VF-PN-XYZG-PINN-SAME5K-E7500-s1234-v2", "PointNet", "xyz+geom", "PINN", "SAME5K-E7500", 7500, "last", "evaluation_same5k_last.json", "test35 SAME5K (same 5000)"),
    Arm("V2-E5", "V2", "E5", "VF-PNPP-XYZ-DATA-SAME5K-E7500-s1234-v2", "PointNet++", "xyz", "data-only", "SAME5K-E7500", 7500, "last", "evaluation_same5k_last.json", "test35 SAME5K (same 5000)"),
    Arm("V2-E6", "V2", "E6", "VF-PNPP-XYZ-PINN-SAME5K-E7500-s1234-v2", "PointNet++", "xyz", "PINN", "SAME5K-E7500", 7500, "last", "evaluation_same5k_last.json", "test35 SAME5K (same 5000)"),
    Arm("V2-E7", "V2", "E7", "VF-PNPP-XYZG-DATA-SAME5K-E7500-s1234-v2", "PointNet++", "xyz+geom", "data-only", "SAME5K-E7500", 7500, "last", "evaluation_same5k_last.json", "test35 SAME5K (same 5000)"),
    Arm("V2-E8", "V2", "E8", "VF-PNPP-XYZG-PINN-SAME5K-E7500-s1234-v2", "PointNet++", "xyz+geom", "PINN", "SAME5K-E7500", 7500, "last", "evaluation_same5k_last.json", "test35 SAME5K (same 5000)"),
    Arm("V3-E1", "V3", "E1", "QSF-XYZ-DATA-s1234-v3", "PointNet smooth query decoder", "xyz", "DATA", "quasi-steady smooth field six-arm", 2500, "best_validation_data", "evaluation_best_validation_data.json", "test35 full strict-volume"),
    Arm("V3-E2", "V3", "E2", "QSF-XYZ-BC-s1234-v3", "PointNet smooth query decoder", "xyz", "DATA+BC", "quasi-steady smooth field six-arm", 2500, "best_validation_data", "evaluation_best_validation_data.json", "test35 full strict-volume"),
    Arm("V3-E3", "V3", "E3", "QSF-XYZ-BCPDE-s1234-v3", "PointNet smooth query decoder", "xyz", "DATA+BC+PDE", "quasi-steady smooth field six-arm", 2500, "best_validation_data", "evaluation_best_validation_data.json", "test35 full strict-volume"),
    Arm("V3-E4", "V3", "E4", "QSF-XYZG-DATA-s1234-v3", "PointNet smooth query decoder", "xyz+geom", "DATA", "quasi-steady smooth field six-arm", 2500, "best_validation_data", "evaluation_best_validation_data.json", "test35 full strict-volume"),
    Arm("V3-E5", "V3", "E5", "QSF-XYZG-BC-s1234-v3", "PointNet smooth query decoder", "xyz+geom", "DATA+BC", "quasi-steady smooth field six-arm", 2500, "best_validation_data", "evaluation_best_validation_data.json", "test35 full strict-volume"),
    Arm("V3-E6", "V3", "E6", "QSF-XYZG-BCPDE-s1234-v3", "PointNet smooth query decoder", "xyz+geom", "DATA+BC+PDE", "quasi-steady smooth field six-arm", 2500, "best_validation_data", "evaluation_best_validation_data.json", "test35 full strict-volume"),
]
ARM_BY_KEY = {arm.key: arm for arm in ARMS}
ARM_BY_ROUTE_AND_ARM = {(arm.route_label, arm.arm): arm for arm in ARMS}


def _json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    temporary.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    os.replace(temporary, path)


def _write_npz(path: Path, arrays: dict[str, np.ndarray]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    with temporary.open("wb") as handle:
        np.savez_compressed(handle, **arrays)
    os.replace(temporary, path)


def _mean_std(values: list[float] | np.ndarray) -> dict[str, float]:
    array = np.asarray(values, dtype=np.float64)
    array = array[np.isfinite(array)]
    if not len(array):
        return {"mean": float("nan"), "std_population": float("nan"), "n": 0}
    return {
        "mean": float(np.mean(array)),
        "std_population": float(np.std(array)),
        "n": int(len(array)),
    }


def _r2(truth: np.ndarray, prediction: np.ndarray) -> float:
    truth = np.asarray(truth, dtype=np.float64)
    prediction = np.asarray(prediction, dtype=np.float64)
    total = float(np.sum(np.square(truth - np.mean(truth))))
    return 1.0 - float(np.sum(np.square(truth - prediction))) / max(total, 1e-30)


def _range_nmae(truth: np.ndarray, prediction: np.ndarray) -> float:
    truth = np.asarray(truth, dtype=np.float64)
    prediction = np.asarray(prediction, dtype=np.float64)
    denominator = float(np.max(truth) - np.min(truth))
    return float(np.mean(np.abs(truth - prediction))) / max(denominator, 1e-30)


class LightVolumeDataset:
    """Volume-only evaluation view that intentionally skips V3 boundary assets."""

    def __init__(self, config: ExperimentConfig):
        aggregate = _json(Path(config["paths"]["sidecar_manifest"]))
        allowed = set(config["data"]["eval_roles"])
        rows = []
        for row in aggregate["cases"]:
            if row["role"] not in allowed:
                continue
            if "source_volume_manifest" in row:
                rows.append(
                    {
                        "canonical_id": row["canonical_id"],
                        "cohort": row["cohort"],
                        "role": row["role"],
                        "manifest": row["source_volume_manifest"],
                        "manifest_sha256": row["source_volume_manifest_sha256"],
                    }
                )
            else:
                rows.append(row)
        self.cases = [
            VolumeCase(row, aggregate_route="volume_uvwp_peak_v1") for row in rows
        ]
        if not self.cases:
            raise ValueError("light dataset has no evaluation cases")
        self.input_variant = config.input_variant
        stats = _json(Path(config["paths"]["field_stats"]))
        self.geometry_mean = np.asarray(stats["geometry"]["mean"], dtype=np.float32)
        self.geometry_std = np.asarray(stats["geometry"]["std"], dtype=np.float32)
        self.curvature_clip = float(stats["geometry"]["curvature_clip_abs"])

    def _features(self, coords: np.ndarray, geometry: np.ndarray) -> np.ndarray:
        xyz = np.asarray(coords, dtype=np.float32)
        if self.input_variant == "xyz":
            return xyz.copy()
        geom = np.asarray(geometry, dtype=np.float32).copy()
        geom[:, 2] = np.clip(geom[:, 2], -self.curvature_clip, self.curvature_clip)
        geom = (geom - self.geometry_mean) / self.geometry_std
        return np.concatenate([xyz, geom], axis=1).astype(np.float32)


def _load_dataset(config: ExperimentConfig) -> LightVolumeDataset:
    return LightVolumeDataset(config)


def _evaluation_indices(case, config: ExperimentConfig, route_label: str) -> np.ndarray:
    strict = np.flatnonzero(~np.asarray(case.is_wall, dtype=bool))
    if route_label == "V3":
        return strict
    rng = np.random.default_rng(
        _case_seed(case.case_id, int(config["sampling"]["seed"]))
    )
    return _take_strict_volume(
        rng, case.is_wall, int(config["sampling"]["support_points"])
    )


def _global_truth_ranges(route_label: str) -> dict[str, Any]:
    arm = next(item for item in ARMS if item.route_label == route_label)
    config = ExperimentConfig.from_json(arm.config_path)
    dataset = _load_dataset(config)
    minima = np.full(3, np.inf, dtype=np.float64)
    maxima = np.full(3, -np.inf, dtype=np.float64)
    pressure_min = np.inf
    pressure_max = -np.inf
    per_case_indices: dict[str, np.ndarray] = {}
    for case in dataset.cases:
        indices = _evaluation_indices(case, config, route_label)
        per_case_indices[case.case_id] = indices
        velocity = np.asarray(case.velocity[indices], dtype=np.float64)
        pressure = np.asarray(case.pressure[indices], dtype=np.float64)
        minima = np.minimum(minima, np.min(velocity, axis=0))
        maxima = np.maximum(maxima, np.max(velocity, axis=0))
        pressure_min = min(pressure_min, float(np.min(pressure)))
        pressure_max = max(pressure_max, float(np.max(pressure)))
    return {
        "velocity_component_min_m_s": minima.tolist(),
        "velocity_component_max_m_s": maxima.tolist(),
        "velocity_component_range_m_s": (maxima - minima).tolist(),
        "pressure_min_pa": pressure_min,
        "pressure_max_pa": pressure_max,
        "pressure_range_pa": pressure_max - pressure_min,
        "case_indices": per_case_indices,
    }


R2_METRICS = {
    "u": ("u", "velocity_u_r2"),
    "v": ("v", "velocity_v_r2"),
    "w": ("w", "velocity_w_r2"),
    "speed": ("speed", "speed_r2"),
    "pressure": ("pressure", "pressure_r2"),
    "near_wall_speed": ("near_wall_speed", None),
    "core_speed": ("core_speed", None),
}


def summarize_evaluation(
    report: dict[str, Any],
    component_range: np.ndarray,
    pressure_range: float,
) -> dict[str, Any]:
    output: dict[str, Any] = {
        "case_count": len(report["cases"]),
        "r2": {},
    }
    for output_name, (metric_name, _) in R2_METRICS.items():
        output["r2"][output_name] = _mean_std(
            [float(row["metrics"][metric_name]["r2"]) for row in report["cases"]]
        )
    velocity_nmae = []
    pressure_nmae = []
    case_rows = []
    for row in report["cases"]:
        component_values = np.asarray(
            [row["metrics"][name]["mae"] for name in ("u", "v", "w")],
            dtype=np.float64,
        ) / np.maximum(component_range, 1e-30)
        velocity_value = float(np.mean(component_values))
        pressure_value = float(row["metrics"]["pressure"]["mae"]) / max(
            pressure_range, 1e-30
        )
        velocity_nmae.append(velocity_value)
        pressure_nmae.append(pressure_value)
        case_rows.append(
            {
                "case_id": row["case_id"],
                "velocity_component_global_range_nmae": {
                    "u": float(component_values[0]),
                    "v": float(component_values[1]),
                    "w": float(component_values[2]),
                    "mean": velocity_value,
                },
                "pressure_global_range_nmae": pressure_value,
            }
        )
    output["velocity_nmae"] = _mean_std(velocity_nmae)
    output["pressure_nmae"] = _mean_std(pressure_nmae)
    output["case_balanced_regression"] = report["case_balanced_regression"]
    output["pooled_regression"] = report.get("pooled_regression", {})
    output["case_balanced"] = report["case_balanced"]
    output["cases"] = case_rows
    return output


def build_field_metrics() -> Path:
    ranges = {route: _global_truth_ranges(route) for route in ("V2", "V3")}
    payload: dict[str, Any] = {
        "schema_version": 1,
        "created_at": utc_now(),
        "status": "completed",
        "definitions": {
            "r2_std": "population std of per-case pointwise R2 across frozen test35",
            "velocity_nmae": (
                "per case: mean over u/v/w of MAE_component divided by the "
                "route/protocol test35-global component range; then case mean/std"
            ),
            "pressure_nmae": (
                "per case pressure MAE divided by the route/protocol test35-global "
                "pressure range; then case mean/std"
            ),
        },
        "routes": {},
        "arms": {},
    }
    for route in ("V2", "V3"):
        route_range = ranges[route]
        payload["routes"][route] = {
            key: value for key, value in route_range.items() if key != "case_indices"
        }
    for arm in ARMS:
        report = _json(arm.evaluation_path)
        route_range = ranges[arm.route_label]
        summary = summarize_evaluation(
            report,
            np.asarray(route_range["velocity_component_range_m_s"]),
            float(route_range["pressure_range_pa"]),
        )
        summary.update(
            {
                "route": arm.route_label,
                "arm": arm.arm,
                "run_id": arm.run_id,
                "evaluation_path": str(arm.evaluation_path),
                "evaluation_sha256": sha256_file(arm.evaluation_path),
                "evaluation_label": arm.evaluation_label,
            }
        )
        payload["arms"][arm.key] = summary
    output = OUTPUT / "field_metrics_v2_v3.json"
    _write_json(output, payload)
    print(json.dumps({"status": "completed", "output": str(output)}, ensure_ascii=False))
    return output


def _load_wss_configs() -> tuple[dict[str, Any], NormalMultiscaleV3Config, SurfaceMLSV4Config, SurfaceMLSV4Config, dict[str, Any]]:
    v1_payload = _json(V1_CONFIG)
    v3_payload = _json(V3_WSS_CONFIG)
    v4_payload = _json(V4_CONFIG)
    v3_config = NormalMultiscaleV3Config.from_mapping(v3_payload["method"])
    v4_mapping = dict(v4_payload["method"])
    if "depth_fractions" in v4_mapping:
        v4_mapping["depth_fractions"] = tuple(v4_mapping["depth_fractions"])
    v4_config = SurfaceMLSV4Config(**v4_mapping)
    v4_config.validate()
    depth3_config = SurfaceMLSV4Config(
        correction_max=1.5,
        parallel_transport=True,
        group_balance_power=0.5,
        depth_degree=3,
    )
    return v1_payload, v3_config, v4_config, depth3_config, v4_payload


def _bundle_arrays(case) -> dict[str, Any]:
    bundle_path = Path(case.manifest["parents"]["bundle"]["path"])
    with np.load(bundle_path, allow_pickle=False) as source:
        peak_step = int(source["peak_step"])
        step_index = int(np.flatnonzero(np.asarray(source["steps"]) == peak_step)[0])
        output = {
            "bundle_path": bundle_path,
            "peak_step": peak_step,
            "coord_scale_mm": float(source["coord_scale"]),
            "int_coords_norm": np.asarray(source["int_coords_norm"], dtype=np.float32),
            "wall_coords_norm": np.asarray(source["wall_coords_norm"], dtype=np.float32),
            "wall_coords_raw": np.asarray(source["wall_coords_raw"], dtype=np.float32),
            "wall_local_radius": np.asarray(source["wall_local_radius"], dtype=np.float64),
            "wall_truth_mag": np.asarray(source["wall_wss"][step_index], dtype=np.float64),
            "wall_truth_vec": np.asarray(source["wall_wss_vec"][step_index], dtype=np.float64),
            "wall_pressure_pa": np.asarray(source["wall_pressure"][step_index], dtype=np.float64),
            "rotation": np.asarray(source["transform_rotation"], dtype=np.float64),
        }
    raw_peak = Path(case.manifest["parents"]["raw_peak"]["path"])
    raw_interior = dl.load_interior(raw_peak)
    if case.manifest["alignment"]["mode"] != "direct_row":
        raise ValueError(f"test35 requires direct-row raw alignment: {case.case_id}")
    output["interior_coords_raw_mm"] = np.asarray(
        raw_interior["coords_mm"], dtype=np.float64
    )
    output["interior_velocity_raw_m_s"] = np.asarray(
        raw_interior["velocity"], dtype=np.float64
    )
    output["alignment_mode"] = "direct_row"
    output["alignment_coordinate_delta_max"] = float(
        case.manifest["alignment"]["coord_delta_max"]
    )
    raw_wall = dl.load_wall(dl.step_path(ROOT / "data_new" / case.case_id, peak_step, "ascii"))
    raw_wall_coords = np.asarray(raw_wall["coords_mm"], dtype=np.float64)
    if len(raw_wall_coords) != len(output["wall_coords_raw"]):
        raise ValueError(f"test35 raw/bundle wall count differs: {case.case_id}")
    wall_indices = np.arange(len(raw_wall_coords), dtype=np.int64)
    output["wall_coords_raw"] = raw_wall_coords[wall_indices]
    output["wall_truth_mag"] = np.asarray(raw_wall["wss_mag"], dtype=np.float64)[wall_indices]
    output["wall_truth_vec_raw"] = np.asarray(raw_wall["wss_vec"], dtype=np.float64)[wall_indices]
    output["wall_pressure_pa"] = np.asarray(raw_wall["pressure"], dtype=np.float64)[wall_indices]
    if output["int_coords_norm"].shape != case.coords.shape or not np.allclose(
        output["int_coords_norm"], np.asarray(case.coords), atol=2e-7, rtol=0
    ):
        raise ValueError(f"bundle/sidecar interior coordinates differ: {case.case_id}")
    # V3 ``case.wall_coords`` is deliberately replaced by the independent
    # boundary-contract wall asset (and can contain more points than the
    # WSS-min cropped wall).  Downstream WSS truth is aligned to the frozen
    # bundle wall, so the bundle coordinates remain authoritative here.
    return output


def _load_runtime(arm: Arm, device_name: str):
    import torch

    config = ExperimentConfig.from_json(arm.config_path)
    dataset = _load_dataset(config)
    field_stats = _json(Path(config["paths"]["field_stats"]))
    device = torch.device(device_name)
    if device.type == "cuda" and not torch.cuda.is_available():
        raise RuntimeError("CUDA was requested but is unavailable")
    dtype = torch.float64 if config["train"]["precision"] == "float64" else torch.float32
    model = build_model(config).to(device=device, dtype=dtype)
    payload = torch.load(arm.checkpoint_path, map_location=device, weights_only=False)
    model.load_state_dict(payload["model"], strict=True)
    model.eval()
    return config, dataset, field_stats, model, device, dtype


def _encode_case(case, dataset, config, model, device, dtype):
    import torch

    rng = np.random.default_rng(
        _case_seed(case.case_id, int(config["sampling"]["seed"]))
    )
    support_idx = _take_strict_volume(
        rng, case.is_wall, int(config["sampling"]["support_points"])
    )
    support_coords = torch.as_tensor(
        np.asarray(case.coords[support_idx]), device=device, dtype=dtype
    )
    support_features = torch.as_tensor(
        dataset._features(case.coords[support_idx], case.geometry[support_idx]),
        device=device,
        dtype=dtype,
    )
    support_batch = torch.zeros(len(support_idx), dtype=torch.long, device=device)
    with torch.no_grad():
        encoded = model.encode_support(
            support_coords,
            support_features,
            support_batch,
            unit_ids=[case.case_id],
            epoch=0,
            global_seed=int(config["train"]["seed"]),
            evaluation=True,
        )
    return encoded, support_idx


def _predict_coords(
    model,
    encoded,
    coords: np.ndarray,
    field_stats: dict[str, Any],
    device,
    dtype,
    chunk_size: int,
) -> tuple[np.ndarray, np.ndarray]:
    import torch

    velocities = []
    pressures = []
    for start in range(0, len(coords), int(chunk_size)):
        query = torch.as_tensor(
            np.array(coords[start : start + chunk_size], copy=True),
            device=device,
            dtype=dtype,
        )
        batch = torch.zeros(len(query), dtype=torch.long, device=device)
        with torch.no_grad():
            normalized = model.decode_query(encoded, query, batch)
            velocity, pressure = denormalize_fields(normalized, field_stats)
        velocities.append(velocity.detach().cpu().numpy().astype(np.float32))
        pressures.append(pressure.detach().cpu().numpy().astype(np.float32))
    return np.concatenate(velocities), np.concatenate(pressures)


def _build_wss_cache_arrays(
    case,
    bundle: dict[str, Any],
    predicted_velocity: np.ndarray,
    arm: Arm,
    sample_count: int,
    v1_payload: dict[str, Any],
    v3_config: NormalMultiscaleV3Config,
    v4_config: SurfaceMLSV4Config,
    depth3_config: SurfaceMLSV4Config,
    velocity_frame: str = "registered",
) -> dict[str, np.ndarray]:
    # Run the frozen calculator in its native raw-CFD millimetre frame.  Model
    # vectors are rotated back with R.T; pressure/magnitudes are invariant.
    interior_mm = np.asarray(bundle["interior_coords_raw_mm"], dtype=np.float64)
    reference_fullwall = (
        ROOT
        / "outputs/wss_mri_calculator/pointcloud_surface_mls_v4/"
        "profile_secant_v3_fullwall/inference/"
        "point_cache_test35_profile_secant_fullwall"
        / (case.case_id.replace("/", "__") + ".npz")
    )
    with np.load(reference_fullwall, allow_pickle=False) as source:
        wall_full_mm = np.asarray(source["wall_mm"], dtype=np.float64)
        full_normals = np.asarray(source["normal"], dtype=np.float64)
        full_local_radius = np.asarray(source["local_radius_mm"], dtype=np.float64)
        full_truth_mag = np.asarray(source["truth_mag"], dtype=np.float64)
        full_truth_vec = np.asarray(source["truth_vec"], dtype=np.float64)
    if velocity_frame == "registered":
        velocity_for_wss = np.asarray(predicted_velocity, dtype=np.float64) @ np.asarray(
            bundle["rotation"], dtype=np.float64
        ).T
    elif velocity_frame == "raw":
        velocity_for_wss = np.asarray(predicted_velocity, dtype=np.float64)
    else:
        raise ValueError(f"unsupported velocity_frame={velocity_frame!r}")
    interior_tree = cKDTree(interior_mm)
    if 0 < int(sample_count) < len(wall_full_mm):
        reference_sample = (
            ROOT
            / "outputs/wss_mri_calculator/pointcloud_surface_mls_v4/"
            f"point_cache_test35_profile_secant_s{int(sample_count)}"
            / (case.case_id.replace("/", "__") + ".npz")
        )
        with np.load(reference_sample, allow_pickle=False) as source:
            sample_index = np.asarray(source["sample_index"], dtype=np.int64)
    else:
        sample_index = np.arange(len(wall_full_mm), dtype=np.int64)
    wall_sample = wall_full_mm[sample_index]
    normal_sample = full_normals[sample_index]

    local_radius = full_local_radius[sample_index]
    radius_source = "frozen_profile_secant_reference_geometry"

    v1_gradient, _, _ = fit_wall_gradient(
        wall_sample,
        normal_sample,
        interior_mm,
        velocity_for_wss,
        interior_tree,
        degree=int(v1_payload["degree"]),
        neighbor_mode="adaptive_cv",
        adaptive_neighbors=tuple(int(value) for value in v1_payload["adaptive_neighbors"]),
        cv_tolerance=float(v1_payload["cv_tolerance"]),
    )
    v3_gradient, _, v3_diagnostics = fit_wall_gradient_normal_multiscale(
        wall_sample,
        normal_sample,
        wall_full_mm,
        full_normals,
        interior_mm,
        velocity_for_wss,
        interior_tree,
        local_radius_mm=local_radius,
        config=v3_config,
        fallback_gradient=v1_gradient,
    )
    _, surface_diagnostics = fit_wall_gradient_surface_mls(
        wall_sample,
        normal_sample,
        wall_full_mm,
        full_normals,
        interior_mm,
        velocity_for_wss,
        interior_tree,
        local_radius_mm=local_radius,
        fallback_gradient=v1_gradient,
        config=v4_config,
    )
    v1_norm = np.linalg.norm(v1_gradient, axis=1)
    v3_correction = np.linalg.norm(v3_gradient, axis=1) / np.maximum(v1_norm, 1e-12)
    combined_correction = np.maximum(v3_correction, surface_diagnostics["correction"])
    base_v4_gradient = v1_gradient * combined_correction[:, None]

    _, depth3_diagnostics = fit_wall_gradient_surface_mls(
        wall_sample,
        normal_sample,
        wall_full_mm,
        full_normals,
        interior_mm,
        velocity_for_wss,
        interior_tree,
        local_radius_mm=local_radius,
        fallback_gradient=v1_gradient,
        config=depth3_config,
    )
    depth3_correction = np.asarray(depth3_diagnostics["correction"], dtype=np.float64)
    fusion_correction = np.maximum(combined_correction, depth3_correction)
    fusion_gradient = v1_gradient * fusion_correction[:, None]

    arrays: dict[str, np.ndarray] = {
        "canonical_id": np.asarray(case.case_id),
        "cohort": np.asarray(case.case_id.split("/", 1)[0]),
        "step": np.asarray(bundle["peak_step"], dtype=np.int32),
        "sample_index": sample_index.astype(np.int64),
        "wall_mm": np.asarray(wall_sample, dtype=np.float32),
        "normal": np.asarray(normal_sample, dtype=np.float32),
        "truth_mag": np.asarray(full_truth_mag[sample_index], dtype=np.float32),
        "truth_vec": np.asarray(full_truth_vec[sample_index], dtype=np.float32),
        "local_radius_mm": np.asarray(local_radius, dtype=np.float32),
        "gradient_v1": np.asarray(v1_gradient, dtype=np.float32),
        "gradient_v3": np.asarray(v3_gradient, dtype=np.float32),
        "gradient_v4": np.asarray(fusion_gradient, dtype=np.float32),
        "v3_correction": np.asarray(v3_correction, dtype=np.float32),
        "combined_correction": np.asarray(fusion_correction, dtype=np.float32),
        "gradient_v4_base": np.asarray(base_v4_gradient, dtype=np.float32),
        "combined_correction_base": np.asarray(combined_correction, dtype=np.float32),
        "depth3_fusion_applied": (fusion_correction > combined_correction + 1e-12).astype(np.int8),
        "upstream_arm": np.asarray(arm.key),
        "upstream_checkpoint_sha256": np.asarray(sha256_file(arm.checkpoint_path)),
        "algorithm_name": np.asarray(ALGORITHM_NAME),
        "radius_source": np.asarray(radius_source),
    }
    for name, values in v3_diagnostics.items():
        arrays[f"v3__{name}"] = np.asarray(values)
    for name, values in surface_diagnostics.items():
        arrays[f"surface__{name}"] = np.asarray(values)
    for name, values in depth3_diagnostics.items():
        arrays[f"depth3__{name}"] = np.asarray(values)
    return arrays


def _cache_valid(path: Path, arm: Arm, sample_count: int) -> bool:
    if not path.is_file():
        return False
    try:
        with np.load(path, allow_pickle=False) as source:
            return (
                str(source["upstream_arm"]) == arm.key
                and str(source["upstream_checkpoint_sha256"])
                == sha256_file(arm.checkpoint_path)
                and len(source["sample_index"])
                == (sample_count if sample_count > 0 else len(source["sample_index"]))
            )
    except Exception:
        return False


def _apply_calibrator(cache_dir: Path) -> tuple[dict[str, Any], dict[str, dict[str, np.ndarray]]]:
    package = joblib.load(CALIBRATOR)
    data = load_cache(cache_dir)
    x, feature_names = build_profile_features(data)
    prediction_data = base_v4_view(data)
    validate_package(package, feature_names)
    groups = np.asarray(data["point_case"])
    truth = np.asarray(data["truth_mag"], dtype=np.float64)
    truth_vec = np.asarray(data["truth_vec"], dtype=np.float64)
    base_vector = wss_from_gradient(
        np.asarray(prediction_data["gradient_v4"], dtype=np.float64), "carreau"
    )
    base_mag = np.linalg.norm(base_vector, axis=1)
    predicted_rank = x[:, feature_names.index("case_pred_rank")]
    current_raw = np.clip(
        np.exp(package["current_model"].predict(x)), *package["current_raw_ratio_clip"]
    )
    tail_raw = np.clip(
        np.exp(package["tail_model"].predict(x)), *package["tail_raw_ratio_clip"]
    )
    cases, case_x = build_case_scale_features(x, feature_names, groups)
    overall_alpha = np.clip(
        np.exp(package["overall_case_model"].predict(case_x)),
        *package["case_alpha_clip"],
    )
    tail_alpha = np.clip(
        np.exp(package["tail_case_model"].predict(case_x)), *package["tail_alpha_clip"]
    )
    peak_alpha = np.clip(
        np.exp(package["peak_case_model"].predict(case_x)), *package["peak_alpha_clip"]
    )
    final_ratio = np.zeros_like(base_mag)
    for case_id, case_alpha, case_tail, case_peak in zip(
        cases, overall_alpha, tail_alpha, peak_alpha
    ):
        mask = groups == case_id
        current_ratio = current_safe_ratio(
            current_raw[mask], base_mag[mask], float(case_alpha)
        )
        final_ratio[mask] = tail_ratio(
            current_ratio,
            tail_raw[mask],
            base_mag[mask],
            predicted_rank[mask],
            float(case_tail),
            float(case_peak),
            gate_start=float(package["tail_gate_start"]),
            use_peak_anchor=float(package["peak_anchor_strength"]),
        )
    final_vector = base_vector * final_ratio[:, None]
    final_mag = np.linalg.norm(final_vector, axis=1)
    case_rows = []
    arrays_by_case: dict[str, dict[str, np.ndarray]] = {}
    for case_id in np.unique(groups):
        mask = groups == case_id
        y = truth[mask]
        p = final_mag[mask]
        nmae = _range_nmae(y, p)
        standard = evaluate_wss(y, truth_vec[mask], final_vector[mask])
        case_rows.append(
            {
                "case_id": str(case_id),
                "points": int(np.sum(mask)),
                "wss_nmae_range": nmae,
                "wss_r2": _r2(y, p),
                "wss_mae_pa": float(np.mean(np.abs(y - p))),
                "truth_range_pa": float(np.max(y) - np.min(y)),
                "calculator_metrics": standard,
            }
        )
        arrays_by_case[str(case_id)] = {
            "wall_mm": np.asarray(data["wall_mm"])[mask],
            "truth_mag": y,
            "truth_vec": truth_vec[mask],
            "pred_mag": p,
            "pred_vec": final_vector[mask],
        }
    case_rows.sort(key=lambda row: row["case_id"])
    nmae_summary = _mean_std([row["wss_nmae_range"] for row in case_rows])
    r2_summary = _mean_std([row["wss_r2"] for row in case_rows])
    scaled_r2_summary = _mean_std(
        [row["calculator_metrics"]["scaled_r2"] for row in case_rows]
    )
    payload = {
        "schema_version": 1,
        "created_at": utc_now(),
        "status": "completed",
        "algorithm": ALGORITHM_NAME,
        "calibrator": str(CALIBRATOR),
        "calibrator_sha256": sha256_file(CALIBRATOR),
        "protocol": ALGORITHM_PROTOCOL,
        "wss_nmae_range": nmae_summary,
        "wss_r2": r2_summary,
        "wss_scaled_r2": scaled_r2_summary,
        "cases": case_rows,
    }
    return payload, arrays_by_case


def run_arm(
    arm: Arm,
    device_name: str,
    sample_count: int,
    query_chunk_size: int,
    case_limit: int,
    force: bool,
) -> Path:
    import torch

    arm_dir = OUTPUT / "arms" / arm.key
    cache_dir = arm_dir / f"point_cache_s{sample_count if sample_count > 0 else 'fullwall'}"
    cache_dir.mkdir(parents=True, exist_ok=True)
    v1_payload, v3_config, v4_config, depth3_config, _ = _load_wss_configs()
    config, dataset, field_stats, model, device, dtype = _load_runtime(arm, device_name)
    cases = dataset.cases[:case_limit] if case_limit > 0 else dataset.cases
    started = time.time()
    for index, case in enumerate(cases, start=1):
        cache_path = cache_dir / (case.case_id.replace("/", "__") + ".npz")
        if not force and _cache_valid(cache_path, arm, sample_count):
            print(
                json.dumps(
                    {"event": "wss_case_cached", "arm": arm.key, "case": case.case_id, "index": index, "total": len(cases)},
                    ensure_ascii=False,
                ),
                flush=True,
            )
            continue
        encoded, _ = _encode_case(case, dataset, config, model, device, dtype)
        velocity, _ = _predict_coords(
            model,
            encoded,
            np.asarray(case.coords),
            field_stats,
            device,
            dtype,
            query_chunk_size,
        )
        bundle = _bundle_arrays(case)
        arrays = _build_wss_cache_arrays(
            case,
            bundle,
            velocity,
            arm,
            sample_count,
            v1_payload,
            v3_config,
            v4_config,
            depth3_config,
        )
        _write_npz(cache_path, arrays)
        del encoded, velocity, arrays
        if device.type == "cuda":
            torch.cuda.empty_cache()
        print(
            json.dumps(
                {"event": "wss_case_completed", "arm": arm.key, "case": case.case_id, "index": index, "total": len(cases), "elapsed_seconds": time.time() - started},
                ensure_ascii=False,
            ),
            flush=True,
        )
    if case_limit > 0:
        print(json.dumps({"status": "case_limit_completed", "cache_dir": str(cache_dir)}, ensure_ascii=False))
        return cache_dir
    report, arrays_by_case = _apply_calibrator(cache_dir)
    report.update(
        {
            "arm": arm.key,
            "run_id": arm.run_id,
            "checkpoint": str(arm.checkpoint_path),
            "checkpoint_sha256": sha256_file(arm.checkpoint_path),
            "sample_count": sample_count,
            "elapsed_seconds": time.time() - started,
        }
    )
    output = arm_dir / "wss_metrics.json"
    _write_json(output, report)
    groups = []
    truth = []
    prediction = []
    truth_vec = []
    prediction_vec = []
    wall_mm = []
    for case_id, arrays in sorted(arrays_by_case.items()):
        groups.append(np.repeat(case_id, len(arrays["truth_mag"])))
        truth.append(arrays["truth_mag"])
        prediction.append(arrays["pred_mag"])
        truth_vec.append(arrays["truth_vec"])
        prediction_vec.append(arrays["pred_vec"])
        wall_mm.append(arrays["wall_mm"])
    _write_npz(
        arm_dir / "wss_predictions_s1200.npz",
        {
            "point_case": np.concatenate(groups),
            "wall_mm": np.concatenate(wall_mm).astype(np.float32),
            "truth_mag": np.concatenate(truth).astype(np.float32),
            "pred_mag": np.concatenate(prediction).astype(np.float32),
            "truth_vec": np.concatenate(truth_vec).astype(np.float32),
            "pred_vec": np.concatenate(prediction_vec).astype(np.float32),
        },
    )
    print(
        json.dumps(
            {"status": "completed", "arm": arm.key, "output": str(output), "wss_nmae": report["wss_nmae_range"]},
            ensure_ascii=False,
        ),
        flush=True,
    )
    return output


def reproduce_truth(case_id: str, sample_count: int) -> dict[str, Any]:
    arm = ARM_BY_KEY["V3-E6"]
    config = ExperimentConfig.from_json(arm.config_path)
    dataset = _load_dataset(config)
    case = next(case for case in dataset.cases if case.case_id == case_id)
    bundle = _bundle_arrays(case)
    v1_payload, v3_config, v4_config, depth3_config, _ = _load_wss_configs()
    arrays = _build_wss_cache_arrays(
        case,
        bundle,
        np.asarray(bundle["interior_velocity_raw_m_s"], dtype=np.float64),
        arm,
        sample_count,
        v1_payload,
        v3_config,
        v4_config,
        depth3_config,
        velocity_frame="raw",
    )
    reference = (
        ROOT
        / "outputs/wss_mri_calculator/pointcloud_surface_mls_v4/"
        f"point_cache_test35_profile_secant_s{sample_count}/"
        / (case_id.replace("/", "__") + ".npz")
    )
    comparisons = {}
    with np.load(reference, allow_pickle=False) as source:
        for key in (
            "sample_index",
            "normal",
            "gradient_v1",
            "gradient_v3",
            "gradient_v4_base",
            "depth3__secant_gradient_p90",
        ):
            actual = np.asarray(arrays[key])
            expected = np.asarray(source[key])
            delta = np.abs(actual.astype(np.float64) - expected.astype(np.float64))
            comparisons[key] = {
                "shape_equal": actual.shape == expected.shape,
                "max_abs_delta": float(np.nanmax(delta)),
                "mean_abs_delta": float(np.nanmean(delta)),
            }
    payload = {
        "status": "completed",
        "case_id": case_id,
        "reference": str(reference),
        "comparisons": comparisons,
    }
    output = OUTPUT / "truth_reproduction_check.json"
    _write_json(output, payload)
    print(json.dumps(payload, indent=2, ensure_ascii=False))
    return payload


def _checkpoint_evaluation(arm: Arm, checkpoint: str) -> Path:
    if arm.route_label == "V2":
        return arm.run_dir / f"evaluation_same5k_{checkpoint}.json"
    return arm.run_dir / f"evaluation_{checkpoint}.json"


def _format_r2(report: dict[str, Any], metric: str) -> str:
    values = [float(row["metrics"][metric]["r2"]) for row in report["cases"]]
    summary = _mean_std(values)
    return f"{summary['mean']:.4f} ± {summary['std_population']:.4f}"


def _format_pct(summary: dict[str, float]) -> str:
    return f"{100.0 * summary['mean']:.2f}% ± {100.0 * summary['std_population']:.2f}%"


def _format_mean_std(summary: dict[str, float], decimals: int = 4) -> str:
    return f"{summary['mean']:.{decimals}f} ± {summary['std_population']:.{decimals}f}"


def _wss_scaled_r2_summary(report: dict[str, Any]) -> dict[str, float]:
    return _mean_std(
        [
            float(row["calculator_metrics"]["scaled_r2"])
            for row in report["cases"]
        ]
    )


def update_workbook() -> Path:
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if not BACKUP.is_file():
        raise FileNotFoundError(f"required pre-edit backup is missing: {BACKUP}")
    if not WSS_R2_BACKUP.is_file():
        shutil.copy2(XLSX, WSS_R2_BACKUP)
    if not WSS_SCALED_R2_BACKUP.is_file():
        shutil.copy2(XLSX, WSS_SCALED_R2_BACKUP)
    field_payload = _json(OUTPUT / "field_metrics_v2_v3.json")
    wss_payloads = {
        arm.key: _json(OUTPUT / "arms" / arm.key / "wss_metrics.json") for arm in ARMS
    }
    wb = load_workbook(XLSX)
    ws = wb["实验矩阵汇总"]
    current_headers = {
        ws.cell(4, column).value: column
        for column in range(1, ws.max_column + 1)
        if ws.cell(4, column).value
    }
    status_column = current_headers.get("状态/判定")
    note_column = current_headers.get("结论/备注")
    old_rows = {
        ws.cell(row, 4).value: {
            "status": ws.cell(row, status_column).value if status_column else None,
            "note": ws.cell(row, note_column).value if note_column else None,
        }
        for row in range(5, ws.max_row + 1)
        if ws.cell(row, 4).value
    }
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    ws.delete_rows(1, ws.max_row)
    ws.delete_cols(1, max(ws.max_column, 32))

    navy = "17365D"
    blue = "2F75B5"
    light_blue = "D9EAF7"
    light_green = "E2F0D9"
    light_orange = "FCE4D6"
    white = "FFFFFF"
    gray = "B7B7B7"
    thin = Side(style="thin", color=gray)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        "路线", "阶段", "编号/臂", "实验ID", "架构/解码器", "输入",
        "训练模式/旋钮", "Split", "Seed", "Epoch/预算", "主checkpoint",
        "场评估集/协议", "u R²_cb (mean±std)", "v R²_cb (mean±std)",
        "w R²_cb (mean±std)", "speed R²_cb (mean±std)",
        "speed R²_pooled (单值)", "speed MAE(m/s)", "speed RMSE(m/s)",
        "velocity NMAE (mean±std)", "p R²_cb (mean±std)",
        "p R²_pooled (单值)", "pressure NMAE (mean±std)",
        "near-wall speed R²_cb (mean±std)", "core speed R²_cb (mean±std)",
        "WSS R²_cb (mean±std)", "WSS scaled R²_cb (mean±std)",
        "WSS NMAE (mean±std)", "WSS下游协议", "continuity RMS",
        "momentum RMS(Pa/m)", "wall RMS(m/s)", "状态/判定", "结论/备注",
    ]
    last_column = len(headers)
    last_letter = get_column_letter(last_column)
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=last_column)
    title = ws.cell(1, 1, "WSS_PINN V2 / V3 实验矩阵与核心指标汇总（2026-08-06 更新）")
    title.font = Font(name="Microsoft YaHei", size=18, bold=True, color=white)
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(horizontal="center", vertical="center")
    for row in (1, 2):
        for column in range(1, last_column + 1):
            ws.cell(row, column).fill = PatternFill("solid", fgColor=navy)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)
    note = ws.cell(
        3,
        1,
        "口径：R²_cb 为 test35 逐病例点级 R² 的 mean ± population std；pooled R² 为全部评估点合并后的单值。"
        "velocity NMAE 按 test35 全局 u/v/w 分量 range 对齐论文口径；pressure 同理。"
        "WSS R²_cb 与 NMAE 均在每例 1200 个原始同点 WSS 上计算；"
        "WSS scaled R² 为逐病例用真值拟合 α 后的诊断值（y≈αp，b=0），不是正式泛化指标；"
        "NMAE=病例内 MAE/(truth max−min)。WSS 由冻结 Profile-Secant V3 对模型预测速度后处理得到，不参与训练/选模。",
    )
    note.fill = PatternFill("solid", fgColor=light_blue)
    note.font = Font(name="Microsoft YaHei", size=10, italic=True, color=navy)
    note.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 54
    for column, header in enumerate(headers, 1):
        cell = ws.cell(4, column, header)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[4].height = 45

    comments = {
        13: "逐病例 u R² 的 mean ± population std，n=35。",
        16: "逐病例 speed R² 的 mean ± population std，n=35。",
        17: "全部评估点合并计算的 pooled R²；不是病例分布，因此不附 std。",
        20: "逐病例 mean(u/v/w component MAE / test35-global component range)，再汇总 mean ± population std。",
        23: "逐病例 pressure MAE / test35-global pressure range，再汇总 mean ± population std。",
        26: "逐病例在 1200 个原始同点 WSS 上计算 R²，再对 test35 汇总 mean ± population std。为未做病例后拟合的 raw R²。",
        27: "逐病例使用该测试病例真值拟合 α=(y·p)/(p·p)，计算 y 与 αp 的 R²，b 强制为 0。仅用于诊断幅值尺度失配，不可代替 raw R²。",
        28: "逐病例 WSS MAE / 病例内 truth WSS range，再汇总 mean ± population std。",
    }
    for column, text in comments.items():
        ws.cell(4, column).comment = Comment(text, "Codex")

    for row_index, arm in enumerate(ARMS, 5):
        field = field_payload["arms"][arm.key]
        wss = wss_payloads[arm.key]
        regression = field["case_balanced_regression"]
        pooled = field.get("pooled_regression", {})
        case_balanced = field["case_balanced"]
        preserved = old_rows.get(arm.run_id, {})
        values = [
            arm.route_label,
            arm.stage_label,
            arm.arm,
            arm.run_id,
            arm.architecture,
            arm.input_label,
            arm.mode_label,
            "train138 / test35" if arm.route_label == "V2" else "train123 / val15 / test35",
            1234,
            arm.epochs,
            "last@7500" if arm.route_label == "V2" else "best_validation_data",
            arm.evaluation_label,
            _format_mean_std(field["r2"]["u"]),
            _format_mean_std(field["r2"]["v"]),
            _format_mean_std(field["r2"]["w"]),
            _format_mean_std(field["r2"]["speed"]),
            pooled.get("speed", {}).get("r2"),
            regression["speed"]["mae"],
            regression["speed"]["rmse"],
            _format_pct(field["velocity_nmae"]),
            _format_mean_std(field["r2"]["pressure"]),
            pooled.get("pressure", {}).get("r2"),
            _format_pct(field["pressure_nmae"]),
            _format_mean_std(field["r2"]["near_wall_speed"]),
            _format_mean_std(field["r2"]["core_speed"]),
            _format_mean_std(wss["wss_r2"]),
            _format_mean_std(_wss_scaled_r2_summary(wss)),
            _format_pct(wss["wss_nmae_range"]),
            "Profile-Secant V3; test35×1200 wall; full predicted interior velocity",
            case_balanced.get("continuity_rms_dimensionless"),
            case_balanced.get("momentum_rms_pa_per_m"),
            case_balanced.get("wall_speed_rms_m_s"),
            preserved.get("status"),
            preserved.get("note"),
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, value)
            cell.font = Font(name="Microsoft YaHei", size=9, color="222222")
            cell.alignment = Alignment(
                horizontal="left" if column in {4, 29, 34} else "center",
                vertical="center",
                wrap_text=column in {4, 7, 12, 29, 33, 34},
            )
            cell.border = border
            if arm.route_label == "V2":
                cell.fill = PatternFill("solid", fgColor="F7FBFF" if row_index % 2 else "EAF3F8")
            else:
                cell.fill = PatternFill("solid", fgColor="FBFDF8" if row_index % 2 else "EEF5E8")
        ws.row_dimensions[row_index].height = 38
        if arm.key in {"V2-E8", "V3-E5"}:
            for column in range(1, last_column + 1):
                ws.cell(row_index, column).fill = PatternFill("solid", fgColor=light_green)
        if arm.key in {"V3-E3", "V3-E6"}:
            ws.cell(row_index, 33).fill = PatternFill("solid", fgColor=light_orange)
        for column in (17, 18, 19, 22, 30, 31, 32):
            ws.cell(row_index, column).number_format = "0.0000"

    widths = {
        1: 8, 2: 22, 3: 8, 4: 39, 5: 27, 6: 13, 7: 22, 8: 22,
        9: 9, 10: 12, 11: 22, 12: 28, 13: 19, 14: 19, 15: 19,
        16: 21, 17: 20, 18: 16, 19: 17, 20: 22, 21: 20, 22: 18,
        23: 22, 24: 25, 25: 22, 26: 21, 27: 25, 28: 20, 29: 43,
        30: 17, 31: 21, 32: 16, 33: 20, 34: 46,
    }
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{last_letter}{4 + len(ARMS)}"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 65
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{last_letter}{4 + len(ARMS)}"

    v3_sheet = wb["V3 Checkpoint敏感性"]
    v3_sheet["A3"] = (
        "V3 保留完整 checkpoint 敏感性；所有 case-balanced R² 已改为 test35 "
        "逐病例 mean ± population std，pooled R² 仍为单值。主结论固定使用 best_validation_data。"
    )
    for column in (5, 9, 11, 12):
        if "mean±std" not in str(v3_sheet.cell(4, column).value):
            v3_sheet.cell(4, column).value += " (mean±std)"
    for column in (6, 10):
        if "单值" not in str(v3_sheet.cell(4, column).value):
            v3_sheet.cell(4, column).value += " (单值)"
    for row in range(5, v3_sheet.max_row + 1):
        arm = ARM_BY_ROUTE_AND_ARM[("V3", str(v3_sheet.cell(row, 1).value))]
        checkpoint = str(v3_sheet.cell(row, 4).value)
        report = _json(_checkpoint_evaluation(arm, checkpoint))
        for column, metric in ((5, "speed"), (9, "pressure"), (11, "near_wall_speed"), (12, "core_speed")):
            v3_sheet.cell(row, column).value = _format_r2(report, metric)

    v2_sheet = wb["V2 Checkpoint敏感性"]
    v2_sheet["A3"] = (
        "best_data / best_total 为 train-selected checkpoint 敏感性诊断；所有 case-balanced "
        "R² 已改为 test35 逐病例 mean ± population std。禁止按 test35 反选 epoch。"
    )
    for column in (6, 9, 10, 11, 15, 16, 17):
        if "mean±std" not in str(v2_sheet.cell(4, column).value):
            v2_sheet.cell(4, column).value += " (mean±std)"
    for row in range(5, v2_sheet.max_row + 1):
        arm = ARM_BY_ROUTE_AND_ARM[("V2", str(v2_sheet.cell(row, 1).value))]
        checkpoint = str(v2_sheet.cell(row, 5).value)
        report = _json(_checkpoint_evaluation(arm, checkpoint))
        for column, metric in (
            (6, "speed"), (9, "pressure"), (10, "near_wall_speed"),
            (11, "core_speed"), (15, "u"), (16, "v"), (17, "w"),
        ):
            v2_sheet.cell(row, column).value = _format_r2(report, metric)

    wb.save(XLSX)
    print(json.dumps({"status": "completed", "workbook": str(XLSX)}, ensure_ascii=False))
    return XLSX


def _select_best_worst(arm: Arm) -> list[tuple[str, str, float]]:
    report = _json(OUTPUT / "arms" / arm.key / "wss_metrics.json")
    rows = sorted(report["cases"], key=lambda row: (row["wss_nmae_range"], row["case_id"]))
    return [
        ("best", rows[0]["case_id"], float(rows[0]["wss_nmae_range"])),
        ("worst", rows[-1]["case_id"], float(rows[-1]["wss_nmae_range"])),
    ]


def _write_vtp(path: Path, points: np.ndarray, arrays: dict[str, np.ndarray], metadata: dict[str, Any]) -> None:
    import pyvista as pv

    mesh = pv.PolyData(np.asarray(points, dtype=np.float32))
    for name, values in arrays.items():
        mesh.point_data[name] = np.asarray(values)
    for name, value in metadata.items():
        if isinstance(value, str):
            mesh.field_data[name] = np.asarray([value])
        elif isinstance(value, (int, float, np.number)):
            mesh.field_data[name] = np.asarray([value])
    path.parent.mkdir(parents=True, exist_ok=True)
    mesh.save(path, binary=True)


def export_vtp(
    arm: Arm,
    device_name: str,
    query_chunk_size: int,
    interior_visual_points: int,
    force: bool,
) -> Path:
    import torch

    if arm.key not in {"V2-E8", "V3-E6"}:
        raise ValueError("VTP export is requested only for V2-E8 and V3-E6")
    selections = _select_best_worst(arm)
    config, dataset, field_stats, model, device, dtype = _load_runtime(arm, device_name)
    by_id = {case.case_id: case for case in dataset.cases}
    v1_payload, v3_config, v4_config, depth3_config, _ = _load_wss_configs()
    fullwall_cache = OUTPUT / "vtp" / arm.key / "fullwall_cache"
    fullwall_cache.mkdir(parents=True, exist_ok=True)
    predicted_by_case: dict[str, dict[str, Any]] = {}
    for label, case_id, score in selections:
        case = by_id[case_id]
        cache_path = fullwall_cache / (case_id.replace("/", "__") + ".npz")
        encoded, _ = _encode_case(case, dataset, config, model, device, dtype)
        velocity, pressure = _predict_coords(
            model,
            encoded,
            np.asarray(case.coords),
            field_stats,
            device,
            dtype,
            query_chunk_size,
        )
        bundle = _bundle_arrays(case)
        wall_velocity, wall_pressure = _predict_coords(
            model,
            encoded,
            np.asarray(bundle["wall_coords_norm"]),
            field_stats,
            device,
            dtype,
            query_chunk_size,
        )
        if force or not _cache_valid(cache_path, arm, 0):
            arrays = _build_wss_cache_arrays(
                case,
                bundle,
                velocity,
                arm,
                0,
                v1_payload,
                v3_config,
                v4_config,
                depth3_config,
            )
            _write_npz(cache_path, arrays)
        predicted_by_case[case_id] = {
            "label": label,
            "sampled_nmae": score,
            "case": case,
            "bundle": bundle,
            "velocity": velocity,
            "pressure": pressure,
            "wall_velocity": wall_velocity,
            "wall_pressure": wall_pressure,
        }
        del encoded
        if device.type == "cuda":
            torch.cuda.empty_cache()
    _, wss_by_case = _apply_calibrator(fullwall_cache)
    manifest_rows = []
    for case_id, item in predicted_by_case.items():
        case = item["case"]
        bundle = item["bundle"]
        wss = wss_by_case[case_id]
        strict = np.flatnonzero(~np.asarray(case.is_wall, dtype=bool))
        take = min(int(interior_visual_points), len(strict))
        seed = int.from_bytes(hashlib.sha256(case_id.encode("utf-8")).digest()[:8], "little")
        selected = np.sort(np.random.default_rng(seed).choice(strict, size=take, replace=False))
        rotation = np.asarray(bundle["rotation"], dtype=np.float64)
        interior_points = np.asarray(bundle["interior_coords_raw_mm"], dtype=np.float64)[selected]
        wall_points = np.asarray(bundle["wall_coords_raw"], dtype=np.float64)
        points = np.concatenate([interior_points, wall_points], axis=0)
        n_interior = len(interior_points)
        n_wall = len(wall_points)

        truth_velocity = np.concatenate(
            [
                (np.asarray(case.velocity[selected], dtype=np.float64) @ rotation.T).astype(np.float32),
                np.zeros((n_wall, 3), dtype=np.float32),
            ],
            axis=0,
        )
        pred_velocity = np.concatenate(
            [
                (np.asarray(item["velocity"][selected], dtype=np.float64) @ rotation.T).astype(np.float32),
                (np.asarray(item["wall_velocity"], dtype=np.float64) @ rotation.T).astype(np.float32),
            ],
            axis=0,
        ).astype(np.float32)
        pressure_reference = float(case.manifest["scales"]["pressure_reference_pa"])
        truth_pressure = np.concatenate(
            [np.asarray(case.pressure[selected], dtype=np.float32), (bundle["wall_pressure_pa"] - pressure_reference).astype(np.float32)],
            axis=0,
        )
        pred_pressure = np.concatenate(
            [item["pressure"][selected], item["wall_pressure"]], axis=0
        ).astype(np.float32)
        zero_wss = np.zeros((n_interior, 3), dtype=np.float32)
        truth_wss_vec = np.concatenate([zero_wss, wss["truth_vec"].astype(np.float32)], axis=0)
        pred_wss_vec = np.concatenate([zero_wss, wss["pred_vec"].astype(np.float32)], axis=0)
        truth_wss_mag = np.concatenate([np.zeros(n_interior, dtype=np.float32), wss["truth_mag"].astype(np.float32)])
        pred_wss_mag = np.concatenate([np.zeros(n_interior, dtype=np.float32), wss["pred_mag"].astype(np.float32)])
        truth_speed = np.linalg.norm(truth_velocity, axis=1).astype(np.float32)
        pred_speed = np.linalg.norm(pred_velocity, axis=1).astype(np.float32)
        point_kind = np.concatenate(
            [np.zeros(n_interior, dtype=np.int8), np.ones(n_wall, dtype=np.int8)]
        )
        field_valid = np.ones(len(points), dtype=np.int8)
        wss_valid = point_kind.copy()
        arrays = {
            "point_kind_0_volume_1_wall": point_kind,
            "field_valid": field_valid,
            "wss_valid": wss_valid,
            "velocity_truth_m_s": truth_velocity,
            "velocity_pred_m_s": pred_velocity,
            "velocity_error_m_s": pred_velocity - truth_velocity,
            "speed_truth_m_s": truth_speed,
            "speed_pred_m_s": pred_speed,
            "speed_abs_error_m_s": np.abs(pred_speed - truth_speed),
            "pressure_truth_relative_pa": truth_pressure,
            "pressure_pred_relative_pa": pred_pressure,
            "pressure_error_signed_pa": pred_pressure - truth_pressure,
            "pressure_abs_error_pa": np.abs(pred_pressure - truth_pressure),
            "wss_truth_vector_pa": truth_wss_vec,
            "wss_pred_vector_pa": pred_wss_vec,
            "wss_truth_pa": truth_wss_mag,
            "wss_pred_pa": pred_wss_mag,
            "wss_error_signed_pa": pred_wss_mag - truth_wss_mag,
            "wss_abs_error_pa": np.abs(pred_wss_mag - truth_wss_mag),
            "source_index": np.concatenate([selected.astype(np.int64), np.arange(n_wall, dtype=np.int64)]),
        }
        safe_case = case_id.replace("/", "__")
        output_path = OUTPUT / "vtp" / arm.key / f"{item['label']}__{safe_case}.vtp"
        _write_vtp(
            output_path,
            points,
            arrays,
            {
                "arm": arm.key,
                "run_id": arm.run_id,
                "case_id": case_id,
                "selection": item["label"],
                "selection_sampled_wss_nmae": item["sampled_nmae"],
                "coordinate_frame": "raw CFD frame, millimetres",
                "wss_algorithm": ALGORITHM_NAME,
            },
        )
        fullwall_nmae = _range_nmae(wss["truth_mag"], wss["pred_mag"])
        manifest_rows.append(
            {
                "selection": item["label"],
                "case_id": case_id,
                "sampled_wss_nmae": item["sampled_nmae"],
                "fullwall_wss_nmae": fullwall_nmae,
                "interior_visual_points": n_interior,
                "wall_points": n_wall,
                "vtp": str(output_path),
                "vtp_sha256": sha256_file(output_path),
            }
        )
        print(json.dumps({"event": "vtp_written", **manifest_rows[-1]}, ensure_ascii=False), flush=True)
    manifest = {
        "schema_version": 1,
        "created_at": utc_now(),
        "status": "completed",
        "arm": arm.key,
        "selection_metric": "sampled test35x1200 per-case WSS range NMAE",
        "coordinate_frame": "raw CFD frame in mm",
        "point_kind": {"0": "strict-volume visualization sample", "1": "all wall points"},
        "wss_valid_rule": "use wss_valid==1 or threshold point_kind==1 for WSS fields",
        "cases": manifest_rows,
    }
    output = OUTPUT / "vtp" / arm.key / "manifest.json"
    _write_json(output, manifest)
    return output


def verify() -> dict[str, Any]:
    from openpyxl import load_workbook
    import pyvista as pv

    failures = []
    checks: dict[str, Any] = {}
    if not BACKUP.is_file():
        failures.append("backup missing")
    else:
        checks["backup_sha256"] = sha256_file(BACKUP)
        checks["current_workbook_sha256"] = sha256_file(XLSX)
        if checks["backup_sha256"] == checks["current_workbook_sha256"]:
            failures.append("workbook is still byte-identical to backup")
    if not WSS_R2_BACKUP.is_file():
        failures.append("pre-WSS-R2 workbook backup missing")
    else:
        checks["pre_wss_r2_backup_sha256"] = sha256_file(WSS_R2_BACKUP)
    if not WSS_SCALED_R2_BACKUP.is_file():
        failures.append("pre-WSS-scaled-R2 workbook backup missing")
    else:
        checks["pre_wss_scaled_r2_backup_sha256"] = sha256_file(
            WSS_SCALED_R2_BACKUP
        )

    wb = load_workbook(XLSX, data_only=False, read_only=False)
    ws = wb["实验矩阵汇总"]
    routes = [ws.cell(row, 1).value for row in range(5, ws.max_row + 1) if ws.cell(row, 1).value]
    checks["first_sheet_routes"] = routes
    checks["first_sheet_rows"] = len(routes)
    if len(routes) != 14 or set(routes) != {"V2", "V3"}:
        failures.append("first sheet does not contain exactly 14 V2/V3 rows")
    header = [ws.cell(4, column).value for column in range(1, ws.max_column + 1)]
    for required in (
        "velocity NMAE (mean±std)",
        "pressure NMAE (mean±std)",
        "WSS R²_cb (mean±std)",
        "WSS scaled R²_cb (mean±std)",
        "WSS NMAE (mean±std)",
    ):
        if required not in header:
            failures.append(f"missing workbook header: {required}")
    for row in range(5, 19):
        for column in (13, 14, 15, 16, 21, 24, 25, 26, 27):
            if "±" not in str(ws.cell(row, column).value):
                failures.append(f"R2 mean±std missing at {ws.cell(row, column).coordinate}")

    arm_reports = []
    for arm in ARMS:
        path = OUTPUT / "arms" / arm.key / "wss_metrics.json"
        if not path.is_file():
            failures.append(f"missing WSS report: {arm.key}")
            continue
        report = _json(path)
        if len(report.get("cases", [])) != 35:
            failures.append(f"WSS report does not contain 35 cases: {arm.key}")
        scaled_summary = _wss_scaled_r2_summary(report)
        workbook_row = 5 + ARMS.index(arm)
        expected_scaled = _format_mean_std(scaled_summary)
        if ws.cell(workbook_row, 27).value != expected_scaled:
            failures.append(
                f"WSS scaled R2 mismatch at {ws.cell(workbook_row, 27).coordinate}"
            )
        arm_reports.append(
            {
                "arm": arm.key,
                "wss_nmae": report.get("wss_nmae_range"),
                "wss_raw_r2": report.get("wss_r2"),
                "wss_scaled_r2": scaled_summary,
                "sha256": sha256_file(path),
            }
        )
    checks["wss_reports"] = arm_reports

    required_arrays = {
        "point_kind_0_volume_1_wall",
        "velocity_truth_m_s",
        "velocity_pred_m_s",
        "pressure_truth_relative_pa",
        "pressure_pred_relative_pa",
        "wss_truth_pa",
        "wss_pred_pa",
        "wss_valid",
    }
    vtp_checks = []
    for arm_key in ("V2-E8", "V3-E6"):
        manifest_path = OUTPUT / "vtp" / arm_key / "manifest.json"
        if not manifest_path.is_file():
            failures.append(f"missing VTP manifest: {arm_key}")
            continue
        manifest = _json(manifest_path)
        if {row["selection"] for row in manifest["cases"]} != {"best", "worst"}:
            failures.append(f"best/worst selection incomplete: {arm_key}")
        for row in manifest["cases"]:
            path = Path(row["vtp"])
            mesh = pv.read(path)
            missing = sorted(required_arrays - set(mesh.point_data.keys()))
            if missing:
                failures.append(f"{path.name} missing arrays: {missing}")
            point_kind = np.asarray(mesh["point_kind_0_volume_1_wall"])
            if not (np.any(point_kind == 0) and np.any(point_kind == 1)):
                failures.append(f"{path.name} lacks volume or wall points")
            vtp_checks.append(
                {"path": str(path), "points": int(mesh.n_points), "arrays": sorted(mesh.point_data.keys()), "sha256": sha256_file(path)}
            )
    checks["vtp"] = vtp_checks
    payload = {
        "schema_version": 1,
        "created_at": utc_now(),
        "status": "passed" if not failures else "failed",
        "failures": failures,
        "checks": checks,
    }
    output = OUTPUT / "verification.json"
    _write_json(output, payload)
    print(json.dumps(payload, indent=2, ensure_ascii=False))
    if failures:
        raise SystemExit(1)
    return payload


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("field-metrics")

    arm_parser = subparsers.add_parser("arm")
    arm_parser.add_argument("--arm", choices=sorted(ARM_BY_KEY), required=True)
    arm_parser.add_argument("--device", default="cuda:0")
    arm_parser.add_argument("--sample-count", type=int, default=1200)
    arm_parser.add_argument("--query-chunk-size", type=int, default=16384)
    arm_parser.add_argument("--case-limit", type=int, default=0)
    arm_parser.add_argument("--force", action="store_true")

    truth_parser = subparsers.add_parser("reproduce-truth")
    truth_parser.add_argument("--case-id", default="AG/fast/FAN_JIAN_MING")
    truth_parser.add_argument("--sample-count", type=int, default=1200)

    vtp_parser = subparsers.add_parser("export-vtp")
    vtp_parser.add_argument("--arm", choices=("V2-E8", "V3-E6"), required=True)
    vtp_parser.add_argument("--device", default="cuda:0")
    vtp_parser.add_argument("--query-chunk-size", type=int, default=16384)
    vtp_parser.add_argument("--interior-visual-points", type=int, default=100000)
    vtp_parser.add_argument("--force", action="store_true")

    subparsers.add_parser("workbook")
    subparsers.add_parser("verify")
    args = parser.parse_args()
    if args.command == "field-metrics":
        build_field_metrics()
    elif args.command == "arm":
        run_arm(
            ARM_BY_KEY[args.arm],
            args.device,
            args.sample_count,
            args.query_chunk_size,
            args.case_limit,
            args.force,
        )
    elif args.command == "reproduce-truth":
        reproduce_truth(args.case_id, args.sample_count)
    elif args.command == "export-vtp":
        export_vtp(
            ARM_BY_KEY[args.arm],
            args.device,
            args.query_chunk_size,
            args.interior_visual_points,
            args.force,
        )
    elif args.command == "workbook":
        update_workbook()
    elif args.command == "verify":
        verify()


if __name__ == "__main__":
    main()
