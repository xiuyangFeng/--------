#!/usr/bin/env python3
"""Add full-point y=a*x+b regression diagnostics for frozen WSS-PINN V2/V3.

The workflow is deliberately downstream-only:

* V2 uses ``last@7500`` on the frozen test35 cases.
* V3 uses ``best_validation_data`` on the same frozen test35 cases.
* Velocity magnitude uses every strict-volume point enclosed by the capped
  anatomical STL, thereby excluding artificial CFD inlet/outlet extensions.
* WSS magnitude uses every frozen wall node and the unchanged Profile-Secant
  V3 downstream calculator.
* Existing raw R2/MAE/RMSE/NMAE columns are preserved.  The new metrics are
  ordinary least-squares slope, intercept and fitted-line R2 for ``y=a*x+b``.

The command is resumable at case granularity.  Full point arrays are consumed
for fitting and density histograms but are not duplicated to disk.  Compact
histograms, metrics and figures are retained instead.
"""

from __future__ import annotations

import argparse
import csv
from copy import copy
from dataclasses import dataclass
import importlib.util
import json
import math
import os
from pathlib import Path
import shutil
import sys
import time
from typing import Any

import numpy as np


ROOT = Path(__file__).resolve().parents[3]
BASE_TOOL_PATH = (
    ROOT / "docs/03-汇报材料/tools/update_wss_pinn_v2_v3_workbook.py"
)
CROP_TOOL_PATH = (
    ROOT
    / "docs/03-汇报材料/tools/"
    "crop_wss_pinn_velocity_pointcloud_to_anatomical_roi.py"
)
OUTPUT = (
    ROOT
    / "outputs/wss_pinn/audits/"
    "v2_v3_linear_regression_fullpoints_20260807"
)
WORKBOOK = (
    ROOT
    / "docs/03-汇报材料/"
    "WSS_PINN_V1_V2_V3_field-v4实验矩阵与指标汇总_2026-08-06.xlsx"
)
WORKBOOK_BACKUP = WORKBOOK.with_name(
    WORKBOOK.stem + "_备份_新增线性回归指标前.xlsx"
)
HISTOGRAM_BINS = 180
ROI_CHUNK_SIZE = 250_000
FILL_HOLE_SIZE_MM = 1_000_000.0
ENCLOSED_POINT_TOLERANCE = 1e-6


def _load_script(name: str, path: Path):
    specification = importlib.util.spec_from_file_location(name, path)
    if specification is None or specification.loader is None:
        raise RuntimeError(f"cannot load Python source: {path}")
    module = importlib.util.module_from_spec(specification)
    sys.modules[name] = module
    specification.loader.exec_module(module)
    return module


BASE = _load_script("wss_pinn_v2_v3_workbook_tool", BASE_TOOL_PATH)
CROP = _load_script("wss_pinn_anatomical_roi_crop_tool", CROP_TOOL_PATH)


@dataclass(frozen=True)
class FitResult:
    count: int
    slope_a: float
    intercept_b: float
    regression_r2: float
    pearson_r: float
    raw_prediction_r2: float
    x_mean: float
    y_mean: float
    x_min: float
    x_max: float
    y_min: float
    y_max: float
    x_variance: float
    y_variance: float
    fitted_sse: float
    fitted_sst_y: float

    def as_dict(self) -> dict[str, Any]:
        return {
            "count": self.count,
            "slope_a": self.slope_a,
            "intercept_b": self.intercept_b,
            "regression_r2": self.regression_r2,
            "pearson_r": self.pearson_r,
            "raw_prediction_r2": self.raw_prediction_r2,
            "x_mean": self.x_mean,
            "y_mean": self.y_mean,
            "x_min": self.x_min,
            "x_max": self.x_max,
            "y_min": self.y_min,
            "y_max": self.y_max,
            "x_variance": self.x_variance,
            "y_variance": self.y_variance,
            "fitted_sse": self.fitted_sse,
            "fitted_sst_y": self.fitted_sst_y,
            "equation": "y = a*x + b",
            "r2_definition": "1 - sum((y-(a*x+b))^2) / sum((y-mean(y))^2)",
        }


def _json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: Any) -> None:
    BASE._write_json(path, payload)


def _safe_case(case_id: str) -> str:
    return case_id.replace("/", "__")


def _sha256(path: Path) -> str:
    return BASE.sha256_file(path)


def _finite_pair(x: np.ndarray, y: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    truth = np.asarray(x, dtype=np.float64).reshape(-1)
    prediction = np.asarray(y, dtype=np.float64).reshape(-1)
    if truth.shape != prediction.shape:
        raise ValueError("regression arrays must have identical shapes")
    finite = np.isfinite(truth) & np.isfinite(prediction)
    if not bool(np.all(finite)):
        truth = truth[finite]
        prediction = prediction[finite]
    if len(truth) < 3:
        raise ValueError("regression requires at least three finite points")
    return truth, prediction


def linear_fit(x: np.ndarray, y: np.ndarray) -> FitResult:
    """Fit prediction ``y`` against CFD truth ``x`` with a free intercept."""
    truth, prediction = _finite_pair(x, y)
    x_mean = float(np.mean(truth))
    y_mean = float(np.mean(prediction))
    x_centered = truth - x_mean
    y_centered = prediction - y_mean
    x_ss = float(np.dot(x_centered, x_centered))
    y_ss = float(np.dot(y_centered, y_centered))
    if x_ss <= 1e-30 or y_ss <= 1e-30:
        raise ValueError("linear-fit R2 is undefined for a constant scalar field")
    cross = float(np.dot(x_centered, y_centered))
    slope = cross / x_ss
    intercept = y_mean - slope * x_mean
    fitted_error = prediction - (slope * truth + intercept)
    fitted_sse = float(np.dot(fitted_error, fitted_error))
    regression_r2 = 1.0 - fitted_sse / y_ss
    pearson_r = cross / math.sqrt(x_ss * y_ss)
    raw_total = x_ss
    raw_error = prediction - truth
    raw_r2 = 1.0 - float(np.dot(raw_error, raw_error)) / raw_total
    return FitResult(
        count=int(len(truth)),
        slope_a=float(slope),
        intercept_b=float(intercept),
        regression_r2=float(regression_r2),
        pearson_r=float(pearson_r),
        raw_prediction_r2=float(raw_r2),
        x_mean=x_mean,
        y_mean=y_mean,
        x_min=float(np.min(truth)),
        x_max=float(np.max(truth)),
        y_min=float(np.min(prediction)),
        y_max=float(np.max(prediction)),
        x_variance=x_ss / len(truth),
        y_variance=y_ss / len(prediction),
        fitted_sse=fitted_sse,
        fitted_sst_y=y_ss,
    )


def density_histogram(
    x: np.ndarray,
    y: np.ndarray,
    bins: int = HISTOGRAM_BINS,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    truth, prediction = _finite_pair(x, y)
    lower = min(0.0, float(np.min(truth)), float(np.min(prediction)))
    upper = max(float(np.max(truth)), float(np.max(prediction)))
    if not np.isfinite(upper) or upper <= lower:
        upper = lower + 1.0
    padding = max((upper - lower) * 0.015, 1e-12)
    value_range = ((lower - padding, upper + padding), (lower - padding, upper + padding))
    counts, x_edges, y_edges = np.histogram2d(
        truth,
        prediction,
        bins=int(bins),
        range=value_range,
    )
    if int(np.sum(counts)) != len(truth):
        raise RuntimeError("density histogram did not retain every regression point")
    return counts.astype(np.int64), x_edges, y_edges


def _plot_density(
    output: Path,
    counts: np.ndarray,
    x_edges: np.ndarray,
    y_edges: np.ndarray,
    fit: FitResult,
    *,
    case_id: str,
    arm_key: str,
    metric_label: str,
    unit: str,
    scope_label: str,
) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.colors import LogNorm

    figure, axis = plt.subplots(figsize=(6.4, 5.7), constrained_layout=True)
    masked = np.ma.masked_less_equal(counts.T, 0)
    mesh = axis.pcolormesh(
        x_edges,
        y_edges,
        masked,
        cmap="Blues",
        norm=LogNorm(vmin=1, vmax=max(int(np.max(counts)), 1)),
        shading="auto",
        rasterized=True,
    )
    lo = float(min(x_edges[0], y_edges[0]))
    hi = float(max(x_edges[-1], y_edges[-1]))
    line_x = np.asarray([lo, hi], dtype=np.float64)
    axis.plot(line_x, line_x, linestyle="--", linewidth=1.2, color="0.45", label="ideal y=x")
    axis.plot(
        line_x,
        fit.slope_a * line_x + fit.intercept_b,
        linewidth=1.8,
        color="#C43C39",
        label="OLS y=ax+b",
    )
    axis.set_xlim(lo, hi)
    axis.set_ylim(lo, hi)
    axis.set_aspect("equal", adjustable="box")
    axis.set_xlabel(f"CFD {metric_label} [{unit}]")
    axis.set_ylabel(f"Predicted {metric_label} [{unit}]")
    axis.set_title(f"{arm_key} | {case_id}")
    axis.text(
        0.035,
        0.965,
        (
            f"y = {fit.slope_a:.4g}x {fit.intercept_b:+.4g}\n"
            f"a = {fit.slope_a:.4g}, b = {fit.intercept_b:.4g}\n"
            f"R² = {fit.regression_r2:.4f}\n"
            f"n = {fit.count:,}"
        ),
        transform=axis.transAxes,
        ha="left",
        va="top",
        fontsize=9,
        bbox={"boxstyle": "round,pad=0.35", "facecolor": "white", "alpha": 0.9, "edgecolor": "0.75"},
    )
    axis.text(
        0.02,
        0.02,
        scope_label,
        transform=axis.transAxes,
        ha="left",
        va="bottom",
        fontsize=7.5,
        color="0.35",
        bbox={"boxstyle": "round,pad=0.2", "facecolor": "white", "alpha": 0.82, "edgecolor": "none"},
    )
    axis.legend(loc="lower right", fontsize=8)
    colorbar = figure.colorbar(mesh, ax=axis, pad=0.02)
    colorbar.set_label("point count per bin")
    output.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output, dpi=220)
    plt.close(figure)


def _resolve_stl(case_id: str) -> Path:
    case_dir = ROOT / "data_new" / case_id
    if case_id == "ILO/ZHANG_JIN_CHUN-1/before":
        preferred = case_dir / "0zhang_jin_chun-sq.stl"
        if preferred.is_file():
            return preferred
    case_name = Path(case_id).parts[-1]
    exact = case_dir / f"{case_name}.stl"
    if exact.is_file():
        return exact
    candidates = sorted(case_dir.glob("*.stl"))
    if len(candidates) == 1:
        return candidates[0]
    raise FileNotFoundError(
        f"cannot resolve a unique anatomical STL for {case_id}: {candidates}"
    )


def _roi_cache_path(case_id: str) -> Path:
    return OUTPUT / "anatomical_roi" / f"{_safe_case(case_id)}.npz"


def _roi_cache_valid(path: Path, case, stl_path: Path) -> bool:
    if not path.is_file():
        return False
    try:
        with np.load(path, allow_pickle=False) as source:
            return (
                str(source["case_id"]) == case.case_id
                and str(source["stl_sha256"]) == _sha256(stl_path)
                and str(source["case_manifest_sha256"])
                == str(case.manifest["parents"]["bundle"]["sha256"])
                and len(source["anatomical_indices"]) > 0
            )
    except Exception:
        return False


def build_roi_cache(case, bundle: dict[str, Any], force: bool = False) -> Path:
    output = _roi_cache_path(case.case_id)
    stl_path = _resolve_stl(case.case_id).resolve()
    if not force and _roi_cache_valid(output, case, stl_path):
        return output
    strict_indices = np.flatnonzero(~np.asarray(case.is_wall, dtype=bool)).astype(np.int64)
    raw_coords = np.asarray(bundle["interior_coords_raw_mm"], dtype=np.float64)
    if len(raw_coords) != len(case.coords):
        raise ValueError(f"raw/registered volume count mismatch: {case.case_id}")
    surface = CROP._read_stl(stl_path)
    boundary_before = CROP._boundary_edge_count(surface)
    capped = CROP._cap_surface(surface, FILL_HOLE_SIZE_MM)
    masks = []
    strict_coords = raw_coords[strict_indices]
    for start in range(0, len(strict_coords), ROI_CHUNK_SIZE):
        masks.append(
            CROP._inside_mask(
                strict_coords[start : start + ROI_CHUNK_SIZE],
                capped,
                ENCLOSED_POINT_TOLERANCE,
            )
        )
    inside = np.concatenate(masks)
    anatomical_indices = strict_indices[inside]
    if len(anatomical_indices) < 3:
        raise RuntimeError(f"anatomical ROI is empty or degenerate: {case.case_id}")
    BASE._write_npz(
        output,
        {
            "case_id": np.asarray(case.case_id),
            "strict_indices": strict_indices,
            "anatomical_indices": anatomical_indices,
            "stl_path": np.asarray(str(stl_path)),
            "stl_sha256": np.asarray(_sha256(stl_path)),
            "case_manifest_sha256": np.asarray(
                str(case.manifest["parents"]["bundle"]["sha256"])
            ),
            "boundary_edges_before_capping": np.asarray(boundary_before, dtype=np.int64),
            "boundary_edges_after_capping": np.asarray(
                CROP._boundary_edge_count(capped), dtype=np.int64
            ),
            "fill_hole_size_mm": np.asarray(FILL_HOLE_SIZE_MM, dtype=np.float64),
            "enclosed_point_tolerance": np.asarray(
                ENCLOSED_POINT_TOLERANCE, dtype=np.float64
            ),
        },
    )
    print(
        json.dumps(
            {
                "event": "roi_completed",
                "case": case.case_id,
                "strict_points": int(len(strict_indices)),
                "anatomical_points": int(len(anatomical_indices)),
                "removed_extension_points": int(len(strict_indices) - len(anatomical_indices)),
            },
            ensure_ascii=False,
        ),
        flush=True,
    )
    return output


def build_all_roi_caches(
    force: bool = False,
    *,
    shard_index: int = 0,
    num_shards: int = 1,
) -> None:
    if num_shards < 1 or shard_index < 0 or shard_index >= num_shards:
        raise ValueError("ROI shard must satisfy 0 <= shard_index < num_shards")
    arm = BASE.ARM_BY_KEY["V3-E1"]
    config = BASE.ExperimentConfig.from_json(arm.config_path)
    dataset = BASE._load_dataset(config)
    selected = [
        case
        for index, case in enumerate(dataset.cases)
        if index % num_shards == shard_index
    ]
    for index, case in enumerate(selected, start=1):
        bundle = BASE._bundle_arrays(case)
        build_roi_cache(case, bundle, force=force)
        print(
            json.dumps(
                {
                    "event": "roi_progress",
                    "shard_index": shard_index,
                    "num_shards": num_shards,
                    "index": index,
                    "total": len(selected),
                },
                ensure_ascii=False,
            ),
            flush=True,
        )


def _load_roi(case_id: str) -> dict[str, Any]:
    path = _roi_cache_path(case_id)
    if not path.is_file():
        raise FileNotFoundError(f"missing anatomical ROI cache: {path}")
    with np.load(path, allow_pickle=False) as source:
        return {name: np.asarray(source[name]) for name in source.files}


def summarize_roi_caches() -> Path:
    paths = sorted((OUTPUT / "anatomical_roi").glob("*.npz"))
    if len(paths) != 35:
        raise RuntimeError(f"anatomical ROI cache count is {len(paths)}; expected 35")
    rows = []
    for path in paths:
        with np.load(path, allow_pickle=False) as source:
            strict = int(len(source["strict_indices"]))
            anatomical = int(len(source["anatomical_indices"]))
            rows.append(
                {
                    "case_id": str(source["case_id"]),
                    "strict_volume_points": strict,
                    "anatomical_points": anatomical,
                    "removed_extension_points": strict - anatomical,
                    "retained_ratio": anatomical / strict,
                    "anatomical_stl": str(source["stl_path"]),
                    "anatomical_stl_sha256": str(source["stl_sha256"]),
                    "roi_cache": str(path),
                    "roi_cache_sha256": _sha256(path),
                }
            )
    strict_total = sum(row["strict_volume_points"] for row in rows)
    anatomical_total = sum(row["anatomical_points"] for row in rows)
    payload = {
        "schema_version": 1,
        "created_at": BASE.utc_now(),
        "status": "completed",
        "case_count": len(rows),
        "method": {
            "name": "capped anatomical STL enclosed-point crop",
            "fill_hole_size_mm": FILL_HOLE_SIZE_MM,
            "enclosed_point_tolerance": ENCLOSED_POINT_TOLERANCE,
            "input_population": "strict-volume non-wall rows",
            "purpose": "exclude artificial CFD inlet/outlet extensions",
        },
        "totals": {
            "strict_volume_points": strict_total,
            "anatomical_points": anatomical_total,
            "removed_extension_points": strict_total - anatomical_total,
            "retained_ratio": anatomical_total / strict_total,
        },
        "cases": sorted(rows, key=lambda row: row["case_id"]),
    }
    output = OUTPUT / "anatomical_roi" / "manifest.json"
    _write_json(output, payload)
    return output


def _case_dir(arm_key: str, case_id: str) -> Path:
    return OUTPUT / "arms" / arm_key / "cases" / _safe_case(case_id)


def _case_complete(path: Path, arm, case_id: str) -> bool:
    if not path.is_file():
        return False
    try:
        payload = _json(path)
        return (
            payload.get("status") == "completed"
            and payload.get("arm") == arm.key
            and payload.get("case_id") == case_id
            and payload.get("checkpoint_sha256") == _sha256(arm.checkpoint_path)
            and payload.get("speed", {}).get("count", 0) > 0
            and payload.get("wss", {}).get("count", 0) > 1200
        )
    except Exception:
        return False


def analyze_case(
    arm,
    case,
    *,
    config,
    dataset,
    field_stats,
    model,
    device,
    dtype,
    wss_configs,
    query_chunk_size: int,
    force: bool,
) -> Path:
    import torch

    case_dir = _case_dir(arm.key, case.case_id)
    metrics_path = case_dir / "metrics.json"
    if not force and _case_complete(metrics_path, arm, case.case_id):
        return metrics_path
    case_dir.mkdir(parents=True, exist_ok=True)
    started = time.time()
    encoded, support_idx = BASE._encode_case(
        case, dataset, config, model, device, dtype
    )
    predicted_velocity, _ = BASE._predict_coords(
        model,
        encoded,
        np.asarray(case.coords),
        field_stats,
        device,
        dtype,
        query_chunk_size,
    )
    bundle = BASE._bundle_arrays(case)
    roi_path = build_roi_cache(case, bundle, force=False)
    roi = _load_roi(case.case_id)
    indices = np.asarray(roi["anatomical_indices"], dtype=np.int64)
    truth_speed = np.linalg.norm(
        np.asarray(case.velocity[indices], dtype=np.float64), axis=1
    )
    predicted_speed = np.linalg.norm(
        np.asarray(predicted_velocity[indices], dtype=np.float64), axis=1
    )
    speed_fit = linear_fit(truth_speed, predicted_speed)
    speed_hist = density_histogram(truth_speed, predicted_speed)

    work_dir = case_dir / "_wss_work"
    work_dir.mkdir(parents=True, exist_ok=True)
    cache_path = work_dir / f"{_safe_case(case.case_id)}.npz"
    if force or not BASE._cache_valid(cache_path, arm, 0):
        v1_payload, v3_config, v4_config, depth3_config, _ = wss_configs
        arrays = BASE._build_wss_cache_arrays(
            case,
            bundle,
            predicted_velocity,
            arm,
            0,
            v1_payload,
            v3_config,
            v4_config,
            depth3_config,
        )
        BASE._write_npz(cache_path, arrays)
        del arrays
    wss_report, wss_arrays = BASE._apply_calibrator(work_dir)
    wss = wss_arrays[case.case_id]
    truth_wss = np.asarray(wss["truth_mag"], dtype=np.float64)
    predicted_wss = np.asarray(wss["pred_mag"], dtype=np.float64)
    wss_fit = linear_fit(truth_wss, predicted_wss)
    wss_hist = density_histogram(truth_wss, predicted_wss)

    hist_path = case_dir / "density_histograms.npz"
    BASE._write_npz(
        hist_path,
        {
            "speed_counts": speed_hist[0],
            "speed_x_edges": speed_hist[1],
            "speed_y_edges": speed_hist[2],
            "wss_counts": wss_hist[0],
            "wss_x_edges": wss_hist[1],
            "wss_y_edges": wss_hist[2],
        },
    )
    speed_plot = case_dir / "speed_regression.png"
    wss_plot = case_dir / "wss_regression.png"
    _plot_density(
        speed_plot,
        *speed_hist,
        speed_fit,
        case_id=case.case_id,
        arm_key=arm.key,
        metric_label="speed magnitude",
        unit="m/s",
        scope_label="all anatomical interior points; CFD extensions excluded",
    )
    _plot_density(
        wss_plot,
        *wss_hist,
        wss_fit,
        case_id=case.case_id,
        arm_key=arm.key,
        metric_label="WSS magnitude",
        unit="Pa",
        scope_label="all frozen wall nodes; Profile-Secant V3",
    )
    wss_case_row = wss_report["cases"][0]
    payload = {
        "schema_version": 1,
        "created_at": BASE.utc_now(),
        "status": "completed",
        "arm": arm.key,
        "route": arm.route_label,
        "run_id": arm.run_id,
        "case_id": case.case_id,
        "checkpoint": str(arm.checkpoint_path),
        "checkpoint_sha256": _sha256(arm.checkpoint_path),
        "evaluation_protocol": arm.evaluation_label,
        "support_points": int(len(support_idx)),
        "support_rule": "frozen deterministic evaluation support",
        "speed_scope": {
            "definition": "norm([u,v,w]) in physical m/s",
            "population": "all strict-volume points enclosed by capped anatomical STL",
            "strict_volume_points": int(len(roi["strict_indices"])),
            "anatomical_points": int(len(indices)),
            "removed_extension_points": int(len(roi["strict_indices"]) - len(indices)),
            "roi_cache": str(roi_path),
            "roi_cache_sha256": _sha256(roi_path),
            "anatomical_stl": str(roi["stl_path"]),
            "anatomical_stl_sha256": str(roi["stl_sha256"]),
        },
        "wss_scope": {
            "definition": "WSS magnitude in Pa",
            "population": "all frozen wall nodes",
            "algorithm": BASE.ALGORITHM_NAME,
            "wall_points": int(len(truth_wss)),
            "calibrator": str(BASE.CALIBRATOR),
            "calibrator_sha256": _sha256(BASE.CALIBRATOR),
            "raw_r2": float(wss_case_row["wss_r2"]),
            "range_nmae": float(wss_case_row["wss_nmae_range"]),
            "mae_pa": float(wss_case_row["wss_mae_pa"]),
        },
        "speed": speed_fit.as_dict(),
        "wss": wss_fit.as_dict(),
        "artifacts": {
            "histograms": str(hist_path),
            "histograms_sha256": _sha256(hist_path),
            "speed_regression_plot": str(speed_plot),
            "speed_regression_plot_sha256": _sha256(speed_plot),
            "wss_regression_plot": str(wss_plot),
            "wss_regression_plot_sha256": _sha256(wss_plot),
        },
        "elapsed_seconds": time.time() - started,
    }
    _write_json(metrics_path, payload)
    if cache_path.is_file():
        cache_path.unlink()
    try:
        work_dir.rmdir()
    except OSError:
        pass
    del encoded, predicted_velocity, truth_speed, predicted_speed, wss_arrays
    if device.type == "cuda":
        torch.cuda.empty_cache()
    return metrics_path


def analyze_arm(
    arm,
    *,
    device_name: str,
    query_chunk_size: int,
    case_id: str | None,
    case_shard_index: int,
    num_case_shards: int,
    force: bool,
) -> None:
    if (
        num_case_shards < 1
        or case_shard_index < 0
        or case_shard_index >= num_case_shards
    ):
        raise ValueError("case shard must satisfy 0 <= index < count")
    config, dataset, field_stats, model, device, dtype = BASE._load_runtime(
        arm, device_name
    )
    wss_configs = BASE._load_wss_configs()
    cases = [
        case
        for index, case in enumerate(dataset.cases)
        if index % num_case_shards == case_shard_index
    ]
    if case_id is not None:
        cases = [case for case in dataset.cases if case.case_id == case_id]
        if not cases:
            raise KeyError(f"case is not in frozen evaluation split: {case_id}")
    for index, case in enumerate(cases, start=1):
        metrics = analyze_case(
            arm,
            case,
            config=config,
            dataset=dataset,
            field_stats=field_stats,
            model=model,
            device=device,
            dtype=dtype,
            wss_configs=wss_configs,
            query_chunk_size=query_chunk_size,
            force=force,
        )
        row = _json(metrics)
        print(
            json.dumps(
                {
                    "event": "case_completed",
                    "arm": arm.key,
                    "case": case.case_id,
                    "index": index,
                    "total": len(cases),
                    "speed_r2": row["speed"]["regression_r2"],
                    "wss_r2": row["wss"]["regression_r2"],
                    "elapsed_seconds": row["elapsed_seconds"],
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
    if case_id is None and num_case_shards == 1:
        summarize_arm(arm)


def _mean_std(values: list[float]) -> dict[str, float | int]:
    return BASE._mean_std(values)


def _representatives(rows: list[dict[str, Any]], metric: str) -> dict[str, Any]:
    ordered = sorted(
        rows,
        key=lambda row: (float(row[metric]["regression_r2"]), row["case_id"]),
    )
    worst = ordered[0]
    best = ordered[-1]
    candidates = ordered[1:-1] or ordered
    values = np.asarray(
        [float(row[metric]["regression_r2"]) for row in rows], dtype=np.float64
    )
    bin_count = max(3, min(8, int(math.ceil(math.sqrt(len(values))))))
    if float(np.max(values) - np.min(values)) <= 1e-15:
        edges = np.linspace(float(values[0]) - 0.5, float(values[0]) + 0.5, bin_count + 1)
    else:
        edges = np.linspace(float(np.min(values)), float(np.max(values)), bin_count + 1)
    counts, edges = np.histogram(values, bins=edges)
    maximum = int(np.max(counts))
    tied = np.flatnonzero(counts == maximum)
    median = float(np.median(values))
    selected_bin = min(
        tied,
        key=lambda index: (abs(float((edges[index] + edges[index + 1]) / 2.0) - median), int(index)),
    )
    centre = float((edges[selected_bin] + edges[selected_bin + 1]) / 2.0)
    frequent = min(
        candidates,
        key=lambda row: (abs(float(row[metric]["regression_r2"]) - centre), row["case_id"]),
    )
    return {
        "worst": worst["case_id"],
        "most_frequent": frequent["case_id"],
        "best": best["case_id"],
        "histogram": {
            "bin_count": bin_count,
            "edges": edges.tolist(),
            "counts": counts.tolist(),
            "selected_bin": int(selected_bin),
            "selected_bin_center": centre,
            "rule": (
                "equal-width sqrt(n) bins over observed R2; select the bin with "
                "maximum count, break ties toward the median, then choose the "
                "non-extreme case nearest the bin centre"
            ),
        },
    }


def _plot_ranking(
    output: Path,
    rows: list[dict[str, Any]],
    metric: str,
    representatives: dict[str, Any],
    arm_key: str,
) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    ordered = sorted(
        rows,
        key=lambda row: (float(row[metric]["regression_r2"]), row["case_id"]),
    )
    labels = [row["case_id"] for row in ordered]
    values = np.asarray(
        [float(row[metric]["regression_r2"]) for row in ordered], dtype=np.float64
    )
    height = max(8.5, 0.32 * len(rows) + 2.0)
    figure, axis = plt.subplots(figsize=(9.2, height), constrained_layout=True)
    colors = []
    for case in labels:
        if case == representatives["worst"]:
            colors.append("#C43C39")
        elif case == representatives["best"]:
            colors.append("#2E7D32")
        elif case == representatives["most_frequent"]:
            colors.append("#E69F00")
        else:
            colors.append("#3B78A8")
    positions = np.arange(len(labels))
    axis.hlines(positions, np.min(values), values, color="0.85", linewidth=0.8)
    axis.scatter(values, positions, c=colors, s=35, zorder=3)
    axis.set_yticks(positions)
    axis.set_yticklabels(labels, fontsize=7.5)
    axis.set_xlabel("Linear-fit R²")
    axis.set_ylabel("Case number / ID (sorted worst → best)")
    axis.set_title(f"{arm_key} | {metric.upper()} case-wise linear-fit R²")
    axis.grid(axis="x", color="0.9", linewidth=0.8)
    axis.set_ylim(-0.8, len(labels) - 0.2)
    output.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output, dpi=220)
    plt.close(figure)


def _plot_distribution(
    output: Path,
    rows: list[dict[str, Any]],
    metric: str,
    representatives: dict[str, Any],
    arm_key: str,
) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    values = np.asarray(
        [float(row[metric]["regression_r2"]) for row in rows], dtype=np.float64
    )
    histogram = representatives["histogram"]
    figure, axis = plt.subplots(figsize=(7.2, 4.7), constrained_layout=True)
    axis.hist(values, bins=np.asarray(histogram["edges"]), color="#72A9CF", edgecolor="white")
    by_case = {row["case_id"]: row for row in rows}
    styles = {
        "worst": ("#C43C39", "worst"),
        "most_frequent": ("#E69F00", "most frequent"),
        "best": ("#2E7D32", "best"),
    }
    for key, (color, label) in styles.items():
        case_id = representatives[key]
        value = float(by_case[case_id][metric]["regression_r2"])
        axis.axvline(value, color=color, linewidth=1.8, label=f"{label}: {case_id}")
    axis.set_xlabel("Linear-fit R²")
    axis.set_ylabel("Case count")
    axis.set_title(f"{arm_key} | {metric.upper()} R² distribution")
    axis.legend(fontsize=7.5)
    output.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output, dpi=220)
    plt.close(figure)


def _plot_representatives(
    output: Path,
    arm_key: str,
    metric: str,
    representatives: dict[str, Any],
) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    labels = ("worst", "most_frequent", "best")
    figure, axes = plt.subplots(1, 3, figsize=(18, 6.2), constrained_layout=True)
    filename = "speed_regression.png" if metric == "speed" else "wss_regression.png"
    for axis, label in zip(axes, labels):
        case_id = representatives[label]
        image_path = _case_dir(arm_key, case_id) / filename
        axis.imshow(plt.imread(image_path))
        axis.axis("off")
        axis.set_title(label.replace("_", " ").title(), fontsize=12, fontweight="bold")
    figure.suptitle(f"{arm_key} | {metric.upper()} representative regressions", fontsize=15)
    output.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output, dpi=180)
    plt.close(figure)


def summarize_arm(arm) -> Path:
    metrics_paths = sorted((OUTPUT / "arms" / arm.key / "cases").glob("*/metrics.json"))
    rows = [_json(path) for path in metrics_paths if _case_complete(path, arm, _json(path)["case_id"])]
    if len(rows) != 35:
        raise RuntimeError(f"{arm.key} has {len(rows)} completed cases; expected 35")
    aggregates: dict[str, Any] = {}
    representatives: dict[str, Any] = {}
    plots: dict[str, Any] = {}
    for metric in ("speed", "wss"):
        aggregates[metric] = {
            "slope_a": _mean_std([float(row[metric]["slope_a"]) for row in rows]),
            "intercept_b": _mean_std([float(row[metric]["intercept_b"]) for row in rows]),
            "regression_r2": _mean_std([float(row[metric]["regression_r2"]) for row in rows]),
            "points": {
                "total": int(sum(int(row[metric]["count"]) for row in rows)),
                "minimum_per_case": int(min(int(row[metric]["count"]) for row in rows)),
                "maximum_per_case": int(max(int(row[metric]["count"]) for row in rows)),
            },
        }
        representatives[metric] = _representatives(rows, metric)
        plot_dir = OUTPUT / "arms" / arm.key / "summary_plots"
        ranking = plot_dir / f"{metric}_case_r2_ranking.png"
        distribution = plot_dir / f"{metric}_r2_distribution.png"
        representative = plot_dir / f"{metric}_representative_regressions.png"
        _plot_ranking(ranking, rows, metric, representatives[metric], arm.key)
        _plot_distribution(distribution, rows, metric, representatives[metric], arm.key)
        _plot_representatives(representative, arm.key, metric, representatives[metric])
        plots[metric] = {
            "case_r2_ranking": str(ranking),
            "r2_distribution": str(distribution),
            "representative_regressions": str(representative),
        }
    summary = {
        "schema_version": 1,
        "created_at": BASE.utc_now(),
        "status": "completed",
        "arm": arm.key,
        "route": arm.route_label,
        "run_id": arm.run_id,
        "checkpoint": str(arm.checkpoint_path),
        "checkpoint_sha256": _sha256(arm.checkpoint_path),
        "evaluation_protocol": arm.evaluation_label,
        "case_count": len(rows),
        "definitions": {
            "equation": "y = a*x + b",
            "x": "CFD scalar truth",
            "y": "deep-learning scalar prediction or velocity-derived WSS",
            "regression_r2": "coefficient of determination of the fitted OLS line with free intercept",
            "speed_population": "all anatomical interior points after capped-STL extension removal",
            "wss_population": "all frozen wall nodes",
        },
        "aggregate": aggregates,
        "representatives": representatives,
        "plots": plots,
        "cases": [str(path) for path in metrics_paths],
    }
    output = OUTPUT / "arms" / arm.key / "summary.json"
    _write_json(output, summary)
    csv_path = OUTPUT / "arms" / arm.key / "case_metrics.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "case_id",
                "speed_points",
                "speed_slope_a",
                "speed_intercept_b_m_s",
                "speed_regression_r2",
                "wss_points",
                "wss_slope_a",
                "wss_intercept_b_pa",
                "wss_regression_r2",
            ]
        )
        for row in sorted(rows, key=lambda item: item["case_id"]):
            writer.writerow(
                [
                    row["case_id"],
                    row["speed"]["count"],
                    row["speed"]["slope_a"],
                    row["speed"]["intercept_b"],
                    row["speed"]["regression_r2"],
                    row["wss"]["count"],
                    row["wss"]["slope_a"],
                    row["wss"]["intercept_b"],
                    row["wss"]["regression_r2"],
                ]
            )
    return output


def summarize_all() -> None:
    summarize_roi_caches()
    for arm in BASE.ARMS:
        summarize_arm(arm)


def _format_mean_std(summary: dict[str, Any]) -> str:
    return f"{float(summary['mean']):.4f} ± {float(summary['std_population']):.4f}"


def update_workbook() -> Path:
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    summaries = {arm.key: _json(OUTPUT / "arms" / arm.key / "summary.json") for arm in BASE.ARMS}
    if not WORKBOOK_BACKUP.is_file():
        shutil.copy2(WORKBOOK, WORKBOOK_BACKUP)
    workbook = load_workbook(WORKBOOK)
    sheet = workbook["实验矩阵汇总"]
    headers = {
        str(sheet.cell(4, column).value): column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(4, column).value is not None
    }
    new_headers = [
        "speed fit a_cb (mean±std)",
        "speed fit b_cb (m/s, mean±std)",
        "speed fit R²_cb (mean±std)",
        "WSS fit a_cb (mean±std)",
        "WSS fit b_cb (Pa, mean±std)",
        "WSS fit R²_cb (mean±std)",
    ]
    for header in new_headers:
        if header not in headers:
            column = sheet.max_column + 1
            headers[header] = column
            source = sheet.cell(4, column - 1)
            target = sheet.cell(4, column, header)
            target._style = copy(source._style)
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            sheet.column_dimensions[get_column_letter(column)].width = 24
    comments = {
        "speed fit a_cb (mean±std)": "逐病例用去延长段后全部解剖体内点拟合 y=a*x+b，a 的 mean ± population std。",
        "speed fit b_cb (m/s, mean±std)": "同上带截距 OLS 拟合的 b，单位 m/s。",
        "speed fit R²_cb (mean±std)": "拟合直线 y=a*x+b 对预测速度标量的决定系数；不替代原 speed raw R²。",
        "WSS fit a_cb (mean±std)": "逐病例使用全部冻结壁面节点拟合 y=a*x+b，a 的 mean ± population std。",
        "WSS fit b_cb (Pa, mean±std)": "同上带截距 OLS 拟合的 b，单位 Pa。",
        "WSS fit R²_cb (mean±std)": "全壁面 Profile-Secant V3 WSS 散点的带截距线性拟合 R²；不替代原 1200 点 raw/scaled R²。",
    }
    for header, text in comments.items():
        sheet.cell(4, headers[header]).comment = Comment(text, "Codex")
    experiment_rows = {
        str(sheet.cell(row, 4).value): row
        for row in range(5, sheet.max_row + 1)
        if sheet.cell(row, 4).value
    }
    for arm in BASE.ARMS:
        row = experiment_rows[arm.run_id]
        summary = summaries[arm.key]["aggregate"]
        values = {
            new_headers[0]: _format_mean_std(summary["speed"]["slope_a"]),
            new_headers[1]: _format_mean_std(summary["speed"]["intercept_b"]),
            new_headers[2]: _format_mean_std(summary["speed"]["regression_r2"]),
            new_headers[3]: _format_mean_std(summary["wss"]["slope_a"]),
            new_headers[4]: _format_mean_std(summary["wss"]["intercept_b"]),
            new_headers[5]: _format_mean_std(summary["wss"]["regression_r2"]),
        }
        for header, value in values.items():
            column = headers[header]
            source = sheet.cell(row, column - 1)
            target = sheet.cell(row, column, value)
            target._style = copy(source._style)
            target.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_fill = copy(sheet["A1"].fill)
    note_fill = copy(sheet["A3"].fill)
    for merged in list(sheet.merged_cells.ranges):
        if (merged.min_row, merged.max_row) in {(1, 2), (3, 3)}:
            sheet.unmerge_cells(str(merged))
    for column in range(1, sheet.max_column + 1):
        sheet.cell(1, column).fill = copy(title_fill)
        sheet.cell(2, column).fill = copy(title_fill)
        sheet.cell(3, column).fill = copy(note_fill)
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=sheet.max_column)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=sheet.max_column)
    note = str(sheet["A3"].value or "")
    addition = (
        " 新增 linear-fit 指标为带截距 OLS y=a·x+b：速度使用去除 CFD 延长段后"
        "全部解剖体内点，WSS 使用全部冻结壁面节点；旧指标保留。"
    )
    if "linear-fit" not in note:
        sheet["A3"] = note + addition
        sheet.row_dimensions[3].height = max(float(sheet.row_dimensions[3].height or 54), 72)
    last_letter = get_column_letter(sheet.max_column)
    sheet.auto_filter.ref = f"A4:{last_letter}{4 + len(BASE.ARMS)}"
    sheet.print_area = f"A1:{last_letter}{4 + len(BASE.ARMS)}"

    for name in ("线性回归汇总", "线性回归逐病例"):
        if name in workbook.sheetnames:
            del workbook[name]
    summary_sheet = workbook.create_sheet("线性回归汇总")
    summary_headers = [
        "路线",
        "臂",
        "实验ID",
        "checkpoint",
        "病例数",
        "speed a",
        "speed b (m/s)",
        "speed fit R²",
        "speed worst",
        "speed most frequent",
        "speed best",
        "speed ranking plot",
        "speed distribution plot",
        "speed representative plot",
        "WSS a",
        "WSS b (Pa)",
        "WSS fit R²",
        "WSS worst",
        "WSS most frequent",
        "WSS best",
        "WSS ranking plot",
        "WSS distribution plot",
        "WSS representative plot",
    ]
    summary_sheet.append(summary_headers)
    for arm in BASE.ARMS:
        summary = summaries[arm.key]
        aggregate = summary["aggregate"]
        representative = summary["representatives"]
        plots = summary["plots"]
        summary_sheet.append(
            [
                arm.route_label,
                arm.arm,
                arm.run_id,
                arm.checkpoint,
                summary["case_count"],
                _format_mean_std(aggregate["speed"]["slope_a"]),
                _format_mean_std(aggregate["speed"]["intercept_b"]),
                _format_mean_std(aggregate["speed"]["regression_r2"]),
                representative["speed"]["worst"],
                representative["speed"]["most_frequent"],
                representative["speed"]["best"],
                plots["speed"]["case_r2_ranking"],
                plots["speed"]["r2_distribution"],
                plots["speed"]["representative_regressions"],
                _format_mean_std(aggregate["wss"]["slope_a"]),
                _format_mean_std(aggregate["wss"]["intercept_b"]),
                _format_mean_std(aggregate["wss"]["regression_r2"]),
                representative["wss"]["worst"],
                representative["wss"]["most_frequent"],
                representative["wss"]["best"],
                plots["wss"]["case_r2_ranking"],
                plots["wss"]["r2_distribution"],
                plots["wss"]["representative_regressions"],
            ]
        )

    case_sheet = workbook.create_sheet("线性回归逐病例")
    case_headers = [
        "路线",
        "臂",
        "实验ID",
        "指标",
        "R²排名(worst=1)",
        "case ID",
        "点数",
        "slope a",
        "intercept b",
        "linear-fit R²",
        "原始 prediction R²",
        "代表类型",
        "数据范围",
        "regression plot",
        "metrics JSON",
    ]
    case_sheet.append(case_headers)
    for arm in BASE.ARMS:
        summary = summaries[arm.key]
        rows = [_json(Path(path)) for path in summary["cases"]]
        for metric in ("speed", "wss"):
            ordered = sorted(
                rows,
                key=lambda row: (float(row[metric]["regression_r2"]), row["case_id"]),
            )
            representative = summary["representatives"][metric]
            representative_by_case = {
                representative["worst"]: "worst",
                representative["most_frequent"]: "most frequent",
                representative["best"]: "best",
            }
            for rank, row in enumerate(ordered, start=1):
                plot_key = f"{metric}_regression_plot"
                scope = (
                    "all anatomical interior points; extensions excluded"
                    if metric == "speed"
                    else "all frozen wall nodes"
                )
                metrics_path = _case_dir(arm.key, row["case_id"]) / "metrics.json"
                case_sheet.append(
                    [
                        arm.route_label,
                        arm.arm,
                        arm.run_id,
                        metric,
                        rank,
                        row["case_id"],
                        row[metric]["count"],
                        row[metric]["slope_a"],
                        row[metric]["intercept_b"],
                        row[metric]["regression_r2"],
                        row[metric]["raw_prediction_r2"],
                        representative_by_case.get(row["case_id"], ""),
                        scope,
                        row["artifacts"][plot_key],
                        str(metrics_path),
                    ]
                )

    header_fill = PatternFill("solid", fgColor="2F75B5")
    for target_sheet in (summary_sheet, case_sheet):
        for cell in target_sheet[1]:
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        target_sheet.freeze_panes = "A2"
        target_sheet.auto_filter.ref = target_sheet.dimensions
        target_sheet.sheet_view.showGridLines = False
        for column in range(1, target_sheet.max_column + 1):
            maximum = max(
                len(str(target_sheet.cell(row, column).value or ""))
                for row in range(1, min(target_sheet.max_row, 80) + 1)
            )
            target_sheet.column_dimensions[get_column_letter(column)].width = min(max(maximum + 2, 10), 42)
    temporary = WORKBOOK.with_name(f".{WORKBOOK.name}.{os.getpid()}.tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, WORKBOOK)
    return WORKBOOK


def verify() -> dict[str, Any]:
    from openpyxl import load_workbook

    failures: list[str] = []
    try:
        roi_manifest = summarize_roi_caches()
    except Exception as error:
        roi_manifest = None
        failures.append(f"ROI manifest failed: {error}")
    arm_checks = []
    total_case_metric_rows = 0
    for arm in BASE.ARMS:
        summary_path = OUTPUT / "arms" / arm.key / "summary.json"
        if not summary_path.is_file():
            failures.append(f"missing summary: {arm.key}")
            continue
        summary = _json(summary_path)
        if summary.get("case_count") != 35:
            failures.append(f"case count is not 35: {arm.key}")
        rows = [_json(Path(path)) for path in summary.get("cases", [])]
        if len(rows) != 35:
            failures.append(f"metrics JSON count is not 35: {arm.key}")
        for row in rows:
            total_case_metric_rows += 2
            for metric in ("speed", "wss"):
                value = float(row[metric]["regression_r2"])
                if not np.isfinite(value) or value < -1e-8 or value > 1.0 + 1e-8:
                    failures.append(
                        f"invalid fitted-line R2: {arm.key} {row['case_id']} {metric}={value}"
                    )
                if int(row[metric]["count"]) < 3:
                    failures.append(f"too few points: {arm.key} {row['case_id']} {metric}")
            if int(row["wss"]["count"]) <= 1200:
                failures.append(f"WSS is not full-wall: {arm.key} {row['case_id']}")
            for key, path in row["artifacts"].items():
                if key.endswith("_sha256"):
                    continue
                if not Path(path).is_file():
                    failures.append(f"missing case artifact: {path}")
        for metric in ("speed", "wss"):
            for path in summary["plots"][metric].values():
                if not Path(path).is_file():
                    failures.append(f"missing summary plot: {path}")
        arm_checks.append(
            {
                "arm": arm.key,
                "summary": str(summary_path),
                "summary_sha256": _sha256(summary_path),
                "speed": summary.get("aggregate", {}).get("speed"),
                "wss": summary.get("aggregate", {}).get("wss"),
            }
        )
    if not WORKBOOK.is_file():
        failures.append("workbook missing")
    else:
        workbook = load_workbook(WORKBOOK, read_only=True, data_only=False)
        required_sheets = {"实验矩阵汇总", "线性回归汇总", "线性回归逐病例"}
        missing = sorted(required_sheets - set(workbook.sheetnames))
        if missing:
            failures.append(f"workbook sheets missing: {missing}")
        else:
            main = workbook["实验矩阵汇总"]
            headers = [main.cell(4, column).value for column in range(1, main.max_column + 1)]
            for required in (
                "speed fit a_cb (mean±std)",
                "speed fit b_cb (m/s, mean±std)",
                "speed fit R²_cb (mean±std)",
                "WSS fit a_cb (mean±std)",
                "WSS fit b_cb (Pa, mean±std)",
                "WSS fit R²_cb (mean±std)",
            ):
                if required not in headers:
                    failures.append(f"workbook header missing: {required}")
            if workbook["线性回归汇总"].max_row != 15:
                failures.append("regression summary sheet does not have 14 data rows")
            if workbook["线性回归逐病例"].max_row != 981:
                failures.append("case regression sheet does not have 980 data rows")
    payload = {
        "schema_version": 1,
        "created_at": BASE.utc_now(),
        "status": "passed" if not failures else "failed",
        "failures": failures,
        "checks": {
            "arms": arm_checks,
            "expected_arms": len(BASE.ARMS),
            "case_metric_rows": total_case_metric_rows,
            "workbook": str(WORKBOOK),
            "workbook_sha256": _sha256(WORKBOOK) if WORKBOOK.is_file() else None,
            "workbook_backup": str(WORKBOOK_BACKUP),
            "workbook_backup_sha256": _sha256(WORKBOOK_BACKUP) if WORKBOOK_BACKUP.is_file() else None,
            "anatomical_roi_manifest": str(roi_manifest) if roi_manifest else None,
            "anatomical_roi_manifest_sha256": _sha256(roi_manifest) if roi_manifest else None,
        },
    }
    output = OUTPUT / "verification.json"
    _write_json(output, payload)
    print(json.dumps(payload, indent=2, ensure_ascii=False))
    if failures:
        raise SystemExit(1)
    return payload


def self_test() -> None:
    x = np.linspace(0.0, 3.0, 1001)
    y = 0.75 * x + 0.2
    fit = linear_fit(x, y)
    assert abs(fit.slope_a - 0.75) < 1e-12
    assert abs(fit.intercept_b - 0.2) < 1e-12
    assert abs(fit.regression_r2 - 1.0) < 1e-12
    counts, _, _ = density_histogram(x, y, bins=32)
    assert int(np.sum(counts)) == len(x)
    print(json.dumps({"status": "passed", "fit": fit.as_dict()}, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    roi_parser = subparsers.add_parser("roi")
    roi_parser.add_argument("--force", action="store_true")
    roi_parser.add_argument("--shard-index", type=int, default=0)
    roi_parser.add_argument("--num-shards", type=int, default=1)
    subparsers.add_parser("roi-summary")

    arm_parser = subparsers.add_parser("arm")
    arm_parser.add_argument("--arm", choices=sorted(BASE.ARM_BY_KEY), required=True)
    arm_parser.add_argument("--device", default="cuda:0")
    arm_parser.add_argument("--query-chunk-size", type=int, default=16384)
    arm_parser.add_argument("--case-id")
    arm_parser.add_argument("--case-shard-index", type=int, default=0)
    arm_parser.add_argument("--num-case-shards", type=int, default=1)
    arm_parser.add_argument("--force", action="store_true")

    summary_parser = subparsers.add_parser("summarize")
    summary_parser.add_argument("--arm", choices=sorted(BASE.ARM_BY_KEY))
    subparsers.add_parser("workbook")
    subparsers.add_parser("verify")
    subparsers.add_parser("self-test")
    args = parser.parse_args()
    if args.command == "roi":
        build_all_roi_caches(
            force=args.force,
            shard_index=args.shard_index,
            num_shards=args.num_shards,
        )
    elif args.command == "roi-summary":
        print(summarize_roi_caches())
    elif args.command == "arm":
        analyze_arm(
            BASE.ARM_BY_KEY[args.arm],
            device_name=args.device,
            query_chunk_size=args.query_chunk_size,
            case_id=args.case_id,
            case_shard_index=args.case_shard_index,
            num_case_shards=args.num_case_shards,
            force=args.force,
        )
    elif args.command == "summarize":
        if args.arm:
            summarize_arm(BASE.ARM_BY_KEY[args.arm])
        else:
            summarize_all()
    elif args.command == "workbook":
        print(update_workbook())
    elif args.command == "verify":
        verify()
    elif args.command == "self-test":
        self_test()


if __name__ == "__main__":
    main()
