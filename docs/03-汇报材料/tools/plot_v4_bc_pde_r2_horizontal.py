#!/usr/bin/env python3
"""Create report-ready horizontal per-case speed-R² scatter plots for V4.

The eight selected arms are the seed1234 ``BC+PDE-FIXED`` and
``BC+PDE-EMA`` rows in the current workbook: steady/transient ×
PointNet/PointNet++ × two PDE controllers.

Two metrics are written deliberately:

* ``speed_r2`` is the official per-case prediction R² used by the workbook;
* ``speed_linear_fit_r2`` is the Pearson/OLS fitted-line R² used by the older
  ``v2_v3_linear_regression_fullpoints_20260807`` figures.  V4 official JSON
  stores sufficient population moments to reconstruct it without reopening
  the large point clouds.

Each figure uses the requested horizontal layout: x = R², y = case rank
(worst → best), with the same worst/most-frequent/best vertical markers as the
older visual style.  The script does not modify the workbook or training
outputs.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[3]
XLSX = ROOT / "docs/03-汇报材料/WSS_PINN_V1_V2_V3_field-v4实验矩阵与指标汇总_2026-08-06.xlsx"
V4_BASE = ROOT / "outputs/wss_pinn/volume_uvwp_bc_rcr_v4"
DEFAULT_OUT = ROOT / "docs/03-汇报材料/V4_BC-PDE_FIX_EMA_横向R2散点图_2026-09-01"

TARGET_MODES = {"BC+PDE-FIXED", "BC+PDE-EMA"}
TIME_DIR = {"steady peak": "steady_peak", "transient 81": "transient_autograd"}

MODE_SHORT = {
    "BC+PDE-FIXED": "BC-PDE-FIXED",
    "BC+PDE-EMA": "BC-PDE-EMA",
}
TIME_SHORT = {"steady peak": "SP", "transient 81": "TR"}
BACKBONE_SHORT = {"PointNet": "PN", "PointNet++": "PNPP"}

COLORS = {
    "line": "#6AAED6",
    "point": "#3B78A8",
    "worst": "#D62F2F",
    "most_frequent": "#E69F00",
    "best": "#2E9D47",
    "grid": "#D9DEE5",
}


def _finite(value: Any) -> bool:
    try:
        return math.isfinite(float(value))
    except (TypeError, ValueError):
        return False


def _safe_name(value: str) -> str:
    return value.replace("/", "__").replace(" ", "_").replace("+", "plus")


def _read_workbook_targets() -> list[dict[str, Any]]:
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["实验矩阵汇总"]
    headers = [cell.value for cell in ws[4]]
    index = {name: i for i, name in enumerate(headers) if name}
    required = {"路线", "阶段", "编号/臂", "实验ID", "架构/解码器", "训练模式/旋钮"}
    missing = required - index.keys()
    if missing:
        raise RuntimeError(f"workbook headers missing: {sorted(missing)}")

    targets: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[index["路线"]] != "V4":
            continue
        mode = row[index["训练模式/旋钮"]]
        if mode not in TARGET_MODES:
            continue
        stage = str(row[index["阶段"]])
        if "steady peak" in stage:
            temporal = "steady peak"
        elif "transient 81" in stage:
            temporal = "transient 81"
        else:
            raise RuntimeError(f"unrecognized V4 temporal stage: {stage}")
        targets.append(
            {
                "index": int(row[index["编号/臂"]]),
                "experiment_id": str(row[index["实验ID"]]),
                "backbone": str(row[index["架构/解码器"]]),
                "mode": str(mode),
                "temporal": temporal,
                "stage": stage,
            }
        )
    targets.sort(key=lambda x: x["index"])
    if len(targets) != 8:
        raise RuntimeError(f"expected 8 V4 target arms, found {len(targets)}")
    return targets


def _linear_fit_r2(speed: dict[str, Any]) -> float:
    """Reconstruct Pearson r² from the official population moments."""
    vx = float(speed["truth_variance"])
    vy = float(speed["prediction_variance"])
    mx = float(speed["truth_mean"])
    my = float(speed["prediction_mean"])
    raw_r2 = float(speed["r2"])
    if not (_finite(vx) and _finite(vy) and _finite(mx) and _finite(my) and _finite(raw_r2)):
        return float("nan")
    if vx <= 0.0 or vy <= 0.0:
        return float("nan")
    mse = (1.0 - raw_r2) * vx
    covariance = (vx + vy + (my - mx) ** 2 - mse) / 2.0
    value = (covariance * covariance) / (vx * vy)
    # Floating-point roundoff can put a perfect fit infinitesimally above 1.
    return float(np.clip(value, 0.0, 1.0))


def _load_cases(target: dict[str, Any]) -> list[dict[str, Any]]:
    run = target["experiment_id"]
    path = V4_BASE / TIME_DIR[target["temporal"]] / run / "evaluation_official_last_converged_full.json"
    if not path.is_file():
        raise FileNotFoundError(path)
    payload = json.loads(path.read_text(encoding="utf-8"))
    cases = []
    for case in payload.get("cases", []):
        speed = case["metrics"]["speed"]
        cases.append(
            {
                "case_id": str(case["case_id"]),
                "speed_r2": float(speed["r2"]),
                "speed_linear_fit_r2": _linear_fit_r2(speed),
            }
        )
    if len(cases) != 35:
        raise RuntimeError(f"{run}: expected 35 cases, found {len(cases)}")
    return cases


def _representatives(cases: list[dict[str, Any]], metric: str) -> dict[str, str]:
    ordered = sorted(cases, key=lambda x: (x[metric], x["case_id"]))
    worst = ordered[0]
    best = ordered[-1]
    candidates = ordered[1:-1] if len(ordered) > 2 else ordered
    values = np.asarray([float(x[metric]) for x in ordered], dtype=float)
    bins = max(2, int(math.sqrt(len(values))))
    edges = np.linspace(float(np.min(values)), float(np.max(values)), bins + 1)
    if np.allclose(edges[0], edges[-1]):
        return {"worst": worst["case_id"], "most_frequent": candidates[0]["case_id"], "best": best["case_id"]}
    counts, edges = np.histogram(values, bins=edges)
    tied = np.flatnonzero(counts == int(np.max(counts)))
    median = float(np.median(values))
    selected = min(tied, key=lambda i: (abs((edges[i] + edges[i + 1]) / 2.0 - median), int(i)))
    centre = float((edges[selected] + edges[selected + 1]) / 2.0)
    frequent = min(candidates, key=lambda x: (abs(float(x[metric]) - centre), x["case_id"]))
    return {"worst": worst["case_id"], "most_frequent": frequent["case_id"], "best": best["case_id"]}


def _plot_horizontal(
    path: Path,
    target: dict[str, Any],
    cases: list[dict[str, Any]],
    metric: str,
    metric_label: str,
    report_title: str,
) -> None:
    ordered = sorted(cases, key=lambda x: (x[metric], x["case_id"]))
    values = np.asarray([x[metric] for x in ordered], dtype=float)
    positions = np.arange(1, len(ordered) + 1)
    reps = _representatives(cases, metric)
    by_id = {x["case_id"]: x for x in cases}

    fig, ax = plt.subplots(figsize=(10.2, 7.0), constrained_layout=True)
    ax.plot(values, positions, color=COLORS["line"], linewidth=2.2, marker="o", markersize=6, zorder=2)
    # Keep the individual markers crisp after the connecting line is drawn.
    ax.scatter(values, positions, s=42, color=COLORS["point"], zorder=3)
    for key in ("worst", "most_frequent", "best"):
        case_id = reps[key]
        x = float(by_id[case_id][metric])
        color = COLORS[key]
        label = f"{key.replace('_', ' ')}: {case_id}"
        ax.axvline(x, color=color, linewidth=2.0, label=label, zorder=1)
        rank = next(i for i, item in enumerate(ordered, 1) if item["case_id"] == case_id)
        ax.scatter([x], [rank], s=74, color=color, edgecolor="white", linewidth=0.8, zorder=4)

    ax.set_xlabel(metric_label, fontsize=12)
    ax.set_ylabel("Case rank (worst → best)", fontsize=12)
    ax.set_title(report_title, fontsize=17, pad=14)
    ax.set_yticks([1, 5, 10, 15, 20, 25, 30, 35])
    ax.set_ylim(0.5, 35.5)
    ax.grid(axis="x", color=COLORS["grid"], linewidth=0.9)
    ax.grid(axis="y", color="#EEF1F4", linewidth=0.6)
    ax.legend(loc="upper left", fontsize=8.2, frameon=True, framealpha=0.92)
    ax.tick_params(labelsize=10)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=220, facecolor="white")
    plt.close(fig)


def _plot_overview(
    path: Path,
    records: list[dict[str, Any]],
    metric: str,
    metric_label: str,
    title: str,
) -> None:
    fig, axes = plt.subplots(4, 2, figsize=(15.5, 21.0), constrained_layout=False)
    axes = np.asarray(axes).reshape(-1)
    for ax, record in zip(axes, records):
        ordered = sorted(record["cases"], key=lambda x: (x[metric], x["case_id"]))
        values = np.asarray([x[metric] for x in ordered], dtype=float)
        positions = np.arange(1, len(ordered) + 1)
        reps = _representatives(record["cases"], metric)
        by_id = {x["case_id"]: x for x in record["cases"]}
        ax.plot(values, positions, color=COLORS["line"], linewidth=1.3, marker="o", markersize=3.3, zorder=2)
        ax.scatter(values, positions, color=COLORS["point"], s=13, zorder=3)
        for key in ("worst", "most_frequent", "best"):
            x = float(by_id[reps[key]][metric])
            color = COLORS[key]
            ax.axvline(x, color=color, linewidth=1.1, zorder=1)
            rank = next(i for i, item in enumerate(ordered, 1) if item["case_id"] == reps[key])
            ax.scatter([x], [rank], s=28, color=color, edgecolor="white", linewidth=0.35, zorder=4)
        ax.set_title(record["panel_title"], fontsize=11, loc="left", pad=5)
        ax.set_ylim(0.5, 35.5)
        ax.set_yticks([1, 10, 20, 30, 35])
        ax.tick_params(labelsize=8)
        ax.grid(axis="x", color=COLORS["grid"], linewidth=0.6)
        ax.grid(axis="y", color="#EEF1F4", linewidth=0.5)
        for spine in ("top", "right"):
            ax.spines[spine].set_visible(False)
    for ax in axes[::2]:
        ax.set_ylabel("Case rank", fontsize=9)
    for ax in axes[-2:]:
        ax.set_xlabel(metric_label, fontsize=9)
    fig.subplots_adjust(left=0.07, right=0.985, bottom=0.045, top=0.965, hspace=0.30, wspace=0.13)
    fig.suptitle(title, fontsize=17, x=0.01, y=0.992, ha="left")
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=180, facecolor="white")
    plt.close(fig)


def _write_pdf(path: Path, images: list[Path]) -> None:
    pages = []
    try:
        for image_path in images:
            pages.append(Image.open(image_path).convert("RGB"))
        if not pages:
            raise RuntimeError("no images supplied for PDF")
        path.parent.mkdir(parents=True, exist_ok=True)
        pages[0].save(path, save_all=True, append_images=pages[1:], resolution=150.0)
    finally:
        for page in pages:
            page.close()


def _write_csv(path: Path, records: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["编号", "实验ID", "阶段", "架构", "训练模式", "case rank", "case ID", "speed R²", "linear-fit R²"])
        for record in records:
            ordered = sorted(record["cases"], key=lambda x: (x["speed_r2"], x["case_id"]))
            for rank, case in enumerate(ordered, 1):
                writer.writerow(
                    [
                        record["index"],
                        record["experiment_id"],
                        record["stage"],
                        record["backbone"],
                        record["mode"],
                        rank,
                        case["case_id"],
                        f'{case["speed_r2"]:.12g}',
                        f'{case["speed_linear_fit_r2"]:.12g}',
                    ]
                )


def _write_readme(path: Path, records: list[dict[str, Any]]) -> None:
    rows = []
    for r in records:
        rows.append(f"| {r['index']} | `{r['experiment_id']}` | {r['stage']} | {r['backbone']} | {r['mode']} |")
    text = (
        "# V4 BC+PDE-FIXED / BC+PDE-EMA 横向 speed R² 图\n\n"
        "生成日期：2026-09-01。数据来自最新工作簿主表的 V4 seed1234 8 个目标臂，以及对应臂的 "
        "`evaluation_official_last_converged_full.json`。\n\n"
        "每张单臂图的横轴是逐病例 speed R²，纵轴是按 speed R² 从 worst 到 best 排序的病例序号；"
        "蓝线连接排序后的散点，红/橙/绿线分别标出 worst / most frequent / best。\n\n"
        "8 个实验各自使用一个独立子文件夹；根目录只保留汇总图、PDF、CSV 和本说明。\n\n"
        "另外输出的 `*_speed_linear_fit_r2_horizontal.png` 与此前 "
        "`v2_v3_linear_regression_fullpoints_20260807` 图的 `Linear-fit R²` 口径一致。V4 JSON "
        "没有保存点级协方差，因此该值由 truth/prediction 的总体均值、方差和官方 raw R² 精确重构；"
        "它不替代工作簿主表的 raw speed R²。\n\n"
        "每个实验子文件夹还包含 `speed_representative_regressions.png` 和 "
        "`wss_representative_regressions.png` 三联图；其原始三张代表病例图保存在该子文件夹的 "
        "`representative_cases/` 下。\n\n"
        "`V4_BC-PDE_FIXED_EMA_speed_r2_horizontal_report.pdf` 和 "
        "`V4_BC-PDE_FIXED_EMA_speed_linear_fit_r2_horizontal_report.pdf` 分别合并了两套 8 张单臂图，"
        "便于直接汇报或打印。\n\n"
        "## 8 个目标臂\n\n"
        "| 编号 | 实验ID | 阶段 | 架构 | 训练模式 |\n"
        "|---:|---|---|---|---|\n"
        + "\n".join(rows)
        + "\n\n横向排序图生成脚本：`docs/03-汇报材料/tools/plot_v4_bc_pde_r2_horizontal.py`。\n"
        + "代表病例回归三联图生成脚本：`docs/03-汇报材料/tools/plot_v4_representative_regressions.py`（需使用 GNN conda 环境）。\n"
    )
    path.write_text(text, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    records: list[dict[str, Any]] = []
    for target in _read_workbook_targets():
        cases = _load_cases(target)
        short = f"V4-{TIME_SHORT[target['temporal']]}-{BACKBONE_SHORT[target['backbone']]}-{MODE_SHORT[target['mode']]}"
        target = {**target, "short": short}
        records.append({**target, "cases": cases, "panel_title": short})

    args.output_dir.mkdir(parents=True, exist_ok=True)
    raw_paths: list[Path] = []
    fit_paths: list[Path] = []
    for record in records:
        stem = _safe_name(record["short"])
        arm_dir = args.output_dir / stem
        raw_path = arm_dir / f"{stem}_speed_r2_horizontal.png"
        fit_path = arm_dir / f"{stem}_speed_linear_fit_r2_horizontal.png"
        _plot_horizontal(
            raw_path,
            record,
            record["cases"],
            "speed_r2",
            "Speed R² (official per-case prediction R²)",
            f"{record['short']} | SPEED R² distribution",
        )
        _plot_horizontal(
            fit_path,
            record,
            record["cases"],
            "speed_linear_fit_r2",
            "Linear-fit R²",
            f"{record['short']} | SPEED R² distribution",
        )
        raw_paths.append(raw_path)
        fit_paths.append(fit_path)

    _plot_overview(
        args.output_dir / "V4_BC-PDE_FIXED_EMA_speed_r2_horizontal_overview.png",
        records,
        "speed_r2",
        "Speed R² (official per-case prediction R²)",
        "V4 · BC+PDE-FIXED / BC+PDE-EMA · horizontal per-case speed R²",
    )
    _plot_overview(
        args.output_dir / "V4_BC-PDE_FIXED_EMA_speed_linear_fit_r2_horizontal_overview.png",
        records,
        "speed_linear_fit_r2",
        "Linear-fit R²",
        "V4 · BC+PDE-FIXED / BC+PDE-EMA · horizontal per-case linear-fit R²",
    )
    _write_csv(args.output_dir / "V4_BC-PDE_FIXED_EMA_per_case_speed_r2.csv", records)
    _write_readme(args.output_dir / "README.md", records)
    _write_pdf(args.output_dir / "V4_BC-PDE_FIXED_EMA_speed_r2_horizontal_report.pdf", raw_paths)
    _write_pdf(args.output_dir / "V4_BC-PDE_FIXED_EMA_speed_linear_fit_r2_horizontal_report.pdf", fit_paths)
    print(json.dumps({"output_dir": str(args.output_dir), "arms": [r["experiment_id"] for r in records]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
