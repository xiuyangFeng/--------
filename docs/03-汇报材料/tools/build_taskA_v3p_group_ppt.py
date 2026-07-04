#!/usr/bin/env python3
"""Build a research-style PPT deck for Task A V3P platform-period report.

The deck is intentionally self-contained: it generates summary figures from
documented metrics, reuses existing postview PNGs, and writes a .pptx that can
be opened and further edited in PowerPoint.
"""

from __future__ import annotations

from pathlib import Path
import math
import textwrap

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
OUT_DIR = ROOT / "docs/03-汇报材料/figures/taskA_v3p_group_20260704"
PPTX_OUT = ROOT / "docs/03-汇报材料/任务A_V3P平台期与下一步_组会汇报_20260704.pptx"

FONT = "Noto Sans CJK SC"
FONT_BOLD = "Noto Sans CJK SC"

NAVY = RGBColor(22, 48, 73)
BLUE = RGBColor(44, 97, 145)
LIGHT_BLUE = RGBColor(225, 236, 247)
PALE = RGBColor(246, 249, 252)
GRAY = RGBColor(97, 112, 128)
MID_GRAY = RGBColor(160, 170, 181)
GREEN = RGBColor(37, 132, 91)
ORANGE = RGBColor(202, 104, 38)
RED = RGBColor(175, 54, 54)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(20, 24, 28)


def setup_matplotlib() -> None:
    candidates = [
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
    ]
    if hasattr(font_manager.fontManager, "addfont"):
        for p in candidates:
            if Path(p).exists():
                font_manager.fontManager.addfont(p)
    else:
        existing = [p for p in candidates if Path(p).exists()]
        if existing and hasattr(font_manager, "createFontList"):
            font_manager.fontManager.ttflist.extend(font_manager.createFontList(existing))
    plt.rcParams["font.family"] = FONT
    plt.rcParams["font.sans-serif"] = [FONT, "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 180


def save_fig(fig: plt.Figure, name: str) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / name
    fig.savefig(path, bbox_inches="tight", dpi=220, facecolor="white")
    plt.close(fig)
    return path


def make_cover_bg() -> Path:
    src = ROOT / "outputs/field/postview/v3p_i6diag_t016_report/GUO_XI_JIANG__result_features_merged-1146/plots/fig_wss_triptych.png"
    out = OUT_DIR / "cover_bg.png"
    img = Image.open(src).convert("RGB")
    w, h = img.size
    img = img.crop((0, int(0.15 * h), w, h))
    w, h = img.size
    target = (1920, 1080)
    scale = max(target[0] / w, target[1] / h)
    img = img.resize((int(w * scale), int(h * scale)))
    left = (img.size[0] - target[0]) // 2
    top = (img.size[1] - target[1]) // 2
    img = img.crop((left, top, left + target[0], top + target[1]))
    img = img.filter(ImageFilter.GaussianBlur(radius=2.2))
    overlay = Image.new("RGBA", target, (12, 31, 49, 178))
    img = Image.alpha_composite(img.convert("RGBA"), overlay)
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 760, 1920, 1080), fill=(12, 31, 49, 120))
    img.convert("RGB").save(out)
    return out


def fig_timeline() -> Path:
    fig, ax = plt.subplots(figsize=(11.8, 3.2))
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    xs = [0.10, 0.36, 0.64, 0.88]
    titles = ["V1 Field", "V2 WSS multitask", "V3P main", "External baselines"]
    dates = ["2026-03~04", "2026-04", "2026-05~07", "2026-06"]
    metrics = [
        "A-Opt-05\nR²_p=0.9234",
        "WSSP-05/06\nWSS<0",
        "I6 0.429\nI6-a 0.446",
        "CROWN p R²<0\nPNCFD p R²=0.9025",
    ]
    colors = ["#2c6191", "#ca6826", "#25845b", "#6c7785"]
    ax.plot([xs[0], xs[-1]], [0.55, 0.55], color="#b8c5d2", lw=4, solid_capstyle="round")
    for x, t, d, m, c in zip(xs, titles, dates, metrics, colors):
        ax.scatter([x], [0.55], s=580, color=c, edgecolor="white", linewidth=2, zorder=5)
        ax.text(x, 0.84, t, ha="center", va="center", fontsize=16, fontweight="bold", color=c)
        ax.text(x, 0.71, d, ha="center", va="center", fontsize=10, color="#617080")
        ax.text(
            x,
            0.23,
            m,
            ha="center",
            va="center",
            fontsize=12,
            color="#1f2933",
            bbox=dict(boxstyle="round,pad=0.35", fc="#f6f9fc", ec="#d5dee8"),
        )
    ax.text(0.5, 0.02, "Storyline: field accuracy -> direct WSS -> plateau evidence -> paradigm shift", ha="center", fontsize=12, color="#163049")
    return save_fig(fig, "timeline_v1_v2_v3p_baseline.png")


def fig_split() -> Path:
    fig, ax = plt.subplots(figsize=(9.5, 4.2))
    ax.axis("equal")
    vals = [57, 8, 16]
    labels = ["train 57", "val 8", "test 16"]
    colors = ["#2c6191", "#91b8d8", "#ca6826"]
    wedges, _ = ax.pie(vals, startangle=90, colors=colors, wedgeprops=dict(width=0.35, edgecolor="white"))
    ax.text(0, 0.08, "81", ha="center", fontsize=30, fontweight="bold", color="#163049")
    ax.text(0, -0.14, "active cases", ha="center", fontsize=13, color="#617080")
    ax.legend(wedges, labels, loc="center left", bbox_to_anchor=(1.0, 0.5), frameon=False, fontsize=12)
    ax.set_title("V3P split_AG_v1: AG-only · post-denylist", fontsize=16, color="#163049", fontweight="bold")
    fig.text(0.52, 0.05, "Discipline: V3P / V3D / external baselines are reported separately", ha="center", fontsize=10.5, color="#617080")
    return save_fig(fig, "split_ag_v1_donut.png")


def fig_route_metrics() -> Path:
    names = ["V1\nA-Opt-05", "V2\nWSSP-05/06", "V3P\nAsymW-a", "I6-diag\nplateau", "I6-a\nsingle"]
    wss = [np.nan, -0.019, 0.394, 0.429, 0.446]
    pressure = [0.9234, 0.061, 0.945, 0.930, 0.941]
    x = np.arange(len(names))
    fig, ax = plt.subplots(figsize=(10.8, 4.6))
    ax.bar(x - 0.18, [0 if math.isnan(v) else v for v in wss], width=0.34, color="#ca6826", label="WSS R²")
    ax.bar(x + 0.18, pressure, width=0.34, color="#2c6191", label="p R²")
    ax.set_ylim(-0.08, 1.02)
    ax.axhline(0.85, ls="--", color="#b53636", lw=1.4, label="M-E target 0.85")
    ax.set_xticks(x)
    ax.set_xticklabels(names)
    ax.set_ylabel("R² / metric value")
    ax.set_title("Route evolution: pressure is learnable, WSS is not engineering-ready", fontsize=15, fontweight="bold", color="#163049")
    ax.grid(axis="y", color="#e5ebf1")
    for key in ["top", "right"]:
        ax.spines[key].set_visible(False)
    ax.legend(frameon=False, ncol=3, loc="upper center", bbox_to_anchor=(0.5, -0.12))
    ax.text(0, 0.05, "V1 not\nWSS-first", ha="center", fontsize=10, color="#617080")
    return save_fig(fig, "route_metrics.png")


def fig_platform_bar() -> Path:
    names = ["AsymW-a\n3 seed", "I6-diag", "I6-a", "G58\nSSL", "J4\n+1", "J5\n+6", "K1\nBranch"]
    vals = [0.394, 0.429, 0.446, 0.438, 0.423, 0.433, 0.400]
    colors = ["#91b8d8", "#2c6191", "#25845b", "#91b8d8", "#ca6826", "#ca6826", "#b53636"]
    fig, ax = plt.subplots(figsize=(11, 4.6))
    bars = ax.bar(np.arange(len(vals)), vals, color=colors, width=0.62)
    ax.axhline(0.459, color="#b53636", ls="--", lw=1.8)
    ax.text(len(vals) - 0.1, 0.462, "Go line 0.459", ha="right", va="bottom", color="#b53636", fontsize=11)
    ax.axhspan(0.42, 0.45, color="#d9eaf7", alpha=0.55, label="plateau band 0.42–0.45")
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + 0.008, f"{v:.3f}", ha="center", fontsize=10, color="#163049")
    ax.set_xticks(np.arange(len(vals)))
    ax.set_xticklabels(names)
    ax.set_ylim(0.35, 0.48)
    ax.set_ylabel("test wss_r2_wss (best_wss)")
    ax.set_title("V3P plateau: probes do not stably cross the Go line", fontsize=15, fontweight="bold", color="#163049")
    ax.grid(axis="y", color="#e5ebf1")
    for key in ["top", "right"]:
        ax.spines[key].set_visible(False)
    return save_fig(fig, "v3p_platform_bar.png")


def fig_nogo_heatmap() -> Path:
    rows = ["A repr", "B near-wall", "F loss/feat", "G57 attention", "G58 SSL", "M-E probes", "J data", "K1 branch"]
    cols = ["WSS", "pressure", "L3/cross", "decision"]
    data = np.array([
        [0.39, 0.95, 0.05, 0],
        [0.40, 0.95, 0.04, 0],
        [0.399, 0.95, 0.04, 0],
        [0.406, 0.94, 0.05, 0],
        [0.438, 0.93, 0.06, 0],
        [0.408, 0.49, 0.05, 0],
        [0.428, 0.94, 0.12, 0],
        [0.400, 0.925, 0.05, 0],
    ])
    fig, ax = plt.subplots(figsize=(8.8, 4.5))
    colors = np.zeros_like(data)
    colors[:, 0] = data[:, 0]
    colors[:, 1] = data[:, 1]
    colors[:, 2] = data[:, 2]
    colors[:, 3] = 0.2
    im = ax.imshow(colors, cmap="RdYlGn", vmin=0, vmax=1, aspect="auto")
    ax.set_xticks(np.arange(len(cols)))
    ax.set_xticklabels(cols)
    ax.set_yticks(np.arange(len(rows)))
    ax.set_yticklabels(rows)
    for i in range(len(rows)):
        for j in range(len(cols)):
            txt = "No-Go" if j == 3 else (f"{data[i,j]:.3f}" if j == 0 else f"{data[i,j]:.2f}")
            ax.text(j, i, txt, ha="center", va="center", fontsize=10, color="#1f2933")
    ax.set_title("No-Go map: pressure survives, WSS/cross components remain closed", fontsize=15, fontweight="bold", color="#163049")
    ax.tick_params(length=0)
    for s in ax.spines.values():
        s.set_visible(False)
    return save_fig(fig, "nogo_heatmap.png")


def fig_j45_l3() -> Path:
    cats = ["all", "fast", "slow", "CHEN"]
    i6 = [0.207, 0.054, 0.250, -0.825]
    j5 = [0.258, 0.198, 0.275, -0.241]
    x = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9.8, 4.3))
    ax.bar(x - 0.18, i6, 0.36, label="I6-diag", color="#2c6191")
    ax.bar(x + 0.18, j5, 0.36, label="J5 +6 AAA", color="#ca6826")
    ax.axhline(0, color="#617080", lw=1)
    ax.set_xticks(x)
    ax.set_xticklabels(cats)
    ax.set_ylabel("L3 subset WSS R²")
    ax.set_title("J5 improves subsets, but CHEN remains negative and global test stays flat", fontsize=14.5, fontweight="bold", color="#163049")
    ax.legend(frameon=False)
    ax.grid(axis="y", color="#e5ebf1")
    for key in ["top", "right"]:
        ax.spines[key].set_visible(False)
    return save_fig(fig, "j5_l3_subset.png")


def fig_external_crown() -> Path:
    methods = ["CROWN\nnon-PINN", "CROWN\nPINN", "PointNetCFD\ngeom_vp", "V3P\nI6-diag"]
    p_r2 = [-2.07, -0.63, 0.9025, 0.930]
    colors = ["#b53636", "#ca6826", "#2c6191", "#25845b"]
    fig, ax = plt.subplots(figsize=(10.2, 4.5))
    ax.bar(np.arange(len(methods)), p_r2, color=colors, width=0.58)
    ax.axhline(0, color="#617080", lw=1.2)
    ax.axhline(0.9, color="#25845b", ls="--", lw=1.2)
    for i, v in enumerate(p_r2):
        ax.text(i, v + (0.08 if v >= 0 else -0.16), f"{v:.2f}", ha="center", fontsize=11, color="#163049")
    ax.set_xticks(np.arange(len(methods)))
    ax.set_xticklabels(methods)
    ax.set_ylabel("p R²")
    ax.set_title("External baselines: xyz(+PINN) cannot replace the V3P route", fontsize=15, fontweight="bold", color="#163049")
    ax.set_ylim(-2.4, 1.15)
    ax.grid(axis="y", color="#e5ebf1")
    for key in ["top", "right"]:
        ax.spines[key].set_visible(False)
    return save_fig(fig, "external_p_r2.png")


def fig_pointnetcfd() -> Path:
    names = ["original_vp", "wall_vp", "geom_vp", "geom_pwss"]
    rmse = [0.8477, 0.7390, 0.6896, 0.7524]
    p = [0.8970, 0.8910, 0.9025, 0.9141]
    wss_xy = [np.nan, np.nan, np.nan, 0.02]
    fig, ax1 = plt.subplots(figsize=(10.3, 4.4))
    x = np.arange(len(names))
    ax1.bar(x - 0.18, rmse, 0.36, color="#91b8d8", label="RMSE")
    ax2 = ax1.twinx()
    ax2.plot(x + 0.18, p, "o-", color="#2c6191", lw=2, label="p R²")
    ax2.scatter([3], [wss_xy[3]], s=90, color="#b53636", label="wss_x/y≈0.02")
    ax1.set_xticks(x)
    ax1.set_xticklabels(names)
    ax1.set_ylabel("normalized RMSE")
    ax2.set_ylabel("R²")
    ax1.set_ylim(0.45, 0.92)
    ax2.set_ylim(0, 1.02)
    ax1.set_title("PointNetCFD: pressure sanity works; direct WSS prediction is No-Go", fontsize=15, fontweight="bold", color="#163049")
    ax1.grid(axis="y", color="#e5ebf1")
    ax1.spines["top"].set_visible(False)
    ax2.spines["top"].set_visible(False)
    lines, labels = ax1.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax1.legend(lines + lines2, labels + labels2, frameon=False, loc="upper center", ncol=3, bbox_to_anchor=(0.5, -0.12))
    return save_fig(fig, "pointnetcfd_matrix.png")


def fig_bottleneck() -> Path:
    fig, ax = plt.subplots(figsize=(9.5, 4.8))
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    pts = np.array([[0.5, 0.86], [0.12, 0.18], [0.88, 0.18]])
    poly = plt.Polygon(pts, closed=True, facecolor="#eef5fb", edgecolor="#2c6191", linewidth=2)
    ax.add_patch(poly)
    labels = [
        ("Representation", "WSS is a near-wall\noperator, not a plain scalar", pts[0]),
        ("Information", "local kNN misses upstream\nbranch / inlet context", pts[1]),
        ("Data", "AG-only expansion is capped\nJ4/J5 did not break plateau", pts[2]),
    ]
    for title, body, (x, y) in labels:
        ax.scatter([x], [y], s=920, color="#2c6191", edgecolor="white", linewidth=2, zorder=3)
        ax.text(x, y + 0.01, title, ha="center", va="center", color="white", fontweight="bold", fontsize=14)
        ax.text(x, y - 0.13 if y > 0.5 else y + 0.18, body, ha="center", va="center", color="#1f2933", fontsize=12)
    ax.text(0.5, 0.44, "Not a hyperparameter issue", ha="center", va="center", fontsize=18, fontweight="bold", color="#b53636")
    ax.text(0.5, 0.36, "Change representation / processor / active-data loop", ha="center", fontsize=12, color="#617080")
    return save_fig(fig, "bottleneck_triangle.png")


def fig_next_roadmap() -> Path:
    fig, ax = plt.subplots(figsize=(10.8, 4.6))
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    items = [
        ("P0", "Differentiable WSS", "profile basis + physics layer"),
        ("P0", "UQ active data", "calibrated uncertainty for selection"),
        ("P1", "Physics attention", "long-range branch context"),
        ("P2", "GINO / multi-fidelity", "external baseline + synthetic data"),
        ("P3", "Diffusion for hotspot", "repair high-WSS regions"),
    ]
    xs = np.linspace(0.08, 0.92, len(items))
    for i, (prio, title, note) in enumerate(items):
        x = xs[i]
        color = "#25845b" if prio == "P0" else ("#2c6191" if prio == "P1" else "#6c7785")
        ax.plot([x, x], [0.28, 0.65], color=color, lw=3)
        ax.scatter([x], [0.65], s=650, color=color, edgecolor="white", linewidth=2)
        ax.text(x, 0.65, prio, ha="center", va="center", color="white", fontweight="bold", fontsize=14)
        ax.text(x, 0.20, title, ha="center", fontsize=12.5, fontweight="bold", color="#163049", wrap=True)
        ax.text(x, 0.08, note, ha="center", fontsize=10, color="#617080", wrap=True)
    ax.plot([xs[0], xs[-1]], [0.65, 0.65], color="#cbd6e2", lw=2, zorder=0)
    ax.text(0.5, 0.91, "Next: move from tuning to testable new paradigms", ha="center", fontsize=16, fontweight="bold", color="#163049")
    ax.text(0.5, 0.80, "Planning only: each line needs a zero-retrain oracle before seed1 short training", ha="center", fontsize=11.5, color="#b53636")
    return save_fig(fig, "next_roadmap.png")


def generate_figures() -> dict[str, Path]:
    setup_matplotlib()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    return {
        "cover": make_cover_bg(),
        "timeline": fig_timeline(),
        "split": fig_split(),
        "route": fig_route_metrics(),
        "platform": fig_platform_bar(),
        "nogo": fig_nogo_heatmap(),
        "j5": fig_j45_l3(),
        "external": fig_external_crown(),
        "pointnetcfd": fig_pointnetcfd(),
        "bottleneck": fig_bottleneck(),
        "roadmap": fig_next_roadmap(),
    }


def set_font(run, size=None, bold=False, color=None):
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_textbox(slide, x, y, w, h, text, size=16, color=BLACK, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    set_font(r, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=14, color=BLACK, leading=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(5)
        p.line_spacing = leading
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, 0.55, 0.28, 9.5, 0.45, title, 25, NAVY, True)
    if subtitle:
        add_textbox(slide, 0.58, 0.78, 10.5, 0.28, subtitle, 11.5, GRAY)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.08), Inches(1.05), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = ORANGE
    line.line.fill.background()


def add_footer(slide, idx):
    add_textbox(slide, 0.55, 7.12, 7.5, 0.22, "任务A · V3P平台期与下一步 · split_AG_v1 / best_wss口径", 8.5, MID_GRAY)
    add_textbox(slide, 12.0, 7.12, 0.7, 0.22, f"{idx:02d}", 8.5, MID_GRAY, align=PP_ALIGN.RIGHT)


def add_conclusion(slide, text, color=BLUE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.75), Inches(1.18), Inches(3.95), Inches(0.58))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    add_textbox(slide, 8.92, 1.31, 3.62, 0.28, text, 12, WHITE, True, align=PP_ALIGN.CENTER)


def add_metric_card(slide, x, y, w, h, label, value, note="", color=BLUE):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid()
    rect.fill.fore_color.rgb = PALE
    rect.line.color.rgb = RGBColor(214, 224, 235)
    add_textbox(slide, x + 0.18, y + 0.15, w - 0.36, 0.24, label, 10.5, GRAY, True)
    add_textbox(slide, x + 0.18, y + 0.45, w - 0.36, 0.42, value, 21, color, True)
    if note:
        add_textbox(slide, x + 0.18, y + 0.92, w - 0.36, 0.32, note, 9.2, GRAY)


def fit_image(slide, path, x, y, w, h):
    img = Image.open(path)
    iw, ih = img.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        width = w
        height = w / img_ratio
        left = x
        top = y + (h - height) / 2
    else:
        height = h
        width = h * img_ratio
        left = x + (w - width) / 2
        top = y
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def add_table(slide, x, y, w, h, headers, rows, font_size=8.8, header_color=BLUE):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for i, head in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = head
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_color
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                set_font(r, size=font_size, bold=True, color=WHITE)
    for r_idx, row in enumerate(rows, 1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(248, 251, 253) if r_idx % 2 else WHITE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if c_idx > 0 else PP_ALIGN.LEFT
                for run in p.runs:
                    set_font(run, size=font_size, color=BLACK)
    return table_shape


def blank_slide(prs, idx, title, subtitle=None, conclusion=None, conclusion_color=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    add_title(slide, title, subtitle)
    add_footer(slide, idx)
    return slide


def build_deck(figs: dict[str, Path]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(figs["cover"]), 0, 0, width=prs.slide_width, height=prs.slide_height)
    add_textbox(slide, 0.68, 0.65, 9.9, 0.5, "任务 A：AG 脑血管 WSS 场重建", 29, WHITE, True)
    add_textbox(slide, 0.72, 1.25, 8.7, 0.4, "V3P 平台期证据、外部 baseline 与下一步换范式路线", 17, RGBColor(225, 236, 247), True)
    add_metric_card(slide, 0.78, 5.42, 2.25, 1.18, "当前 WSS 平台", "0.42–0.45", "I6-diag 0.429", ORANGE)
    add_metric_card(slide, 3.20, 5.42, 2.25, 1.18, "压力 R²", "0.93–0.96", "V3P 主线", BLUE)
    add_metric_card(slide, 5.62, 5.42, 2.25, 1.18, "工程目标", "0.85", "WSS R²", RED)
    add_textbox(slide, 9.2, 6.6, 3.35, 0.28, "导师组会 · 15–20 min · 2026-07-04", 10.5, RGBColor(225, 236, 247), align=PP_ALIGN.RIGHT)

    # 2 problem
    slide = blank_slide(prs, 2, "问题与指标定义", "主指标固定为 V3P test wss_r2_wss（best_wss）", "压力能学 ≠ WSS 梯度可用")
    add_metric_card(slide, 0.75, 1.45, 2.65, 1.25, "主指标", "WSS R²", "test wss_r2_wss · best_wss", ORANGE)
    add_metric_card(slide, 3.65, 1.45, 2.65, 1.25, "副指标", "p R²", "test r2_p", BLUE)
    add_metric_card(slide, 6.55, 1.45, 2.65, 1.25, "工程目标", "≈0.85", "M-E 替代 CFD", RED)
    add_metric_card(slide, 9.45, 1.45, 2.65, 1.25, "当前平台", "≈0.43", "距目标约 0.42", ORANGE)
    add_bullets(slide, 0.9, 3.15, 5.25, 2.8, [
        "输入：壁面/近壁点云，几何特征，入口/出口边界条件。",
        "输出：壁面 WSS 与压力场；V3P 重点优化 WSS。",
        "评价：pooled test R²，同时检查 L3/CHEN/fast 子集。",
        "纪律：不同 split、不同点集、不同目标不做同榜排名。",
    ], 14)
    fit_image(slide, figs["split"], 6.5, 3.0, 5.7, 3.25)

    # 3 split discipline
    slide = blank_slide(prs, 3, "数据与口径纪律", "先统一 split，后谈指标", "V3P / V3D / baseline 必须分表")
    fit_image(slide, figs["split"], 0.7, 1.35, 5.7, 4.7)
    rows = [
        ["V3P", "split_AG_v1", "AG-only", "81 活跃例，57/8/16"],
        ["V3D", "split_data_new_v3", "AAA+AG+ILO", "禁止与 V3P 同表对比"],
        ["CROWN", "split_AG_v1", "仅 xyz", "全点 NMAE 主指标"],
        ["PointNetCFD", "split_AG_v1", "点云 VP/WSS", "归一化 RMSE/R²"],
    ]
    add_table(slide, 6.55, 1.55, 5.95, 2.6, ["对象", "split", "输入/域", "汇报口径"], rows, 9.5)
    add_bullets(slide, 6.75, 4.45, 5.4, 1.5, [
        "V3P 主表只看 `split_AG_v1`。",
        "CROWN 的 wall/interior 是体点云诊断，不等同 V3P 壁面点。",
        "外部 baseline 只说明输入与归纳偏置边界。",
    ], 12.5, GRAY)

    # 4 timeline
    slide = blank_slide(prs, 4, "方法演进时间线", "从场重建到 WSS 直接监督，再到外部方法边界", "演进逻辑：场精度 → WSS 平台 → 换范式")
    fit_image(slide, figs["timeline"], 0.72, 1.45, 11.9, 4.35)
    add_bullets(slide, 1.0, 6.15, 11.0, 0.75, [
        "V1 证明压力可学；V2 排除简单 WSS 多任务；V3P 建立当前最强主线；external baseline 排除“仅 xyz(+PINN)”替代方案。"
    ], 12.5, GRAY)

    # 5 V1
    slide = blank_slide(prs, 5, "V1：场重建成功，但不是 WSS 终点", "A-Opt-05 是场重建母版；压力 R²≈0.92+", "V1 回答了压力，不足以回答 WSS")
    fit_image(slide, figs["route"], 0.75, 1.45, 6.55, 4.3)
    rows = [
        ["A-Opt-05", "RMSE_|v| 1.0399±0.0082", "R²_p 0.9234±0.0005"],
        ["A-Opt-07", "RMSE_|v| 1.0433±0.0120", "R²_p 0.9224±0.0035"],
        ["A-Opt-05-wss-multi", "壁面 WSS R² 0.463±0.004", "R²_p 0.922±0.003"],
    ]
    add_table(slide, 7.65, 1.7, 4.75, 2.2, ["实验", "速度/WSS", "压力"], rows, 8.5)
    add_bullets(slide, 7.75, 4.3, 4.35, 1.7, [
        "压力和低频场量可稳定学习。",
        "WSS 依赖近壁速度梯度，不能只用场 RMSE 代替。",
        "后续路线必须把 WSS 作为主目标或结构算子。"
    ], 12.5)

    # 6 V2
    slide = blank_slide(prs, 6, "V2：WSS 多任务全线 No-Go", "WSS loss 易劫持梯度，简单重加权不能解决", "排除了“直接加 WSS loss”的路线")
    rows = [
        ["WSSP-01", "p+WSS", "wall.r2_wss≈0.44", "单 seed，速度未训"],
        ["WSSP-02", "全场+WSS 0.5", "-0.018", "p R²=0.004，崩溃"],
        ["WSSP-03", "轻量 WSS 0.01", "-0.017", "压力略修复"],
        ["WSSP-05", "p+WSS 三 seed", "-0.025/-0.022/-0.016", "No-Go"],
        ["WSSP-06", "Huber WSS", "-0.016/-0.015/-0.013", "未稳定优于 MSE"],
    ]
    add_table(slide, 0.78, 1.55, 7.2, 3.05, ["实验", "配置", "WSS R²", "判读"], rows, 8.6)
    add_bullets(slide, 8.35, 1.55, 3.9, 2.8, [
        "WSS loss 尺度远大于 data loss 时，压力被拖垮。",
        "降低 WSS 权重可缓解压力，但 WSS R² 不转正。",
        "V2 结论推动 V3P 采用 wall-rich + AsymW + best_wss 选优。"
    ], 13)
    fit_image(slide, ROOT / "outputs/field/field_v2_pointnext_wssp02_geom_full_supervision_wss_split_AG_v1_seed1_20260425_093041/predictions_test/regional_eval/fig_A5_regional_bar_rmse_p.png", 8.4, 4.4, 3.6, 2.15)

    # 7 architecture
    slide = blank_slide(prs, 7, "V3P 架构：PointNeXt 双域 + AsymW-a", "wall-rich 采样、几何+BC、p/WSS 双输出", "当前主线：让模型直接面向 WSS")
    # Draw architecture blocks
    blocks = [
        (0.9, 2.0, 2.1, 0.9, "壁面/近壁点\nwall13000+near2000", BLUE),
        (0.9, 3.35, 2.1, 0.9, "几何 + BC\ncoord/t/branch/radius", BLUE),
        (3.7, 2.55, 2.25, 1.15, "PointNeXt\n双域局部池化", GREEN),
        (6.65, 1.85, 2.2, 0.95, "Pressure head\np", BLUE),
        (6.65, 3.25, 2.2, 0.95, "WSS head\nAsymW-a", ORANGE),
        (9.65, 1.85, 2.25, 0.95, "副指标\nr2_p≈0.93–0.96", BLUE),
        (9.65, 3.25, 2.25, 0.95, "主指标\nwss_r2_wss≈0.43", ORANGE),
    ]
    for x, y, w, h, txt, col in blocks:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = col
        shape.line.fill.background()
        add_textbox(slide, x + 0.08, y + 0.18, w - 0.16, h - 0.25, txt, 12.5, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for x1, y1, x2, y2 in [(3.0,2.45,3.7,3.0),(3.0,3.8,3.7,3.12),(5.95,3.05,6.65,2.35),(5.95,3.05,6.65,3.75),(8.85,2.35,9.65,2.35),(8.85,3.75,9.65,3.75)]:
        slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)).line.color.rgb = MID_GRAY
    add_bullets(slide, 1.15, 5.25, 10.8, 1.0, [
        "AsymW-a 使用非对称 WSS 权重，优先保护主分量；best_wss 选优避免 best_model 与 WSS 主目标错位。",
        "压力长期保持高 R²，核心难点集中在壁面剪切尤其横向/周向信息。"
    ], 12.5, GRAY)

    # 8 master
    slide = blank_slide(prs, 8, "母版确立：AsymW-a 抬升到平台带宽", "4957/4999/5000 → post5463/I6", "AsymW 有效，但不是工程突破")
    fit_image(slide, ROOT / "docs/03-汇报材料/figures/v3p_asymw_seed_consistency.png", 0.75, 1.35, 5.5, 3.85)
    fit_image(slide, figs["platform"], 6.55, 1.35, 5.75, 3.85)
    rows = [
        ["AsymW-a 3 seed", "0.394±0.005", "~0.93–0.96", "升 PW 母版"],
        ["post5463 band", "0.425±0.012", "~0.93–0.96", "平台参考"],
        ["I6-diag", "0.429", "0.930", "200ep 诊断"],
    ]
    add_table(slide, 1.0, 5.65, 11.3, 0.95, ["锚点", "WSS R²", "p R²", "判读"], rows, 8.2)

    # 9 platform
    slide = blank_slide(prs, 9, "平台期关键实验", "I6-diag / I6-a / G58 / J4 / J5 / K1", "多条探针未稳定跨过 Go 线", ORANGE)
    fit_image(slide, figs["platform"], 0.82, 1.35, 7.25, 4.55)
    rows = [
        ["I6-diag", "0.429", "0.930", "平台参考"],
        ["I6-a", "0.446", "0.941", "单点最高"],
        ["J5 +6 AAA", "0.433", "0.951", "弱 No-Go"],
        ["K1 Branch", "0.400", "0.925", "No-Go"],
    ]
    add_table(slide, 8.35, 1.65, 3.8, 2.5, ["实验", "WSS", "p", "判读"], rows, 8.8)
    add_bullets(slide, 8.45, 4.55, 3.45, 1.2, [
        "Go 门槛：I6-diag +0.03 = 0.459。",
        "I6-a 未扩 seed，不能写成稳定突破。",
        "压力提升不自动带来 WSS 突破。"
    ], 11.5, GRAY)

    # 10 nogo map
    slide = blank_slide(prs, 10, "系统性探索与 No-Go 地图", "路径 A–G、M-E 探针、J/K 平台期探针", "平台期来自系统性封口，不是单点失败")
    fit_image(slide, figs["nogo"], 0.78, 1.35, 7.1, 4.5)
    add_bullets(slide, 8.25, 1.55, 3.95, 3.25, [
        "路径 A/B：表示和近壁几何约 0.37–0.41。",
        "路径 F：loss/特征探针多数 No-Go；PGradFeat 三 seed 0.399±0.010。",
        "路径 G57：0.392–0.420；G58 SSL 0.435–0.440 仍未过线。",
        "M-E WSS-only/PWeak/E-J：0.401–0.411，压力或迁移性不足。"
    ], 12)
    add_textbox(slide, 8.35, 5.35, 3.7, 0.45, "判读：继续旧范式横扫的边际收益很低", 14, RED, True, align=PP_ALIGN.CENTER)

    # 11 J4/J5
    slide = blank_slide(prs, 11, "数据探针：扩池不能直接破平台", "J4 +1 / J5 +6 AAA / L3 子集", "局部改善不等于全 test 突破")
    fit_image(slide, figs["j5"], 0.78, 1.35, 6.7, 4.25)
    rows = [
        ["J4 +1", "0.423", "-0.006", "No-Go"],
        ["J5 +6 AAA", "0.433", "+0.004", "弱 No-Go"],
        ["L3 fast", "0.054→0.198", "+0.145", "改善但绝对低"],
        ["L3 CHEN", "-0.825→-0.241", "+0.584", "仍为负"],
    ]
    add_table(slide, 7.8, 1.55, 4.55, 2.65, ["对象", "指标", "Δ", "判读"], rows, 8.6)
    add_bullets(slide, 8.0, 4.65, 4.0, 1.2, [
        "J5 说明主动选数有局部信号，但全 test band 未破 0.45。",
        "因此 J6 GPU 不自动开启，Batch-2 只保留采购/记录价值。"
    ], 12, GRAY)

    # 12 K1 oracle vs gpu
    slide = blank_slide(prs, 12, "结构探针：K1 oracle Go，GPU No-Go", "branch/topology 信号存在，但手工 concat 不足", "oracle 有信号 ≠ 模型可用")
    rows = [
        ["K1 CPU", "branch error explained 0.683", "Go", "超过 0.30 阈值"],
        ["K3 CPU", "mode R² 0.170 < abs 0.235", "No-Go", "非幅值标定主因"],
        ["K5 CPU", "BC explained -0.43", "No-Go", "BC/相位不解释误差"],
        ["K1 GPU 5886", "WSS 0.400 / p 0.925", "No-Go", "未过 0.459"],
    ]
    add_table(slide, 0.85, 1.55, 6.4, 2.9, ["探针", "关键数", "判读", "说明"], rows, 8.7)
    # simple oracle to gpu funnel
    for x, label, val, col in [(8.0, "CPU oracle", "0.683", GREEN), (10.3, "GPU probe", "0.400", RED)]:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(1.7), Inches(1.05))
        shape.fill.solid()
        shape.fill.fore_color.rgb = col
        shape.line.fill.background()
        add_textbox(slide, x+0.1, 2.18, 1.5, 0.25, label, 10.5, WHITE, True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x+0.1, 2.52, 1.5, 0.28, val, 20, WHITE, True, align=PP_ALIGN.CENTER)
    slide.shapes.add_connector(1, Inches(9.75), Inches(2.52), Inches(10.3), Inches(2.52)).line.color.rgb = MID_GRAY
    add_bullets(slide, 8.0, 3.85, 3.95, 1.65, [
        "下一步若继续补分支信息，应换 processor 或全局 token，而不是再加手工特征。",
        "这直接指向 Transolver / physics-attention。"
    ], 12.2)

    # 13 CROWN
    slide = blank_slide(prs, 13, "外部 baseline：CROWN 复现成功但 AG 不可用", "非 PINN / PINN 均已完成 paper_full evaluate", "仅 xyz(+PINN) 不能替代 V3P")
    fit_image(slide, figs["external"], 0.78, 1.32, 6.7, 4.25)
    rows = [
        ["非 PINN", "p NMAE 16.0%", "p R² -2.07", "No-Go"],
        ["PINN", "p NMAE 13.2%", "p R² -0.63", "改善但仍 No-Go"],
        ["V3P 对照", "-", "p R² 0.930", "主模型"],
    ]
    add_table(slide, 7.75, 1.55, 4.6, 2.2, ["方法", "NMAE", "R²", "判读"], rows, 8.8)
    add_bullets(slide, 7.95, 4.2, 4.15, 1.8, [
        "CROWN 主指标是 NMAE；R² 是本项目补充诊断。",
        "PINN 降低压力 NMAE约18%，但没有学到足够空间结构。",
        "CROWN 无 WSS 输出，不能作为 WSS baseline。"
    ], 11.8, GRAY)

    # 14 PointNetCFD
    slide = blank_slide(prs, 14, "外部 baseline：PointNetCFD 的边界", "压力 sanity baseline 可用，WSS 直预测不够", "信息量比模型名更关键")
    fit_image(slide, figs["pointnetcfd"], 0.75, 1.3, 6.9, 4.4)
    rows = [
        ["original_vp", "0.8477", "0.8970", "速度 R²<0.1"],
        ["wall_vp", "0.7390", "0.8910", "w R² 0.471"],
        ["geom_vp", "0.6896", "0.9025", "VP sanity"],
        ["geom_pwss", "0.7524", "0.9141", "wss_x/y≈0.02"],
    ]
    add_table(slide, 7.95, 1.55, 4.4, 2.65, ["实验", "RMSE", "p R²", "判读"], rows, 8.5)
    add_bullets(slide, 8.05, 4.55, 4.05, 1.3, [
        "加 wall/几何能显著改善速度，说明输入信息是关键缺口。",
        "但 WSS 直预测的横向分量仍接近 0。"
    ], 12, GRAY)

    # 15 visualization
    slide = blank_slide(prs, 15, "定性可视化：merged-1146 同相位对照", "GUO / ZHANG / CHEN · 收缩期上升段 · 仅作定性", "CROWN 压力结构弱，且无 WSS")
    base = ROOT / "outputs/field/postview"
    imgs = [
        ("V3P p", base / "v3p_i6diag_t016_report/GUO_XI_JIANG__result_features_merged-1146/plots/fig_p_triptych.png"),
        ("V3P WSS", base / "v3p_i6diag_t016_report/GUO_XI_JIANG__result_features_merged-1146/plots/fig_wss_triptych.png"),
        ("CROWN PINN p", base / "crown_pinn_t016_report/GUO_XI_JIANG__result_features_merged-1146/plots/fig_p_triptych.png"),
        ("CROWN PINN |v|", base / "crown_pinn_t016_report/GUO_XI_JIANG__result_features_merged-1146/plots/fig_vel_mag_triptych.png"),
    ]
    coords = [(0.75,1.45),(6.85,1.45),(0.75,4.15),(6.85,4.15)]
    for (label, path), (x, y) in zip(imgs, coords):
        fit_image(slide, path, x, y, 5.6, 2.05)
        add_textbox(slide, x, y-0.25, 2.8, 0.2, label, 9.8, NAVY, True)
    add_textbox(slide, 0.9, 6.55, 11.2, 0.25, "注：单帧壁面点对照不同于 CROWN paper_full 全点指标；V3P 与 CROWN 不做 WSS 同榜排名。", 9.5, GRAY)

    # 16 bottleneck
    slide = blank_slide(prs, 16, "瓶颈判断：表示 / 信息 / 数据三重约束", "平台期不是超参问题", "下一步要换范式")
    fit_image(slide, figs["bottleneck"], 0.9, 1.4, 6.45, 4.45)
    rows = [
        ["表示", "WSS=近壁微分算子", "direct head 难学横向剪切"],
        ["信息", "局部 kNN / 手工 concat", "上游与分支条件不足"],
        ["数据", "AG-only 图内扩池", "J4/J5 未破平台"],
    ]
    add_table(slide, 7.75, 1.75, 4.45, 2.25, ["瓶颈", "证据", "含义"], rows, 8.8)
    add_bullets(slide, 7.95, 4.4, 3.95, 1.45, [
        "旧方向的 No-Go 不是坏消息，而是把后续投入从“继续微调”收束到“换假设”。",
        "新方向必须有 0 重训 oracle 和 L3 不退化门禁。"
    ], 12)

    # 17 MN vs ME
    slide = blank_slide(prs, 17, "分轨结论：M-N vs M-E", "论文/叙事可收口；工程替代 CFD 未达标", "诚实平台期 + 明确下一步")
    rows = [
        ["M-N 论文/方法叙事", "压力可用；病例级 Pa p95 Spearman 0.559；G5 区域 surrogate", "可阶段性收口"],
        ["M-E 工程替代 CFD", "WSS R²≈0.43，距 0.85 差约 0.42；CHEN/L3 仍弱", "不可宣称可用"],
        ["共同要求", "继续报告 best_wss、L3、横向/hotspot 与压力不退化", "防止单指标误判"],
    ]
    add_table(slide, 0.85, 1.6, 11.55, 2.45, ["轨道", "证据", "结论"], rows, 10)
    add_metric_card(slide, 1.0, 4.55, 2.8, 1.2, "M-N 可述", "0.559", "Pa p95 Spearman", GREEN)
    add_metric_card(slide, 4.25, 4.55, 2.8, 1.2, "M-E 当前", "0.429", "I6-diag WSS R²", ORANGE)
    add_metric_card(slide, 7.5, 4.55, 2.8, 1.2, "工程缺口", "~0.42", "0.85 - 0.43", RED)

    # 18 roadmap
    slide = blank_slide(prs, 18, "下一步前沿方向（规划层 / 未跑实验）", "只保留能改变旧范式隐含假设的方向", "P0 先做可微 WSS 与 UQ 主动选数")
    fit_image(slide, figs["roadmap"], 0.75, 1.35, 11.75, 4.75)
    add_textbox(slide, 1.0, 6.35, 11.0, 0.35, "不建议重跑：AsymW 微调、手工 branch concat、有限差分 vel_diff、盲目扩池、CROWN paper-original。", 11, RED, True, align=PP_ALIGN.CENTER)

    # 19 summary
    slide = blank_slide(prs, 19, "总结：三句话版", "给导师组会的决策锚点", "Go/No-Go 清楚，下一步不散")
    add_metric_card(slide, 0.9, 1.55, 3.55, 1.35, "1. 压力场", "可学", "V3P p R² 0.93–0.96", GREEN)
    add_metric_card(slide, 4.9, 1.55, 3.55, 1.35, "2. WSS", "平台期", "0.42–0.45，未过 0.459", ORANGE)
    add_metric_card(slide, 8.9, 1.55, 3.55, 1.35, "3. 下一步", "换范式", "可微 WSS / UQ / attention", BLUE)
    add_bullets(slide, 1.15, 3.65, 10.8, 1.8, [
        "已完成 V1→V2→V3P 的路线演进，并用外部 baseline 验证输入信息与 inductive bias 的边界。",
        "平台期判断有多条独立证据支撑：loss、表示、数据扩池、branch 特征与 PINN baseline 均没有稳定突破。",
        "M-N 可以围绕压力与病例级 surrogate 叙事；M-E 必须进入可微 WSS、长程 processor 和主动选数闭环。"
    ], 14)

    # 20 backup
    slide = blank_slide(prs, 20, "备份：实验规模与证据索引", "数字来自 xlsx / experiment_index.csv / 跟踪日志 / baseline 记录", "所有核心数字可追溯")
    rows = [
        ["field run 汇总", "outputs/field/experiment_index.csv", "169 条 run"],
        ["实验记录表", "docs/00-规范与记录/实验记录表.xlsx", "160+ 行，2026-07-04 零缺口"],
        ["V3P 日志", "V3_实验执行跟踪日志.md", "顶部平台期/J/K/M-E 记录"],
        ["CROWN", "CROWN_非PINN与PINN复现汇报_合并.md", "5751/5774, 5757/5806"],
        ["PointNetCFD", "docs/paper_reproduction/papers/pointnetcfd/梳理记录.md", "四组矩阵"],
        ["下一步", "WSS精度突破的前沿方向与见解_2026-07-04.md", "规划层"],
    ]
    add_table(slide, 0.85, 1.55, 11.7, 3.35, ["材料", "路径", "用途/规模"], rows, 9)
    add_bullets(slide, 1.05, 5.35, 11.0, 0.9, [
        "Q&A 中若被问到某个 Job，可回到备份页说明：主表只讲 V3P 口径，外部 baseline 分表，规划方向不混作实验结论。"
    ], 12.5, GRAY)

    PPTX_OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX_OUT)


def main() -> None:
    figs = generate_figures()
    build_deck(figs)
    print(PPTX_OUT)
    print(OUT_DIR)


if __name__ == "__main__":
    main()
