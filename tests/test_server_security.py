"""v0.14 security review (S1–S11): sessions die with credentials, claim takeover, read-only admin trash,
formula injection, resource limits, headers / CSP, access + audit logs, staged report refresh, error detail,
gzip + ETag, health / ready."""
from __future__ import annotations

import gzip
import hashlib
import http.client
import io
import json
import logging
import threading
import time
import zipfile
from pathlib import Path

import pytest

from wss_deploy import server as S
from wss_deploy import report_freshness as F
from wss_deploy.export_table import COLUMNS, to_csv, to_xlsx
from wss_deploy.jobs import JobError
from wss_deploy.registry import ReleaseError
from wss_deploy.server import LoginThrottle, SessionStore, stream_multipart
from wss_deploy.users import MAX_TRACKED_FAILURES, UserStore
from tests._c_helpers import Service, call, finished, multipart


class Clock:
    def __init__(self, now=1_800_000_000.0): self.now = now
    def __call__(self): return self.now


def _users(root: Path) -> UserStore:
    store = UserStore(root / "users.json")
    store.add("alice", "alice-secret")
    store.add("root", "root-secret", admin=True)
    return store


def _cookie(sid: str) -> str:
    return f"{S.COOKIE_NAME}={sid}"


def _login(api, body, cookie=None):
    headers = {"Content-Type": "application/json"}
    if cookie:
        headers["Cookie"] = cookie
    api.request("POST", "/api/session", body=json.dumps(body), headers=headers)
    response = api.getresponse(); payload = json.loads(response.read())
    cookie = (response.getheader("Set-Cookie") or "").split(";", 1)[0]
    return response.status, payload, cookie, {"Cookie": cookie, "X-CSRF-Token": payload.get("csrf_token") or "", "Content-Type": "application/json"}


def _stl(faces: int) -> bytes:
    """Binary STL with ``faces`` triangles (header + count + 50 bytes per face)."""
    import struct
    body = b"".join(struct.pack("<12fH", 0, 0, 1, i, 0, 0, i + 1, 0, 0, i, 1, 0, 0) for i in range(faces))
    return b"wss test".ljust(80, b" ") + struct.pack("<I", faces) + body


# ----------------------------------------------------------------------------------------------- S1
def test_sessions_die_with_the_token_and_the_password(tmp_path):
    users = UserStore(tmp_path / "users.json")
    store = SessionStore(tmp_path / "sessions.json", shared=True, token="token-one", users=users)
    sid, row = store.create("token-one")
    assert row["token_fingerprint"] and "token-one" not in json.dumps(row)
    assert store.lookup(_cookie(sid))[1]["role"] == "legacy"
    # restart with a rotated token: the persisted row no longer authenticates
    rotated = SessionStore(tmp_path / "sessions.json", shared=True, token="token-two", users=users)
    assert rotated.lookup(_cookie(sid)) is None and sid not in rotated.sessions
    # users appear → the token window closes → legacy rows stop authenticating (but stay for the claim offer)
    again = SessionStore(tmp_path / "sessions.json", shared=True, token="token-one", users=users)
    sid, _ = again.create("token-one")
    users.add("alice", "alice-secret"); users.add("root", "root-secret", admin=True)
    assert again.lookup(_cookie(sid)) is None and again.lookup(_cookie(sid), for_claim=True) is not None
    # password sessions: passwd / disable / role change each end the user's sessions
    a1, _ = again.create(username="alice", password="alice-secret")
    r1, _ = again.create(username="root", password="root-secret")
    assert again.lookup(_cookie(a1)) and again.lookup(_cookie(r1))
    users.set_password("alice", "alice-secret-2")
    assert again.lookup(_cookie(a1)) is None and again.lookup(_cookie(r1)) is not None
    a2, _ = again.create(username="alice", password="alice-secret-2")
    users.set_role("alice", "admin")
    assert again.lookup(_cookie(a2)) is None
    a3, row = again.create(username="alice", password="alice-secret-2")
    assert row["role"] == "admin" and row["credential_generation"] == 2
    users.disable("alice")
    assert again.lookup(_cookie(a3)) is None
    assert users.credential("alice") == {"generation": 3, "role": "admin", "disabled": True}


def test_rows_written_before_v014_keep_working_until_the_credential_changes(tmp_path):
    users = _users(tmp_path)
    now = time.time()
    old = {"owner": "alice", "csrf_token": "c", "expires_at": now + 3600, "shared": True, "username": "alice", "role": "user",
           "display_name": "alice", "claimable_owners": []}
    legacy = {"owner": "anon", "csrf_token": "c", "expires_at": now + 3600, "shared": True, "username": None, "role": "legacy"}
    (tmp_path / "sessions.json").write_text(json.dumps({"old": old, "legacy": legacy}), encoding="utf-8")
    store = SessionStore(tmp_path / "sessions.json", shared=True, token="t0ken", users=users)
    assert store.lookup(_cookie("old"))[1]["username"] == "alice"       # no generation recorded = generation 0
    assert store.lookup(_cookie("legacy"), for_claim=True) is None      # legacy rows without a token fingerprint are dropped
    users.set_password("alice", "alice-secret-2")
    assert store.lookup(_cookie("old")) is None


def test_idle_timeout_last_seen_throttle_and_runtime_pruning(tmp_path, monkeypatch):
    monkeypatch.setenv(S.SESSION_IDLE_ENV, "1")
    clock = Clock()
    store = SessionStore(tmp_path / "sessions.json", shared=True, token="t0ken", clock=clock)
    writes = []
    real = store._persist
    monkeypatch.setattr(store, "_persist", lambda: (writes.append(clock.now), real())[1])
    sid, _ = store.create("t0ken")
    assert len(writes) == 1
    for _ in range(50):                     # many requests within one minute: no extra write
        clock.now += 1
        assert store.lookup(_cookie(sid))
    assert len(writes) == 1
    clock.now += 61
    assert store.lookup(_cookie(sid)) and len(writes) == 2
    clock.now += 3599                       # just inside the 1 h idle window
    assert store.lookup(_cookie(sid))
    clock.now += 3601
    assert store.lookup(_cookie(sid)) is None and sid not in store.sessions
    # absolute lifetime still applies to a busy session
    sid, _ = store.create("t0ken")
    for _ in range(210):                    # 210 × 3000 s > 7 days, never idle for an hour
        clock.now += 3000
        store.lookup(_cookie(sid))
    assert sid not in store.sessions
    # expired rows are swept at runtime, not only at start-up
    dead = [store.create("t0ken")[0] for _ in range(5)]
    clock.now += 2 * 3600
    live, _ = store.create("t0ken")
    for _ in range(S.SESSION_PRUNE_EVERY):
        store.lookup(_cookie(live))
    assert not any(key in store.sessions for key in dead) and live in store.sessions
    assert set(json.loads((tmp_path / "sessions.json").read_text(encoding="utf-8"))) == {live}


def test_loopback_cookieless_requests_do_not_write_sessions(tmp_path):
    service = Service(tmp_path)
    path = tmp_path / "sessions.json"
    api = service.api()
    try:
        for _ in range(5):   # S7: the page that opens a local session is the workspace (``/`` redirects there)
            status, _, response = call(api, "GET", "/v2/", {})
            assert status == 200 and response.getheader("Set-Cookie")
        assert not path.exists() or json.loads(path.read_text(encoding="utf-8")) == {}
        assert len(service.sessions.sessions) == 5
        api.request("GET", "/api/session"); response = api.getresponse(); response.read()
        cookie = response.getheader("Set-Cookie").split(";", 1)[0]
        status, _, _ = call(api, "GET", "/api/jobs", {"Cookie": cookie})      # the browser came back: now persisted
        assert status == 200 and list(json.loads(path.read_text(encoding="utf-8"))) == [cookie.split("=", 1)[1]]
    finally:
        api.close(); service.close()


def test_api_password_change_renews_this_session_and_ends_the_others(tmp_path):
    users = _users(tmp_path)
    service = Service(tmp_path, shared=True, users=users)
    api = service.api()
    try:
        _, _, other_cookie, _ = _login(api, {"username": "alice", "password": "alice-secret"})
        status, payload, cookie, headers = _login(api, {"username": "alice", "password": "alice-secret"})
        status, body, response = call(api, "POST", "/api/session/password", headers, {"old_password": "alice-secret", "new_password": "alice-secret-2"})
        assert status == 200 and body["changed"] is True
        renewed = response.getheader("Set-Cookie").split(";", 1)[0]
        assert renewed != cookie
        assert call(api, "GET", "/api/jobs", {"Cookie": renewed})[0] == 200
        assert call(api, "GET", "/api/jobs", {"Cookie": cookie})[0] == 401
        assert call(api, "GET", "/api/jobs", {"Cookie": other_cookie})[0] == 401
        # the renewed session keeps the page's CSRF token
        status, _, _ = call(api, "PUT", "/api/preferences", {**headers, "Cookie": renewed}, {"a": 1})
        assert status == 200
    finally:
        api.close(); service.close()


# ----------------------------------------------------------------------------------------------- S2
def test_named_previous_session_is_never_offered_for_claim(tmp_path):
    users = _users(tmp_path)
    service = Service(tmp_path, shared=True, users=users)
    api = service.api()
    try:
        _, _, root_cookie, _ = _login(api, {"username": "root", "password": "root-secret"})
        finished(service.manager, owner="root", case_id="ROOTS")
        # alice logs in on root's browser: root's jobs must not become claimable
        status, payload, _, headers = _login(api, {"username": "alice", "password": "alice-secret"}, cookie=root_cookie)
        assert status == 200 and payload["claimable_owners"] == [] and "claimable" not in payload
        status, body, _ = call(api, "POST", "/api/jobs/claim", headers, {"owner": "root"})
        assert status == 403
        # not even an administrator moves a registered user's jobs through the API
        _, _, _, admin_headers = _login(api, {"username": "root", "password": "root-secret"})
        finished(service.manager, owner="alice", case_id="ALICES")
        status, body, _ = call(api, "POST", "/api/jobs/claim", admin_headers, {"owner": "alice"})
        assert status == 403 and "命令行" in body["error"]["message"]
        assert [job["case_id"] for job in service.manager.list("alice")] == ["ALICES"]
    finally:
        api.close(); service.close()


# ----------------------------------------------------------------------------------------------- S3
def test_admin_trash_view_is_read_only_and_mutations_are_audited(tmp_path, caplog):
    users = _users(tmp_path)
    service = Service(tmp_path, shared=True, users=users)
    api = service.api()
    try:
        _, _, _, alice = _login(api, {"username": "alice", "password": "alice-secret"})
        _, _, _, root = _login(api, {"username": "root", "password": "root-secret"})
        job = finished(service.manager, owner="alice", case_id="A")
        other = finished(service.manager, owner="alice", case_id="B")
        with caplog.at_level(logging.INFO, logger="wss_deploy.audit"):
            status, _, _ = call(api, "POST", f"/api/jobs/{job['id']}/delete", alice, {"version": job["version"]})
            assert status == 200
            status, _, _ = call(api, "POST", "/api/jobs/delete", alice, {"jobs": [{"id": other["id"], "version": other["version"]}]})
            assert status == 200
            status, listing, _ = call(api, "GET", "/api/trash?all=1", {"Cookie": root["Cookie"]})
            assert status == 200 and {item["id"] for item in listing["items"]} == {job["id"], other["id"]}
            assert call(api, "POST", f"/api/trash/{job['id']}/restore", root, {})[0] == 404
            assert call(api, "POST", f"/api/trash/{job['id']}/purge", root, {})[0] == 404
            assert call(api, "POST", f"/api/trash/{job['id']}/restore", alice, {})[0] == 200
            assert call(api, "POST", f"/api/trash/{other['id']}/purge", alice, {})[0] == 200
        lines = [json.loads(record.getMessage()[5:]) for record in caplog.records
                 if record.name == "wss_deploy.audit" and record.getMessage().startswith("http ")]
        assert {(line["action"], line["actor"], tuple(line["jobs"])) for line in lines} >= {
            ("delete", "alice", (job["id"],)), ("delete", "alice", (other["id"],)), ("restore", "alice", (job["id"],)), ("purge", "alice", (other["id"],))}
        assert all(line["ip"] == "127.0.0.1" for line in lines) and not any(line["actor"] == "root" for line in lines)
    finally:
        api.close(); service.close()


# ----------------------------------------------------------------------------------------------- S4
def test_exports_neutralise_formulas():
    rows = [{column: None for column in COLUMNS} | {"case_id": "=1+1", "patient_id": "@SUM(A1)", "scan_label": "\tx", "tags": "+cmd",
                                                     "reviewer": "-2", "wss_p99_pa": -1.5, "job_id": "plain"}]
    line = to_csv(rows).decode("utf-8-sig").splitlines()[1].split(",")
    assert line[COLUMNS.index("case_id")] == "'=1+1" and line[COLUMNS.index("patient_id")] == "'@SUM(A1)"
    assert line[COLUMNS.index("tags")] == "'+cmd" and line[COLUMNS.index("reviewer")] == "'-2"
    assert line[COLUMNS.index("wss_p99_pa")] == "-1.5" and line[COLUMNS.index("job_id")] == "plain"
    from openpyxl import load_workbook
    data = to_xlsx(rows)
    assert b"<f>" not in zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml")
    cell = load_workbook(io.BytesIO(data)).active.cell(row=2, column=COLUMNS.index("case_id") + 1)
    assert cell.value == "=1+1" and cell.data_type == "s"


# ----------------------------------------------------------------------------------------------- S5
def _parts(boundary: bytes, chunk: int, body: bytes):
    parts, current = [], {}
    def start(name, filename):
        current = {"name": name, "filename": filename, "data": bytearray()}; parts.append(current)
        return current["data"].extend
    stream = io.BytesIO(body)
    stream_multipart(lambda n: stream.read(min(n, chunk)), len(body), boundary, start)
    return [(p["name"], p["filename"], bytes(p["data"])) for p in parts]


@pytest.mark.parametrize("chunk", [1, 3, 7, 64, 1 << 20])
def test_streaming_multipart_parser_handles_split_delimiters(chunk):
    tricky = b"\r\n--bn" + b"\r\n--bnd" * 3 + b"\r\n-" + bytes(range(256)) * 4
    body, _ = multipart("bndx", [("stl", tricky, "病例 A.stl"), ("notes", "中文 备注", None), ("empty", b"", None)])
    parts = _parts(b"bndx", chunk, b"preamble\r\n" + body + b"epilogue")
    assert parts == [("stl", "病例 A.stl", tricky), ("notes", None, "中文 备注".encode()), ("empty", None, b"")]
    rfc = b'--bndx\r\nContent-Disposition: form-data; name="stl"; filename*=UTF-8\'\'%E7%97%85.stl\r\n\r\nX\r\n--bndx--'
    assert _parts(b"bndx", chunk, rfc) == [("stl", "病.stl", b"X")]
    with pytest.raises(JobError):
        _parts(b"bndx", chunk, b"--bndx\r\nContent-Disposition: form-data; name=\"a\"\r\n\r\nno closing delimiter")


def _upload(api, headers, parts, path="/api/jobs"):
    body, ctype = multipart("sec-boundary", parts)
    api.request("POST", path, body=body, headers={**headers, "Content-Type": ctype})
    response = api.getresponse(); raw = response.read()
    return response.status, json.loads(raw)


def test_upload_streams_to_disk_with_limits(tmp_path, monkeypatch):
    service = Service(tmp_path)
    api = service.api(timeout=30)
    try:
        cookie, headers, owner = service.session(api)
        headers = {k: v for k, v in headers.items() if k != "Content-Type"}
        # the manager receives open file objects (content_file), never the upload as bytes
        sources = []
        real_create = service.manager.create
        def spy(*args, **kwargs):
            source = kwargs.get("content_file")
            sources.append((kwargs.get("content"), kwargs.get("content_path"), type(source).__name__,
                            getattr(source, "_rolled", None), source is not None and not source.closed))
            return real_create(*args, **kwargs)
        monkeypatch.setattr(service.manager, "create", spy)
        stl = _stl(200_000)                                                # 10 MB: spills past the 8 MiB spool
        status, payload = _upload(api, headers, [("stl", stl, "big.stl"), ("device", "cpu", None), ("notes", "备注" * 500, None)])
        assert status == 201, payload
        assert sources == [(None, None, "SpooledTemporaryFile", True, True)]
        assert (service.manager.root / payload["job_id"] / "input.stl").read_bytes() == stl
        assert payload["job"]["notes"] == "备注" * 500 and payload["job"]["source_filename"] == "big.stl"
        assert payload["job"]["input_sha256"] == hashlib.sha256(stl).hexdigest()
        assert not any(service.server.tmp_dir.iterdir())                  # spooled and staged files are gone
        # small (in-memory) spools stream the same way; batch items too, and client metadata can never name a
        # server-side file as the STL source
        secret = tmp_path / "server_side.stl"; secret.write_bytes(b"server secret")
        sources.clear()
        status, payload = _upload(api, headers, [("stl", b"solid one", "one.stl"), ("stl", b"solid two", "two.stl"),
                                                 ("metadata_json", json.dumps([{"content_path": str(secret), "content_file": None,
                                                                                "content": "x", "patient_id": "P1"}, {}]), None)],
                                  path="/api/jobs/batch")
        assert status == 201 and payload["created_count"] == 2, payload
        assert [(s[0], s[1], s[2], s[4]) for s in sources] == [(None, None, "SpooledTemporaryFile", True)] * 2
        created = {job["source_filename"]: job for job in payload["jobs"]}
        assert (service.manager.root / created["one.stl"]["id"] / "input.stl").read_bytes() == b"solid one"
        assert created["one.stl"]["patient_id"] == "P1"
        assert (service.manager.root / created["two.stl"]["id"] / "input.stl").read_bytes() == b"solid two"
        monkeypatch.setattr(service.manager, "create", real_create)
        # per-STL size limit while streaming
        monkeypatch.setattr(S, "MAX_UPLOAD_BYTES", 1000)
        status, payload = _upload(api, headers, [("stl", b"x" * 5000, "a.stl")])
        assert status == 413
        monkeypatch.setattr(S, "MAX_UPLOAD_BYTES", 128 * 1024 * 1024)
        # field limit
        status, payload = _upload(api, headers, [("stl", b"solid", "a.stl"), ("case_id", "c" * 2000, None)])
        assert status == 400 and "过长" in payload["error"]["message"]
        # concurrency slots
        slots = service.server.upload_slots
        slots.acquire(); slots.acquire()
        try:
            status, payload = _upload(api, headers, [("stl", b"solid", "a.stl")])
            assert status == 429 and "上传" in payload["error"]["message"]
        finally:
            slots.release(); slots.release()
        # per-owner cap on queued + running jobs (the test manager never starts its worker: jobs stay queued)
        counts = service.manager.queue_counts(owner)
        cap = service.server.max_active_jobs = counts["queued"] + counts["running"] + 1
        status, payload = _upload(api, headers, [("stl", b"solid four", "b.stl")])
        assert status == 201
        status, payload = _upload(api, headers, [("stl", b"solid three", "c.stl")])
        assert status == 429 and payload["max_active_jobs"] == cap
        status, payload = _upload(api, headers, [("stl", b"x1", "x1.stl"), ("stl", b"x2", "x2.stl")], path="/api/jobs/batch")
        assert status == 429
    finally:
        api.close(); service.close()


def test_login_throttle_and_failure_table_bound(tmp_path, caplog):
    users = _users(tmp_path)
    service = Service(tmp_path, shared=True, token="t0ken-value", users=users)
    service.server.login_throttle = LoginThrottle(3)
    api = service.api()
    try:
        with caplog.at_level(logging.INFO, logger="wss_deploy.audit"):
            assert _login(api, {"username": "alice", "password": "alice-secret"})[0] == 200      # success refunds
            assert _login(api, {"username": "alice", "password": "wrong-secret"})[0] == 401
            assert _login(api, {"username": "Secret-Typed-Here", "password": "whatever-1"})[0] == 401
        audit = [json.loads(r.getMessage()[5:]) for r in caplog.records if r.name == "wss_deploy.audit" and r.getMessage().startswith("http ")]
        assert [(a["action"], a.get("actor") or a.get("username")) for a in audit] == [("login", "alice"), ("login_failed", "alice"), ("login_failed", "<unknown>")]
        assert "Secret-Typed-Here" not in caplog.text and "wrong-secret" not in caplog.text
        assert _login(api, {"username": "nobody", "password": "whatever-1"})[0] == 401
        status, payload, _, _ = _login(api, {"username": "alice", "password": "alice-secret"})
        assert status == 429 and "频繁" in payload["error"]["message"]
        status, _, _, _ = _login(api, {"token": "t0ken-value"})
        assert status == 429                                                                 # token logins share the budget
    finally:
        api.close(); service.close()
    bucket = LoginThrottle(2, clock=Clock(0.0))
    assert bucket.take("a") and bucket.take("a") and not bucket.take("a") and bucket.take("b")
    bucket.clock.now += 30
    assert bucket.take("a") and not bucket.take("a")
    store = UserStore(tmp_path / "users.json")
    for index in range(MAX_TRACKED_FAILURES + 50):
        store._record_failure(f"user{index}")
    assert len(store._failures) == MAX_TRACKED_FAILURES and "user0" not in store._failures
    with pytest.raises(JobError):
        store.verify("never-seen", "whatever-1")
    assert list(store._failures)[-1] == "never-seen" and len(store._failures) == MAX_TRACKED_FAILURES


def test_unknown_user_costs_one_scrypt(tmp_path):
    users = _users(tmp_path)
    users.verify("alice", "alice-secret")          # warm the dummy hash once
    def timed(name):
        start = time.perf_counter()
        with pytest.raises(JobError):
            users.verify(name, "wrong-password")
        return time.perf_counter() - start
    known = min(timed("alice") for _ in range(2)); users._failures.clear()
    unknown = min(timed("ghost-user") for _ in range(2))
    assert unknown > 0.4 * known


def test_sse_streams_per_session_are_capped(tmp_path):
    service = Service(tmp_path)
    service.server.max_sse_per_session = 1
    try:
        api = service.api(timeout=10)
        cookie, _, _ = service.session(api); api.close()
        first = http.client.HTTPConnection("127.0.0.1", service.port, timeout=10)
        first.request("GET", "/api/events", headers={"Cookie": cookie})
        response = first.getresponse(); assert response.status == 200
        assert response.fp.readline().startswith(b"event: hello")
        second = http.client.HTTPConnection("127.0.0.1", service.port, timeout=10)
        second.request("GET", "/api/events", headers={"Cookie": cookie})
        other = second.getresponse(); assert other.status == 200
        started = time.time()
        rest = response.read()                    # the older stream is closed at once, not after a keepalive
        assert time.time() - started < 5 and b"event:" not in rest.split(b"\n\n", 1)[-1]
        first.close(); second.close()
    finally:
        service.close()


def test_bundles_are_built_on_disk_with_a_cap_and_one_at_a_time(tmp_path, monkeypatch):
    service = Service(tmp_path)
    api = service.api(timeout=30)
    try:
        cookie, headers, owner = service.session(api)
        jobs = [finished(service.manager, owner=owner, case_id=f"C{i}") for i in range(2)]
        monkeypatch.setattr("wss_deploy.bundle.job_bundle", lambda *a, **k: pytest.fail("in-memory bundle used"))
        status, data, response = call(api, "GET", f"/api/jobs/{jobs[0]['id']}/bundle.zip", {"Cookie": cookie})
        assert status == 200 and int(response.getheader("Content-Length")) == len(data)
        assert "summary.json" in zipfile.ZipFile(io.BytesIO(data)).namelist()
        status, data, response = call(api, "POST", "/api/jobs/bundle", headers, {"ids": [j["id"] for j in jobs]})
        assert status == 200 and len(zipfile.ZipFile(io.BytesIO(data)).namelist()) == 2
        assert not any(service.server.tmp_dir.iterdir())
        service.server.max_bundle_bytes = 100
        status, payload, _ = call(api, "POST", "/api/jobs/bundle", headers, {"ids": [j["id"] for j in jobs]})
        assert status == 413 and "GiB" in payload["error"]["message"]
        service.server.max_bundle_bytes = S.DEFAULT_MAX_BUNDLE_BYTES
        monkeypatch.setattr(S, "BUNDLE_WAIT_SECONDS", 0.1)
        service.server.bundle_slots.acquire()
        try:
            status, payload, _ = call(api, "GET", f"/api/jobs/{jobs[0]['id']}/bundle.zip", {"Cookie": cookie})
            assert status == 429
        finally:
            service.server.bundle_slots.release()
        assert not any(service.server.tmp_dir.iterdir())
    finally:
        api.close(); service.close()


# ----------------------------------------------------------------------------------------------- S6 / S10
def test_headers_csp_gzip_and_etags(tmp_path, monkeypatch):
    monkeypatch.setenv(F.ENV_SWITCH, "0")
    service = Service(tmp_path)
    api = service.api(timeout=30)
    try:
        cookie, headers, owner = service.session(api)
        job = finished(service.manager, owner=owner, case_id="G")
        report = service.manager.root / job["id"] / "report.html"
        page = "<html><body>" + "<p>WSS 报告</p>" * 4000 + "</body></html>"
        report.write_text(page, encoding="utf-8")
        # static workbench file: no inline script, gzip, ETag + revalidation (S7: the classic app.js is retired; a
        # workspace script without ``?v=`` is served the same way)
        api.request("GET", "/static/v2/ws_shell.js", headers={"Accept-Encoding": "gzip, deflate"})
        response = api.getresponse(); body = response.read()
        assert response.status == 200 and response.getheader("Content-Encoding") == "gzip" and response.getheader("Server") is None
        assert gzip.decompress(body) == (S.V2_DIR / "ws_shell.js").read_bytes()
        assert response.getheader("Cache-Control") == S.STATIC_CACHE and response.getheader("Vary") == "Accept-Encoding"
        etag = response.getheader("ETag")
        csp = response.getheader("Content-Security-Policy")
        assert "script-src 'self';" in csp and "unsafe-inline'; style" not in csp
        api.request("GET", "/static/v2/ws_shell.js", headers={"If-None-Match": etag})
        response = api.getresponse(); assert response.status == 304 and response.read() == b""
        status, index, response = call(api, "GET", "/v2/", {"Cookie": cookie})
        assert "script-src 'self';" in response.getheader("Content-Security-Policy") and "frame-ancestors 'none'" in response.getheader("Content-Security-Policy")
        # report: gzip, private no-cache + ETag, inline scripts allowed, embeddable by our origin
        api.request("GET", f"/api/jobs/{job['id']}/report", headers={"Cookie": cookie, "Accept-Encoding": "gzip"})
        response = api.getresponse(); body = response.read()
        assert gzip.decompress(body).decode("utf-8") == page and response.getheader("Cache-Control") == S.REPORT_CACHE
        assert "'unsafe-inline'" in response.getheader("Content-Security-Policy") and "frame-ancestors 'self'" in response.getheader("Content-Security-Policy")
        etag = response.getheader("ETag")
        api.request("GET", f"/api/jobs/{job['id']}/report", headers={"Cookie": cookie, "If-None-Match": etag})
        response = api.getresponse(); assert response.status == 304 and response.read() == b""
        report.write_text(page + "<!-- reviewed -->", encoding="utf-8")
        api.request("GET", f"/api/jobs/{job['id']}/report", headers={"Cookie": cookie, "If-None-Match": etag})
        response = api.getresponse(); assert response.status == 200 and response.read().endswith(b"<!-- reviewed -->")
        # without Accept-Encoding the file is streamed as is
        status, raw, response = call(api, "GET", f"/api/jobs/{job['id']}/report", {"Cookie": cookie})
        assert response.getheader("Content-Encoding") is None and int(response.getheader("Content-Length")) == len(report.read_bytes())
        # one-page report: only the print handler may run
        status, _, response = call(api, "GET", f"/api/jobs/{job['id']}/onepage", {"Cookie": cookie})
        csp = response.getheader("Content-Security-Policy")
        assert status == 200 and "script-src 'unsafe-hashes' " + S.ONEPAGE_HANDLER_HASH + ";" in csp
        etag = response.getheader("ETag")
        assert call(api, "GET", f"/api/jobs/{job['id']}/onepage", {"Cookie": cookie, "If-None-Match": etag})[0] == 304
        # API JSON: no-store, gzip once large
        for index in range(25):
            service.manager.create(owner, content=b"solid %d" % index, filename=f"q{index}.stl", notes="x" * 200)
        api.request("GET", "/api/jobs", headers={"Cookie": cookie, "Accept-Encoding": "gzip"})
        response = api.getresponse(); body = response.read()
        assert response.getheader("Cache-Control") == "no-store" and response.getheader("Content-Encoding") == "gzip"
        assert len(json.loads(gzip.decompress(body))["jobs"]) == 26
        # the geometry ETag keeps its format and now revalidates
        status, geometry, response = call(api, "GET", f"/api/jobs/{job['id']}/geometry", {"Cookie": cookie})
        etag = response.getheader("ETag")
        assert status == 200 and etag.startswith('W/"' + job["id"]) and response.getheader("Cache-Control") == S.REPORT_CACHE
        assert call(api, "GET", f"/api/jobs/{job['id']}/geometry", {"Cookie": cookie, "If-None-Match": etag})[0] == 304
    finally:
        api.close(); service.close()


def test_accept_encoding_and_etag_parsing():
    assert S.accepts_gzip("gzip") and S.accepts_gzip("br, gzip;q=0.5") and S.accepts_gzip("*")
    assert not S.accepts_gzip("gzip;q=0") and not S.accepts_gzip("identity") and not S.accepts_gzip(None)
    assert not S.accepts_gzip("*, gzip;q=0")
    assert S.etag_matches('W/"a", W/"b"', 'W/"b"') and S.etag_matches('"b"', 'W/"b"') and S.etag_matches("*", 'W/"x"')
    assert not S.etag_matches('W/"c"', 'W/"b"') and not S.etag_matches(None, 'W/"b"')


# ----------------------------------------------------------------------------------------------- S7
def test_access_log_has_paths_only_with_patient_ids_redacted(tmp_path, caplog):
    service = Service(tmp_path)
    api = service.api()
    try:
        cookie, _, _ = service.session(api)
        access = logging.getLogger("wss_deploy.access")
        access.addHandler(caplog.handler)          # cli ``serve`` may have turned propagation off in this process
        try:
            with caplog.at_level(logging.INFO, logger="wss_deploy.access"):
                call(api, "GET", "/api/patients/ZHANG%20SAN/timeline?all=1&secret=PHI", {"Cookie": cookie})
                call(api, "GET", "/api/jobs?q=ZHANG", {"Cookie": cookie})
                # Each line is written after its response has been sent: wait for the /api/jobs line itself
                # (counting lines races with the /api/session line of service.session()).
                def seen():
                    lines = [r.getMessage() for r in caplog.records if r.name == "wss_deploy.access"]
                    return any(" /api/jobs " in line for line in lines) and any("/api/patients/" in line for line in lines)
                deadline = time.time() + 10
                while time.time() < deadline and not seen():
                    time.sleep(0.05)
        finally:
            access.removeHandler(caplog.handler)
        lines = list(dict.fromkeys(record.getMessage() for record in caplog.records if record.name == "wss_deploy.access"))
        assert any("/api/patients/<id>/timeline" in line for line in lines)
        assert any(" /api/jobs 200" in line for line in lines)
        assert not any("ZHANG" in line or "secret" in line or "?" in line for line in lines)
        assert not any(record.name == "wss_deploy.service" and "/api/jobs" in record.getMessage() for record in caplog.records)
    finally:
        api.close(); service.close()
    assert S.access_path("/api/patients/abc/timeline?x=1") == "/api/patients/<id>/timeline"
    assert "PHI" not in S.redact_log_text("Bad request ('GET /api/patients/PHI/timeline?q=PHI HTTP/1.1')")


# ----------------------------------------------------------------------------------------------- S8
def _stale_job(tmp_path) -> Path:
    job_dir = tmp_path / "20260101_000000_abcdef"
    job_dir.mkdir()
    (job_dir / "report.html").write_text("<html>old template</html>", encoding="utf-8")
    (job_dir / "run_manifest.json").write_text(json.dumps({"outputs": {"report.html": {"path": "report.html"}}}), encoding="utf-8")
    return job_dir


def test_ensure_fresh_renders_outside_the_manager_lock(tmp_path, monkeypatch):
    monkeypatch.delenv(F.ENV_SWITCH, raising=False)
    monkeypatch.setattr(F, "code_changed", lambda: False)
    F._failed.clear()
    job_dir = _stale_job(tmp_path)
    guard = threading.Lock()
    seen = []

    def refresher(stage):
        seen.append((Path(stage).name, Path(stage) != job_dir, guard.acquire(blocking=False)))
        guard.release()
        page = Path(stage) / "report.html"
        page.write_text(page.read_text(encoding="utf-8").replace("old", "new"), encoding="utf-8")
        return {}
    assert F.ensure_fresh(job_dir, guard=guard, refresher=refresher) == "refreshed"
    assert seen == [(job_dir.name, True, True)]
    assert (job_dir / "report.html").read_text(encoding="utf-8") == "<html>new template</html>"
    from wss_deploy.io_utils import file_sha256
    assert json.loads((job_dir / "run_manifest.json").read_text())["outputs"]["report.html"]["sha256"] == file_sha256(job_dir / "report.html")
    assert not F.is_stale(job_dir) and not list(job_dir.glob(F.STAGE_PREFIX + "*"))


def test_ensure_fresh_discards_a_render_when_the_report_changed_meanwhile(tmp_path, monkeypatch):
    monkeypatch.delenv(F.ENV_SWITCH, raising=False)
    monkeypatch.setattr(F, "code_changed", lambda: False)
    F._failed.clear()
    job_dir = _stale_job(tmp_path)
    calls = []

    def refresher(stage):
        calls.append(1)
        if len(calls) == 1:     # a review decision rewrites the embedded metadata while we render
            (job_dir / "report.html").write_text("<html>old template reviewed</html>", encoding="utf-8")
        page = Path(stage) / "report.html"
        page.write_text(page.read_text(encoding="utf-8").replace("old", "new"), encoding="utf-8")
        return {}
    assert F.ensure_fresh(job_dir, guard=threading.Lock(), refresher=refresher) == "refreshed"
    assert len(calls) == 2 and (job_dir / "report.html").read_text(encoding="utf-8") == "<html>new template reviewed</html>"
    assert not list(job_dir.glob(F.STAGE_PREFIX + "*"))


# ----------------------------------------------------------------------------------------------- S9
def test_admin_detail_is_for_administrators_only(tmp_path):
    users = _users(tmp_path)
    service = Service(tmp_path, shared=True, users=users)
    api = service.api()
    try:
        _, _, alice_cookie, _ = _login(api, {"username": "alice", "password": "alice-secret"})
        _, _, root_cookie, _ = _login(api, {"username": "root", "password": "root-secret"})
        job = service.manager.create("alice", content=b"solid", filename="a.stl")
        with service.manager.lock:
            service.manager.jobs[job["id"]].update(status="failed", error={
                "message": "中心线提取失败。", "diagnostic_id": "abc", "category": "toolchain", "retryable": False,
                "admin_detail": "vmtk stderr: /public/secret/path"})
        status, body, _ = call(api, "GET", f"/api/jobs/{job['id']}", {"Cookie": alice_cookie})
        assert status == 200 and body["error"] == {"message": "中心线提取失败。", "diagnostic_id": "abc", "category": "toolchain", "retryable": False}
        status, listing, _ = call(api, "GET", "/api/jobs", {"Cookie": alice_cookie})
        assert "admin_detail" not in json.dumps(listing)
        status, body, _ = call(api, "GET", f"/api/jobs/{job['id']}", {"Cookie": root_cookie})
        assert body["error"]["admin_detail"].startswith("vmtk stderr")
    finally:
        api.close(); service.close()


def _sse_events(port, path, cookie, *, count=None, timeout=10) -> list[tuple[str, dict]]:
    """Read server-sent events until the stream ends (or ``count`` events arrived)."""
    connection = http.client.HTTPConnection("127.0.0.1", port, timeout=timeout)
    connection.request("GET", path, headers={"Cookie": cookie})
    response = connection.getresponse()
    events, name = [], None
    try:
        while count is None or len(events) < count:
            line = response.fp.readline()
            if not line:
                break
            line = line.decode("utf-8").rstrip("\n")
            if line.startswith("event: "):
                name = line[7:]
            elif line.startswith("data: "):
                events.append((name, json.loads(line[6:])))
    finally:
        connection.close()
    return events


def test_unknown_exception_admin_detail_is_stripped_in_json_and_sse(tmp_path):
    """jobs.py records an unexpected exception as the generic message + ``admin_detail: "RuntimeError: …"``."""
    from tests._c_helpers import manager as make_manager
    def broken_stage_a(*_args, **_kwargs):
        raise RuntimeError("worker exploded at /public/newhome/secret/path")
    mgr = make_manager(tmp_path / "jobs", stage_a=broken_stage_a)
    users = _users(tmp_path)
    service = Service(tmp_path, shared=True, users=users, mgr=mgr)
    mgr.start()
    api = service.api()
    try:
        _, _, alice_cookie, _ = _login(api, {"username": "alice", "password": "alice-secret"})
        _, _, root_cookie, _ = _login(api, {"username": "root", "password": "root-secret"})
        job = mgr.create("alice", content=b"solid", filename="a.stl")
        deadline = time.time() + 20
        while time.time() < deadline and mgr.get(job["id"], "alice")["status"] != "failed":
            time.sleep(0.05)
        record = mgr.get(job["id"], "alice")["error"]
        assert record["admin_detail"].startswith("RuntimeError: worker exploded") and record["retryable"] is True
        for path in (f"/api/jobs/{job['id']}", "/api/jobs", "/api/jobs?status=failed&page_size=5", "/api/cases"):
            status, body, _ = call(api, "GET", path, {"Cookie": alice_cookie})
            assert status == 200 and "admin_detail" not in json.dumps(body) and "RuntimeError" not in json.dumps(body), path
        status, body, _ = call(api, "GET", f"/api/jobs/{job['id']}", {"Cookie": alice_cookie})
        assert body["error"]["message"] and body["error"]["diagnostic_id"] == record["diagnostic_id"]
        # SSE: the snapshot of a failed job, as the owner (non-admin) and as the administrator
        events = _sse_events(service.port, f"/api/jobs/{job['id']}/events", alice_cookie)
        assert events[0][0] == "snapshot" and events[0][1]["error"]["diagnostic_id"] == record["diagnostic_id"]
        assert "admin_detail" not in json.dumps(events) and "RuntimeError" not in json.dumps(events)
        events = _sse_events(service.port, f"/api/jobs/{job['id']}/events", root_cookie)
        assert events[0][1]["error"]["admin_detail"].startswith("RuntimeError")
        # a live event carrying a record (owner stream) is stripped the same way
        results = []
        reader = threading.Thread(target=lambda: results.append(_sse_events(service.port, "/api/events", alice_cookie, count=2)))
        reader.start()
        deadline = time.time() + 10
        while time.time() < deadline and not mgr._owner_subscribers.get("alice"):
            time.sleep(0.02)
        mgr._offer(mgr._owner_subscribers.get("alice"), {"job_id": job["id"], "action": "failed", "final": True, "error": dict(record)})
        reader.join(timeout=15)
        assert results and results[0][1][0] == "job" and results[0][1][1]["error"]["diagnostic_id"] == record["diagnostic_id"]
        assert "admin_detail" not in json.dumps(results) and "RuntimeError" not in json.dumps(results)
    finally:
        api.close(); service.close()


def test_release_errors_reach_the_client_without_server_paths(tmp_path, monkeypatch):
    service = Service(tmp_path)
    api = service.api()
    try:
        cookie, headers, _ = service.session(api)
        def boom(*_a, **_k):
            raise ReleaseError("发布根目录没有可用 release.json：/public/newhome/cy/Digital_twin/GNN/outputs/releases")
        monkeypatch.setattr(service.manager, "create", boom)
        body, ctype = multipart("b", [("stl", b"solid", "a.stl")])
        api.request("POST", "/api/jobs", body=body, headers={"Cookie": cookie, "X-CSRF-Token": headers["X-CSRF-Token"], "Content-Type": ctype})
        response = api.getresponse(); payload = json.loads(response.read())
        assert response.status == 422 and "/public" not in payload["error"]["message"] and payload["error"]["message"].endswith("：releases")
    finally:
        api.close(); service.close()


# ----------------------------------------------------------------------------------------------- S11
def test_health_merges_manager_checks_and_ready_follows_ok(tmp_path):
    service = Service(tmp_path)
    state = {"ok": True}
    service.manager.health = lambda: {"ok": state["ok"], "checks": {"disk": {"ok": state["ok"], "free_gb": 1.5}}, "cached_at": "2026-09-24T00:00:00+08:00"}
    api = service.api()
    try:
        status, payload, _ = call(api, "GET", "/api/health", {})
        assert status == 200 and payload == {"ok": True, "version": S.__version__}
        assert call(api, "GET", "/api/ready", {})[0] == 200
        state["ok"] = False
        status, payload, _ = call(api, "GET", "/api/health", {})
        assert status == 200 and payload == {"ok": False, "version": S.__version__}
        status, payload, _ = call(api, "GET", "/api/ready", {})
        assert status == 503 and payload == {"ok": False, "version": S.__version__}
        cookie, _, _ = service.session(api)
        status, payload, _ = call(api, "GET", "/api/ready", {"Cookie": cookie})
        assert status == 503 and payload["checks"]["disk"]["free_gb"] == 1.5 and payload["cached_at"] and "queue" in payload
        def broken():
            raise RuntimeError("probe")
        service.manager.health = broken
        status, payload, _ = call(api, "GET", "/api/ready", {})
        assert status == 503 and payload["ok"] is False
    finally:
        api.close(); service.close()


def test_ensure_fresh_skips_the_renderer_for_pages_without_embedded_arrays(tmp_path, monkeypatch):
    """A page without embedded arrays can never be refreshed: fail fast, without importing / calling the renderer."""
    monkeypatch.delenv(F.ENV_SWITCH, raising=False)
    monkeypatch.setattr(F, "code_changed", lambda: False)
    F._failed.clear()
    from wss_deploy import rebuild_report as RB
    monkeypatch.setattr(RB, "refresh_ui_only", lambda *_a, **_k: pytest.fail("renderer called"))
    job_dir = _stale_job(tmp_path)
    before = (job_dir / "report.html").read_bytes()
    assert F.ensure_fresh(job_dir, guard=threading.Lock()) == "failed"
    assert (job_dir / "report.html").read_bytes() == before and not list(job_dir.glob(F.STAGE_PREFIX + "*"))
    assert F.ensure_fresh(job_dir, guard=threading.Lock()) == "failed"      # not retried for this template version
    F._failed.clear()
