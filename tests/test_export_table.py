"""C14: the summary table has fixed columns, blanks for missing values, and the route is owner-scoped."""
from __future__ import annotations

import csv
import io
import json

from wss_deploy.export_table import COLUMNS, MORPHOLOGY_COLUMNS, build_row, to_csv, to_xlsx
from tests._c_helpers import Service, call, finished


def test_columns_and_rows_for_both_families():
    wall_job = {"id": "j1", "case_id": "A", "patient_id": "P", "tags": ["AAA", "随访"], "status": "done", "review": {"status": "reviewed", "by": "R"},
                "model_release": {"id": "REL_A"}, "family": "wall", "input_sha256": "a" * 64, "run_identity": "r" * 64, "timing": {"compute_s": 31.2}}
    wall = {"fields": {"wss": {}}, "peak": {"p99_pa": 16.6, "max_pa": 48.6}, "wss_field_pa": {"mean": 2.1, "area_frac_low": 0.43, "area_low_mm2": 14780.0, "area_frac_high": 0.17, "area_high_mm2": 5764.0, "area_frac_very_high": 0.097},
            "per_branch": {"主动脉": {"wss_p99_pa": 5.09}, "左髂内": {"wss_p99_pa": 37.1}, "未知": {"wss_p99_pa": 1}},
            "reference_assessment": {"population": {"percentile": 32.35}}, "quality": {"level": "good", "label": "模型集成稳定"},
            "findings": {"items": [{"severity": "attention"}, {"severity": "attention"}, {"severity": "note"}, {"severity": "info"}]}}
    row = build_row(wall_job, wall)
    assert list(row) == list(COLUMNS)
    assert row["wss_p99_pa"] == 16.6 and row["branch_p99_左髂内"] == 37.1 and row["branch_p99_右髂外"] is None and row["area_very_high_mm2"] is None
    assert row["population_percentile"] == 32.35 and row["quality_grade"] == "模型集成稳定" and row["compute_total_s"] == 31.2
    assert row["tags"] == "AAA, 随访" and row["review_status"] == "reviewed" and row["reviewer"] == "R" and row["input_sha256_12"] == "a" * 12
    assert row["findings_attention"] == 2 and row["findings_note"] == 1 and row["speed_p99_m_s"] is None
    volume_job = {"id": "j2", "case_id": "V", "status": "done", "model_release": {"id": "REL_V"}, "family": "volume", "review": {}}
    volume = {"fields": {"pressure": {}, "velocity": {}}, "volume_statistics": {"speed_m_s": {"p99": 1.37, "max": 1.57}, "pressure_interior_pa": {"min": -1404.9, "max": 130.5}},
              "findings": {"items": [{"kind": "pressure_drop", "branch": "左髂内", "value": 981.7, "severity": "note"}, {"kind": "low_speed_region", "severity": "note"},
                                     {"kind": "low_speed_region", "severity": "note"}]}, "timing_s": {"total": 40.9}}
    row = build_row(volume_job, volume)
    assert row["speed_p99_m_s"] == 1.37 and row["pressure_min_pa"] == -1404.9 and row["dp_左髂内"] == 981.7 and row["dp_主动脉"] is None
    assert row["low_speed_regions"] == 2 and row["wss_p99_pa"] is None and row["compute_total_s"] == 40.9 and row["review_status"] == "unreviewed"
    assert row["max_diameter_mm"] is None and row["sac_present"] is None   # no morphology in this summary
    queued = build_row({"id": "j3", "status": "queued", "family": None}, None)
    assert queued["job_id"] == "j3" and queued["family"] is None and queued["findings_attention"] is None
    data = to_csv([build_row(wall_job, wall), row, queued])
    assert data.startswith("﻿".encode("utf-8"))
    parsed = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert parsed[0] == list(COLUMNS) and parsed[1][0] == "j1" and parsed[3][COLUMNS.index("wss_p99_pa")] == ""
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(to_xlsx([build_row(wall_job, wall), row])))
    sheet = book.active
    assert [cell.value for cell in sheet[1]] == list(COLUMNS) and sheet.cell(row=2, column=COLUMNS.index("wss_p99_pa") + 1).value == 16.6


def test_morphology_columns_sit_before_the_wall_columns_and_fill_for_both_families():
    assert COLUMNS[COLUMNS.index("compute_total_s") + 1:][:8] == MORPHOLOGY_COLUMNS
    assert COLUMNS.index("lumen_volume_ml") < COLUMNS.index("wss_p99_pa") < COLUMNS.index("speed_p99_m_s")
    morphology = {"lumen_volume_ml": 158.2,
                  "aorta": {"max": {"max_diameter_mm": 63.4, "s_from_root_mm": 120.0},
                            "sac": {"present": True, "length_mm": 85.0, "volume_ml": 96.3},
                            "neck": {"present": True, "length_mm": 80.0, "diameter_mean_mm": 19.1}}}
    wall = build_row({"id": "j1", "family": "wall"}, {"fields": {"wss": {}}, "morphology": morphology, "peak": {"p99_pa": 1.0}})
    volume = build_row({"id": "j2", "family": "volume"}, {"fields": {"pressure": {}, "velocity": {}}, "morphology": morphology})
    for row in (wall, volume):
        assert row["max_diameter_mm"] == 63.4 and row["max_diameter_s_mm"] == 120.0 and row["sac_present"] == 1
        assert row["sac_length_mm"] == 85.0 and row["sac_volume_ml"] == 96.3
        assert row["neck_diameter_mm"] == 19.1 and row["neck_length_mm"] == 80.0 and row["lumen_volume_ml"] == 158.2
    absent = build_row({"id": "j3", "family": "wall"}, {"fields": {"wss": {}},
                                                       "morphology": {"aorta": {"max": {"max_diameter_mm": 21.0},
                                                                                "sac": {"present": False}, "neck": {"present": False}},
                                                                      "lumen_volume_ml": 70.0}})
    assert absent["sac_present"] == 0 and absent["sac_length_mm"] is None and absent["neck_diameter_mm"] is None
    assert absent["max_diameter_mm"] == 21.0 and absent["lumen_volume_ml"] == 70.0


def test_http_export_route_isolated_by_owner(tmp_path):
    service = Service(tmp_path)
    api = service.api()
    try:
        cookie, headers, owner = service.session(api)
        a = finished(service.manager, owner=owner, case_id="A")
        b = finished(service.manager, owner=owner, case_id="B", protocol="single_frame_volume", summary_extra={"fields": {"pressure": {}, "velocity": {}}, "volume_statistics": {"speed_m_s": {"p99": 0.5}}})
        other = finished(service.manager, owner="someone-else", case_id="C")
        status, data, response = call(api, "GET", f"/api/jobs/export?ids={a['id']},{b['id']}&format=csv", {"Cookie": cookie})
        assert status == 200 and response.getheader("Content-Type").startswith("text/csv") and "wss_summary_" in response.getheader("Content-Disposition")
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        assert rows[0] == list(COLUMNS) and rows[1][1] == "A" and rows[2][1] == "B"
        assert rows[1][COLUMNS.index("wss_p99_pa")] == "1.0" and rows[2][COLUMNS.index("speed_p99_m_s")] == "0.5" and rows[2][COLUMNS.index("dp_主动脉")] == "266.6"
        status, data, response = call(api, "GET", f"/api/jobs/export?ids={a['id']}&format=xlsx", {"Cookie": cookie})
        assert status == 200 and data[:2] == b"PK" and "spreadsheetml" in response.getheader("Content-Type")
        status, payload, _ = call(api, "GET", f"/api/jobs/export?ids={a['id']},{other['id']}", {"Cookie": cookie})
        assert status == 404
        status, payload, _ = call(api, "GET", "/api/jobs/export?ids=&format=csv", {"Cookie": cookie})
        assert status == 400
    finally:
        api.close(); service.close()


def test_population_blocks_only_for_releases_that_declare_one(tmp_path):
    from wss_deploy.export_table import population_block, population_blocks

    class FakeRegistry:
        def __init__(self, paths): self.paths = paths
        def describe(self, release_id):
            if release_id not in self.paths: raise KeyError(release_id)
            return type("R", (), {"path": self.paths[release_id]})()

    with_pop, without_pop = tmp_path / "REL_A", tmp_path / "REL_V"
    with_pop.mkdir(); without_pop.mkdir()
    (with_pop / "reference.json").write_text(json.dumps({
        "population_reference": {"metric": "p99_pa", "values_pa": [1.0, 2.0, float("3.5")], "reference_set": "train136 CV3 OOF",
                                 "thresholds_pa": [0.4, 4.0, 7.0]}}), encoding="utf-8")
    (without_pop / "reference.json").write_text(json.dumps({"geometry_reference": {}}), encoding="utf-8")
    assert population_block(without_pop) is None and population_block(tmp_path / "missing") is None
    blocks = population_blocks(FakeRegistry({"REL_A": with_pop, "REL_V": without_pop}), ["REL_A", "REL_V", "REL_A", None, "REL_X"])
    assert set(blocks) == {"REL_A"}
    assert blocks["REL_A"] == {"metric": "p99_pa", "values_pa": [1.0, 2.0, 3.5], "case_count": 3,
                               "reference_set": "train136 CV3 OOF", "thresholds_pa": [0.4, 4.0, 7.0]}
    assert population_blocks(None, ["REL_A"]) == {}


def test_http_export_json_scope_done_carries_rows_and_population(tmp_path):
    release_dir = tmp_path / "releases" / "REL_A"
    release_dir.mkdir(parents=True)
    (release_dir / "reference.json").write_text(json.dumps({
        "population_reference": {"metric": "p99_pa", "values_pa": [1.5, 2.5], "reference_set": "OOF"}}), encoding="utf-8")

    class FakeRegistry:
        def describe(self, release_id):
            if release_id != "REL_A": raise ValueError("unknown release")
            return type("R", (), {"path": release_dir})()
        def list(self): return [{"id": "REL_A", "release": "REL_A"}]

    service = Service(tmp_path)
    api = service.api()
    try:
        cookie, headers, owner = service.session(api)
        a = finished(service.manager, owner=owner, case_id="A")
        finished(service.manager, owner="someone-else", case_id="C")
        service.manager.create(owner, content=b"queued", filename="q.stl", case_id="Q")
        service.manager.registry = FakeRegistry()   # jobs are bound to the injected release, as in production
        status, payload, response = call(api, "GET", "/api/jobs/export?format=json&scope=done", {"Cookie": cookie})
        assert status == 200 and response.getheader("Content-Type").startswith("application/json")
        assert payload["columns"] == list(COLUMNS)
        assert [row["case_id"] for row in payload["rows"]] == ["A"]  # other owners and unfinished jobs excluded
        assert payload["rows"][0]["wss_p99_pa"] == 1.0 and payload["rows"][0]["release_id"] == "REL_A"
        assert payload["population"] == {"REL_A": {"metric": "p99_pa", "values_pa": [1.5, 2.5], "case_count": 2,
                                                   "reference_set": "OOF", "thresholds_pa": [0.4, 4.0, 7.0]}}
        status, payload, _ = call(api, "GET", f"/api/jobs/export?format=json&ids={a['id']}", {"Cookie": cookie})
        assert status == 200 and len(payload["rows"]) == 1 and payload["rows"][0]["job_id"] == a["id"]
        status, payload, _ = call(api, "GET", "/api/jobs/export?format=json", {"Cookie": cookie})
        assert status == 400  # scope=ids still requires ids
        status, payload, _ = call(api, "GET", "/api/jobs/export?format=json&scope=all", {"Cookie": cookie})
        assert status == 400
        status, payload, _ = call(api, "GET", "/api/jobs/export?format=pdf", {"Cookie": cookie})
        assert status == 400
    finally:
        api.close(); service.close()
