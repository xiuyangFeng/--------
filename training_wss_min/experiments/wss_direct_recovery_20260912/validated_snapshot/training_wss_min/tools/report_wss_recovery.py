"""Idempotent evidence-based tracker and 168-column workbook recovery report.

Only owns the new fourteen recovery rows and tracker section 19. All historical
cells, including their styles/comments, and both other sheets are verified after
saving. Pending or unaccepted runs never acquire numerical result cells.
"""
from __future__ import annotations

import argparse
import copy
from datetime import datetime
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook
from openpyxl.comments import Comment

from training_wss_min.tools.tracker_section import replace_or_append
from training_wss_min.tools.wss_recovery_common import ROOT, EXP, CONFIGS, RUNS

BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
TRACKER = ROOT / "docs/02-推进与变更/WSS_PINN/WSS_V5_训练实验跟踪.md"
GROUP = "WSS直接回归纠正矩阵｜2026-09-12（14臂）"
HEADING = "## 19. 直接回归14臂纠正矩阵：修复、预检与恢复执行（2026-09-12）"
MISSING = "—"
DESCRIPTIONS = {
    "E0": "M2独立query同期控制",
    "E1": "A5 SAME-query同期控制",
    "E2": "局部切平面轴向/周向分额邻域",
    "E3": "完整壁面K16 query patch＋粗层上下文FiLM残差",
    "E4": "max＋距离加权mean＋std池化，输出宽度不变",
    "E5": "固定毫米邻域＋局部半径比例邻域",
    "E6": "物理毫米范围/法向及切平面相容近邻差分损失",
    "E7": "逐病例独立配额的pairwise ranking loss",
    "E8": "70%空间覆盖＋15%高曲率＋15%近分叉support采样",
    "E9": "真实入口面几何面积＋公共名义Q构造Q/A输入",
    "C1": "E2＋E3",
    "C2": "E2＋E6",
    "C3": "E3＋E7",
    "C4": "E2–E5中动态选出的结构＋E9，仅训练一个C4",
}

# Existing workbook columns, one-based. Neither labels nor semantics are changed.
METRIC_COLUMNS = {
    6: "field_casebalanced.r2", 7: "field.rmse", 8: "field.mae", 9: "field.nmae_range",
    10: "aggregate.r2_casemed", 11: "aggregate.r2_casep10",
    15: "field.r2", 16: "aggregate.r2_casemean", 17: "field.nrmse_range",
    18: "aggregate.nrmse_casemean", 19: "aggregate.nmae_casemean",
    20: "regional_field.high_wss.r2", 21: "calibration.top10_pred_true_ratio",
    22: "hotspot.top10_iou_casemean", 23: "calibration.p99_pred_true_ratio",
    24: "normalized.field_casebalanced.r2", 25: "normalized.field.r2",
    26: "normalized.aggregate.r2_casemean", 27: "normalized.aggregate.r2_casemed",
    28: "normalized.aggregate.r2_casep10", 29: "normalized.field.mae", 30: "normalized.field.rmse",
    31: "normalized.field.nrmse_range", 32: "normalized.aggregate.nrmse_casemean",
    33: "normalized.field.nmae_range", 34: "normalized.aggregate.nmae_casemean",
    35: "normalized.hotspot.spearman_all_casemean", 36: "normalized.hotspot.spearman_high_wss_casemean",
    37: "normalized.calibration.top10_pred_true_ratio", 38: "normalized.calibration.p99_pred_true_ratio",
    39: "field_casebalanced.r2_linear_fit", 40: "field_casebalanced.linear_fit_slope",
    41: "field_casebalanced.linear_fit_intercept", 42: "field.r2_linear_fit",
    43: "field.linear_fit_slope", 44: "field.linear_fit_intercept",
    45: "normalized.field_casebalanced.r2_linear_fit", 46: "normalized.field_casebalanced.linear_fit_slope",
    47: "normalized.field_casebalanced.linear_fit_intercept", 48: "normalized.field.r2_linear_fit",
    49: "normalized.field.linear_fit_slope", 50: "normalized.field.linear_fit_intercept",
    51: "regional_field.high_wss.mae", 52: "regional_field.high_wss.rmse",
    53: "regional_field.high_wss.nrmse_range", 54: "field_casebalanced.mae",
    55: "field_casebalanced.rmse", 56: "normalized.field_casebalanced.mae",
    57: "normalized.field_casebalanced.rmse", 58: "normalized.regional_field.high_wss.r2",
    59: "hotspot.spearman_all_casemean", 60: "hotspot.spearman_high_wss_casemean",
    62: "group_casebalanced.AG.r2", 63: "group_casebalanced.AAA.r2", 64: "group_casebalanced.ILO.r2",
}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path, default=None):
    path = Path(path)
    return json.loads(path.read_text()) if path.is_file() else default


def save_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(dir=path.parent, prefix=path.name + ".", suffix=".tmp")
    try:
        with os.fdopen(fd, "w") as out:
            json.dump(value, out, ensure_ascii=False, indent=2, allow_nan=False)
            out.write("\n")
        os.replace(temp, path)
    finally:
        if Path(temp).exists():
            Path(temp).unlink()


def get_metric(metrics, dotted):
    value = metrics
    for key in dotted.split("."):
        if not isinstance(value, dict) or key not in value:
            return None
        value = value[key]
    if isinstance(value, (int, float)) and math.isfinite(value):
        return value
    return None


def metric_values(metrics):
    return {column: get_metric(metrics, path) for column, path in METRIC_COLUMNS.items()}


def finite_summary(value):
    if isinstance(value, dict):
        return {k: finite_summary(v) for k, v in value.items() if k != "per_case"}
    if isinstance(value, list):
        return [finite_summary(v) for v in value]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def history_evidence(run, epochs):
    path = run / "history.jsonl"
    rows = []
    errors = []
    if path.is_file():
        lines = path.read_text().splitlines()
        for i, line in enumerate(lines):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                # A live writer may be midway through its final line.
                errors.append(f"history line {i + 1} is not complete JSON")
                break
            rows.append(row)
    actual = [r.get("epoch") for r in rows]
    complete = actual == list(range(epochs))
    finite_loss = all(isinstance(r.get("train_loss"), (int, float)) and math.isfinite(r["train_loss"]) for r in rows)
    return {"rows": len(rows), "last_epoch": actual[-1] if actual else None,
            "epochs_complete": complete, "finite_training_loss": finite_loss,
            "elapsed_seconds": rows[-1].get("elapsed_s") if rows else None, "errors": errors}


def run_evidence(arm, config_path, state):
    config = read_json(config_path)
    if config["train"]["seed"] != 1234 or config["train"]["epochs"] != 400:
        raise ValueError("report only accepts the authorized seed1234/400epoch protocol")
    run = RUNS / config["name"]
    hist = history_evidence(run, int(config["train"]["epochs"]))
    acceptance = read_json(run / "acceptance.json", {})
    failures = []
    outputs = {}
    accepted_checkpoints = {}
    for checkpoint in ("best", "last"):
        checkpoint_failure_start = len(failures)
        path = run / "eval" / f"ckpt_{checkpoint}" / "metrics.json"
        acc = read_json(run / f"acceptance_{checkpoint}.json", {})
        accepted_checkpoints[checkpoint] = acc.get("passed") is True
        if not (run / f"ckpt_{checkpoint}.pt").is_file():
            failures.append(f"missing ckpt_{checkpoint}.pt")
        if acc.get("passed") is not True:
            failures.append(f"{checkpoint} independent acceptance not passed")
        payload = read_json(path)
        if payload is None:
            failures.append(f"missing {checkpoint} metrics")
            accepted_checkpoints[checkpoint] = False
            continue
        metrics = payload.get("test", {})
        current_hash = sha(path)
        if acc.get("passed") is True and acc.get("files_sha256", {}).get(str(path.resolve())) != current_hash:
            failures.append(f"{checkpoint} metrics changed since independent acceptance")
        if acceptance.get("passed") is True and acceptance.get("files_sha256", {}).get(str(path.resolve())) != current_hash:
            failures.append(f"{checkpoint} metrics differ from whole-run acceptance")
        checkpoint_path = run / f"ckpt_{checkpoint}.pt"
        accepted_ckpt = acceptance.get("training", {}).get("checkpoints", {}).get(checkpoint, {})
        if acceptance.get("passed") is True and (not checkpoint_path.is_file()
                or accepted_ckpt.get("sha256") != sha(checkpoint_path)):
            failures.append(f"{checkpoint} checkpoint differs from whole-run acceptance")
        core = ["field_casebalanced.r2", "field.mae", "field.rmse", "normalized.field_casebalanced.r2"]
        if any(get_metric(metrics, key) is None for key in core):
            failures.append(f"nonfinite/incomplete {checkpoint} core metrics")
        expected_cases = set(read_json(config["data"]["split_path"])["test_cases"])
        if set(metrics.get("per_case", {})) != expected_cases or set(metrics.get("normalized", {}).get("per_case", {})) != expected_cases:
            failures.append(f"wrong {checkpoint} evaluated case set")
        outputs[checkpoint] = {"path": str(path), "sha256": current_hash, "metrics": finite_summary(metrics)}
        accepted_checkpoints[checkpoint] = acc.get("passed") is True and len(failures) == checkpoint_failure_start
    if not hist["epochs_complete"] or not hist["finite_training_loss"] or hist["errors"]:
        failures.append("400-epoch finite contiguous history incomplete")
    if acceptance.get("passed") is not True:
        failures.append("whole-run acceptance not passed")
    elif acceptance.get("training", {}).get("epochs") != 400:
        failures.append("whole-run acceptance does not prove 400 epochs")
    for path in (run / "history.jsonl", run / "config.json"):
        recorded = acceptance.get("files_sha256", {}).get(str(path.resolve()))
        if recorded is not None and (not path.is_file() or sha(path) != recorded):
            failures.append(f"{path.name} changed since whole-run acceptance")
    accepted = not failures
    status = state.get("status", "pending")
    if accepted:
        status_text = "已完成并独立验收"
    elif status in {"failed", "blocked", "interrupted", "skipped"}:
        status_text = {"failed": "失败", "blocked": "阻塞", "interrupted": "中断", "skipped": "跳过"}[status]
    elif status == "complete":
        status_text = "队列结束，验收证据待齐全"
    elif hist["rows"]:
        status_text = f"训练/评估中（{hist['rows']}/400轮）"
    elif status in {"train", "eval_best", "eval_last", "ready_eval_best", "ready_eval_last"}:
        status_text = {"train": "训练启动中", "eval_best": "best评估中", "eval_last": "last评估中",
                       "ready_eval_best": "等待best评估", "ready_eval_last": "等待last评估"}[status]
    else:
        status_text = "等待执行"
    return {**arm, "config": str(config_path), "run_name": config["name"], "run_dir": str(run),
            "config_sha256": sha(config_path), "seed": config["train"]["seed"], "epochs": config["train"]["epochs"],
            "query_mode": config["data"]["query_mode"], "input_dim": len(config["data"]["input_features"]),
            "support_n_points": config["data"]["support_n_points"], "history": hist,
            "queue_status": status, "status": status_text, "accepted": accepted,
            "acceptance_failures": failures, "accepted_checkpoints": accepted_checkpoints,
            "best": outputs.get("best") if accepted else None, "last": outputs.get("last") if accepted else None,
            "diagnostics": read_json(run / "training_diagnostics.json", {}), "queue_record": state}


def build_report(experiment=EXP, configs=CONFIGS):
    matrix = read_json(configs / "matrix.json")
    if {a["id"] for a in matrix["arms"]} != {*(f"E{i}" for i in range(10)), *(f"C{i}" for i in range(1, 5))}:
        raise ValueError("expected exactly the authorized E0–E9 + C1–C4 matrix")
    queue = read_json(experiment / "queue_status.json", {})
    submission = read_json(experiment / "submission.json", {})
    preflight = read_json(experiment / "runtime_preflight.json", {})
    arms = []
    for arm in matrix["arms"]:
        state = queue.get("arms", {}).get(arm["id"], {})
        config_path = Path(state.get("config", configs / arm["config"]))
        item = run_evidence(arm, config_path, state)
        if not submission.get("job_id") and not queue.get("job_id") and item["status"] == "等待执行":
            item["status"] = "尚未提交（预检阶段）"
        arms.append(item)
    selection = read_json(experiment / "combination_selection.json", queue.get("combination_selection", {}))
    return {"generated_at": datetime.now().astimezone().isoformat(), "experiment": str(experiment),
            "submission": submission, "job_id": queue.get("job_id") or submission.get("job_id"),
            "queue_status": queue.get("status", "not_submitted"),
            "preflight": {key: preflight.get(key) for key in ("passed", "slurm_job_id", "cli_smoke_passed")},
            "training_complete": sum(a["history"]["epochs_complete"] and a["history"]["finite_training_loss"] for a in arms),
            "best_accepted": sum(a["accepted_checkpoints"]["best"] for a in arms),
            "last_accepted": sum(a["accepted_checkpoints"]["last"] for a in arms),
            "complete": sum(a["accepted"] for a in arms), "total": 14, "selection": selection,
            "arms": arms, "interpretation": "Single seed 1234; test34 used for adaptive development selection, not untouched validation."}


def cell_record(cell):
    return (cell.value, copy.copy(cell._style), cell.number_format,
            (cell.comment.text, cell.comment.author) if cell.comment else None,
            cell.hyperlink.target if cell.hyperlink else None)


def paired_references(arm, selection):
    if arm["id"] == "E0":
        return []
    pairs = {"C1": ["E2", "E3"], "C2": ["E2", "E6"], "C3": ["E3", "E7"],
             "C4": [selection.get("selected_structure"), "E9"]}.get(arm["id"], [])
    return list(dict.fromkeys([x for x in pairs + ["E0"] if x]))


def update_workbook(report, book=BOOK):
    before = sha(book)
    backups = Path(report["experiment"]) / "report_backups"
    backups.mkdir(parents=True, exist_ok=True)
    backup = backups / f"workbook_before_{before[:16]}.xlsx"
    if not backup.exists():
        shutil.copy2(book, backup)
    wb = load_workbook(book)
    if wb.sheetnames != ["WSS实验矩阵", "速度与压力实验矩阵", "指标说明"]:
        raise ValueError("unexpected workbook sheet schema")
    ws = wb["WSS实验矩阵"]
    if ws.max_column != 168 or ws["FI5"].value != "原始实验 ID" or "WSSResults" not in ws.tables:
        raise ValueError("unexpected workbook 168-column schema")
    owned = [r for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value == GROUP]
    if owned and len(owned) != 14:
        raise ValueError("owned recovery rows must number exactly 14")
    locations = ({ws.cell(r, 165).value: r for r in owned} if owned else
                 {a["run_name"]: ws.max_row + i + 1 for i, a in enumerate(report["arms"])})
    if set(locations) != {a["run_name"] for a in report["arms"]}:
        raise ValueError("recovery run identities differ from workbook")
    preserved = {(s.title, c.coordinate): cell_record(c) for s in wb for row in s for c in row
                 if not (s.title == ws.title and c.row in owned)}
    dimensions = {s.title: (copy.deepcopy(s.column_dimensions), copy.deepcopy(s.row_dimensions)) for s in wb}
    merges = {s.title: tuple(map(str, s.merged_cells.ranges)) for s in wb}
    lookup = {a["id"]: a for a in report["arms"]}
    for i, arm in enumerate(report["arms"]):
        row = locations[arm["run_name"]]
        values = [MISSING] * 168
        values[:5] = [GROUP, f"{arm['id']}｜{DESCRIPTIONS[arm['id']]}｜{arm['status']}", "WSS · Pa",
                      "best主结果；last见逐指标批注" if arm["accepted"] else "尚无完整验收结果",
                      "V5 train138/test34；峰值1162；test34开发筛选"]
        refs = paired_references(arm, report["selection"])
        values[13] = "；".join(refs) or "M2同期控制，不计算自参照增益"
        note = (f"Job {report['job_id'] or '尚未提交'}；{arm['status']}。"
                f"有效history {arm['history']['rows']}/400；best/last分别独立验收。"
                "单seed1234；仅开发筛选，不作统计显著/稳定泛化结论。"
                "E1为真正SAME A5；E2–E9和组合统一以E0独立query M2为父臂，避免改变query协议的混杂。")
        if arm["id"] == "C4":
            note += " C4动态选择记录：" + json.dumps(report["selection"], ensure_ascii=False)
        if not arm["accepted"]:
            note += " 尚缺：" + "；".join(arm["acceptance_failures"])
        source_paths = [arm["config"], str(Path(arm["run_dir"]) / "history.jsonl"),
                        str(Path(arm["run_dir"]) / "acceptance.json"), str(Path(report["experiment"]) / "recovery_report.json")]
        last_values = {}
        if arm["accepted"]:
            best, last = arm["best"]["metrics"], arm["last"]["metrics"]
            for col, value in metric_values(best).items():
                if value is not None:
                    values[col - 1] = value
            values[11] = f"{int(best['aggregate']['r2_negative_cases'])}/{int(best['aggregate']['n_cases'])}"
            values[60] = arm["diagnostics"].get("parameter_count", MISSING)
            values[112] = get_metric(last, "field_casebalanced.r2") - get_metric(best, "field_casebalanced.r2")
            values[139] = 1
            values[147] = arm["diagnostics"].get("parameter_count", MISSING)
            values[148] = arm["diagnostics"].get("elapsed_seconds", arm["history"]["elapsed_seconds"])
            last_values = metric_values(last)
            source_paths += [arm["best"]["path"], arm["last"]["path"]]
            details = {"best": best, "last": last, "diagnostics": arm["diagnostics"]}
            note += "\n完整best/last汇总（每病例细节见源metrics）：\n" + json.dumps(details, ensure_ascii=False, indent=2)
            pair_paths = ["field_casebalanced.r2", "normalized.field_casebalanced.r2", "field.mae",
                          "hotspot.spearman_all_casemean", "hotspot.top10_iou_casemean"]
            for j, ref in enumerate(refs[:4]):
                parent = lookup.get(ref)
                if not parent or not parent["accepted"]:
                    continue
                base = 64 + j * 6
                values[base] = ref + "；" + parent["best"]["path"]
                for k, dotted in enumerate(pair_paths, 1):
                    own, other = get_metric(best, dotted), get_metric(parent["best"]["metrics"], dotted)
                    if own is not None and other is not None:
                        values[base + k] = own - other
                if ref == "E0":
                    values[12] = get_metric(best, "field_casebalanced.r2") - get_metric(parent["best"]["metrics"], "field_casebalanced.r2")
            values[88] = "；".join(refs) + "；best同checkpoint配对；单seed，无跨seed统计"
        sources = "\n".join(source_paths)
        values[159:] = ["Pa", "PointNeXt-R＋LocalGeoPE＋L-SA2＋local branch", DESCRIPTIONS[arm["id"]],
                        f"seed1234 / 400ep / {arm['input_dim']}D / support5000 / query5000 {arm['query_mode']}",
                        arm["support_n_points"], arm["run_name"], note[:32000], sources,
                        "2026-09-12纠正恢复矩阵；保留第18节失败记录"]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col)
            cell._style = copy.copy(ws.cell(310 if i % 2 else 311, col)._style)
            cell.value = value
            cell.comment = None
            cell.hyperlink = None
            if arm["accepted"] and col in last_values:
                cell.comment = Comment(f"本单元格=best；同口径last={last_values[col] if last_values[col] is not None else MISSING}\n"
                                       f"路径={METRIC_COLUMNS[col]}\nbest来源={arm['best']['path']}\nlast来源={arm['last']['path']}", "Codex recovery report")
        ws.cell(row, 166).comment = Comment(note[:32000], "Codex recovery report")
        ws.cell(row, 167).comment = Comment(sources, "Codex recovery report")
        ws.row_dimensions[row].height = 72
    ws.tables["WSSResults"].ref = f"A5:FL{ws.max_row}"
    ws.tables["WSSResults"].autoFilter.ref = f"A5:FL{ws.max_row}"
    ws.print_area = f"A1:N{ws.max_row}"
    fd, temp = tempfile.mkstemp(dir=book.parent, prefix=".wss_recovery_", suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(temp)
        checked = load_workbook(temp)
        for (sheet, coordinate), record in preserved.items():
            if cell_record(checked[sheet][coordinate]) != record:
                raise AssertionError(f"historical cell changed: {sheet}!{coordinate}")
        assert {s.title: tuple(map(str, s.merged_cells.ranges)) for s in checked} == merges
        for sheet in checked:
            old_cols, old_rows = dimensions[sheet.title]
            assert {k: dict(v) for k, v in sheet.column_dimensions.items()} == {k: dict(v) for k, v in old_cols.items()}
            for key, value in old_rows.items():
                if sheet.title != ws.title or key not in owned:
                    assert dict(sheet.row_dimensions[key]) == dict(value)
        for arm in report["arms"]:
            if not arm["accepted"]:
                for column in [*range(6, 14), *range(15, 160)]:
                    assert checked[ws.title].cell(locations[arm["run_name"]], column).value == MISSING
        checked.close()
        if sha(book) != before:
            raise RuntimeError("workbook changed concurrently; refusing replacement")
        os.replace(temp, book)
    finally:
        if Path(temp).exists():
            Path(temp).unlink()
    evidence = {"passed": True, "before_sha256": before, "after_sha256": sha(book), "backup": str(backup),
                "rows": locations, "records": 14, "historical_cells_verified": len(preserved),
                "historical_values_styles_comments_preserved": True, "other_two_sheets_preserved": True,
                "table_ref": ws.tables["WSSResults"].ref, "accepted_result_rows": report["complete"]}
    save_json(Path(report["experiment"]) / "xlsx_acceptance.json", evidence)
    return evidence


def write_documents(report, workbook, tracker=TRACKER):
    preflight = report["preflight"]
    job = str(report["job_id"]) if report["job_id"] else "尚未提交正式矩阵（预检阶段）"
    lines = [HEADING, "", f"**实际状态：{report['complete']}/14完整验收，训练{report['training_complete']}/14，"
             f"best独立验收{report['best_accepted']}/14，last独立验收{report['last_accepted']}/14。** "
             f"正式作业：{job}；队列状态：`{report['queue_status']}`。更新于{report['generated_at']}。", "",
             "第18节的14079失败与原16臂偏差记录保留。本次另建14臂纠正矩阵：E0–E9＋C1–C4，全部seed1234、400epoch；"
             "独立输出目录，旧失败记录不覆盖。E1保持原A5的SAME query；E2–E9与组合统一以E0/M2独立query为父臂。", "",
             "已落实：真实法向/atlas切向局部方向邻域；完整壁面K16 query patch与FiLM；统计池化；毫米/局部半径邻域；"
             "E6物理距离、法向/切平面相容点对差分与E7逐病例排序配额；空间覆盖/高曲率/分叉support配额。"
             "E6属于几何相容近邻差分，未宣称取得原始曲面网格测地邻接。", "",
             "E9的172/172例输入审计已通过：公共Q=0.000107269066557m³/s；真实入口面积231.450–999.260mm²；"
             "Q/A=0.107349–0.463466m/s。只读取冻结入口面几何和公共名义波形；无CFD解、实测flux或RCR输入。"
             "新增通道统计只用train138，每例一次；旧25D统计保留。该特征提供几何相关参考速度，不增加患者特异流量信息。", "",
             f"GPU预检记录：passed={preflight.get('passed')}，Slurm预检作业={preflight.get('slurm_job_id')}，"
             f"train→全壁面评估CLI smoke={preflight.get('cli_smoke_passed')}。没有通过证据时不会标记已通过。"
             "正式队列按每卡共享显存槽位并行，训练→best评估→last评估；同卡评估串行，独立验收通过后才填成绩。", "",
             "| 臂 | 实际变量 | 执行状态 | history轮数 | best Pa R²cb | last Pa R²cb |", "|---|---|---|---:|---:|---:|"]
    for arm in report["arms"]:
        b = get_metric(arm["best"]["metrics"], "field_casebalanced.r2") if arm["accepted"] else None
        l = get_metric(arm["last"]["metrics"], "field_casebalanced.r2") if arm["accepted"] else None
        fmt = lambda x: f"{x:.6f}" if x is not None else MISSING
        lines.append(f"| {arm['id']} | {DESCRIPTIONS[arm['id']]} | {arm['status']} | {arm['history']['rows']} | {fmt(b)} | {fmt(l)} |")
    selection = report["selection"]
    lines += ["", "C4选择：" + (json.dumps(selection, ensure_ascii=False) if selection else
              "尚未执行；第一阶段完成后按已登记保护线筛E2–E5，再按归一化R²cb、Pa R²cb、ID排序。无合格结构时仍按用户要求运行排名最高结构＋E9，并标记forced_exploratory；只训练1个C4。"), "",
              "结论边界：单seed只用于开发筛选；test34参与自适应组合选择，不能称未触碰的确认集。"
              "完成状态须同时满足连续400轮有限训练记录、best/last checkpoint、test34全病例评估、独立重算验收与完整验收文件。"
              "在这些证据齐全前成绩一律为“—”，不使用历史成绩、0或预检精度代替。", "",
              f"工作簿新增/更新14条记录（第{min(workbook['rows'].values())}–{max(workbook['rows'].values())}行），保留历史348行及另外两页的值、样式与批注。"
              "best主指标沿既有168列表头回填，last同口径值写入逐指标批注及完整结果JSON；回填前自动备份并验收。", "",
              "[恢复执行明细](../../../training_wss_min/experiments/wss_direct_recovery_20260912/recovery_report.json) · "
              "[入口尺度与冻结几何审计](../../../training_wss_min/experiments/wss_direct_recovery_20260912/inlet_audit/README.md) · "
              "[工作簿验收](../../../training_wss_min/experiments/wss_direct_recovery_20260912/xlsx_acceptance.json) · "
              "[工作簿](../../03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx)", ""]
    before = tracker.read_text()
    backups = Path(report["experiment"]) / "report_backups"
    backup = backups / f"tracker_before_{hashlib.sha256(before.encode()).hexdigest()[:16]}.md"
    if not backup.exists():
        backup.write_text(before)
    after = replace_or_append(before, HEADING, "\n".join(lines))
    if HEADING not in before and not after.startswith(before.rstrip("\n")):
        raise AssertionError("tracker append changed historical content")
    tracker.write_text(after)
    readme = ["# WSS直接回归14臂纠正矩阵", "", *lines[2:]]
    readme_text = "\n".join(readme).replace("../../../training_wss_min/", str(ROOT / "training_wss_min") + "/")
    readme_text = readme_text.replace("../../03-汇报材料/", str(ROOT / "docs/03-汇报材料") + "/")
    (Path(report["experiment"]) / "README.md").write_text(readme_text)


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--experiment-dir", type=Path, default=EXP)
    p.add_argument("--config-dir", type=Path, default=CONFIGS)
    p.add_argument("--workbook", type=Path, default=BOOK)
    p.add_argument("--tracker", type=Path, default=TRACKER)
    args = p.parse_args(argv)
    report = build_report(args.experiment_dir, args.config_dir)
    save_json(args.experiment_dir / "recovery_report.json", report)
    workbook = update_workbook(report, args.workbook)
    write_documents(report, workbook, args.tracker)
    print(json.dumps({"job_id": report["job_id"], "complete": report["complete"], "total": 14,
                      "rows": list(workbook["rows"].values()), "backup": workbook["backup"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
