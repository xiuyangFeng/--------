#!/usr/bin/env python3
"""Read-only, independent transcription audit of the A5 follow-up delivery.

No reporter/writer functions are imported. Every expected metric comes from a
run's original metrics.json. Only the explicitly named audit JSON is written.
This checks faithful reporting, not whether the underlying evaluator is correct.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
EXP = Path(__file__).resolve().parent
RUNS = ROOT / "training_wss_min/runs"
MATRIX = ROOT / "training_wss_min/configs/v6_followup_20260909/matrix.json"
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
BACKUP = EXP / "WSS_PointNet实验矩阵与结果汇总last_backup_before_followup_20260909.xlsx"
PREFIX = "V6-A5后续20260909"
FIELDS = {
    "physical_r2_cb": "field_casebalanced.r2",
    "normalized_r2_cb": "normalized.field_casebalanced.r2",
    "case_mean": "aggregate.r2_casemean",
    "case_p10": "aggregate.r2_casep10",
    "mae": "field.mae",
    "rmse": "field.rmse",
    "high_wss_r2": "regional_field.high_wss.r2",
    "top10_ratio": "calibration.top10_pred_true_ratio",
    "p99_ratio": "calibration.p99_pred_true_ratio",
    "top10_iou": "hotspot.top10_iou_casemean",
    "spearman": "hotspot.spearman_all_casemean",
    "linear_fit_r2_cb": "field_casebalanced.r2_linear_fit",
    "linear_fit_slope_cb": "field_casebalanced.linear_fit_slope",
    "linear_fit_intercept_cb": "field_casebalanced.linear_fit_intercept",
}
OVERVIEW = {
    "G": "field_casebalanced.r2", "H": "field.r2",
    "I": "aggregate.r2_casemean", "J": "aggregate.r2_casemed",
    "K": "aggregate.r2_casep10", "M": "field.rmse", "N": "field.mae",
    "O": "field.nrmse_range", "P": "aggregate.nrmse_casemean",
    "Q": "field.nmae_range", "R": "aggregate.nmae_casemean",
    "S": "regional_field.high_wss.r2", "T": "calibration.top10_pred_true_ratio",
    "U": "hotspot.top10_iou_casemean", "V": "normalized.field_casebalanced.r2",
    "W": "normalized.field.r2", "X": "normalized.aggregate.r2_casemean",
    "Y": "normalized.aggregate.r2_casemed", "Z": "normalized.aggregate.r2_casep10",
    "AA": "normalized.field.mae", "AB": "normalized.field.rmse",
    "AC": "normalized.field.nrmse_range", "AD": "normalized.aggregate.nrmse_casemean",
    "AE": "normalized.field.nmae_range", "AF": "normalized.aggregate.nmae_casemean",
    "AG": "hotspot.spearman_all_casemean", "AH": "hotspot.spearman_high_wss_casemean",
    "AI": "normalized.calibration.top10_pred_true_ratio",
    "AJ": "normalized.calibration.p99_pred_true_ratio",
    "AL": "field_casebalanced.r2_linear_fit", "AM": "field_casebalanced.linear_fit_slope",
    "AN": "field_casebalanced.linear_fit_intercept", "AO": "field.r2_linear_fit",
    "AP": "field.linear_fit_slope", "AQ": "field.linear_fit_intercept",
    "AR": "normalized.field_casebalanced.r2_linear_fit",
    "AS": "normalized.field_casebalanced.linear_fit_slope",
    "AT": "normalized.field_casebalanced.linear_fit_intercept",
    "AU": "normalized.field.r2_linear_fit", "AV": "normalized.field.linear_fit_slope",
    "AW": "normalized.field.linear_fit_intercept",
}
TEACHER = dict(zip(
    "EFGHIJKLMNOPQRS",
    ["field_casebalanced.r2", "aggregate.r2_casemean", "field.rmse", "field.nmae_range",
     "aggregate.nmae_casemean", "regional_field.high_wss.r2", "normalized.field_casebalanced.r2",
     "normalized.field.nmae_range", "hotspot.spearman_all_casemean", "hotspot.top10_iou_casemean",
     "normalized.calibration.p99_pred_true_ratio", "field_casebalanced.r2_linear_fit",
     "field.r2_linear_fit", "field.linear_fit_slope", "field.linear_fit_intercept"], strict=True))
COMPARISON = {
    "D": "field_casebalanced.r2", "G": "normalized.field_casebalanced.r2",
    "H": "aggregate.r2_casemean", "I": "field.mae", "J": "regional_field.high_wss.r2",
    "K": "calibration.top10_pred_true_ratio", "L": "calibration.p99_pred_true_ratio",
    "M": "hotspot.top10_iou_casemean",
}


def dig(obj, path):
    for key in path.split("."):
        obj = obj[key]
    return obj


class Audit:
    def __init__(self):
        self.inputs = {}
        self.counts = Counter()
        self.differences = []

    def read_bytes(self, path):
        path = Path(path)
        data = path.read_bytes()
        self.inputs[str(path)] = hashlib.sha256(data).hexdigest()
        return data

    def read_json(self, path):
        return json.loads(self.read_bytes(path))

    def check(self, category, label, actual, expected, *, exact=False):
        self.counts[category] += 1
        if not exact and isinstance(expected, (float, int)) and not isinstance(expected, bool):
            ok = (isinstance(actual, (float, int)) and not isinstance(actual, bool)
                  and math.isfinite(actual) and math.isfinite(expected)
                  and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12))
        else:
            ok = actual == expected
        if not ok:
            self.differences.append({"category": category, "location": label,
                                     "actual": actual, "expected": expected})


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", type=Path, default=EXP / "delivery_verification_partial.json")
    args = ap.parse_args()
    audit = Audit()
    matrix = audit.read_json(MATRIX)
    arms = matrix["arms"]
    ids = [a["id"] for a in arms]
    anchor = matrix["reference"]["run_name"]
    summaries = {ck: audit.read_json(EXP / f"matrix_summary_{ck}.json") for ck in ("best", "last")}
    update = audit.read_json(EXP / "workbook_update.json")
    sources = {}
    missing_sources = []
    for ck in summaries:
        sources[ck] = {}
        for aid, run_name in [("B0", anchor)] + [(a["id"], a["run_name"]) for a in arms]:
            path = RUNS / run_name / "eval" / f"ckpt_{ck}" / "metrics.json"
            if not path.is_file():
                missing_sources.append(str(path))
                continue
            m = audit.read_json(path)["test"]
            sources[ck][aid] = m
            audit.check("source_test_set", f"{ck}/{aid}/n_cases", m["aggregate"]["n_cases"], 34)
            audit.check("source_test_set", f"{ck}/{aid}/per_case_count", len(m["per_case"]), 34)
        base = sources[ck]["B0"]
        for aid, m in sources[ck].items():
            audit.check("source_test_set", f"{ck}/{aid}/patients",
                        sorted(m["per_case"]), sorted(base["per_case"]), exact=True)

    checkpoint_results = {}
    for ck, summary in summaries.items():
        base = sources[ck]["B0"]
        audit.check("summary_metadata", f"{ck}/checkpoint", summary["checkpoint"], ck)
        audit.check("summary_metadata", f"{ck}/reference", summary["reference"]["run_name"], anchor)
        audit.check("summary_metadata", f"{ck}/reference_checkpoint", summary["reference"]["checkpoint"], ck)
        audit.check("summary_metadata", f"{ck}/arm_ids", sorted(summary["arms"]), sorted(ids), exact=True)
        for key, path in FIELDS.items():
            audit.check("summary_reference", f"{ck}/B0/{key}", summary["reference"].get(key), dig(base, path))
        valid = [aid for aid in ids if summary["arms"][aid]["evidence"]["valid_scientific_result"]]
        for arm in arms:
            aid = arm["id"]
            entry = summary["arms"][aid]
            m = sources[ck].get(aid)
            label = f"{ck}/{aid}"
            parent = arm.get("baseline", arm.get("parent", "B0"))
            audit.check("summary_metadata", label + "/run_name", entry["run_name"], arm["run_name"])
            audit.check("summary_metadata", label + "/parent", entry["parent"], parent)
            if "raw_metrics" in entry:
                audit.check("summary_metadata", label + "/raw_source_exists", m is not None, True)
                audit.check("summary_metadata", label + "/raw_metrics_path", entry.get("raw_metrics_path"),
                            str(RUNS / arm["run_name"] / "eval" / f"ckpt_{ck}" / "metrics.json"))
                if m is not None:
                    for key, path in FIELDS.items():
                        audit.check("summary_raw", label + "/" + key, entry["raw_metrics"].get(key), dig(m, path))
            if aid not in valid:
                for key in list(FIELDS) + ["delta_physical", "delta_normalized", "delta_parent_physical",
                                         "delta_parent_normalized", "per_case_delta", "cases_improved"]:
                    audit.check("summary_pending_blank", label + "/" + key, entry.get(key), None)
                continue
            audit.check("summary_metadata", label + "/valid_source_exists", m is not None, True)
            if m is None:
                continue
            for key, path in FIELDS.items():
                audit.check("summary_valid", label + "/" + key, entry.get(key), dig(m, path))
            pm = base if parent in {"B0", "A5", "B0=A5"} else sources[ck].get(parent) if parent in valid else None
            for suffix, path in (("physical", FIELDS["physical_r2_cb"]), ("normalized", FIELDS["normalized_r2_cb"])):
                audit.check("summary_delta", label + "/delta_" + suffix,
                            entry.get("delta_" + suffix), dig(m, path) - dig(base, path))
                audit.check("summary_delta", label + "/delta_parent_" + suffix,
                            entry.get("delta_parent_" + suffix), dig(m, path) - dig(pm, path) if pm else None)
            delta = {case: m["per_case"][case]["overall"]["r2"] - base["per_case"][case]["overall"]["r2"]
                     for case in base["per_case"]}
            audit.check("summary_per_case", label + "/delta_keys", sorted(entry.get("per_case_delta", {})), sorted(delta), exact=True)
            for case, expected in delta.items():
                audit.check("summary_per_case", label + "/" + case, entry.get("per_case_delta", {}).get(case), expected)
            audit.check("summary_per_case", label + "/cases_improved", entry.get("cases_improved"), sum(v > 0 for v in delta.values()))
        audit.check("summary_metadata", ck + "/completed_count", summary["completed_count"], len(valid))
        audit.check("summary_metadata", ck + "/expected_count", summary["expected_count"], len(arms))
        audit.check("summary_metadata", ck + "/raw_count", summary["raw_metrics_count"], sum("raw_metrics" in e for e in summary["arms"].values()))
        checkpoint_results[ck] = {"valid_count": len(valid), "valid_ids": valid,
                                  "pending_ids": [aid for aid in ids if aid not in valid],
                                  "matrix_finalized_reported": summary["matrix_finalized"]}

    wb = load_workbook(BytesIO(audit.read_bytes(BOOK)), data_only=False)
    old = load_workbook(BytesIO(audit.read_bytes(BACKUP)), data_only=False)
    audit.check("workbook_metadata", "sheetnames", wb.sheetnames, old.sheetnames, exact=True)
    audit.check("workbook_metadata", "sheet_count", len(wb.sheetnames), 5)
    markers = {}
    for name in ("实验矩阵总览", "教师汇报视图", "汇总对比"):
        rows = [c.row for cells in wb[name].iter_rows(min_col=1, max_col=1) for c in cells
                if isinstance(c.value, str) and PREFIX + "｜" in c.value]
        audit.check("workbook_metadata", name + "/marker_count", len(rows), 1)
        if len(rows) != 1:
            raise RuntimeError(f"Cannot identify own section: {name}: {rows}")
        markers[name] = rows[0]
        audit.check("workbook_metadata", name + "/explicit_best", "checkpoint=best" in wb[name].cell(rows[0], 1).value, True)
    preserved = {}
    for name in old.sheetnames:
        before, after = old[name], wb[name]
        end = markers[name] - 1 if name in markers else max(before.max_row, after.max_row)
        audit.check("legacy_boundary", name + "/old_rows_within_cutoff", before.max_row <= end, True)
        count = formulas = 0
        for row in range(1, max(end, before.max_row) + 1):
            for col in range(1, max(before.max_column, after.max_column) + 1):
                expected, actual = before.cell(row, col).value, after.cell(row, col).value
                if expected is None and actual is None:
                    continue
                count += expected is not None
                formulas += isinstance(expected, str) and expected.startswith("=")
                audit.check("legacy_cell", f"{name}!{before.cell(row, col).coordinate}", actual, expected, exact=True)
        preserved[name] = {"last_original_row": end, "backup_max_row": before.max_row,
                           "original_nonempty_cells": count, "original_formula_cells": formulas}
    best = summaries["best"]
    valid = checkpoint_results["best"]["valid_ids"]
    audit.check("workbook_metadata", "update_checkpoint", update["checkpoint"], "best")
    audit.check("workbook_metadata", "update_completed", update["completed_best"], len(valid))
    audit.check("workbook_metadata", "update_finalized", update["matrix_finalized"], best["matrix_finalized"])
    audit.check("workbook_metadata", "update_reference", update["reference"], anchor)

    def cells(ws, row, mapping, metric):
        for col, path in mapping.items():
            audit.check("workbook_metric" if metric else "workbook_pending_blank", f"{ws.title}!{col}{row}",
                        ws[f"{col}{row}"].value, dig(metric, path) if metric else None)

    for index, arm in enumerate(arms, 1):
        aid = arm["id"]
        m = sources["best"].get(aid) if aid in valid else None
        ov, te, su = (wb[name] for name in ("实验矩阵总览", "教师汇报视图", "汇总对比"))
        ro, rt, rs = markers[ov.title] + index, markers[te.title] + index, markers[su.title] + index + 2
        audit.check("workbook_metadata", aid + "/overview_row", update["overview_rows"][aid], ro)
        audit.check("workbook_metadata", f"{ov.title}!A{ro}", ov[f"A{ro}"].value, f"{PREFIX}·{aid} {Path(arm['config']).stem}")
        audit.check("workbook_metadata", f"{ov.title}!AK{ro}", ov[f"AK{ro}"].value, arm["run_name"])
        audit.check("workbook_metadata", f"{te.title}!B{rt}/id", te[f"B{rt}"].value.startswith(aid + " "), True)
        audit.check("workbook_metadata", f"{su.title}!A{rs}", su[f"A{rs}"].value, aid)
        audit.check("workbook_metadata", f"{su.title}!O{rs}", su[f"O{rs}"].value, arm["run_name"])
        for ws, row, col in ((ov, ro, "E"), (te, rt, "D")):
            audit.check("workbook_metadata", f"{ws.title}!{col}{row}/best", "checkpoint=best" in ws[f"{col}{row}"].value, True)
        cells(ov, ro, OVERVIEW, m)
        cells(te, rt, TEACHER, m)
        cells(su, rs, COMPARISON, m)
        negative = f"{m['aggregate']['r2_negative_cases']}/{m['aggregate']['n_cases']}" if m else None
        audit.check("workbook_metric" if m else "workbook_pending_blank", f"{ov.title}!L{ro}", ov[f"L{ro}"].value, negative)
        parent = arm.get("baseline", arm.get("parent", "B0"))
        pm = sources["best"]["B0"] if parent in {"B0", "A5", "B0=A5"} else sources["best"].get(parent) if parent in valid else None
        for col, reference in (("E", sources["best"]["B0"]), ("F", pm)):
            expected = m["field_casebalanced"]["r2"] - reference["field_casebalanced"]["r2"] if m and reference else None
            audit.check("workbook_delta" if m else "workbook_pending_blank", f"{su.title}!{col}{rs}", su[f"{col}{rs}"].value, expected)
    su = wb["汇总对比"]
    reference_row = markers[su.title] + 2
    audit.check("workbook_metadata", f"{su.title}!A{reference_row}", su[f"A{reference_row}"].value, "B0=A5")
    cells(su, reference_row, COMPARISON, sources["best"]["B0"])
    wb.close()
    old.close()
    drift = [path for path, expected in audit.inputs.items()
             if not Path(path).is_file() or hashlib.sha256(Path(path).read_bytes()).hexdigest() != expected]
    audit.check("snapshot", "input_hash_drift", drift, [], exact=True)
    passed = not audit.differences
    complete = all(r["valid_count"] == len(arms) and r["matrix_finalized_reported"] for r in checkpoint_results.values())
    output = args.output.resolve()
    if str(output) in audit.inputs or output.suffix != ".json":
        raise ValueError("Output must be a separate JSON audit file")
    result = {
        "schema": "v6_followup_delivery_verification_v1",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "passed": passed, "scope": "final" if complete else "partial",
        "all_15_best_and_last_reported_finalized": complete,
        "limitations": ["Checks faithful transcription from original metrics; does not recompute evaluator outputs.",
                        "Matrix-finalized flags are recorded from summaries; queue/source provenance is audited separately.",
                        "Pending metrics are checked as absent/blank and never converted to zero."],
        "checkpoints": checkpoint_results, "workbook_checkpoint": "best",
        "metric_semantics": {"overview_AA_AB": "log_z pooled MAE/RMSE", "overview_M_N": "Pa pooled RMSE/MAE",
                             "teacher_O": "log_z pooled p99_pred_true_ratio", "comparison_L": "Pa pooled p99_pred_true_ratio"},
        "original_workbook_preservation": preserved,
        "checks_by_category": dict(audit.counts), "total_checks": sum(audit.counts.values()),
        "difference_count": len(audit.differences), "differences": audit.differences,
        "input_files": [{"path": path, "sha256": digest} for path, digest in sorted(audit.inputs.items())],
        "input_hashes_unchanged_during_audit": not drift, "missing_original_metric_files": missing_sources,
    }
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"output": str(output), "passed": passed, "scope": result["scope"],
                      "best": checkpoint_results["best"]["valid_count"], "last": checkpoint_results["last"]["valid_count"],
                      "checks": result["total_checks"], "differences": len(audit.differences)}, ensure_ascii=False))
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
