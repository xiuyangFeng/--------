#!/usr/bin/env python3
"""Verify the actual xlsx, then export representative Calc-rendered previews."""
import copy
import json
import math
import re
from pathlib import Path

import uno
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter as letter
from PIL import ImageFont

HERE=Path(__file__).resolve().parent
SOURCE=HERE/'重构待验收.xlsx'


def prop(name,value):
 p=uno.createUnoStruct('com.sun.star.beans.PropertyValue');p.Name=name;p.Value=value
 return p


def main():
 wb=load_workbook(SOURCE)
 manifest=json.loads((HERE/'build_manifest.json').read_text())
 normalized=json.loads((HERE/'normalized_rows.json').read_text())
 report={'sheet_names':wb.sheetnames,'errors':[],'numeric_cells':0,'preserved_values':0,'preview_samples':[]}
 assert wb.sheetnames==['WSS实验矩阵','速度与压力实验矩阵','指标说明']
 font=ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',14)
 for sheet in wb:
  assert len(sheet._tables)==1
  assert sheet.freeze_panes
  for row in sheet:
   for c in row:
    if c.data_type=='e' or (c.data_type=='f' and '#REF!' in c.value):report['errors'].append([sheet.title,c.coordinate,c.value])
    if isinstance(c.value,str) and re.fullmatch(r'#{3,}',c.value):report['errors'].append([sheet.title,c.coordinate,c.value])
    if isinstance(c.value,(int,float)):
     report['numeric_cells']+=1
     assert math.isfinite(c.value)
     precision=6 if '0.000000' in c.number_format else 4
     s=f'{c.value:.{precision}f}'
     if '%' in c.number_format:s=f'{c.value*100:.2f}%'
     available=sheet.column_dimensions[c.column_letter].width*7-12
     if font.getsize(s)[0]>available:report['errors'].append([sheet.title,c.coordinate,'numeric width',s,available])
 for item in manifest['sheets']:
  sheet=wb[item['sheet']]
  assert sheet.max_row==item['records']+5
  assert len(set(c[1] for c in item['columns']))==len(item['columns'])
  for rn,row in enumerate(normalized[sheet.title],6):
   for cn,col in enumerate(item['columns'],1):
    v=row.get(col[0])
    if isinstance(v,(int,float)) and math.isfinite(v):
     actual=sheet.cell(rn,cn).value
     if not isinstance(actual,(int,float)) or not math.isclose(actual,v,rel_tol=1e-14,abs_tol=1e-14):report['errors'].append([sheet.title,rn,col[0],v,actual])
     report['preserved_values']+=1
 # Small export workbook uses the actual result cells, formats, widths and heights.
 preview=Workbook();preview.remove(preview.active)
 samples=[]
 for name,ids,cp in [('WSS实验矩阵',None,None),('速度与压力实验矩阵',['R5P','P00','P01','P02','P03'],None),('速度与压力实验矩阵',['R5V','V00','V01','V06','V07'],None)]:
  src=wb[name]
  selected=[]
  for rn in range(6,src.max_row+1):
   if ids is None:
    if len(selected)<8:selected.append(rn)
   elif src.cell(rn,2).value in ids and '注意力' in str(src.cell(rn,1).value):selected.append(rn)
  samples.append((src,selected,'WSS核心' if ids is None else ('压力核心' if 'P02' in ids else '速度核心'),list(range(1,15))))
 # Supplement preview explicitly opens the velocity / longwave columns.
 src=wb['速度与压力实验矩阵'];info=manifest['sheets'][1]
 keys=['name','task','checkpoint','extra:vector_rmse','extra:axial_r2_cb','extra:radial_r2_cb','extra:circ_r2_cb','extra:longwave_20mm_mse','extra:longwave_40mm_mse']
 selected_cols=[next(i for i,col in enumerate(info['columns'],1) if col[0]==k) for k in keys]
 selected=[rn for rn in range(6,src.max_row+1) if src.cell(rn,2).value in ['P02','V07'] and '注意力' in str(src.cell(rn,1).value)]
 samples.append((src,selected,'向量与长波',selected_cols))
 for src,selected,title,colnums in samples:
  dst=preview.create_sheet(title)
  for destrow,source_row in enumerate([1,2,3,4,5]+selected,1):
   for destcol,source_col in enumerate(colnums,1):
    old=src.cell(source_row,source_col);new=dst.cell(destrow,destcol,old.value)
    if old.has_style:
     new.font=copy.copy(old.font);new.fill=copy.copy(old.fill)
     new.border=copy.copy(old.border);new.number_format=old.number_format
     new.protection=copy.copy(old.protection)
    new.alignment=copy.copy(old.alignment)
   dst.row_dimensions[destrow].height=src.row_dimensions[source_row].height
  for rn in [1,2,3]:
   dst.merge_cells(start_row=rn,start_column=1,end_row=rn,end_column=len(colnums))
   old=src.cell(rn,1);new=dst.cell(rn,1)
   new.font=copy.copy(old.font);new.fill=copy.copy(old.fill);new.alignment=copy.copy(old.alignment)
  if title!='向量与长波':
   dst.merge_cells('A4:E4');dst.merge_cells('F4:N4')
  if title=='向量与长波':
   dst.cell(1,1,'速度与压力  |  向量与长波补充指标')
   dst.cell(2,1,'长波 σ=20/40 mm；压力仅内部，单位 Pa²；速度仅轴向分量，单位 (m/s)²。')
   dst.cell(3,1,'来自最终工作簿实际单元格；— 表示不适用，表中没有用 WSS 名称代替体场指标。')
  for ci,oldcol in enumerate(colnums,1):
   dst.column_dimensions[letter(ci)].width=src.column_dimensions[letter(oldcol)].width
  dst.sheet_view.showGridLines=False
  dst.sheet_properties.pageSetUpPr.fitToPage=True
  dst.page_setup.orientation='landscape';dst.page_setup.paperSize=dst.PAPERSIZE_A3
  dst.page_setup.fitToWidth=1;dst.page_setup.fitToHeight=1
  dst.print_area=f'A1:{letter(len(colnums))}{dst.max_row}'
  dst.page_margins=copy.copy(src.page_margins)
  report['preview_samples'].append({'sheet':src.title,'rows':selected,'columns':colnums})
 previewpath=HERE/'版式核查样张.xlsx';preview.save(previewpath)
 local=uno.getComponentContext()
 resolver=local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver',local)
 ctx=resolver.resolve('uno:socket,host=localhost,port=2097;urp;StarOffice.ComponentContext')
 desktop=ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop',ctx)
 doc=desktop.loadComponentFromURL(SOURCE.as_uri(),'_blank',0,(prop('Hidden',True),prop('ReadOnly',True),prop('UpdateDocMode',0)))
 try:
  report['calc_opened']=doc is not None
  report['calc_sheet_names']=list(doc.Sheets.ElementNames)
  report['calc_sample_display']={}
  for item in manifest['sheets']:
   sheet=doc.Sheets.getByName(item['sheet'])
   report['calc_sample_display'][item['sheet']]=[sheet.getCellByPosition(c,5).String for c in range(5,13)]
   data=sheet.getCellRangeByPosition(0,5,len(item['columns'])-1,item['records']+4).getDataArray()
   for rn,row in enumerate(data,6):
    for cn,value in enumerate(row,1):
     if isinstance(value,str) and (re.fullmatch(r'#{3,}',value) or value in ['#REF!','#VALUE!','#DIV/0!','#NAME?']):report['errors'].append([item['sheet'],rn,cn,value])
 finally:doc.close(True)
 doc=desktop.loadComponentFromURL(previewpath.as_uri(),'_blank',0,(prop('Hidden',True),prop('ReadOnly',True),prop('UpdateDocMode',0)))
 try:doc.storeToURL((HERE/'版式核查样张.pdf').as_uri(),(prop('FilterName','calc_pdf_Export'),prop('Overwrite',True)))
 finally:doc.close(True)
 (HERE/'validation_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
 print(json.dumps({k:v for k,v in report.items() if k!='preview_samples'},ensure_ascii=False))
 assert not report['errors'],report['errors']


if __name__=='__main__':main()
