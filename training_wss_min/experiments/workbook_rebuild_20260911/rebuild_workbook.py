#!/usr/bin/env python3
"""Build the reviewed workbook from the preserved extraction snapshots.

Run in this directory after extracting wss_audit.json, volume_records.json and
metric_definitions_audit.json. The original workbook is never used as a template:
each result sheet has exactly one header/schema for all its experiment matrices.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter as letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
OUTPUT = HERE / '重构待验收.xlsx'
NAVY, TEAL, BLUE = '17324D', '167C80', '326496'
INK, MUTED, LIGHT, WHITE = '243748', '63788B', 'F1F6FA', 'FFFFFF'
FONT = 'Noto Sans CJK SC'
FLOAT = '0.0000;[Red]-0.0000;0.0000'
ERROR = '#,##0.0000;[Red]-#,##0.0000;0.0000'
PERCENT = '0.00%;[Red]-0.00%;0.00%'
DELTA = '+0.0000;[Red]-0.0000;0.0000'

# (key, display label, width, numeric format); identical order on both result sheets.
COMMON = [
 ('group','实验矩阵',24,None), ('name','实验 / 模型',40,None),
 ('task','预测对象',13,None), ('checkpoint','检查点 / 汇总',17,None),
 ('scope','评估范围 / 划分',24,None),
 ('r2_cb','物理 R²\n病例等权 cb',16,FLOAT),
 ('rmse','标量 RMSE\npooled · 行内单位',19,ERROR),
 ('mae','标量 MAE\npooled · 行内单位',19,ERROR),
 ('nmae_pooled','标量 NMAE\npooled · 极差归一',19,PERCENT),
 ('case_median_r2','病例 R²\n中位数',16,FLOAT),
 ('case_p10_r2','病例 R²\nP10',16,FLOAT),
 ('negative','负 R² 病例\n数量 / 总数',17,None),
 ('delta_r2','ΔR² cb\n相对指定参照',17,DELTA),
 ('reference','对照实验 / 口径',35,None),
 ('r2_pooled','物理 R²\npooled',16,FLOAT),
 ('case_mean_r2','病例 R²\n算术均值',16,FLOAT),
 ('nrmse_pooled','range-NRMSE\npooled',19,PERCENT),
 ('nrmse_case_mean','range-NRMSE\ncase-mean',19,PERCENT),
 ('nmae_case_mean','NMAE\ncase-mean',19,PERCENT),
 ('high_r2','物理高值区\nR² · top10%',19,FLOAT),
 ('top10_ratio','物理 top10\n幅值比 · pooled',19,FLOAT),
 ('top10_iou','top10 IoU\ncase-mean',19,FLOAT),
 ('p99_ratio','物理 p99 比\npooled',19,FLOAT),
]
OLD_KEYS = {'G':'r2_cb','H':'r2_pooled','I':'case_mean_r2','J':'case_median_r2',
 'K':'case_p10_r2','L':'negative','M':'rmse','N':'mae','O':'nrmse_pooled',
 'P':'nrmse_case_mean','Q':'nmae_pooled','R':'nmae_case_mean','S':'high_r2',
 'T':'top10_ratio','U':'top10_iou'}
WSS_DETAIL = [
 ('norm_r2_cb','归一化 R²\n病例等权 cb',18,FLOAT),
 ('norm_r2_pooled','归一化 R²\npooled',18,FLOAT),
 ('norm_case_mean','归一化病例 R²\n均值',19,FLOAT),
 ('norm_case_med','归一化病例 R²\n中位数',19,FLOAT),
 ('norm_case_p10','归一化病例 R²\nP10',19,FLOAT),
 ('norm_mae','归一化 MAE\npooled · 无量纲',21,FLOAT),
 ('norm_rmse','归一化 RMSE\npooled · 无量纲',21,FLOAT),
 ('norm_nrmse_pool','归一化 range-NRMSE\npooled',24,PERCENT),
 ('norm_nrmse_case','归一化 range-NRMSE\ncase-mean',24,PERCENT),
 ('norm_nmae_pool','归一化 NMAE\npooled',20,PERCENT),
 ('norm_nmae_case','归一化 NMAE\ncase-mean',20,PERCENT),
 ('norm_spearman','归一化 Spearman\ncase-mean',23,FLOAT),
 ('norm_high_spearman','归一化高值区 Spearman\ncase-mean',27,FLOAT),
 ('norm_top10','归一化 top10\n幅值比 · pooled',23,FLOAT),
 ('norm_p99','归一化 p99 比\npooled',20,FLOAT),
]
OLD_NORM = dict(zip(['V','W','X','Y','Z','AA','AB','AC','AD','AE','AF','AG','AH','AI','AJ'],[c[0] for c in WSS_DETAIL]))
FIT = []
for space, unit in [('phys','行内单位'),('norm','无量纲')]:
 for agg in ['cb','pooled']:
  for metric in ['r2','a','b']:
   title = f'{"物理" if space=="phys" else "归一化"} fit {"R²" if metric=="r2" else metric}\n{agg}'
   if metric == 'b': title += f' · {unit}'
   FIT.append((f'{space}_fit_{metric}_{agg}',title,24 if metric=='b' else 20,FLOAT))
OLD_FIT = dict(zip(['AL','AM','AN','AO','AP','AQ','AR','AS','AT','AU','AV','AW'],[c[0] for c in FIT]))
TAIL = [('units','标量单位',14,None),('model','骨干 / 模型',45,None),
 ('architecture','结构 / 唯一变量',70,None),('sampling','采样 / 训练协议',58,None),
 ('points','每例支撑点数',18,None),('run_id','原始实验 ID',65,None),
 ('notes','实验备注（完整说明见批注）',76,None),('sources','结果来源（完整路径见批注）',68,None),
 ('old_location','原表位置',36,None)]


def load(name):
 return json.loads((HERE/name).read_text())


def text_value(v):
 if isinstance(v,(dict,list)): return json.dumps(v,ensure_ascii=False)
 return str(v) if v is not None else ''


def short(v, n=200):
 s=text_value(v)
 return s if len(s)<=n else s[:n-1]+'…'


def result_rows_wss(audit):
 rows=[]
 for record in audit['records']:
  v=record['values']
  row={new:v.get(old) for old,new in {**OLD_KEYS,**OLD_NORM,**OLD_FIT}.items()}
  kind=record.get('record_type','run')
  cp=record.get('checkpoint') or '历史未标注'
  if kind not in ['run','single_run']:
   cp += ' · '+(str(record.get('extra_metrics',{}).get('seed数',''))+' seed均值' if kind=='seed_mean' else {'comparison':'对比统计'}.get(kind,kind))
  row.update(group=record.get('group','历史矩阵'),name=v.get('A',''),task='WSS · Pa',
   checkpoint=cp,scope=v.get('B',''),units='Pa',model=v.get('C'),architecture=v.get('D'),
   sampling=v.get('E'),points=v.get('F'),run_id=v.get('AK'),
   notes='；'.join(map(text_value,record.get('notes',[]))),
   sources='\n'.join(map(text_value,record.get('source_paths',[]))),
   old_location=f"{record.get('source_sheet','')}!{record.get('source_row','')}",
   _original=record,_kind=kind)
  for k,val in record.get('extra_metrics',{}).items():
   if k=='Pa p99比':row['p99_ratio']=val
   else:row['extra:'+k]=val
  comparisons=record.get('comparisons',[])
  if comparisons:
   blocks=[]
   for pair_index,comparison in enumerate(comparisons,1):
    named=comparison.get('named_values',{})
    values={re.sub(r'\s*\[[A-Z]+\]$','',k):v for k,v in named.items()}
    reference=comparison.get('reference')
    if reference in ('unknown','未知','—'):reference=None
    if not reference:
     reference=next((values[k] for k in ['对照臂','参照','父臂'] if values.get(k) not in [None,'','—']),None)
    reference=comparison.get('primary_delta_reference') or reference
    if reference in ['unknown','未知','—']:reference=None
    delta=comparison.get('primary_delta')
    if reference and delta is not None and row.get('delta_r2') is None:
     row['reference']=text_value(reference);row['delta_r2']=delta
    context=f"{comparison.get('source_sheet','')}!{comparison.get('source_row','')}｜参照：{text_value(reference) if reference else '原对比上下文，未明确配对ID'}"
    lines=[context]
    row[f'extra:配对{pair_index}｜参照与来源']=context
    for label,val in named.items():
     source_column=re.search(r'\[([A-Z]+)\]$',label)
     ref=comparison.get('reference_by_column',{}).get(source_column[1] if source_column else '',reference)
     clean=re.sub(r'\s*\[[A-Z]+\]$','',label)
     if clean=='Δ' and source_column:
      headers=list(comparison.get('headers',{}).items())
      pos=next((i for i,x in enumerate(headers) if x[0]==source_column[1]),0)
      clean='Δ'+(headers[pos-1][1] if pos else '原表指标')
     if 'Δ' in clean or any(k in clean for k in ['95%','胜/负','学到的','判定','last−','last-','sd','std','标准差','秒','时长','耗时','参数量','seed数','种子数']):
      key=f'extra:配对{pair_index}｜{clean}'
      if key in row:key+=' '+(source_column[0] if source_column else '')
      row[key]=val
      row.setdefault('_cell_notes',{})[key]=f'{context}\n本指标参照：{ref}\n原表头：{label}\n'+comparison.get('reference_evidence','')
     lines.append(f'{label}：{text_value(val)}'+(f'（参照：{ref}）' if 'Δ' in label else ''))
    blocks.append('\n'.join(lines))
   row['extra:配对比较（逐项注明参照）']='\n\n'.join(blocks)
  rows.append(row)
 order={}
 for row in rows:
  group=row['group'].replace('｜多seed汇总','')
  row['group']=group
  order.setdefault(group,len(order))
 return sorted(rows,key=lambda r:(order[r['group']],r.get('_kind')=='seed_mean',r.get('run_id') or r['name'],r['checkpoint'].startswith('last')))


def result_rows_volume(data):
 records=data['records'] if isinstance(data,dict) else data
 rows=[]
 for record in records:
  m=record.get('metrics',{})
  row=dict(m)
  task=record.get('task','')
  units=record.get('units') or record.get('unit') or ('Pa' if '压' in task or task=='pressure' else 'm/s')
  notes=record.get('notes',[])
  ex=record.get('extras',{})
  row.update(group=record.get('group','体场矩阵'),name=record.get('name') or record.get('id'),
   task=('压力 · Pa' if units=='Pa' else '速度 · m/s'),checkpoint=record.get('snapshot') or record.get('checkpoint'),
   scope=' · '.join(filter(None,[record.get('split'),record.get('scope')])),units=units,model=record.get('model'),
   architecture=record.get('architecture') or ex.get('hypothesis'),sampling=record.get('protocol'),points=record.get('points') or 5000,
   run_id=record.get('run_dir') or record.get('run_id') or record.get('id'),reference=record.get('reference') or ex.get('parent'),
   delta_r2=record.get('delta_r2') if record.get('delta_r2') is not None else ex.get('delta_r2_parent'),
   notes='；'.join(map(text_value,notes)) if isinstance(notes,list) else text_value(notes),
   sources='\n'.join(map(text_value,record.get('sources',[]))),
   old_location=text_value(record.get('old_locations',[])),_original=record,_kind='run')
  neg=m.get('negative_count'); total=m.get('test_count')
  if neg is not None: row['negative']=f'{int(neg)} / {int(total)}' if total is not None else f'{int(neg)} / 待核'
  skip={'parent','delta_r2_parent','hypothesis','normalization','query_groups','longwave_summary','longwave_method','gate_details','legacy_summary','changes_vs_parent'}
  for k,val in ex.items():
   if k not in skip:row['extra:'+k]=val
  # Preserve any source metric beyond the common schema; never silently discard it.
  known={c[0] for c in COMMON}|{'negative_count','test_count','delta_r2'}
  for k,val in m.items():
   if k not in known:row['extra:'+k]=val
  rows.append(row)
 return sorted(rows,key=lambda r:('历史' not in r['group'],r['group'],r['task'],r['name'],r['checkpoint']!='best'))


def extra_label(name):
 labels={'vector_rmse':'速度向量 RMSE\npooled · m/s','vector_rmse_cb':'速度向量 RMSE\n病例等权 · m/s',
  'vector_rmse_diagnostic':'诊断向量 RMSE\nm/s','direction_cosine':'方向余弦\ncase-mean',
  'direction_cosine_speedweighted':'速度加权方向余弦','component_normalized_r2_pooled':'三分量归一化 R²\npooled',
  'direction_speed_floor_m_s':'方向统计速度下限\nm/s','parameters':'参数量',
  'eval_seconds':'全场评估耗时\ns','train_seconds':'训练耗时\ns',
  'train_peak_cuda_allocated_mib':'训练峰值显存\nMiB','eval_peak_cuda_allocated_mib':'评估峰值显存\nMiB',
  'gate':'开发筛选判定','point_count':'评估点数',
  'original_pooled_normalized_mae':'历史归一化 pooled MAE\n原完整指标',
  'original_pooled_normalized_rmse':'历史归一化 pooled RMSE\n原完整指标'}
 if name in labels:return labels[name]
 if name.startswith('longwave_'):
  match=re.match(r'longwave_(\d+)mm_mse(_raw)?',name)
  if match:return f'长波 σ={match[1]} mm\n{"未去偏" if match[2] else "去偏"} MSE · 行内单位²'
 norm=name.startswith('normalized_');key=name[len('normalized_'):] if norm else name
 prefix='归一化 ' if norm else '物理 '
 for axis,title in [('axial','轴向'),('radial','径向'),('circ','周向'),('u','u 分量'),('v','v 分量'),('w','w 分量')]:
  if key.startswith(axis+'_'):
   metric=key[len(axis)+1:].replace('r2','R²').replace('rmse','RMSE').replace('_cb','\n病例等权').replace('_pooled','\npooled')
   return title+' '+metric+(' · m/s' if 'RMSE' in metric else '')
 terms={'r2_cb':'R² cb','r2_pooled':'R² pooled','r2_case_mean':'病例 R² 均值','r2_case_median':'病例 R² 中位数','r2_case_p10':'病例 R² P10',
  'mae':'MAE pooled','rmse':'RMSE pooled','mae_cb':'MAE 病例等权','rmse_cb':'RMSE 病例等权',
  'mae_case_mean':'MAE case-mean','rmse_case_mean':'RMSE case-mean','nmae_pooled':'NMAE pooled','nmae_case_mean':'NMAE case-mean',
  'nrmse_pooled':'range-NRMSE pooled','nrmse_case_mean':'range-NRMSE case-mean','high_value_r2':'高值区 R²',
  'spearman':'Spearman case-mean','spearman_high_value':'高值区 Spearman','top10_ratio':'top10 幅值比 pooled',
  'top10_iou':'top10 IoU case-mean','p99_ratio':'p99 比 pooled'}
 if key.startswith('fit_'):
  title=key.replace('fit_r2','fit R²').replace('fit_slope','fit a').replace('fit_intercept','fit b').replace('_',' ')
 else:title=terms.get(key,key)
 unit=''
 if any(k in key for k in ['mae','rmse','intercept']) and not any(k in key for k in ['nmae','nrmse']):unit=' · 无量纲' if norm else ' · 行内单位'
 return prefix+title.replace(' ','\n',1)+unit


def attach_volume_conclusions(rows,inventory):
 counts=Counter()
 for old in inventory['体场注意力结论']['rows']:
  v=old['values'];rn=old['row']
  if 4<=rn<=10:
   for row in rows:
    if '注意力' in row['group'] and row['name']==v['A'] and row['checkpoint']=='best' and ((v['B']=='压力')==(row['units']=='Pa')):
     row['extra:汇报定位']=v.get('G')
  if not 13<=rn<=24:continue
  target,reference=[x.strip() for x in v['A'].split('/')]
  matches=[row for row in rows if '注意力' in row['group'] and row['name']==target and row['checkpoint']==v['C'] and ((v['B']=='压力')==(row['units']=='Pa'))]
  assert len(matches)==1,(rn,matches)
  row=matches[0];identity=(row['name'],row['task'],row['checkpoint']);counts[identity]+=1
  prefix=f'结论配对{counts[identity]}｜'
  ref=f"{reference}｜{v['B']}｜{v['C']}"
  row['extra:'+prefix+'参照']=ref
  for c,label in [('D','ΔR² cb'),('E','主要误差降低'),('F','σ20mm MSE降低'),('G','σ40mm MSE降低'),('H','精度门槛')]:
   key='extra:'+prefix+label;row[key]=v.get(c)
   row.setdefault('_cell_notes',{})[key]=f'原体场注意力结论!{c}{rn}；参照={ref}；压力主要误差=MAE，速度主要误差=向量RMSE；误差降低=(参照−当前)/参照。'
  row['_original'].setdefault('conclusion_comparisons',[]).append(old)


def extra_columns(rows, definitions=None):
 definitions=definitions or {}
 keys=list(dict.fromkeys(k for r in rows for k in r if k.startswith('extra:')))
 out=[]
 for key in keys:
  name=key[6:]
  meta=definitions.get(name,{}) if isinstance(definitions,dict) else {}
  label=meta.get('label',extra_label(name)) if isinstance(meta,dict) else extra_label(name)
  numeric=all(isinstance(r.get(key),(int,float)) or r.get(key) is None for r in rows)
  fmt=FLOAT if numeric else None
  if numeric and any(s in name for s in ['参数量','parameters','count','epoch','seed数']):fmt='#,##0'
  if numeric and ('百分比' in name or 'reduction' in name):fmt=PERCENT
  if numeric and '降低' in name:fmt=PERCENT
  if numeric and ('Δ' in name or 'delta' in name):fmt=DELTA
  if numeric and ('nmae' in name.lower() or 'nrmse' in name.lower() or 'range-NRMSE' in name):fmt=PERCENT
  if numeric and name.startswith('longwave_'):fmt='0.000000;[Red]-0.000000;0.000000'
  if re.search('[\u4e00-\u9fff]',name):label=name
  out.append((key,label,max(22,min(42,len(label)*1.15)),fmt))
 return out


def add_comment(cell, value):
 content=text_value(value)
 if content:
  cell.comment=Comment(content[:32000],'来源核对 · 2026-09-11')
  cell.comment.width=640;cell.comment.height=280


def init_sheet(ws,title,subtitle,info,cols):
 ws.sheet_view.showGridLines=False
 ws.sheet_view.zoomScale=80
 ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.sheet_properties.outlinePr.summaryRight=True
 ws.sheet_properties.tabColor=TEAL if ws.title.startswith('WSS') else BLUE
 ws.freeze_panes='D6'
 ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=min(14,len(cols)))
 ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=min(14,len(cols)))
 ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=min(14,len(cols)))
 for rn,txt in [(1,title),(2,subtitle),(3,info)]:
  c=ws.cell(rn,1,txt);c.fill=PatternFill('solid',fgColor=NAVY if rn==1 else LIGHT)
  c.font=Font(name=FONT,size=20 if rn==1 else 11,bold=rn==1,color=WHITE if rn==1 else MUTED)
  c.alignment=Alignment(vertical='center',wrap_text=True,indent=1)
  ws.row_dimensions[rn].height={1:42,2:32,3:28}[rn]
 ws.row_dimensions[4].height=27;ws.row_dimensions[5].height=58
 for i,(key,label,width,fmt) in enumerate(cols,1):
  ws.column_dimensions[letter(i)].width=width
  c=ws.cell(5,i,label);c.fill=PatternFill('solid',fgColor=NAVY)
  c.font=Font(name=FONT,size=11,bold=True,color=WHITE)
  c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
 ws.print_options.horizontalCentered=True
 ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
 ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_title_rows='4:5';ws.print_title_cols='A:C'
 ws.page_margins.left=0.25;ws.page_margins.right=0.25
 ws.page_margins.top=0.35;ws.page_margins.bottom=0.35
 ws.oddFooter.center.text='&A  |  第 &P / &N 页'
 ws.oddFooter.right.text='2026-09-11'
 return ws


def band(ws,start,end,title,color):
 if start>end:return
 ws.merge_cells(start_row=4,start_column=start,end_row=4,end_column=end)
 c=ws.cell(4,start,title);c.fill=PatternFill('solid',fgColor=color)
 c.font=Font(name=FONT,size=11,bold=True,color=WHITE)
 c.alignment=Alignment(horizontal='left',vertical='center',indent=1)


def make_result_sheet(wb,name,rows,detail,extra_defs=None):
 extras=extra_columns(rows,extra_defs)
 cols=COMMON+detail+extras+TAIL
 ws=wb.create_sheet(name)
 n=len(rows);groups=len(set(r.get('group') for r in rows))
 init_sheet(ws,name+'  |  实验矩阵与结果',
  ('每行一个实验 / 检查点；按实验矩阵筛选。物理 WSS 单位 Pa；归一化与物理空间分列；种子均值明确标注。' if name.startswith('WSS') else '每行一个实验 / 检查点；按矩阵与对象筛选。压力为相对压力（Pa）；速度标量为 |u|（m/s）；向量与长波指标见右侧扩展列。'),
  f'{groups} 组 · {n} 条记录  |  — = 未产出或不适用；数字保留完整精度，显示 4 位。右侧 + 展开详细指标与来源；打印为 A:N 核心结果。',cols)
 band(ws,1,5,'01  实验与评估协议',TEAL)
 band(ws,6,14,'02  主要结果与指定参照',BLUE)
 band(ws,15,23,'03  物理标量：完整误差、拟合与高值区指标',TEAL)
 enddetail=23+len(detail)
 if detail:band(ws,24,enddetail,'04  归一化与线性拟合 · 展开查看',BLUE)
 if extras:band(ws,enddetail+1,enddetail+len(extras),'05  补充诊断与历史对比 · 展开查看',TEAL)
 tailstart=enddetail+len(extras)+1
 band(ws,tailstart,len(cols),'06  设置、备注与可追溯来源 · 展开查看',BLUE)
 border=Side(style='thin',color='D9E4EC')
 previous=None
 for rn,row in enumerate(rows,6):
  new_group=row['group']!=previous;previous=row['group']
  ws.row_dimensions[rn].height=72
  for cn,(key,label,width,fmt) in enumerate(cols,1):
   val=row.get(key)
   if isinstance(val,float) and not math.isfinite(val):val=None
   original=val
   if val is None or val=='':val='—'
   if isinstance(val,(dict,list)):val=text_value(val)
   if isinstance(val,str) and len(val)>220:val=short(val,220)
   c=ws.cell(rn,cn,val)
   c.fill=PatternFill('solid',fgColor=WHITE if rn%2 else LIGHT)
   c.font=Font(name=FONT,size=10,color=INK,bold=cn in (2,6))
   c.border=Border(bottom=border,top=Side(style='medium',color='A4BECF') if new_group else Side())
   c.alignment=Alignment(horizontal='right' if isinstance(val,(float,int)) else ('center' if cn in [3,4,12] else 'left'),vertical='center',wrap_text=not isinstance(val,(float,int)),indent=1 if cn in [1,2] else 0)
   if fmt and isinstance(val,(int,float)):c.number_format=fmt
   if cn==3:
    c.font=Font(name=FONT,size=10,bold=True,color=TEAL if row['units']=='Pa' else BLUE)
   if row.get('_kind')=='seed_mean':c.fill=PatternFill('solid',fgColor='E8F4F3')
   if isinstance(original,(dict,list)) or (isinstance(original,str) and (len(original)>220 or key in ['sources','notes','scope','architecture'])):add_comment(c,original)
   if key in row.get('_cell_notes',{}):add_comment(c,row['_cell_notes'][key])
   if key=='name':add_comment(c,row.get('_original',{}))
  # Every table has one consistent numeric/label schema, no inserted mini-tables.
  ws.cell(rn,1).alignment=Alignment(wrap_text=True,vertical='center',indent=1)
  for key in ['group','scope']:
   ci=next(i for i,c in enumerate(cols,1) if c[0]==key)
   c=ws.cell(rn,ci)
   if len(text_value(row.get(key)))>65:
    add_comment(c,row[key]);c.value=short(row[key],65)
  source=row.get('sources','')
  candidates=re.findall(r'/[^\n;；]+',source)
  if candidates:
   path=Path(candidates[0])
   if path.exists():
    ci=next(i for i,c in enumerate(cols,1) if c[0]=='sources')
    ws.cell(rn,ci).hyperlink=path.as_uri()
 last=5+n
 table=Table(displayName='WSSResults' if name.startswith('WSS') else 'VolumeResults',ref=f'A5:{letter(len(cols))}{last}')
 table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=False,showColumnStripes=False)
 ws.add_table(table)
 for start,end in [(24,enddetail),(enddetail+1,enddetail+len(extras)),(tailstart,len(cols))]:
  if start<=end:
   ws.column_dimensions.group(letter(start),letter(end),outline_level=1,hidden=True)
   # openpyxl groups ranges in one ColumnDimension; set individual widths again.
   for i in range(start,end+1):
    d=ws.column_dimensions[letter(i)];d.width=cols[i-1][2];d.hidden=True;d.outlineLevel=1
   if end<len(cols):ws.column_dimensions[letter(end+1)].collapsed=True
 ws.column_dimensions[letter(len(cols)+1)].width=3
 ws.column_dimensions[letter(len(cols)+1)].collapsed=True
 ws.print_area=f'A1:N{last}'
 audited={d['key']:d for d in load('metric_definitions_audit.json')['definitions']}
 for c in ws[5]:
  key=cols[c.column-1][0]
  d=audited.get(key, audited.get(key.replace('extra:',''),{}))
  definition='\n'.join(text_value(d.get(k)) for k in ['definition','unit','note'] if d.get(k))
  if isinstance(extra_defs,dict):definition+='\n'+text_value(extra_defs.get(key.replace('extra:',''),''))
  add_comment(c,f'字段：{key}\n{definition}\n显示精度不会改变存储数值。详见「指标说明」。')
 ci=next(i for i,c in enumerate(cols,1) if c[0]=='delta_r2')
 if n:
  ws.conditional_formatting.add(f'{letter(ci)}6:{letter(ci)}{last}',CellIsRule(operator='greaterThan',formula=['0'],font=Font(color='127666')))
  ws.conditional_formatting.add(f'{letter(ci)}6:{letter(ci)}{last}',CellIsRule(operator='lessThan',formula=['0'],font=Font(color='B44646')))
 return {'sheet':name,'records':n,'groups':groups,'columns':cols,'rows':rows}


def make_definitions(wb,old_inventory,definition_audit,wss,volume):
 ws=wb.create_sheet('指标说明')
 cols=[('category','类别',22,None),('metric','指标 / 主题',34,None),('definition','计算定义 / 阅读说明',76,None),('direction','方向 / 单位',30,None),('notes','备注与边界',76,None)]
 init_sheet(ws,'指标说明  |  统一口径与阅读指南','WSS 与体场采用相同的标量报告框架；专属向量、长波及历史统计单独注明。','数据截止 2026-09-11；保留历史未知项；不跨划分、空间、检查点或种子汇总方式直接计算增益。',cols)
 ws.freeze_panes='C6';ws.sheet_view.zoomScale=85
 band(ws,1,5,'口径、来源与使用规则',TEAL)
 rows=[
  ['阅读指南','工作簿结构','WSS实验矩阵：全部 WSS 实验及可识别的种子汇总。速度与压力实验矩阵：压力和速度同页、按对象筛选，best / last 分行。','3 个页面','已合并原总览、教师视图、汇总对比及体场注意力结论；删除 RCR Oracle 独立页，相关历史实验仍保留在 WSS 矩阵。'],
  ['阅读指南','详细指标与来源','结果页冻结前 3 列与表头；列 O:W 为物理补充指标；右上方 + 可展开归一化、线性拟合、额外诊断与来源。表头筛选支持实验组、对象、检查点。','显示 4 位小数','原始数值未截断；百分比保存为比例。— 表示未产出或不适用；历史数据缺项未填 0。长文字在批注中完整保留。'],
  ['阅读指南','对照与 Δ','只使用源记录明确指定的父臂/参照。没有可靠参照的 Δ 保留 —。同 seed、同划分、同评估区域、同 checkpoint 才能直接配对。','Δ = 当前 − 指定参照','R² 的正 Δ 为改善；不得按行号或整个测试集首行自动选参照。种子均值不是一条单 seed run。'],
  ['阅读指南','best / last','best 保留各实验原有选模规则；last 为训练终点。联合任务 best 按联合损失选择，与单任务 best 选择规则不同。','不从两者择优','联合对单任务主比较采用 last；历史未标注 checkpoint 明确保留未知状态。'],
  ['阅读指南','打印与电子表格','结果页设置 A3 横向，打印 A:N 核心结果，重复表头；全部明细以电子表格形式保留。','打印核心列','需要打印扩展指标时，可先展开对应列并重新设置打印区域。'],
 ]
 for d in definition_audit['definitions']:
  rows.append(['统一指标口径',d['title'],d['definition'],d.get('unit','')+'；'+d.get('direction',''),d.get('note','')])
 rows += [
  ['体场定义','压力标量','p_rel = p − 病例峰值时刻的体积加权平均压力；常规主指标在壁面∪内部评估；历史 wall/internal 分域独立保留。','Pa','压力不是绝对压力；同一病例的压力参考常数固定。'],
  ['体场定义','速度标量 |u|','先由预测三分量计算模长 ||u_pred||₂，再与真实速度模长比较；常规评估范围为内部。','m/s','标量 RMSE 与三分量向量 RMSE 的定义不同，必须分列。'],
  ['体场定义','向量 RMSE','sqrt(mean(||u_pred − u_true||₂²))，评估三分量向量误差。','m/s · 越低越好','轴向/径向/周向 R² 是额外分量指标，不能用模长 R² 替代。'],
  ['体场定义','线性 z 归一化','压力和速度标量使用线性 z 空间时，R² 与对应物理空间一致；MAE / RMSE 的单位及数值随尺度变化。','R² 无量纲','WSS log-z 是非线性变换，其物理 R² 与归一化 R² 不能互换。'],
  ['长波诊断','20 / 40 mm','在分支内按 5 mm 分箱并以体积加权，去掉病例体积均值误差后，按 Gaussian σ=20/40 mm 平滑残差，计算 MSE 后对病例等权。','压力 Pa²；速度 (m/s)²','压力长波只在内部；速度长波只用轴向分量。20/40 mm 是高斯标准差，不是严格波长或频带。'],
  ['长波诊断','额外改善比例','(参照误差 − 当前误差) / 参照误差；正数表示误差下降。','百分比','压力主误差是 MAE；速度主误差是向量 RMSE。长波筛选必须同时满足全场精度保护要求。'],
  ['数据来源','重构与备份','重构输入快照、完整提取记录、原工作簿备份和验证报告保存于 training_wss_min/experiments/workbook_rebuild_20260911/。','可追溯','旧脚本依赖旧页面名；不要再用旧版逐页追加器覆盖本文件。后续须沿用本次统一列定义写入。'],
 ]
 # Original final conclusions are narrative evidence, not a fourth duplicate score table.
 for r in old_inventory['体场注意力结论']['rows']:
  if r['row'] in [26,28,30,32,34,36,38]:
   title=next(x['values']['A'] for x in old_inventory['体场注意力结论']['rows'] if x['row']==r['row']-1)
   rows.append(['体场矩阵结论',title,r['values']['A'],'单 seed 开发筛选','完整数字见速度与压力实验矩阵；旧表结论按原实验报告保留。'])
 for rn,row in enumerate(rows,6):
  ws.row_dimensions[rn].height=max(82,min(225,26*math.ceil(max(len(str(row[2]))/42,len(str(row[4]))/42))))
  for cn,v in enumerate(row,1):
   c=ws.cell(rn,cn,v);c.fill=PatternFill('solid',fgColor=LIGHT if rn%2==0 else WHITE)
   c.font=Font(name=FONT,size=11,color=INK,bold=cn==2)
   c.alignment=Alignment(wrap_text=True,vertical='center')
   c.border=Border(bottom=Side(style='thin',color='D9E4EC'))
 table=Table(displayName='MetricDefinitions',ref=f'A5:E{ws.max_row}')
 table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False)
 ws.add_table(table);ws.print_area=f'A1:E{ws.max_row}'
 return len(rows)


def show_fit_columns(wb):
 """Keep all physical/normalized fit R², a and b visible by default."""
 for ws in wb.worksheets[:2]:
  fitcols={c.column for c in ws[5] if c.value and 'fit' in str(c.value).lower()}
  # Explicit single-column dimensions prevent overlapping group spans from
  # hiding the fit columns in spreadsheet readers that resolve spans differently.
  for key,dimension in ws.column_dimensions.items():
   from openpyxl.utils import column_index_from_string
   index=column_index_from_string(key)
   dimension.min=index;dimension.max=index
   if index in fitcols:
    dimension.hidden=False;dimension.outlineLevel=0;dimension.collapsed=False
  for merged in list(ws.merged_cells.ranges):
   if merged.min_row!=4 or merged.max_row!=4:continue
   if not any(c in fitcols for c in range(merged.min_col,merged.max_col+1)):continue
   original=ws.cell(4,merged.min_col).value or '补充指标'
   start,end=merged.min_col,merged.max_col
   ws.unmerge_cells(str(merged))
   def category(c):
    if c not in fitcols:return 'other'
    return 'norm' if '归一化' in str(ws.cell(5,c).value) else 'physical'
   run=start
   for c in range(start+1,end+2):
    if c<=end and category(c)==category(run):continue
    kind=category(run)
    title={'physical':'线性拟合｜物理空间：R² / a / b（直接显示）',
           'norm':'线性拟合｜归一化空间：R² / a / b（直接显示）'}.get(kind,original)
    band(ws,run,c-1,title,BLUE if kind=='norm' else TEAL)
    run=c
  ws.cell(3,1).value=str(ws.cell(3,1).value).replace(
   '右侧 + 展开详细指标与来源', 'fit R² / a / b 直接显示；其余扩展指标与来源用 + 展开')
 if '指标说明' in wb.sheetnames:
  ws=wb['指标说明']
  for row in ws.iter_rows(min_row=6):
   if row[1].value=='详细指标与来源':
    row[2].value='结果页冻结前 3 列与表头；列 O:W 为物理补充指标。物理及归一化的 fit R²、fit a、fit b 全部直接显示；其余扩展指标和来源可用 + 展开。可按实验组、对象和检查点筛选。'


def main():
 wss=load('wss_audit.json');volume=load('volume_records.json')
 definitions=load('metric_definitions_audit.json')
 inventory=load('original_inventory.json')
 wb=Workbook();wb.remove(wb.active)
 wr=result_rows_wss(wss);vr=result_rows_volume(volume)
 attach_volume_conclusions(vr,inventory)
 reports=[make_result_sheet(wb,'WSS实验矩阵',wr,WSS_DETAIL+FIT),
          make_result_sheet(wb,'速度与压力实验矩阵',vr,[],volume.get('field_definitions',{}))]
 n=make_definitions(wb,inventory,definitions,wss,volume)
 show_fit_columns(wb)
 wb.active=0
 wb.properties.title='WSS、速度场与压力场实验矩阵及结果汇总'
 wb.properties.subject='按任务与实验矩阵统一指标、检查点、单位及来源'
 wb.properties.creator='GNN Research';wb.properties.description='2026-09-11 重构；完整精度、单一表头、来源可追溯。'
 wb.save(OUTPUT)
 manifest=[]
 for report in reports:
  manifest.append({k:v for k,v in report.items() if k!='rows'})
 (HERE/'build_manifest.json').write_text(json.dumps({'sheets':manifest,'definitions':n,'output':str(OUTPUT)},ensure_ascii=False,indent=2))
 # Normalized rows make all original numeric mappings independently auditable.
 (HERE/'normalized_rows.json').write_text(json.dumps({r['sheet']:r['rows'] for r in reports},ensure_ascii=False,indent=2))
 print(json.dumps({'output':str(OUTPUT),'counts':[(r['sheet'],r['records'],len(r['columns'])) for r in reports]},ensure_ascii=False))


if __name__=='__main__':main()
