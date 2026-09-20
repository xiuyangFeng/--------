"""Publish the final interpretation without changing validated scientific results."""
from __future__ import annotations

import copy
import fcntl
import hashlib
import json
import math
import os
from pathlib import Path
import tempfile
import sys

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from training_wss_min.tools.volume_attention_common import EXP, save_json, sha256, stamp
from training_wss_min.tools import update_volume_attention_xlsx as U
from training_wss_min.tools.tracker_section import replace_or_append
from finalize import TRACKER, CHANGELOG, guarded_write

SUMMARY_SHEET = "体场注意力结论"
TITLE = "体场注意力26臂｜最终结论（2026-09-11）"


def make_summary(report):
    rows = {(r["id"], r["task"], r["checkpoint"]): r for r in report["rows"]}
    pairs = {(p["id"], p["task"], p["reference"], p["checkpoint"]): p for p in report["comparisons"]}
    get = lambda aid, task, ck="best": rows[(aid, task, ck)]
    pair = lambda aid, task, ref, ck="best": pairs[(aid, task, ref, ck)]
    p2, p3, p9 = (get(a, "pressure") for a in ("P02", "P03", "P09"))
    v7, v11 = (get(a, "velocity") for a in ("V07", "V11"))
    p2d = pair("P02", "pressure", "P00")
    v7d = pair("V07", "velocity", "V06")
    jv = pair("J01", "velocity", "J00")
    jr, j0 = get("J01", "velocity"), get("J00", "velocity")
    conclusions = [
        ("压力优先P02", f"P02（原R5结构＋直接瓶颈全局注意力）best/last R²_cb={p2['r2_cb']:.5f}/{get('P02','pressure','last')['r2_cb']:.5f}，best MAE={p2['mae']:.3f} Pa、RMSE={p2['rmse']:.3f} Pa。相对同期P00，best ΔR²={p2d['delta_r2']:+.5f}、MAE降低{100*p2d['mae_reduction']:.2f}%；20/40mm去病例偏差的平滑残差MSE降低{100*p2d['lw20_reduction']:.2f}%/{100*p2d['lw40_reduction']:.2f}%，精度和长波门槛均通过。相对历史R5P，best R²增加{pair('P02','pressure','R5P')['delta_r2']:.5f}。"),
        ("局部分支与多半径未进一步提高压力R²", f"P03的best/last MAE={p3['mae']:.3f}/{get('P03','pressure','last')['mae']:.3f} Pa，为本轮压力误差较低的备选；best R²={p3['r2_cb']:.5f}低于P02，且相对P01的40mm改善仅{100*pair('P03','pressure','P01')['lw40_reduction']:.2f}%。P03优于条件编码P04和同容量FFN P05，支持这一配置中跨token交互的价值。P09相对P08通过精度门槛，但best R²={p9['r2_cb']:.5f}仍低于P00的{get('P00','pressure')['r2_cb']:.5f}；P10同半径三支的R²还略高于P09，未证实不同半径带来额外收益。"),
        ("速度V07可作候选，模块增益未过父臂门槛", f"V07的幅值best/last R²_cb={v7['r2_cb']:.5f}/{get('V07','velocity','last')['r2_cb']:.5f}，best向量RMSE={v7['vector_rmse']:.6f} m/s。相对V00和历史R5V的全场精度均通过筛选；但相对直接父臂V06，向量RMSE仅降低{100*v7d['vector_rmse_reduction']:.2f}%，未达到3%。V07相对V06的20/40mm残差虽下降{100*v7d['lw20_reduction']:.2f}%/{100*v7d['lw40_reduction']:.2f}%，仍不能越过全场保护门槛宣称该模块通过长波筛选。新增瓶颈采用逐token FFN的V11幅值R²={v11['r2_cb']:.5f}，与V07接近；V09相对V08幅值R²下降、长波残差增加。所有速度臂均未通过各自父臂的完整预登记门槛。"),
        ("联合模型未达到两任务共同提升", f"J01相对J00的压力通过精度和长波门槛；速度幅值best R²增加{jv['delta_r2']:.5f}，但向量RMSE仅降低{100*jv['vector_rmse_reduction']:.2f}%，周向R²下降{j0['circ_r2']-jr['circ_r2']:.5f}，超过0.01保护线。按联合对单任务的last400主口径，J01压力/速度R²={get('J01','pressure','last')['r2_cb']:.5f}/{get('J01','velocity','last')['r2_cb']:.5f}，均低于P03/V03的{get('P03','pressure','last')['r2_cb']:.5f}/{get('V03','velocity','last')['r2_cb']:.5f}；本轮不建议用联合模型替代对应单任务模型。"),
        ("病例和分域仍有差异", "P02相对P00有25/34例R²提高、27/34例MAE降低；病例R²中位数0.75536→0.79597，P10由0.33443→0.64119，负R²病例2→0。AAA/ILO域R²分别由0.76416/0.62343提高至0.82043/0.67799；AG域由0.81313降至0.80082，不能称所有域均改善。V07相对V00有28/34例幅值R²提高、31/34例向量RMSE降低。完整逐病例表和分域指标保留，低压差病例同时按Pa误差判读。"),
        ("资源与验收", "26次400轮训练全部完成，14040耗时12小时26分51秒；14041汇总与14056监控均COMPLETED/0:0。单任务阶段单卡四路、联合阶段两路共享显存；单任务训练约104–113分钟，联合约60分钟，不能用不同并发条件下的耗时比归因模型加速。联合真实预检8例平均9968.375个query/例（每任务5000），仅代表该batch；全场联合查询15,353,714点，两个独立任务合计29,476,133点。60组任务/checkpoint指标、26条history验收通过，6936行逐病例配对差值与2040张病例剖面图齐全。"),
        ("结论边界", "固定seed1234、18维输入、train138/test34、峰值1162；test34已暴露，仅用于开发筛选。长波是分支内5mm分箱及5/10/20/40mm高斯尺度诊断，不是严格频带或精度上限。父臂模块增益、相对同期基线的模型效果和相对历史模型的差值分开判读；不据单seed声称统计显著或稳定泛化。"),
    ]
    focus = [("P02", "pressure", "压力R²优先"), ("P03", "pressure", "压力MAE备选"),
             ("P09", "pressure", "多半径主候选"), ("V07", "velocity", "速度幅值最高；父臂门槛未过"),
             ("V11", "velocity", "无BT容量对照"), ("J01", "pressure", "联合压力"), ("J01", "velocity", "联合速度")]
    lines = ["# 压力/速度全局注意力：26臂最终汇总", "", f"复核日期：{stamp()}。全部26次训练、60组任务/checkpoint结果及工作簿已完成核验。", "",
             "| 重点模型 | 任务 | best R²_cb | last R²_cb | best MAE | best向量RMSE | 定位 |",
             "|---|---|---:|---:|---:|---:|---|"]
    model_rows = []
    for aid, task, note in focus:
        b, l = get(aid, task), get(aid, task, "last")
        unit = "Pa" if task == "pressure" else "m/s（幅值）"
        values = [aid, "压力" if task == "pressure" else "速度", b["r2_cb"], l["r2_cb"], b["mae"], b["vector_rmse"], note, unit]
        model_rows.append(values)
        lines.append(f"| {aid} | {values[1]} | {b['r2_cb']:.5f} | {l['r2_cb']:.5f} | {b['mae']:.6f} {unit} | {b['vector_rmse']:.6f} | {note} |" if b["vector_rmse"] is not None else f"| {aid} | 压力 | {b['r2_cb']:.5f} | {l['r2_cb']:.5f} | {b['mae']:.3f} Pa | — | {note} |")
    for title, body in conclusions:
        lines.extend(["", "## " + title, "", body])
    lines.extend(["", "[全部best/last结果](results.md) · [预设配对比较](comparisons.csv) · [逐病例差值](per_case_comparisons.csv) · [参数/显存/耗时](resources.csv) · [联合query成本](joint_query_costs.csv)", "",
                  "![全部模型best精度](comparison.png)", "", "原始科学报告[report.json](report.json)保持不变；工作簿使用[最终展示副本](final_delivery_report.json)，仅将历史参照行的门槛文字改为‘历史参照，不适用’，不修改数值。", ""])
    compare_rows = []
    for aid, task, ref, ck in (("P02","pressure","P00","best"), ("P02","pressure","R5P","best"),
                              ("P03","pressure","P01","best"), ("P09","pressure","P08","best"),
                              ("V07","velocity","V06","best"), ("V07","velocity","V00","best"),
                              ("V07","velocity","R5V","best"), ("V09","velocity","V08","best"),
                              ("J01","pressure","J00","best"), ("J01","velocity","J00","best"),
                              ("J01","pressure","P03","last"), ("J01","velocity","V03","last")):
        p = pair(aid, task, ref, ck)
        error = p["mae_reduction"] if task == "pressure" else p["vector_rmse_reduction"]
        compare_rows.append([aid+" / "+ref, "压力" if task=="pressure" else "速度", ck, p["delta_r2"],
                             error, p["lw20_reduction"], p["lw40_reduction"], "精度通过" if p["qualified"] else "精度未通过"])
    return "\n".join(lines), conclusions, model_rows, compare_rows


def write_summary_sheet(workbook, conclusions, model_rows, compare_rows):
    ws = workbook.create_sheet(SUMMARY_SHEET)
    expected = {}
    def row(values, *, merge=False, header=False):
        index = ws.max_row + 1 if expected else 1
        for col, value in enumerate(values, 1):
            cell = ws.cell(index, col, value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.font = Font(name="Calibri", size=11, bold=header, color="FFFFFF" if header else "172B4D")
            if header:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.number_format = "0.000000" if isinstance(value, float) else "General"
            expected[cell.coordinate] = value
        if merge:
            ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=8)
        ws.row_dimensions[index].height = 30 if header else 45
        return index
    row([TITLE], merge=True, header=True)
    row(["26个新训练｜seed1234｜单seed、暴露test34开发筛选｜best与last分别报告"], merge=True)
    row(["模型", "任务", "best R²_cb", "last R²_cb", "best MAE", "best向量RMSE m/s", "定位", "MAE单位"], header=True)
    for values in model_rows:
        row(values)
    row(["配对比较：压力主要误差=MAE，速度主要误差=向量RMSE；正百分比表示误差降低；20/40mm均为去病例偏差后残差"], merge=True, header=True)
    row(["候选 / 参照", "任务", "checkpoint", "ΔR²_cb", "主要误差降低", "20mm MSE降低", "40mm MSE降低", "完整精度门槛"], header=True)
    for values in compare_rows:
        index = row(values)
        for col in (5, 6, 7):
            ws.cell(index, col).number_format = "0.00%"
    for title, body in conclusions:
        row([title], merge=True, header=True)
        index = row([body], merge=True)
        ws.row_dimensions[index].height = 100
    link_row = row(["详细数值见原三张结果表；点击查看完整结论、诊断与逐病例资料。历史R5门槛不适用。"], merge=True)
    ws.cell(link_row, 1).hyperlink = "../../training_wss_min/experiments/volume_attention_20260910/final_summary.md"
    ws.freeze_panes = "C4"
    for letter, width in zip("ABCDEFGH", (24, 12, 17, 17, 18, 21, 43, 25)):
        ws.column_dimensions[letter].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:H{ws.max_row}"
    return expected


def publish_workbook(display, conclusions, model_rows, compare_rows):
    book = U.BOOK
    original = book.read_bytes()
    before = hashlib.sha256(original).hexdigest()
    backup = EXP / f"{book.stem}_before_final_summary_{before[:16]}.xlsx"
    if not backup.exists():
        backup.write_bytes(original)
    elif backup.read_bytes() != original:
        raise ValueError("Backup collision")
    old_acceptance = EXP / "xlsx_acceptance.json"
    if old_acceptance.exists():
        evidence = EXP / ("xlsx_acceptance_before_delivery_" + sha256(old_acceptance)[:16] + ".json")
        if not evidence.exists():
            evidence.write_bytes(old_acceptance.read_bytes())
    fd, temp = tempfile.mkstemp(prefix=".volume_final_summary_", suffix=".xlsx", dir=book.parent)
    os.close(fd)
    candidate = Path(temp)
    try:
        candidate.write_bytes(original)
        workbook = load_workbook(candidate, data_only=False)
        if SUMMARY_SHEET in workbook.sheetnames:
            if workbook[SUMMARY_SHEET]["A1"].value != TITLE:
                raise ValueError("Refusing to replace an unowned conclusion sheet")
            del workbook[SUMMARY_SHEET]
        positions, historical = U.update_in_memory(workbook, display)
        expected = write_summary_sheet(workbook, conclusions, model_rows, compare_rows)
        workbook.save(candidate)
        workbook.close()
        checked = load_workbook(candidate, data_only=False)
        if checked.sheetnames != historical["sheetnames"] + [SUMMARY_SHEET]:
            raise ValueError("Unexpected worksheet order")
        for address, value in expected.items():
            cell = checked[SUMMARY_SHEET][address]
            equal = math.isclose(cell.value, value, rel_tol=1e-12, abs_tol=1e-12) if type(value) in (float, int) else cell.value == value
            if not equal or cell.data_type == "f":
                raise ValueError(f"Summary transcription mismatch: {address}")
        del checked[SUMMARY_SHEET]  # Only in-memory: verify all original sheets separately.
        verification = U.verify_readback(checked, display, positions, historical)
        checked.close()
        if sha256(book) != before:
            raise RuntimeError("Workbook changed concurrently")
        os.chmod(candidate, book.stat().st_mode & 0o7777)
        candidate.replace(book)
    finally:
        candidate.unlink(missing_ok=True)
    acceptance = dict(book=str(book), report=str(EXP / "final_delivery_report.json"), backup=str(backup),
                      before_sha256=before, after_sha256=sha256(book), report_sha256=sha256(EXP / "final_delivery_report.json"),
                      scientific_report_sha256=sha256(EXP / "report.json"), complete=True, positions=positions,
                      best_rows_overview=30, best_rows_teacher=30, best_last_rows_summary=60,
                      joint_vs_single_primary_checkpoint="last", all_result_cells_literal=True,
                      row_insertions=0, row_deletions=0, **verification,
                      summary_sheet=SUMMARY_SHEET, summary_cell_checks=len(expected), summary_source=str(EXP / "final_summary.md"),
                      timestamp=stamp(), presentation_only_changes="4 historical report rows / 8 workbook gate cells relabeled; new owned summary sheet")
    save_json(EXP / "xlsx_acceptance.json", acceptance)
    return acceptance


def main():
    with (EXP / ".reporting.lock").open("a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        report = json.loads((EXP / "report.json").read_text())
        if not report["complete"] or len(report["validation"]) != 60 or len(report["history_checks"]) != 26:
            raise ValueError("Complete validated matrix required")
        if not all(v["passed"] for v in report["validation"].values()) or not all(v["passed"] for v in report["history_checks"].values()):
            raise ValueError("Scientific acceptance failed")
        original_report_hash = sha256(EXP / "report.json")
        display = copy.deepcopy(report)
        for row in display["rows"]:
            if row["id"].startswith("R5"):
                row["gate"] = "历史参照，不适用"
        display["delivery_presentation"] = dict(source_sha256=original_report_hash, numerical_changes=0,
                                               reason="Historical references have no candidate screening gate", updated_at=stamp())
        save_json(EXP / "final_delivery_report.json", display)
        summary, conclusions, model_rows, compare_rows = make_summary(report)
        (EXP / "final_summary.md").write_text(summary)
        acceptance = publish_workbook(display, conclusions, model_rows, compare_rows)
        heading = "### 17.1 最终判读与选择（2026-09-11）"
        section = heading + "\n\n" + "\n\n".join(f"**{title}。** {body}" for title, body in conclusions)
        section += "\n\n完整[结论摘要](../../../training_wss_min/experiments/volume_attention_20260910/final_summary.md)与[工作簿](../../03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx)已同步；工作簿新增“体场注意力结论”页，原三张表完整保留30/30/60行结果，历史参照门槛已标为不适用。\n"
        old = TRACKER.read_text()
        guarded_write(TRACKER, old, replace_or_append(old, heading, section))
        readme = EXP / "README.md"
        old = readme.read_text()
        short = "## 最终结论（2026-09-11）\n\n压力优先P02：best/last R²_cb 0.74911/0.75283，相对P00通过精度和长波门槛；P03保留为低MAE备选。速度V07幅值R²最高（0.77625），相对V00及历史R5V精度提高，但相对V06的模块增益未通过向量RMSE门槛。联合J01未达到两任务共同提升。\n\n详见[完整结论与证据](final_summary.md)。工作簿新增“体场注意力结论”页，最终回读证据见[xlsx_acceptance.json](xlsx_acceptance.json)。全部26臂仅seed1234，结论限定为暴露test34开发筛选。\n"
        guarded_write(readme, old, replace_or_append(old, "## 最终结论（2026-09-11）", short))
        old = CHANGELOG.read_text()
        old_heading = "## 2026-09-10｜压力/速度26臂完成、独立验收与工作簿回填"
        new_heading = "## 2026-09-11｜压力/速度26臂完成、最终判读与工作簿回填"
        updated = old.replace(old_heading, new_heading)
        log_section = new_heading + "\n\n26臂400轮、60组目标/checkpoint和26条history验收均通过；14040耗时12:26:51，14041与14056完成退出0。压力P02通过精度和长波门槛，P03为低MAE备选；速度V07相对同期/历史基线更好，但各速度臂均未通过直接父臂完整门槛；联合J01不构成两任务共同提升。三张原结果表完整回填，新增“体场注意力结论”页，修正历史参照的待完成显示。原始科学报告不改数值；源表3000项及结论页" + str(acceptance["summary_cell_checks"]) + "项转录复核通过，历史公式/合并保留。详见[跟踪§17](WSS_PINN/WSS_V5_训练实验跟踪.md)与[完整结论](../../training_wss_min/experiments/volume_attention_20260910/final_summary.md)。单seed1234、暴露test34边界保留。\n"
        guarded_write(CHANGELOG, old, replace_or_append(updated, new_heading, log_section))
        handoff = json.loads((EXP / "handoff_status.json").read_text())
        handoff.update(timestamp=stamp(), implementation="complete", formal_status="complete", completed_arms=26, active=[],
                       finalizer_status="complete", monitor_status="complete", notes="26 runs and final metrics complete; final interpretation and workbook independently checked.",
                       final_summary=str(EXP / "final_summary.md"), workbook_acceptance=str(EXP / "xlsx_acceptance.json"))
        save_json(EXP / "handoff_status.json", handoff)
        if sha256(EXP / "report.json") != original_report_hash:
            raise ValueError("Scientific report unexpectedly changed")
        save_json(EXP / "final_publication.json", dict(status="complete", timestamp=stamp(), scientific_report_unchanged=True,
                  scientific_report_sha256=original_report_hash, source_sha256=sha256(__file__),
                  workbook_sha256=acceptance["after_sha256"], result_cells_checked=acceptance["transcription_checks"],
                  summary_cells_checked=acceptance["summary_cell_checks"], historical_formulas=acceptance["historical_formulas"],
                  files=[str(EXP / "final_summary.md"), str(TRACKER), str(EXP / "README.md"), str(CHANGELOG), str(U.BOOK)]))
        print(json.dumps(acceptance, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
