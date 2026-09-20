"""Append/update the owned pressure/velocity attention sections without moving history.

The report is the only result source. Overview and teacher sheets contain 30
best rows; the comparison sheet contains all 60 best/last rows. All result cells
are literals. The candidate workbook must pass independent readback and history
checks before an atomic replacement of the original file.
"""
from __future__ import annotations

import argparse
from copy import copy
import hashlib
import json
import math
import os
from pathlib import Path
import tempfile

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from training_wss_min.tools.update_v5_rerun_xlsx import BOOK


ROOT = Path(__file__).resolve().parents[2]
EXP = ROOT / "training_wss_min/experiments/volume_attention_20260910"
REPORT = EXP / "report.json"
NAME = "volume_attention_20260910"
TITLE = "2026-09-10｜体场速度/压力全局注意力矩阵（26个新训练，seed1234）"
END = "volume_attention_20260910｜本节结束"
SHEETS = ("实验矩阵总览", "教师汇报视图", "汇总对比")
PAIR_ORDER = ([('R5P', 'pressure')] + [(f'P{i:02d}', 'pressure') for i in range(12)]
              + [('R5V', 'velocity')] + [(f'V{i:02d}', 'velocity') for i in range(12)]
              + [(aid, task) for aid in ('J00', 'J01') for task in ('pressure', 'velocity')])
EXPECTED_KEYS = {(aid, task, ck) for aid, task in PAIR_ORDER for ck in ("best", "last")}
COLUMNS = (
    ("id", "实验ID"), ("task", "预测任务"), ("unit", "标量误差单位"),
    ("checkpoint", "checkpoint"), ("status", "状态"), ("parent", "父臂"),
    ("r2_cb", "物理 R²_cb"), ("delta_r2_parent", "ΔR²_cb vs 同任务/同checkpoint父臂"),
    ("mae", "标量 pooled MAE（本行单位）"), ("rmse", "标量 pooled RMSE（本行单位）"),
    ("case_median", "病例 R² 中位"), ("case_p10", "病例 R² P10"),
    ("negative_cases", "负 R² 病例数 / 34"), ("vector_rmse", "速度向量 RMSE（m/s）"),
    ("axial_r2", "轴向速度 R²_cb"), ("radial_r2", "径向速度 R²_cb"),
    ("circ_r2", "周向速度 R²_cb"), ("lw20", "长波20mm残差MSE（Pa² / 轴向速度(m/s)²）"),
    ("lw40", "长波40mm残差MSE（Pa² / 轴向速度(m/s)²）"), ("parameters", "模型参数量"),
    ("train_seconds", "训练秒数"), ("eval_seconds", "评估秒数"),
    ("gate", "预设门槛判定"), ("run_dir", "结果目录"), ("comparison_note", "比较口径"),
)
WIDTH = len(COLUMNS)
NUMERIC_FIELDS = {key for key, _ in COLUMNS} - {
    "id", "task", "unit", "checkpoint", "status", "parent", "gate", "run_dir", "comparison_note"}
NOTE = ("train138/test34；单seed1234探索；test34为已暴露开发集。压力=p−病例峰值体积均压，"
        "壁面∪内部；速度标量=预测三分量的幅值，仅内部。物理与线性z的标量R²相同。"
        "总览/教师仅best；汇总保留best和last。联合与单任务的主比较看last，"
        "joint的best由联合总损失选择，不与单任务best混作相同选模准则。"
        "长波按分支5mm体积分箱，去病例体积均值偏差，Gaussian σ=20/40mm残差MSE后病例等权；"
        "压力单位Pa²，速度仅轴向分量、单位(m/s)²。空值表示未产出/不适用；门槛以report为准。")


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _tag(kind: str) -> str:
    return f"{NAME}|{kind}"


def _keys(sheet: str) -> list[tuple[str, str, str]]:
    return [(aid, task, ck) for aid, task in PAIR_ORDER
            for ck in (("best", "last") if sheet == "汇总对比" else ("best",))]


def validate_report(report: dict) -> dict[tuple[str, str, str], dict]:
    if type(report.get("complete")) is not bool or not isinstance(report.get("rows"), list):
        raise ValueError("report requires boolean complete and a rows list")
    entries = {}
    for row in report["rows"]:
        key = (row.get("id"), row.get("task"), row.get("checkpoint"))
        if key in entries:
            raise ValueError(f"duplicate report row: {key}")
        if key not in EXPECTED_KEYS:
            raise ValueError(f"unexpected report row: {key}")
        if not isinstance(row.get("status"), str):
            raise ValueError(f"report row lacks status: {key}")
        for field in NUMERIC_FIELDS:
            value = row.get(field)
            if value is not None and (type(value) not in (int, float) or not math.isfinite(value)):
                raise ValueError(f"nonfinite or nonnumeric report value: {key}/{field}")
        entries[key] = row
    if entries.keys() != EXPECTED_KEYS:
        raise ValueError(f"report must include all 60 rows; missing={sorted(EXPECTED_KEYS - entries.keys())}")
    return entries


def _literal(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)
    return value


def row_values(row: dict) -> list:
    task = row["task"]
    derived = {"task": "相对压力（壁面∪内部）" if task == "pressure" else "速度幅值（内部）",
               "unit": "Pa" if task == "pressure" else "m/s",
               "comparison_note": ("联合与单任务主比较看last；best按联合总损失选模"
                                   if row["id"].startswith("J") else "同任务、同checkpoint比较")}
    return [_literal(derived.get(key, row.get(key))) for key, _ in COLUMNS]


def _markers(sheet: str) -> dict[int, str]:
    count = len(_keys(sheet))
    return {0: "title", 1: "header", **{i + 2: f"slot:{i}" for i in range(count)},
            count + 2: "note", count + 3: "end"}


def locate_section(ws) -> int | None:
    starts = [cell.row for cell in ws["A"] if cell.value == TITLE]
    ends = [cell.row for cell in ws["A"] if cell.value == END]
    if not starts and not ends:
        return None
    height = len(_keys(ws.title)) + 4
    if len(starts) != 1 or len(ends) != 1 or ends[0] - starts[0] != height - 1:
        raise ValueError(f"ambiguous or resized owned volume section: {ws.title}")
    first = starts[0]
    for offset, kind in _markers(ws.title).items():
        comment = ws.cell(first + offset, 1).comment
        if comment is None or comment.text != _tag(kind):
            raise ValueError(f"unowned or changed volume row: {ws.title}!{first + offset}")
    return first


def _inside(row, col, rect) -> bool:
    return rect is not None and rect[0] <= row <= rect[1] and 1 <= col <= WIDTH


def snapshot_history(workbook, excluded: dict) -> dict:
    cells, merges = {}, {}
    for ws in workbook.worksheets:
        rect = excluded.get(ws.title)
        for (row, col), cell in ws._cells.items():
            if not _inside(row, col, rect):
                cells[(ws.title, cell.coordinate)] = (copy(cell.value), cell.data_type)
        merges[ws.title] = {str(rng) for rng in ws.merged_cells.ranges
                            if not (rect and rect[0] <= rng.min_row <= rng.max_row <= rect[1]
                                    and rng.max_col <= WIDTH)}
    return {"sheetnames": list(workbook.sheetnames), "cells": cells, "merges": merges}


def _write_row(ws, index: int, values: list, *, kind: str, header=False):
    for col in range(1, WIDTH + 1):
        cell = ws.cell(index, col)
        cell.value = values[col - 1] if col <= len(values) else None
        if isinstance(cell.value, str):
            cell.data_type = "s"  # Explicit literals, even if a supplied label begins with '='.
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.font = Font(name="Calibri", size=10, bold=header)
        cell.fill = PatternFill("solid", fgColor="D9EAF7" if header else "FFFFFF")
        cell.border = Border(bottom=Side(style="thin", color="D9E1F2"))
        cell.number_format = "0.000000" if isinstance(cell.value, float) else "General"
        cell.comment = None
    ws.cell(index, 1).comment = Comment(_tag(kind), "Volume attention results")
    ws.row_dimensions[index].height = 45 if header else 32


def update_in_memory(workbook, report: dict):
    entries = validate_report(report)
    if not set(SHEETS) <= set(workbook.sheetnames):
        raise ValueError("workbook lacks one of the three required result sheets")
    old_starts = {name: locate_section(workbook[name]) for name in SHEETS}
    excluded = {name: (first, first + len(_keys(name)) + 3)
                for name, first in old_starts.items() if first is not None}
    history = snapshot_history(workbook, excluded)
    positions = {}
    for name in SHEETS:
        ws, keys = workbook[name], _keys(name)
        first = old_starts[name] or ws.max_row + 2
        last = first + len(keys) + 3
        for rng in list(ws.merged_cells.ranges):
            overlaps = not (rng.max_row < first or rng.min_row > last or rng.min_col > WIDTH)
            if overlaps:
                if not (first <= rng.min_row <= rng.max_row <= last and rng.max_col <= WIDTH):
                    raise ValueError(f"historical merge crosses owned rectangle: {name}!{rng}")
                ws.unmerge_cells(str(rng))
        _write_row(ws, first, [TITLE], kind="title", header=True)
        ws.cell(first, 1).font = Font(bold=True, color="FFFFFF", size=12)
        ws.cell(first, 1).fill = PatternFill("solid", fgColor="1F4E78")
        ws.merge_cells(start_row=first, start_column=1, end_row=first, end_column=WIDTH)
        _write_row(ws, first + 1, [label for _, label in COLUMNS], kind="header", header=True)
        for slot, key in enumerate(keys):
            _write_row(ws, first + slot + 2, row_values(entries[key]), kind=f"slot:{slot}")
        state = "整批完成" if report["complete"] else "进行中；空值不表示零"
        _write_row(ws, last - 1, ["说明", f"{state}。{NOTE}"], kind="note")
        ws.merge_cells(start_row=last - 1, start_column=2, end_row=last - 1, end_column=WIDTH)
        ws.row_dimensions[last - 1].height = 65
        _write_row(ws, last, [END], kind="end")
        positions[name] = {"first_row": first, "last_row": last, "width": WIDTH,
                           "header_row": first + 1, "data_rows": list(range(first + 2, last - 1))}
    return positions, history


def verify_readback(workbook, report: dict, positions: dict, history: dict) -> dict:
    """Reconstruct expected cells directly from the report, separately from the writer."""
    if workbook.sheetnames != history["sheetnames"]:
        raise AssertionError("historical worksheet names/order changed")
    for (sheet, coordinate), expected in history["cells"].items():
        cell = workbook[sheet][coordinate]
        if (cell.value, cell.data_type) != expected:
            raise AssertionError(f"historical value/formula changed: {sheet}!{coordinate}")
    for sheet, expected in history["merges"].items():
        current = {str(rng) for rng in workbook[sheet].merged_cells.ranges}
        position = positions.get(sheet)
        if position:
            first, last = position["first_row"], position["last_row"]
            owned = {str(rng) for rng in workbook[sheet].merged_cells.ranges
                     if first <= rng.min_row <= rng.max_row <= last and rng.max_col <= WIDTH}
            current -= owned
        if expected != current:
            raise AssertionError(f"historical merge changed: {sheet}")
    entries = {(row["id"], row["task"], row["checkpoint"]): row for row in report["rows"]}
    numeric_checks, literal_checks = 0, 0
    for sheet, position in positions.items():
        ws = workbook[sheet]
        if locate_section(ws) != position["first_row"]:
            raise AssertionError("owned section markers changed during serialization")
        for column, (_, label) in enumerate(COLUMNS, 1):
            if ws.cell(position["header_row"], column).value != label:
                raise AssertionError(f"volume header changed: {sheet}/{column}")
        for index, key in zip(position["data_rows"], _keys(sheet)):
            source = entries[key]
            for column, (field, _) in enumerate(COLUMNS, 1):
                expected = source.get(field)
                if field == "task":
                    expected = "相对压力（壁面∪内部）" if key[1] == "pressure" else "速度幅值（内部）"
                elif field == "unit":
                    expected = "Pa" if key[1] == "pressure" else "m/s"
                elif field == "comparison_note":
                    expected = ("联合与单任务主比较看last；best按联合总损失选模"
                                if key[0].startswith("J") else "同任务、同checkpoint比较")
                elif isinstance(expected, (dict, list, tuple)):
                    expected = json.dumps(expected, ensure_ascii=False, sort_keys=True, allow_nan=False)
                cell = ws.cell(index, column)
                actual = cell.value
                if cell.data_type == "f":
                    raise AssertionError(f"new result is a formula: {sheet}!{cell.coordinate}")
                if type(expected) in (int, float):
                    same = type(actual) in (int, float) and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12)
                    numeric_checks += 1
                else:
                    same = actual == expected
                    literal_checks += 1
                if not same:
                    raise AssertionError(f"report transcription mismatch: {sheet}!{cell.coordinate}")
    values = [value for value, _ in history["cells"].values()]
    return {"transcription_checks": numeric_checks + literal_checks, "numeric_checks": numeric_checks,
            "literal_or_null_checks": literal_checks, "historical_cell_checks": len(values),
            "historical_nonempty_cells": sum(value is not None for value in values),
            "historical_formulas": sum(dtype == "f" for _, dtype in history["cells"].values()),
            "historical_merges": sum(map(len, history["merges"].values())),
            "all_historical_values_formulas_merges_preserved": True}


def _atomic_json(path: Path, payload: dict):
    handle, name = tempfile.mkstemp(prefix=".volume_xlsx_acceptance_", suffix=".json", dir=path.parent)
    os.close(handle)
    candidate = Path(name)
    try:
        candidate.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
        candidate.replace(path)
    finally:
        candidate.unlink(missing_ok=True)


def update_workbook(book=BOOK, report_path=REPORT, experiment_dir=EXP, require_complete=True):
    book, report_path, experiment_dir = Path(book), Path(report_path), Path(experiment_dir)
    report_bytes = report_path.read_bytes()
    report = json.loads(report_bytes)
    validate_report(report)
    if require_complete and not report["complete"]:
        raise RuntimeError("final workbook update requires a complete report; use --allow-pending to register blank rows")
    original = book.read_bytes()
    original_hash, report_hash = hashlib.sha256(original).hexdigest(), hashlib.sha256(report_bytes).hexdigest()
    experiment_dir.mkdir(parents=True, exist_ok=True)
    backup = experiment_dir / f"{book.stem}_backup_before_volume_attention_{original_hash[:16]}{book.suffix}"
    if backup.exists():
        if backup.read_bytes() != original:
            raise RuntimeError("existing backup does not match the original workbook")
    else:
        with backup.open("xb") as handle:
            handle.write(original)
    descriptor, candidate_name = tempfile.mkstemp(prefix=".volume_attention_", suffix=book.suffix, dir=book.parent)
    os.close(descriptor)
    candidate = Path(candidate_name)
    try:
        candidate.write_bytes(original)
        workbook = load_workbook(candidate, data_only=False)
        try:
            positions, history = update_in_memory(workbook, report)
            workbook.save(candidate)
        finally:
            workbook.close()
        check = load_workbook(candidate, data_only=False)
        try:
            checks = verify_readback(check, report, positions, history)
        finally:
            check.close()
        if _hash(book) != original_hash or _hash(report_path) != report_hash:
            raise RuntimeError("workbook or report changed concurrently; refusing replacement")
        os.chmod(candidate, book.stat().st_mode & 0o7777)
        candidate.replace(book)
    finally:
        candidate.unlink(missing_ok=True)
    acceptance = {"book": str(book.resolve()), "report": str(report_path.resolve()),
                  "backup": str(backup.resolve()), "before_sha256": original_hash,
                  "after_sha256": _hash(book), "report_sha256": report_hash,
                  "complete": report["complete"], "positions": positions,
                  "best_rows_overview": 30, "best_rows_teacher": 30, "best_last_rows_summary": 60,
                  "joint_vs_single_primary_checkpoint": "last", "all_result_cells_literal": True,
                  "row_insertions": 0, "row_deletions": 0, **checks}
    _atomic_json(experiment_dir / "xlsx_acceptance.json", acceptance)
    return acceptance


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--book", type=Path, default=BOOK)
    parser.add_argument("--report", type=Path, default=REPORT)
    parser.add_argument("--experiment-dir", type=Path, default=EXP)
    parser.add_argument("--allow-pending", action="store_true")
    args = parser.parse_args(argv)
    print(json.dumps(update_workbook(args.book, args.report, args.experiment_dir,
                                    require_complete=not args.allow_pending), ensure_ascii=False))


if __name__ == "__main__":
    main()
