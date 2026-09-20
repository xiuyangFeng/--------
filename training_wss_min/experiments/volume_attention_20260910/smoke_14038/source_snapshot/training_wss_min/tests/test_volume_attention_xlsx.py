"""Workbook literal transcription, fixed ownership and preservation guarantees."""
from __future__ import annotations

import copy
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from training_wss_min.tools import update_volume_attention_xlsx as W


def report_fixture(complete=True):
    rows = []
    for index, (aid, task) in enumerate(W.PAIR_ORDER):
        for ck in ("best", "last"):
            delta = 0.001 if ck == "last" else 0.0
            row = {"id": aid, "task": task, "checkpoint": ck, "status": "complete",
                   "r2_cb": .65 + index * .001 + delta, "mae": 130.5 if task == "pressure" else .095,
                   "rmse": 320.1 if task == "pressure" else .17,
                   "case_median": .62, "case_p10": .21, "negative_cases": 2,
                   "vector_rmse": .21 if task == "velocity" else None,
                   "axial_r2": .71 if task == "velocity" else None,
                   "radial_r2": .3 if task == "velocity" else None,
                   "circ_r2": .1 if task == "velocity" else None,
                   "lw20": 3600.0 if task == "pressure" else .01,
                   "lw40": 900.0 if task == "pressure" else .005,
                   "parameters": 2100000, "train_seconds": 2200.2, "eval_seconds": 131.0,
                   "run_dir": f"/work/runs/{aid}", "parent": None if aid.startswith("R5") else "P00" if task == "pressure" else "V00",
                   "delta_r2_parent": None if aid.startswith("R5") else .014,
                   "gate": "历史参照" if aid.startswith("R5") else "通过开发筛选"}
            rows.append(row)
    return {"complete": complete, "rows": rows}


def setup_files(tmp_path):
    book, report_path, exp = tmp_path / "results.xlsx", tmp_path / "report.json", tmp_path / "exp"
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("汇总对比", "教师汇报视图", "RCR Oracle", "指标说明", "实验矩阵总览"):
        ws = wb.create_sheet(name)
        ws["A1"] = "历史标题"
        ws.merge_cells("A1:C1")
        ws["D2"] = "=SUM(A3:A4)"
        ws["A3"] = .123456789012345
        ws["A4"] = 2
        ws.column_dimensions["A"].width = 42
        ws.row_dimensions[3].height = 23
    wb.save(book)
    wb.close()
    report_path.write_text(json.dumps(report_fixture(), ensure_ascii=False))
    return book, report_path, exp


def column(field):
    return [key for key, _ in W.COLUMNS].index(field) + 1


def row_for(acceptance, sheet, aid, task, ck="best"):
    keys = W._keys(sheet)
    return acceptance["positions"][sheet]["data_rows"][keys.index((aid, task, ck))]


def test_atomic_update_writes_30_best_30_best_60_all_with_correct_units(tmp_path):
    book, report_path, exp = setup_files(tmp_path)
    before = book.read_bytes()
    evidence = W.update_workbook(book, report_path, exp)
    assert Path(evidence["backup"]).read_bytes() == before
    assert evidence["transcription_checks"] == (30 + 30 + 60) * W.WIDTH
    assert evidence["all_historical_values_formulas_merges_preserved"]
    assert evidence["historical_formulas"] == 5
    assert evidence["row_insertions"] == evidence["row_deletions"] == 0
    assert json.loads((exp / "xlsx_acceptance.json").read_text())["after_sha256"] == W._hash(book)
    wb = load_workbook(book, data_only=False)
    try:
        for sheet, count in zip(W.SHEETS, (30, 30, 60)):
            assert len(evidence["positions"][sheet]["data_rows"]) == count
            assert wb[sheet]["D2"].value == "=SUM(A3:A4)"
            assert "A1:C1" in {str(rng) for rng in wb[sheet].merged_cells.ranges}
            assert wb[sheet].column_dimensions["A"].width == 42
            assert wb[sheet].row_dimensions[3].height == 23
            for task, aid, unit in (("pressure", "P01", "Pa"), ("velocity", "V01", "m/s")):
                index = row_for(evidence, sheet, aid, task)
                assert wb[sheet].cell(index, column("unit")).value == unit
            joint = row_for(evidence, sheet, "J00", "velocity")
            assert "last" in wb[sheet].cell(joint, column("comparison_note")).value
            header = evidence["positions"][sheet]["header_row"]
            headers = [wb[sheet].cell(header, i).value for i in range(1, W.WIDTH + 1)]
            assert not any("WSS" in label or "log_z" in label for label in headers)
            assert "Pa²" in headers[column("lw20") - 1]
        assert {wb["汇总对比"].cell(i, column("checkpoint")).value
                for i in evidence["positions"]["汇总对比"]["data_rows"]} == {"best", "last"}
    finally:
        wb.close()


def test_idempotent_update_clears_pending_numbers_and_preserves_later_sections(tmp_path):
    book, report_path, exp = setup_files(tmp_path)
    first = W.update_workbook(book, report_path, exp)
    wb = load_workbook(book)
    later = {}
    for sheet in W.SHEETS:
        row = first["positions"][sheet]["last_row"] + 3
        wb[sheet].cell(row, 1, "后续实验")
        wb[sheet].merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        wb[sheet].cell(row + 1, 1, "=A3+A4")
        # Same rows, outside this tool's rectangle: must also remain untouched.
        wb[sheet].cell(first["positions"][sheet]["header_row"], W.WIDTH + 2, "侧边历史")
        later[sheet] = row
    wb.save(book)
    wb.close()
    report = report_fixture(False)
    row = next(r for r in report["rows"] if (r["id"], r["checkpoint"]) == ("P01", "best"))
    for field in W.NUMERIC_FIELDS:
        row[field] = None
    row["status"] = "pending"
    report_path.write_text(json.dumps(report, ensure_ascii=False))
    second = W.update_workbook(book, report_path, exp, require_complete=False)
    assert second["positions"] == first["positions"]
    wb = load_workbook(book, data_only=False)
    try:
        for sheet in W.SHEETS:
            assert sum(cell.value == W.TITLE for cell in wb[sheet]["A"]) == 1
            index = row_for(second, sheet, "P01", "pressure")
            assert wb[sheet].cell(index, column("r2_cb")).value is None
            assert wb[sheet].cell(index, column("status")).value == "pending"
            assert wb[sheet].cell(later[sheet], 1).value == "后续实验"
            assert wb[sheet].cell(later[sheet] + 1, 1).value == "=A3+A4"
            assert wb[sheet].cell(second["positions"][sheet]["header_row"], W.WIDTH + 2).value == "侧边历史"
    finally:
        wb.close()


def test_supplied_formula_looking_text_is_written_as_a_literal(tmp_path):
    book, report_path, exp = setup_files(tmp_path)
    report = report_fixture()
    report["rows"][0]["gate"] = '=SUM(1,2)'
    report_path.write_text(json.dumps(report))
    evidence = W.update_workbook(book, report_path, exp)
    wb = load_workbook(book, data_only=False)
    try:
        row = row_for(evidence, "实验矩阵总览", "R5P", "pressure")
        cell = wb["实验矩阵总览"].cell(row, column("gate"))
        assert cell.value == '=SUM(1,2)' and cell.data_type == "s"
    finally:
        wb.close()


@pytest.mark.parametrize("mutation", ["duplicate", "missing", "nan", "wrong_task"])
def test_invalid_report_is_rejected_before_workbook_mutation(tmp_path, mutation):
    book, report_path, exp = setup_files(tmp_path)
    report = report_fixture()
    if mutation == "duplicate":
        report["rows"].append(copy.deepcopy(report["rows"][0]))
    elif mutation == "missing":
        report["rows"].pop()
    elif mutation == "nan":
        report["rows"][0]["mae"] = float("nan")
    else:
        report["rows"][0]["task"] = "wss"
    report_path.write_text(json.dumps(report))
    original = book.read_bytes()
    with pytest.raises(ValueError):
        W.update_workbook(book, report_path, exp)
    assert book.read_bytes() == original
    assert not exp.exists()


def test_pending_requires_explicit_registration_option(tmp_path):
    book, report_path, exp = setup_files(tmp_path)
    report_path.write_text(json.dumps(report_fixture(False)))
    original = book.read_bytes()
    with pytest.raises(RuntimeError, match="complete report"):
        W.update_workbook(book, report_path, exp)
    assert book.read_bytes() == original


def test_independent_readback_catches_writer_error_before_replacement(tmp_path, monkeypatch):
    book, report_path, exp = setup_files(tmp_path)
    original = book.read_bytes()
    writer = W.row_values
    def corrupt(row):
        values = writer(row)
        values[column("r2_cb") - 1] += .1
        return values
    monkeypatch.setattr(W, "row_values", corrupt)
    with pytest.raises(AssertionError, match="transcription mismatch"):
        W.update_workbook(book, report_path, exp)
    assert book.read_bytes() == original
    assert not (exp / "xlsx_acceptance.json").exists()
    assert not list(tmp_path.glob(".volume_attention_*.xlsx"))


def test_concurrent_report_change_is_not_published(tmp_path, monkeypatch):
    book, report_path, exp = setup_files(tmp_path)
    original = book.read_bytes()
    verifier = W.verify_readback
    def change_after_check(*args):
        result = verifier(*args)
        report_path.write_text(report_path.read_text() + "\n")
        return result
    monkeypatch.setattr(W, "verify_readback", change_after_check)
    with pytest.raises(RuntimeError, match="changed concurrently"):
        W.update_workbook(book, report_path, exp)
    assert book.read_bytes() == original


@pytest.mark.parametrize("kind", ["formula", "merge"])
def test_historical_corruption_prevents_atomic_replacement(tmp_path, monkeypatch, kind):
    book, report_path, exp = setup_files(tmp_path)
    original = book.read_bytes()
    writer = W.update_in_memory
    def corrupt_history(workbook, report):
        result = writer(workbook, report)
        if kind == "formula":
            workbook["RCR Oracle"]["D2"] = "=1"
        else:
            workbook["RCR Oracle"].unmerge_cells("A1:C1")
        return result
    monkeypatch.setattr(W, "update_in_memory", corrupt_history)
    with pytest.raises(AssertionError, match="historical"):
        W.update_workbook(book, report_path, exp)
    assert book.read_bytes() == original


def test_concurrent_workbook_change_is_preserved(tmp_path, monkeypatch):
    book, report_path, exp = setup_files(tmp_path)
    verifier = W.verify_readback
    external = {}
    def change_after_check(*args):
        result = verifier(*args)
        wb = load_workbook(book)
        wb["RCR Oracle"]["A4"] = 777
        wb.save(book)
        wb.close()
        external["bytes"] = book.read_bytes()
        return result
    monkeypatch.setattr(W, "verify_readback", change_after_check)
    with pytest.raises(RuntimeError, match="changed concurrently"):
        W.update_workbook(book, report_path, exp)
    assert book.read_bytes() == external["bytes"]


def test_changed_ownership_marker_is_not_overwritten(tmp_path):
    book, report_path, exp = setup_files(tmp_path)
    evidence = W.update_workbook(book, report_path, exp)
    wb = load_workbook(book)
    index = evidence["positions"]["实验矩阵总览"]["data_rows"][0]
    wb["实验矩阵总览"].cell(index, 1).comment = None
    wb.save(book)
    wb.close()
    original = book.read_bytes()
    with pytest.raises(ValueError, match="unowned or changed"):
        W.update_workbook(book, report_path, exp)
    assert book.read_bytes() == original
