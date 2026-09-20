#!/usr/bin/env python3
"""Read-only, independent transcription audit of the MS1–MS6 delivery.

No reporter/writer functions are imported. Every expected metric comes from a
run's original metrics.json. Only the explicitly named audit JSON is written.
This checks faithful reporting, not whether the underlying evaluator is correct.
"""
from __future__ import annotations

import argparse
from collections import Counter, OrderedDict
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import math
import pickle
from pathlib import Path
import zipfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
EXP = Path(__file__).resolve().parent
RUNS = ROOT / "training_wss_min/runs"
MATRIX = ROOT / "training_wss_min/configs/v6_multiradius_bt_20260909/matrix.json"
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
BACKUP = EXP / "WSS_PointNet实验矩阵与结果汇总last_backup_before_multiradius_20260909.xlsx"
PREFIX = "V6-多半径BT20260909"
TRACKER = ROOT / "docs/02-推进与变更/WSS_PINN/WSS_V5_训练实验跟踪.md"
EXPECTED_IDS = [f"MS{i}" for i in range(1, 7)]
PARENTS = dict(zip(EXPECTED_IDS, ("MS0", "MS1", "MS1", "MS3", "MS4", "MS4")))
FIELDS = {
    "physical_r2_cb": "field_casebalanced.r2",
    "normalized_r2_cb": "normalized.field_casebalanced.r2",
    "case_mean": "aggregate.r2_casemean",
    "case_p10": "aggregate.r2_casep10",
    "mae": "field.mae",
    "rmse": "field.rmse",
    "high_wss_r2": "regional_field.high_wss.r2",
    "top10_ratio": "calibration.top10_pred_true_ratio",
    "p99_ratio": "calibration.p99_pred_true_ratio",
    "top10_iou": "hotspot.top10_iou_casemean",
    "spearman": "hotspot.spearman_all_casemean",
    "linear_fit_r2_cb": "field_casebalanced.r2_linear_fit",
    "linear_fit_slope_cb": "field_casebalanced.linear_fit_slope",
    "linear_fit_intercept_cb": "field_casebalanced.linear_fit_intercept",
}
OVERVIEW = {
    "G": "field_casebalanced.r2", "H": "field.r2",
    "I": "aggregate.r2_casemean", "J": "aggregate.r2_casemed",
    "K": "aggregate.r2_casep10", "M": "field.rmse", "N": "field.mae",
    "O": "field.nrmse_range", "P": "aggregate.nrmse_casemean",
    "Q": "field.nmae_range", "R": "aggregate.nmae_casemean",
    "S": "regional_field.high_wss.r2", "T": "calibration.top10_pred_true_ratio",
    "U": "hotspot.top10_iou_casemean", "V": "normalized.field_casebalanced.r2",
    "W": "normalized.field.r2", "X": "normalized.aggregate.r2_casemean",
    "Y": "normalized.aggregate.r2_casemed", "Z": "normalized.aggregate.r2_casep10",
    "AA": "normalized.field.mae", "AB": "normalized.field.rmse",
    "AC": "normalized.field.nrmse_range", "AD": "normalized.aggregate.nrmse_casemean",
    "AE": "normalized.field.nmae_range", "AF": "normalized.aggregate.nmae_casemean",
    "AG": "hotspot.spearman_all_casemean", "AH": "hotspot.spearman_high_wss_casemean",
    "AI": "normalized.calibration.top10_pred_true_ratio",
    "AJ": "normalized.calibration.p99_pred_true_ratio",
    "AL": "field_casebalanced.r2_linear_fit", "AM": "field_casebalanced.linear_fit_slope",
    "AN": "field_casebalanced.linear_fit_intercept", "AO": "field.r2_linear_fit",
    "AP": "field.linear_fit_slope", "AQ": "field.linear_fit_intercept",
    "AR": "normalized.field_casebalanced.r2_linear_fit",
    "AS": "normalized.field_casebalanced.linear_fit_slope",
    "AT": "normalized.field_casebalanced.linear_fit_intercept",
    "AU": "normalized.field.r2_linear_fit", "AV": "normalized.field.linear_fit_slope",
    "AW": "normalized.field.linear_fit_intercept",
}
TEACHER = dict(zip(
    "EFGHIJKLMNOPQRS",
    ["field_casebalanced.r2", "aggregate.r2_casemean", "field.rmse", "field.nmae_range",
     "aggregate.nmae_casemean", "regional_field.high_wss.r2", "normalized.field_casebalanced.r2",
     "normalized.field.nmae_range", "hotspot.spearman_all_casemean", "hotspot.top10_iou_casemean",
     "normalized.calibration.p99_pred_true_ratio", "field_casebalanced.r2_linear_fit",
     "field.r2_linear_fit", "field.linear_fit_slope", "field.linear_fit_intercept"]))
COMPARISON = {
    "D": "field_casebalanced.r2", "G": "normalized.field_casebalanced.r2",
    "H": "aggregate.r2_casemean", "I": "field.mae", "J": "regional_field.high_wss.r2",
    "K": "calibration.top10_pred_true_ratio", "L": "calibration.p99_pred_true_ratio",
    "M": "hotspot.top10_iou_casemean",
}


def dig(obj, path):
    for key in path.split("."):
        obj = obj[key]
    return obj


class Audit:
    def __init__(self):
        self.inputs = {}
        self.counts = Counter()
        self.differences = []
        self.unstable_inputs = set()

    def read_bytes(self, path):
        path = Path(path)
        data = path.read_bytes()
        self.inputs[str(path)] = hashlib.sha256(data).hexdigest()
        return data

    def read_json(self, path):
        return json.loads(self.read_bytes(path))

    def check(self, category, label, actual, expected, *, exact=False):
        self.counts[category] += 1
        if not exact and isinstance(expected, (float, int)) and not isinstance(expected, bool):
            ok = (isinstance(actual, (float, int)) and not isinstance(actual, bool)
                  and math.isfinite(actual) and math.isfinite(expected)
                  and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12))
        else:
            ok = actual == expected
        if not ok:
            self.differences.append({"category": category, "location": label,
                                     "actual": actual, "expected": expected})


class TensorStoragePlaceholder:
    pass


def discard_tensor(*args):
    return None


class MetadataUnpickler(pickle.Unpickler):
    """Read scalar metadata without importing torch or rebuilding any tensor.

    Only the exact constructors emitted by the local state_dict writer are
    accepted; every tensor/storage object is replaced with an inert placeholder.
    """
    def find_class(self, module, name):
        if (module, name) == ("collections", "OrderedDict"):
            return OrderedDict
        if module == "torch._utils" and name in {"_rebuild_tensor_v2", "_rebuild_tensor"}:
            return discard_tensor
        if module == "torch" and name in {"FloatStorage", "DoubleStorage", "HalfStorage", "BFloat16Storage", "LongStorage", "IntStorage", "ShortStorage", "ByteStorage", "BoolStorage"}:
            return TensorStoragePlaceholder
        raise pickle.UnpicklingError(f"Unsupported checkpoint metadata global: {module}.{name}")

    def persistent_load(self, pid):
        if not isinstance(pid, tuple) or not pid or pid[0] != "storage":
            raise pickle.UnpicklingError("Unsupported persistent id")
        return None


def checkpoint_metadata(audit, path):
    data = audit.read_bytes(path)
    with zipfile.ZipFile(BytesIO(data)) as archive:
        audit.check("checkpoint_archive", str(path) + "/crc", archive.testzip(), None)
        metadata_files = [n for n in archive.namelist() if n.endswith("/data.pkl")]
        if len(metadata_files) != 1:
            raise ValueError(f"Unexpected checkpoint archive: {path}")
        state = MetadataUnpickler(BytesIO(archive.read(metadata_files[0]))).load()
    if not isinstance(state, dict) or not isinstance(state.get("model"), OrderedDict):
        raise ValueError(f"Invalid checkpoint payload: {path}")
    return {"epoch": state.get("epoch"), "metric": state.get("metric"),
            "cfg_name": state.get("cfg_name"), "state_tensor_count": len(state["model"])}


def check_subset(audit, category, label, actual, expected):
    if isinstance(expected, dict):
        audit.check(category, label + "/is_dict", isinstance(actual, dict), True)
        for key, value in expected.items():
            check_subset(audit, category, label + "/" + key,
                         actual.get(key) if isinstance(actual, dict) else None, value)
    else:
        audit.check(category, label, actual, expected, exact=True)


def execution_and_configuration(audit, arms, summaries, require_final):
    """Independently qualify finished stages; partial runs never prove finality."""
    anchor_path = ROOT / "training_wss_min/configs/v6_followup_20260909/M2_a5_independent_k3_s1234.json"
    anchor_config = audit.read_json(anchor_path)
    freeze = audit.read_json(EXP / "grouping_freeze.json")
    audit.check("freeze", "cap", freeze.get("cap"), 32)
    audit.check("freeze", "radii", freeze.get("radii"), [.1, .2, .4], exact=True)
    audit.check("freeze", "labels_used", freeze.get("labels_used"), False)
    audit.check("freeze", "test_used", freeze.get("test_used"), False)
    audit.check("freeze", "n_cases", freeze.get("n_cases"), 138)
    audit.check("freeze", "all_arms", freeze.get("all_arms"), EXPECTED_IDS, exact=True)
    for filename, digest in freeze["final_config_sha256"].items():
        actual = hashlib.sha256(audit.read_bytes(MATRIX.parent / filename)).hexdigest()
        audit.check("freeze_hash", filename, actual, digest)
    qp = EXP / "queue_status.json"
    queue = audit.read_json(qp) if qp.is_file() else {}
    queue_finished = queue.get("status") == "complete" and bool(queue.get("ended_at"))
    if not queue_finished:
        audit.unstable_inputs.add(str(qp))
    if require_final:
        audit.check("queue", "complete", queue_finished, True)
        audit.check("queue", "arm_ids", sorted(queue.get("arms", {})), EXPECTED_IDS, exact=True)
    start, end = queue.get("source_sha256", {}), queue.get("source_sha256_end", {})
    if queue_finished or require_final:
        audit.check("queue", "start_hash_available", bool(start), True)
        audit.check("queue", "end_hash_available", bool(end), True)
        audit.check("queue", "start_end_hash_equal", start, end, exact=True)
        audit.check("queue", "source_changed", queue.get("source_changed"), False)
    # Archive validity and current worktree drift are distinct. No extra training
    # may be inferred from a later authorized edit of the checked-in source.
    current_drift = []
    for relative, digest in start.items():
        p = ROOT / relative
        actual = hashlib.sha256(audit.read_bytes(p)).hexdigest() if p.is_file() else None
        if actual != digest:
            current_drift.append(relative)
        if not queue_finished:
            audit.check("source_live", relative, actual, digest)
    required_model_common = {"sa_nsample": [64, 16, 32], "sa_center_counts": [125, 125, 32],
        "sa_radius": [.05, .1, .2], "local_branch_nsample": [16, 16, 16],
        "local_branch_radii": [.015, .03, .06], "local_branch": True,
        "fp_knn": 3, "query_decoder": "interpolate", "out_dim": 1,
        "bottleneck_transformer": False, "bottleneck_layers": 1,
        "bottleneck_heads": 8, "bottleneck_ffn_ratio": 2,
        "bottleneck_dropout": 0., "bottleneck_pos_enc": "input_features",
        "bottleneck_geo_bias": True, "bottleneck_residual_scale_init": .001}
    changed_model_keys = {"multiradius_values", "multiradius_bottleneck", "multiradius_ffn", "sa_nsample"} | {k for k in required_model_common if k.startswith("bottleneck_")}
    records = {}
    for arm in arms:
        aid, run = arm["id"], RUNS / arm["run_name"]
        cfg = audit.read_json(MATRIX.parent / arm["config"])
        audit.check("configuration", aid + "/name", cfg.get("name"), arm["run_name"])
        audit.check("configuration", aid + "/parent", arm.get("baseline"), PARENTS[aid])
        for part in ("data", "train", "eval"):
            audit.check("configuration_unchanged", aid + "/" + part, cfg.get(part), anchor_config[part], exact=True)
        check_subset(audit, "configuration", aid + "/model", cfg["model"], required_model_common)
        audit.check("configuration", aid + "/multiradius_values", cfg["model"].get("multiradius_values"), [.2] if aid in {"MS1", "MS2"} else [.2, .2, .2] if aid == "MS5" else [.1, .2, .4], exact=True)
        audit.check("configuration", aid + "/multiradius_bottleneck", cfg["model"].get("multiradius_bottleneck"), aid in {"MS2", "MS4", "MS5"})
        audit.check("configuration", aid + "/multiradius_ffn", cfg["model"].get("multiradius_ffn"), aid == "MS6")
        for key in set(anchor_config["model"]) | set(cfg["model"]):
            if key not in changed_model_keys:
                audit.check("configuration_unchanged", aid + "/model/" + key, cfg["model"].get(key), anchor_config["model"].get(key), exact=True)
        q = queue.get("arms", {}).get(aid, {})
        stages = q.get("stages", {})
        train_done = stages.get("train", {}).get("returncode") == 0
        claimed = any(s["arms"][aid]["evidence"].get("valid_scientific_result") for s in summaries.values())
        if require_final or claimed:
            audit.check("queue_stage", aid + "/train_code", stages.get("train", {}).get("returncode"), 0)
            audit.check("queue_stage", aid + "/run_name", q.get("run_name"), arm["run_name"])
        if require_final:
            audit.check("queue_stage", aid + "/status", q.get("status"), "complete")
        for ck in ("best", "last"):
            if require_final or summaries[ck]["arms"][aid]["evidence"].get("valid_scientific_result"):
                audit.check("queue_stage", f"{aid}/eval_{ck}_code", stages.get(f"eval_{ck}", {}).get("returncode"), 0)
        runtime_cfg = run / "config.json"
        if runtime_cfg.is_file():
            check_subset(audit, "runtime_configuration", aid, audit.read_json(runtime_cfg), cfg)
        elif require_final or claimed:
            audit.check("runtime_configuration", aid + "/exists", False, True)
        rec = {"train_done": train_done, "queue_status": q.get("status"), "checkpoints": {}}
        if train_done or require_final or claimed:
            hp = run / "history.jsonl"
            audit.check("history", aid + "/exists", hp.is_file(), True)
            history = [json.loads(line) for line in audit.read_bytes(hp).splitlines() if line.strip()] if hp.is_file() else []
            audit.check("history", aid + "/epochs", [h.get("epoch") for h in history], list(range(400)), exact=True)
            for h in history:
                for key in ("train_loss", "selection_score", "train_mse_norm", "train_mae_norm", "train_rmse_norm"):
                    value = h.get(key)
                    audit.check("history_finite", f"{aid}/{h.get('epoch')}/{key}", isinstance(value, (float, int)) and math.isfinite(value), True)
                if isinstance(h.get("selection_score"), (int, float)) and isinstance(h.get("train_loss"), (int, float)):
                    audit.check("history_selection", f"{aid}/{h.get('epoch')}", h["selection_score"], -h["train_loss"])
            for ck in ("best", "last"):
                cp = run / f"ckpt_{ck}.pt"
                audit.check("checkpoint", f"{aid}/{ck}/exists", cp.is_file(), True)
                if not cp.is_file():
                    continue
                meta = checkpoint_metadata(audit, cp)
                rec["checkpoints"][ck] = meta
                audit.check("checkpoint", f"{aid}/{ck}/cfg_name", meta["cfg_name"], arm["run_name"])
                if history:
                    winning = max(history, key=lambda h: h["selection_score"])
                    audit.check("checkpoint", f"{aid}/{ck}/epoch", meta["epoch"], winning["epoch"] if ck == "best" else 399)
                    if ck == "best":
                        audit.check("checkpoint", f"{aid}/{ck}/selection_metric", meta["metric"], winning["selection_score"])
            initialization = run / "initialization.json"
            if initialization.is_file():
                audit.check("initialization", aid, audit.read_json(initialization).get("mode"), "random_initialization")
            rec["history_rows"] = len(history)
        records[aid] = rec
    return {"queue_finished": queue_finished, "start_end_hash_equal": bool(start and end and start == end),
            "current_source_drift": current_drift, "arms": records}


def tracker_section(text, heading):
    """Slice the tracker block owned by `heading`, independently of the reporter.

    Section boundaries are the heading itself and the next heading of the same or
    shallower level; the tracker carries no anchor comments (a rich-text editor
    would refuse them).
    """
    def level(line):
        depth = len(line) - len(line.lstrip("#"))
        return depth if 1 <= depth <= 6 and line[depth:depth + 1] == " " else 0

    lines = text.split("\n")
    hits = [i for i, line in enumerate(lines) if line == heading]
    if len(hits) != 1:
        return None
    start = hits[0]
    end = next((i for i in range(start + 1, len(lines))
                if 0 < level(lines[i]) <= level(heading)), len(lines))
    return "\n".join(lines[start:end])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--final", action="store_true", help="Require 6/6 finalized best and last; fail on incomplete work")
    ap.add_argument("--output", type=Path, default=None)
    args = ap.parse_args()
    audit = Audit()
    matrix = audit.read_json(MATRIX)
    arms = matrix["arms"]
    ids = [a["id"] for a in arms]
    audit.check("matrix", "exact_authorized_arms", ids, EXPECTED_IDS, exact=True)
    anchor = matrix["reference"]["run_name"]
    canonical_split = audit.read_json(ROOT / "data_wss_v5/views/wss_min_view_v1/split_V5_train138_test34.json")
    expected_patients = sorted(canonical_split["test_cases"])
    audit.check("source_test_set", "canonical_test_count", len(expected_patients), 34)
    summaries = {ck: audit.read_json(EXP / f"matrix_summary_{ck}.json") for ck in ("best", "last")}
    execution = execution_and_configuration(audit, arms, summaries, args.final)
    update = audit.read_json(EXP / "workbook_update.json")
    sources = {}
    missing_sources = []
    for ck in summaries:
        sources[ck] = {}
        for aid, run_name in [("MS0", anchor)] + [(a["id"], a["run_name"]) for a in arms]:
            path = RUNS / run_name / "eval" / f"ckpt_{ck}" / "metrics.json"
            if not path.is_file():
                missing_sources.append(str(path))
                continue
            m = audit.read_json(path)["test"]
            sources[ck][aid] = m
            audit.check("source_test_set", f"{ck}/{aid}/n_cases", m["aggregate"]["n_cases"], 34)
            audit.check("source_test_set", f"{ck}/{aid}/per_case_count", len(m["per_case"]), 34)
            audit.check("source_test_set", f"{ck}/{aid}/canonical_patients", sorted(m["per_case"]), expected_patients, exact=True)
            audit.check("source_test_set", f"{ck}/{aid}/normalized_patients", sorted(m["normalized"]["per_case"]), expected_patients, exact=True)
            audit.check("source_test_set", f"{ck}/{aid}/normalized_n_cases", m["normalized"]["aggregate"]["n_cases"], 34)
            audit.check("source_protocol", f"{ck}/{aid}/surface_metric_mode", m.get("surface_metric_mode"), "legacy_vertex")
            if aid != "MS0":
                cfg = json.loads((MATRIX.parent / next(a["config"] for a in arms if a["id"] == aid)).read_text())
                audit.check("source_protocol", f"{ck}/{aid}/evaluation_split_path", m.get("evaluation_split_path"), cfg["data"]["split_path"])
        base = sources[ck]["MS0"]
        for aid, m in sources[ck].items():
            audit.check("source_test_set", f"{ck}/{aid}/patients",
                        sorted(m["per_case"]), sorted(base["per_case"]), exact=True)

    checkpoint_results = {}
    if args.final:
        audit.check("final_completeness", "missing_metrics", missing_sources, [], exact=True)
    for ck, summary in summaries.items():
        base = sources[ck]["MS0"]
        audit.check("summary_metadata", f"{ck}/checkpoint", summary["checkpoint"], ck)
        audit.check("summary_metadata", f"{ck}/reference", summary["reference"]["run_name"], anchor)
        audit.check("summary_metadata", f"{ck}/arm_ids", sorted(summary["arms"]), sorted(ids), exact=True)
        for key, path in FIELDS.items():
            audit.check("summary_reference", f"{ck}/MS0/{key}", summary["reference"].get(key), dig(base, path))
        valid = [aid for aid in ids if summary["arms"][aid]["evidence"]["valid_scientific_result"]]
        for arm in arms:
            aid = arm["id"]
            entry = summary["arms"][aid]
            m = sources[ck].get(aid)
            label = f"{ck}/{aid}"
            parent = arm.get("baseline", arm.get("parent", "MS0"))
            audit.check("summary_metadata", label + "/run_name", entry["run_name"], arm["run_name"])
            audit.check("summary_metadata", label + "/parent", entry["parent"], parent)
            if "raw_metrics" in entry:
                audit.check("summary_metadata", label + "/raw_source_exists", m is not None, True)
                audit.check("summary_metadata", label + "/raw_metrics_path", entry.get("raw_metrics_path"),
                            str(RUNS / arm["run_name"] / "eval" / f"ckpt_{ck}" / "metrics.json"))
                if m is not None:
                    for key, path in FIELDS.items():
                        audit.check("summary_raw", label + "/" + key, entry["raw_metrics"].get(key), dig(m, path))
            if aid not in valid:
                for key in list(FIELDS) + ["delta_physical", "delta_normalized", "delta_parent_physical",
                                         "delta_parent_normalized", "per_case_delta", "cases_improved"]:
                    audit.check("summary_pending_blank", label + "/" + key, entry.get(key), None)
                continue
            audit.check("summary_metadata", label + "/valid_source_exists", m is not None, True)
            if m is None:
                continue
            for key, path in FIELDS.items():
                audit.check("summary_valid", label + "/" + key, entry.get(key), dig(m, path))
            pm = base if parent in {"MS0", "M2", "MS0=M2"} else sources[ck].get(parent) if parent in valid else None
            for suffix, path in (("physical", FIELDS["physical_r2_cb"]), ("normalized", FIELDS["normalized_r2_cb"])):
                audit.check("summary_delta", label + "/delta_" + suffix,
                            entry.get("delta_" + suffix), dig(m, path) - dig(base, path))
                audit.check("summary_delta", label + "/delta_parent_" + suffix,
                            entry.get("delta_parent_" + suffix), dig(m, path) - dig(pm, path) if pm else None)
            for key, path in FIELDS.items():
                audit.check("summary_all_metric_delta", label + "/ms0/" + key, entry.get("deltas_vs_ms0", {}).get(key), dig(m, path) - dig(base, path))
                audit.check("summary_all_metric_delta", label + "/parent/" + key, (entry.get("deltas_vs_parent") or {}).get(key), dig(m, path) - dig(pm, path) if pm else None)
            delta = {case: m["per_case"][case]["overall"]["r2"] - base["per_case"][case]["overall"]["r2"]
                     for case in base["per_case"]}
            audit.check("summary_per_case", label + "/delta_keys", sorted(entry.get("per_case_delta", {})), sorted(delta), exact=True)
            for case, expected in delta.items():
                audit.check("summary_per_case", label + "/" + case, entry.get("per_case_delta", {}).get(case), expected)
            audit.check("summary_per_case", label + "/cases_improved", entry.get("cases_improved"), sum(v > 0 for v in delta.values()))
        audit.check("summary_metadata", ck + "/completed_count", summary["completed_count"], len(valid))
        audit.check("summary_metadata", ck + "/expected_count", summary["expected_count"], len(arms))
        audit.check("summary_metadata", ck + "/raw_count", summary["raw_metrics_count"], sum("raw_metrics" in e for e in summary["arms"].values()))
        if args.final:
            audit.check("final_completeness", ck + "/qualified", len(valid), 6)
            audit.check("final_completeness", ck + "/matrix_finalized", summary.get("matrix_finalized"), True)
        if summary.get("matrix_finalized"):
            audit.check("summary_finalization", ck + "/qualified", len(valid), 6)
            audit.check("summary_finalization", ck + "/queue_finished", execution["queue_finished"], True)
            audit.check("summary_finalization", ck + "/source_hashes", execution["start_end_hash_equal"], True)
        if {"MS1", "MS2", "MS3", "MS4"} <= set(valid):
            for key, path in FIELDS.items():
                expected = (dig(sources[ck]["MS4"], path) - dig(sources[ck]["MS3"], path)) - (dig(sources[ck]["MS2"], path) - dig(sources[ck]["MS1"], path))
                audit.check("summary_interaction", ck + "/" + key, summary.get("factorial_interaction", {}).get(key), expected)
        for child, parent in (("MS4", "MS2"), ("MS4", "MS3"), ("MS4", "MS5"), ("MS4", "MS6")):
            if child in valid and parent in valid:
                for key, path in FIELDS.items():
                    expected = dig(sources[ck][child], path) - dig(sources[ck][parent], path)
                    audit.check("summary_contrast", f"{ck}/{child}-{parent}/{key}", summary.get("planned_contrasts", {}).get(child + "-" + parent, {}).get(key), expected)
        checkpoint_results[ck] = {"valid_count": len(valid), "valid_ids": valid,
                                  "pending_ids": [aid for aid in ids if aid not in valid],
                                  "matrix_finalized_reported": summary["matrix_finalized"]}

    # Verify each displayed rounded table cell from the original numeric source,
    # rather than importing the formatting/report function being audited.
    markdowns = {ck: audit.read_bytes(EXP / f"matrix_tables_{ck}.md").decode() for ck in ("best", "last")}
    for ck, markdown in markdowns.items():
        table_rows = [line for line in markdown.splitlines() if line.startswith("| MS")]
        audit.check("markdown", ck + "/row_count", len(table_rows), 7)
        for aid in ["MS0"] + ids:
            display_id = "MS0=M2" if aid == "MS0" else aid
            rows = [r for r in table_rows if r.split("|")[1].strip() == display_id]
            audit.check("markdown", f"{ck}/{aid}/identity_count", len(rows), 1)
            if len(rows) != 1:
                continue
            cells_text = [v.strip() for v in rows[0].split("|")[1:-1]]
            audit.check("markdown", f"{ck}/{aid}/columns", len(cells_text), 14)
            if len(cells_text) != 14:
                continue
            valid = aid == "MS0" or aid in checkpoint_results[ck]["valid_ids"]
            metric = sources[ck].get(aid) if valid else None
            for col, path, decimals in ((3, FIELDS["physical_r2_cb"], 4), (6, FIELDS["normalized_r2_cb"], 4),
                (7, FIELDS["case_mean"], 4), (8, FIELDS["mae"], 3), (9, FIELDS["high_wss_r2"], 4),
                (10, FIELDS["top10_ratio"], 3), (11, FIELDS["p99_ratio"], 3), (12, FIELDS["top10_iou"], 3), (13, FIELDS["spearman"], 3)):
                audit.check("markdown", f"{ck}/{aid}/column{col+1}", cells_text[col], f"{dig(metric, path):.{decimals}f}" if metric else "—")
            parent = None if aid == "MS0" else PARENTS[aid]
            pm = sources[ck].get(parent) if parent == "MS0" or parent in checkpoint_results[ck]["valid_ids"] else None
            for col, reference in ((4, sources[ck]["MS0"]), (5, pm)):
                expected = f"{dig(metric, FIELDS['physical_r2_cb']) - dig(reference, FIELDS['physical_r2_cb']):.4f}" if metric and reference and aid != "MS0" else "—"
                audit.check("markdown", f"{ck}/{aid}/delta_column{col+1}", cells_text[col], expected)
    tracker = audit.read_bytes(TRACKER).decode()
    heading = "### 14.2 执行与结果（best；由原始指标及队列证据自动回填）"
    audit.check("tracker", "anchor_heading_count", tracker.count(heading + "\n"), 1)
    block = tracker_section(tracker, heading)
    if block is not None:
        audit.check("tracker", "section_14.2", block.startswith("### 14.2 "), True)
        audit.check("tracker", "best_body_equal", markdowns["best"].split("\n", 1)[1].strip() in block, True)

    wb = load_workbook(BytesIO(audit.read_bytes(BOOK)), data_only=False)
    old = load_workbook(BytesIO(audit.read_bytes(BACKUP)), data_only=False)
    audit.check("workbook_metadata", "sheetnames", wb.sheetnames, old.sheetnames, exact=True)
    audit.check("workbook_metadata", "sheet_count", len(wb.sheetnames), 5)
    markers = {}
    for name in ("实验矩阵总览", "教师汇报视图", "汇总对比"):
        rows = [c.row for cells in wb[name].iter_rows(min_col=1, max_col=1) for c in cells
                if isinstance(c.value, str) and PREFIX + "｜" in c.value]
        audit.check("workbook_metadata", name + "/marker_count", len(rows), 1)
        if len(rows) != 1:
            raise RuntimeError(f"Cannot identify own section: {name}: {rows}")
        markers[name] = rows[0]
        audit.check("workbook_metadata", name + "/explicit_best", "checkpoint=best" in wb[name].cell(rows[0], 1).value, True)
    preserved = {}
    for name in old.sheetnames:
        before, after = old[name], wb[name]
        end = markers[name] - 1 if name in markers else max(before.max_row, after.max_row)
        audit.check("legacy_boundary", name + "/old_rows_within_cutoff", before.max_row <= end, True)
        count = formulas = 0
        for row in range(1, max(end, before.max_row) + 1):
            for col in range(1, max(before.max_column, after.max_column) + 1):
                expected, actual = before.cell(row, col).value, after.cell(row, col).value
                if expected is None and actual is None:
                    continue
                count += expected is not None
                formulas += isinstance(expected, str) and expected.startswith("=")
                audit.check("legacy_cell", f"{name}!{before.cell(row, col).coordinate}", actual, expected, exact=True)
        preserved[name] = {"last_original_row": end, "backup_max_row": before.max_row,
                           "original_nonempty_cells": count, "original_formula_cells": formulas}
    best = summaries["best"]
    valid = checkpoint_results["best"]["valid_ids"]
    audit.check("workbook_metadata", "update_checkpoint", update["checkpoint"], "best")
    audit.check("workbook_metadata", "update_completed", update["completed_best"], len(valid))
    audit.check("workbook_metadata", "update_finalized", update["matrix_finalized"], best["matrix_finalized"])
    audit.check("workbook_metadata", "update_reference", update["reference"], anchor)

    def cells(ws, row, mapping, metric):
        for col, path in mapping.items():
            audit.check("workbook_metric" if metric else "workbook_pending_blank", f"{ws.title}!{col}{row}",
                        ws[f"{col}{row}"].value, dig(metric, path) if metric else None)

    for index, arm in enumerate(arms, 1):
        aid = arm["id"]
        m = sources["best"].get(aid) if aid in valid else None
        ov, te, su = (wb[name] for name in ("实验矩阵总览", "教师汇报视图", "汇总对比"))
        ro, rt, rs = markers[ov.title] + index, markers[te.title] + index, markers[su.title] + index + 2
        audit.check("workbook_metadata", aid + "/overview_row", update["overview_rows"][aid], ro)
        audit.check("workbook_metadata", aid + "/teacher_row", update["teacher_rows"][aid], rt)
        audit.check("workbook_metadata", aid + "/summary_row", update["summary_rows"][aid], rs)
        audit.check("workbook_metadata", f"{ov.title}!A{ro}", ov[f"A{ro}"].value, f"{PREFIX}·{aid} {Path(arm['config']).stem}")
        audit.check("workbook_metadata", f"{ov.title}!AK{ro}", ov[f"AK{ro}"].value, arm["run_name"])
        audit.check("workbook_metadata", f"{te.title}!B{rt}/id", te[f"B{rt}"].value.startswith(aid + " "), True)
        audit.check("workbook_metadata", f"{su.title}!A{rs}", su[f"A{rs}"].value, aid)
        audit.check("workbook_metadata", f"{su.title}!O{rs}", su[f"O{rs}"].value, arm["run_name"])
        for ws, row, col in ((ov, ro, "E"), (te, rt, "D")):
            audit.check("workbook_metadata", f"{ws.title}!{col}{row}/best", "checkpoint=best" in ws[f"{col}{row}"].value, True)
        cells(ov, ro, OVERVIEW, m)
        cells(te, rt, TEACHER, m)
        cells(su, rs, COMPARISON, m)
        negative = f"{m['aggregate']['r2_negative_cases']}/{m['aggregate']['n_cases']}" if m else None
        audit.check("workbook_metric" if m else "workbook_pending_blank", f"{ov.title}!L{ro}", ov[f"L{ro}"].value, negative)
        parent = arm.get("baseline", arm.get("parent", "MS0"))
        pm = sources["best"]["MS0"] if parent in {"MS0", "M2", "MS0=M2"} else sources["best"].get(parent) if parent in valid else None
        for col, reference in (("E", sources["best"]["MS0"]), ("F", pm)):
            expected = m["field_casebalanced"]["r2"] - reference["field_casebalanced"]["r2"] if m and reference else None
            audit.check("workbook_delta" if m else "workbook_pending_blank", f"{su.title}!{col}{rs}", su[f"{col}{rs}"].value, expected)
    su = wb["汇总对比"]
    reference_row = markers[su.title] + 2
    audit.check("workbook_metadata", f"{su.title}!A{reference_row}", su[f"A{reference_row}"].value, "MS0=M2")
    cells(su, reference_row, COMPARISON, sources["best"]["MS0"])
    wb.close()
    old.close()
    drift = [path for path, expected in audit.inputs.items()
             if not Path(path).is_file() or hashlib.sha256(Path(path).read_bytes()).hexdigest() != expected]
    strict_drift = [p for p in drift if args.final or p not in audit.unstable_inputs]
    audit.check("snapshot", "input_hash_drift", strict_drift, [], exact=True)
    passed = not audit.differences
    complete = execution["queue_finished"] and execution["start_end_hash_equal"] and all(r["valid_count"] == len(arms) and r["matrix_finalized_reported"] for r in checkpoint_results.values())
    output = (args.output or EXP / ("delivery_verification_final.json" if args.final else "delivery_verification_partial.json")).resolve()
    if str(output) in audit.inputs or output.suffix != ".json":
        raise ValueError("Output must be a separate JSON audit file")
    result = {
        "schema": "v6_multiradius_delivery_verification_v1",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "passed": passed, "scope": "final" if complete else "partial",
        "all_6_best_and_last_reported_finalized": complete,
        "limitations": ["Checks faithful transcription from original metrics; does not recompute evaluator outputs.",
                        "Checkpoint ZIP CRC and scalar metadata are checked with inert tensor placeholders; no torch or training package is imported.",
                        "Only archived start/end hashes establish training source integrity; later current-source drift is reported separately.",
                        "Pending metrics are checked as absent/blank and never converted to zero."],
        "execution_and_configuration": execution,
        "checkpoints": checkpoint_results, "workbook_checkpoint": "best",
        "metric_semantics": {"overview_AA_AB": "log_z pooled MAE/RMSE", "overview_M_N": "Pa pooled RMSE/MAE",
                             "teacher_O": "log_z pooled p99_pred_true_ratio", "comparison_L": "Pa pooled p99_pred_true_ratio"},
        "original_workbook_preservation": preserved,
        "checks_by_category": dict(audit.counts), "total_checks": sum(audit.counts.values()),
        "difference_count": len(audit.differences), "differences": audit.differences,
        "input_files": [{"path": path, "sha256": digest} for path, digest in sorted(audit.inputs.items())],
        "input_hashes_unchanged_during_audit": not drift, "missing_original_metric_files": missing_sources,
    }
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"output": str(output), "passed": passed, "scope": result["scope"],
                      "best": checkpoint_results["best"]["valid_count"], "last": checkpoint_results["last"]["valid_count"],
                      "checks": result["total_checks"], "differences": len(audit.differences)}, ensure_ascii=False))
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
