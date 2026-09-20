"""Result qualification and workbook ownership checks without touching real outputs."""
from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from training_wss_min.tools import report_v6_followup as R, update_v6_followup_xlsx as W


def fake_metric(score=.6):
    m = {"aggregate": {"n_cases": 34}, "per_case": {str(i): {"overall": {"r2": score}} for i in range(34)}}
    for path in R.PATHS.values():
        block = m
        for key in path[:-1]:
            block = block.setdefault(key, {})
        block[path[-1]] = score
    return m


class QualificationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.patch_runs = patch.object(R, "RUNS", self.root / "runs")
        self.patch_runs.start()
        self.arm = {"id": "D1", "run_name": "formal/D1", "config": "D1.json", "hypothesis": "test", "baseline": "B0"}
        self.run = R.RUNS / self.arm["run_name"]
        self.run.mkdir(parents=True)

    def tearDown(self):
        self.patch_runs.stop()
        self.tmp.cleanup()

    def artifacts(self, checkpoint="best", epochs=400):
        (self.run / "history.jsonl").write_text("".join(json.dumps({"epoch": i}) + "\n" for i in range(epochs)))
        (self.run / f"ckpt_{checkpoint}.pt").write_bytes(b"checkpoint existence fixture")
        path = self.run / "eval" / f"ckpt_{checkpoint}"
        path.mkdir(parents=True)
        (path / "metrics.json").write_text(json.dumps({"test": fake_metric()}))

    def queue(self, status="complete", failed_stage=None):
        q = {"run_name": self.arm["run_name"], "status": status,
             "stages": {"train": {"returncode": 0}, "eval_best": {"returncode": 0}}}
        if failed_stage:
            q["failed_stage"] = failed_stage
            q["stages"][failed_stage] = {"returncode": 1}
        return {"arms": {"D1": q}}

    def test_submission_pending_queue_pending_failure_and_interruption_are_distinct(self):
        _, submitted = R.run_evidence(self.arm, "best", submission={"formal_job_id": "1"})
        self.assertIn("作业已提交", submitted["status"])
        for queue_status, text in [("pending", "队列待运行"), ("failed", "失败"), ("interrupted", "已中断")]:
            _, evidence = R.run_evidence(self.arm, "best", self.queue(queue_status))
            self.assertIn(text, evidence["status"])
            self.assertFalse(evidence["valid_scientific_result"])

    def test_failed_training_raw_metric_is_not_qualified(self):
        self.artifacts(epochs=18)
        metric, evidence = R.run_evidence(self.arm, "best", self.queue("failed", "train"))
        self.assertIsNotNone(metric)
        self.assertFalse(evidence["valid_scientific_result"])
        self.assertIn("400 complete epochs not verified", evidence["issues"])

    def test_best_remains_usable_when_only_last_evaluation_failed(self):
        self.artifacts()
        _, evidence = R.run_evidence(self.arm, "best", self.queue("failed", "eval_last"))
        self.assertTrue(evidence["valid_scientific_result"])
        self.assertIn("eval_last", evidence["status"])
        self.assertIn("best评估可用", evidence["status"])

    def test_no_queue_provenance_or_unfinished_eval_never_counts(self):
        self.artifacts()
        _, evidence = R.run_evidence(self.arm, "best")
        self.assertFalse(evidence["valid_scientific_result"])
        state = self.queue("eval_best")
        state["arms"]["D1"]["stages"]["eval_best"] = {}
        _, evidence = R.run_evidence(self.arm, "best", state)
        self.assertFalse(evidence["valid_scientific_result"])

    def test_live_source_mutation_detected_before_runner_exits(self):
        source = self.root / "source.py"
        source.write_text("old")
        expected = hashlib.sha256(source.read_bytes()).hexdigest()
        source.write_text("changed")
        state = {"source_changed": False, "source_sha256": {str(source): expected}}
        self.assertTrue(R.source_integrity(state)["changed"])

    def test_archived_matching_hashes_preserve_results_after_later_source_edit(self):
        source = self.root / "source.py"
        source.write_text("old")
        hashes = {str(source): hashlib.sha256(source.read_bytes()).hexdigest()}
        state = {"status": "complete", "ended_at": "2026-09-09T12:00:00+0800",
                 "source_sha256": hashes, "source_sha256_end": dict(hashes), "source_changed": False}
        source.write_text("authorized later work")
        evidence = R.source_integrity(state)
        self.assertFalse(evidence["changed"])
        self.assertTrue(evidence["current_drift"])
        self.assertTrue(evidence["archived"])
        self.assertEqual(evidence["basis"], "archived_start_end")

    def test_archived_mismatch_is_detected_without_runner_boolean(self):
        state = {"status": "complete", "ended_at": "done", "source_sha256": {"source.py": "old"},
                 "source_sha256_end": {"source.py": "changed"}, "source_changed": False}
        evidence = R.source_integrity(state)
        self.assertTrue(evidence["changed"])
        self.assertFalse(evidence["archived_hashes_match"])

    def test_matrix_completion_requires_real_end_hashes_and_all_stage_success(self):
        source = self.root / "source.py"
        source.write_text("frozen")
        hashes = {str(source): hashlib.sha256(source.read_bytes()).hexdigest()}
        state = self.queue()
        state.update(status="complete", ended_at="done", source_sha256=hashes)
        state["arms"]["D1"]["stages"]["eval_last"] = {"returncode": 0}
        check = lambda: R.matrix_finalization(state, ["D1"], R.source_integrity(state))
        self.assertFalse(check()["finalized"])
        state["source_sha256_end"] = dict(hashes)
        self.assertTrue(check()["finalized"])
        state["arms"]["D1"]["stages"]["eval_last"] = {"returncode": 1}
        self.assertFalse(check()["finalized"])

    def test_matrix_does_not_finalize_in_runner_exit_window_or_wrong_arm_set(self):
        state = self.queue()
        state.update(status="running", source_sha256={"source.py": "same"},
                     source_sha256_end={"source.py": "same"})
        state["arms"]["D1"]["stages"]["eval_last"] = {"returncode": 0}
        self.assertFalse(R.matrix_finalization(state, ["D1"], R.source_integrity(state))["finalized"])
        state.update(status="complete", ended_at="done")
        self.assertFalse(R.matrix_finalization(state, ["D1", "D2"], R.source_integrity(state))["finalized"])

    def test_best_and_last_use_corresponding_anchor_and_parent(self):
        exp = self.root / "experiment"
        exp.mkdir()
        matrix = self.root / "matrix.json"
        child = {**self.arm, "id": "M3", "run_name": "formal/M3", "baseline": "D1"}
        matrix.write_text(json.dumps({"arms": [self.arm, child]}))
        valid = {"valid_scientific_result": True, "status": "已评估", "issues": []}
        for checkpoint, anchor, parent, child_score in [("best", .5, .6, .8), ("last", .4, .7, .75)]:
            with patch.object(R, "EXP", exp), patch.object(R, "MATRIX", matrix), \
                 patch.object(R, "load_metrics", return_value=fake_metric(anchor)) as load, \
                 patch.object(R, "run_evidence", side_effect=[(fake_metric(parent), valid), (fake_metric(child_score), valid)]) as evidence:
                markdown, summary = R.report(checkpoint)
            load.assert_called_once_with(R.ANCHOR, checkpoint)
            self.assertTrue(all(call.args[1] == checkpoint for call in evidence.call_args_list))
            self.assertAlmostEqual(summary["arms"]["M3"]["delta_parent_physical"], child_score - parent)
            self.assertAlmostEqual(summary["arms"]["M3"]["delta_physical"], child_score - anchor)
            self.assertIn("Spearman", markdown)
            self.assertIn("linear_fit_r2_cb", summary["arms"]["M3"])

    def test_summary_preserves_raw_but_excludes_changed_source_and_parent_delta(self):
        exp = self.root / "experiment"
        exp.mkdir()
        matrix = self.root / "matrix.json"
        child = {**self.arm, "id": "M3", "run_name": "formal/M3", "baseline": "D1"}
        matrix.write_text(json.dumps({"arms": [self.arm, child]}))
        invalid = {"valid_scientific_result": False, "status": "源文件变化", "issues": ["source_changed"]}
        valid = {"valid_scientific_result": True, "status": "已评估", "issues": []}
        with patch.object(R, "EXP", exp), patch.object(R, "MATRIX", matrix), \
             patch.object(R, "load_metrics", return_value=fake_metric(.5)), \
             patch.object(R, "run_evidence", side_effect=[(fake_metric(.6), invalid), (fake_metric(.7), valid)]):
            _, summary = R.report("best")
        self.assertEqual(summary["completed_count"], 1)
        self.assertEqual(summary["raw_metrics_count"], 2)
        self.assertNotIn("physical_r2_cb", summary["arms"]["D1"])
        self.assertEqual(summary["arms"]["D1"]["raw_metrics"]["physical_r2_cb"], .6)
        self.assertIsNone(summary["arms"]["M3"]["delta_parent_physical"])


class WorkbookOwnershipTests(unittest.TestCase):
    ARMS = [{"id": "D1", "config": "D1_a5_s1234.json", "run_name": "formal/D1", "hypothesis": "test"}]

    def sheet(self, name):
        w = Workbook()
        s = w.active
        s.title = name
        s.append(["Prior section", "=1+1"])
        title = W.TEACHER_TITLE if name == "教师汇报视图" else W.TITLE
        s.append([title])
        arm = self.ARMS[0]
        if name == "实验矩阵总览":
            s.cell(3, 1, f"{W.PREFIX}·D1 D1_a5_s1234")
            s.cell(3, 37, arm["run_name"])
        elif name == "教师汇报视图":
            s.append(["V6-A5后续", "D1 test｜训练中"])
        else:
            s.append(["ID", "父臂"])
            s.cell(4, 1, "B0=A5"); s.cell(4, 15, W.ANCHOR)
            s.cell(5, 1, "D1"); s.cell(5, 15, arm["run_name"])
            s.append(["说明", "known note"])
        return s, title

    def test_exact_own_sections_remove_cleanly_preserving_previous_formulas(self):
        for name in ("实验矩阵总览", "教师汇报视图", "汇总对比"):
            s, title = self.sheet(name)
            W.drop_own_tail(s, title, self.ARMS)
            self.assertEqual(s.max_row, 1)
            self.assertEqual(s.cell(1, 2).value, "=1+1")

    def test_exact_legacy_title_migrates_without_duplicating_sections(self):
        for name in ("实验矩阵总览", "教师汇报视图", "汇总对比"):
            s, title = self.sheet(name)
            s.cell(2, 1, ("ⅩⅣ " if name == "教师汇报视图" else "") + W.LEGACY_TITLE)
            W.drop_own_tail(s, title, self.ARMS)
            self.assertEqual(s.max_row, 1)
            self.assertEqual(s.cell(1, 2).value, "=1+1")

    def test_unknown_following_titles_unlabeled_rows_and_side_data_refused(self):
        for extra in ("unrecognized heading", None):
            s, title = self.sheet("实验矩阵总览")
            s.append([extra, "colleague data"])
            with self.assertRaisesRegex(RuntimeError, "Unowned"):
                W.drop_own_tail(s, title, self.ARMS)
            self.assertEqual(s.cell(4, 2).value, "colleague data")
        s, title = self.sheet("教师汇报视图")
        s.cell(3, 20, "outside owned columns")
        with self.assertRaisesRegex(RuntimeError, "Unowned"):
            W.drop_own_tail(s, title, self.ARMS)

    def test_changed_run_identity_refused_before_any_deletion(self):
        s, title = self.sheet("实验矩阵总览")
        s.cell(3, 37, "another-run")
        with self.assertRaisesRegex(RuntimeError, "Unknown/missing"):
            W.drop_own_tail(s, title, self.ARMS)
        self.assertEqual(s.cell(2, 1).value, title)

    def test_teacher_p99_uses_normalized_space(self):
        m = fake_metric()
        m["normalized"]["field"] = {}
        m["normalized"]["calibration"] = {"p99_pred_true_ratio": .88}
        values = W.teacher_metric_values({"p99_ratio": .61}, m)
        self.assertEqual(values[10], .88)  # E:S element 10 == column O

    def test_overview_normalized_mae_rmse_match_existing_pooled_headers(self):
        m = {"normalized": {"field": {"mae": .11, "rmse": .22},
                            "field_casebalanced": {"mae": .33, "rmse": .44}}}
        with patch.object(W, "overview_values", return_value=[None] * 49):
            values = W.followup_overview_values(self.ARMS[0], m)
        self.assertEqual(values[26:28], [.11, .22])  # AA/AB, not casebalanced values


if __name__ == "__main__":
    unittest.main()
