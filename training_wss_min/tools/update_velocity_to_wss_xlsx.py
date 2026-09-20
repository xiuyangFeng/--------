"""Append the completed VELWSS1 prediction result without touching old sections.

Uses only metrics.json['test']; oracle and physics_only are deliberately excluded.
Run with /usr/bin/python3 -m training_wss_min.tools.update_velocity_to_wss_xlsx.
--book and --experiment-dir permit isolated tests; --verify-only audits an update.
No archived updater/reporter is imported or invoked.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import tempfile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
EXP = ROOT / "training_wss_min/experiments/v5_velocity_to_wss_20260909"
SHEETS = ["汇总对比", "教师汇报视图", "RCR Oracle", "指标说明", "实验矩阵总览"]
RUN = "v5_rerun_20260906/outputs/r5v_velocity_qad_s1234"
RESULT = "v5_velocity_to_wss_20260909"
PREFIX = "V5-速度派生WSS20260909"
TITLE = "ⅩⅥ " + PREFIX + "｜VELWSS1；R5V best预测速度→冻结Profile-Secant V3；test34；seed1234"
LABEL = PREFIX + "·VELWSS1 R5V best预测速度→Profile-Secant V3"
MODEL = "PointNeXt-R + LocalGeoPE + L-SA2 + QAD-lite（18D几何/点类型输入，预测3分量速度）→Profile-Secant V3"
CHANGE = "复用已训练R5V best及其预测速度；冻结Profile-Secant V3派生WSS；原速度分量MSE训练，无新增WSS监督"
PROTOCOL = "原R5V support随机5000 / INDEPENDENT；seed1234；checkpoint=best；峰值1162；V5有效全壁面WSS评估"
NOTE = ("本节仅记录预测速度派生WSS成绩，oracle/physics_only仅作审计，不混入本表。"
        "使用既有R5V速度best（按训练损失选模），未新增训练或WSS监督；单seed1234。"
        "冻结校准器历史使用train138 WSS监督。本轮沿用旧半径解析的坐标尺度问题，详见实验报告。"
        "Pa与log_z均为WSS指标；总览M/N为Pa pooled RMSE/MAE，AA/AB为log_z pooled MAE/RMSE；"
        "教师O为log_z p99比，汇总L为Pa p99比。文件名last沿用历史命名。")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def atomic_json(path, payload):
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent,
                                     prefix=path.name + ".", suffix=".tmp", delete=False) as f:
        staged = Path(f.name)
        json.dump(payload, f, ensure_ascii=False, indent=2, allow_nan=False)
        f.write("\n")
    staged.replace(path)


def load_result(exp):
    gate = read_json(exp / "evaluation_gate.json")
    if gate.get("passed") is not True or gate.get("n_cases") != 34 or gate.get("checkpoint") != "best":
        raise ValueError("Requires passed evaluation_gate, n_cases=34, checkpoint=best")
    m = read_json(exp / "metrics.json")["test"]
    if m["aggregate"]["n_cases"] != 34 or m["normalized"]["aggregate"]["n_cases"] != 34:
        raise ValueError("Both physical and log_z metrics must contain exactly test34")
    return m


def snapshot(wb):
    return {(ws.title, cell.coordinate): cell.value for ws in wb for row in ws
            for cell in row if cell.value is not None}


def remove_owned_tail(ws):
    starts = [c.row for c in ws["A"] if c.value == TITLE]
    if not starts:
        # Refuse accidental duplicate identities with an absent/edited marker.
        if any("VELWSS1" in str(c.value or "") for row in ws for c in row):
            raise ValueError(f"Unrecognized VELWSS1 section in {ws.title}")
        return
    if len(starts) != 1:
        raise ValueError(f"Duplicate section in {ws.title}")
    first = starts[0]
    identities = {
        "实验矩阵总览": [{1: LABEL, 37: RESULT}],
        "教师汇报视图": [{1: "V5 速度派生WSS", 2: "VELWSS1｜已评估；预测速度派生"}],
        "汇总对比": [{1: "ID", 2: "速度来源"}, {1: "VELWSS1", 2: "R5V best", 15: RESULT}, {1: "说明"}],
    }[ws.title]
    width = {"实验矩阵总览": 49, "教师汇报视图": 19, "汇总对比": 15}[ws.title]
    for offset, identity in enumerate(identities, 1):
        for column, expected in identity.items():
            if ws.cell(first + offset, column).value != expected:
                raise ValueError(f"Unrecognized owned row {ws.title}!{first + offset}")
    last = first + len(identities)
    if any(c.value is not None and (c.row > last or c.column > width)
           for row in ws.iter_rows(min_row=first) for c in row):
        raise ValueError(f"Unowned data after/beside section in {ws.title}")
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= first:
            ws.unmerge_cells(str(merged))
    ws.delete_rows(first, ws.max_row - first + 1)


def append_title(ws):
    row = ws.max_row + 1
    width = {"实验矩阵总览": 20, "教师汇报视图": 19, "汇总对比": 15}[ws.title]
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    c = ws.cell(row, 1, TITLE)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 42


def append_values(ws, values, header=False):
    row = ws.max_row + 1
    for col, value in enumerate(values, 1):
        if value is not None:
            c = ws.cell(row, col, value)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            if header:
                c.fill = PatternFill("solid", fgColor="D9EAF7")
                c.font = Font(bold=True)
            elif isinstance(value, float):
                c.number_format = "0.0000"
    ws.row_dimensions[row].height = 48 if header else 72
    return row


def populate(wb, m):
    n, f, a = m["normalized"], m["field"], m["aggregate"]
    for name in ("实验矩阵总览", "教师汇报视图", "汇总对比"):
        append_title(wb[name])
    overview = [LABEL, "V5 train138/test34；18D；峰值1162；seed1234", MODEL, CHANGE, PROTOCOL, 5000,
        m["field_casebalanced"]["r2"], f["r2"], a["r2_casemean"], a["r2_casemed"], a["r2_casep10"],
        f'{a["r2_negative_cases"]}/{a["n_cases"]}', f["rmse"], f["mae"], f["nrmse_range"], a["nrmse_casemean"],
        f["nmae_range"], a["nmae_casemean"], m["regional_field"]["high_wss"]["r2"],
        m["calibration"]["top10_pred_true_ratio"], m["hotspot"]["top10_iou_casemean"],
        n["field_casebalanced"]["r2"], n["field"]["r2"], n["aggregate"]["r2_casemean"],
        n["aggregate"]["r2_casemed"], n["aggregate"]["r2_casep10"], n["field"]["mae"], n["field"]["rmse"],
        n["field"]["nrmse_range"], n["aggregate"]["nrmse_casemean"], n["field"]["nmae_range"],
        n["aggregate"]["nmae_casemean"], m["hotspot"]["spearman_all_casemean"],
        m["hotspot"]["spearman_high_wss_casemean"], n["calibration"]["top10_pred_true_ratio"],
        n["calibration"]["p99_pred_true_ratio"], RESULT,
        m["field_casebalanced"]["r2_linear_fit"], m["field_casebalanced"]["linear_fit_slope"],
        m["field_casebalanced"]["linear_fit_intercept"], f["r2_linear_fit"], f["linear_fit_slope"],
        f["linear_fit_intercept"], n["field_casebalanced"]["r2_linear_fit"],
        n["field_casebalanced"]["linear_fit_slope"], n["field_casebalanced"]["linear_fit_intercept"],
        n["field"]["r2_linear_fit"], n["field"]["linear_fit_slope"], n["field"]["linear_fit_intercept"]]
    teacher = ["V5 速度派生WSS", "VELWSS1｜已评估；预测速度派生", MODEL, PROTOCOL,
        m["field_casebalanced"]["r2"], a["r2_casemean"], f["rmse"], f["nmae_range"], a["nmae_casemean"],
        m["regional_field"]["high_wss"]["r2"], n["field_casebalanced"]["r2"], n["field"]["nmae_range"],
        m["hotspot"]["spearman_all_casemean"], m["hotspot"]["top10_iou_casemean"],
        n["calibration"]["p99_pred_true_ratio"], m["field_casebalanced"]["r2_linear_fit"],
        f["r2_linear_fit"], f["linear_fit_slope"], f["linear_fit_intercept"]]
    su = wb["汇总对比"]
    append_values(su, ["ID", "速度来源", "状态", "物理R²_cb", "物理R²_pool", "log_z R²_cb", "case mean",
                      "RMSE Pa", "MAE Pa", "high-WSS R²", "top10比", "p99比", "IoU", "变化", "结果目录"], True)
    summary = ["VELWSS1", "R5V best", "已评估；无新增训练", m["field_casebalanced"]["r2"], f["r2"],
        n["field_casebalanced"]["r2"], a["r2_casemean"], f["rmse"], f["mae"],
        m["regional_field"]["high_wss"]["r2"], m["calibration"]["top10_pred_true_ratio"],
        m["calibration"]["p99_pred_true_ratio"], m["hotspot"]["top10_iou_casemean"], CHANGE, RESULT]
    rows = {"实验矩阵总览": append_values(wb["实验矩阵总览"], overview),
            "教师汇报视图": append_values(wb["教师汇报视图"], teacher),
            "汇总对比": append_values(su, summary)}
    append_values(su, ["说明", NOTE, "原速度run：" + RUN, "主结果：预测速度→WSS；真值速度oracle仅审计"])
    return rows


# Separate, column-addressed transcription contract: does not reuse populate().
METRIC_COLUMNS = {
    "实验矩阵总览": {
        "G": "field_casebalanced.r2", "H": "field.r2", "I": "aggregate.r2_casemean",
        "J": "aggregate.r2_casemed", "K": "aggregate.r2_casep10", "M": "field.rmse", "N": "field.mae",
        "O": "field.nrmse_range", "P": "aggregate.nrmse_casemean", "Q": "field.nmae_range",
        "R": "aggregate.nmae_casemean", "S": "regional_field.high_wss.r2",
        "T": "calibration.top10_pred_true_ratio", "U": "hotspot.top10_iou_casemean",
        "V": "normalized.field_casebalanced.r2", "W": "normalized.field.r2",
        "X": "normalized.aggregate.r2_casemean", "Y": "normalized.aggregate.r2_casemed",
        "Z": "normalized.aggregate.r2_casep10", "AA": "normalized.field.mae", "AB": "normalized.field.rmse",
        "AC": "normalized.field.nrmse_range", "AD": "normalized.aggregate.nrmse_casemean",
        "AE": "normalized.field.nmae_range", "AF": "normalized.aggregate.nmae_casemean",
        "AG": "hotspot.spearman_all_casemean", "AH": "hotspot.spearman_high_wss_casemean",
        "AI": "normalized.calibration.top10_pred_true_ratio", "AJ": "normalized.calibration.p99_pred_true_ratio",
        "AL": "field_casebalanced.r2_linear_fit", "AM": "field_casebalanced.linear_fit_slope",
        "AN": "field_casebalanced.linear_fit_intercept", "AO": "field.r2_linear_fit",
        "AP": "field.linear_fit_slope", "AQ": "field.linear_fit_intercept",
        "AR": "normalized.field_casebalanced.r2_linear_fit", "AS": "normalized.field_casebalanced.linear_fit_slope",
        "AT": "normalized.field_casebalanced.linear_fit_intercept", "AU": "normalized.field.r2_linear_fit",
        "AV": "normalized.field.linear_fit_slope", "AW": "normalized.field.linear_fit_intercept"},
    "教师汇报视图": {
        "E": "field_casebalanced.r2", "F": "aggregate.r2_casemean", "G": "field.rmse", "H": "field.nmae_range",
        "I": "aggregate.nmae_casemean", "J": "regional_field.high_wss.r2", "K": "normalized.field_casebalanced.r2",
        "L": "normalized.field.nmae_range", "M": "hotspot.spearman_all_casemean", "N": "hotspot.top10_iou_casemean",
        "O": "normalized.calibration.p99_pred_true_ratio", "P": "field_casebalanced.r2_linear_fit",
        "Q": "field.r2_linear_fit", "R": "field.linear_fit_slope", "S": "field.linear_fit_intercept"},
    "汇总对比": {"D": "field_casebalanced.r2", "E": "field.r2", "F": "normalized.field_casebalanced.r2",
        "G": "aggregate.r2_casemean", "H": "field.rmse", "I": "field.mae", "J": "regional_field.high_wss.r2",
        "K": "calibration.top10_pred_true_ratio", "L": "calibration.p99_pred_true_ratio",
        "M": "hotspot.top10_iou_casemean"},
}


def verify(book, preserved, m, rows, historical_merges):
    wb = load_workbook(book, data_only=False)
    try:
        if wb.sheetnames != SHEETS:
            raise AssertionError("Workbook sheet order changed")
        checks = []
        for (sheet, coordinate), expected in preserved.items():
            if wb[sheet][coordinate].value != expected:
                raise AssertionError(f"Historical cell changed: {sheet}!{coordinate}")
        for name, merges in historical_merges.items():
            if not merges.issubset({str(r) for r in wb[name].merged_cells.ranges}):
                raise AssertionError(f"Historical merged ranges changed in {name}")
        for sheet, columns in METRIC_COLUMNS.items():
            for column, source in columns.items():
                expected = m
                for key in source.split("."):
                    expected = expected[key]
                actual = wb[sheet][f"{column}{rows[sheet]}"].value
                if (not isinstance(expected, (float, int)) or not math.isfinite(expected)
                        or not isinstance(actual, (float, int))
                        or not math.isclose(actual, expected, rel_tol=1e-13, abs_tol=1e-13)):
                    raise AssertionError(f"Metric mismatch: {sheet}!{column}{rows[sheet]} ← test.{source}")
                checks.append({"cell": f"{sheet}!{column}{rows[sheet]}", "source": "test." + source,
                               "expected": expected, "actual": actual})
        for sheet in rows:
            if sum(c.value == TITLE for c in wb[sheet]["A"]) != 1:
                raise AssertionError(f"Section missing/duplicated in {sheet}")
        if wb["实验矩阵总览"].cell(rows["实验矩阵总览"], 37).value != RESULT:
            raise AssertionError("Wrong derived-result identity")
        return {"passed": True, "source_result": "test only (oracle/physics_only excluded)",
                "preserved_nonempty_cells": len(preserved),
                "preserved_formulas": sum(isinstance(v, str) and v.startswith("=") for v in preserved.values()),
                "metric_checks": len(checks), "differences": 0, "transcription": checks}
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--book", type=Path, default=BOOK)
    parser.add_argument("--experiment-dir", type=Path, default=EXP)
    parser.add_argument("--verify-only", action="store_true")
    args = parser.parse_args()
    book, exp = args.book.resolve(), args.experiment_dir.resolve()
    m = load_result(exp)
    source_hashes = {name: sha(exp / name) for name in ("metrics.json", "evaluation_gate.json")}
    backup = exp / "WSS_PointNet实验矩阵与结果汇总last_backup_before_velocity_to_wss_20260909.xlsx"
    if args.verify_only:
        update = read_json(exp / "workbook_update.json")
        original = load_workbook(backup, data_only=False)
        preserved = snapshot(original)
        merges = {ws.title: {str(r) for r in ws.merged_cells.ranges} for ws in original}
        original.close()
        verified = verify(book, preserved, m, update["rows"], merges)
        if update["source_sha256"] != source_hashes or update["book_sha256"] != sha(book):
            raise AssertionError("Source or workbook hash changed since update")
        atomic_json(exp / "workbook_transcription_verification.json", verified)
        print(json.dumps({k: v for k, v in verified.items() if k != "transcription"}, ensure_ascii=False))
        return
    original_hash = sha(book)
    original_bytes = book.read_bytes()
    if hashlib.sha256(original_bytes).hexdigest() != original_hash:
        raise RuntimeError("Workbook changed while reading")
    with tempfile.TemporaryDirectory(prefix="velocity_wss_workbook_") as td:
        candidate = Path(td) / book.name
        candidate.write_bytes(original_bytes)
        wb = load_workbook(candidate, data_only=False)
        if wb.sheetnames != SHEETS:
            raise ValueError(f"Unexpected sheets: {wb.sheetnames}")
        for name in ("实验矩阵总览", "教师汇报视图", "汇总对比"):
            remove_owned_tail(wb[name])
        preserved = snapshot(wb)
        merges = {ws.title: {str(r) for r in ws.merged_cells.ranges} for ws in wb}
        if backup.exists():
            old = load_workbook(backup, data_only=False)
            old_cells = snapshot(old)
            old.close()
            if old_cells != preserved:
                raise RuntimeError("Historical workbook differs from initial backup; refusing replacement")
        else:
            with backup.open("xb") as f:
                f.write(original_bytes)
        rows = populate(wb, m)
        wb.save(candidate)
        wb.close()
        verified = verify(candidate, preserved, m, rows, merges)
        if sha(book) != original_hash or any(sha(exp / n) != h for n, h in source_hashes.items()):
            raise RuntimeError("Workbook or evaluation changed concurrently; refusing replacement")
        fd, staged_name = tempfile.mkstemp(prefix=book.stem + ".velocity_to_wss.", suffix=".xlsx", dir=book.parent)
        staged = Path(staged_name)
        try:
            with os.fdopen(fd, "wb") as f:
                f.write(candidate.read_bytes())
                f.flush()
                os.fsync(f.fileno())
            os.chmod(staged, book.stat().st_mode & 0o7777)
            if sha(book) != original_hash:
                raise RuntimeError("Workbook changed before atomic replacement")
            staged.replace(book)
        finally:
            staged.unlink(missing_ok=True)
    metadata = {"book": str(book), "backup": str(backup), "experiment": "VELWSS1", "checkpoint": "best",
                "n_cases": 34, "source_run": RUN, "source_result": "test", "rows": rows,
                "overview_rows": {"VELWSS1": rows["实验矩阵总览"]},
                "teacher_rows": {"VELWSS1": rows["教师汇报视图"]},
                "summary_rows": {"VELWSS1": rows["汇总对比"]},
                "source_sha256": source_hashes, "book_sha256_before": original_hash, "book_sha256": sha(book),
                "backup_sha256": sha(backup), "preserved_nonempty_cells": len(preserved),
                "preserved_formulas": verified["preserved_formulas"], "metric_checks": verified["metric_checks"]}
    atomic_json(exp / "workbook_update.json", metadata)
    atomic_json(exp / "workbook_transcription_verification.json", verified)
    print(json.dumps(metadata, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
