#!/usr/bin/env python3
"""Update the compact WSS workbook with the LSA2-H2 log-radius screen."""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from training_wss_min.tools.update_s3_rootcause_xlsx import copy_row_style


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = (
    ROOT
    / "training_wss_min/preflight/"
    "lsa2_h2_logradius_matrix_20260729_results_analysis.json"
)
SECTION_TITLE = (
    "2026-07-29｜LSA2 SAME-H2：并发复现 / +log(local_radius)"
)
CONTROL = "lsa2_h2_same_repro_s1234"
TREATMENT = "lsa2_h2_logradius_s1234"
ORDER = (CONTROL, TREATMENT)
LABELS = {
    CONTROL: "LSA2 H2 concurrent repro s1234",
    TREATMENT: "★ LSA2 H2 + log(radius) s1234",
}
CHANGES = {
    CONTROL: "原6D输入；SAME-H2精确并发复现",
    TREATMENT: "原6D末尾新增train-only标准化log(local_radius)",
}
EXPECTED_SHEETS = [
    "汇总对比",
    "教师汇报视图",
    "RCR Oracle",
    "指标说明",
    "实验矩阵总览",
]

NAVY = "1F4E78"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def raw_metrics(record: dict) -> dict:
    return read_json(
        Path(record["run_dir"]) / "eval/ckpt_best/metrics.json"
    )["test"]


def overview_values(record: dict) -> list:
    metrics = raw_metrics(record)
    normalized = metrics["normalized"]
    experiment_id = record["experiment_id"]
    return [
        LABELS[experiment_id],
        "138/0/36（AG/AAA/ILO mixed；seed=1234）",
        "PointNeXt-R + LocalGeoPE + L-SA2",
        CHANGES[experiment_id],
        "random5000 / SAME；FPS center",
        5000,
        metrics["field_casebalanced"]["r2"],
        metrics["field"]["r2"],
        metrics["aggregate"]["r2_casemean"],
        metrics["aggregate"]["r2_casemed"],
        metrics["aggregate"]["r2_casep10"],
        (
            f"{metrics['aggregate']['r2_negative_cases']}/"
            f"{metrics['aggregate']['n_cases']}"
        ),
        metrics["field"]["rmse"],
        metrics["field"]["mae"],
        metrics["field"]["nrmse_range"],
        metrics["aggregate"]["nrmse_casemean"],
        metrics["field"].get("nmae_range"),
        metrics["aggregate"].get("nmae_casemean"),
        metrics["regional_field"]["high_wss"]["r2"],
        metrics["calibration"]["top10_pred_true_ratio"],
        metrics["hotspot"]["top10_iou_casemean"],
        normalized["field_casebalanced"]["r2"],
        normalized["field"]["r2"],
        normalized["aggregate"]["r2_casemean"],
        normalized["aggregate"]["r2_casemed"],
        normalized["aggregate"]["r2_casep10"],
        normalized["field_casebalanced"]["mae"],
        normalized["field_casebalanced"]["rmse"],
        normalized["field"]["nrmse_range"],
        normalized["aggregate"]["nrmse_casemean"],
        normalized["field"].get("nmae_range"),
        normalized["aggregate"].get("nmae_casemean"),
        metrics["hotspot"]["spearman_all_casemean"],
        metrics["hotspot"]["spearman_high_wss_casemean"],
        normalized["calibration"]["top10_pred_true_ratio"],
        normalized["calibration"]["p99_pred_true_ratio"],
        experiment_id,
    ]


def teacher_values(record: dict) -> list:
    metrics = raw_metrics(record)
    normalized = metrics["normalized"]
    return [
        "LSA2 H2 log-radius",
        LABELS[record["experiment_id"]],
        "PointNeXt-R + LocalGeoPE + L-SA2",
        "random5000 / SAME",
        metrics["field_casebalanced"]["r2"],
        metrics["aggregate"]["r2_casemean"],
        metrics["field"]["rmse"],
        metrics["field"].get("nmae_range"),
        metrics["aggregate"].get("nmae_casemean"),
        metrics["regional_field"]["high_wss"]["r2"],
        normalized["field_casebalanced"]["r2"],
        normalized["field"].get("nmae_range"),
        metrics["hotspot"]["spearman_all_casemean"],
        metrics["hotspot"]["top10_iou_casemean"],
        normalized["calibration"]["p99_pred_true_ratio"],
    ]


def add_title(sheet, row: int) -> int:
    sheet.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=15
    )
    cell = sheet.cell(row, 1, SECTION_TITLE)
    cell.font = Font(bold=True, color=WHITE, size=12)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 23
    return row + 1


def add_header(sheet, row: int) -> int:
    headers = [
        "臂",
        "唯一变化",
        "R²_cb",
        "ΔR²_cb",
        "normalized R²_cb",
        "Δnormalized",
        "MAE_cb (Pa)",
        "ΔMAE",
        "high-WSS nRMSE",
        "ΔnRMSE",
        "top10 IoU",
        "ΔIoU",
        "AG / AAA / ILO R²",
        "预注册 Gate",
        "判定",
    ]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    sheet.row_dimensions[row].height = 36
    return row + 1


def add_summary_row(
    sheet,
    row: int,
    record: dict,
    paired: dict,
    gate: dict,
) -> None:
    experiment_id = record["experiment_id"]
    if experiment_id == TREATMENT:
        delta = paired["aggregate_treatment_minus_control"]
        extra = paired["extra_treatment_minus_control"]
        gate_text = (
            f"ΔR²={delta['physical_r2_casebalanced']:+.4f}；"
            f"ΔnRMSE={extra['high_wss_nrmse_range']:+.5f}"
        )
        decision = "Go；新单seed开发锚点"
    else:
        delta = {
            "physical_r2_casebalanced": None,
            "normalized_r2_casebalanced": None,
            "physical_mae": None,
            "physical_top10_iou": None,
        }
        extra = {"high_wss_nrmse_range": None}
        gate_text = "正式并发对照"
        decision = "并发基准"
    values = [
        LABELS[experiment_id],
        CHANGES[experiment_id],
        record["best"]["physical_r2_casebalanced"],
        delta["physical_r2_casebalanced"],
        record["best"]["normalized_r2_casebalanced"],
        delta["normalized_r2_casebalanced"],
        record["best"]["physical_mae"],
        delta["physical_mae"],
        record["extra"]["high_wss_nrmse_range"],
        extra["high_wss_nrmse_range"],
        record["best"]["physical_top10_iou"],
        delta["physical_top10_iou"],
        (
            f"{record['domains']['AG']:.3f} / "
            f"{record['domains']['AAA']:.3f} / "
            f"{record['domains']['ILO']:.3f}"
        ),
        gate_text,
        decision,
    ]
    fill = GREEN if gate["decision"].startswith("go_") else YELLOW
    if experiment_id == CONTROL:
        fill = BLUE
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if column in range(3, 13):
            cell.number_format = (
                "+0.0000;-0.0000;0.0000"
                if column in (4, 6, 8, 10, 12)
                else "0.0000"
            )
    sheet.row_dimensions[row].height = 44


def render_and_count_pages(workbook_path: Path, out_dir: Path) -> tuple[Path, int]:
    subprocess.run(
        [
            "/usr/bin/libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(workbook_path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    pdf = out_dir / workbook_path.with_suffix(".pdf").name
    output = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    for line in output.splitlines():
        if line.startswith("Pages:"):
            return pdf, int(line.split(":", 1)[1].strip())
    raise RuntimeError("pdfinfo did not report page count")


def main() -> None:
    analysis = read_json(ANALYSIS)
    records = analysis["records"]
    paired = analysis["paired_comparison"]
    gate = analysis["registered_gate"]
    if gate["decision"] != "go_promote_logradius_as_single_seed_development_anchor":
        raise RuntimeError("unexpected registered Gate decision")
    for experiment_id in ORDER:
        if records[experiment_id]["integrity"] != "passed":
            raise RuntimeError(f"integrity failed: {experiment_id}")

    with tempfile.TemporaryDirectory(prefix="wss_lsa2_h2_logradius_xlsx_") as temp_dir:
        temp = Path(temp_dir)
        candidate = temp / BOOK.name
        shutil.copy2(BOOK, candidate)
        workbook = load_workbook(candidate, data_only=False)
        if workbook.sheetnames != EXPECTED_SHEETS:
            raise RuntimeError(
                f"unexpected sheet set/order: {workbook.sheetnames}"
            )

        overall = workbook["实验矩阵总览"]
        old_rows = [
            cell.row for cell in overall["AK"] if cell.value in set(ORDER)
        ]
        for row in sorted(old_rows, reverse=True):
            overall.delete_rows(row, 1)
        if overall.max_row != 192:
            raise RuntimeError(f"unexpected overview base rows: {overall.max_row}")
        first_overall = overall.max_row + 1
        for experiment_id in ORDER:
            row = overall.max_row + 1
            copy_row_style(overall, row - 1, row, 37)
            for column, value in enumerate(
                overview_values(records[experiment_id]), 1
            ):
                overall.cell(row, column, value)

        teacher = workbook["教师汇报视图"]
        old_rows = [
            cell.row
            for cell in teacher["B"]
            if cell.value in set(LABELS.values())
        ]
        for row in sorted(old_rows, reverse=True):
            teacher.delete_rows(row, 1)
        if teacher.max_row != 166:
            raise RuntimeError(f"unexpected teacher base rows: {teacher.max_row}")
        first_teacher = teacher.max_row + 1
        for experiment_id in ORDER:
            row = teacher.max_row + 1
            copy_row_style(teacher, row - 1, row, 15)
            for column, value in enumerate(
                teacher_values(records[experiment_id]), 1
            ):
                teacher.cell(row, column, value)

        summary = workbook["汇总对比"]
        old_title_rows = [
            cell.row for cell in summary["A"] if cell.value == SECTION_TITLE
        ]
        if len(old_title_rows) > 1:
            raise RuntimeError("duplicate LSA2 H2 log-radius summary sections")
        if old_title_rows:
            summary.delete_rows(old_title_rows[0], 2 + len(ORDER))
        if summary.max_row != 265:
            raise RuntimeError(f"unexpected summary base rows: {summary.max_row}")
        summary_start = summary.max_row + 1
        row = add_title(summary, summary_start)
        row = add_header(summary, row)
        for experiment_id in ORDER:
            add_summary_row(
                summary,
                row,
                records[experiment_id],
                paired,
                gate,
            )
            row += 1
        for column, width in {
            "A": 30,
            "B": 34,
            "C": 9,
            "D": 9,
            "E": 12,
            "F": 10,
            "G": 11,
            "H": 9,
            "I": 12,
            "J": 9,
            "K": 10,
            "L": 9,
            "M": 18,
            "N": 25,
            "O": 19,
        }.items():
            summary.column_dimensions[column].width = width
        summary.print_area = (
            f"A{summary_start}:O{summary_start + len(ORDER) + 1}"
        )
        summary.page_setup.fitToWidth = 1
        summary.page_setup.fitToHeight = 1
        summary.sheet_view.selection[0].activeCell = f"A{summary_start}"
        summary.sheet_view.selection[0].sqref = f"A{summary_start}"

        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(candidate)
        workbook.close()

        check = load_workbook(candidate, read_only=True, data_only=False)
        if check.sheetnames != EXPECTED_SHEETS:
            raise RuntimeError("readback sheet set changed")
        saved_ids = {
            values[0]
            for values in check["实验矩阵总览"].iter_rows(
                min_col=37, max_col=37, values_only=True
            )
        }
        saved_labels = {
            values[0]
            for values in check["教师汇报视图"].iter_rows(
                min_col=2, max_col=2, values_only=True
            )
        }
        saved_titles = {
            values[0]
            for values in check["汇总对比"].iter_rows(
                min_col=1, max_col=1, values_only=True
            )
        }
        if not set(ORDER).issubset(saved_ids):
            raise RuntimeError("overview readback missing log-radius rows")
        if not set(LABELS.values()).issubset(saved_labels):
            raise RuntimeError("teacher readback missing log-radius rows")
        if SECTION_TITLE not in saved_titles:
            raise RuntimeError("summary readback missing log-radius section")
        if check["实验矩阵总览"].max_row != 194:
            raise RuntimeError("overview row count is not idempotent")
        if check["教师汇报视图"].max_row != 168:
            raise RuntimeError("teacher row count is not idempotent")
        if check["汇总对比"].max_row != 269:
            raise RuntimeError("summary row count is not idempotent")
        check.close()

        render_dir = temp / "render"
        render_dir.mkdir()
        pdf, pages = render_and_count_pages(candidate, render_dir)
        if pages != 6:
            raise RuntimeError(
                f"workbook render grew from the compact 6-page layout to {pages}"
            )
        shutil.copy2(candidate, BOOK)

    print(
        json.dumps(
            {
                "status": "updated_and_rendered",
                "workbook": str(BOOK),
                "sheet_count": len(EXPECTED_SHEETS),
                "pdf_pages": pages,
                "overall_rows": [first_overall, first_overall + len(ORDER) - 1],
                "teacher_rows": [first_teacher, first_teacher + len(ORDER) - 1],
                "summary_rows": [
                    summary_start,
                    summary_start + len(ORDER) + 1,
                ],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
