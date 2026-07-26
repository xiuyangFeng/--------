#!/usr/bin/env python3
"""Append the completed S3-GEOPE root-cause matrix to the WSS results workbook.

Appends the 12 new arms to 实验矩阵总览 / 教师汇报视图, and a per-domain
"vs S3 parent" comparison block (with three-seed means) to 汇总对比.  Mirrors the
D2-K64 updater: temp copy -> style-preserving append -> readback -> overwrite.
"""
from __future__ import annotations

import json
import shutil
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = ROOT / "training_wss_min/preflight/s3_rootcause_results_analysis.json"
SECTION_TITLE = "2026-07-24｜S3-GEOPE 根因矩阵（口径/域条件/尾部；vs S3 父模板）"

# experiment_id -> (label, capacity/change)  — all arms share model / split shape / sampling
ORDER = (
    "caliber_pooled138_s1234",
    "caliber_casebal_s1234", "caliber_casebal_s7", "caliber_casebal_s2025",
    "caliber_cohortbal_s1234", "caliber_cohortbal_s7", "caliber_cohortbal_s2025",
    "cohort_onehot_s1234", "cohort_onehot_s7", "cohort_onehot_s2025",
    "rawhuber02_s1234", "cohortbal_onehot_s1234",
)
CHANGE = {
    "caliber_pooled138": ("RC0 S3+口径pooled138", "S3 + 目标统计 train138 pooled 重算"),
    "caliber_casebal": ("RC1 S3+口径casebal(A1)", "S3 + 逐病例等权 log-z 统计"),
    "caliber_cohortbal": ("RC2 S3+口径cohortbal(A1b)", "S3 + 逐父系等权 log-z 统计"),
    "cohort_onehot": ("RC3 S3+cohort条件", "S3 + cohort one-hot（input_dim 6→9）"),
    "rawhuber02": ("RC4 S3+rawHuber0.2", "S3 + 物理 raw-Huber λ=0.2"),
    "cohortbal_onehot": ("RC5 S3+cohortbal+onehot", "S3 + cohortbal + cohort one-hot（非归因）"),
}


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def label_of(rec: dict) -> str:
    base, _ = CHANGE[rec["variant"]]
    return f"{base} s{rec['seed']}"


def copy_row_style(ws, source_row: int, target_row: int, columns: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, columns + 1):
        source, target = ws.cell(source_row, col), ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def matrix_values(rec: dict) -> list:
    _, capacity = CHANGE[rec["variant"]]
    raw = load_json(Path(rec["run_dir"]) / "eval/ckpt_best/metrics.json")["test"]
    norm = raw["normalized"]
    split = f"138/0/36（AG/AAA/ILO mixed；seed={rec['seed']}）"
    return [
        label_of(rec), split, "PointNeXt-R + LocalGeoPE", capacity, "random / SAME；FPS center", 5000,
        raw["field_casebalanced"]["r2"], raw["field"]["r2"], raw["aggregate"]["r2_casemean"],
        raw["aggregate"]["r2_casemed"], raw["aggregate"]["r2_casep10"],
        f"{raw['aggregate']['r2_negative_cases']}/{raw['aggregate']['n_cases']}",
        raw["field"]["rmse"], raw["field"]["mae"], raw["field"]["nrmse_range"],
        raw["aggregate"]["nrmse_casemean"], raw["field"].get("nmae_range"), raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"], raw["calibration"]["top10_pred_true_ratio"],
        raw["hotspot"]["top10_iou_casemean"], norm["field_casebalanced"]["r2"], norm["field"]["r2"],
        norm["aggregate"]["r2_casemean"], norm["aggregate"]["r2_casemed"], norm["aggregate"]["r2_casep10"],
        norm["field_casebalanced"]["mae"], norm["field_casebalanced"]["rmse"], norm["field"]["nrmse_range"],
        norm["aggregate"]["nrmse_casemean"], norm["field"].get("nmae_range"), norm["aggregate"].get("nmae_casemean"),
        raw["hotspot"]["spearman_all_casemean"], raw["hotspot"]["spearman_high_wss_casemean"],
        norm["calibration"]["top10_pred_true_ratio"], norm["calibration"]["p99_pred_true_ratio"],
        rec["experiment_id"],
    ]


def teacher_values(rec: dict) -> list:
    raw = load_json(Path(rec["run_dir"]) / "eval/ckpt_best/metrics.json")["test"]
    norm = raw["normalized"]
    return [
        "S3根因矩阵", label_of(rec), "PointNeXt-R + LocalGeoPE", "5000 / SAME",
        raw["field_casebalanced"]["r2"], raw["aggregate"]["r2_casemean"], raw["field"]["rmse"],
        raw["field"].get("nmae_range"), raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"], norm["field_casebalanced"]["r2"], norm["field"].get("nmae_range"),
        raw["hotspot"]["spearman_all_casemean"], raw["hotspot"]["top10_iou_casemean"],
        norm["calibration"]["p99_pred_true_ratio"],
    ]


def append_summary(summary, analysis: dict) -> int:
    records = {r["experiment_id"]: r for r in analysis["records"].values()} \
        if isinstance(analysis["records"], dict) else {r["experiment_id"]: r for r in analysis["records"]}
    row = summary.max_row + 1
    span = 11
    summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    title = summary.cell(row, 1, SECTION_TITLE)
    title.font = Font(bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E78")
    start = row
    row += 1
    summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    summary.cell(row, 1, "每臂对照三种子确认的 S3-GEOPE 父模板（同种子配对）；"
                         "分域 Δ 正=回补/负=退化；单种子臂仅作筛选。").font = Font(bold=True)
    row += 1
    headers = ["处理臂", "种子", "R²_cb", "ΔR²_cb", "ΔAG R²", "ΔAAA R²", "ΔILO R²",
               "high-WSS R²", "top10 IoU", "病例R²胜/负", "病例均值ΔR² 95%CI"]
    for col, value in enumerate(headers, 1):
        cell = summary.cell(row, col, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    pairs = [c for c in analysis["paired_comparisons"] if c["comparison"].endswith("_vs_s3parent")]
    for c in pairs:
        row += 1
        t = records[c["treatment"]]
        agg = c["aggregate_treatment_minus_control"]
        dom = c["domain_r2_delta"]
        st = c["paired_case_stats"]["case_r2"]
        values = [
            label_of(t), t["seed"], t["best"]["physical_r2_casebalanced"], agg["physical_r2_casebalanced"],
            dom["AG"], dom["AAA"], dom["ILO"], t["best"]["high_wss_r2"], t["best"]["physical_top10_iou"],
            f"{st['wins']}/{st['losses']}",
            f"{st['mean_delta']:+.4f} [{st['ci95_low']:+.4f}, {st['ci95_high']:+.4f}]",
        ]
        for col, value in enumerate(values, 1):
            summary.cell(row, col, value)
    # three-seed mean block
    row += 1
    summary.cell(row, 1, "三种子 vs S3 父（同种子）ΔR²_cb 均值").font = Font(bold=True, italic=True)
    for ts in analysis["three_seed_vs_s3parent"]:
        row += 1
        per = "; ".join(f"s{s}:{v:+.4f}" for s, v in ts["delta_r2cb_per_seed"].items())
        summary.cell(row, 1, CHANGE[ts["variant"]][0])
        summary.cell(row, 3, ts["mean"])
        summary.cell(row, 4, f"{ts['wins']}胜/{ts['losses']}负")
        summary.cell(row, 5, per)
    return start


def main() -> None:
    analysis = load_json(ANALYSIS)
    recs = analysis["records"]
    records = recs if isinstance(recs, dict) else {r["experiment_id"]: r for r in recs}
    for exp_id in ORDER:
        if exp_id not in records:
            raise RuntimeError(f"analysis missing {exp_id}")
        if records[exp_id]["integrity"] != "passed":
            raise RuntimeError(f"integrity failed: {exp_id}")

    with tempfile.TemporaryDirectory(prefix="wss_s3rc_xlsx_") as temp_dir:
        candidate = Path(temp_dir) / XLSX.name
        shutil.copy2(XLSX, candidate)
        workbook = load_workbook(candidate, data_only=False)

        overall = workbook["实验矩阵总览"]
        existing_ids = {cell.value for cell in overall["AK"]}
        dups = set(ORDER) & existing_ids
        if dups:
            raise RuntimeError(f"refusing duplicate append: {sorted(dups)}")
        first_overall = overall.max_row + 1
        for exp_id in ORDER:
            row = overall.max_row + 1
            copy_row_style(overall, row - 1, row, 37)
            for col, value in enumerate(matrix_values(records[exp_id]), 1):
                overall.cell(row, col, value)

        teacher = workbook["教师汇报视图"]
        existing_teacher = {cell.value for cell in teacher["B"]}
        labels = [label_of(records[e]) for e in ORDER]
        if any(lbl in existing_teacher for lbl in labels):
            raise RuntimeError("refusing duplicate append in 教师汇报视图")
        first_teacher = teacher.max_row + 1
        for exp_id in ORDER:
            row = teacher.max_row + 1
            copy_row_style(teacher, row - 1, row, 15)
            for col, value in enumerate(teacher_values(records[exp_id]), 1):
                teacher.cell(row, col, value)

        summary = workbook["汇总对比"]
        if SECTION_TITLE in {cell.value for cell in summary["A"]}:
            raise RuntimeError("refusing duplicate append in 汇总对比")
        summary_start = append_summary(summary, analysis)
        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(candidate)

        check = load_workbook(candidate, read_only=True, data_only=False)
        check_ids = {r[0].value for r in check["实验矩阵总览"].iter_rows(min_col=37, max_col=37)}
        if not set(ORDER).issubset(check_ids):
            raise RuntimeError("saved workbook failed 实验矩阵总览 readback")
        teacher_labels = {r[0].value for r in check["教师汇报视图"].iter_rows(min_col=2, max_col=2)}
        if not all(lbl in teacher_labels for lbl in labels):
            raise RuntimeError("saved workbook failed 教师汇报视图 readback")
        summary_titles = {r[0].value for r in check["汇总对比"].iter_rows(min_col=1, max_col=1)}
        if SECTION_TITLE not in summary_titles:
            raise RuntimeError("saved workbook failed 汇总对比 readback")
        check.close()
        shutil.copy2(candidate, XLSX)

    print(json.dumps({"status": "updated", "workbook": str(XLSX),
                      "overall_rows": [first_overall, first_overall + len(ORDER) - 1],
                      "teacher_rows": [first_teacher, first_teacher + len(ORDER) - 1],
                      "summary_start_row": summary_start, "rows_added": len(ORDER)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
