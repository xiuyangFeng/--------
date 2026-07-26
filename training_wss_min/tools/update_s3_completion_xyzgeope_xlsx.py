#!/usr/bin/env python3
"""Append completed S3 regularization and xyz-LocalGeoPE results to two workbook views."""
from __future__ import annotations

import json, shutil, subprocess, tempfile
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from training_wss_min.tools.update_s3_regularization_xlsx import matrix_values
from training_wss_min.tools.update_s3_rootcause_xlsx import copy_row_style

ROOT=Path(__file__).resolve().parents[2]
BOOK=ROOT/"docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
REG=ROOT/"training_wss_min/preflight/s3_regularization_completion_20260725_results_analysis.json"
XYZ=ROOT/"training_wss_min/preflight/xyz_local_geope_20260725_results_analysis.json"
REG_TITLE="2026-07-25｜S3 正则化强度补齐与两两交叉（seed1234）"
XYZ_TITLE="2026-07-25｜S2 xyz-only + 4D LocalGeoPE（seed1234）"

def overview_values(record,label,change):
    v=matrix_values({**record,"variant":"head_dropout01","seed":1234})
    v[0:4]=[label,"138/0/36（AG/AAA/ILO mixed；seed=1234）","PointNeXt-R + LocalGeoPE",change]
    return v
def title(ws,row,text):
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=12);c=ws.cell(row,1,text);c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="1F4E78");return row+1
def header(ws,row,vals):
    for i,v in enumerate(vals,1):c=ws.cell(row,i,v);c.font=Font(bold=True);c.fill=PatternFill("solid",fgColor="D9EAF7")
    return row+1
def main():
    reg=json.loads(REG.read_text(encoding="utf-8"));xyz=json.loads(XYZ.read_text(encoding="utf-8"))
    completion=[r["experiment_id"] for r in json.loads((ROOT/"training_wss_min/preflight/s3_regularization_completion_20260725_prepared.json").read_text())["configs"]]
    ids=set(completion)|{"s2_d2k64_pnxr_mixed_xyz_geope"}
    with tempfile.TemporaryDirectory(prefix="wss_completion_xlsx_") as tmp:
        candidate=Path(tmp)/BOOK.name;shutil.copy2(BOOK,candidate);wb=load_workbook(candidate,data_only=False)
        overall=wb["实验矩阵总览"]
        for row in sorted([cell.row for cell in overall["AK"] if cell.value in ids],reverse=True):overall.delete_rows(row,1)
        for eid in completion:
            row=overall.max_row+1;copy_row_style(overall,row-1,row,37);record=reg["records"][eid]
            fields=json.loads((Path(record["run_dir"])/"config.json").read_text())["model"]
            change=" + ".join(f"{k.replace('drop_path_rate','DropPath').replace('neighbor_drop_rate','NeighborDrop').replace('dropout','HeadDrop')}={fields[k]:.2f}" for k in ("dropout","drop_path_rate","neighbor_drop_rate") if float(fields.get(k,0))>0)
            for col,value in enumerate(overview_values(record,eid,change),1):overall.cell(row,col,value)
        geo=xyz["records"]["s2_d2k64_pnxr_mixed_xyz_geope"];row=overall.max_row+1;copy_row_style(overall,row-1,row,37)
        v=matrix_values({**geo,"variant":"head_dropout01","seed":1234});v[0:4]=["S2-PNXR xyz + LocalGeoPE","138/0/36（AG/AAA/ILO mixed；seed=1234）","PointNeXt-R + XYZ-LocalGeoPE","4D Δxyz/r + distance；无语义属性差分"]
        for col,value in enumerate(v,1):overall.cell(row,col,value)
        summary=wb["汇总对比"]
        existing={c.value:c.row for c in summary["A"] if c.value in {REG_TITLE,XYZ_TITLE}}
        if existing: raise RuntimeError("completion summary already exists")
        row=summary.max_row+1;row=title(summary,row,REG_TITLE);row=header(summary,row,["变体", "R²_cb", "ΔR²_cb vs S3", "ΔMAE", "ΔRMSE", "Δhigh-WSS", "pair interaction ΔR²", "结论"])
        parent=reg["records"]["s3_geope_s1234"]["best"]["physical_r2_casebalanced"]
        pairs={x["experiment_id"]:x for x in reg["pairwise_interactions"]}
        deltas={x["treatment"]:x["aggregate_treatment_minus_control"] for x in reg["paired_comparisons"]}
        for eid in completion:
            d=deltas[eid];r=reg["records"][eid];pair=pairs.get(eid)
            conclusion="pairwise" if pair else "single-rate"
            summary.append([eid,r["best"]["physical_r2_casebalanced"],d["physical_r2_casebalanced"],d["physical_mae"],d["physical_rmse"],d["high_wss_r2"],None if pair is None else pair["interaction_r2cb"],conclusion])
        row=summary.max_row+1;row=title(summary,row,XYZ_TITLE);row=header(summary,row,["比较", "R²_cb", "ΔR²_cb", "ΔMAE", "ΔRMSE", "Δhigh-WSS", "病例 ΔR² 95%CI", "结论"])
        for result in xyz["paired_comparisons"]:
            d,c=result["aggregate_treatment_minus_control"],result["paired_case_stats"]["case_r2"]
            treatment=xyz["records"][result["treatment"]]
            summary.append([result["comparison"],treatment["best"]["physical_r2_casebalanced"],d["physical_r2_casebalanced"],d["physical_mae"],d["physical_rmse"],d["high_wss_r2"],f"{c['mean_delta']:+.4f} [{c['ci95_low']:+.4f},{c['ci95_high']:+.4f}]","vs xyz" if result["control"]=="s2_pnxr_xyz" else "vs xyz+geom"])
        wb.save(candidate);wb.close()
        check=load_workbook(candidate,read_only=True,data_only=False);saved={r[0] for r in check["实验矩阵总览"].iter_rows(min_col=37,max_col=37,values_only=True)}
        if not ids.issubset(saved) or not {REG_TITLE,XYZ_TITLE}.issubset({r[0] for r in check["汇总对比"].iter_rows(min_col=1,max_col=1,values_only=True)}):raise RuntimeError("xlsx readback failed")
        check.close();render=Path(tmp)/"render";render.mkdir();subprocess.run(["/usr/bin/libreoffice","--headless","--convert-to","pdf","--outdir",str(render),str(candidate)],check=True,capture_output=True,text=True)
        if not (render/candidate.with_suffix(".pdf").name).is_file():raise RuntimeError("LibreOffice render check failed")
        shutil.copy2(candidate,BOOK)
    print(json.dumps({"status":"updated_and_rendered","overview_rows":len(ids),"summary_sections":2,"workbook":str(BOOK)},ensure_ascii=False))
if __name__=="__main__":main()
