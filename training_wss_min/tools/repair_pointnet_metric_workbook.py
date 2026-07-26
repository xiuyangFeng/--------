#!/usr/bin/env python3
"""Repair scope and copy errors in the teacher-facing PointNet metric workbook.

The four legacy GLOBAL rows were populated with case-balanced MAE/RMSE while
the identically named v4/Phase-V columns use pooled field values.  The sheet
does not label those columns as case-balanced, so this script normalizes them
to the pooled definition and repairs the E4 normalized p99 copy error.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"

SOURCES = {
    "E2-GLOBAL": ROOT / "training_wss_min/runs/pointnet_distribution_matrix/outputs/e2_global/eval/ckpt_best/metrics.json",
    "E3-GLOBAL": ROOT / "training_wss_min/runs/pointnet_distribution_matrix/outputs/e3_global/eval/ckpt_best/metrics.json",
    "E23-GLOBAL": ROOT / "training_wss_min/runs/pointnet_distribution_matrix/outputs/e23_global/eval/ckpt_best/metrics.json",
    "E4-DEEP-GLOBAL": ROOT / "training_wss_min/runs/pointnet_deeper/outputs/e4_deep_global/eval/ckpt_best/metrics.json",
}

NORMALIZED_ONLY_SOURCES = {
    "E2-CASE": ROOT / "training_wss_min/runs/pointnet_distribution_matrix/outputs/e2_case/eval/ckpt_best/metrics.json",
    "E3-CASE": ROOT / "training_wss_min/runs/pointnet_distribution_matrix/outputs/e3_case/eval/ckpt_best/metrics.json",
}


def _row_by_id(ws, experiment_id: str) -> int:
    for row in range(3, ws.max_row + 1):
        if ws.cell(row, 1).value == experiment_id:
            return row
    raise KeyError(f"experiment ID not found: {experiment_id}")


def _teacher_row_by_id(ws, experiment_id: str) -> int:
    for row in range(5, ws.max_row + 1):
        value = str(ws.cell(row, 2).value or "").replace("★ ", "")
        if value == experiment_id:
            return row
    raise KeyError(f"teacher-view experiment ID not found: {experiment_id}")


def main() -> None:
    wb = load_workbook(WORKBOOK)
    detail = wb["实验矩阵总览"]
    teacher = wb["教师汇报视图"]

    detail["M2"] = "RMSE (Pa, pooled)"
    detail["N2"] = "MAE (Pa, pooled)"
    detail["AA2"] = "MAE (pooled)"
    detail["AB2"] = "RMSE (pooled)"
    definition = "pooled：拼接该 test split 的全部点后计算；不等同于 case-balanced 指标。"
    for coord in ("M2", "N2", "AA2", "AB2"):
        detail[coord].comment = Comment(definition, "Codex")
    teacher["G3"] = "RMSE (Pa, pooled)"

    for experiment_id, source in SOURCES.items():
        metrics = json.loads(source.read_text())["test"]
        normalized = metrics["normalized"]
        row = _row_by_id(detail, experiment_id)
        detail.cell(row, 13).value = metrics["field"]["rmse"]
        detail.cell(row, 14).value = metrics["field"]["mae"]
        detail.cell(row, 27).value = normalized["field"]["mae"]
        detail.cell(row, 28).value = normalized["field"]["rmse"]
        teacher_row = _teacher_row_by_id(teacher, experiment_id)
        teacher.cell(teacher_row, 7).value = metrics["field"]["rmse"]

    for experiment_id, source in NORMALIZED_ONLY_SOURCES.items():
        metrics = json.loads(source.read_text())["test"]
        row = _row_by_id(detail, experiment_id)
        detail.cell(row, 27).value = metrics["field"]["mae"]
        detail.cell(row, 28).value = metrics["field"]["rmse"]

    e4 = json.loads(SOURCES["E4-DEEP-GLOBAL"].read_text())["test"]["normalized"]
    detail.cell(_row_by_id(detail, "E4-DEEP-GLOBAL"), 36).value = e4["calibration"]["p99_pred_true_ratio"]
    teacher.cell(_teacher_row_by_id(teacher, "E4-DEEP-GLOBAL"), 15).value = e4["calibration"]["p99_pred_true_ratio"]

    wb.save(WORKBOOK)
    print(WORKBOOK)


if __name__ == "__main__":
    main()
