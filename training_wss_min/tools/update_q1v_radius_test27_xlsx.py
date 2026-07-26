#!/usr/bin/env python3
"""Add submitted Q1V/SAME radius-only arms to the experiment workbook.

The rows intentionally contain no metrics.  They are planning/submission
records only and are kept separate from the completed Q2V historical-test27
exploration until all jobs have passed their post-training checks.
"""

from __future__ import annotations

import argparse
import copy
import json
from pathlib import Path

from openpyxl import load_workbook


REPO = Path(__file__).resolve().parents[2]
WORKBOOK = REPO / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
SUBMISSION = REPO / "training_wss_min/preflight/q1v_radius_test27_submission.json"


ARMS = (
    ("q1v_radius_r80_same", "Q1V半径探索(0.8×，已提交)", "0.04/0.08/0.16"),
    ("q1v_radius_r120_same", "Q1V半径探索(1.2×，已提交)", "0.06/0.12/0.24"),
    ("q1v_radius_r150_same", "Q1V半径探索(1.5×，已提交)", "0.075/0.15/0.30"),
)


def copy_row(ws, source: int, target: int) -> None:
    for col in range(1, ws.max_column + 1):
        src, dst = ws.cell(source, col), ws.cell(target, col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.fill:
            dst.fill = copy.copy(src.fill)
        if src.font:
            dst.font = copy.copy(src.font)
        if src.border:
            dst.border = copy.copy(src.border)
        dst.value = None
    ws.row_dimensions[target].height = ws.row_dimensions[source].height


def add_overview(ws, job_id: str) -> None:
    title = "Q1V/SAME 半径探索（2026-07-19；已提交，非确认性）"
    if any(ws.cell(row, 1).value == title for row in range(1, ws.max_row + 1)):
        return
    section_row = ws.max_row + 1
    copy_row(ws, 35, section_row)
    ws.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=36)
    ws.cell(section_row, 1).value = title
    for offset, (_, label, radii) in enumerate(ARMS, start=1):
        row = section_row + offset
        copy_row(ws, 37, row)
        ws.cell(row, 1).value = label
        ws.cell(row, 2).value = "106/0/27（历史test27探索；已提交）"
        ws.cell(row, 3).value = "PointNet++ SA3"
        ws.cell(row, 4).value = f"5000→500→125→32；r={radii}；nsample=16；width=32"
        ws.cell(row, 5).value = "vertex random5000 / SAME；FPS center"
        ws.cell(row, 6).value = 5000
        for col in range(7, 37):
            ws.cell(row, col).value = None


def add_teacher_view(ws, job_id: str) -> None:
    title = "Ⅶ Q1V/SAME 半径探索（2026-07-19；已提交，非确认性）"
    if any(ws.cell(row, 1).value == title for row in range(1, ws.max_row + 1)):
        return
    section_row = ws.max_row + 1
    copy_row(ws, 45, section_row)
    ws.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=15)
    ws.cell(section_row, 1).value = title
    for offset, (_, label, _) in enumerate(ARMS, start=1):
        row = section_row + offset
        copy_row(ws, 47, row)
        ws.cell(row, 1).value = f"探索性test27（Job {job_id}）"
        ws.cell(row, 2).value = label
        ws.cell(row, 3).value = "PointNet++ SA3"
        ws.cell(row, 4).value = "vertex random5000 / SAME；FPS center"
        for col in range(5, 16):
            ws.cell(row, col).value = None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--submission", type=Path, default=SUBMISSION)
    args = parser.parse_args()
    submission = json.loads(args.submission.read_text(encoding="utf-8"))
    if submission.get("status") != "submitted" or len(submission.get("configs", [])) != 3:
        raise RuntimeError("refusing to write workbook without the exact submitted Q1V radius matrix")
    jobs = {item["role"]: item["job_id"] for item in submission["jobs"]}
    job_id = jobs.get("training_test27_exploratory_array")
    if not job_id:
        raise RuntimeError("submitted matrix lacks the training array job id")
    wb = load_workbook(args.workbook)
    add_overview(wb["实验矩阵总览"], job_id)
    add_teacher_view(wb["教师汇报视图"], job_id)
    wb.save(args.workbook)
    check = load_workbook(args.workbook, read_only=True, data_only=False)
    for sheet, title in (("实验矩阵总览", "Q1V/SAME 半径探索"), ("教师汇报视图", "Ⅶ Q1V/SAME 半径探索")):
        if not any(str(check[sheet].cell(row, 1).value).startswith(title) for row in range(1, check[sheet].max_row + 1)):
            raise RuntimeError(f"workbook verification failed for {sheet}")
    print(json.dumps({"status": "updated", "workbook": str(args.workbook), "job_id": job_id}, ensure_ascii=False))


if __name__ == "__main__":
    main()
