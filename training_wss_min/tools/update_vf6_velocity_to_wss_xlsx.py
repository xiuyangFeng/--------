"""Append / refresh the VELWSS2 rows in the 168-column WSS实验矩阵 sheet.

Rows owned by GROUP: three VF6 seeds with the legacy radius (VELWSS1 protocol), their three-seed mean,
and the three-seed means for legacy_exact and v5atlas. CFD-velocity oracle and uncalibrated-physics
diagnostics stay in the experiment README (same policy as VELWSS1). Historical cells are verified
untouched; the workbook is backed up before replacement.
"""
from __future__ import annotations

import argparse
import copy
import json
import os
import tempfile
import time
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.comments import Comment

from training_wss_min.tools.report_wss_recovery import METRIC_COLUMNS, cell_record, get_metric, metric_values
from training_wss_min.tools.vf6_velocity_to_wss import EXP, RADIUS_SOURCES, ROOT, RUNS, SEEDS, VELWSS1, read, sha
from training_wss_min.tools.report_vf6_velocity_to_wss import X5_REFS

BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
GROUP = "V5速度派生WSS·VELWSS2｜2026-09-16（VF6三seed→冻结Profile-Secant V3；6行）"
MISSING = "—"
PROTOCOL = "V5 train138/test34；峰值1162；V5有效全壁面1,231,295节点；VF6 best（train_loss选模）预测速度→冻结算子→病例内校准"
STRUCTURE = ("VF6（V07速度底座＋F6 Murray先验：PointNeXt-R＋LocalGeoPE＋L-SA2＋local branch＋多半径BT＋qad_lite，"
             "输出三分量速度）→ 冻结Profile-Secant V3（V1自适应梯度＋V3法向多尺度＋V4表面MLS＋depth3 profile＋Carreau–Yasuda＋84特征校准器）")
RADIUS_TEXT = {"legacy": "legacy半径（VELWSS1同口径：旧bundle半径按旧坐标最近邻映射，含unit_factor错位）",
               "legacy_exact": "legacy_exact半径（同旧半径场，旧坐标乘1000/unit_factor后逐节点精确对应，只修映射）",
               "v5atlas": "v5atlas半径（V5 bundle wall_local_radius，无映射）"}
PAIR_PATHS = ["field_casebalanced.r2", "normalized.field_casebalanced.r2", "field.mae",
              "hotspot.spearman_all_casemean", "hotspot.top10_iou_casemean"]


def mean_metrics(metric_dicts):
    """Column-wise mean of the workbook metric paths over seeds (None if any seed is missing)."""
    out = {}
    for column, path in METRIC_COLUMNS.items():
        values = [get_metric(m, path) for m in metric_dicts]
        out[column] = float(np.mean(values)) if all(v is not None for v in values) else None
    return out


def x5_metrics():
    out = {}
    for seed, (path, arm) in X5_REFS.items():
        source = read(path)["arms"][arm]["checkpoints"]["best"]["source"]
        out[seed] = (source, read(source)["test"])
    return out


def build_rows():
    provenance = read(EXP / "provenance.json")
    metrics = {}
    for radius in RADIUS_SOURCES:
        gate = read(EXP / "metrics" / f"{radius}_evaluation_gate.json")
        assert gate["passed"] and gate["n_cases"] == 34
        metrics[radius] = read(EXP / "metrics" / f"{radius}.json")
    reps = {seed: read(EXP / "velocity" / f"{seed}_reproduction.json") for seed in SEEDS}
    velwss1 = read(VELWSS1 / "metrics.json")["test"]
    x5 = x5_metrics()
    x5_mean = mean_metrics([m for _, m in x5.values()])
    results = read(EXP / "results.json") if (EXP / "results.json").is_file() else None
    rows = []

    def paired(values, refs, own_values):
        for j, (name, ref_values, ref_path) in enumerate(refs[:4]):
            base = 64 + j * 6
            values[base] = f"{name}；{ref_path}"
            for k, dotted in enumerate(PAIR_PATHS, 1):
                col = next(c for c, p in METRIC_COLUMNS.items() if p == dotted)
                own, theirs = own_values.get(col), ref_values.get(col)
                if own is not None and theirs is not None:
                    values[base + k] = own - theirs
        values[12] = own_values[6] - refs[0][1][6] if own_values.get(6) is not None else MISSING
        values[13] = "；".join(r[0] for r in refs) + "（ΔR²列相对第一个参照）"
        values[88] = "；".join(r[0] for r in refs) + "；best同checkpoint配对"

    velwss1_values = metric_values(velwss1)
    for radius in RADIUS_SOURCES:
        seed_values = {seed: metric_values(metrics[radius][f"{seed}_calibrated"]) for seed in SEEDS}
        seed_rows = SEEDS if radius == "legacy" else ()
        for seed in seed_rows:
            values = [MISSING] * 168
            m = metrics[radius][f"{seed}_calibrated"]
            values[:5] = [GROUP, f"VELWSS2_{radius}_{seed}｜VF6_{seed}预测速度→冻结Profile-Secant V3（{RADIUS_TEXT[radius].split('（')[0]}）｜已完成",
                          "WSS · Pa（速度派生）", "校准后主结果；CFD速度oracle/未校准物理核见实验README", PROTOCOL]
            for col, value in seed_values[seed].items():
                if value is not None:
                    values[col - 1] = value
            values[11] = f"{int(m['aggregate']['r2_negative_cases'])}/{int(m['aggregate']['n_cases'])}"
            values[139] = 1
            values[148] = reps[seed]["elapsed_seconds"]
            refs = [("VELWSS1（R5V速度→同算子）", velwss1_values, str(VELWSS1 / "metrics.json")),
                    (f"X5直接WSS同seed{seed[1:]}", metric_values(x5[seed][1]), x5[seed][0])]
            paired(values, refs, seed_values[seed])
            note = (f"速度来源 {RUNS[seed].relative_to(ROOT)}（best，epoch {reps[seed]['checkpoint_epoch']}，速度幅值R²_cb "
                    f"{reps[seed]['saved_metrics']['physical_r2_cb']:.4f}）；导出 {reps[seed]['execution']}；"
                    f"{RADIUS_TEXT[radius]}；冻结校准器历史用train138 WSS监督；不新增训练。")
            sources = [str(RUNS[seed]), str(EXP / "velocity" / f"{seed}_reproduction.json"),
                       str(EXP / "metrics" / f"{radius}.json"), str(EXP / "README.md")]
            values[159:] = ["Pa", STRUCTURE, f"VF6_{seed} 预测速度→冻结 Profile-Secant V3→WSS；{RADIUS_TEXT[radius]}",
                            f"seed{seed[1:]} / VF6 400ep / support5000 / 全内部query / 全壁面{provenance['n_wall']}节点", 5000,
                            f"vf6_velocity_to_wss_20260916/{radius}/VF6_{seed}", note[:32000], "\n".join(sources), GROUP]
            rows.append(dict(values=values, other={}))
        # three-seed mean row
        values = [MISSING] * 168
        mean_vals = mean_metrics([metrics[radius][f"{seed}_calibrated"] for seed in SEEDS])
        values[:5] = [GROUP, f"VELWSS2_{radius}_mean｜VF6三seed均值（1234/7/2025）预测速度→冻结Profile-Secant V3（{RADIUS_TEXT[radius].split('（')[0]}）｜已完成",
                      "WSS · Pa（速度派生）", "三seed逐列算术均值（校准后）；CFD速度oracle/未校准物理核见实验README", PROTOCOL]
        for col, value in mean_vals.items():
            if value is not None:
                values[col - 1] = value
        neg = [int(metrics[radius][f"{s}_calibrated"]["aggregate"]["r2_negative_cases"]) for s in SEEDS]
        values[11] = f"{'/'.join(map(str, neg))} /34（逐seed）"
        values[139] = 1
        refs = [("VELWSS1（R5V速度→同算子）", velwss1_values, str(VELWSS1 / "metrics.json")),
                ("X5直接WSS三seed均值", x5_mean, "；".join(p for p, _ in x5.values()))]
        paired(values, refs, mean_vals)
        sd = float(np.std([seed_values[s][6] for s in SEEDS], ddof=1))
        note = (f"三seed物理R²_cb：" + "、".join(f"{s} {seed_values[s][6]:.4f}" for s in SEEDS) + f"；sd {sd:.4f}。{RADIUS_TEXT[radius]}。"
                f"CFD真值速度→同算子 R²_cb {metrics[radius]['cfd_calibrated']['field_casebalanced']['r2']:.4f}（诊断，不入表）。")
        values[159:] = ["Pa", STRUCTURE, f"VF6 三seed均值 预测速度→冻结 Profile-Secant V3→WSS；{RADIUS_TEXT[radius]}",
                        f"seed1234/7/2025 均值 / VF6 400ep / support5000 / 全壁面{provenance['n_wall']}节点", 5000,
                        f"vf6_velocity_to_wss_20260916/{radius}/mean3", note[:32000],
                        "\n".join([str(EXP / "metrics" / f"{radius}.json"), str(EXP / "results.json"), str(EXP / "README.md")]), GROUP]
        rows.append(dict(values=values, other={s: seed_values[s] for s in SEEDS}))
    return rows


def update_workbook(book: Path, dry_run: bool = False):
    rows = build_rows()
    before = sha(book)
    wb = load_workbook(book)
    if wb.sheetnames[:3] != ["WSS实验矩阵", "速度与压力实验矩阵", "指标说明"]:
        raise ValueError("unexpected workbook sheet schema")
    ws = wb["WSS实验矩阵"]
    if ws.max_column != 168 or ws["FI5"].value != "原始实验 ID" or "WSSResults" not in ws.tables:
        raise ValueError("unexpected workbook 168-column schema")
    owned = [r for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value == GROUP]
    identities = [r["values"][164] for r in rows]
    if owned and len(owned) != len(rows):
        raise ValueError(f"owned rows must number exactly {len(rows)}")
    locations = ({ws.cell(r, 165).value: r for r in owned} if owned else
                 {ident: ws.max_row + i + 1 for i, ident in enumerate(identities)})
    if set(locations) != set(identities):
        raise ValueError("row identities differ from the workbook's owned rows")
    preserved = {(s.title, c.coordinate): cell_record(c) for s in wb for row in s for c in row
                 if not (s.title == ws.title and c.row in owned)}
    merges = {s.title: tuple(map(str, s.merged_cells.ranges)) for s in wb}
    for i, record in enumerate(rows):
        row = locations[record["values"][164]]
        for col, value in enumerate(record["values"], 1):
            cell = ws.cell(row, col)
            cell._style = copy.copy(ws.cell(310 if i % 2 else 311, col)._style)
            cell.value = value
            cell.comment = None
            cell.hyperlink = None
            if col in METRIC_COLUMNS and record["other"]:
                text = "；".join(f"{s}={vals[col] if vals[col] is not None else MISSING}" for s, vals in record["other"].items())
                cell.comment = Comment(f"本单元格=三seed均值；逐seed {text}\n路径={METRIC_COLUMNS[col]}", "VELWSS2 report")
        ws.cell(row, 166).comment = Comment(record["values"][165][:32000], "VELWSS2 report")
        ws.cell(row, 167).comment = Comment(record["values"][166], "VELWSS2 report")
        ws.row_dimensions[row].height = 72
    ws.tables["WSSResults"].ref = f"A5:FL{ws.max_row}"
    ws.tables["WSSResults"].autoFilter.ref = f"A5:FL{ws.max_row}"
    ws.print_area = f"A1:N{ws.max_row}"
    if dry_run:
        print(json.dumps({"rows": locations, "dry_run": True}, ensure_ascii=False))
        return
    fd, temp = tempfile.mkstemp(dir=book.parent, prefix=".vf6_velocity_to_wss_", suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(temp)
        checked = load_workbook(temp)
        for (sheet, coordinate), record in preserved.items():
            if cell_record(checked[sheet][coordinate]) != record:
                raise AssertionError(f"historical cell changed: {sheet}!{coordinate}")
        assert {s.title: tuple(map(str, s.merged_cells.ranges)) for s in checked} == merges
        for record in rows:
            r = locations[record["values"][164]]
            for col, path in METRIC_COLUMNS.items():
                expected = record["values"][col - 1]
                actual = checked[ws.title].cell(r, col).value
                if isinstance(expected, float):
                    assert isinstance(actual, float) and abs(actual - expected) <= 1e-12, (r, col)
        checked.close()
        if sha(book) != before:
            raise RuntimeError("workbook changed concurrently; refusing replacement")
        backup = EXP / f"{book.stem}_backup_before_vf6_velocity_to_wss_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        backup.write_bytes(book.read_bytes())
        os.replace(temp, book)
    finally:
        if Path(temp).exists():
            Path(temp).unlink()
    evidence = {"passed": True, "before_sha256": before, "after_sha256": sha(book), "backup": str(backup),
                "rows": locations, "records": len(rows), "historical_cells_verified": len(preserved),
                "table_ref": ws.tables["WSSResults"].ref, "policy": "calibrated prediction rows only; oracle/physics in README"}
    (EXP / "xlsx_acceptance.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
    print(json.dumps(evidence, ensure_ascii=False, indent=2))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=BOOK)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    update_workbook(args.workbook, dry_run=args.dry_run)


if __name__ == "__main__":
    main()
