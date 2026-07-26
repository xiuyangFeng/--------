#!/usr/bin/env python3
"""Append the submitted xyz-only controls to the PointNet result workbook."""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
MANIFEST = ROOT / "training_wss_min/preflight/xyz_input_ablation_20260724_prepared.json"
SUBMISSION = ROOT / "training_wss_min/preflight/xyz_input_ablation_20260724_submission.json"
SHEET = "D2K64纯XYZ待运行"


def main() -> None:
    matrix = json.loads(MANIFEST.read_text(encoding="utf-8"))
    submission = json.loads(SUBMISSION.read_text(encoding="utf-8"))
    if matrix.get("status") != "prepared" or len(matrix.get("configs", [])) != 3:
        raise RuntimeError("missing exact three-arm xyz-only manifest")
    if submission.get("status") != "submitted":
        raise RuntimeError("xyz-only matrix was not formally submitted")
    with tempfile.TemporaryDirectory(prefix="d2xyz_xlsx_") as tmp:
        candidate = Path(tmp) / BOOK.name
        shutil.copy2(BOOK, candidate)
        workbook = load_workbook(candidate)
        if SHEET in workbook.sheetnames:
            raise RuntimeError(f"refusing duplicate sheet: {SHEET}")
        sheet = workbook.create_sheet(SHEET)
        sheet.merge_cells("A1:H1")
        sheet["A1"] = "D2-K64 三锚点纯 XYZ 输入对照（已提交；结果待回填）"
        sheet["A1"].font = Font(bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        sheet.append([
            "新臂", "锚定配置", "唯一变化", "split", "PointNeXt-R blocks",
            "seed / epochs", "作业链", "状态",
        ])
        for cell in sheet[2]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        labels = {
            "d2_c125_k64_xyz": "D2 c125×k64",
            "m1_d2k64_mixed138_test36_xyz": "M1/S0 D2-K64 mixed138/test36",
            "s2_d2k64_pnxr_mixed_xyz": "S2 D2-K64 + PointNeXt-R",
        }
        for row in matrix["configs"]:
            config = json.loads(Path(row["config"]).read_text(encoding="utf-8"))
            model = config["model"]
            sheet.append([
                row["experiment_id"], labels[row["experiment_id"]],
                "input_features: xyz+geom → xyz（其余字段冻结）",
                Path(config["data"]["split_path"]).name.replace("split_", "").replace(".json", ""),
                "/".join(map(str, model["sa_blocks"])),
                f"{config['train']['seed']} / {config['train']['epochs']}",
                "10903 → 10904_[0-2]%3",
                "GPU 门禁排队；afterok 后训练+best/last test 评估",
            ])
        sheet.append([])
        sheet.merge_cells("A7:H7")
        sheet["A7"] = (
            "静态逐字段审计 3/3 通过；仅允许 name、notes、data.input_features 变更。"
        )
        sheet["A7"].font = Font(italic=True)
        for column, width in enumerate((36, 38, 44, 46, 22, 16, 26, 46), 1):
            sheet.column_dimensions[chr(64 + column)].width = width
        workbook.save(candidate)
        workbook.close()
        check = load_workbook(candidate, read_only=True, data_only=False)
        if SHEET not in check.sheetnames or check[SHEET].max_row != 7:
            raise RuntimeError("xlsx readback failed")
        check.close()
        render = Path(tmp) / "render"
        render.mkdir()
        subprocess.run(
            ["/usr/bin/libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(render), str(candidate)],
            check=True, capture_output=True, text=True,
        )
        if not (render / candidate.with_suffix(".pdf").name).is_file():
            raise RuntimeError("LibreOffice render check failed")
        shutil.copy2(candidate, BOOK)
    print(json.dumps({"status": "updated_and_readback_rendered", "workbook": str(BOOK), "sheet": SHEET}, ensure_ascii=False))


if __name__ == "__main__":
    main()
