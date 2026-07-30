#!/usr/bin/env python3
"""Update existing workbook views with the completed O0 hotspot/tail screen."""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from training_wss_min.tools.update_s3_rootcause_xlsx import copy_row_style


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = (
    ROOT
    / "training_wss_min/preflight/"
    "o0_hotspot_tail_matrix_20260728_results_analysis.json"
)
SECTION_TITLE = "2026-07-28｜O0 hotspot BCE / q90 pinball 单变量筛选"
ORDER = (
    "h1_hotspot_bce_q90_lam020_s1234",
    "h2_pinball_q90_lam020_s1234",
)
LABELS = {
    "h1_hotspot_bce_q90_lam020_s1234": "O0-H1 hotspot BCE q90 λ0.20 s1234",
    "h2_pinball_q90_lam020_s1234": "O0-H2 pinball q90 λ0.20 s1234",
}
CHANGES = {
    "h1_hotspot_bce_q90_lam020_s1234": (
        "第二输出通道 + 病例相对 q90 balanced BCE（λ=0.20）"
    ),
    "h2_pinball_q90_lam020_s1234": "MSE + q90 pinball（λ=0.20）",
}
EXPECTED_SHEETS = [
    "汇总对比",
    "教师汇报视图",
    "RCR Oracle",
    "指标说明",
    "实验矩阵总览",
]

NAVY = "1F4E78"
BLUE = "D9EAF7"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def raw_metrics(record: dict) -> dict:
    return read_json(
        Path(record["run_dir"]) / "eval/ckpt_best/metrics.json"
    )["test"]


def overview_values(record: dict) -> list:
    raw = raw_metrics(record)
    norm = raw["normalized"]
    experiment_id = record["experiment_id"]
    return [
        LABELS[experiment_id],
        "138/0/36（AG/AAA/ILO mixed；seed=1234）",
        "PointNeXt-R + LocalGeoPE",
        CHANGES[experiment_id],
        "random / SAME；FPS center",
        5000,
        raw["field_casebalanced"]["r2"],
        raw["field"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["aggregate"]["r2_casemed"],
        raw["aggregate"]["r2_casep10"],
        f"{raw['aggregate']['r2_negative_cases']}/{raw['aggregate']['n_cases']}",
        raw["field"]["rmse"],
        raw["field"]["mae"],
        raw["field"]["nrmse_range"],
        raw["aggregate"]["nrmse_casemean"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        raw["calibration"]["top10_pred_true_ratio"],
        raw["hotspot"]["top10_iou_casemean"],
        norm["field_casebalanced"]["r2"],
        norm["field"]["r2"],
        norm["aggregate"]["r2_casemean"],
        norm["aggregate"]["r2_casemed"],
        norm["aggregate"]["r2_casep10"],
        norm["field_casebalanced"]["mae"],
        norm["field_casebalanced"]["rmse"],
        norm["field"]["nrmse_range"],
        norm["aggregate"]["nrmse_casemean"],
        norm["field"].get("nmae_range"),
        norm["aggregate"].get("nmae_casemean"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["spearman_high_wss_casemean"],
        norm["calibration"]["top10_pred_true_ratio"],
        norm["calibration"]["p99_pred_true_ratio"],
        experiment_id,
    ]


def teacher_values(record: dict) -> list:
    raw = raw_metrics(record)
    norm = raw["normalized"]
    experiment_id = record["experiment_id"]
    return [
        "O0 hotspot/tail",
        LABELS[experiment_id],
        "PointNeXt-R + LocalGeoPE",
        "random5000 / SAME",
        raw["field_casebalanced"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["field"]["rmse"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        norm["field_casebalanced"]["r2"],
        norm["field"].get("nmae_range"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["top10_iou_casemean"],
        norm["calibration"]["p99_pred_true_ratio"],
    ]


def add_title(sheet, row: int) -> int:
    sheet.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=15
    )
    cell = sheet.cell(row, 1, SECTION_TITLE)
    cell.font = Font(bold=True, color=WHITE, size=12)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 23
    return row + 1


def add_header(sheet, row: int) -> int:
    headers = [
        "臂",
        "唯一变化",
        "R²_cb",
        "ΔR²_cb",
        "normalized R²_cb",
        "Δnormalized",
        "MAE_cb (Pa)",
        "ΔMAE",
        "high-WSS nRMSE",
        "ΔnRMSE",
        "top10 IoU",
        "ΔIoU",
        "主 Gate",
        "共同保护线",
        "判定",
    ]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    sheet.row_dimensions[row].height = 31
    return row + 1


def add_summary_row(sheet, row: int, record: dict, comparison: dict, decision: dict) -> None:
    delta = comparison["aggregate_treatment_minus_control"]
    extra = comparison["extra_treatment_minus_control"]
    experiment_id = record["experiment_id"]
    primary = decision["primary"]
    values = [
        LABELS[experiment_id],
        CHANGES[experiment_id],
        record["best"]["physical_r2_casebalanced"],
        delta["physical_r2_casebalanced"],
        record["best"]["normalized_r2_casebalanced"],
        delta["normalized_r2_casebalanced"],
        record["best"]["physical_mae"],
        delta["physical_mae"],
        record["extra"]["high_wss_nrmse_range"],
        extra["high_wss_nrmse_range"],
        record["best"]["physical_top10_iou"],
        delta["physical_top10_iou"],
        (
            f"ΔIoU≥+0.020；实测{primary['delta']:+.4f}"
            if experiment_id.startswith("h1_")
            else f"ΔnRMSE≤-0.002；实测{primary['delta']:+.5f}"
        ),
        "normalized R² / MAE 均通过",
        "No-Go，不组合",
    ]
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.fill = PatternFill("solid", fgColor=YELLOW)
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if column in range(3, 13):
            cell.number_format = (
                "+0.0000;-0.0000;0.0000"
                if column in (4, 6, 8, 10, 12)
                else "0.0000"
            )
    sheet.row_dimensions[row].height = 33


def render_and_count_pages(workbook_path: Path, out_dir: Path) -> int:
    subprocess.run(
        [
            "/usr/bin/libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(workbook_path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    pdf = out_dir / workbook_path.with_suffix(".pdf").name
    output = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    for line in output.splitlines():
        if line.startswith("Pages:"):
            return int(line.split(":", 1)[1].strip())
    raise RuntimeError("pdfinfo did not report page count")


def main() -> None:
    analysis = read_json(ANALYSIS)
    records = analysis["records"]
    comparisons = {
        item["treatment"]: item for item in analysis["paired_comparisons"]
    }
    decisions = analysis["registered_gate_decisions"]
    for experiment_id in ORDER:
        if records[experiment_id]["integrity"] != "passed":
            raise RuntimeError(f"integrity failed: {experiment_id}")
        if decisions[experiment_id]["decision"] != "no_go":
            raise RuntimeError(f"unexpected Gate decision: {experiment_id}")

    with tempfile.TemporaryDirectory(prefix="wss_o0_hotspot_xlsx_") as temp_dir:
        temp = Path(temp_dir)
        candidate = temp / BOOK.name
        shutil.copy2(BOOK, candidate)
        workbook = load_workbook(candidate, data_only=False)
        if workbook.sheetnames != EXPECTED_SHEETS:
            raise RuntimeError(
                f"unexpected sheet set/order: {workbook.sheetnames}"
            )

        overall = workbook["实验矩阵总览"]
        old_rows = [
            cell.row for cell in overall["AK"] if cell.value in set(ORDER)
        ]
        for row in sorted(old_rows, reverse=True):
            overall.delete_rows(row, 1)
        first_overall = overall.max_row + 1
        for experiment_id in ORDER:
            row = overall.max_row + 1
            copy_row_style(overall, row - 1, row, 37)
            for column, value in enumerate(
                overview_values(records[experiment_id]), 1
            ):
                overall.cell(row, column, value)

        teacher = workbook["教师汇报视图"]
        old_rows = [
            cell.row
            for cell in teacher["B"]
            if cell.value in set(LABELS.values())
        ]
        for row in sorted(old_rows, reverse=True):
            teacher.delete_rows(row, 1)
        first_teacher = teacher.max_row + 1
        for experiment_id in ORDER:
            row = teacher.max_row + 1
            copy_row_style(teacher, row - 1, row, 15)
            for column, value in enumerate(
                teacher_values(records[experiment_id]), 1
            ):
                teacher.cell(row, column, value)

        summary = workbook["汇总对比"]
        old_title_rows = [
            cell.row for cell in summary["A"] if cell.value == SECTION_TITLE
        ]
        if len(old_title_rows) > 1:
            raise RuntimeError("duplicate O0 hotspot/tail summary sections")
        if old_title_rows:
            summary.delete_rows(old_title_rows[0], 4)
        summary_start = summary.max_row + 1
        row = add_title(summary, summary_start)
        row = add_header(summary, row)
        for experiment_id in ORDER:
            add_summary_row(
                summary,
                row,
                records[experiment_id],
                comparisons[experiment_id],
                decisions[experiment_id],
            )
            row += 1

        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(candidate)
        workbook.close()

        check = load_workbook(candidate, read_only=True, data_only=False)
        if check.sheetnames != EXPECTED_SHEETS:
            raise RuntimeError("readback sheet set changed")
        saved_ids = {
            values[0]
            for values in check["实验矩阵总览"].iter_rows(
                min_col=37, max_col=37, values_only=True
            )
        }
        saved_labels = {
            values[0]
            for values in check["教师汇报视图"].iter_rows(
                min_col=2, max_col=2, values_only=True
            )
        }
        saved_titles = {
            values[0]
            for values in check["汇总对比"].iter_rows(
                min_col=1, max_col=1, values_only=True
            )
        }
        if not set(ORDER).issubset(saved_ids):
            raise RuntimeError("overview readback missing O0 hotspot/tail rows")
        if not set(LABELS.values()).issubset(saved_labels):
            raise RuntimeError("teacher readback missing O0 hotspot/tail rows")
        if SECTION_TITLE not in saved_titles:
            raise RuntimeError("summary readback missing O0 hotspot/tail section")
        if check["实验矩阵总览"].max_row != 186:
            raise RuntimeError("overview row count is not idempotent")
        if check["教师汇报视图"].max_row != 160:
            raise RuntimeError("teacher row count is not idempotent")
        check.close()

        render_dir = temp / "render"
        render_dir.mkdir()
        pages = render_and_count_pages(candidate, render_dir)
        if pages != 6:
            raise RuntimeError(
                f"workbook render grew from the compact 6-page layout to {pages}"
            )
        shutil.copy2(candidate, BOOK)

    print(
        json.dumps(
            {
                "status": "updated_and_rendered",
                "workbook": str(BOOK),
                "sheet_count": len(EXPECTED_SHEETS),
                "pdf_pages": pages,
                "overall_rows": [first_overall, first_overall + 1],
                "teacher_rows": [first_teacher, first_teacher + 1],
                "summary_rows": [summary_start, summary_start + 3],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
