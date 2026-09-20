"""Append/update the authorized A5 follow-up in the existing three workbook views.

Pending rows contain no invented metrics. Only this script's marked trailing
sections are replaced; all earlier sheets, formulas and sections are retained.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .report_v6_followup import EXP, MATRIX, ANCHOR, load_metrics, report
from .update_v5_rerun_xlsx import BOOK, EXPECTED_SHEETS, LSA2, NAVY, BLUE, WHITE
from .update_v6_matrix_xlsx import overview_values

PREFIX = "V6-A5后续20260909"
LEGACY_TITLE = PREFIX + "｜D1–D7 / L1–L3 / M1–M5；仅WSS；各seed1234；同seed A5参照"
TITLE = LEGACY_TITLE + "；本节checkpoint=best（训练总损失选模）"
TEACHER_TITLE = "ⅩⅣ " + TITLE


def drop_own_tail(ws, title, arms):
    # Accept only the exact prior marker when migrating to an explicit best label.
    titles = {title}
    if title == TITLE:
        titles.add(LEGACY_TITLE)
    elif title == TEACHER_TITLE:
        titles.add("ⅩⅣ " + LEGACY_TITLE)
    rows = [cell.row for cell in ws["A"] if cell.value in titles]
    if not rows:
        return
    if len(rows) != 1:
        raise ValueError(f"Duplicate section marker in {ws.title}")
    first = rows[0]
    # Own rows have exact known identities. A title-prefix heuristic can silently
    # delete a colleague's section, or even unlabeled data appended afterwards.
    if ws.title == "实验矩阵总览":
        expected = [{1: f"{PREFIX}·{a['id']} {Path(a['config']).stem}", 37: a["run_name"]} for a in arms]
        width = 49
    elif ws.title == "教师汇报视图":
        expected = [{1: "V6-A5后续", 2: f"{a['id']} {a['hypothesis']}｜"} for a in arms]
        width = 19
    elif ws.title == "汇总对比":
        expected = [{1: "ID", 2: "父臂"}, {1: "B0=A5", 15: ANCHOR}]
        expected += [{1: a["id"], 15: a["run_name"]} for a in arms]
        expected += [{1: "说明"}]
        width = 15
    else:
        raise ValueError(f"No ownership schema for sheet {ws.title}")
    for offset, identity in enumerate(expected, 1):
        row = first + offset
        for column, value in identity.items():
            actual = ws.cell(row, column).value
            matches = (isinstance(actual, str) and actual.startswith(value)) if (
                ws.title == "教师汇报视图" and column == 2
            ) else actual == value
            if not matches:
                raise RuntimeError(f"Unknown/missing section row {ws.title}!{row}; refusing to delete")
    last = first + len(expected)
    for row in ws.iter_rows(min_row=first):
        if any(cell.value is not None and (cell.row > last or cell.column > width) for cell in row):
            raise RuntimeError(f"Unowned data after/beside our section in {ws.title}; refusing to delete")
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= first:
            ws.unmerge_cells(str(merged))
    ws.delete_rows(first, ws.max_row - first + 1)


def append_title(ws, title, width):
    row = ws.max_row + 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = ws.cell(row, 1, title)
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 40


def append_values(ws, values, header=False):
    row = ws.max_row + 1
    for col, value in enumerate(values, 1):
        if value is None:
            continue
        cell = ws.cell(row, col, value)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if header:
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(bold=True)
        elif isinstance(value, float):
            cell.number_format = "0.0000"
    ws.row_dimensions[row].height = 48 if header else 60
    return row


def teacher_metric_values(entry, m):
    """Columns E:S follow the existing teacher header, including log_z p99 in O."""
    get = entry.get
    return [get("physical_r2_cb"), get("case_mean"), get("rmse"),
            m["field"].get("nmae_range") if m else None,
            m["aggregate"].get("nmae_casemean") if m else None,
            get("high_wss_r2"), get("normalized_r2_cb"),
            m["normalized"]["field"].get("nmae_range") if m else None,
            get("spearman"), get("top10_iou"),
            m["normalized"]["calibration"].get("p99_pred_true_ratio") if m else None,
            m["field_casebalanced"].get("r2_linear_fit") if m else None,
            m["field"].get("r2_linear_fit") if m else None,
            m["field"].get("linear_fit_slope") if m else None,
            m["field"].get("linear_fit_intercept") if m else None]


def followup_overview_values(arm, m):
    """Honor this workbook's AA/AB pooled headers without rewriting old sections."""
    values = overview_values(arm, m)
    values[26] = m["normalized"]["field"]["mae"]
    values[27] = m["normalized"]["field"]["rmse"]
    return values


def main():
    _, summary = report("best")
    matrix = json.loads(MATRIX.read_text())
    arms = matrix["arms"]
    original = BOOK.read_bytes()
    original_hash = hashlib.sha256(original).hexdigest()
    backup = EXP / "WSS_PointNet实验矩阵与结果汇总last_backup_before_followup_20260909.xlsx"
    if not backup.exists():
        backup.write_bytes(original)
    with tempfile.TemporaryDirectory(prefix="v6_followup_workbook_") as td:
        candidate = Path(td) / BOOK.name
        candidate.write_bytes(original)
        wb = load_workbook(candidate)
        if wb.sheetnames != EXPECTED_SHEETS:
            raise ValueError(f"Unexpected sheets: {wb.sheetnames}")
        for name, title, width in [("实验矩阵总览", TITLE, 20), ("教师汇报视图", TEACHER_TITLE, 19), ("汇总对比", TITLE, 15)]:
            drop_own_tail(wb[name], title, arms)
            append_title(wb[name], title, width)
        ov, te, su = wb["实验矩阵总览"], wb["教师汇报视图"], wb["汇总对比"]
        append_values(su, ["ID", "父臂", "状态", "物理R²_cb", "Δ vs A5", "Δ vs父臂", "log_z R²_cb", "case mean", "MAE Pa", "high-WSS R²", "top10比", "p99比", "IoU", "变化", "结果目录"], header=True)
        base = summary["reference"]
        append_values(su, ["B0=A5", "—", "已有best；seed1234", base["physical_r2_cb"], None, None, base["normalized_r2_cb"], base["case_mean"], base["mae"], base["high_wss_r2"], base["top10_ratio"], base["p99_ratio"], base["top10_iou"], "原始25D与原损失", ANCHOR])
        rows = {}
        for arm in arms:
            cfg = json.loads((MATRIX.parent / arm["config"]).read_text())
            entry = summary["arms"][arm["id"]]
            m = load_metrics(arm["run_name"], "best") if entry["evidence"]["valid_scientific_result"] else None
            label = f"{PREFIX}·{arm['id']} {Path(arm['config']).stem}"
            protocol = f"random5000 / {cfg['data']['query_mode'].upper()}；seed1234；checkpoint=best"
            model = LSA2 + (" + pointwise容量对照" if arm["id"] == "M1" else " + local wall branch")
            if m:
                vals = followup_overview_values(arm, m)
                vals[0], vals[2], vals[4], vals[36] = label, model, protocol, arm["run_name"]
            else:
                vals = [None] * 49
                vals[:6] = [label, "V5 train138/test34；峰值单帧；seed1234", model,
                            arm["hypothesis"] + "｜" + entry["evidence"]["status"], protocol, 5000]
                vals[36] = arm["run_name"]
            rows[arm["id"]] = append_values(ov, vals)
            get = lambda key: entry.get(key)
            append_values(te, ["V6-A5后续", f"{arm['id']} {arm['hypothesis']}｜{entry['evidence']['status']}", model, protocol]
                          + teacher_metric_values(entry, m))
            append_values(su, [arm["id"], entry["parent"], entry["evidence"]["status"], get("physical_r2_cb"), get("delta_physical"),
                get("delta_parent_physical"), get("normalized_r2_cb"), get("case_mean"), get("mae"), get("high_wss_r2"),
                get("top10_ratio"), get("p99_ratio"), get("top10_iou"), arm["hypothesis"], arm["run_name"]])
        append_values(su, ["说明", "本节所有指标为best；文件名last为历史命名，不表示本节checkpoint。best由各臂自身训练总损失选模，L臂须同时参考matrix_tables_last.md。全部一个seed；仅描述变化，不作显著性/稳定性结论；P新几何等待用户确认。",
                           "整批完结归档已核验" if summary["matrix_finalized"] else "整批完结归档尚未核验",
                           "R²_cb病例等权；总览物理M/N列和归一化AA/AB列的RMSE/MAE按各自表头均为顶点pool；high-WSS为各病例真值q90选区后pool；top10比为全局pool真值q90选区；p99为全局预测/真值各自分位数比；IoU与Spearman为病例均值。教师视图O列沿用log_z p99，汇总对比L列为Pa p99。"])
        wb.save(candidate)
        check = load_workbook(candidate, read_only=True, data_only=False)
        assert check.sheetnames == EXPECTED_SHEETS
        for arm in arms:
            assert check["实验矩阵总览"].cell(rows[arm["id"]], 37).value == arm["run_name"]
            expected = summary["arms"][arm["id"]].get("physical_r2_cb")
            actual = check["实验矩阵总览"].cell(rows[arm["id"]], 7).value
            assert actual is None if expected is None else abs(actual - expected) < 1e-12
        check.close()
        if hashlib.sha256(BOOK.read_bytes()).hexdigest() != original_hash:
            raise RuntimeError("Workbook changed during generation; refusing to overwrite concurrent edits")
        # Atomic replacement on the same filesystem.
        staged = BOOK.with_suffix(".followup.tmp.xlsx")
        shutil.copy2(candidate, staged)
        staged.replace(BOOK)
    (EXP / "workbook_update.json").write_text(json.dumps({"book": str(BOOK), "backup": str(backup),
        "checkpoint": "best", "matrix_finalized": summary["matrix_finalized"],
        "completed_best": summary["completed_count"], "overview_rows": rows, "reference": ANCHOR}, indent=2, ensure_ascii=False) + "\n")
    print(f"Updated workbook: {summary['completed_count']}/15 evaluated; rows={rows}")


if __name__ == "__main__":
    main()
