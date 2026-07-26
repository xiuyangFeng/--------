#!/usr/bin/env python3
"""Append the completed Q1V/SAME 0.6x-radius follow-up to the workbook."""

from __future__ import annotations

import argparse
import copy
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


REPO = Path(__file__).resolve().parents[2]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from training_wss_min.tools.update_q1v_radius_test27_results_xlsx import METRIC_KEYS


WORKBOOK = REPO / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = REPO / "training_wss_min/preflight/q1v_radius_r60_same_results_analysis.json"
TITLE_OVERVIEW = "Q1V/SAME 0.6× 半径补充（2026-07-20；完成，非确认性）"
TITLE_TEACHER = "Ⅸ Q1V/SAME 0.6× 半径补充（2026-07-20；完成，非确认性）"
LABEL = "Q1V半径探索(0.6×补充)"


def copy_row(ws, source: int, target: int) -> None:
    for col in range(1, ws.max_column + 1):
        src, dst = ws.cell(source, col), ws.cell(target, col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        dst.value = None
    ws.row_dimensions[target].height = ws.row_dimensions[source].height


def find_row(ws, prefix: str) -> int | None:
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, str) and value.startswith(prefix):
            return row
    return None


def update_overview(ws, record: dict) -> None:
    section = find_row(ws, TITLE_OVERVIEW)
    if section is None:
        section = ws.max_row + 1
        template = find_row(ws, "Q1V/SAME 半径探索") or 1
        copy_row(ws, template, section)
        ws.merge_cells(start_row=section, start_column=1, end_row=section, end_column=36)
        ws.cell(section, 1).value = TITLE_OVERVIEW
        row = section + 1
        copy_row(ws, find_row(ws, "Q1V半径探索(0.8×)") or 1, row)
    else:
        row = next((i for i in range(section + 1, ws.max_row + 1) if ws.cell(i, 1).value == LABEL), None)
        if row is None:
            row = section + 1
            copy_row(ws, find_row(ws, "Q1V半径探索(0.8×)") or 1, row)
    metric = record["best"]
    values = (LABEL, "106/0/27（历史test27探索）", "PointNet++ SA3",
              "5000→500→125→32；r=0.03/0.06/0.12；nsample=16；width=32",
              "vertex random5000 / SAME；FPS center", 5000)
    for col, value in enumerate(values, start=1):
        ws.cell(row, col).value = value
    for col, key in enumerate(METRIC_KEYS, start=7):
        ws.cell(row, col).value = metric.get(key)


def update_teacher(ws, record: dict) -> None:
    section = find_row(ws, TITLE_TEACHER)
    if section is None:
        section = ws.max_row + 1
        template = find_row(ws, "Ⅶ Q1V/SAME 半径探索") or 1
        copy_row(ws, template, section)
        ws.merge_cells(start_row=section, start_column=1, end_row=section, end_column=15)
        ws.cell(section, 1).value = TITLE_TEACHER
        row = section + 1
        copy_row(ws, find_row(ws, "探索性test27（10546_0") or 1, row)
    else:
        row = next((i for i in range(section + 1, ws.max_row + 1) if ws.cell(i, 2).value == LABEL), None)
        if row is None:
            row = section + 1
            copy_row(ws, find_row(ws, "探索性test27（10546_0") or 1, row)
    metric = record["best"]
    values = ("探索性test27（10556；已完成）", LABEL, "PointNet++ SA3",
              "vertex random5000 / SAME；FPS center", metric["physical_r2_casebalanced"],
              metric["physical_r2_case_mean"], metric["physical_rmse"], metric["physical_nmae_pooled"],
              metric["physical_nmae_case_mean"], metric["high_wss_r2"],
              metric["normalized_r2_casebalanced"], metric["normalized_nmae_pooled"],
              metric["spearman_case_mean"], metric["physical_top10_iou"], metric["physical_p99_amplitude_ratio"])
    for col, value in enumerate(values, start=1):
        ws.cell(row, col).value = value


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--analysis", type=Path, default=ANALYSIS)
    args = parser.parse_args()
    analysis = json.loads(args.analysis.read_text(encoding="utf-8"))
    record = analysis["record"]
    if analysis.get("test27_policy") != "user-authorized historical anchor; not independent confirmation" or record.get("integrity") != "passed":
        raise RuntimeError("refusing workbook update without integrity-passed exploratory result")
    wb = load_workbook(args.workbook)
    update_overview(wb["实验矩阵总览"], record)
    update_teacher(wb["教师汇报视图"], record)
    wb.save(args.workbook)
    check = load_workbook(args.workbook, read_only=True, data_only=False)
    for sheet, column in (("实验矩阵总览", 1), ("教师汇报视图", 2)):
        if not any(check[sheet].cell(row, column).value == LABEL for row in range(1, check[sheet].max_row + 1)):
            raise RuntimeError(f"workbook verification failed for {sheet}")
    print(json.dumps({"status": "updated", "workbook": str(args.workbook)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
