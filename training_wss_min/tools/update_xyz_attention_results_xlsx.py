#!/usr/bin/env python3
"""Append xyz and attention results only to the workbook's two overview sheets."""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from training_wss_min.tools.update_s3_regularization_xlsx import matrix_values
from training_wss_min.tools.update_s3_rootcause_xlsx import copy_row_style


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
XYZ = ROOT / "training_wss_min/preflight/xyz_input_ablation_20260724_results_analysis.json"
ATTN = ROOT / "training_wss_min/preflight/s3_sa3_coarse_attention_single_results_analysis.json"
LEGACY_SHEETS = ("D2K64纯XYZ待运行", "S3注意力单种子")
XYZ_IDS = (
    "d2_c125_k64_xyz",
    "m1_d2k64_mixed138_test36_xyz",
    "s2_d2k64_pnxr_mixed_xyz",
)
ATTN_ID = "sa3_coarse_attention_s1234"
XYZ_TITLE = "2026-07-25｜D2-K64 三锚点纯 XYZ 输入对照（vs 各自 xyz+geom 父模型）"
ATTN_TITLE = "2026-07-25｜S3-GEOPE + SA3 coarse attention（seed1234，单种子）"


def overview_values(record: dict, *, label: str, split: str, backbone: str, change: str) -> list:
    """Use the exact 37-column extractor already used by this workbook page."""
    # The shared extractor only uses variant/seed to make its display label;
    # replace that display value below while retaining its metric mapping.
    values = matrix_values({**record, "variant": "head_dropout01", "seed": 1234})
    values[0:4] = [label, split, backbone, change]
    return values


def summary_title(sheet, row: int, text: str) -> int:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    cell = sheet.cell(row, 1, text)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    return row + 1


def write_header(sheet, row: int, values: list[str]) -> int:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    return row + 1


def main() -> None:
    xyz = json.loads(XYZ.read_text(encoding="utf-8"))
    attention = json.loads(ATTN.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory(prefix="wss_xyz_attention_overview_") as temp_dir:
        candidate = Path(temp_dir) / BOOK.name
        shutil.copy2(BOOK, candidate)
        workbook = load_workbook(candidate, data_only=False)
        for name in LEGACY_SHEETS:
            if name in workbook.sheetnames:
                del workbook[name]

        overall = workbook["实验矩阵总览"]
        ids = {*XYZ_IDS, ATTN_ID}
        row_ids = {cell.value: cell.row for cell in overall["AK"] if cell.value in ids}
        for row in sorted(row_ids.values(), reverse=True):
            overall.delete_rows(row, 1)
        xyz_labels = {
            "d2_c125_k64_xyz": ("D2 c125×k64 · xyz", "106/0/27（AG/AAA；seed=1234）", "PointNet++", "xyz-only；vs D2 xyz+geom"),
            "m1_d2k64_mixed138_test36_xyz": ("M1/S0 D2-K64 · xyz", "138/0/36（AG/AAA/ILO mixed；seed=1234）", "PointNet++", "xyz-only；vs M1 xyz+geom"),
            "s2_d2k64_pnxr_mixed_xyz": ("S2-PNXR D2-K64 · xyz", "138/0/36（AG/AAA/ILO mixed；seed=1234）", "PointNeXt-R", "xyz-only；vs S2 xyz+geom"),
        }
        for experiment_id in XYZ_IDS:
            row = overall.max_row + 1
            copy_row_style(overall, row - 1, row, 37)
            record = xyz["records"][experiment_id]
            label, split, backbone, change = xyz_labels[experiment_id]
            for column, value in enumerate(overview_values(record, label=label, split=split, backbone=backbone, change=change), 1):
                overall.cell(row, column, value)
        row = overall.max_row + 1
        copy_row_style(overall, row - 1, row, 37)
        record = attention["records"][ATTN_ID]
        for column, value in enumerate(overview_values(record, label="S3+SA3 attention s1234", split="138/0/36（AG/AAA/ILO mixed；seed=1234）", backbone="PointNeXt-R + LocalGeoPE", change="+SA3 coarse attention（32 centers / 4 heads）"), 1):
            overall.cell(row, column, value)

        summary = workbook["汇总对比"]
        title_rows = {cell.value: cell.row for cell in summary["A"] if cell.value in {XYZ_TITLE, ATTN_TITLE}}
        if title_rows:
            first = min(title_rows.values())
            expected_last = summary.max_row
            if set(title_rows) != {XYZ_TITLE, ATTN_TITLE} or expected_last - first + 1 != 8:
                raise RuntimeError("existing xyz/attention summary sections are not the expected trailing block")
            summary.delete_rows(first, 8)
        row = summary.max_row + 1
        row = summary_title(summary, row, XYZ_TITLE)
        row = write_header(summary, row, ["对照", "父 R²_cb", "XYZ R²_cb", "ΔR²_cb", "ΔMAE Pa", "ΔRMSE Pa", "Δhigh-WSS R²", "Δtop10 IoU", "病例 ΔR² 95%CI", "结论"])
        for result in xyz["paired_comparisons"]:
            control, treatment = xyz["records"][result["control"]], xyz["records"][result["treatment"]]
            delta, case = result["aggregate_treatment_minus_control"], result["paired_case_stats"]["case_r2"]
            summary.append([treatment["experiment_id"], control["best"]["physical_r2_casebalanced"], treatment["best"]["physical_r2_casebalanced"], delta["physical_r2_casebalanced"], delta["physical_mae"], delta["physical_rmse"], delta["high_wss_r2"], delta["physical_top10_iou"], f"{case['mean_delta']:+.4f} [{case['ci95_low']:+.4f},{case['ci95_high']:+.4f}]", "No-Go：去除几何特征后整体退化"])
        row = summary.max_row + 1
        row = summary_title(summary, row, ATTN_TITLE)
        row = write_header(summary, row, ["比较", "父 R²_cb", "attention R²_cb", "ΔR²_cb", "ΔMAE Pa", "ΔRMSE Pa", "Δhigh-WSS R²", "ΔAG / AAA / ILO", "病例 ΔR² 95%CI", "判定"])
        paired = attention["paired_comparison"]
        parent, child = attention["records"][paired["control"]], attention["records"][paired["treatment"]]
        delta, domains, case = paired["aggregate_treatment_minus_control"], paired["domain_r2_delta"], paired["paired_case_stats"]["case_r2"]
        summary.append(["SA3 32 centers / 4 heads", parent["best"]["physical_r2_casebalanced"], child["best"]["physical_r2_casebalanced"], delta["physical_r2_casebalanced"], delta["physical_mae"], delta["physical_rmse"], delta["high_wss_r2"], f"{domains['AG']:+.4f} / {domains['AAA']:+.4f} / {domains['ILO']:+.4f}", f"{case['mean_delta']:+.4f} [{case['ci95_low']:+.4f},{case['ci95_high']:+.4f}]", "Weak-Go 单种子信号；需三种子确认"])
        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(candidate)
        workbook.close()
        check = load_workbook(candidate, read_only=True, data_only=False)
        if any(name in check.sheetnames for name in LEGACY_SHEETS):
            raise RuntimeError("legacy result sheets survived removal")
        ids_after = {row[0] for row in check["实验矩阵总览"].iter_rows(min_col=37, max_col=37, values_only=True)}
        if not ids.issubset(ids_after):
            raise RuntimeError("overview readback missing xyz or attention rows")
        titles_after = {row[0] for row in check["汇总对比"].iter_rows(min_col=1, max_col=1, values_only=True)}
        if not {XYZ_TITLE, ATTN_TITLE}.issubset(titles_after):
            raise RuntimeError("summary readback missing xyz or attention sections")
        check.close()
        render = Path(temp_dir) / "render"
        render.mkdir()
        subprocess.run(["/usr/bin/libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(render), str(candidate)], check=True, capture_output=True, text=True)
        if not (render / candidate.with_suffix(".pdf").name).is_file():
            raise RuntimeError("LibreOffice render check failed")
        shutil.copy2(candidate, BOOK)
    print(json.dumps({"status": "updated_and_rendered", "workbook": str(BOOK), "overview_rows_added": 4, "summary_sections_added": 2, "removed_sheets": list(LEGACY_SHEETS)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
