#!/usr/bin/env python3
"""Append the completed D2 fixed-split PNXR+LocalGeoPE control to existing workbook views."""
from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from training_wss_min.tools.update_s3_regularization_xlsx import matrix_values
from training_wss_min.tools.update_s3_rootcause_xlsx import copy_row_style


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = ROOT / "training_wss_min/preflight/d2_c125_k64_pnxr_geope_20260726_results_analysis.json"
EXPERIMENT_ID = "d2_c125_k64_pnxr_geope"
TITLE = "2026-07-26｜D2 c125×k64 固定 106/0/27：PointNeXt-R + LocalGeoPE（seed1234）"


def write_title(ws, row: int, text: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    cell = ws.cell(row, 1, text)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    return row + 1


def write_header(ws, row: int, values: list[str]) -> int:
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    return row + 1


def main() -> None:
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    record = analysis["records"][EXPERIMENT_ID]
    paired = analysis["paired_comparison"]
    if record["integrity"] != "passed":
        raise RuntimeError("analysis integrity failed")
    with tempfile.TemporaryDirectory(prefix="wss_d2c125geo_xlsx_") as tmp:
        candidate = Path(tmp) / BOOK.name
        shutil.copy2(BOOK, candidate)
        workbook = load_workbook(candidate, data_only=False)
        overall = workbook["实验矩阵总览"]
        if EXPERIMENT_ID in {cell.value for cell in overall["AK"]}:
            raise RuntimeError(f"duplicate experiment row: {EXPERIMENT_ID}")
        row = overall.max_row + 1
        copy_row_style(overall, row - 1, row, 37)
        values = matrix_values({**record, "variant": "head_dropout01", "seed": 1234})
        values[0:4] = [
            "D2 c125×k64 + PointNeXt-R + LocalGeoPE",
            "106/0/27（AG/AAA；seed=1234）",
            "PointNeXt-R + LocalGeoPE",
            "D2 + sa_blocks=1/1/0 + 7D LocalGeoPE；其余字段冻结",
        ]
        for col, value in enumerate(values, 1):
            overall.cell(row, col, value)

        summary = workbook["汇总对比"]
        if TITLE in {cell.value for cell in summary["A"]}:
            raise RuntimeError("duplicate summary section")
        row = write_title(summary, summary.max_row + 1, TITLE)
        row = write_header(summary, row, [
            "严格比较", "D2 R²_cb", "PNXR+GeoPE R²_cb", "ΔR²_cb", "ΔMAE", "ΔRMSE",
            "Δhigh-WSS R²", "Δtop10 IoU", "ΔAG", "ΔAAA", "病例 ΔR² 95%CI", "结论",
        ])
        control = analysis["records"][paired["control"]]
        delta = paired["aggregate_treatment_minus_control"]
        case = paired["paired_case_stats"]["case_r2"]
        domain = analysis["domain_r2_delta"]
        summary.append([
            "D2 c125×k64 → +PointNeXt-R + LocalGeoPE",
            control["best"]["physical_r2_casebalanced"], record["best"]["physical_r2_casebalanced"],
            delta["physical_r2_casebalanced"], delta["physical_mae"], delta["physical_rmse"],
            delta["high_wss_r2"], delta["physical_top10_iou"], domain["AG"], domain["AAA"],
            f"{case['mean_delta']:+.4f} [{case['ci95_low']:+.4f},{case['ci95_high']:+.4f}]；{case['wins']}/{case['losses']}",
            "single-seed fixed test27 screening",
        ])
        workbook.save(candidate)
        workbook.close()
        check = load_workbook(candidate, read_only=True, data_only=False)
        saved = {row[0] for row in check["实验矩阵总览"].iter_rows(min_col=37, max_col=37, values_only=True)}
        titles = {row[0] for row in check["汇总对比"].iter_rows(min_col=1, max_col=1, values_only=True)}
        if EXPERIMENT_ID not in saved or TITLE not in titles:
            raise RuntimeError("xlsx readback failed")
        check.close()
        render = Path(tmp) / "render"
        render.mkdir()
        subprocess.run(["/usr/bin/libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(render), str(candidate)], check=True, capture_output=True, text=True)
        if not (render / candidate.with_suffix(".pdf").name).is_file():
            raise RuntimeError("LibreOffice render check failed")
        shutil.copy2(candidate, BOOK)
    print(json.dumps({"status": "updated_and_rendered", "workbook": str(BOOK), "experiment_id": EXPERIMENT_ID}, ensure_ascii=False))


if __name__ == "__main__":
    main()
