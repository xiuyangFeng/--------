#!/usr/bin/env python3
"""Recompute affine-fit R² for every historical PointNet workbook experiment.

The workflow is intentionally resumable:

1. ``prepare`` matches workbook rows to the exact historical metrics/checkpoint.
2. ``worker`` performs inference and writes one small JSON per unique source.
3. ``finalize`` updates the workbook and writes an auditable CSV/JSON summary.

The fitted line is always ``prediction = a * truth + b``.  Both point-pooled
and case-balanced shared-line fits are reported in physical and normalized
spaces.  Historical raw R² values are reproduced as a guard before a result is
accepted for workbook backfill.
"""

from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import json
import math
import re
import shutil
import time
from pathlib import Path

import numpy as np


REPO = Path(__file__).resolve().parents[2]
WORKBOOK = REPO / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
OUTPUT_DIR = REPO / "outputs/wss_min/linear_fit_r2_backfill_20260807"
AGGREGATE_ROWS = {
    52: [46, 48, 50],
    53: [47, 49, 51],
    60: [20, 56, 58],
    61: [55, 57, 59],
}
FIT_KEYS = (
    "physical_r2_fit_casebalanced",
    "physical_fit_casebalanced_slope",
    "physical_fit_casebalanced_intercept",
    "physical_r2_fit_pooled",
    "physical_fit_pooled_slope",
    "physical_fit_pooled_intercept",
    "normalized_r2_fit_casebalanced",
    "normalized_fit_casebalanced_slope",
    "normalized_fit_casebalanced_intercept",
    "normalized_r2_fit_pooled",
    "normalized_fit_pooled_slope",
    "normalized_fit_pooled_intercept",
)


def _finite(value) -> bool:
    return isinstance(value, (int, float)) and math.isfinite(float(value))


def _nested(mapping: dict, *path: str):
    value = mapping
    for key in path:
        if not isinstance(value, dict) or key not in value:
            return None
        value = value[key]
    return float(value) if _finite(value) else None


def _metric_vector(result: dict) -> dict[str, float | None]:
    normalized = result.get("normalized", result)
    return {
        "physical_r2_casebalanced": _nested(result, "field_casebalanced", "r2"),
        "physical_r2_pooled": _nested(result, "field", "r2"),
        "physical_r2_case_mean": _nested(result, "aggregate", "r2_casemean"),
        "normalized_r2_casebalanced": _nested(
            normalized, "field_casebalanced", "r2"
        ),
        "normalized_r2_pooled": _nested(normalized, "field", "r2"),
        "normalized_r2_case_mean": _nested(
            normalized, "aggregate", "r2_casemean"
        ),
    }


def _workbook_vector(sheet, row: int) -> dict[str, float]:
    columns = {
        "physical_r2_casebalanced": 7,
        "physical_r2_pooled": 8,
        "physical_r2_case_mean": 9,
        "normalized_r2_casebalanced": 22,
        "normalized_r2_pooled": 23,
        "normalized_r2_case_mean": 24,
    }
    return {
        key: float(sheet.cell(row, column).value)
        for key, column in columns.items()
        if _finite(sheet.cell(row, column).value)
    }


def _find_run_dir(metrics_path: Path) -> Path:
    for parent in (metrics_path.parent, *metrics_path.parents):
        if (parent / "config.json").is_file():
            return parent
    raise FileNotFoundError(f"no run config above {metrics_path}")


def _eval_partition(result_key: str) -> str:
    key = result_key.lower()
    if key.startswith("val"):
        return "val"
    if key.startswith("train"):
        return "train"
    if key.startswith("test"):
        return "test"
    raise ValueError(f"cannot infer partition from result key {result_key!r}")


def _source_id(metrics_path: Path, result_key: str, checkpoint: str) -> str:
    raw = f"{metrics_path.resolve()}::{result_key}::{checkpoint}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:20]


def _candidate_sources() -> list[dict]:
    sources = []
    for metrics_path in REPO.glob("training_wss_min/runs/**/metrics.json"):
        try:
            payload = json.loads(metrics_path.read_text(encoding="utf-8"))
        except Exception:
            continue
        for result_key, result in payload.items():
            if not isinstance(result, dict) or "field" not in result:
                continue
            checkpoint = "last" if "ckpt_last" in metrics_path.parts else "best"
            try:
                partition = _eval_partition(result_key)
                run_dir = _find_run_dir(metrics_path)
            except (ValueError, FileNotFoundError):
                continue
            sources.append({
                "metrics_path": str(metrics_path.resolve()),
                "result_key": result_key,
                "partition": partition,
                "checkpoint": checkpoint,
                "run_dir": str(run_dir.resolve()),
                "old_values": _metric_vector(result),
                "evaluation_split_path": result.get("evaluation_split_path"),
            })
    return sources


def _match_source(target: dict[str, float], raw_id: str, candidates: list[dict]) -> dict:
    scored = []
    raw_token = re.sub(r"[^a-z0-9]+", "", raw_id.lower())
    for candidate in candidates:
        old = candidate["old_values"]
        common = [key for key in target if old.get(key) is not None]
        if len(common) < min(3, len(target)):
            continue
        diffs = [abs(target[key] - float(old[key])) for key in common]
        path_token = re.sub(r"[^a-z0-9]+", "", candidate["run_dir"].lower())
        alias_penalty = 0 if raw_token and raw_token in path_token else 1
        checkpoint_penalty = 0 if candidate["checkpoint"] == "best" else 1
        score = (
            max(diffs), alias_penalty, checkpoint_penalty,
            sum(diffs) / len(diffs), -len(common), candidate["metrics_path"],
        )
        scored.append((score, candidate))
    if not scored:
        raise RuntimeError("no historical metrics candidate has enough common fields")
    scored.sort(key=lambda item: item[0])
    score, source = scored[0]
    if score[0] > 5e-4:
        raise RuntimeError(f"best historical match exceeds tolerance: {score}")
    return dict(source, match_max_abs_delta=score[0], match_common_fields=-score[4])


def prepare(workbook: Path, output_dir: Path, workers: int) -> None:
    from openpyxl import load_workbook
    from training_wss_min import config as C

    candidates = _candidate_sources()
    book = load_workbook(workbook, read_only=True, data_only=True)
    sheet = book["实验矩阵总览"]
    rows: dict[str, dict] = {}
    sources: dict[str, dict] = {}
    failures = []

    for row in range(3, sheet.max_row + 1):
        experiment = sheet.cell(row, 1).value
        if not experiment:
            continue
        raw_id = str(sheet.cell(row, 37).value or "")
        target = _workbook_vector(sheet, row)
        if row in AGGREGATE_ROWS:
            rows[str(row)] = {
                "row": row, "experiment": experiment, "kind": "aggregate",
                "component_rows": AGGREGATE_ROWS[row],
            }
            continue
        if len(target) < 2:
            rows[str(row)] = {"row": row, "experiment": experiment, "kind": "section"}
            continue
        try:
            if row == 3:
                # E0 predates the dual-space JSON schema: its historical file
                # stores physical metrics at the top level, while normalized
                # R² was produced by the companion full-point report.  Pin the
                # documented run explicitly and use the workbook normalized
                # value as the reproduction guard.
                source = next(
                    dict(candidate) for candidate in candidates
                    if candidate["result_key"] == "test"
                    and "pointnet_trainloss_e400/outputs/pointnet_xyzgeom/eval/metrics.json"
                    in candidate["metrics_path"]
                )
                source["old_values"]["normalized_r2_casebalanced"] = target.get(
                    "normalized_r2_casebalanced"
                )
                source["old_values"]["normalized_r2_pooled"] = None
                source["old_values"]["normalized_r2_case_mean"] = None
                source.update(match_max_abs_delta=0.0, match_common_fields=4)
            else:
                source = _match_source(target, raw_id, candidates)
            source_key = _source_id(
                Path(source["metrics_path"]), source["result_key"], source["checkpoint"]
            )
            if source_key not in sources:
                cfg = C.ExpConfig.from_json(Path(source["run_dir"]) / "config.json")
                split_path = source["evaluation_split_path"] or cfg.data.split_path
                group_key = json.dumps([
                    str(Path(cfg.data.data_root).resolve()),
                    cfg.data.required_frame_version,
                    cfg.data.target,
                    str(Path(split_path).resolve()),
                    source["partition"],
                ], ensure_ascii=False)
                source.update({
                    "source_key": source_key,
                    "split_path": str(Path(split_path).resolve()),
                    "group_key": group_key,
                    "workbook_rows": [],
                })
                sources[source_key] = source
            sources[source_key]["workbook_rows"].append(row)
            rows[str(row)] = {
                "row": row, "experiment": experiment, "raw_id": raw_id,
                "kind": "source", "source_key": source_key,
            }
        except Exception as exc:
            failures.append({"row": row, "experiment": experiment, "error": str(exc)})

    group_members: dict[str, list[str]] = {}
    for key, source in sources.items():
        group_members.setdefault(source["group_key"], []).append(key)
    worker_loads = [0] * workers
    for group_key, members in sorted(group_members.items(), key=lambda item: -len(item[1])):
        worker = min(range(workers), key=lambda index: worker_loads[index])
        for source_key in members:
            sources[source_key]["worker"] = worker
        worker_loads[worker] += len(members)

    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "rows").mkdir(exist_ok=True)
    manifest = {
        "schema_version": 1,
        "created": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "workbook": str(workbook.resolve()),
        "workers": workers,
        "worker_loads": worker_loads,
        "rows": rows,
        "sources": sources,
        "prepare_failures": failures,
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps({
        "rows": len(rows), "unique_sources": len(sources),
        "worker_loads": worker_loads, "failures": failures,
        "manifest": str((output_dir / "manifest.json").resolve()),
    }, ensure_ascii=False))


def _result_path(output_dir: Path, source_key: str) -> Path:
    return output_dir / "rows" / f"{source_key}.json"


def _space_fit(true_by_case, pred_by_case) -> dict:
    from training_wss_min import metrics as M

    pooled = M.linear_fit_metrics(np.concatenate(true_by_case), np.concatenate(pred_by_case))
    casebalanced = M.casebalanced_linear_fit_metrics(true_by_case, pred_by_case)
    return {
        "pooled": pooled,
        "casebalanced": casebalanced,
        "raw_r2_pooled": M.r2_score(
            np.concatenate(true_by_case), np.concatenate(pred_by_case)
        ),
        "raw_r2_casebalanced": M.casebalanced_field_metrics(
            true_by_case, pred_by_case
        )["r2"],
    }


def worker(output_dir: Path, worker_index: int, force: bool) -> None:
    import torch
    from training_wss_min import config as C, dataset as D
    from training_wss_min.evaluate import (
        load_model_from_run, load_wss_stats_for_run, predict_case_norm,
    )

    manifest = json.loads((output_dir / "manifest.json").read_text(encoding="utf-8"))
    selected = [
        source for source in manifest["sources"].values()
        if int(source["worker"]) == worker_index
    ]
    selected.sort(key=lambda source: (source["group_key"], min(source["workbook_rows"])))
    base_cases: dict[tuple, dict] = {}
    split_labels: dict[tuple[str, str], list[tuple[str, str]]] = {}
    case_feature_tables: dict[str, dict] = {}
    device = "cuda" if torch.cuda.is_available() else "cpu"

    for index, source in enumerate(selected, start=1):
        destination = _result_path(output_dir, source["source_key"])
        if destination.is_file() and not force:
            existing = json.loads(destination.read_text(encoding="utf-8"))
            if existing.get("status") == "ok" and existing.get("reproduction_ok"):
                print(f"[{worker_index}] skip {index}/{len(selected)} {source['workbook_rows']}", flush=True)
                continue
        started = time.time()
        result = {
            "source_key": source["source_key"],
            "workbook_rows": source["workbook_rows"],
            "run_dir": source["run_dir"],
            "metrics_path": source["metrics_path"],
            "result_key": source["result_key"],
            "checkpoint": source["checkpoint"],
            "partition": source["partition"],
        }
        try:
            run_dir = Path(source["run_dir"])
            cfg = C.ExpConfig.from_json(run_dir / "config.json")
            wss_stats = load_wss_stats_for_run(run_dir)
            labels_key = (source["split_path"], source["partition"])
            labels = split_labels.setdefault(
                labels_key, D.load_split_cases(*labels_key)
            )
            feature_table = None
            if cfg.data.case_features_path:
                feature_path = str(Path(cfg.data.case_features_path).resolve())
                feature_table = case_feature_tables.setdefault(
                    feature_path, D.load_case_feature_table(feature_path)
                )
            cases = []
            effective_data_root = Path(cfg.data.data_root)
            legacy_snapshot = REPO / "data_wss_min/_snapshots/AG_legacy_v3_20260716_signedoff"
            if (cfg.data.required_frame_version is None
                    and Path(source["split_path"]).name == "split_AG_wss_min_v1_traintest.json"
                    and legacy_snapshot.is_dir()):
                effective_data_root = legacy_snapshot
            data_root = str(effective_data_root.resolve())
            for cohort, case_name in labels:
                cache_key = (
                    data_root, cfg.data.required_frame_version, cfg.data.target,
                    cohort, case_name,
                )
                if cache_key not in base_cases:
                    base_cases[cache_key] = D.load_case(
                        cohort, case_name, wss_stats,
                        target=cfg.data.target,
                        target_normalization=cfg.data.target_normalization,
                        data_root=effective_data_root,
                        required_frame_version=cfg.data.required_frame_version,
                    )
                case = dict(base_cases[cache_key])
                case["_linear_fit_base_cache_key"] = cache_key
                case["y_norm"], case["target_normalization_meta"] = D.normalize_target(
                    case["y_raw"], wss_stats, cfg.data.target_normalization
                )
                case["y_norm"] = case["y_norm"].astype(np.float32)
                case["target_normalization"] = cfg.data.target_normalization
                if feature_table is not None:
                    unit_id = f"{cohort}/{case_name}"
                    for key, value in feature_table[unit_id].items():
                        case[key] = float(value)
                cases.append(case)

            cfg_loaded, feat_stats, model, _ = load_model_from_run(
                run_dir, device, source["checkpoint"]
            )
            if source["result_key"] == "test_all27_vertex_metrics":
                # This one historical diagnostic was produced by
                # evaluate_mapping_blocked_run.py before that helper called
                # model.eval().  Reproduce its recorded prediction state
                # exactly for a like-for-like workbook backfill.
                model.train()
            else:
                model.eval()
            support_n = int(cfg_loaded.data.support_n_points or cfg_loaded.data.wall_n_points)
            support_sampling = cfg_loaded.data.support_sampling or cfg_loaded.data.sampling
            if (cfg_loaded.eval.fixed_support and support_sampling == "fps_multistart"
                    and not cfg_loaded.data.fps_cache_dir):
                # Evaluation always selects pool entry 0 (run_seed=case_index=0
                # in predict_case_norm).  Build only that exact entry instead
                # of the seven unused training-time alternatives, then retain
                # it in the base-case cache for later historical models.
                for case in cases:
                    k_eff = min(len(case["pos"]), support_n)
                    if ("_fps_pool" not in case
                            or case.get("_fps_pool_k") != k_eff):
                        case["_fps_pool"] = D.farthest_point_sample(
                            case["pos"], k_eff, 17
                        )[None, :]
                        case["_fps_pool_k"] = k_eff
                        case["_fps_pool_size"] = 1
            pred_norm = [
                np.asarray(predict_case_norm(
                    model, case, cfg_loaded.data.input_features, feat_stats, device,
                    cfg=cfg_loaded,
                ), dtype=np.float64)
                for case in cases
            ]
            for case in cases:
                base = base_cases[case["_linear_fit_base_cache_key"]]
                for key, value in case.items():
                    if key.startswith("_fps_"):
                        base[key] = value
            true_norm = [np.asarray(case["y_norm"], dtype=np.float64) for case in cases]
            normalized = _space_fit(true_norm, pred_norm)
            physical = None
            if cfg.data.target_normalization != "case_max":
                pred_raw = []
                for prediction in pred_norm:
                    values = D.denormalize_wss(prediction, wss_stats)
                    if wss_stats.get("method") == "log_z":
                        values = np.clip(values, 0, None)
                    pred_raw.append(np.asarray(values, dtype=np.float64))
                true_raw = [np.asarray(case["y_raw"], dtype=np.float64) for case in cases]
                physical = _space_fit(true_raw, pred_raw)

            checks = []
            old = source["old_values"]
            for space_name, current in (("physical", physical), ("normalized", normalized)):
                if current is None:
                    continue
                for aggregation in ("pooled", "casebalanced"):
                    old_key = f"{space_name}_r2_{aggregation}"
                    old_value = old.get(old_key)
                    new_value = current[f"raw_r2_{aggregation}"]
                    if old_value is not None:
                        checks.append({
                            "key": old_key, "old": old_value, "new": new_value,
                            "abs_delta": abs(float(old_value) - float(new_value)),
                        })
            max_delta = max((item["abs_delta"] for item in checks), default=0.0)
            result.update({
                "status": "ok",
                "device": device,
                "n_cases": len(cases),
                "physical": physical,
                "normalized": normalized,
                "reproduction_checks": checks,
                "reproduction_max_abs_delta": max_delta,
                "reproduction_ok": max_delta <= 5e-4,
                "elapsed_seconds": time.time() - started,
            })
            del model
            if device == "cuda":
                torch.cuda.empty_cache()
        except Exception as exc:
            result.update({
                "status": "error", "error_type": type(exc).__name__,
                "error": str(exc), "elapsed_seconds": time.time() - started,
            })
        destination.write_text(
            json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        print(
            f"[{worker_index}] {index}/{len(selected)} rows={source['workbook_rows']} "
            f"status={result['status']} elapsed={result['elapsed_seconds']:.1f}s",
            flush=True,
        )


def _flatten_fit(result: dict) -> dict[str, float | None]:
    values = {key: None for key in FIT_KEYS}
    for space in ("physical", "normalized"):
        section = result.get(space)
        if section is None:
            continue
        for aggregation in ("casebalanced", "pooled"):
            fit = section[aggregation]
            values[f"{space}_r2_fit_{aggregation}"] = fit["r2_linear_fit"]
            values[f"{space}_fit_{aggregation}_slope"] = fit["linear_fit_slope"]
            values[f"{space}_fit_{aggregation}_intercept"] = fit["linear_fit_intercept"]
    return values


def _mean_rows(rows: list[dict[str, float | None]]) -> dict[str, float | None]:
    return {
        key: (
            float(np.mean([row[key] for row in rows if row.get(key) is not None]))
            if any(row.get(key) is not None for row in rows) else None
        )
        for key in FIT_KEYS
    }


def _copy_style(source, target) -> None:
    target._style = copy.copy(source._style)
    target.number_format = source.number_format
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.protection = copy.copy(source.protection)


def _overview_row_from_formula(value) -> int | None:
    if not isinstance(value, str) or not value.startswith("="):
        return None
    match = re.search(r"实验矩阵总览!\$?A\$?(\d+)", value)
    return int(match.group(1)) if match else None


def _update_workbook(workbook: Path, row_values: dict[int, dict]) -> None:
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter

    book = load_workbook(workbook)
    overview = book["实验矩阵总览"]
    start_col = 38
    headers = (
        "R² fit cb", "fit a cb", "fit b cb (Pa)",
        "R² fit pooled", "fit a pooled", "fit b pooled (Pa)",
        "R² fit cb", "fit a cb", "fit b cb",
        "R² fit pooled", "fit a pooled", "fit b pooled",
    )
    for merged in list(overview.merged_cells.ranges):
        if merged.min_row <= 1 <= merged.max_row and merged.max_col >= start_col:
            overview.unmerge_cells(str(merged))
    overview.merge_cells(start_row=1, start_column=38, end_row=1, end_column=43)
    overview.merge_cells(start_row=1, start_column=44, end_row=1, end_column=49)
    overview.cell(1, 38).value = "线性拟合补充｜物理 WSS（Pa）"
    overview.cell(1, 44).value = "线性拟合补充｜归一化空间"
    for column in range(38, 44):
        _copy_style(overview.cell(1, 7), overview.cell(1, column))
    for column in range(44, 50):
        _copy_style(overview.cell(1, 22), overview.cell(1, column))
    for offset, header in enumerate(headers):
        cell = overview.cell(2, start_col + offset)
        _copy_style(overview.cell(2, 7 if offset < 6 else 22), cell)
        cell.value = header
        cell.number_format = "0.000000"
    overview.cell(2, 38).comment = Comment(
        "对所有病例共享一条 prediction=a×truth+b 拟合线；病例总权重相等。",
        "Codex",
    )
    overview.cell(2, 41).comment = Comment(
        "全部点 pooled 后拟合 prediction=a×truth+b；R² 等价于 Pearson r²，需与 a、b 同报。",
        "Codex",
    )
    for row in range(3, overview.max_row + 1):
        values = row_values.get(row, {})
        for offset, key in enumerate(FIT_KEYS):
            column = start_col + offset
            cell = overview.cell(row, column)
            _copy_style(overview.cell(row, 7 if offset < 6 else 22), cell)
            cell.value = values.get(key)
            cell.number_format = "0.000000"
    for column in range(38, 50):
        overview.column_dimensions[get_column_letter(column)].width = 15
    if overview.auto_filter.ref:
        overview.auto_filter.ref = f"A2:AW{overview.max_row}"
    overview.print_area = f"$A$1:$AW${overview.max_row}"

    id_to_row = {
        str(overview.cell(row, 1).value): row
        for row in range(3, overview.max_row + 1)
        if overview.cell(row, 1).value
    }
    summary_specs = (
        ("汇总对比", 4, 2),
        ("教师汇报视图", 3, 2),
        ("RCR Oracle", 4, 1),
    )
    summary_headers = ("R² fit cb", "R² fit pooled", "fit a pooled", "fit b pooled (Pa)")
    overview_columns = (38, 41, 42, 43)
    for sheet_name, header_row, id_column in summary_specs:
        sheet = book[sheet_name]
        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row == 1 and merged.min_col == 1:
                sheet.unmerge_cells(str(merged))
                sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=19)
                break
        group_row = header_row - 1
        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row <= group_row <= merged.max_row and merged.max_col >= 16:
                sheet.unmerge_cells(str(merged))
        sheet.merge_cells(start_row=group_row, start_column=16, end_row=group_row, end_column=19)
        group_cell = sheet.cell(group_row, 16)
        _copy_style(sheet.cell(group_row, 5), group_cell)
        group_cell.value = "线性拟合补充（物理）"
        for offset, header in enumerate(summary_headers):
            cell = sheet.cell(header_row, 16 + offset)
            _copy_style(sheet.cell(header_row, 5), cell)
            cell.value = header
            cell.number_format = "0.000000"
        for row in range(header_row + 1, sheet.max_row + 1):
            overview_row = _overview_row_from_formula(sheet.cell(row, id_column).value)
            if overview_row is None:
                identifier = sheet.cell(row, id_column).value
                overview_row = id_to_row.get(str(identifier)) if identifier else None
            for offset, overview_column in enumerate(overview_columns):
                cell = sheet.cell(row, 16 + offset)
                _copy_style(sheet.cell(row, 5), cell)
                if overview_row is None:
                    cell.value = None
                else:
                    ref = f"实验矩阵总览!{get_column_letter(overview_column)}{overview_row}"
                    cell.value = f'=IF({ref}="","",{ref})'
                cell.number_format = "0.000000"
        for column in range(16, 20):
            sheet.column_dimensions[get_column_letter(column)].width = 15
        sheet.print_area = f"$A$1:$S${sheet.max_row}"

    notes = book["指标说明"]
    note_rows = [
        (
            "整体拟合", "R² fit pooled",
            "把全部点合并，拟合 prediction=a×truth+b 后计算回归 R²；等价于 Pearson r²。",
            "越高表示线性趋势越一致；1 最佳。",
            "允许整体缩放和偏移后，预测是否保持真实值的线性排序与相对变化。",
            "“线性拟合 R² 为 X；同时 a=Y、b=Z，说明趋势一致性与幅值偏差分别为……”",
            "不能替代 raw R²；a≠1 或 b≠0 时仍存在系统性校准偏差。",
        ),
        (
            "整体拟合", "R² fit cb",
            "每个病例总权重相等，在全部病例上共享一条 prediction=a×truth+b 拟合线。",
            "越高越好。",
            "避免大病例按点数主导后的病例等权线性趋势一致性。",
            "“病例等权线性拟合 R² 为 X。”",
            "不是每例各拟合一条线后再平均；所有病例共享同一 a、b。",
        ),
        (
            "校准参数", "fit a / fit b",
            "线性拟合 prediction=a×truth+b 的斜率与截距；物理空间 b 的单位为 Pa。",
            "理想为 a=1、b=0。",
            "区分比例压缩/放大与整体偏移。",
            "“a<1 表明预测动态范围收缩；b 表示整体偏移。”",
            "R² fit 高但 a、b 偏离理想值时，只能说明趋势好，不能说明绝对值准确。",
        ),
    ]
    existing = {str(notes.cell(row, 2).value) for row in range(1, notes.max_row + 1)}
    for values in note_rows:
        if values[1] in existing:
            row = next(row for row in range(1, notes.max_row + 1) if notes.cell(row, 2).value == values[1])
        else:
            row = notes.max_row + 1
        for column, value in enumerate(values, start=1):
            _copy_style(notes.cell(11, column), notes.cell(row, column))
            notes.cell(row, column).value = value
        notes.row_dimensions[row].height = 48
    notes.print_area = f"$A$1:$G${notes.max_row}"

    try:
        book.calculation.fullCalcOnLoad = True
        book.calculation.forceFullCalc = True
        book.calculation.calcMode = "auto"
    except AttributeError:
        pass
    temporary = workbook.with_suffix(".linear_fit_tmp.xlsx")
    book.save(temporary)
    shutil.move(temporary, workbook)


def finalize(workbook: Path, output_dir: Path) -> None:
    manifest = json.loads((output_dir / "manifest.json").read_text(encoding="utf-8"))
    row_values: dict[int, dict] = {}
    audit_rows = []
    failures = list(manifest.get("prepare_failures", []))

    for row_key, row_info in sorted(manifest["rows"].items(), key=lambda item: int(item[0])):
        row = int(row_key)
        if row_info["kind"] == "section":
            continue
        if row_info["kind"] == "aggregate":
            components = [row_values.get(component) for component in row_info["component_rows"]]
            if all(component is not None for component in components):
                row_values[row] = _mean_rows(components)
                audit_rows.append({
                    "row": row, "experiment": row_info["experiment"],
                    "status": "aggregate_mean", "source": str(row_info["component_rows"]),
                    **row_values[row],
                })
            else:
                failures.append({"row": row, "experiment": row_info["experiment"], "error": "missing aggregate component"})
            continue
        source_key = row_info["source_key"]
        result_path = _result_path(output_dir, source_key)
        if not result_path.is_file():
            failures.append({"row": row, "experiment": row_info["experiment"], "error": "worker result missing"})
            continue
        result = json.loads(result_path.read_text(encoding="utf-8"))
        if result.get("status") != "ok" or not result.get("reproduction_ok"):
            failures.append({
                "row": row, "experiment": row_info["experiment"],
                "error": result.get("error", "historical raw R² reproduction failed"),
                "reproduction_max_abs_delta": result.get("reproduction_max_abs_delta"),
            })
            continue
        row_values[row] = _flatten_fit(result)
        audit_rows.append({
            "row": row, "experiment": row_info["experiment"], "status": "recomputed",
            "source": result["run_dir"], "checkpoint": result["checkpoint"],
            "partition": result["partition"],
            "reproduction_max_abs_delta": result["reproduction_max_abs_delta"],
            **row_values[row],
        })

    _update_workbook(workbook, row_values)
    csv_path = output_dir / "linear_fit_metrics_by_overview_row.csv"
    fieldnames = sorted({key for row in audit_rows for key in row})
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(audit_rows)
    summary = {
        "status": "complete" if not failures else "partial",
        "workbook": str(workbook.resolve()),
        "filled_experiment_rows": len(row_values),
        "audit_csv": str(csv_path.resolve()),
        "failures": failures,
    }
    (output_dir / "backfill_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=("prepare", "worker", "finalize"))
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--workers", type=int, default=2)
    parser.add_argument("--worker-index", type=int, default=0)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    if args.command == "prepare":
        prepare(args.workbook, args.output_dir, args.workers)
    elif args.command == "worker":
        worker(args.output_dir, args.worker_index, args.force)
    else:
        finalize(args.workbook, args.output_dir)


if __name__ == "__main__":
    main()
