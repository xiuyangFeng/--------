#!/usr/bin/env python3
"""Idempotently backfill Slurm 10711 norm/loss matrix results into the xlsx."""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
A0_METRICS = ROOT / "training_wss_min/runs/pointnetpp_q2v_ilo/outputs/q2v_ilo_d3_pool2025_refit/eval/ckpt_best/metrics.json"
OUTPUT_ROOT = ROOT / "training_wss_min/runs/pointnetpp_norm_loss/outputs"

ARMS = (
    ("A0", "Q2V pool2025（原 point-pooled 对照，复用）", "q2v_ilo_d3_pool2025_refit", A0_METRICS, "复用 10487_5"),
    ("A1", "逐病例等权 log-z + MSE", "norm_loss_a1_casebal_mse", OUTPUT_ROOT / "norm_loss_a1_casebal_mse/eval/ckpt_best/metrics.json", "COMPLETED | 10711_0"),
    ("A1b", "逐父系等权 log-z + MSE", "norm_loss_a1b_cohortbal_mse", OUTPUT_ROOT / "norm_loss_a1b_cohortbal_mse/eval/ckpt_best/metrics.json", "COMPLETED | 10711_1"),
    ("A2", "A1 + cohort one-hot", "norm_loss_a2_casebal_cohortfeat", OUTPUT_ROOT / "norm_loss_a2_casebal_cohortfeat/eval/ckpt_best/metrics.json", "COMPLETED | 10711_2"),
    ("A3", "A1 + raw-Huber λ=0.2", "norm_loss_a3_casebal_rawhuber", OUTPUT_ROOT / "norm_loss_a3_casebal_rawhuber/eval/ckpt_best/metrics.json", "COMPLETED | 10711_3"),
    ("A4", "A1 + cohort one-hot + raw-Huber λ=0.2", "norm_loss_a4_casebal_cohortfeat_rawhuber", OUTPUT_ROOT / "norm_loss_a4_casebal_cohortfeat_rawhuber/eval/ckpt_best/metrics.json", "COMPLETED | 10711_4"),
)


def clone_row(ws, src: int, dst: int) -> None:
    for col in range(1, ws.max_column + 1):
        source, target = ws.cell(src, col), ws.cell(dst, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)
        if source.protection:
            target.protection = copy(source.protection)
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def metrics(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))["test"]


def values(m: dict) -> list:
    physical, normalized = m["field"], m["normalized"]["field"]
    p_cb, n_cb, a, h = m["field_casebalanced"], m["normalized"]["field_casebalanced"], m["aggregate"], m["hotspot"]
    return [
        p_cb["r2"], physical["r2"], a["r2_casemean"], a["r2_casemed"], a["r2_casep10"], f'{a["r2_negative_cases"]}/{a["n_cases"]}',
        physical["rmse"], physical["mae"], physical["nrmse_range"], a["nrmse_casemean"], physical["nmae_range"], a["nmae_casemean"],
        m["regional_field"]["high_wss"]["r2"], m["calibration"]["top10_pred_true_ratio"], h["top10_iou"],
        n_cb["r2"], normalized["r2"], m["normalized"]["aggregate"]["r2_casemean"], m["normalized"]["aggregate"]["r2_casemed"], m["normalized"]["aggregate"]["r2_casep10"],
        normalized["mae"], normalized["rmse"], normalized["nrmse_range"], m["normalized"]["aggregate"]["nrmse_casemean"], normalized["nmae_range"], m["normalized"]["aggregate"]["nmae_casemean"],
        h["spearman_all_casemean"], h["spearman_high_wss_casemean"], m["normalized"]["calibration"]["top10_pred_true_ratio"], m["calibration"]["p99_pred_true_ratio"],
    ]


def replace_or_append(ws, id_col: int, identifier: str, template: int) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, id_col).value == identifier:
            return row
    row = ws.max_row + 1
    clone_row(ws, template, row)
    return row


def update_overview(ws, rows: list[tuple]) -> None:
    for arm, label, exp, path, status in rows:
        identifier = f"NormLoss {arm}｜{label}"
        row = replace_or_append(ws, 1, identifier, 84)
        m = metrics(path)
        base = [identifier, "138/0/36（pool2025 historical test36；seed=1234）", "PointNet++ SA3",
                "5000→500→125→32；nsample=16；width=32", "vertex random5000 / SEP；FPS center", 5000]
        for col, value in enumerate(base + values(m) + [exp], 1):
            ws.cell(row, col).value = value


def update_teacher(ws, rows: list[tuple]) -> None:
    header = "Ⅺ Q2V-pool2025 归一化口径×尾部损失（2026-07-20；10711 单seed筛查）"
    header_row = next((r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == header), None)
    if header_row is None:
        header_row = ws.max_row + 1
        clone_row(ws, 83, header_row)
        ws.merge_cells(start_row=header_row, start_column=1, end_row=header_row, end_column=15)
        ws.cell(header_row, 1).value = header
    for arm, label, exp, path, status in rows:
        identifier = f"NormLoss {arm}｜{label}"
        row = replace_or_append(ws, 2, identifier, 83)
        if row <= header_row:
            row = ws.max_row + 1
            clone_row(ws, 83, row)
        m = metrics(path)
        a, h = m["aggregate"], m["hotspot"]
        v = ["归一化/损失", identifier, "PointNet++ SA3", "random5000 / SEP；FPS center", m["field_casebalanced"]["r2"], a["r2_casemean"], m["field"]["rmse"], m["field"]["nmae_range"], a["nmae_casemean"], m["regional_field"]["high_wss"]["r2"], m["normalized"]["field_casebalanced"]["r2"], m["normalized"]["field"]["nmae_range"], h["spearman_all_casemean"], h["top10_iou"], m["calibration"]["p99_pred_true_ratio"]]
        for col, value in enumerate(v, 1):
            ws.cell(row, col).value = value


def update_summary(ws, rows: list[tuple]) -> None:
    header = "Q2V-pool2025 norm/loss（10711；Δ 相对 A0，single seed=1234）"
    header_row = next((r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == header), None)
    if header_row is None:
        header_row = ws.max_row + 1
        clone_row(ws, 92, header_row)
        ws.cell(header_row, 1).value = header
        row = header_row + 1
        clone_row(ws, 93, row)
        for col, value in enumerate(["配置", "对照", "物理R²_cb", "ΔR²_cb", "归一化R²_cb", "Δ", "病例均值R²", "负例", "MAE", "ΔMAE", "Spearman", "Δ", "top10 IoU", "Δ", "high-WSS R²"], 1):
            ws.cell(row, col).value = value
    a0 = metrics(rows[0][3])
    for arm, label, exp, path, status in rows:
        identifier = f"NormLoss {arm}｜{label}"
        row = replace_or_append(ws, 1, identifier, 99)
        m = metrics(path)
        v = [identifier, exp, m["field_casebalanced"]["r2"], m["field_casebalanced"]["r2"] - a0["field_casebalanced"]["r2"], m["normalized"]["field_casebalanced"]["r2"], m["normalized"]["field_casebalanced"]["r2"] - a0["normalized"]["field_casebalanced"]["r2"], m["aggregate"]["r2_casemean"], m["aggregate"]["r2_negative_cases"], m["field"]["mae"], m["field"]["mae"] - a0["field"]["mae"], m["hotspot"]["spearman_all_casemean"], m["hotspot"]["spearman_all_casemean"] - a0["hotspot"]["spearman_all_casemean"], m["hotspot"]["top10_iou"], m["hotspot"]["top10_iou"] - a0["hotspot"]["top10_iou"], m["regional_field"]["high_wss"]["r2"]]
        for col, value in enumerate(v, 1):
            ws.cell(row, col).value = value


def main() -> None:
    for _, _, _, path, _ in ARMS:
        if not path.is_file():
            raise FileNotFoundError(path)
    wb = load_workbook(WORKBOOK)
    update_overview(wb["实验矩阵总览"], ARMS)
    update_teacher(wb["教师汇报视图"], ARMS)
    update_summary(wb["汇总对比"], ARMS)
    wb.save(WORKBOOK)
    print(WORKBOOK)


if __name__ == "__main__":
    main()
