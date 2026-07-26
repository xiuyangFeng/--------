#!/usr/bin/env python3
"""Append the completed D2-K64 ILO/structure screen to the WSS results workbook."""

from __future__ import annotations

import json
import shutil
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = ROOT / "training_wss_min/preflight/d2_k64_ilo_structure_results_analysis.json"
SECTION_TITLE = "2026-07-23｜D2 c125×k64：ILO 两协议与结构模块单种子筛选"

ORDER = (
    "f1_d2k64_ilo41_fixed_test27",
    "m0_d2k64_zeroshot_mixed36",
    "m1_d2k64_mixed138_test36",
    "s1_d2k64_mfeat_mixed",
    "s2_d2k64_pnxr_mixed",
    "s3_d2k64_pnxr_geope_mixed",
    "s4_d2k64_pnxr_attn_mixed",
    "s5c_d2k64_pnxr_independent_interp_mixed",
    "s5_d2k64_pnxr_sep_mixed",
)

META = {
    "f1_d2k64_ilo41_fixed_test27": (
        "F1 D2-K64 + ILO41（fixed test27）",
        "147/0/27（AG/AAA test27固定；single seed=1234）",
        "PointNet++ SA3",
        "5000→125/125/32；64/16/16；width32；base",
        "random / SAME；FPS center",
    ),
    "m0_d2k64_zeroshot_mixed36": (
        "M0 D2-K64 ILO zero-shot（mixed test36）",
        "106/0/36（AG/AAA/ILO mixed test36；single seed=1234）",
        "PointNet++ SA3",
        "5000→125/125/32；64/16/16；width32；base",
        "random / SAME；FPS center",
    ),
    "m1_d2k64_mixed138_test36": (
        "M1/S0 D2-K64 mixed138（mixed test36）",
        "138/0/36（AG/AAA/ILO mixed；single seed=1234）",
        "PointNet++ SA3",
        "5000→125/125/32；64/16/16；width32；base",
        "random / SAME；FPS center",
    ),
    "s1_d2k64_mfeat_mixed": (
        "S1 D2-K64 + M-FEAT",
        "138/0/36（AG/AAA/ILO mixed；single seed=1234）",
        "PointNet++ SA3",
        "M1 + radius_gradient + coord_scale（8D）",
        "random / SAME；FPS center",
    ),
    "s2_d2k64_pnxr_mixed": (
        "S2 D2-K64 + PointNeXt-R",
        "138/0/36（AG/AAA/ILO mixed；single seed=1234）",
        "PointNeXt-R",
        "M1 + sa_blocks=1/1/0；无Drop",
        "random / SAME；FPS center",
    ),
    "s3_d2k64_pnxr_geope_mixed": (
        "S3 D2-K64 PointNeXt-R + LocalGeoPE",
        "138/0/36（AG/AAA/ILO mixed；single seed=1234）",
        "PointNeXt-R + LocalGeoPE",
        "S2 + LocalGeoPE v1",
        "random / SAME；FPS center",
    ),
    "s4_d2k64_pnxr_attn_mixed": (
        "S4 D2-K64 PointNeXt-R + coarse attention",
        "138/0/36（AG/AAA/ILO mixed；single seed=1234）",
        "PointNeXt-R + Attention",
        "S2 + SA3 coarse attention",
        "random / SAME；FPS center",
    ),
    "s5c_d2k64_pnxr_independent_interp_mixed": (
        "S5C D2-K64 independent + 3NN control",
        "138/0/36（AG/AAA/ILO mixed；single seed=1234）",
        "PointNeXt-R",
        "S2 + independent query + 3NN interpolate",
        "random / independent；FPS center",
    ),
    "s5_d2k64_pnxr_sep_mixed": (
        "S5 D2-K64 Conservative SEP-Kernel",
        "138/0/36（AG/AAA/ILO mixed；single seed=1234）",
        "PointNeXt-R + SEP",
        "S5C + Conservative SEP-Kernel",
        "random / independent SEP；FPS center",
    ),
}


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def copy_row_style(ws, source_row: int, target_row: int, columns: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, columns + 1):
        source, target = ws.cell(source_row, col), ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def matrix_values(item: dict) -> list:
    exp_id = item["experiment_id"]
    label, split, model, capacity, sampling = META[exp_id]
    raw = load_json(Path(item["run_dir"]) / "eval/ckpt_best/metrics.json")["test"]
    norm = raw["normalized"]
    return [
        label,
        split,
        model,
        capacity,
        sampling,
        5000,
        raw["field_casebalanced"]["r2"],
        raw["field"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["aggregate"]["r2_casemed"],
        raw["aggregate"]["r2_casep10"],
        f"{raw['aggregate']['r2_negative_cases']}/{raw['aggregate']['n_cases']}",
        raw["field"]["rmse"],
        raw["field"]["mae"],
        raw["field"]["nrmse_range"],
        raw["aggregate"]["nrmse_casemean"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        raw["calibration"]["top10_pred_true_ratio"],
        raw["hotspot"]["top10_iou_casemean"],
        norm["field_casebalanced"]["r2"],
        norm["field"]["r2"],
        norm["aggregate"]["r2_casemean"],
        norm["aggregate"]["r2_casemed"],
        norm["aggregate"]["r2_casep10"],
        norm["field_casebalanced"]["mae"],
        norm["field_casebalanced"]["rmse"],
        norm["field"]["nrmse_range"],
        norm["aggregate"]["nrmse_casemean"],
        norm["field"].get("nmae_range"),
        norm["aggregate"].get("nmae_casemean"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["spearman_high_wss_casemean"],
        norm["calibration"]["top10_pred_true_ratio"],
        norm["calibration"]["p99_pred_true_ratio"],
        exp_id,
    ]


def teacher_values(item: dict) -> list:
    exp_id = item["experiment_id"]
    label, _, model, _, sampling = META[exp_id]
    raw = load_json(Path(item["run_dir"]) / "eval/ckpt_best/metrics.json")["test"]
    norm = raw["normalized"]
    return [
        "D2-K64 ILO+结构",
        label,
        model,
        f"5000 / {'SEP' if 'SEP' in sampling else ('independent' if 'independent' in sampling else 'SAME')}",
        raw["field_casebalanced"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["field"]["rmse"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        norm["field_casebalanced"]["r2"],
        norm["field"].get("nmae_range"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["top10_iou_casemean"],
        norm["calibration"]["p99_pred_true_ratio"],
    ]


def append_summary(summary, analysis: dict) -> int:
    row = summary.max_row + 1
    summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=15)
    title = summary.cell(row, 1, SECTION_TITLE)
    title.font = Font(bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E78")
    start = row
    row += 1
    summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=15)
    summary.cell(
        row,
        1,
        "严格配对差：F1−F0、M1−M0、S1−M1、S2−M1、S3−S2、S4−S2、S5−S5C；单种子仅作筛选。",
    ).font = Font(bold=True)
    row += 1
    headers = [
        "比较",
        "处理臂",
        "对照臂",
        "R²_cb",
        "ΔR²_cb",
        "MAE_cb",
        "ΔMAE",
        "RMSE_cb",
        "ΔRMSE",
        "high-WSS R²",
        "Δ",
        "top10 IoU",
        "Δ",
        "病例R²胜/负",
        "病例均值ΔR² 95%CI",
    ]
    for col, value in enumerate(headers, 1):
        cell = summary.cell(row, col, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    by_id = {item["experiment_id"]: item for item in analysis["records"]}
    for pair in analysis["paired_comparisons"]:
        row += 1
        treatment = by_id[pair["treatment"]]
        best = treatment["best"]
        delta = pair["aggregate_treatment_minus_control"]
        stats = pair["paired_case_stats"]["case_r2"]
        values = [
            pair["comparison"],
            treatment["label"],
            by_id[pair["control"]]["label"],
            best["physical_r2_casebalanced"],
            delta["physical_r2_casebalanced"],
            best["physical_mae"],
            delta["physical_mae"],
            best["physical_rmse"],
            delta["physical_rmse"],
            best["high_wss_r2"],
            delta["high_wss_r2"],
            best["physical_top10_iou"],
            delta["physical_top10_iou"],
            f"{stats['wins']}/{stats['losses']}",
            f"{stats['mean_delta']:+.4f} [{stats['ci95_low']:+.4f}, {stats['ci95_high']:+.4f}]",
        ]
        for col, value in enumerate(values, 1):
            summary.cell(row, col, value)
    return start


def main() -> None:
    analysis = load_json(ANALYSIS)
    records = {item["experiment_id"]: item for item in analysis["records"]}
    if any(exp_id not in records for exp_id in ORDER):
        raise RuntimeError("analysis is missing one or more expected completed runs")
    if any(records[exp_id]["integrity"] != "passed" for exp_id in ORDER):
        raise RuntimeError("refusing workbook update because a run failed integrity audit")

    with tempfile.TemporaryDirectory(prefix="wss_d2k64_xlsx_") as temp_dir:
        candidate = Path(temp_dir) / XLSX.name
        shutil.copy2(XLSX, candidate)
        workbook = load_workbook(candidate, data_only=False)
        overall = workbook["实验矩阵总览"]
        existing_ids = {cell.value for cell in overall["AK"]}
        duplicates = set(ORDER) & existing_ids
        if duplicates:
            raise RuntimeError(f"refusing duplicate append: {sorted(duplicates)}")
        first_overall = overall.max_row + 1
        for exp_id in ORDER:
            row = overall.max_row + 1
            copy_row_style(overall, row - 1, row, 37)
            for col, value in enumerate(matrix_values(records[exp_id]), 1):
                overall.cell(row, col, value)

        teacher = workbook["教师汇报视图"]
        existing_teacher = {cell.value for cell in teacher["B"]}
        if any(META[exp_id][0] in existing_teacher for exp_id in ORDER):
            raise RuntimeError("refusing duplicate append in 教师汇报视图")
        first_teacher = teacher.max_row + 1
        for exp_id in ORDER:
            row = teacher.max_row + 1
            copy_row_style(teacher, row - 1, row, 15)
            for col, value in enumerate(teacher_values(records[exp_id]), 1):
                teacher.cell(row, col, value)

        summary = workbook["汇总对比"]
        if SECTION_TITLE in {cell.value for cell in summary["A"]}:
            raise RuntimeError("refusing duplicate append in 汇总对比")
        summary_start = append_summary(summary, analysis)
        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(candidate)

        check = load_workbook(candidate, read_only=True, data_only=False)
        check_ids = {
            row[0].value
            for row in check["实验矩阵总览"].iter_rows(
                min_col=37, max_col=37
            )
        }
        if not set(ORDER).issubset(check_ids):
            raise RuntimeError("saved workbook failed 实验矩阵总览 readback")
        teacher_labels = {
            row[0].value
            for row in check["教师汇报视图"].iter_rows(min_col=2, max_col=2)
        }
        if not all(META[exp_id][0] in teacher_labels for exp_id in ORDER):
            raise RuntimeError("saved workbook failed 教师汇报视图 readback")
        summary_titles = {
            row[0].value
            for row in check["汇总对比"].iter_rows(min_col=1, max_col=1)
        }
        if SECTION_TITLE not in summary_titles:
            raise RuntimeError("saved workbook failed 汇总对比 readback")
        check.close()
        shutil.copy2(candidate, XLSX)

    print(
        json.dumps(
            {
                "status": "updated",
                "workbook": str(XLSX),
                "overall_rows": [first_overall, first_overall + len(ORDER) - 1],
                "teacher_rows": [first_teacher, first_teacher + len(ORDER) - 1],
                "summary_start_row": summary_start,
                "rows_added": len(ORDER),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
