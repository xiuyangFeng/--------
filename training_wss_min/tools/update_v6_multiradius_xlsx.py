"""Append evidence-qualified MS results while preserving the historical workbook."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook

from training_wss_min.tools.report_v6_multiradius import (
    ANCHOR, EXP, LABELS, matrix_arms, report, load_metrics,
)
from training_wss_min.tools.update_v5_rerun_xlsx import BOOK, EXPECTED_SHEETS, LSA2
from training_wss_min.tools.update_v6_followup_xlsx import (
    append_title, append_values, teacher_metric_values, followup_overview_values,
)

PREFIX = "V6-多半径BT20260909"
TITLE = PREFIX + "｜MS1–MS6；M2底座；仅WSS；seed1234；本节checkpoint=best（训练总损失选模）"
TEACHER_TITLE = "ⅩⅤ " + TITLE


def label(arm):
    return f"{PREFIX}·{arm['id']} {Path(arm['config']).stem}"


def model_text(arm):
    mode = arm["resolved_config"]["model"].get("multiradius_context", "none")
    # Older prepared manifests use explicit boolean switches; retain compatibility.
    if arm["resolved_config"]["model"].get("multiradius_bottleneck"):
        mode = "bt"
    if arm["resolved_config"]["model"].get("multiradius_ffn"):
        mode = "ffn"
    return LSA2 + " + local wall branch + MS-SA3" + {"bt": " + BT", "ffn": " + token FFN"}.get(mode, "")


def remove_owned_tail(ws, title, arms):
    starts = [c.row for c in ws["A"] if c.value == title]
    if not starts:
        return
    if len(starts) != 1:
        raise ValueError(f"Duplicate MS section marker in {ws.title}")
    first = starts[0]
    if ws.title == "实验矩阵总览":
        expected = [{1: label(a), 37: a["run_name"]} for a in arms]
        width = 49
    elif ws.title == "教师汇报视图":
        expected = [{1: "V6 多半径BT", 2: f"{a['id']} {LABELS[a['id']]}｜"} for a in arms]
        width = 19
    elif ws.title == "汇总对比":
        expected = [{1: "ID", 2: "父臂"}, {1: "MS0=M2", 15: ANCHOR}]
        expected += [{1: a["id"], 15: a["run_name"]} for a in arms]
        expected += [{1: "说明"}]
        width = 15
    else:
        raise ValueError(f"No MS ownership schema for {ws.title}")
    for offset, identity in enumerate(expected, 1):
        row = first + offset
        for column, value in identity.items():
            actual = ws.cell(row, column).value
            matches = actual == value
            if ws.title == "教师汇报视图" and column == 2:
                matches = isinstance(actual, str) and actual.startswith(value)
            if not matches:
                raise RuntimeError(f"Unknown MS row {ws.title}!{row}; refusing to remove it")
    last = first + len(expected)
    for row in ws.iter_rows(min_row=first):
        if any(c.value is not None and (c.row > last or c.column > width) for c in row):
            raise RuntimeError(f"Unowned data after/beside MS section in {ws.title}")
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= first:
            ws.unmerge_cells(str(merged))
    ws.delete_rows(first, ws.max_row - first + 1)


def retained_cells(wb):
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws for c in row if c.value is not None}


def main():
    _, summary = report("best")
    arms = matrix_arms()
    original = BOOK.read_bytes()
    original_hash = hashlib.sha256(original).hexdigest()
    backup = EXP / "WSS_PointNet实验矩阵与结果汇总last_backup_before_multiradius_20260909.xlsx"
    if not backup.exists():
        backup.write_bytes(original)
    with tempfile.TemporaryDirectory(prefix="v6_multiradius_workbook_") as td:
        candidate = Path(td) / BOOK.name
        candidate.write_bytes(original)
        wb = load_workbook(candidate, data_only=False)
        if wb.sheetnames != EXPECTED_SHEETS:
            raise ValueError(f"Unexpected sheets: {wb.sheetnames}")
        for name, title in (("实验矩阵总览", TITLE), ("教师汇报视图", TEACHER_TITLE), ("汇总对比", TITLE)):
            remove_owned_tail(wb[name], title, arms)
        preserved = retained_cells(wb)
        for name, title, width in (("实验矩阵总览", TITLE, 20), ("教师汇报视图", TEACHER_TITLE, 19), ("汇总对比", TITLE, 15)):
            append_title(wb[name], title, width)
        ov, te, su = wb["实验矩阵总览"], wb["教师汇报视图"], wb["汇总对比"]
        append_values(su, ["ID", "父臂", "状态", "物理R²_cb", "Δ vs MS0", "Δ vs父臂", "log_z R²_cb", "case mean", "MAE Pa", "high-WSS R²", "top10比", "p99比", "IoU", "变化", "结果目录"], header=True)
        base = summary["reference"]
        append_values(su, ["MS0=M2", "—", "已有best；seed1234", base["physical_r2_cb"], None, None, base["normalized_r2_cb"], base["case_mean"], base["mae"], base["high_wss_r2"], base["top10_ratio"], base["p99_ratio"], base["top10_iou"], "历史M2；原25D、原A5损失、independent query/3NN", ANCHOR])
        rows, teacher_rows, summary_rows = {}, {}, {}
        for arm in arms:
            entry = summary["arms"][arm["id"]]
            m = load_metrics(arm["run_name"], "best") if entry["evidence"]["valid_scientific_result"] else None
            model = model_text(arm)
            protocol = "random5000 / INDEPENDENT；seed1234；checkpoint=best"
            if m:
                vals = followup_overview_values(arm, m)
                vals[0], vals[2], vals[3], vals[4], vals[36] = label(arm), model, arm["hypothesis"], protocol, arm["run_name"]
            else:
                vals = [None] * 49
                vals[:6] = [label(arm), "V5 train138/test34；峰值单帧；seed1234", model,
                            arm["hypothesis"] + "｜" + entry["evidence"]["status"], protocol, 5000]
                vals[36] = arm["run_name"]
            rows[arm["id"]] = append_values(ov, vals)
            teacher_rows[arm["id"]] = append_values(te, ["V6 多半径BT", f"{arm['id']} {LABELS[arm['id']]}｜{entry['evidence']['status']}", model, protocol] + teacher_metric_values(entry, m))
            get = entry.get
            summary_rows[arm["id"]] = append_values(su, [arm["id"], entry["parent"], entry["evidence"]["status"], get("physical_r2_cb"), get("delta_physical"), get("delta_parent_physical"), get("normalized_r2_cb"), get("case_mean"), get("mae"), get("high_wss_r2"), get("top10_ratio"), get("p99_ratio"), get("top10_iou"), arm["hypothesis"], arm["run_name"]])
        append_values(su, ["说明", "本节全部best；文件名last沿用历史命名。best按训练总损失选模，last表另存matrix_tables_last.md。单seed1234、已暴露test34，仅探索性比较；G/M6/新几何待用户验收。",
                           "整批完结归档已核验" if summary["matrix_finalized"] else "整批完结归档尚未核验",
                           "总览M/N为Pa pooled RMSE/MAE；AA/AB为log_z pooled MAE/RMSE。教师O为log_z p99比，汇总L为Pa p99比。"])
        wb.save(candidate)
        wb.close()
        check = load_workbook(candidate, read_only=False, data_only=False)
        assert check.sheetnames == EXPECTED_SHEETS
        for (sheet, coordinate), value in preserved.items():
            if check[sheet][coordinate].value != value:
                raise AssertionError(f"Historical cell changed: {sheet}!{coordinate}")
        for arm in arms:
            r = rows[arm["id"]]
            assert check["实验矩阵总览"].cell(r, 37).value == arm["run_name"]
            expected = summary["arms"][arm["id"]].get("physical_r2_cb")
            actual = check["实验矩阵总览"].cell(r, 7).value
            assert actual is None if expected is None else abs(actual - expected) < 1e-12
        check.close()
        if hashlib.sha256(BOOK.read_bytes()).hexdigest() != original_hash:
            raise RuntimeError("Workbook changed concurrently; refusing to overwrite")
        staged = BOOK.with_suffix(".multiradius.tmp.xlsx")
        shutil.copy2(candidate, staged)
        staged.replace(BOOK)
    (EXP / "workbook_update.json").write_text(json.dumps({"book": str(BOOK), "backup": str(backup), "checkpoint": "best", "matrix_finalized": summary["matrix_finalized"], "completed_best": summary["completed_count"], "overview_rows": rows, "teacher_rows": teacher_rows, "summary_rows": summary_rows, "reference": ANCHOR, "preserved_nonempty_cells": len(preserved), "preserved_formulas": sum(isinstance(v, str) and v.startswith("=") for v in preserved.values())}, indent=2, ensure_ascii=False) + "\n")
    print(f"Updated MS workbook: {summary['completed_count']}/6 evaluated; rows={rows}")


if __name__ == "__main__":
    main()
