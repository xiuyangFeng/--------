#!/usr/bin/env python3
"""Backfill the WSS workbook with the V6 single-frame matrix (geometry / decoder / calibration arms).

Own section, separate from the V5 rerun, volume-target and bottleneck-transformer sections:
* 实验矩阵总览: a title row + one row per arm (每臂一个 seed，一行一个 run);
* 教师汇报视图: section ⅩⅢ, one row per arm plus the screening note;
* 汇总对比: a dated section comparing每个臂 against the R4 seed mean with the correct difference bands.

Statistics and the verdict text come from tools/report_v6_matrix.py / report_bt_matrix.verdict_for, so the
workbook can never disagree with the markdown tables. Every arm here has ONE seed, so no arm gets a verdict:
the section ranks candidates for a three-seed confirmation round. Idempotent; run with /usr/bin/python3.
"""
from __future__ import annotations

import json
import math
import shutil
import tempfile
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill

from training_wss_min.tools.update_v5_rerun_xlsx import (
    BLUE, BOOK, EXPECTED_SHEETS, GREEN, LSA2, NAVY, RUNS, THIN, WHITE, YELLOW,
    copy_row_style, delete_section, styled, unmerge_from,
)
from training_wss_min.tools.report_bt_matrix import BASE_ARM, NORM, PHYS, dig, seeds_of, verdict_for

ROOT = Path(__file__).resolve().parents[2]
MATRIX = ROOT / "training_wss_min/configs/v6_singleframe_20260909/matrix.json"
EXP = ROOT / "training_wss_min/experiments/v6_singleframe_20260909"
BACKUP = EXP / "WSS_PointNet实验矩阵与结果汇总last_backup_before_v6_20260909.xlsx"
OVERVIEW_TITLE = ("V6 单帧 WSS 矩阵（2026-09-09；峰值单帧、不加时间维 t）—— 每行一个 run（每臂 1 个 seed=1234）；"
                  "参照为 R4 全部 seed 的均值，判定带见 教师汇报视图 ⅩⅢ 节与 汇总对比")
TEACHER_TITLE = ("ⅩⅢ V6 单帧 WSS 矩阵（2026-09-09；局部分支 / 解码器 / 病例级幅值 / 局部微分几何 / NLL+Jensen；"
                 "每臂只跑 seed 1234，按既定统计合同一律不判定，本节用于排出补 seed 的优先级）")
TEACHER_TITLE_PREFIX = "ⅩⅢ V6 单帧 WSS 矩阵"
SECTION_TITLE = "2026-09-09｜V6 单帧 WSS 矩阵：几何/解码/校准候选筛选（参照 = R4 全部 seed 的均值；每臂单 seed 不判定）"
SPLIT = "V5 train138/test34；support 壁面 5000；峰值单帧；seed=1234"
STAGE = "V6 单帧矩阵"


def load_arm(run_name: str, checkpoint: str = "best"):
    path = RUNS / run_name / "eval" / f"ckpt_{checkpoint}" / "metrics.json"
    if not path.is_file():
        return None
    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload.get("test", payload)


def change_text(arm: dict) -> str:
    parts = [f"{k}: {v['from']} → {v['to']}" for k, v in arm["changes"].items()]
    return f"{arm['hypothesis']}｜唯一变化 " + "；".join(parts)


def overview_values(arm: dict, m: dict) -> list:
    n = m["normalized"]
    model = LSA2 + {"A1": "（+ 多尺度局部壁面分支）", "A2b": "（+ 16 邻域局部 query 解码器）",
                    "A3": "（+ 病例级幅值头）"}.get(arm["id"], "")
    protocol = ("random5000 / INDEPENDENT；FPS center" if "data.query_mode" in arm["changes"]
                else "random5000 / SAME；FPS center")
    return [
        f"V6·{arm['id']} {arm['config'][:-5]}", SPLIT, model, change_text(arm), protocol, 5000,
        m["field_casebalanced"]["r2"], m["field"]["r2"], m["aggregate"]["r2_casemean"], m["aggregate"]["r2_casemed"],
        m["aggregate"]["r2_casep10"], f"{m['aggregate']['r2_negative_cases']}/{m['aggregate']['n_cases']}",
        m["field"]["rmse"], m["field"]["mae"], m["field"]["nrmse_range"], m["aggregate"]["nrmse_casemean"],
        m["field"].get("nmae_range"), m["aggregate"].get("nmae_casemean"),
        m["regional_field"]["high_wss"]["r2"], m["calibration"]["top10_pred_true_ratio"], m["hotspot"]["top10_iou_casemean"],
        n["field_casebalanced"]["r2"], n["field"]["r2"], n["aggregate"]["r2_casemean"], n["aggregate"]["r2_casemed"],
        n["aggregate"]["r2_casep10"], n["field_casebalanced"]["mae"], n["field_casebalanced"]["rmse"],
        n["field"]["nrmse_range"], n["aggregate"]["nrmse_casemean"], n["field"].get("nmae_range"),
        n["aggregate"].get("nmae_casemean"),
        m["hotspot"]["spearman_all_casemean"], m["hotspot"]["spearman_high_wss_casemean"],
        n["calibration"]["top10_pred_true_ratio"], n["calibration"]["p99_pred_true_ratio"],
        f"v6_singleframe/{arm['config'][:-5]}",
        m["field_casebalanced"].get("r2_linear_fit"), m["field_casebalanced"].get("linear_fit_slope"),
        m["field_casebalanced"].get("linear_fit_intercept"),
        m["field"].get("r2_linear_fit"), m["field"].get("linear_fit_slope"), m["field"].get("linear_fit_intercept"),
        n["field_casebalanced"].get("r2_linear_fit"), n["field_casebalanced"].get("linear_fit_slope"),
        n["field_casebalanced"].get("linear_fit_intercept"),
        n["field"].get("r2_linear_fit"), n["field"].get("linear_fit_slope"), n["field"].get("linear_fit_intercept"),
    ]


def main() -> None:
    matrix = json.loads(MATRIX.read_text(encoding="utf-8"))
    arms = []
    for arm in matrix["arms"]:
        m = load_arm(arm["run_name"])
        if m is None:
            raise SystemExit(f"missing evaluation for {arm['id']} ({arm['run_name']}); run the queue first")
        arms.append((arm, m))
    base = seeds_of(BASE_ARM[1])
    if not base:
        raise SystemExit("R4 baseline runs missing")
    bp = np.array([dig(m, PHYS) for m in base.values()], float)
    bn = np.array([dig(m, NORM) for m in base.values()], float)
    n_base = len(bp)
    sd_p, sd_n = float(bp.std(ddof=1)), float(bn.std(ddof=1))
    EXP.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(BOOK, BACKUP)
    ids = [f"v6_singleframe/{arm['config'][:-5]}" for arm, _ in arms]
    label_prefixes = [f"V6·{arm['id']} " for arm, _ in arms]

    with tempfile.TemporaryDirectory(prefix="wss_v6_xlsx_") as tmp:
        cand = Path(tmp) / BOOK.name
        shutil.copy2(BOOK, cand)
        wb = load_workbook(cand, data_only=False)
        if wb.sheetnames != EXPECTED_SHEETS:
            raise RuntimeError(f"unexpected sheets {wb.sheetnames}")
        # ---- 实验矩阵总览
        ov = wb["实验矩阵总览"]
        stale = [c.row for c in ov["AK"] if c.value in set(ids)] + [c.row for c in ov["A"] if c.value == OVERVIEW_TITLE]
        if stale:
            unmerge_from(ov, min(stale))
        for r in sorted(set(stale), reverse=True):
            ov.delete_rows(r, 1)
        base_ov = ov.max_row
        r = ov.max_row + 1
        ov.merge_cells(start_row=r, start_column=1, end_row=r, end_column=20)
        cell = ov.cell(r, 1, OVERVIEW_TITLE)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ov.row_dimensions[r].height = 34
        overview_rows = {}
        for arm, m in arms:
            r = ov.max_row + 1
            copy_row_style(ov, base_ov, r, 49)
            for c, v in enumerate(overview_values(arm, m), 1):
                ov.cell(r, c, v)
            overview_rows[arm["id"]] = r
        # ---- 教师汇报视图 ⅩⅢ
        te = wb["教师汇报视图"]
        stale = [c.row for c in te["B"]
                 if isinstance(c.value, str) and any(c.value.startswith(p) for p in label_prefixes)]
        stale += [c.row for c in te["A"] if isinstance(c.value, str) and c.value.startswith(TEACHER_TITLE_PREFIX)]
        if stale:
            unmerge_from(te, min(stale))
        stale += [c.row for c in te["A"] if c.value == STAGE and te.cell(c.row, 2).value is None]
        for r in sorted(set(stale), reverse=True):
            te.delete_rows(r, 1)
        base_te = te.max_row
        header_src = max(c.row for c in te["A"] if isinstance(c.value, str) and c.value[:1] in "ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩⅪⅫ")
        r = te.max_row + 1
        copy_row_style(te, header_src, r, 19)
        te.merge_cells(start_row=r, start_column=1, end_row=r, end_column=15)
        te.cell(r, 1, TEACHER_TITLE)
        for arm, m in arms:
            n = m["normalized"]
            r = te.max_row + 1
            copy_row_style(te, base_te, r, 19)
            values = [STAGE, f"V6·{arm['id']} {arm['hypothesis']}", LSA2, "random5000 / SAME",
                      m["field_casebalanced"]["r2"], m["aggregate"]["r2_casemean"], m["field"]["rmse"],
                      m["field"].get("nmae_range"), m["aggregate"].get("nmae_casemean"),
                      m["regional_field"]["high_wss"]["r2"], n["field_casebalanced"]["r2"],
                      n["field"].get("nmae_range"), m["hotspot"]["spearman_all_casemean"],
                      m["hotspot"]["top10_iou_casemean"], n["calibration"]["p99_pred_true_ratio"],
                      m["field_casebalanced"].get("r2_linear_fit"), m["field"].get("r2_linear_fit"),
                      m["field"].get("linear_fit_slope"), m["field"].get("linear_fit_intercept")]
            for c, v in enumerate(values, 1):
                te.cell(r, c, v)
        # ---- 汇总对比
        su = wb["汇总对比"]
        old = [c.row for c in su["A"] if c.value == SECTION_TITLE]
        if old:
            unmerge_from(su, old[0])
            delete_section(su, old[0])
        r = su.max_row + 1
        su.merge_cells(start_row=r, start_column=1, end_row=r, end_column=15)
        cell = su.cell(r, 1, SECTION_TITLE)
        cell.font = Font(bold=True, color=WHITE, size=12)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center")
        su.row_dimensions[r].height = 23
        r += 1
        band_note = (f"噪声由 R4 {n_base} 个 seed 现场标定：物理 sd={sd_p:.4f}、归一化 sd={sd_n:.4f}；"
                     f"单 seed 对 {n_base} seed 参照的 95% 带 ±{1.96 * sd_p * math.sqrt(1 + 1 / n_base):.3f}（物理）/"
                     f"±{1.96 * sd_n * math.sqrt(1 + 1 / n_base):.3f}（归一化）。"
                     "本节每臂只有 1 个 seed，按既定合同不判定；Δ 只用于排出补 seed 的优先级。"
                     "判据指标 = 归一化 R²_cb，物理 R²_cb 作一致性检查。")
        su.merge_cells(start_row=r, start_column=1, end_row=r, end_column=15)
        cell = su.cell(r, 1, band_note)
        cell.font = Font(italic=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        su.row_dimensions[r].height = 30
        r += 1
        headers = ["臂", "唯一变化", "物理 R²_cb", "Δ物理", "归一化 R²_cb", "Δ归一化", "逐例均值", "MAE (Pa)",
                   "high-WSS R²", "Δhigh-WSS", "top10 幅值比", "p99 比", "top10 IoU", "回变换", "筛选状态"]
        for c, h in enumerate(headers, 1):
            cell = su.cell(r, c, h)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        su.row_dimensions[r].height = 36
        r += 1
        ref_high = float(np.mean([dig(m, ("regional_field", "high_wss", "r2")) for m in base.values()]))
        values = ["R4 基准（参照）", f"{n_base} seed 均值", float(bp.mean()), None, float(bn.mean()), None,
                  float(np.mean([m["aggregate"]["r2_casemean"] for m in base.values()])),
                  float(np.mean([m["field"]["mae"] for m in base.values()])), ref_high, None,
                  float(np.mean([m["calibration"]["top10_pred_true_ratio"] for m in base.values()])),
                  float(np.mean([m["calibration"]["p99_pred_true_ratio"] for m in base.values()])),
                  float(np.mean([m["hotspot"]["top10_iou_casemean"] for m in base.values()])), "exp(mu)", "参照"]
        for c, v in enumerate(values, 1):
            styled(su.cell(r, c, v), BLUE, {3: "0.0000", 5: "0.0000", 7: "0.0000", 8: "0.000", 9: "0.0000",
                                            11: "0.000", 12: "0.000", 13: "0.000"}.get(c))
        su.row_dimensions[r].height = 40
        r += 1
        for arm, m in arms:
            phys, norm = dig(m, PHYS), dig(m, NORM)
            dp, dn = phys - bp.mean(), norm - bn.mean()
            verdict, band_p, band_n = verdict_for(dp, dn, 1, n_base, sd_p, sd_n)
            verdict = verdict.replace("**", "")
            ratio = dn / band_n if band_n else float("nan")
            promising = ratio >= 2.0 and dp > 0
            fill = GREEN if promising else YELLOW
            values = [f"V6·{arm['id']} {arm['config'][:-5]}", change_text(arm), phys, dp, norm, dn,
                      m["aggregate"]["r2_casemean"], m["field"]["mae"],
                      dig(m, ("regional_field", "high_wss", "r2")),
                      dig(m, ("regional_field", "high_wss", "r2")) - ref_high,
                      m["calibration"]["top10_pred_true_ratio"], m["calibration"]["p99_pred_true_ratio"],
                      m["hotspot"]["top10_iou_casemean"], m.get("back_transform", "exp(mu)"),
                      f"Δ归一化 = {ratio:+.1f}× 单 seed 带；"
                      + ("优先补 3 seed 复核；" if promising else "") + verdict]
            for c, v in enumerate(values, 1):
                fmt = {3: "0.0000", 5: "0.0000", 7: "0.0000", 8: "0.000", 9: "0.0000",
                       11: "0.000", 12: "0.000", 13: "0.000"}.get(c)
                if c in (4, 6, 10):
                    fmt = "+0.0000;-0.0000;0.0000"
                styled(su.cell(r, c, v), fill, fmt)
            su.row_dimensions[r].height = 44
            r += 1
        if hasattr(wb, "calculation"):
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        wb.save(cand)
        wb.close()
        chk = load_workbook(cand, read_only=True, data_only=False)
        got = {v[0] for v in chk["实验矩阵总览"].iter_rows(min_col=37, max_col=37, values_only=True)}
        assert set(ids) <= got, sorted(set(ids) - got)[:3]
        counts = (chk["实验矩阵总览"].max_row, chk["教师汇报视图"].max_row, chk["汇总对比"].max_row)
        chk.close()
        shutil.copy2(cand, BOOK)
    print(json.dumps({"arms_written": len(ids), "overview_rows": overview_rows,
                      "sheet_max_rows": counts, "backup": str(BACKUP)}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
