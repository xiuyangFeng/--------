#!/usr/bin/env python3
"""Update the WSS PointNet workbook with the completed RCR Oracle matrix."""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from training_wss_min.tools.update_s3_rootcause_xlsx import copy_row_style


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = (
    ROOT
    / "training_wss_min/preflight/"
    "rcr_oracle_matrix_20260728_results_analysis.json"
)
SECTION_TITLE = "2026-07-28｜S3 PointNeXt-R + LocalGeoPE · RCR Oracle"
SHEET_TITLE = "RCR Oracle"
REDUNDANT_SHEET = "D2K64三种子确认"
ORDER = (
    "o0_geometry_s1234",
    "o1a_true_rcr_s1234",
    "o2_shuffled_rcr_s1234",
)
LABELS = {
    "o0_geometry_s1234": "RCR-O0 geometry s1234",
    "o1a_true_rcr_s1234": "RCR-O1a true RCR s1234",
    "o2_shuffled_rcr_s1234": "RCR-O2 shuffled RCR s1234",
}
INPUTS = {
    "o0_geometry_s1234": "geometry-only（6D）",
    "o1a_true_rcr_s1234": "geometry + 4出口真实 log(R1/R2/C)（18D）",
    "o2_shuffled_rcr_s1234": "geometry + 分层内打乱 log-RCR（18D）",
}

NAVY = "1F4E78"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
GRAY = "E7E6E6"
RED = "FCE4D6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def raw_metrics(record: dict) -> dict:
    return read_json(
        Path(record["run_dir"]) / "eval/ckpt_best/metrics.json"
    )["test"]


def overview_values(record: dict) -> list:
    raw = raw_metrics(record)
    norm = raw["normalized"]
    experiment_id = record["experiment_id"]
    return [
        LABELS[experiment_id],
        "138/0/36（AG/AAA/ILO mixed；seed=1234）",
        "PointNeXt-R + LocalGeoPE",
        INPUTS[experiment_id],
        "random / SAME；FPS center",
        5000,
        raw["field_casebalanced"]["r2"],
        raw["field"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["aggregate"]["r2_casemed"],
        raw["aggregate"]["r2_casep10"],
        f"{raw['aggregate']['r2_negative_cases']}/{raw['aggregate']['n_cases']}",
        raw["field"]["rmse"],
        raw["field"]["mae"],
        raw["field"]["nrmse_range"],
        raw["aggregate"]["nrmse_casemean"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        raw["calibration"]["top10_pred_true_ratio"],
        raw["hotspot"]["top10_iou_casemean"],
        norm["field_casebalanced"]["r2"],
        norm["field"]["r2"],
        norm["aggregate"]["r2_casemean"],
        norm["aggregate"]["r2_casemed"],
        norm["aggregate"]["r2_casep10"],
        norm["field_casebalanced"]["mae"],
        norm["field_casebalanced"]["rmse"],
        norm["field"]["nrmse_range"],
        norm["aggregate"]["nrmse_casemean"],
        norm["field"].get("nmae_range"),
        norm["aggregate"].get("nmae_casemean"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["spearman_high_wss_casemean"],
        norm["calibration"]["top10_pred_true_ratio"],
        norm["calibration"]["p99_pred_true_ratio"],
        experiment_id,
    ]


def teacher_values(record: dict) -> list:
    raw = raw_metrics(record)
    norm = raw["normalized"]
    experiment_id = record["experiment_id"]
    return [
        "RCR Oracle",
        LABELS[experiment_id],
        "PointNeXt-R + LocalGeoPE",
        "random5000 / SAME",
        raw["field_casebalanced"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["field"]["rmse"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        norm["field_casebalanced"]["r2"],
        norm["field"].get("nmae_range"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["top10_iou_casemean"],
        norm["calibration"]["p99_pred_true_ratio"],
    ]


def fill_row(sheet, row: int, values: list, fill: str | None = None) -> None:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)


def add_title(sheet, row: int, text: str, columns: int = 15) -> int:
    sheet.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=columns
    )
    cell = sheet.cell(row, 1, text)
    cell.font = Font(bold=True, color=WHITE, size=12)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 23
    return row + 1


def add_header(sheet, row: int, values: list[str]) -> int:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    sheet.row_dimensions[row].height = 31
    return row + 1


def build_oracle_sheet(workbook, analysis: dict) -> None:
    if SHEET_TITLE in workbook.sheetnames:
        workbook.remove(workbook[SHEET_TITLE])
    sheet = workbook.create_sheet(SHEET_TITLE, 2)
    records = analysis["records"]
    comparisons = {
        item["comparison"]: item for item in analysis["paired_comparisons"]
    }

    add_title(sheet, 1, "RCR Oracle：真实边界条件信息上限")
    sheet.merge_cells("A2:O2")
    sheet["A2"] = (
        "同一 S3 PointNeXt-R + LocalGeoPE、mixed 138/0/36、seed1234、"
        "400 epoch、train-loss 选模、legacy_vertex。test36 为复用开发集，"
        "以下是工程筛选证据，不是无偏最终测试结论。"
    )
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet["A2"].fill = PatternFill("solid", fgColor=YELLOW)
    sheet.row_dimensions[2].height = 35

    add_header(
        sheet,
        4,
        [
            "臂",
            "输入",
            "维度",
            "best epoch",
            "physical R²_cb",
            "ΔR²_cb vs O0",
            "normalized R²_cb",
            "physical MAE (Pa)",
            "ΔMAE vs O0",
            "high-WSS nRMSE",
            "ΔnRMSE vs O0",
            "top10 IoU",
            "ΔIoU vs O0",
            "top10 幅值比",
            "p99 幅值比",
        ],
    )
    o0 = records["o0_geometry_s1234"]
    pair_by_treatment = {
        item["treatment"]: item
        for item in analysis["paired_comparisons"]
        if item["control"] == "o0_geometry_s1234"
    }
    for row, experiment_id in enumerate(ORDER, 5):
        record = records[experiment_id]
        delta = (
            None
            if experiment_id == "o0_geometry_s1234"
            else pair_by_treatment[experiment_id]
        )
        agg_delta = (
            {} if delta is None else delta["aggregate_treatment_minus_control"]
        )
        extra_delta = (
            {} if delta is None else delta["extra_treatment_minus_control"]
        )
        values = [
            LABELS[experiment_id],
            INPUTS[experiment_id],
            record["input_dim"],
            record["history"]["best_epoch"],
            record["best"]["physical_r2_casebalanced"],
            None if delta is None else agg_delta["physical_r2_casebalanced"],
            record["best"]["normalized_r2_casebalanced"],
            record["best"]["physical_mae"],
            None if delta is None else agg_delta["physical_mae"],
            record["extra"]["high_wss_nrmse_range"],
            None if delta is None else extra_delta["high_wss_nrmse_range"],
            record["best"]["physical_top10_iou"],
            None if delta is None else agg_delta["physical_top10_iou"],
            record["best"]["physical_top10_amplitude_ratio"],
            record["best"]["physical_p99_amplitude_ratio"],
        ]
        fill_row(
            sheet,
            row,
            values,
            GREEN
            if experiment_id == "o1a_true_rcr_s1234"
            else GRAY
            if experiment_id == "o0_geometry_s1234"
            else RED,
        )

    add_title(sheet, 9, "配对证据（treatment − control）")
    add_header(
        sheet,
        10,
        [
            "比较",
            "ΔR²_cb",
            "病例均值 ΔR²",
            "病例均值 ΔR² 95%CI",
            "R² 胜/负",
            "ΔMAE (Pa)",
            "Δhigh-WSS nRMSE",
            "Δtop10 IoU",
            "ΔAG R²",
            "ΔAAA R²",
            "ΔILO R²",
            "Δp99 幅值比",
            "判读",
        ],
    )
    pair_order = (
        "o1a_true_rcr_vs_o0_geometry",
        "o2_shuffled_rcr_vs_o0_geometry",
        "o1a_true_rcr_vs_o2_shuffled_rcr",
    )
    pair_labels = {
        pair_order[0]: "O1a true RCR − O0 geometry",
        pair_order[1]: "O2 shuffled RCR − O0 geometry",
        pair_order[2]: "O1a true RCR − O2 shuffled RCR",
    }
    pair_decisions = {
        pair_order[0]: "真实 RCR 明显改善",
        pair_order[1]: "打乱 RCR 与 O0 基本持平",
        pair_order[2]: "真实信息显著优于等维负对照",
    }
    for row, name in enumerate(pair_order, 11):
        item = comparisons[name]
        agg = item["aggregate_treatment_minus_control"]
        extra = item["extra_treatment_minus_control"]
        domains = item["domain_r2_delta"]
        cases = item["paired_case_stats"]["case_r2"]
        fill_row(
            sheet,
            row,
            [
                pair_labels[name],
                agg["physical_r2_casebalanced"],
                cases["mean_delta"],
                f"[{cases['ci95_low']:+.4f}, {cases['ci95_high']:+.4f}]",
                f"{cases['wins']}/{cases['losses']}",
                agg["physical_mae"],
                extra["high_wss_nrmse_range"],
                agg["physical_top10_iou"],
                domains["AG"],
                domains["AAA"],
                domains["ILO"],
                extra["physical_p99_ratio"],
                pair_decisions[name],
            ],
            GREEN if name != pair_order[1] else YELLOW,
        )

    add_title(sheet, 15, "结论与边界")
    sheet.merge_cells("A16:O16")
    sheet["A16"] = (
        "结论：O1a 同时优于 O0 与 O2，且 O2 在主指标上贴近 O0，"
        "支持“四出口 RCR 含有几何之外的有效 WSS 信息”。下一步应把 RCR 或可部署代理"
        "作为输入方向继续验证；正式晋级前仍需补种子或新协议确认。"
    )
    sheet["A16"].fill = PatternFill("solid", fgColor=GREEN)
    sheet["A16"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[16].height = 38
    sheet.merge_cells("A17:O17")
    sheet["A17"] = (
        "边界：单种子；test36 多轮复用；checkpoint 按 train loss 选择；"
        "真实 RCR 目前不是常规可部署输入；O1b 流量占比仅 58/174 直接覆盖且属于求解器输出，继续暂缓。"
    )
    sheet["A17"].fill = PatternFill("solid", fgColor=YELLOW)
    sheet["A17"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[17].height = 36

    add_title(sheet, 19, "运行与产物")
    provenance = [
        ("Slurm", "11008 preflight；11009_[0-2] 三臂训练 + best/last test36；均 COMPLETED (0:0)"),
        ("分析", "training_wss_min/preflight/rcr_oracle_matrix_20260728_results_analysis.json"),
        ("配置", "training_wss_min/configs/pointnetpp_rcr_oracle_20260728/"),
        ("运行目录", "training_wss_min/runs/pointnetpp_rcr_oracle/outputs/"),
    ]
    for row, (key, value) in enumerate(provenance, 20):
        sheet.cell(row, 1, key).font = Font(bold=True)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=15)
        sheet.cell(row, 2, value)
        for column in range(1, 16):
            sheet.cell(row, column).border = Border(bottom=THIN)
            sheet.cell(row, column).alignment = Alignment(
                vertical="center", wrap_text=True
            )

    for column in range(3, 16):
        for row in range(5, 8):
            sheet.cell(row, column).number_format = "0.0000"
    for row in range(5, 8):
        sheet.cell(row, 3).number_format = "0"
        sheet.cell(row, 4).number_format = "0"
    for row in range(11, 14):
        for column in (2, 3, 6, 7, 8, 9, 10, 11, 12):
            sheet.cell(row, column).number_format = "+0.0000;-0.0000;0.0000"
        sheet.row_dimensions[row].height = 29
        sheet.cell(row, 4).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        sheet.cell(row, 13).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    widths = {
        "A": 30,
        "B": 39,
        "C": 10,
        "D": 24,
        "E": 15,
        "F": 15,
        "G": 17,
        "H": 16,
        "I": 14,
        "J": 17,
        "K": 17,
        "L": 13,
        "M": 30,
        "N": 15,
        "O": 15,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = "A4:O7"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 2
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = "A1:O23"
    sheet.sheet_properties.tabColor = "70AD47"


def main() -> None:
    analysis = read_json(ANALYSIS)
    records = analysis["records"]
    for experiment_id in ORDER:
        if experiment_id not in records:
            raise RuntimeError(f"analysis missing {experiment_id}")
        if records[experiment_id]["integrity"] != "passed":
            raise RuntimeError(f"integrity failed: {experiment_id}")

    with tempfile.TemporaryDirectory(prefix="wss_rcr_oracle_xlsx_") as temp_dir:
        candidate = Path(temp_dir) / BOOK.name
        shutil.copy2(BOOK, candidate)
        workbook = load_workbook(candidate, data_only=False)
        if REDUNDANT_SHEET in workbook.sheetnames:
            workbook.remove(workbook[REDUNDANT_SHEET])

        overall = workbook["实验矩阵总览"]
        existing_rows = [
            cell.row for cell in overall["AK"] if cell.value in set(ORDER)
        ]
        for row in sorted(existing_rows, reverse=True):
            overall.delete_rows(row, 1)
        first_overall = overall.max_row + 1
        for experiment_id in ORDER:
            row = overall.max_row + 1
            copy_row_style(overall, row - 1, row, 37)
            for column, value in enumerate(
                overview_values(records[experiment_id]), 1
            ):
                overall.cell(row, column, value)

        teacher = workbook["教师汇报视图"]
        existing_teacher_rows = [
            cell.row for cell in teacher["B"] if cell.value in set(LABELS.values())
        ]
        for row in sorted(existing_teacher_rows, reverse=True):
            teacher.delete_rows(row, 1)
        first_teacher = teacher.max_row + 1
        for experiment_id in ORDER:
            row = teacher.max_row + 1
            copy_row_style(teacher, row - 1, row, 15)
            for column, value in enumerate(
                teacher_values(records[experiment_id]), 1
            ):
                teacher.cell(row, column, value)

        summary = workbook["汇总对比"]
        title_rows = [
            cell.row for cell in summary["A"] if cell.value == SECTION_TITLE
        ]
        if len(title_rows) > 1:
            raise RuntimeError("duplicate RCR Oracle summary sections")
        if title_rows:
            summary.delete_rows(title_rows[0], 5)
        summary_start = summary.max_row + 1
        row = add_title(summary, summary_start, SECTION_TITLE)
        row = add_header(
            summary,
            row,
            [
                "臂",
                "输入",
                "R²_cb",
                "ΔR²_cb",
                "normalized R²_cb",
                "MAE (Pa)",
                "ΔMAE",
                "high-WSS nRMSE",
                "ΔnRMSE",
                "top10 IoU",
                "ΔIoU",
                "top10 幅值比",
                "p99 幅值比",
                "AG / AAA / ILO R²",
                "判定",
            ],
        )
        pair_by_treatment = {
            item["treatment"]: item
            for item in analysis["paired_comparisons"]
            if item["control"] == "o0_geometry_s1234"
        }
        for experiment_id in ORDER:
            record = records[experiment_id]
            delta = pair_by_treatment.get(experiment_id)
            agg = {} if delta is None else delta["aggregate_treatment_minus_control"]
            extra = {} if delta is None else delta["extra_treatment_minus_control"]
            domains = record["domains"]
            values = [
                LABELS[experiment_id],
                INPUTS[experiment_id],
                record["best"]["physical_r2_casebalanced"],
                None if delta is None else agg["physical_r2_casebalanced"],
                record["best"]["normalized_r2_casebalanced"],
                record["best"]["physical_mae"],
                None if delta is None else agg["physical_mae"],
                record["extra"]["high_wss_nrmse_range"],
                None if delta is None else extra["high_wss_nrmse_range"],
                record["best"]["physical_top10_iou"],
                None if delta is None else agg["physical_top10_iou"],
                record["best"]["physical_top10_amplitude_ratio"],
                record["best"]["physical_p99_amplitude_ratio"],
                f"{domains['AG']:.3f} / {domains['AAA']:.3f} / {domains['ILO']:.3f}",
                (
                    "几何基准"
                    if experiment_id == "o0_geometry_s1234"
                    else "真实 RCR 有效"
                    if experiment_id == "o1a_true_rcr_s1234"
                    else "负对照≈O0"
                ),
            ]
            fill_row(
                summary,
                row,
                values,
                GREEN
                if experiment_id == "o1a_true_rcr_s1234"
                else GRAY
                if experiment_id == "o0_geometry_s1234"
                else YELLOW,
            )
            row += 1

        build_oracle_sheet(workbook, analysis)
        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(candidate)
        workbook.close()

        check = load_workbook(candidate, read_only=True, data_only=False)
        saved_ids = {
            values[0]
            for values in check["实验矩阵总览"].iter_rows(
                min_col=37, max_col=37, values_only=True
            )
        }
        saved_labels = {
            values[0]
            for values in check["教师汇报视图"].iter_rows(
                min_col=2, max_col=2, values_only=True
            )
        }
        saved_titles = {
            values[0]
            for values in check["汇总对比"].iter_rows(
                min_col=1, max_col=1, values_only=True
            )
        }
        if not set(ORDER).issubset(saved_ids):
            raise RuntimeError("overview readback missing RCR Oracle rows")
        if not set(LABELS.values()).issubset(saved_labels):
            raise RuntimeError("teacher readback missing RCR Oracle rows")
        if SECTION_TITLE not in saved_titles or SHEET_TITLE not in check.sheetnames:
            raise RuntimeError("RCR Oracle summary/sheet readback failed")
        if REDUNDANT_SHEET in check.sheetnames:
            raise RuntimeError("redundant D2K64 confirmation sheet was not removed")
        if check[SHEET_TITLE]["E6"].value != records[
            "o1a_true_rcr_s1234"
        ]["best"]["physical_r2_casebalanced"]:
            raise RuntimeError("RCR Oracle numeric readback mismatch")
        check.close()

        render_dir = Path(temp_dir) / "render"
        render_dir.mkdir()
        subprocess.run(
            [
                "/usr/bin/libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(render_dir),
                str(candidate),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        if not (render_dir / candidate.with_suffix(".pdf").name).is_file():
            raise RuntimeError("LibreOffice render check failed")
        shutil.copy2(candidate, BOOK)

    print(
        json.dumps(
            {
                "status": "updated_and_rendered",
                "workbook": str(BOOK),
                "overall_rows": [first_overall, first_overall + len(ORDER) - 1],
                "teacher_rows": [first_teacher, first_teacher + len(ORDER) - 1],
                "summary_start_row": summary_start,
                "oracle_sheet": SHEET_TITLE,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
