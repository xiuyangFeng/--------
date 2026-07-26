#!/usr/bin/env python3
"""Apply the presentation layout for the WSS PointNet result workbook.

The workbook intentionally keeps two different reading modes:
* ``教师汇报视图``: an at-a-glance, one-row-per-experiment scorecard.
* ``实验矩阵总览``: the complete numeric matrix, optimized for filtering and audit.

The obsolete provenance/comment columns are removed once, by their headers, so
future result-updaters cannot silently reintroduce them.
"""

from __future__ import annotations

import argparse
import shutil
import tempfile
import zipfile
from pathlib import Path

import uno
from openpyxl import load_workbook
from com.sun.star.awt import FontWeight
from com.sun.star.lang import Locale
from com.sun.star.table import CellRangeAddress


REPO = Path(__file__).resolve().parents[2]
WORKBOOK = REPO / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"

NAVY = 0x17324D
BLUE = 0x285E8E
TEAL = 0x1B7F8C
PALE_BLUE = 0xEDF4FA
PALE_TEAL = 0xE8F4F5
PALE_GOLD = 0xFFF2CC
SECTION = 0xD9EAF7
WHITE = 0xFFFFFF
INK = 0x17202A
MUTED = 0x5D6D7E


def prop(name: str, value):
    item = uno.createUnoStruct("com.sun.star.beans.PropertyValue")
    item.Name = name
    item.Value = value
    return item


def last_used_row(sheet, col: int = 0, end: int = 400) -> int:
    last = 0
    for row in range(end):
        if sheet.getCellByPosition(col, row).String.strip():
            last = row
    return last


def set_range_style(rng, *, background=None, color=None, weight=None, align=None, wrap=None) -> None:
    if background is not None:
        rng.CellBackColor = background
    if color is not None:
        rng.CharColor = color
    if weight is not None:
        rng.CharWeight = weight
    if align is not None:
        rng.HoriJustify = align
    if wrap is not None:
        rng.IsTextWrapped = wrap
    rng.VertJustify = 2  # CENTER


def number_format(doc, pattern: str) -> int:
    locale = Locale(Language="en", Country="US")
    key = doc.NumberFormats.queryKey(pattern, locale, True)
    return key if key != -1 else doc.NumberFormats.addNew(pattern, locale)


def remove_obsolete_columns(teacher, overview) -> None:
    # P:U in the teacher scorecard were execution/provenance notes; those now
    # live in the linked markdown and JSON audit rather than the presentation.
    if teacher.getCellByPosition(15, 2).String == "状态 / Job":
        teacher.Columns.removeByIndex(15, 6)
    # AK:AP were the same notes on the full matrix.
    if overview.getCellByPosition(36, 1).String == "状态 / Job":
        overview.Columns.removeByIndex(36, 6)


def set_print_area(sheet, end_col: int, end_row: int) -> None:
    sheet.setPrintAreas((CellRangeAddress(
        Sheet=sheet.RangeAddress.Sheet,
        StartColumn=0,
        EndColumn=end_col,
        StartRow=0,
        EndRow=end_row,
    ),))
    # Keep the metric labels visible on continuation pages.  The presentation
    # sheet repeats its title plus two header rows; the audit sheet repeats its
    # two matrix header rows on each horizontal print segment.
    header_end = 2 if sheet.Name == "教师汇报视图" else 1
    sheet.setPrintTitleRows(True)
    sheet.setTitleRows(CellRangeAddress(
        Sheet=sheet.RangeAddress.Sheet,
        StartColumn=0,
        EndColumn=end_col,
        StartRow=0,
        EndRow=header_end,
    ))


def style_teacher(doc, sheet) -> None:
    last = last_used_row(sheet)
    four_decimals = number_format(doc, "0.0000")
    three_decimals = number_format(doc, "0.000")

    # Title and the two-level header separate setup from outcome metrics.
    set_range_style(sheet.getCellRangeByPosition(0, 0, 14, 0), background=NAVY, color=WHITE,
                    weight=FontWeight.BOLD, align=2, wrap=True)
    sheet.Rows.getByIndex(0).Height = 780
    set_range_style(sheet.getCellRangeByPosition(0, 1, 3, 1), background=TEAL, color=WHITE,
                    weight=FontWeight.BOLD, align=2, wrap=True)
    set_range_style(sheet.getCellRangeByPosition(4, 1, 14, 1), background=BLUE, color=WHITE,
                    weight=FontWeight.BOLD, align=2, wrap=True)
    set_range_style(sheet.getCellRangeByPosition(0, 2, 14, 2), background=NAVY, color=WHITE,
                    weight=FontWeight.BOLD, align=2, wrap=True)
    for col, header in {
        7: "NMAE\npooled (%)",
        8: "NMAE\ncase-mean (%)",
        11: "NMAE\npooled (%)",
    }.items():
        sheet.getCellByPosition(col, 2).String = header
    sheet.Rows.getByIndex(1).Height = 620
    sheet.Rows.getByIndex(2).Height = 880

    for row in range(3, last + 1):
        stage = sheet.getCellByPosition(0, row).String.strip()
        experiment = sheet.getCellByPosition(1, row).String.strip()
        if stage and not experiment:
            set_range_style(sheet.getCellRangeByPosition(0, row, 14, row), background=SECTION,
                            color=NAVY, weight=FontWeight.BOLD, align=1, wrap=True)
            sheet.Rows.getByIndex(row).Height = 520
            continue
        if not experiment:
            continue
        fill = PALE_BLUE if row % 2 else WHITE
        if "★" in experiment:
            fill = PALE_GOLD
        elif experiment == "q2v_to_q1v_same_finetune":
            fill = PALE_TEAL
        set_range_style(sheet.getCellRangeByPosition(0, row, 3, row), background=fill,
                        color=INK, align=1, wrap=True)
        set_range_style(sheet.getCellRangeByPosition(4, row, 14, row), background=fill,
                        color=INK, align=2, wrap=False)
        sheet.Rows.getByIndex(row).Height = 560

    # Text is deliberately roomy; numeric metrics are fixed to comparable precision.
    for col, width in enumerate((2600, 4900, 2800, 6200, 1500, 1700, 1800, 1700, 1800,
                                 1750, 1700, 1800, 1700, 1500, 1800)):
        sheet.Columns.getByIndex(col).Width = width
    sheet.getCellRangeByPosition(4, 3, 14, last).NumberFormat = four_decimals
    sheet.getCellRangeByPosition(6, 3, 6, last).NumberFormat = three_decimals
    set_print_area(sheet, 14, last)


def style_overview(doc, sheet) -> None:
    last = last_used_row(sheet)
    four_decimals = number_format(doc, "0.0000")
    three_decimals = number_format(doc, "0.000")
    integer = number_format(doc, "0")

    # Full matrix: headers stay visually strong while all metric columns remain
    # machine-readable numeric cells for later filtering or comparison.
    set_range_style(sheet.getCellRangeByPosition(0, 0, 35, 0), background=NAVY, color=WHITE,
                    weight=FontWeight.BOLD, align=2, wrap=True)
    set_range_style(sheet.getCellRangeByPosition(0, 1, 35, 1), background=BLUE, color=WHITE,
                    weight=FontWeight.BOLD, align=2, wrap=True)
    for col, header in {
        14: "物理 range-NRMSE\npooled (%)",
        15: "物理 range-NRMSE\ncase-mean (%)",
        16: "物理 NMAE\npooled (%)",
        17: "物理 NMAE\ncase-mean (%)",
        28: "归一化 range-NRMSE\npooled (%)",
        29: "归一化 range-NRMSE\ncase-mean (%)",
        30: "归一化 NMAE\npooled (%)",
        31: "归一化 NMAE\ncase-mean (%)",
    }.items():
        sheet.getCellByPosition(col, 1).String = header
    sheet.Rows.getByIndex(0).Height = 720
    sheet.Rows.getByIndex(1).Height = 1050

    candidates = {"P2V-10474", "Q1V-10476"}
    for row in range(2, last + 1):
        experiment = sheet.getCellByPosition(0, row).String.strip()
        split = sheet.getCellByPosition(1, row).String.strip()
        if experiment and not split:
            set_range_style(sheet.getCellRangeByPosition(0, row, 35, row), background=SECTION,
                            color=NAVY, weight=FontWeight.BOLD, align=1, wrap=True)
            sheet.Rows.getByIndex(row).Height = 500
            continue
        if not experiment:
            continue
        fill = PALE_BLUE if row % 2 else WHITE
        if experiment in candidates:
            fill = PALE_GOLD
        elif experiment == "q2v_to_q1v_same_finetune":
            fill = PALE_TEAL
        set_range_style(sheet.getCellRangeByPosition(0, row, 5, row), background=fill,
                        color=INK, align=1, wrap=True)
        set_range_style(sheet.getCellRangeByPosition(6, row, 35, row), background=fill,
                        color=INK, align=2, wrap=False)
        sheet.Rows.getByIndex(row).Height = 470

    widths = (4300, 2600, 2600, 6400, 5200, 1500) + (1550,) * 30
    for col, width in enumerate(widths):
        sheet.Columns.getByIndex(col).Width = width
    sheet.getCellRangeByPosition(6, 2, 35, last).NumberFormat = four_decimals
    sheet.getCellRangeByPosition(11, 2, 11, last).NumberFormat = integer  # negative-case count
    sheet.getCellRangeByPosition(12, 2, 13, last).NumberFormat = three_decimals  # RMSE / MAE
    set_print_area(sheet, 35, last)


def remove_excel_currency_locale_markers(path: Path) -> None:
    """Write plain numeric formats that do not expose Excel's ``[$-409]`` tag.

    LibreOffice's number-format service serializes an en-US locale as
    ``[$-409]0.0000``.  It is a locale marker rather than a currency in Calc,
    but some spreadsheet viewers display the leading ``$`` confusingly.  The
    final xlsx pass writes explicit, locale-free numeric formats for metrics.
    """
    book = load_workbook(path)
    teacher = book["教师汇报视图"]
    for row in range(4, teacher.max_row + 1):
        if not teacher.cell(row, 2).value:
            continue
        for col in range(5, 16):
            cell = teacher.cell(row, col)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                if col in (8, 9, 12):
                    cell.number_format = "0.00%"
                else:
                    cell.number_format = "0.000" if col == 7 else "0.0000"
    overview = book["实验矩阵总览"]
    for row in range(3, overview.max_row + 1):
        if not overview.cell(row, 1).value:
            continue
        for col in range(7, 37):
            cell = overview.cell(row, col)
            if not isinstance(cell.value, (int, float)) or isinstance(cell.value, bool):
                continue
            if col in (15, 16, 17, 18, 29, 30, 31, 32):
                cell.number_format = "0.00%"
            elif col in (13, 14):
                cell.number_format = "0.000"
            elif col == 12:
                cell.number_format = "0"
            else:
                cell.number_format = "0.0000"
    book.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument(
        "--uno-url", default="uno:socket,host=localhost,port=2085;urp;StarOffice.ComponentContext"
    )
    args = parser.parse_args()
    temp_dir = Path(tempfile.mkdtemp(prefix="pointnet_xlsx_style_"))
    temp_book = temp_dir / args.workbook.name
    shutil.copy2(args.workbook, temp_book)
    local = uno.getComponentContext()
    resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
    ctx = resolver.resolve(args.uno_url)
    desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
    doc = desktop.loadComponentFromURL(
        uno.systemPathToFileUrl(str(temp_book.resolve())), "_blank", 0,
        (prop("Hidden", True), prop("ReadOnly", False)),
    )
    try:
        teacher = doc.Sheets.getByName("教师汇报视图")
        overview = doc.Sheets.getByName("实验矩阵总览")
        remove_obsolete_columns(teacher, overview)
        style_teacher(doc, teacher)
        style_overview(doc, overview)
        styles = doc.StyleFamilies.getByName("PageStyles")
        for sheet, pages_x, pages_y in ((teacher, 1, 2), (overview, 3, 1)):
            page = styles.getByName(sheet.PageStyle)
            page.IsLandscape = True
            page.ScaleToPagesX = pages_x
            page.ScaleToPagesY = pages_y
            page.LeftMargin = 900
            page.RightMargin = 900
            page.TopMargin = 800
            page.BottomMargin = 800
        doc.store()
    finally:
        doc.close(True)
    if not zipfile.is_zipfile(temp_book):
        raise RuntimeError("formatted workbook is not a valid xlsx zip container")
    remove_excel_currency_locale_markers(temp_book)
    if not zipfile.is_zipfile(temp_book):
        raise RuntimeError("locale-format repair produced an invalid xlsx container")
    shutil.copy2(temp_book, args.workbook)
    shutil.rmtree(temp_dir)
    print({"status": "styled", "workbook": str(args.workbook), "teacher_columns": "A:O", "overview_columns": "A:AJ"})


if __name__ == "__main__":
    main()
