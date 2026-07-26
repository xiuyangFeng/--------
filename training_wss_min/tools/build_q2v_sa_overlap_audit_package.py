#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Create a tidy, reproducible Q2V SA overlap-audit package.

Input traces are exact no-training PointNet++ evaluation traces.  Ball-query
statistics use the recorded ``torch_cluster.radius(..., max_num_neighbors=16)``
assignments.  KNN statistics are geometry-only counterfactuals over the exact
same SA1 support points and FPS centres.
"""

from __future__ import annotations

import argparse
import csv
import json
import shutil
from collections import defaultdict
from pathlib import Path

import numpy as np
from scipy.spatial import cKDTree


PROJECT = Path(__file__).resolve().parents[2]
ROOT = PROJECT / "例子/06_PointNet++_SA三层采样与分组"
PACKAGE = ROOT / "汇总_重叠审计_Q2V-10477_2026-07-19"
TRACE_ROOT = PACKAGE / "_traces"
DEFAULT_OUTPUT = PACKAGE

BALL_SETTINGS = (
    ("Ball 0.6×", 0.6, "ball_r60"),
    ("Ball 0.8×", 0.8, "ball_r80"),
    ("Ball 1.0× (baseline)", 1.0, "ball_r100"),
    ("Ball 1.2×", 1.2, "ball_r120"),
    ("Ball 1.5×", 1.5, "ball_r150"),
)
KNN_VALUES = (6, 8, 10, 16)


def load_stage(trace_root: Path, stage: int) -> dict:
    with (trace_root / "sa_centers_marked.csv").open(newline="", encoding="utf-8") as handle:
        marked = list(csv.DictReader(handle))
    point_xyz = {
        int(row["full_wall_index"]): np.asarray(
            [float(row["x_mm"]), float(row["y_mm"]), float(row["z_mm"])], dtype=float
        )
        for row in marked
    }
    centres: dict[int, np.ndarray] = {}
    with (trace_root / "sa_centers_by_stage.csv").open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            if int(row["sa_stage"]) == stage:
                centres[int(row["center_id"])] = np.asarray(
                    [float(row["x_mm"]), float(row["y_mm"]), float(row["z_mm"])], dtype=float
                )
    members: dict[int, set[int]] = defaultdict(set)
    with (trace_root / f"sa{stage}_assignments.csv").open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            members[int(row["center_id"])].add(int(row["source_full_wall_index"]))
    manifest = json.loads((trace_root / "manifest_sa3.json").read_text(encoding="utf-8"))
    stage_info = manifest["stages"][stage - 1]
    return {
        "trace_root": trace_root,
        "stage": stage,
        "point_xyz": point_xyz,
        "full_xyz": np.asarray(list(point_xyz.values()), dtype=float),
        "centres": centres,
        "members": members,
        "radius_mm": float(stage_info["radius_mm"] if "radius_mm" in stage_info else stage_info["radius_mm_for_this_case"]),
        "radius_norm": float(stage_info["radius_norm"] if "radius_norm" in stage_info else stage_info["ball_query_radius_norm"]),
    }


def nearest_pairs(data: dict) -> list[tuple[int, int]]:
    ids = np.asarray(sorted(data["centres"]), dtype=np.int64)
    xyz = np.asarray([data["centres"][int(key)] for key in ids], dtype=float)
    _, neighbour = cKDTree(xyz).query(xyz, k=2)
    return sorted({tuple(sorted((int(ids[row]), int(ids[index]))))
                   for row, index in enumerate(neighbour[:, 1])})


def global_nearest_pair(data: dict) -> tuple[int, int]:
    ids = np.asarray(sorted(data["centres"]), dtype=np.int64)
    xyz = np.asarray([data["centres"][int(key)] for key in ids], dtype=float)
    distances, neighbours = cKDTree(xyz).query(xyz, k=2)
    row = int(np.argmin(distances[:, 1]))
    return tuple(sorted((int(ids[row]), int(ids[neighbours[row, 1]]))))


def radius_reference_pair(ball_data: dict[float, dict]) -> tuple[int, int]:
    """Choose one genuine nearest-neighbour relation that makes radius effects visible.

    The absolute closest pair happens to have 0 shared retained members at all
    five radii, so it is useful for the closest-pair audit but not for a radius
    comparison.  This selector stays within nearest-neighbour pairs and picks
    a pair whose five-radius mean is closest to 1/3; ties prefer more visible
    radius-dependent change.
    """
    reference = ball_data[1.0]
    candidates = []
    for pair in nearest_pairs(reference):
        rates = [pair_detail(data, data["members"], pair)["overlap_rate"]
                 for data in ball_data.values()]
        candidates.append((abs(float(np.mean(rates)) - 1 / 3), -float(np.ptp(rates)), pair))
    return min(candidates)[2]


def knn_members(data: dict, k: int) -> dict[int, set[int]]:
    source_ids = np.asarray(sorted(data["point_xyz"]), dtype=np.int64)
    source_xyz = np.asarray([data["point_xyz"][int(key)] for key in source_ids], dtype=float)
    centre_ids = np.asarray(sorted(data["centres"]), dtype=np.int64)
    centre_xyz = np.asarray([data["centres"][int(key)] for key in centre_ids], dtype=float)
    _, indices = cKDTree(source_xyz).query(centre_xyz, k=k)
    indices = np.asarray(indices).reshape((len(centre_ids), k))
    return {
        int(centre_id): {int(source_ids[index]) for index in indices[row]}
        for row, centre_id in enumerate(centre_ids)
    }


def pair_detail(data: dict, members: dict[int, set[int]], pair: tuple[int, int]) -> dict:
    a, b = pair
    first, second = members[a], members[b]
    shared = first & second
    smaller = min(len(first), len(second))
    return {
        "center_a_id": a,
        "center_b_id": b,
        "center_distance_mm": float(np.linalg.norm(data["centres"][a] - data["centres"][b])),
        "group_a_points": len(first),
        "group_b_points": len(second),
        "shared_points": len(shared),
        "overlap_rate": len(shared) / smaller,
        "jaccard": len(shared) / len(first | second),
        "only_a": first - second,
        "only_b": second - first,
        "shared": shared,
    }


def summarise(label: str, grouping: str, data: dict, members: dict[int, set[int]], *, stage: int = 1,
              multiplier: float | None = None, k: int | None = None) -> tuple[dict, list[dict]]:
    pairs = nearest_pairs(data)
    details = [pair_detail(data, members, pair) for pair in pairs]
    rates = np.asarray([row["overlap_rate"] for row in details], dtype=float)
    row = {
        "method": label,
        "grouping": grouping,
        "stage": stage,
        "radius_multiplier": multiplier,
        "knn_k": k,
        "nearest_pair_count": len(details),
        "median_overlap": float(np.median(rates)),
        "p10_overlap": float(np.quantile(rates, 0.10)),
        "fraction_le_1_4": float(np.mean(rates <= 1 / 4)),
        "fraction_le_1_3": float(np.mean(rates <= 1 / 3)),
        "fraction_le_1_2": float(np.mean(rates <= 1 / 2)),
        "fraction_exact_100": float(np.mean(rates == 1.0)),
        "count_exact_100": int(np.sum(rates == 1.0)),
    }
    detailed_rows = [{
        "method": label, "grouping": grouping, "stage": stage,
        "radius_multiplier": multiplier if multiplier is not None else "",
        "knn_k": k if k is not None else "", **detail,
    } for detail in details]
    return row, detailed_rows


def _points(data: dict, ids: set[int]) -> np.ndarray:
    return np.asarray([data["point_xyz"][key] for key in sorted(ids)], dtype=float).reshape((-1, 3))


def draw_local(ax, data: dict, members: dict[int, set[int]], pair: tuple[int, int], *, title: str,
               show_radius: bool) -> None:
    from matplotlib.patches import Circle

    detail = pair_detail(data, members, pair)
    ca, cb = data["centres"][pair[0]], data["centres"][pair[1]]
    midpoint = (ca + cb) / 2
    scale = max((data["radius_mm"] * 1.45 if show_radius else 14.0), detail["center_distance_mm"] * 1.40)
    context = data["full_xyz"][np.linalg.norm(data["full_xyz"] - midpoint, axis=1) <= scale]
    ax.scatter(context[:, 0], context[:, 2], c="#D0D0D0", s=7.5, alpha=0.38, linewidths=0,
               label="local support")
    for key, colour, label, size in (
        ("only_a", "#1976D2", "A only", 52),
        ("only_b", "#F57C00", "B only", 52),
        ("shared", "#8E24AA", "shared", 64),
    ):
        xyz = _points(data, detail[key])
        if len(xyz):
            ax.scatter(xyz[:, 0], xyz[:, 2], c=colour, s=size, edgecolors="white", linewidths=0.42,
                       label=f"{label} ({len(xyz)})", zorder=3)
    ax.scatter(ca[0], ca[2], marker="*", s=200, c="#0D47A1", edgecolors="black", linewidths=0.52,
               label=f"center A ({pair[0]})", zorder=5)
    ax.scatter(cb[0], cb[2], marker="*", s=200, c="#E65100", edgecolors="black", linewidths=0.52,
               label=f"center B ({pair[1]})", zorder=5)
    ax.plot([ca[0], cb[0]], [ca[2], cb[2]], c="#555555", lw=0.9, ls="--", zorder=2)
    if show_radius:
        ax.add_patch(Circle((ca[0], ca[2]), data["radius_mm"], fill=False, ls="--", lw=1.05, ec="#0D47A1"))
        ax.add_patch(Circle((cb[0], cb[2]), data["radius_mm"], fill=False, ls="--", lw=1.05, ec="#E65100"))
    half = max(float(np.ptp(context[:, (0, 2)], axis=0).max()) / 2, 7.0)
    ax.set_xlim(midpoint[0] - half, midpoint[0] + half)
    ax.set_ylim(midpoint[2] - half, midpoint[2] + half)
    ax.set_aspect("equal")
    ax.set_xlabel("X (mm)")
    ax.set_ylabel("Z (mm)")
    rule = "≤1/3" if detail["overlap_rate"] <= 1 / 3 else ">1/3"
    radius_text = f"; r={data['radius_mm']:.1f} mm" if show_radius else "; KNN"
    ax.set_title(
        f"{title}\nshared={detail['shared_points']}/{min(detail['group_a_points'], detail['group_b_points'])} "
        f"= {detail['overlap_rate']:.1%} ({rule}); d={detail['center_distance_mm']:.2f} mm{radius_text}",
        fontsize=9.0,
    )


def write_figures(output: Path, ball_data: dict[float, dict], knn_data: dict[int, dict[int, set[int]]],
                  method_rows: list[dict], stage_rows: list[dict]) -> list[dict]:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import PercentFormatter

    figures: list[dict] = []
    baseline = ball_data[1.0]
    # Chart contract: each SA layer's actual closest pair, showing that stage level changes the source set.
    fig, axes = plt.subplots(1, 3, figsize=(17, 5.8), constrained_layout=True)
    for stage, ax in enumerate(axes, start=1):
        data = load_stage(baseline["trace_root"], stage)
        draw_local(ax, data, data["members"], global_nearest_pair(data), title=f"Q2V baseline | SA{stage} closest pair",
                   show_radius=True)
        if stage == 1:
            ax.legend(loc="upper left", fontsize=6.8, frameon=True)
    fig.suptitle("Q2V baseline ball query: local closest-centre group overlap by SA stage", fontsize=14)
    path = output / "01_SA1_SA2_SA3_最近center局部重叠.png"
    fig.savefig(path, dpi=240)
    plt.close(fig)
    figures.append({"file": path.name, "purpose": "SA1/SA2/SA3 真实 ball-query 最近 center 局部对照"})

    # Ball radius metrics: two small, readable comparisons rather than a dense table plot.
    ball_rows = [row for row in method_rows if row["grouping"] == "ball"]
    fig, axes = plt.subplots(1, 2, figsize=(13.8, 5.2), constrained_layout=True)
    xs = np.arange(len(ball_rows))
    labels = [f"{row['radius_multiplier']:.1f}×" for row in ball_rows]
    axes[0].plot(xs, [row["median_overlap"] for row in ball_rows], marker="o", color="#285E8E", label="median")
    axes[0].plot(xs, [row["p10_overlap"] for row in ball_rows], marker="o", color="#8E24AA", label="P10")
    axes[0].axhline(1 / 3, color="#555555", ls="--", lw=1, label="teacher 1/3")
    axes[0].set_xticks(xs)
    axes[0].set_xticklabels(labels)
    axes[0].set_ylim(0, 1)
    axes[0].yaxis.set_major_formatter(PercentFormatter(1.0, decimals=0))
    axes[0].set_title("SA1 nearest-pair overlap: distribution position")
    axes[0].set_ylabel("overlap rate")
    axes[0].legend(fontsize=8)
    width = 0.22
    for offset, key, colour, label in ((-width, "fraction_le_1_4", "#1976D2", "≤1/4"),
                                       (0, "fraction_le_1_3", "#F57C00", "≤1/3"),
                                       (width, "fraction_le_1_2", "#8E24AA", "≤1/2")):
        axes[1].bar(xs + offset, [row[key] for row in ball_rows], width=width, color=colour, label=label)
    axes[1].set_xticks(xs)
    axes[1].set_xticklabels(labels)
    axes[1].set_ylim(0, 1)
    axes[1].yaxis.set_major_formatter(PercentFormatter(1.0, decimals=0))
    axes[1].set_title("SA1 nearest pairs passing overlap thresholds")
    axes[1].set_ylabel("share of nearest pairs")
    axes[1].legend(fontsize=8)
    path = output / "02_SA1_ball半径_最近邻重叠统计对比.png"
    fig.savefig(path, dpi=240)
    plt.close(fig)
    figures.append({"file": path.name, "purpose": "Ball 半径 0.6×–1.5× 的全部最近邻 pair 分布统计"})

    # A genuine nearest-neighbour relation, selected to reveal the radius effect.
    pair = radius_reference_pair(ball_data)
    fig, axes = plt.subplots(1, len(BALL_SETTINGS), figsize=(24, 5.2), constrained_layout=True)
    for ax, (label, multiplier, _) in zip(axes, BALL_SETTINGS):
        data = ball_data[multiplier]
        draw_local(ax, data, data["members"], pair, title=f"{label}", show_radius=True)
    axes[0].legend(loc="upper left", fontsize=6.5, frameon=True)
    fig.suptitle(
        "SA1 same nearest-neighbour pair: exact ball-query radius counterfactual "
        "(pair selected with five-radius mean closest to 1/3)", fontsize=14,
    )
    path = output / "03_SA1_ball半径_同一最近center局部比较.png"
    fig.savefig(path, dpi=240)
    plt.close(fig)
    figures.append({"file": path.name, "purpose": "同一最近邻 SA1 center 对在五种 ball 半径下的局部成员比较；展示pair的五半径均值最接近1/3。"})

    # KNN figures and summary metrics.
    knn_rows = [row for row in method_rows if row["grouping"] == "knn"]
    fig, axes = plt.subplots(1, 2, figsize=(13.8, 5.2), constrained_layout=True)
    xs = np.arange(len(knn_rows))
    labels = [f"K={row['knn_k']}" for row in knn_rows]
    axes[0].plot(xs, [row["median_overlap"] for row in knn_rows], marker="o", color="#285E8E", label="median")
    axes[0].plot(xs, [row["p10_overlap"] for row in knn_rows], marker="o", color="#8E24AA", label="P10")
    axes[0].axhline(1 / 3, color="#555555", ls="--", lw=1, label="teacher 1/3")
    axes[0].set_xticks(xs)
    axes[0].set_xticklabels(labels)
    axes[0].set_ylim(0, 1)
    axes[0].yaxis.set_major_formatter(PercentFormatter(1.0, decimals=0))
    axes[0].set_title("SA1 nearest-pair overlap: KNN distribution position")
    axes[0].set_ylabel("overlap rate")
    axes[0].legend(fontsize=8)
    axes[1].bar(xs, [row["fraction_exact_100"] for row in knn_rows], color="#8E24AA", label="100% overlap")
    axes[1].plot(xs, [row["fraction_le_1_3"] for row in knn_rows], marker="o", color="#F57C00", label="≤1/3")
    axes[1].set_xticks(xs)
    axes[1].set_xticklabels(labels)
    axes[1].set_ylim(0, 1)
    axes[1].yaxis.set_major_formatter(PercentFormatter(1.0, decimals=0))
    axes[1].set_title("SA1 KNN: threshold pass and exact-100% tail")
    axes[1].set_ylabel("share of nearest pairs")
    axes[1].legend(fontsize=8)
    path = output / "04_SA1_KNN连接数_最近邻重叠统计对比.png"
    fig.savefig(path, dpi=240)
    plt.close(fig)
    figures.append({"file": path.name, "purpose": "KNN 连接数 6/8/10/16 的全部最近邻 pair 分布统计"})

    fig, axes = plt.subplots(1, len(KNN_VALUES), figsize=(19.2, 5.0), constrained_layout=True)
    closest_pair = global_nearest_pair(baseline)
    for ax, k in zip(axes, KNN_VALUES):
        draw_local(ax, baseline, knn_data[k], closest_pair, title=f"KNN-{k}", show_radius=False)
    axes[0].legend(loc="upper left", fontsize=6.5, frameon=True)
    fig.suptitle("SA1 same absolute-nearest FPS-centre pair: KNN connection-count counterfactual", fontsize=14)
    path = output / "05_SA1_KNN连接数_同一最近center局部比较.png"
    fig.savefig(path, dpi=240)
    plt.close(fig)
    figures.append({"file": path.name, "purpose": "同一真正最近 SA1 center 对在 KNN-6/8/10/16 下的局部成员比较"})
    return figures


def style_sheet(ws, *, percent_columns: set[int], title: str, widths: list[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    navy, blue, pale, white = "17324D", "285E8E", "EDF4FA", "FFFFFF"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(widths))
    header = ws.cell(1, 1, title)
    header.font = Font(bold=True, color=white, size=13)
    header.fill = PatternFill("solid", fgColor=navy)
    header.alignment = Alignment(horizontal="center")
    for cell in ws[2]:
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(3, ws.max_row + 1):
        fill = pale if row % 2 else white
        for cell in ws[row]:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(vertical="center", horizontal="center" if cell.column > 2 else "left",
                                       wrap_text=True)
            if cell.column in percent_columns and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00%"
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(widths))}{ws.max_row}"
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 34


def write_workbook(output: Path, method_rows: list[dict], stage_rows: list[dict], radius_stage_rows: list[dict],
                   pair_rows: list[dict], figures: list[dict]) -> Path:
    from openpyxl import Workbook

    book = Workbook()
    overview = book.active
    overview.title = "统计总览"
    overview.append(["Q2V-10477 SA group overlap audit | no-training geometry comparison"])
    headers = ["方法", "Grouping", "SA层", "半径倍率", "KNN连接数", "最近邻pair数", "中位重叠率", "P10重叠率",
               "≤1/4比例", "≤1/3比例", "≤1/2比例", "100%重叠比例", "100%重叠pair数"]
    overview.append(headers)
    for row in method_rows:
        overview.append([row["method"], row["grouping"], row["stage"], row["radius_multiplier"], row["knn_k"],
                         row["nearest_pair_count"], row["median_overlap"], row["p10_overlap"],
                         row["fraction_le_1_4"], row["fraction_le_1_3"], row["fraction_le_1_2"],
                         row["fraction_exact_100"], row["count_exact_100"]])
    style_sheet(overview, percent_columns={7, 8, 9, 10, 11, 12},
                title="Q2V-10477 | SA1 最近 center-group 重叠统计", widths=[24, 18, 8, 12, 12, 14, 14, 14, 14, 14, 14, 16, 17])

    stages = book.create_sheet("SA1-SA3分层")
    stages.append(["Q2V baseline ball-query | layer-wise nearest-pair overlap"])
    stages.append(headers)
    for row in stage_rows:
        stages.append([row["method"], row["grouping"], row["stage"], row["radius_multiplier"], row["knn_k"],
                       row["nearest_pair_count"], row["median_overlap"], row["p10_overlap"],
                       row["fraction_le_1_4"], row["fraction_le_1_3"], row["fraction_le_1_2"],
                       row["fraction_exact_100"], row["count_exact_100"]])
    style_sheet(stages, percent_columns={7, 8, 9, 10, 11, 12},
                title="Q2V-10477 | baseline ball-query 的 SA1 / SA2 / SA3 分层统计", widths=[24, 18, 8, 12, 12, 14, 14, 14, 14, 14, 14, 16, 17])

    radius_stages = book.create_sheet("半径×SA分层")
    radius_stages.append(["Q2V ball-query | layer-wise nearest-pair overlap at every radius multiplier"])
    radius_stages.append(headers)
    for row in radius_stage_rows:
        radius_stages.append([row["method"], row["grouping"], row["stage"], row["radius_multiplier"], row["knn_k"],
                              row["nearest_pair_count"], row["median_overlap"], row["p10_overlap"],
                              row["fraction_le_1_4"], row["fraction_le_1_3"], row["fraction_le_1_2"],
                              row["fraction_exact_100"], row["count_exact_100"]])
    style_sheet(radius_stages, percent_columns={7, 8, 9, 10, 11, 12},
                title="Q2V-10477 | 五种 ball 半径下 SA1 / SA2 / SA3 的最近 center-group 重叠统计",
                widths=[24, 18, 8, 12, 12, 14, 14, 14, 14, 14, 14, 16, 17])

    pairs = book.create_sheet("逐pair明细")
    pairs.append(["Nearest-neighbour pair detail | overlap = shared unique source IDs / smaller group"])
    pair_headers = ["方法", "Grouping", "SA层", "半径倍率", "KNN连接数", "center A", "center B", "center距离(mm)",
                    "A组点数", "B组点数", "共享点数", "重叠率", "Jaccard"]
    pairs.append(pair_headers)
    for row in pair_rows:
        pairs.append([row[key] for key in ("method", "grouping", "stage", "radius_multiplier", "knn_k", "center_a_id",
                                             "center_b_id", "center_distance_mm", "group_a_points", "group_b_points",
                                             "shared_points", "overlap_rate", "jaccard")])
    style_sheet(pairs, percent_columns={12, 13}, title="所有最近邻 center pair 的可追溯明细", widths=[24, 18, 8, 12, 12, 11, 11, 16, 12, 12, 12, 13, 13])
    for row in range(3, pairs.max_row + 1):
        pairs.cell(row, 8).number_format = "0.000"

    index = book.create_sheet("图件与口径")
    index.append(["Read me | visual index and metric definition"])
    index.append(["项目", "内容"])
    notes = [
        ("病例 / 输入", "Q2V-10477；AG/fast/RAN_QING_BO；固定 vertex-random5000 support；固定 FPS centres 500→125→32。"),
        ("最近邻pair", "对每个 center 取空间最近的另一个 center，去重得到 nearest-pair 分布；不是任意远距离 pair。"),
        ("重叠率", "共享的不同 source-point ID 数 ÷ 两个 group 中较小的组大小。ball trace 的重复边不重复计数。"),
        ("P10", "所有最近邻 pair 重叠率的第10百分位，不是第10个 pair。"),
        ("Ball半径", "0.6×/0.8×/1.0×/1.2×/1.5× 分别为基线 r=(0.05,0.10,0.20) 的同比缩放；每层仍最多16个成员。"),
        ("半径×SA分层表", "“半径×SA分层”sheet 为五种 ball 半径各自的 SA1/SA2/SA3 最近邻 pair 分布；阈值列是累计比例，不能相加为100%。"),
        ("KNN", "KNN-6/8/10/16 仅做无训练几何反事实：固定同一支持点和中心，直接取欧氏距离最近的 K 个源点。替换正式训练 grouping 必须重训。"),
        ("教师参考", "每组恰为16点时，≤1/3 等价于共享点≤5（31.25%）；共享6点即37.5%。"),
    ]
    notes.extend((item["file"], item["purpose"]) for item in figures)
    for note in notes:
        index.append(note)
    style_sheet(index, percent_columns=set(), title="口径、反事实边界与图件索引", widths=[42, 110])
    index.auto_filter.ref = "A2:B2"

    path = output / "SA_group重叠统计汇总.xlsx"
    book.save(path)
    return path


def write_readme(output: Path, figures: list[dict]) -> None:
    text = "# Q2V-10477 SA group 重叠审计汇总\n\n"
    text += "- 病例：`AG/fast/RAN_QING_BO`；固定 `vertex-random5000` support 和 FPS centre `500→125→32`。\n"
    text += "- `SA_group重叠统计汇总.xlsx`：统计总览、baseline SA1–SA3分层、五种半径×三层分层、逐pair明细和指标定义。\n"
    text += "- `01`–`15`：给教师看的图件；`_traces/`：五种 ball 半径的原始可追溯 SA trace。\n"
    text += "- 所有重叠均按不同 source-point ID 计算；KNN 是无训练反事实，不可直接代替模型性能结论。\n\n"
    text += "## 图件\n\n"
    for figure in figures:
        text += f"- `{figure['file']}`：{figure['purpose']}。\n"
    (output / "README.md").write_text(text, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="build Q2V SA overlap audit workbook and figures")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output = args.output
    output.mkdir(parents=True, exist_ok=True)
    missing = [TRACE_ROOT / dirname for _, _, dirname in BALL_SETTINGS if not (TRACE_ROOT / dirname).is_dir()]
    if missing:
        raise FileNotFoundError("missing exact ball traces: " + ", ".join(map(str, missing)))

    ball_data = {multiplier: load_stage(TRACE_ROOT / dirname, 1) for _, multiplier, dirname in BALL_SETTINGS}
    # Deterministic FPS centre positions must remain frozen across every layer
    # of the radius-only counterfactual; only the ball-query memberships vary.
    for stage in (1, 2, 3):
        reference = load_stage(TRACE_ROOT / "ball_r100", stage)
        reference_ids = sorted(reference["centres"])
        reference_xyz = np.asarray([reference["centres"][key] for key in reference_ids])
        for _, multiplier, dirname in BALL_SETTINGS:
            current_data = ball_data[multiplier] if stage == 1 else load_stage(TRACE_ROOT / dirname, stage)
            current = np.asarray([current_data["centres"][key] for key in reference_ids])
            if not np.array_equal(reference_xyz, current):
                raise RuntimeError(f"radius {multiplier} changed SA{stage} FPS centres")

    method_rows: list[dict] = []
    pair_rows: list[dict] = []
    for label, multiplier, _ in BALL_SETTINGS:
        summary, details = summarise(label, "ball", ball_data[multiplier], ball_data[multiplier]["members"],
                                     multiplier=multiplier)
        method_rows.append(summary)
        pair_rows.extend(details)
    knn_data = {k: knn_members(ball_data[1.0], k) for k in KNN_VALUES}
    for k, members in knn_data.items():
        summary, details = summarise(f"KNN-{k}", "knn", ball_data[1.0], members, k=k)
        method_rows.append(summary)
        pair_rows.extend(details)

    stage_rows: list[dict] = []
    for stage in (1, 2, 3):
        data = load_stage(TRACE_ROOT / "ball_r100", stage)
        summary, details = summarise(f"Ball baseline SA{stage}", "ball", data, data["members"], stage=stage,
                                     multiplier=1.0)
        stage_rows.append(summary)
        pair_rows.extend(details)

    # Each radius replays all three SA layers over the frozen input/FPS protocol.
    # Keep this as a separate sheet so the baseline-only layer comparison stays readable.
    radius_stage_rows: list[dict] = []
    for label, multiplier, dirname in BALL_SETTINGS:
        for stage in (1, 2, 3):
            data = load_stage(TRACE_ROOT / dirname, stage)
            summary, _ = summarise(f"{label} SA{stage}", "ball", data, data["members"], stage=stage,
                                   multiplier=multiplier)
            radius_stage_rows.append(summary)

    figures = write_figures(output, ball_data, knn_data, method_rows, stage_rows)
    optional_figures = (
        ("06_SA1_完整血管_相邻group与局部放大_跨支风险审计.png",
         "完整血管表面中的相邻 SA1 group 定位 + 右侧局部放大；红叉为表面测地/欧氏距离比≥2.5的拓扑风险筛查标记"),
        ("07_SA1_ball_完整点云与局部放大.png",
         "baseline SA1 ball-query：灰色完整点云定位 + 右侧局部放大；蓝/黄虚线圈为实际 ball 半径"),
        ("08_SA1_KNN8_完整点云与局部放大.png",
         "KNN-8：灰色完整点云定位 + 右侧局部放大；蓝/黄虚线圈仅为 baseline ball 半径空间参照，不是 KNN 边界"),
        ("09_SA2_ball_完整点云与局部放大.png",
         "baseline SA2 ball-query：左侧为真正的 500 个 SA1-center 源点云，右侧为最近 center pair 的局部放大"),
        ("10_SA3_ball_完整点云与局部放大.png",
         "baseline SA3 ball-query：左侧为真正的 125 个 SA2-center 源点云，右侧为最近 center pair 的局部放大"),
        ("11_SA1_ball_r0.6x_完整点云与局部放大.png",
         "SA1 ball 0.6×：与 0.8×–1.5× 固定同一最近邻 center pair 的完整点云定位与局部放大"),
        ("12_SA1_ball_r0.8x_完整点云与局部放大.png",
         "SA1 ball 0.8×：与其他半径固定同一最近邻 center pair 的完整点云定位与局部放大"),
        ("13_SA1_ball_r1.0x_完整点云与局部放大.png",
         "SA1 baseline ball 1.0×：与其他半径固定同一最近邻 center pair 的完整点云定位与局部放大"),
        ("14_SA1_ball_r1.2x_完整点云与局部放大.png",
         "SA1 ball 1.2×：与其他半径固定同一最近邻 center pair 的完整点云定位与局部放大"),
        ("15_SA1_ball_r1.5x_完整点云与局部放大.png",
         "SA1 ball 1.5×：与其他半径固定同一最近邻 center pair 的完整点云定位与局部放大"),
    )
    for filename, purpose in optional_figures:
        if (output / filename).is_file():
            figures.append({"file": filename, "purpose": purpose})
    workbook = write_workbook(output, method_rows, stage_rows, radius_stage_rows, pair_rows, figures)
    write_readme(output, figures)
    summary_csv = output / "SA1_方法统计汇总.csv"
    with summary_csv.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(method_rows[0]))
        writer.writeheader()
        writer.writerows(method_rows)
    print(workbook)
    for figure in figures:
        print(output / figure["file"])


if __name__ == "__main__":
    main()
