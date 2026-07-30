#!/usr/bin/env python3
"""Update the WSS PointNet workbook with REG-P10 Transformer results."""
from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from training_wss_min.tools.update_s3_regularization_xlsx import matrix_values
from training_wss_min.tools.update_s3_rootcause_xlsx import copy_row_style


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = (
    ROOT
    / "training_wss_min/preflight/"
    "regp10_transformer_matrix_20260726_results_analysis.json"
)
SECTION_TITLE = (
    "2026-07-26｜REG-P10 局部 SA Transformer 完整矩阵 + SA3 全局对照"
)
ORDER = (
    "localtf_sa1_s1234",
    "localtf_sa2_s1234",
    "localtf_sa3_s1234",
    "localtf_sa12_s1234",
    "localtf_sa13_s1234",
    "localtf_sa23_s1234",
    "localtf_sa123_s1234",
    "globaltf_sa3_s1234",
)
LABELS = {
    "localtf_sa1_s1234": "REG-P10 + L-SA1 s1234",
    "localtf_sa2_s1234": "REG-P10 + L-SA2 s1234",
    "localtf_sa3_s1234": "REG-P10 + L-SA3 s1234",
    "localtf_sa12_s1234": "REG-P10 + L-SA12 s1234",
    "localtf_sa13_s1234": "REG-P10 + L-SA13 s1234",
    "localtf_sa23_s1234": "REG-P10 + L-SA23 s1234",
    "localtf_sa123_s1234": "REG-P10 + L-SA123 s1234",
    "globaltf_sa3_s1234": "REG-P10 + G-SA3 s1234",
}
CHANGES = {
    "localtf_sa1_s1234": "local Transformer stages=[1]",
    "localtf_sa2_s1234": "local Transformer stages=[2]",
    "localtf_sa3_s1234": "local Transformer stages=[3]",
    "localtf_sa12_s1234": "local Transformer stages=[1,2]",
    "localtf_sa13_s1234": "local Transformer stages=[1,3]",
    "localtf_sa23_s1234": "local Transformer stages=[2,3]",
    "localtf_sa123_s1234": "local Transformer stages=[1,2,3]",
    "globaltf_sa3_s1234": "SA3 coarse global attention（32 centers / 4 heads）",
}
DECISIONS = {
    "go_to_three_seed_confirmation": "晋级 seeds 7/2025 确认",
    "below_primary_gate": "正向但未过 ΔR²_cb 主门槛",
    "no_go_single_seed": "单种子 No-Go",
}


def overview_values(record: dict) -> list:
    values = matrix_values({**record, "variant": "head_dropout01", "seed": 1234})
    experiment_id = record["experiment_id"]
    values[0:4] = [
        LABELS[experiment_id],
        "138/0/36（AG/AAA/ILO mixed；seed=1234）",
        (
            "PointNeXt-R + LocalGeoPE + local Transformer"
            if record["variant"] == "local"
            else "PointNeXt-R + LocalGeoPE + SA3 global attention"
        ),
        CHANGES[experiment_id],
    ]
    return values


def teacher_values(record: dict) -> list:
    raw = json.loads(
        (Path(record["run_dir"]) / "eval/ckpt_best/metrics.json").read_text(
            encoding="utf-8"
        )
    )["test"]
    normalized = raw["normalized"]
    return [
        "REG-P10 Transformer",
        LABELS[record["experiment_id"]],
        (
            "PointNeXt-R + LocalGeoPE + local-TF"
            if record["variant"] == "local"
            else "PointNeXt-R + LocalGeoPE + global-TF"
        ),
        "random5000 / SAME",
        raw["field_casebalanced"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["field"]["rmse"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        normalized["field_casebalanced"]["r2"],
        normalized["field"].get("nmae_range"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["top10_iou_casemean"],
        normalized["calibration"]["p99_pred_true_ratio"],
    ]


def title(sheet, row: int, text: str) -> int:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=15)
    cell = sheet.cell(row, 1, text)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    return row + 1


def header(sheet, row: int, values: list[str]) -> int:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    return row + 1


def main() -> None:
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    records = analysis["records"]
    comparisons = {
        item["treatment"]: item for item in analysis["paired_comparisons"]
    }
    for experiment_id in ORDER:
        if experiment_id not in records or experiment_id not in comparisons:
            raise RuntimeError(f"analysis missing {experiment_id}")
        if records[experiment_id]["integrity"] != "passed":
            raise RuntimeError(f"integrity failed: {experiment_id}")

    with tempfile.TemporaryDirectory(prefix="wss_regp10tf_xlsx_") as temp_dir:
        candidate = Path(temp_dir) / BOOK.name
        shutil.copy2(BOOK, candidate)
        workbook = load_workbook(candidate, data_only=False)

        overall = workbook["实验矩阵总览"]
        existing_rows = [
            cell.row for cell in overall["AK"] if cell.value in set(ORDER)
        ]
        for row in sorted(existing_rows, reverse=True):
            overall.delete_rows(row, 1)
        first_overall = overall.max_row + 1
        for experiment_id in ORDER:
            row = overall.max_row + 1
            copy_row_style(overall, row - 1, row, 37)
            for column, value in enumerate(
                overview_values(records[experiment_id]), 1
            ):
                overall.cell(row, column, value)

        teacher = workbook["教师汇报视图"]
        existing_teacher_rows = [
            cell.row for cell in teacher["B"] if cell.value in set(LABELS.values())
        ]
        for row in sorted(existing_teacher_rows, reverse=True):
            teacher.delete_rows(row, 1)
        first_teacher = teacher.max_row + 1
        for experiment_id in ORDER:
            row = teacher.max_row + 1
            copy_row_style(teacher, row - 1, row, 15)
            for column, value in enumerate(
                teacher_values(records[experiment_id]), 1
            ):
                teacher.cell(row, column, value)

        summary = workbook["汇总对比"]
        title_rows = [
            cell.row for cell in summary["A"] if cell.value == SECTION_TITLE
        ]
        if len(title_rows) > 1:
            raise RuntimeError("duplicate REG-P10 Transformer summary sections")
        if title_rows:
            start = title_rows[0]
            summary.delete_rows(start, 2 + len(ORDER))
        summary_start = summary.max_row + 1
        row = title(summary, summary_start, SECTION_TITLE)
        row = header(
            summary,
            row,
            [
                "处理臂",
                "模块位置",
                "R²_cb",
                "ΔR²_cb",
                "ΔMAE Pa",
                "ΔRMSE Pa",
                "Δhigh-WSS R²",
                "Δtop10 IoU",
                "ΔAG",
                "ΔAAA",
                "ΔILO",
                "病例均值Δ 95%CI",
                "病例胜/负",
                "判定",
                "last−best R²_cb",
            ],
        )
        for experiment_id in ORDER:
            record = records[experiment_id]
            result = comparisons[experiment_id]
            delta = result["aggregate_treatment_minus_control"]
            domains = result["domain_r2_delta"]
            cases = result["paired_case_stats"]["case_r2"]
            values = [
                LABELS[experiment_id],
                CHANGES[experiment_id],
                record["best"]["physical_r2_casebalanced"],
                delta["physical_r2_casebalanced"],
                delta["physical_mae"],
                delta["physical_rmse"],
                delta["high_wss_r2"],
                delta["physical_top10_iou"],
                domains["AG"],
                domains["AAA"],
                domains["ILO"],
                (
                    f"{cases['mean_delta']:+.4f} "
                    f"[{cases['ci95_low']:+.4f},{cases['ci95_high']:+.4f}]"
                ),
                f"{cases['wins']}/{cases['losses']}",
                DECISIONS[result["gate"]["decision"]],
                record["last_minus_best"]["physical_r2_casebalanced"],
            ]
            for column, value in enumerate(values, 1):
                summary.cell(row, column, value)
            row += 1

        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(candidate)
        workbook.close()

        check = load_workbook(candidate, read_only=True, data_only=False)
        saved_ids = {
            values[0]
            for values in check["实验矩阵总览"].iter_rows(
                min_col=37, max_col=37, values_only=True
            )
        }
        if not set(ORDER).issubset(saved_ids):
            raise RuntimeError("overview readback missing REG-P10 Transformer rows")
        saved_labels = {
            values[0]
            for values in check["教师汇报视图"].iter_rows(
                min_col=2, max_col=2, values_only=True
            )
        }
        if not set(LABELS.values()).issubset(saved_labels):
            raise RuntimeError("teacher readback missing REG-P10 Transformer rows")
        saved_titles = {
            values[0]
            for values in check["汇总对比"].iter_rows(
                min_col=1, max_col=1, values_only=True
            )
        }
        if SECTION_TITLE not in saved_titles:
            raise RuntimeError("summary readback missing REG-P10 Transformer section")
        check.close()

        render_dir = Path(temp_dir) / "render"
        render_dir.mkdir()
        subprocess.run(
            [
                "/usr/bin/libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(render_dir),
                str(candidate),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        if not (render_dir / candidate.with_suffix(".pdf").name).is_file():
            raise RuntimeError("LibreOffice render check failed")
        shutil.copy2(candidate, BOOK)

    print(
        json.dumps(
            {
                "status": "updated_and_rendered",
                "workbook": str(BOOK),
                "overall_rows": [first_overall, first_overall + len(ORDER) - 1],
                "teacher_rows": [first_teacher, first_teacher + len(ORDER) - 1],
                "summary_start_row": summary_start,
                "rows_added": len(ORDER),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
