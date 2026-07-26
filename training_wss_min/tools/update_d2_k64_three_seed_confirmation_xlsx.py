#!/usr/bin/env python3
"""Append the frozen M1/S2/S3 three-seed confirmation summary to the workbook."""
from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = ROOT / "training_wss_min/preflight/d2_k64_ilo_structure_three_seed_confirmation_analysis.json"
SHEET = "D2K64三种子确认"


def main() -> None:
    data = json.loads(ANALYSIS.read_text())
    if not data["decision"]["s3_advances_to_drop_single_variable"]:
        raise RuntimeError("analysis does not pass the core S3 decision rule")
    with tempfile.TemporaryDirectory(prefix="d2k64_confirm_xlsx_") as tmp:
        candidate = Path(tmp) / BOOK.name
        shutil.copy2(BOOK, candidate)
        wb = load_workbook(candidate)
        if SHEET in wb.sheetnames:
            raise RuntimeError(f"refusing duplicate sheet: {SHEET}")
        ws = wb.create_sheet(SHEET)
        ws.append(["D2-K64 M1/S2/S3 三种子严格配对确认（test36：历史工程筛选集，非独立外部测试）"])
        ws.merge_cells("A1:J1"); ws["A1"].font = Font(bold=True, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        ws.append(["比较", "seed", "ΔR²_cb", "ΔMAE Pa", "ΔRMSE Pa", "Δhigh-WSS R²", "Δp99 ratio", "Δtop10 IoU", "ΔSpearman", "备注"])
        for cell in ws[2]: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for comparison in data["comparisons"]:
            for row in comparison["per_seed"]:
                delta = row["best_delta"]
                ws.append([comparison["comparison"], row["seed"], delta["physical_r2_casebalanced"], delta["physical_mae"], delta["physical_rmse"], delta["high_wss_r2"], delta["physical_p99_amplitude_ratio"], delta["physical_top10_iou"], delta["spearman_case_mean"], "ckpt_best(train_loss)"])
            stat = comparison["best_seed_summary"]["physical_r2_casebalanced"]
            boot = comparison["paired_case_bootstrap"]["r2"]
            ws.append([comparison["comparison"] + " summary", "mean±sd", stat["mean"], None, None, None, None, None, None, f"R²_cb: {stat['mean']:+.4f}±{stat['std']:.4f}; pos/neg={stat['positive']}/{stat['negative']}; case bootstrap {boot['wins']}/{boot['losses']} [{boot['ci95_low']:+.4f},{boot['ci95_high']:+.4f}]"])
        ws.append([])
        ws.append(["决策", "S3 通过；仅提出下一阶段单变量：DropPath 0.05、DropPath 0.10、NeighborDrop 0.05；本批未自动训练。"])
        ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=10)
        for column, width in enumerate((24, 14, 14, 14, 14, 16, 16, 16, 16, 72), 1): ws.column_dimensions[chr(64 + column)].width = width
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=3, max_col=9):
            for cell in row: cell.number_format = "0.0000"
        wb.save(candidate); wb.close()
        check = load_workbook(candidate, read_only=True, data_only=False)
        if SHEET not in check.sheetnames or check[SHEET].max_row < 13: raise RuntimeError("xlsx readback failed")
        check.close()
        render_dir = Path(tmp) / "render"; render_dir.mkdir()
        subprocess.run(["/usr/bin/libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(render_dir), str(candidate)], check=True, capture_output=True, text=True)
        if not (render_dir / candidate.with_suffix(".pdf").name).is_file(): raise RuntimeError("LibreOffice render check failed")
        shutil.copy2(candidate, BOOK)
    print(json.dumps({"status": "updated_and_readback_rendered", "workbook": str(BOOK), "sheet": SHEET}, ensure_ascii=False))


if __name__ == "__main__": main()
