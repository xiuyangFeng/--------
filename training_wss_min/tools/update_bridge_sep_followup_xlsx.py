#!/usr/bin/env python3
"""Append the bridge SEP control follow-up row to the PointNet results xlsx.

Mirrors the column construction in analyze_pointnetpp_sa1_scale_results.py
(excel_values / teacher block / summary-comparison block) for a single new
row, without re-running the full 17-arm matrix analysis.
"""
from __future__ import annotations

import json
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = ROOT / "training_wss_min/preflight/bridge_sep_followup_results_analysis.json"

PARENT_LABEL = "bridge random5000 + 500c + k64"
CHILD_LABEL = "bridge random5000 + 500c + k64 SEP（对照 SAME）"
CHILD_EXP_ID = "bridge_rand5000_fixed500_k64_sep"
SECTION_TITLE = "2026-07-22｜bridge SEP 对照臂（Δ 相对 bridge_rand5000_fixed500_k64，唯一变量 SAME→SEP）"


def load(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def copy_row_style(ws, source_row: int, target_row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        source, target = ws.cell(source_row, col), ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def excel_values(run_dir: Path, m: dict[str, Any]) -> list[Any]:
    raw = load(run_dir / "eval/ckpt_best/metrics.json")["test"]
    norm = raw["normalized"]
    return [
        CHILD_LABEL, "106/0/27（historical test27；single seed=1234）", "PointNet++ SA3",
        "5000→500/125/32；nsample=64；width=32；SA1=knn_cover；batch=8",
        "random / SEP；FPS center", 5000,
        m["physical_r2_cb"], m["physical_r2"], m["physical_casemean_r2"], raw["aggregate"]["r2_casemed"], raw["aggregate"]["r2_casep10"],
        f"{m['physical_r2_negative_cases']}/27", raw["field"]["rmse"], raw["field"]["mae"], raw["field"]["nrmse_range"], raw["aggregate"]["nrmse_casemean"],
        raw["field"].get("nmae_range"), raw["aggregate"].get("nmae_casemean"), m["physical_high_wss_r2"], raw["calibration"]["top10_pred_true_ratio"],
        m["physical_top10_iou"], m["normalized_r2_cb"], m["normalized_r2"], m["normalized_casemean_r2"], norm["aggregate"]["r2_casemed"], norm["aggregate"]["r2_casep10"],
        norm["field_casebalanced"]["mae"], norm["field_casebalanced"]["rmse"], norm["field"]["nrmse_range"], norm["aggregate"]["nrmse_casemean"],
        norm["field"].get("nmae_range"), norm["aggregate"].get("nmae_casemean"), m["physical_spearman"], raw["hotspot"]["spearman_high_wss_casemean"],
        norm["calibration"]["top10_pred_true_ratio"], norm["calibration"]["p99_pred_true_ratio"], CHILD_EXP_ID,
    ]


def teacher_values(run_dir: Path, m: dict[str, Any]) -> list[Any]:
    raw = load(run_dir / "eval/ckpt_best/metrics.json")["test"]
    return [
        "SA1-scale", CHILD_LABEL, "PointNet++ SA3", "5000 / SEP",
        m["physical_r2_cb"], m["physical_casemean_r2"], raw["field"]["rmse"],
        raw["field"]["nrmse_range"], raw["normalized"]["field"]["nrmse_range"],
        m["physical_high_wss_r2"], m["normalized_r2_cb"],
        raw["normalized"]["field"]["nrmse_range"],
        m["physical_spearman"], m["physical_top10_iou"], m["normalized_r2"],
    ]


def main() -> None:
    analysis = load(ANALYSIS)
    child = analysis["child"]
    parent_best = analysis["parent"]["metrics_best"]
    child_best = child["metrics_best"]
    deltas = analysis["paired_delta_vs_parent_best"]
    run_dir = Path(child["run_dir"])

    wb = load_workbook(XLSX, data_only=False)

    overall = wb["实验矩阵总览"]
    if CHILD_EXP_ID in [c.value for c in overall["AK"]]:
        raise RuntimeError("bridge SEP row already present in 实验矩阵总览; refusing duplicate append")
    row = overall.max_row + 1
    copy_row_style(overall, row - 1, row, 37)
    for col, value in enumerate(excel_values(run_dir, child_best), 1):
        overall.cell(row, col, value)

    teacher = wb["教师汇报视图"]
    trow = teacher.max_row + 1
    copy_row_style(teacher, trow - 1, trow, 15)
    for col, value in enumerate(teacher_values(run_dir, child_best), 1):
        teacher.cell(trow, col, value)

    summary = wb["汇总对比"]
    srow = summary.max_row + 1
    summary.merge_cells(start_row=srow, start_column=1, end_row=srow, end_column=15)
    cell = summary.cell(srow, 1, SECTION_TITLE)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    srow += 1
    summary.merge_cells(start_row=srow, start_column=1, end_row=srow, end_column=15)
    summary.cell(srow, 1, "bridge_sep_followup（Δ 相对 reference_id）").font = Font(bold=True)
    srow += 1
    headers = ["配置", "对照", "物理R²_cb", "ΔR²_cb", "归一化R²_cb", "Δ", "病例均值R²", "负例", "MAE", "ΔMAE", "Spearman", "Δ", "top10 IoU", "Δ", "high-WSS R²"]
    for col, value in enumerate(headers, 1):
        c = summary.cell(srow, col, value)
        c.font = Font(bold=True)
    srow += 1
    values = [
        CHILD_LABEL, "bridge_rand5000_fixed500_k64",
        child_best["physical_r2_cb"], deltas["physical_r2_cb"],
        child_best["normalized_r2_cb"], deltas["normalized_r2_cb"],
        child_best["physical_casemean_r2"], child_best["physical_r2_negative_cases"],
        child_best["physical_mae"], deltas["physical_mae"],
        child_best["physical_spearman"], deltas["physical_spearman"],
        child_best["physical_top10_iou"], deltas["physical_top10_iou"],
        child_best["physical_high_wss_r2"],
    ]
    for col, value in enumerate(values, 1):
        summary.cell(srow, col, value)

    wb.save(XLSX)

    check = load_workbook(XLSX, data_only=False)
    assert CHILD_EXP_ID in [c.value for c in check["实验矩阵总览"]["AK"]]
    assert CHILD_LABEL in [c.value for c in check["教师汇报视图"]["B"]]
    assert SECTION_TITLE in [c.value for c in check["汇总对比"]["A"]]
    print(json.dumps({
        "status": "updated", "overall_row": row, "teacher_row": trow, "summary_section_row": srow - 3,
        "parent_physical_r2_cb": parent_best["physical_r2_cb"], "child_physical_r2_cb": child_best["physical_r2_cb"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
