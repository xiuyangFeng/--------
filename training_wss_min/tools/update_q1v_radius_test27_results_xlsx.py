#!/usr/bin/env python3
"""Backfill completed Q1V/SAME historical-test27 radius exploration metrics."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


REPO = Path(__file__).resolve().parents[2]
WORKBOOK = REPO / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = REPO / "training_wss_min/preflight/q1v_radius_test27_results_analysis.json"

METRIC_KEYS = (
    "physical_r2_casebalanced", "physical_r2_pooled", "physical_r2_case_mean",
    "physical_r2_case_median", "physical_r2_case_p10", "physical_r2_negative_cases",
    "physical_rmse", "physical_mae", "physical_nrmse_range_pooled",
    "physical_nrmse_range_case_mean", "physical_nmae_pooled", "physical_nmae_case_mean",
    "high_wss_r2", "physical_top10_amplitude_ratio", "physical_top10_iou",
    "normalized_r2_casebalanced", "normalized_r2_pooled", "normalized_r2_case_mean",
    "normalized_r2_case_median", "normalized_r2_case_p10", "normalized_mae", "normalized_rmse",
    "normalized_nrmse_range_pooled", "normalized_nrmse_range_case_mean",
    "normalized_nmae_pooled", "normalized_nmae_case_mean", "spearman_case_mean",
    "high_wss_spearman_case_mean", "normalized_top10_amplitude_ratio",
    "physical_p99_amplitude_ratio",
)

META = {
    "q1v_radius_r80_same": ("Q1V半径探索(0.8×)", "0.04/0.08/0.16", "10546_0"),
    "q1v_radius_r120_same": ("Q1V半径探索(1.2×)", "0.06/0.12/0.24", "10546_1"),
    "q1v_radius_r150_same": ("Q1V半径探索(1.5×)", "0.075/0.15/0.30", "10546_2"),
}


def find_row(ws, predicate) -> int:
    matches = [row for row in range(1, ws.max_row + 1) if predicate(ws.cell(row, 1).value, ws.cell(row, 2).value)]
    if len(matches) != 1:
        raise RuntimeError(f"expected exactly one matching row in {ws.title}, found {matches}")
    return matches[0]


def update_overview(ws, records: dict[str, dict]) -> None:
    section = find_row(ws, lambda col_a, _: isinstance(col_a, str) and col_a.startswith("Q1V/SAME 半径探索"))
    ws.cell(section, 1).value = "Q1V/SAME 半径探索（2026-07-19；3/3完成，非确认性）"
    for experiment_id, record in records.items():
        label, radii, job_id = META[experiment_id]
        row = find_row(ws, lambda col_a, _: isinstance(col_a, str) and col_a.startswith(label[:-1]))
        metric = record["best"]
        values = (label, "106/0/27（历史test27探索）", "PointNet++ SA3",
                  f"5000→500→125→32；r={radii}；nsample=16；width=32",
                  "vertex random5000 / SAME；FPS center", 5000)
        for col, value in enumerate(values, start=1):
            ws.cell(row, col).value = value
        for col, key in enumerate(METRIC_KEYS, start=7):
            ws.cell(row, col).value = metric.get(key)


def update_teacher(ws, records: dict[str, dict]) -> None:
    section = find_row(ws, lambda col_a, _: isinstance(col_a, str) and col_a.startswith("Ⅶ Q1V/SAME 半径探索"))
    ws.cell(section, 1).value = "Ⅶ Q1V/SAME 半径探索（2026-07-19；3/3完成，非确认性）"
    for experiment_id, record in records.items():
        label, _, job_id = META[experiment_id]
        row = find_row(ws, lambda _, col_b: isinstance(col_b, str) and col_b.startswith(label[:-1]))
        metric = record["best"]
        values = (
            f"探索性test27（{job_id}；已完成）", label, "PointNet++ SA3",
            "vertex random5000 / SAME；FPS center", metric["physical_r2_casebalanced"],
            metric["physical_r2_case_mean"], metric["physical_rmse"], metric["physical_nmae_pooled"],
            metric["physical_nmae_case_mean"], metric["high_wss_r2"],
            metric["normalized_r2_casebalanced"], metric["normalized_nmae_pooled"],
            metric["spearman_case_mean"], metric["physical_top10_iou"], metric["physical_p99_amplitude_ratio"],
        )
        for col, value in enumerate(values, start=1):
            ws.cell(row, col).value = value


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--analysis", type=Path, default=ANALYSIS)
    args = parser.parse_args()
    analysis = json.loads(args.analysis.read_text(encoding="utf-8"))
    if analysis.get("test27_policy") != "user-authorized historical anchor; not independent confirmation":
        raise RuntimeError("refusing to write a non-exploratory historical-test27 result")
    records = {item["experiment_id"]: item for item in analysis["records"]}
    if set(records) != set(META) or any(item["integrity"] != "passed" for item in records.values()):
        raise RuntimeError("analysis does not contain the exact three integrity-passed Q1V radius records")
    wb = load_workbook(args.workbook)
    update_overview(wb["实验矩阵总览"], records)
    update_teacher(wb["教师汇报视图"], records)
    wb.save(args.workbook)
    check = load_workbook(args.workbook, read_only=True, data_only=False)
    for experiment_id, (label, _, _) in META.items():
        for sheet in ("实验矩阵总览", "教师汇报视图"):
            found = any(
                check[sheet].cell(row, 1).value == label or check[sheet].cell(row, 2).value == label
                for row in range(1, check[sheet].max_row + 1)
            )
            if not found:
                raise RuntimeError(f"workbook verification failed for {experiment_id} in {sheet}")
    print(json.dumps({"status": "updated", "workbook": str(args.workbook), "rows": len(records)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
