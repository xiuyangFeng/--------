#!/usr/bin/env python3
"""Update the WSS workbook with fixed D2 local-Transformer SA1/SA2 results."""
from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from training_wss_min.tools.update_s3_regularization_xlsx import matrix_values
from training_wss_min.tools.update_s3_rootcause_xlsx import copy_row_style


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = (
    ROOT
    / "training_wss_min/preflight/"
    "d2_c125_k64_pnxr_geope_transformer_20260726_results_analysis.json"
)
TITLE = (
    "2026-07-27｜固定 D2 PNXR+GeoPE：local-SA1 / local-SA2（seed1234）"
)
ORDER = (
    "d2_pnxr_geope_localtf_sa1_s1234",
    "d2_pnxr_geope_localtf_sa2_s1234",
)
LABELS = {
    "d2_pnxr_geope_localtf_sa1_s1234": "D2 PNXR+GeoPE + L-SA1 s1234",
    "d2_pnxr_geope_localtf_sa2_s1234": "D2 PNXR+GeoPE + L-SA2 s1234",
}
CHANGES = {
    "d2_pnxr_geope_localtf_sa1_s1234": "local Transformer stages=[1]",
    "d2_pnxr_geope_localtf_sa2_s1234": "local Transformer stages=[2]",
}
DECISIONS = {
    "no_go": "No-Go",
    "flat_no_promotion": "持平，不晋级",
    "positive_screen": "单种子正信号",
}


def raw_metrics(record: dict) -> dict:
    return json.loads(
        (Path(record["run_dir"]) / "eval/ckpt_best/metrics.json").read_text(
            encoding="utf-8"
        )
    )["test"]


def overview_values(record: dict) -> list:
    values = matrix_values({**record, "variant": "head_dropout01", "seed": 1234})
    experiment_id = record["experiment_id"]
    values[0:4] = [
        LABELS[experiment_id],
        "106/0/27（AG/AAA；seed=1234）",
        "PointNeXt-R + LocalGeoPE + local Transformer",
        CHANGES[experiment_id],
    ]
    return values


def teacher_values(record: dict) -> list:
    raw = raw_metrics(record)
    normalized = raw["normalized"]
    return [
        "D2 fixed Transformer",
        LABELS[record["experiment_id"]],
        "PointNeXt-R + LocalGeoPE + local-TF",
        "random5000 / SAME",
        raw["field_casebalanced"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["field"]["rmse"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        normalized["field_casebalanced"]["r2"],
        normalized["field"].get("nmae_range"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["top10_iou_casemean"],
        normalized["calibration"]["p99_pred_true_ratio"],
    ]


def title(sheet, row: int) -> int:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=15)
    cell = sheet.cell(row, 1, TITLE)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    return row + 1


def header(sheet, row: int, values: list[str]) -> int:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    return row + 1


def main() -> None:
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    records = analysis["records"]
    comparisons = {
        item["treatment"]: item
        for item in analysis["paired_comparisons_vs_parent"]
    }
    for experiment_id in ORDER:
        if records[experiment_id]["integrity"] != "passed":
            raise RuntimeError(f"integrity failed: {experiment_id}")

    with tempfile.TemporaryDirectory(prefix="wss_d2tf_xlsx_") as temp_dir:
        candidate = Path(temp_dir) / BOOK.name
        shutil.copy2(BOOK, candidate)
        workbook = load_workbook(candidate, data_only=False)

        overall = workbook["实验矩阵总览"]
        for row in sorted(
            [cell.row for cell in overall["AK"] if cell.value in set(ORDER)],
            reverse=True,
        ):
            overall.delete_rows(row, 1)
        first_overall = overall.max_row + 1
        for experiment_id in ORDER:
            row = overall.max_row + 1
            copy_row_style(overall, row - 1, row, 37)
            for column, value in enumerate(
                overview_values(records[experiment_id]), 1
            ):
                overall.cell(row, column, value)

        teacher = workbook["教师汇报视图"]
        for row in sorted(
            [
                cell.row
                for cell in teacher["B"]
                if cell.value in set(LABELS.values())
            ],
            reverse=True,
        ):
            teacher.delete_rows(row, 1)
        first_teacher = teacher.max_row + 1
        for experiment_id in ORDER:
            row = teacher.max_row + 1
            copy_row_style(teacher, row - 1, row, 15)
            for column, value in enumerate(
                teacher_values(records[experiment_id]), 1
            ):
                teacher.cell(row, column, value)

        summary = workbook["汇总对比"]
        title_rows = [cell.row for cell in summary["A"] if cell.value == TITLE]
        if len(title_rows) > 1:
            raise RuntimeError("duplicate D2 Transformer summary sections")
        if title_rows:
            summary.delete_rows(title_rows[0], 2 + len(ORDER))
        summary_start = summary.max_row + 1
        row = title(summary, summary_start)
        row = header(
            summary,
            row,
            [
                "处理臂",
                "物理 R²_cb",
                "Δ物理 R²_cb",
                "归一化 R²_cb",
                "Δ归一化 R²_cb",
                "ΔMAE_cb Pa",
                "ΔRMSE_cb Pa",
                "Δhigh-WSS R²",
                "Δtop10 IoU",
                "ΔAG",
                "ΔAAA",
                "病例均值Δ 95%CI",
                "胜/负",
                "last−parent last R²_cb",
                "主结果判定",
            ],
        )
        for experiment_id in ORDER:
            record = records[experiment_id]
            result = comparisons[experiment_id]
            delta = result["aggregate_treatment_minus_control"]
            domains = result["domain_r2_delta"]
            case = result["paired_case_stats"]["case_r2"]
            values = [
                LABELS[experiment_id],
                record["best"]["physical_r2_casebalanced"],
                delta["physical_r2_casebalanced"],
                record["best"]["normalized_r2_casebalanced"],
                delta["normalized_r2_casebalanced"],
                delta["physical_mae"],
                delta["physical_rmse"],
                delta["high_wss_r2"],
                delta["physical_top10_iou"],
                domains["AG"],
                domains["AAA"],
                (
                    f"{case['mean_delta']:+.4f} "
                    f"[{case['ci95_low']:+.4f},{case['ci95_high']:+.4f}]"
                ),
                f"{case['wins']}/{case['losses']}",
                result["last_treatment_minus_parent_last"][
                    "physical_r2_casebalanced"
                ],
                DECISIONS[result["decision"]],
            ]
            for column, value in enumerate(values, 1):
                summary.cell(row, column, value)
            row += 1

        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(candidate)
        workbook.close()

        check = load_workbook(candidate, read_only=True, data_only=False)
        saved_ids = {
            row[0]
            for row in check["实验矩阵总览"].iter_rows(
                min_col=37, max_col=37, values_only=True
            )
        }
        saved_labels = {
            row[0]
            for row in check["教师汇报视图"].iter_rows(
                min_col=2, max_col=2, values_only=True
            )
        }
        saved_titles = {
            row[0]
            for row in check["汇总对比"].iter_rows(
                min_col=1, max_col=1, values_only=True
            )
        }
        if (
            not set(ORDER).issubset(saved_ids)
            or not set(LABELS.values()).issubset(saved_labels)
            or TITLE not in saved_titles
        ):
            raise RuntimeError("xlsx readback failed")
        check.close()

        render_dir = Path(temp_dir) / "render"
        render_dir.mkdir()
        subprocess.run(
            [
                "/usr/bin/libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(render_dir),
                str(candidate),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        if not (render_dir / candidate.with_suffix(".pdf").name).is_file():
            raise RuntimeError("LibreOffice render check failed")
        shutil.copy2(candidate, BOOK)

    print(
        json.dumps(
            {
                "status": "updated_and_rendered",
                "workbook": str(BOOK),
                "overall_rows": [
                    first_overall,
                    first_overall + len(ORDER) - 1,
                ],
                "teacher_rows": [
                    first_teacher,
                    first_teacher + len(ORDER) - 1,
                ],
                "summary_start_row": summary_start,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
