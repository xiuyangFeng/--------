#!/usr/bin/env python3
"""Append the completed S3 regularization matrix to the PointNet workbook."""
from __future__ import annotations

import json
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from training_wss_min.tools.update_s3_rootcause_xlsx import (
    copy_row_style,
    load_json,
)


ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
ANALYSIS = ROOT / "training_wss_min/preflight/s3_regularization_results_analysis.json"
SECTION_TITLE = "2026-07-24｜S3-GEOPE 正则化三种子矩阵（vs 同 seed S3 父模板）"
VARIANTS = (
    "head_dropout01",
    "droppath005",
    "droppath010",
    "neighbordrop005",
)
SEEDS = (1234, 7, 2025)
ORDER = tuple(f"{variant}_s{seed}" for variant in VARIANTS for seed in SEEDS)
CHANGE = {
    "head_dropout01": ("REG-H S3+HeadDrop0.10", "S3 + head dropout=0.10"),
    "droppath005": ("REG-P05 S3+DropPath0.05", "S3 + linear DropPath max=0.05"),
    "droppath010": ("REG-P10 S3+DropPath0.10", "S3 + linear DropPath max=0.10"),
    "neighbordrop005": (
        "REG-N05 S3+NeighborDrop0.05",
        "S3 + preserve-nearest NeighborDrop=0.05",
    ),
}


def label_of(record: dict) -> str:
    return f"{CHANGE[record['variant']][0]} s{record['seed']}"


def matrix_values(record: dict) -> list:
    raw = load_json(Path(record["run_dir"]) / "eval/ckpt_best/metrics.json")["test"]
    norm = raw["normalized"]
    return [
        label_of(record),
        f"138/0/36（AG/AAA/ILO mixed；seed={record['seed']}）",
        "PointNeXt-R + LocalGeoPE",
        CHANGE[record["variant"]][1],
        "random / SAME；FPS center",
        5000,
        raw["field_casebalanced"]["r2"],
        raw["field"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["aggregate"]["r2_casemed"],
        raw["aggregate"]["r2_casep10"],
        f"{raw['aggregate']['r2_negative_cases']}/{raw['aggregate']['n_cases']}",
        raw["field"]["rmse"],
        raw["field"]["mae"],
        raw["field"]["nrmse_range"],
        raw["aggregate"]["nrmse_casemean"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        raw["calibration"]["top10_pred_true_ratio"],
        raw["hotspot"]["top10_iou_casemean"],
        norm["field_casebalanced"]["r2"],
        norm["field"]["r2"],
        norm["aggregate"]["r2_casemean"],
        norm["aggregate"]["r2_casemed"],
        norm["aggregate"]["r2_casep10"],
        norm["field_casebalanced"]["mae"],
        norm["field_casebalanced"]["rmse"],
        norm["field"]["nrmse_range"],
        norm["aggregate"]["nrmse_casemean"],
        norm["field"].get("nmae_range"),
        norm["aggregate"].get("nmae_casemean"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["spearman_high_wss_casemean"],
        norm["calibration"]["top10_pred_true_ratio"],
        norm["calibration"]["p99_pred_true_ratio"],
        record["experiment_id"],
    ]


def teacher_values(record: dict) -> list:
    raw = load_json(Path(record["run_dir"]) / "eval/ckpt_best/metrics.json")["test"]
    norm = raw["normalized"]
    return [
        "S3正则化",
        label_of(record),
        "PointNeXt-R + LocalGeoPE",
        "5000 / SAME",
        raw["field_casebalanced"]["r2"],
        raw["aggregate"]["r2_casemean"],
        raw["field"]["rmse"],
        raw["field"].get("nmae_range"),
        raw["aggregate"].get("nmae_casemean"),
        raw["regional_field"]["high_wss"]["r2"],
        norm["field_casebalanced"]["r2"],
        norm["field"].get("nmae_range"),
        raw["hotspot"]["spearman_all_casemean"],
        raw["hotspot"]["top10_iou_casemean"],
        norm["calibration"]["p99_pred_true_ratio"],
    ]


def append_summary(sheet, analysis: dict) -> int:
    records = analysis["records"]
    row = sheet.max_row + 1
    start = row
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    title = sheet.cell(row, 1, SECTION_TITLE)
    title.font = Font(bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E78")
    row += 1
    headers = [
        "处理臂",
        "seed",
        "R²_cb",
        "ΔR²_cb",
        "ΔMAE",
        "Δhigh-WSS R²",
        "ΔAG",
        "ΔAAA",
        "ΔILO",
        "病例胜/负",
        "病例均值Δ 95%CI",
        "口径",
    ]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    pairs = {
        result["treatment"]: result for result in analysis["paired_comparisons"]
    }
    for experiment_id in ORDER:
        row += 1
        record = records[experiment_id]
        result = pairs[experiment_id]
        delta = result["aggregate_treatment_minus_control"]
        domain = result["domain_r2_delta"]
        case = result["paired_case_stats"]["case_r2"]
        values = [
            label_of(record),
            record["seed"],
            record["best"]["physical_r2_casebalanced"],
            delta["physical_r2_casebalanced"],
            delta["physical_mae"],
            delta["high_wss_r2"],
            domain["AG"],
            domain["AAA"],
            domain["ILO"],
            f"{case['wins']}/{case['losses']}",
            (
                f"{case['mean_delta']:+.4f} "
                f"[{case['ci95_low']:+.4f}, {case['ci95_high']:+.4f}]"
            ),
            "ckpt_best(train_loss)",
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(row, column, value)
    row += 1
    sheet.cell(row, 1, "三种子汇总").font = Font(bold=True, italic=True)
    row += 1
    summary_headers = [
        "变体",
        "mean ΔR²_cb",
        "sd",
        "正/负",
        "mean ΔMAE",
        "mean Δhigh-WSS",
        "mean ΔAG",
        "mean ΔAAA",
        "mean ΔILO",
        "决策",
    ]
    for column, value in enumerate(summary_headers, 1):
        sheet.cell(row, column, value).font = Font(bold=True)
    summaries = {row["variant"]: row for row in analysis["three_seed_summaries"]}
    for variant in VARIANTS:
        row += 1
        summary = summaries[variant]
        means = summary["metric_delta_means"]
        domains = summary["domain_delta_means"]
        values = [
            CHANGE[variant][0],
            summary["delta_r2cb_mean"],
            summary["delta_r2cb_sd"],
            f"{summary['r2_wins']}/{summary['r2_losses']}",
            means["physical_mae"],
            means["high_wss_r2"],
            domains["AG"],
            domains["AAA"],
            domains["ILO"],
            summary["decision"],
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(row, column, value)
    return start


def main() -> None:
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    records = analysis["records"]
    for experiment_id in ORDER:
        if experiment_id not in records:
            raise RuntimeError(f"analysis missing {experiment_id}")
        if records[experiment_id]["integrity"] != "passed":
            raise RuntimeError(f"integrity failed: {experiment_id}")

    with tempfile.TemporaryDirectory(prefix="wss_s3reg_xlsx_") as temp_dir:
        candidate = Path(temp_dir) / XLSX.name
        shutil.copy2(XLSX, candidate)
        workbook = load_workbook(candidate, data_only=False)

        overall = workbook["实验矩阵总览"]
        existing_ids = {cell.value for cell in overall["AK"]}
        duplicates = set(ORDER) & existing_ids
        if duplicates:
            raise RuntimeError(f"refusing duplicate append: {sorted(duplicates)}")
        first_overall = overall.max_row + 1
        for experiment_id in ORDER:
            row = overall.max_row + 1
            copy_row_style(overall, row - 1, row, 37)
            for column, value in enumerate(
                matrix_values(records[experiment_id]), 1
            ):
                overall.cell(row, column, value)

        teacher = workbook["教师汇报视图"]
        labels = [label_of(records[experiment_id]) for experiment_id in ORDER]
        existing_labels = {cell.value for cell in teacher["B"]}
        if any(label in existing_labels for label in labels):
            raise RuntimeError("refusing duplicate append in 教师汇报视图")
        first_teacher = teacher.max_row + 1
        for experiment_id in ORDER:
            row = teacher.max_row + 1
            copy_row_style(teacher, row - 1, row, 15)
            for column, value in enumerate(
                teacher_values(records[experiment_id]), 1
            ):
                teacher.cell(row, column, value)

        summary = workbook["汇总对比"]
        if SECTION_TITLE in {cell.value for cell in summary["A"]}:
            raise RuntimeError("refusing duplicate append in 汇总对比")
        summary_start = append_summary(summary, analysis)
        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(candidate)

        check = load_workbook(candidate, read_only=True, data_only=False)
        check_ids = {
            row[0].value
            for row in check["实验矩阵总览"].iter_rows(
                min_col=37, max_col=37
            )
        }
        if not set(ORDER).issubset(check_ids):
            raise RuntimeError("saved workbook failed matrix readback")
        check_labels = {
            row[0].value
            for row in check["教师汇报视图"].iter_rows(min_col=2, max_col=2)
        }
        if not set(labels).issubset(check_labels):
            raise RuntimeError("saved workbook failed teacher readback")
        check_titles = {
            row[0].value
            for row in check["汇总对比"].iter_rows(min_col=1, max_col=1)
        }
        if SECTION_TITLE not in check_titles:
            raise RuntimeError("saved workbook failed summary readback")
        check.close()
        shutil.copy2(candidate, XLSX)

    print(json.dumps({
        "status": "updated",
        "workbook": str(XLSX),
        "overall_rows": [first_overall, first_overall + len(ORDER) - 1],
        "teacher_rows": [first_teacher, first_teacher + len(ORDER) - 1],
        "summary_start_row": summary_start,
        "rows_added": len(ORDER),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
