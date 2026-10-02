"""Idempotent workbook backfill for pf6vf6_v52d_retrain_20261002 (PF6 pressure / VF6 velocity on v5.2d; 01 block §43–§44).

Owns exactly the rows whose column A equals GROUP on sheet ``速度与压力实验矩阵`` (135 columns, table ``VolumeResults``):

* per-run rows: 36 arms x {best, last} = 72 rows. CV5 arms read their held-out fold, full265 arms read recover8.
  Column semantics follow the existing pressure / velocity rows (6-125 = peak-frame metrics of evaluate.py).
* summary rows (column D = 「汇总」), from ``readout_ckpt_best.json`` and ``compare_old_test34/compare_old_new.json``:
  per target CV5 out-of-fold pooled per seed + three-seed ensemble (261 units), recover8 full265 ensemble / E3 15-model
  ensemble / E5 deployed v5.0 ensemble, and the paired old-vs-new comparison on test34 (same units, same v5.2d labels).
  Only the test34 pair fills ΔR² (column M); every other row leaves it empty because the evaluation sets differ.

    python -m training_wss_min.tools.update_pf6vf6_v52d_xlsx [--workbook PATH]
"""
from __future__ import annotations

import argparse
import copy
import json
import os
import tempfile
import time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

from training_wss_min import config as C
from training_wss_min.tools.report_wss_recovery import cell_record, get_metric
from training_wss_min.tools.update_wss_local_wave2_volume_xlsx import COLUMNS, sha

ROOT = C.PROJECT_ROOT
NAME = "pf6vf6_v52d_retrain_20261002"
CONFIGS = ROOT / "training_wss_min/configs" / NAME
EXP = ROOT / "training_wss_min/experiments" / NAME
RUNS = ROOT / "training_wss_min/runs" / NAME
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
SHEET = "速度与压力实验矩阵"
GROUP = "体场 PF6/VF6 v5.2d 复验｜2026-10-02（CV5×3 seed + full265×3 seed，recover8；10-03 上线）"
MISSING = "—"
TRACK = "docs/02-推进与变更/01-X5D主线与新数据/X5D主线_实验跟踪.md §43–§44"
REQUIRED_SHEETS = ("WSS实验矩阵", "速度与压力实验矩阵", "指标说明")


def read_json(path, default=None):
    path = Path(path)
    return json.loads(path.read_text()) if path.is_file() else default


def run_rows():
    matrix = read_json(CONFIGS / "matrix.json")
    master = read_json(EXP / "queue_status.json", {"arms": {}})
    node04 = read_json(EXP / "queue_status_node04.json", {"arms": {}})
    rows = []
    for arm in matrix["arms"]:
        aid, kind = arm["id"], arm["arm"][0]
        config = read_json(CONFIGS / arm["config"])
        on_node04 = aid in node04["arms"]
        if on_node04:
            rec = node04["arms"][aid]
            status = "complete" if rec.get("next") is None else str(rec.get("next"))
            host = "node04 A100（Slurm 之外）"
        else:
            rec = master["arms"].get(aid, {})
            status = rec.get("status", "pending")
            host = f"master 4090（Slurm {master.get('job_id', '—')}）"
        for ck in arm["evaluate"]:
            path = RUNS / aid / f"eval/ckpt_{ck}/metrics.json"
            payload = read_json(path)
            rows.append(dict(key=(config["name"], ck), id=aid, kind=kind, arm=arm, checkpoint=ck, status=status, host=host,
                             metrics=payload["test"] if payload else None, metrics_path=str(path),
                             config_path=str(CONFIGS / arm["config"]), run_dir=str(RUNS / aid),
                             diagnostics=read_json(RUNS / aid / "training_diagnostics.json", {}),
                             accepted=status == "complete" and payload is not None, summary=False))
    return matrix, rows


def summary_rows():
    r = read_json(EXP / "readout_ckpt_best.json")
    cmp_ = read_json(EXP / "compare_old_test34/compare_old_new.json")
    out = []
    for t in ("PF6", "VF6"):
        kind = t[0]
        key = "r2cb_all" if t == "PF6" else "r2cb_speed"
        for s, v in r["A"][t]["per_seed"].items():
            out.append(dict(key=(f"{NAME}/summary/{t}_cv5_oof_s{s}", "汇总"), kind=kind, label=f"{t} CV5 折外 pooled · seed {s}",
                            scope="v5.2d CV5 五折折外（261 例）", s=v["summary"], note=f"seed {s} 的五个折模型各自的留出折合并（{v['folds_done']}/5 折）。"))
        e = r["A"][t]["ensemble3"]
        sd = r["A"][t]["seed_mean_sd"][key]
        out.append(dict(key=(f"{NAME}/summary/{t}_cv5_oof_ensemble3", "汇总"), kind=kind, label=f"{t} CV5 折外 · 三 seed 集成",
                        scope="v5.2d CV5 五折折外（261 例）", s=e,
                        note=f"每例由三个没见过它的折模型（seed 1234 / 7 / 2025）逐点平均"
                             + ("（速度先平均向量再取速率）" if t == "VF6" else "") + f"；单 seed 均值 {sd['mean']:.4f} ± {sd['sd']:.4f}。"))
        rows = {x["model"]: x for x in r["B"][t]["rows"]}
        for label, model, note in (("recover8 · full265 三 seed 集成", "**full265 三 seed 集成**", "全量 265 例三 seed（上线权重 PF6_VF6_v52d_3seed_20261003）。"),
                                   ("recover8 · E3 15 个 CV5 折模型集成", "**E3 15 模型集成**", "CV5 五折 × 三 seed 的 15 个折模型逐点平均。"),
                                   ("recover8 · E5 已部署 v5.0 三 seed 集成", "**E5 已部署 v5.0 三 seed 集成**", "旧发布包 PF6_VF6_peak_3seed_20260920（v5.0，train138，10-03 已下线）。")):
            out.append(dict(key=(f"{NAME}/summary/{t}_{model.strip('*').split()[0]}_recover8", "汇总"), kind=kind, label=f"{t} {label}",
                            scope="recover8（8 例，只作描述）", s=rows[model], note=note + "recover8 只有 8 例，只作描述。"))
        c = cmp_[t]["test34"]
        old_v, new_v = c["old_ens"][key], c["new_ens"][key]
        pr = c["paired"]
        out.append(dict(key=(f"{NAME}/summary/{t}_test34_old_v50", "汇总"), kind=kind, label=f"{t} test34 · 旧 v5.0 三 seed 集成（v5.2d 标签重评）",
                        scope="test34（34 例，旧模型没见过）· v5.2d 标签", s=c["old_ens"],
                        note="已部署 v5.0 权重在 v5.2d 标签上重评（作业 16973）；test34 是旧模型的开发暴露集，对旧模型略偏乐观。"))
        out.append(dict(key=(f"{NAME}/summary/{t}_test34_new_cv5_oof", "汇总"), kind=kind, label=f"{t} test34 · 新 v5.2d CV5 折外三 seed 集成",
                        scope="test34（34 例）· v5.2d 标签 · 同单元配对", s=c["new_ens"], delta=(new_v - old_v, f"{t} test34 旧 v5.0 三 seed 集成（同单元同标签）"),
                        note=f"同一批 34 例、同一套 v5.2d 标签的配对比较：{pr['better']}/{pr['n']} 例改善，逐例 R² 中位差 {pr['median_delta_unit_r2']:+.4f}；"
                             "差值混着训练数据量（每折约 209 例对 138 例）与数据修正，不单独归因。"))
    return out


def run_values(item):
    kind, aid, arm = item["kind"], item["id"], item["arm"]
    values = [MISSING] * 135
    values[0] = GROUP
    fold = arm.get("fold")
    target = "压力" if kind == "P" else "速度"
    proto = (f"CV5 fold{fold}（训练其余四折，读留出折）" if fold is not None else "full265（训练 265 例，读 recover8）")
    values[1] = f"{aid}｜{target} {arm['arm']} 原配方只换 v5.2d 数据：{proto}，seed {arm['seed']}｜{item['status']}"
    values[2] = "压力 · Pa" if kind == "P" else "速度 · m/s"
    values[3] = item["checkpoint"]
    values[4] = ((f"v5.2d CV5 fold{fold} 留出折" if fold is not None else "recover8（8 例）")
                 + (" · 壁面∪内部" if kind == "P" else " · 内部"))
    sources = [item["config_path"], str(Path(item["run_dir"]) / "history.jsonl"), str(EXP / f"readout_ckpt_{item['checkpoint']}.md")]
    note = [f"{item['host']}；{item['status']}；配方逐字不变（seed 1234 源自 wave2、7 / 2025 源自 wave3），只改数据路径、分区与统计。",
            "数据 = v5.2d 体场数据根 wss_min_volview_v1；"
            + (f"CV5 患者分组五折 fold{fold}，统计与特征归一化用该折训练分区。" if fold is not None
               else "全量 265 例训练（无验证集，按 train_loss 选模），recover8 8 例从未参与训练，只作描述。"),
            "与 v5.0 test34、v5.1 cv3 的旧体场行不是同一评估集，只并列、不相减（同口径的新旧对比见本组 test34 汇总行）。",
            f"已知数据限制：150 个每步迭代提前退出的单元标签保持原样。来源 {TRACK}。"]
    if item["accepted"]:
        m = item["metrics"]
        for col, dotted in COLUMNS.items():
            value = get_metric(m, dotted)
            if value is not None:
                values[col - 1] = value
        values[11] = f"{int(m['aggregate']['r2_negative_cases'])} / {int(m['aggregate']['n_cases'])}"
        values[89] = item["diagnostics"].get("elapsed_seconds", MISSING)
        values[90] = item["diagnostics"].get("peak_cuda_allocated_mb", MISSING)
        peak = get_metric(m, "efficiency.peak_cuda_memory_bytes")
        values[91] = peak / 2 ** 20 if peak is not None else MISSING
        values[13] = "无同评估集参照（新数据基线；旧体场数字评估集不同，不相减）"
        values[92] = "建立基线，无过门判据"
        values[93] = "上线权重（full265）" if fold is None else "CV5 折模型（折外读数）"
        if kind == "P" and "query_groups" in m:
            note.append(f"壁面 R²_cb {m['query_groups']['wall']['field_casebalanced']['r2']:.4f}、"
                        f"内部 {m['query_groups']['interior']['field_casebalanced']['r2']:.4f}。")
    else:
        note.append("本行尚无可接受的评估产物。")
    values[126] = "Pa" if kind == "P" else "m/s"
    values[127] = ("PointNeXt-R + LocalGeoPE + L-SA2 + QAD-lite；"
                   + ("PF6：原 SA3 后全局 BT ＋ Murray 先验" if kind == "P" else "VF6：新 SA3 单半径 0.20 加 BT ＋ Murray 先验"))
    values[128] = f"{arm['arm']} 原配方（峰值帧），只换 v5.2d 数据"
    values[129] = "400 epoch；support 随机5000；独立 query；20D 输入；" + ("相对压力 p−体积均值、全局线性 z、MSE" if kind == "P" else "速度三分量、全局线性 z、MSE")
    values[130] = 5000
    values[131] = item["key"][0]
    sources.append(item["metrics_path"])
    values[132] = " ".join(note)[:32000]
    values[133] = "\n".join(sources)
    values[134] = "[]"
    return values, " ".join(note), sources


def summary_values(item):
    kind, s = item["kind"], item["s"]
    values = [MISSING] * 135
    values[0] = GROUP
    values[1] = item["label"]
    values[2] = "压力 · Pa" if kind == "P" else "速度 · m/s"
    values[3] = "汇总"
    values[4] = item["scope"] + (" · 壁面∪内部" if kind == "P" else " · 内部")
    values[5] = s["r2cb_all"] if kind == "P" else s["r2cb_speed"]
    values[9] = s.get("unit_median_r2", MISSING)
    if kind == "V":
        values[25] = s.get("vector_rmse_m_s", MISSING)
        values[26] = s.get("direction_cosine_casemean", MISSING)
        values[27] = s.get("direction_cosine_speedweighted_casemean", MISSING)
        for col, comp in ((30, "u"), (33, "v"), (36, "w")):
            values[col] = s.get(f"r2cb_{comp}", MISSING)
    if "delta" in item:
        values[12], values[13] = item["delta"]
    else:
        values[13] = "无（不同评估集不相减）"
    values[92] = "建立基线，无过门判据"
    values[93] = "汇总（读数工具重算，与 metrics.json 一致性检查 0 差）"
    note = [item["note"], "指标由 tools/report_pf6vf6_v52d_retrain.py 从保存的预测按病例等权重算（只填 R²_cb、逐例中位"
            + ("、向量误差与方向余弦、分量 R²_cb" if kind == "V" else "") + "）。"]
    if kind == "P" and "r2cb_wall" in s:
        note.append(f"壁面 R²_cb {s['r2cb_wall']:.4f}、内部 {s['r2cb_interior']:.4f}。")
    values[126] = "Pa" if kind == "P" else "m/s"
    values[131] = item["key"][0]
    sources = [str(EXP / "readout_ckpt_best.json"), str(EXP / "compare_old_test34/compare_old_new.json")]
    values[132] = " ".join(note)[:32000]
    values[133] = "\n".join(sources)
    values[134] = "[]"
    return values, " ".join(note), sources


def update_workbook(book: Path):
    matrix, runs = run_rows()
    items = runs + summary_rows()
    before = sha(book)
    wb = load_workbook(book)
    if any(name not in wb.sheetnames for name in REQUIRED_SHEETS):
        raise ValueError("unexpected workbook sheet schema")
    ws = wb[SHEET]
    if ws.max_column != 135 or ws["EB5"].value != "原始实验 ID" or "VolumeResults" not in ws.tables:
        raise ValueError("unexpected volume sheet 135-column schema")
    owned = [r for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value == GROUP]
    if owned and len(owned) != len(items):
        raise ValueError(f"owned rows must number exactly {len(items)}")
    keys = [it["key"] for it in items]
    if len(set(keys)) != len(keys):
        raise ValueError("duplicate row identities")
    locations = ({(ws.cell(r, 132).value, ws.cell(r, 4).value): r for r in owned} if owned else
                 {key: ws.max_row + i + 1 for i, key in enumerate(keys)})
    if set(locations) != set(keys):
        raise ValueError("row identities differ from the workbook's owned rows")
    preserved = {(s.title, c.coordinate): cell_record(c) for s in wb for row in s for c in row
                 if not (s.title == ws.title and c.row in owned)}
    dimensions = {s.title: (copy.deepcopy(s.column_dimensions), copy.deepcopy(s.row_dimensions)) for s in wb}
    merges = {s.title: tuple(map(str, s.merged_cells.ranges)) for s in wb}
    for it in items:
        row = locations[it["key"]]
        values, note, sources = summary_values(it) if "label" in it else run_values(it)
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col)
            cell._style = copy.copy(ws.cell(6 if row % 2 == 0 else 7, col)._style)
            cell.value = value
            cell.comment = None
            cell.hyperlink = None
        ws.cell(row, 133).comment = Comment(note[:32000], "pf6vf6 v5.2d report")
        ws.cell(row, 134).comment = Comment("\n".join(sources), "pf6vf6 v5.2d report")
        ws.row_dimensions[row].height = 72
    ws.tables["VolumeResults"].ref = f"A5:EE{ws.max_row}"
    ws.tables["VolumeResults"].autoFilter.ref = f"A5:EE{ws.max_row}"
    ws.print_area = f"A1:N{ws.max_row}"
    fd, temp = tempfile.mkstemp(dir=book.parent, prefix=f".{NAME}_", suffix=".xlsx")
    os.close(fd)
    backup = None
    try:
        wb.save(temp)
        checked = load_workbook(temp)
        for (sheet, coordinate), record in preserved.items():
            if cell_record(checked[sheet][coordinate]) != record:
                raise AssertionError(f"historical cell changed: {sheet}!{coordinate}")
        assert {s.title: tuple(map(str, s.merged_cells.ranges)) for s in checked} == merges
        for sheet in checked:
            old_cols, old_rows = dimensions[sheet.title]
            assert {k: dict(v) for k, v in sheet.column_dimensions.items()} == {k: dict(v) for k, v in old_cols.items()}
            for key, value in old_rows.items():
                if sheet.title != ws.title or key not in owned:
                    assert dict(sheet.row_dimensions[key]) == dict(value)
        for it in items:
            r = locations[it["key"]]
            assert checked[SHEET].cell(r, 1).value == GROUP and checked[SHEET].cell(r, 132).value == it["key"][0]
        checked.close()
        if sha(book) != before:
            raise RuntimeError("workbook changed concurrently; refusing replacement")
        backup = book.parent / f"{book.stem}_backup_before_{NAME}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        backup.write_bytes(book.read_bytes())
        os.replace(temp, book)
    finally:
        if Path(temp).exists():
            Path(temp).unlink()
    evidence = {"passed": True, "before_sha256": before, "after_sha256": sha(book), "backup": str(backup),
                "records": len(items), "run_rows": len(runs), "summary_rows": len(items) - len(runs),
                "accepted_run_rows": sum(r["accepted"] for r in runs), "historical_cells_verified": len(preserved),
                "table_ref": ws.tables["VolumeResults"].ref, "row_span": [min(locations.values()), max(locations.values())]}
    (EXP / "xlsx_acceptance.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
    return evidence


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=BOOK)
    args = parser.parse_args(argv)
    print(json.dumps(update_workbook(args.workbook), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
