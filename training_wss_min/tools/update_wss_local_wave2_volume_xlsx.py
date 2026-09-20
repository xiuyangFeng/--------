"""Idempotent workbook backfill for the wave-2 volume arms (F6 Murray prior into pressure/velocity).

Owns exactly the 8 rows (4 arms × best/last) whose column A equals GROUP on sheet
``速度与压力实验矩阵`` (135 columns, table ``VolumeResults``).  Column semantics follow the
existing pressure/velocity rows (same header order as the 2026-09-11 rebuild); the long-wave
diagnostic and the axial/radial/circumferential decomposition are not produced by the
standard evaluation and stay ``—``.  Every historical cell (value, style, comment, hyperlink),
merged range and dimension of every other row/sheet is verified unchanged after saving.

    python -m training_wss_min.tools.update_wss_local_wave2_volume_xlsx [--workbook PATH] [--name NAME --group GROUP]

With ``--name`` (another experiment, e.g. wave 3) the arm list, target kind, parent pairing and descriptions
are derived from that experiment's ``matrix.json`` (fields ``base``/``kind``/``parent``/``modules``/``seed``).
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import tempfile
import time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

from training_wss_min import config as C
from training_wss_min.tools.report_wss_recovery import cell_record, get_metric

ROOT = C.PROJECT_ROOT
NAME = "wss_local_wave2_20260912"
CONFIGS = ROOT / "training_wss_min/configs" / NAME
EXP = ROOT / "training_wss_min/experiments" / NAME
RUNS = ROOT / "training_wss_min/runs"
HIST = RUNS / "volume_attention_20260910"
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
SHEET = "速度与压力实验矩阵"
GROUP = "波2 F6接体场｜2026-09-12（P02/V07＋Murray分支流量先验，含同期对照）"
MISSING = "—"
ARMS = ["PF6_s1234", "P02r_s1234", "VF6_s1234", "V07r_s1234"]
DESCRIPTIONS = {
    "PF6_s1234": "P02（压力最优：SA3后全局BT）＋F6 Murray分支流量先验（内部单元按所在血管段广播log Q份额、log τ0用单元处atlas半径）",
    "P02r_s1234": "P02同配置同期重跑（对照）",
    "VF6_s1234": "V07（速度最优：新SA3单半径0.20＋BT）＋F6 Murray分支流量先验",
    "V07r_s1234": "V07同配置同期重跑（对照）",
}
STRUCTURE = {"PF6_s1234": "P02＋F6（输入+2维）", "P02r_s1234": "P02原样重跑", "VF6_s1234": "V07＋F6（输入+2维）", "V07r_s1234": "V07原样重跑"}
KIND = {"PF6_s1234": "P", "P02r_s1234": "P", "VF6_s1234": "V", "V07r_s1234": "V"}
# (reference label, run dir) — first entry drives the ΔR² column
REFERENCES = {
    "PF6_s1234": [("P02r_s1234（同期对照）", RUNS / NAME / "P02r_s1234"), ("P02历史", HIST / "P02_s1234")],
    "P02r_s1234": [("P02历史", HIST / "P02_s1234")],
    "VF6_s1234": [("V07r_s1234（同期对照）", RUNS / NAME / "V07r_s1234"), ("V07历史", HIST / "V07_s1234")],
    "V07r_s1234": [("V07历史", HIST / "V07_s1234")],
}
STATUS_TEXT = {"complete": "已完成（队列退出码0；best/last评估齐全）", "failed": "失败", "blocked": "阻塞",
               "cancelled": "取消", "waiting": "等待依赖", "pending": "等待执行"}

# one-based workbook column → dotted path inside metrics.json["test"]
COLUMNS = {
    6: "field_casebalanced.r2", 7: "field.rmse", 8: "field.mae", 9: "field.nmae_range",
    10: "aggregate.r2_casemed", 11: "aggregate.r2_casep10",
    15: "field.r2", 16: "aggregate.r2_casemean", 17: "field.nrmse_range", 18: "aggregate.nrmse_casemean",
    19: "aggregate.nmae_casemean", 20: "regional_field.high_wss.r2", 21: "calibration.top10_pred_true_ratio",
    22: "hotspot.top10_iou_casemean", 23: "calibration.p99_pred_true_ratio",
    24: "efficiency.parameters", 25: "efficiency.elapsed_seconds",
    26: "vector.vector_rmse_m_s", 27: "vector.direction_cosine_casemean",
    28: "vector.direction_cosine_speedweighted_casemean", 29: "vector.component_normalized_r2_pooled",
    30: "vector.speed_floor_m_s",
    31: "vector.components.u.r2_casebalanced", 32: "vector.components.u.r2_pooled", 33: "vector.components.u.rmse",
    34: "vector.components.v.r2_casebalanced", 35: "vector.components.v.r2_pooled", 36: "vector.components.v.rmse",
    37: "vector.components.w.r2_casebalanced", 38: "vector.components.w.r2_pooled", 39: "vector.components.w.rmse",
    40: "field.n", 41: "field_casebalanced.mae", 42: "field_casebalanced.rmse",
    43: "aggregate.mae_casemean", 44: "aggregate.rmse_casemean",
    45: "hotspot.spearman_all_casemean", 46: "hotspot.spearman_high_wss_casemean",
    47: "field_casebalanced.r2_linear_fit", 48: "field_casebalanced.linear_fit_slope",
    49: "field_casebalanced.linear_fit_intercept", 50: "field.r2_linear_fit", 51: "field.linear_fit_slope",
    52: "field.linear_fit_intercept",
    53: "normalized.field_casebalanced.r2", 54: "normalized.field.r2", 55: "normalized.aggregate.r2_casemean",
    56: "normalized.aggregate.r2_casemed", 57: "normalized.aggregate.r2_casep10", 58: "normalized.field.mae",
    59: "normalized.field.rmse", 60: "normalized.field_casebalanced.mae", 61: "normalized.field_casebalanced.rmse",
    62: "normalized.aggregate.mae_casemean", 63: "normalized.aggregate.rmse_casemean",
    64: "normalized.field.nrmse_range", 65: "normalized.aggregate.nrmse_casemean",
    66: "normalized.field.nmae_range", 67: "normalized.aggregate.nmae_casemean",
    68: "normalized.regional_field.high_wss.r2", 69: "normalized.calibration.top10_pred_true_ratio",
    70: "normalized.calibration.p99_pred_true_ratio", 71: "normalized.hotspot.top10_iou_casemean",
    72: "normalized.hotspot.spearman_all_casemean", 73: "normalized.hotspot.spearman_high_wss_casemean",
    74: "normalized.field_casebalanced.r2_linear_fit", 75: "normalized.field_casebalanced.linear_fit_slope",
    76: "normalized.field_casebalanced.linear_fit_intercept", 77: "normalized.field.r2_linear_fit",
    78: "normalized.field.linear_fit_slope", 79: "normalized.field.linear_fit_intercept",
    80: "normalized.field.mae", 81: "normalized.field.rmse",
    107: "vector.vector_rmse_m_s",
}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path, default=None):
    path = Path(path)
    return json.loads(path.read_text()) if path.is_file() else default


def metrics_of(run_dir: Path, checkpoint: str):
    path = run_dir / f"eval/ckpt_{checkpoint}/metrics.json"
    payload = read_json(path)
    return (payload["test"] if payload else None), str(path)


def primary_error(kind: str, metrics: dict):
    """Pressure: pooled MAE (Pa); velocity: vector RMSE (m/s) — the same choice as §17."""
    return get_metric(metrics, "field.mae" if kind == "P" else "vector.vector_rmse_m_s")


def derive_from_matrix(matrix: dict) -> None:
    """Fill ARMS / KIND / DESCRIPTIONS / STRUCTURE / REFERENCES from a wave-3-style matrix.json."""
    ARMS.clear()
    for arm in matrix["arms"]:
        aid, base, kind, seed = arm["id"], arm["base"], arm.get("kind") or arm["base"][0], arm.get("seed", 1234)
        ARMS.append(aid)
        KIND[aid] = kind
        with_f6 = "F6" in arm.get("modules", [])
        target = "压力最优：SA3后全局BT" if base == "P02" else "速度最优：新SA3单半径0.20＋BT"
        DESCRIPTIONS[aid] = (f"{base}（{target}）＋F6 Murray分支流量先验，seed {seed}" if with_f6
                             else f"{base}同配置同seed对照（seed {seed}，与F6臂共享继承初始权重）")
        STRUCTURE[aid] = f"{base}＋F6（输入+2维）" if with_f6 else f"{base}原样（seed {seed}）"
        refs = []
        if with_f6 and arm.get("parent"):
            refs.append((f"{arm['parent']}（同seed对照）", RUNS / NAME / arm["parent"]))
        refs.append((f"{base}历史（seed 1234）", HIST / f"{base}_s1234"))
        REFERENCES[aid] = refs


def arm_rows():
    matrix = read_json(CONFIGS / "matrix.json")
    if NAME != "wss_local_wave2_20260912":
        derive_from_matrix(matrix)
    queue = read_json(EXP / "queue_status.json", {"arms": {}})
    arms = {a["id"]: a for a in matrix["arms"]}
    rows = []
    for aid in ARMS:
        arm = arms[aid]
        config = read_json(CONFIGS / arm["config"])
        record = queue["arms"].get(aid, {})
        run = RUNS / config["name"]
        diagnostics = read_json(run / "training_diagnostics.json", {})
        kind = KIND[aid]
        for checkpoint in ("best", "last"):
            stage = record.get("stages", {}).get(f"eval_{checkpoint}", {})
            metrics, path = metrics_of(run, checkpoint)
            accepted = record.get("status") == "complete" and stage.get("returncode") == 0 and metrics is not None
            refs = []
            for label, ref_dir in REFERENCES[aid]:
                ref_metrics, ref_path = metrics_of(ref_dir, checkpoint)
                if ref_metrics is not None:
                    refs.append((label, ref_metrics, ref_path))
            rows.append(dict(id=aid, arm=arm, kind=kind, checkpoint=checkpoint, run_name=config["name"], run_dir=str(run),
                             config_path=str(CONFIGS / arm["config"]), metrics=metrics, metrics_path=path, refs=refs,
                             accepted=accepted, diagnostics=diagnostics, epochs=config["train"]["epochs"],
                             input_dim=len(config["data"]["input_features"]),
                             status=STATUS_TEXT.get(record.get("status", "pending"), record.get("status", "pending")),
                             queue_status=record.get("status", "pending")))
    return matrix, queue, rows


def row_values(item: dict, job_id) -> tuple[list, str, list[str]]:
    kind, aid = item["kind"], item["id"]
    values = [MISSING] * 135
    values[0] = GROUP
    values[1] = f"{aid}｜{DESCRIPTIONS[aid]}｜{item['status']}"
    values[2] = "压力 · Pa" if kind == "P" else "速度 · m/s"
    values[3] = item["checkpoint"]
    values[4] = "V5 train138/test34 · 壁面∪内部" if kind == "P" else "V5 train138/test34 · 内部"
    sources = [item["config_path"], str(Path(item["run_dir"]) / "history.jsonl"), str(EXP / "results.json")]
    seed = item["arm"].get("seed", 1234)
    note = (f"Job {job_id}；{item['status']}；{item['arm']['title']}。单seed{seed}、已暴露test34：只作筛选，不作显著性或泛化结论。"
            f"init_reference_config={'P02' if kind == 'P' else 'V07'}_s{seed}（继承张量初始权重逐位相同，新增F6输入列零初始化）；"
            "压力标量为 p−病例峰值步体积均值（Pa）；速度标量为 |u|（m/s）。长波诊断与轴向/径向/周向分解未在本轮评估中产出。")
    if item["accepted"]:
        m = item["metrics"]
        for col, dotted in COLUMNS.items():
            value = get_metric(m, dotted)
            if value is not None:
                values[col - 1] = value
        values[11] = f"{int(m['aggregate']['r2_negative_cases'])} / {int(m['aggregate']['n_cases'])}"
        values[89] = item["diagnostics"].get("elapsed_seconds", MISSING)
        values[90] = item["diagnostics"].get("peak_cuda_allocated_mb", MISSING)
        peak = get_metric(m, "efficiency.peak_cuda_memory_bytes")
        values[91] = peak / 2 ** 20 if peak is not None else MISSING
        own_r2 = get_metric(m, "field_casebalanced.r2")
        own_err = primary_error(kind, m)
        for j, (label, ref_metrics, ref_path) in enumerate(item["refs"][:2]):
            base = 94 + j * 6
            ref_r2 = get_metric(ref_metrics, "field_casebalanced.r2")
            ref_err = primary_error(kind, ref_metrics)
            values[base] = f"{label}｜{'压力' if kind == 'P' else '速度'}｜{item['checkpoint']}"
            values[base + 1] = own_r2 - ref_r2 if (own_r2 is not None and ref_r2 is not None) else MISSING
            values[base + 2] = ((ref_err - own_err) / ref_err if (own_err is not None and ref_err) else MISSING)
            values[base + 5] = "单seed筛选，无门槛判定"
            sources.append(ref_path)
            if j == 0:
                values[12] = values[base + 1]
                values[13] = label
        values[92] = "单seed筛选（同checkpoint配对同期对照与历史臂）"
        values[93] = ("压力：F6接入" if aid.startswith("PF6") else "速度：F6接入" if aid.startswith("VF6")
                      else "同期/同seed对照")
        sources.append(item["metrics_path"])
    values[126] = "Pa" if kind == "P" else "m/s"
    values[127] = ("PointNeXt-R + LocalGeoPE + L-SA2 + QAD-lite；" +
                   ("P02：原SA3后全局BT" if kind == "P" else "V07：新SA3单半径0.20加BT") +
                   ("＋F6 Murray先验" if aid.startswith(("PF6", "VF6")) else "（同配置对照）"))
    values[128] = STRUCTURE[aid]
    values[129] = (f"{item['epochs']} epoch；support 随机5000；独立 query；{item['input_dim']}D 输入；"
                   + ("相对压力线性 z、MSE" if kind == "P" else "速度三分量逐分量线性 z、MSE"))
    values[130] = 5000
    values[131] = item["run_name"]
    values[132] = note[:32000]
    values[133] = "\n".join(sources)
    values[134] = "[]"
    return values, note, sources


def update_workbook(book: Path):
    matrix, queue, rows = arm_rows()
    before = sha(book)
    wb = load_workbook(book)
    if wb.sheetnames[:3] != ["WSS实验矩阵", "速度与压力实验矩阵", "指标说明"]:
        raise ValueError("unexpected workbook sheet schema")
    ws = wb[SHEET]
    if ws.max_column != 135 or ws["EB5"].value != "原始实验 ID" or "VolumeResults" not in ws.tables:
        raise ValueError("unexpected volume sheet 135-column schema")
    owned = [r for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value == GROUP]
    if owned and len(owned) != len(rows):
        raise ValueError(f"owned rows must number exactly {len(rows)}")
    keys = [(item["run_name"], item["checkpoint"]) for item in rows]
    locations = ({(ws.cell(r, 132).value, ws.cell(r, 4).value): r for r in owned} if owned else
                 {key: ws.max_row + i + 1 for i, key in enumerate(keys)})
    if set(locations) != set(keys):
        raise ValueError("run identities differ from the workbook's owned rows")
    preserved = {(s.title, c.coordinate): cell_record(c) for s in wb for row in s for c in row
                 if not (s.title == ws.title and c.row in owned)}
    dimensions = {s.title: (copy.deepcopy(s.column_dimensions), copy.deepcopy(s.row_dimensions)) for s in wb}
    merges = {s.title: tuple(map(str, s.merged_cells.ranges)) for s in wb}
    job_id = queue.get("job_id", "尚未提交")
    for item in rows:
        row = locations[(item["run_name"], item["checkpoint"])]
        values, note, sources = row_values(item, job_id)
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col)
            cell._style = copy.copy(ws.cell(6 if row % 2 == 0 else 7, col)._style)
            cell.value = value
            cell.comment = None
            cell.hyperlink = None
        ws.cell(row, 133).comment = Comment(note[:32000], "wave2 volume report")
        ws.cell(row, 134).comment = Comment("\n".join(sources), "wave2 volume report")
        ws.row_dimensions[row].height = 72
    ws.tables["VolumeResults"].ref = f"A5:EE{ws.max_row}"
    ws.tables["VolumeResults"].autoFilter.ref = f"A5:EE{ws.max_row}"
    ws.print_area = f"A1:N{ws.max_row}"
    fd, temp = tempfile.mkstemp(dir=book.parent, prefix=f".{NAME}_volume_", suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(temp)
        checked = load_workbook(temp)
        for (sheet, coordinate), record in preserved.items():
            if cell_record(checked[sheet][coordinate]) != record:
                raise AssertionError(f"historical cell changed: {sheet}!{coordinate}")
        assert {s.title: tuple(map(str, s.merged_cells.ranges)) for s in checked} == merges
        for sheet in checked:
            old_cols, old_rows = dimensions[sheet.title]
            assert {k: dict(v) for k, v in sheet.column_dimensions.items()} == {k: dict(v) for k, v in old_cols.items()}
            for key, value in old_rows.items():
                if sheet.title != ws.title or key not in owned:
                    assert dict(sheet.row_dimensions[key]) == dict(value)
        for item in rows:
            r = locations[(item["run_name"], item["checkpoint"])]
            assert checked[SHEET].cell(r, 1).value == GROUP and checked[SHEET].cell(r, 132).value == item["run_name"]
            if not item["accepted"]:
                for column in [*range(6, 14), *range(15, 127)]:
                    assert checked[SHEET].cell(r, column).value == MISSING
        checked.close()
        if sha(book) != before:
            raise RuntimeError("workbook changed concurrently; refusing replacement")
        backup = EXP / f"{book.stem}_backup_before_{NAME}_volume_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        backup.write_bytes(book.read_bytes())
        os.replace(temp, book)
    finally:
        if Path(temp).exists():
            Path(temp).unlink()
    evidence = {"passed": True, "before_sha256": before, "after_sha256": sha(book), "backup": str(backup),
                "rows": {f"{k[0]}@{k[1]}": v for k, v in locations.items()}, "records": len(rows),
                "accepted_result_rows": sum(r["accepted"] for r in rows),
                "historical_cells_verified": len(preserved), "table_ref": ws.tables["VolumeResults"].ref}
    (EXP / "xlsx_acceptance_volume.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
    return evidence


def main(argv=None):
    global NAME, CONFIGS, EXP, GROUP
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=BOOK)
    parser.add_argument("--name", default=NAME, help="experiment name (configs/<name>, experiments/<name>)")
    parser.add_argument("--group", default=GROUP, help="workbook column-A group label that owns the rows")
    args = parser.parse_args(argv)
    NAME, GROUP = args.name, args.group
    CONFIGS, EXP = ROOT / "training_wss_min/configs" / NAME, ROOT / "training_wss_min/experiments" / NAME
    print(json.dumps(update_workbook(args.workbook), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
