#!/usr/bin/env python3
"""Idempotent workbook sync for the V5.2 velocity-cycle development screen.

Read existing metrics only. Preserve every historical cell/style/formula except
three explicitly updated reading-index cells; atomically install after reopen
validation. The dedicated sheet uses per-unit means, never historical R2_cb.
"""
from __future__ import annotations

import argparse
import copy
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import statistics
import tempfile
import os

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[2]
EXP = ROOT / "training_wss_min/experiments/velocity_phase_v52_20260930"
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
SHEET = "V52速度周期优化"
INDEX = "指标说明"
ARMS = ("Iu", "U0", "D1", "D2", "A0", "A1", "G00", "G01", "G10", "G11")
STAGES = {"cycle": "整周期", "peak": "峰值窗", "trough": "谷底窗", "decel": "减速段"}
DESCRIPTIONS = {
    "Iu": "第一轮独立速度头；引用历史结果，未重训",
    "U0": "第二轮速度控制；引用历史结果，未重训",
    "D1": "U0＋有界方向辅助损失（训练侧梯度标定λ）",
    "D2": "U0＋每相位训练速度q80尾部双权重",
    "A0": "局部坐标架输入＋全局XYZ残差",
    "A1": "A0同输入/容量＋局部有符号残差旋回XYZ",
    "G00": "A0＋静态点式体内锚点",
    "G01": "A0＋相位条件点式体内锚点",
    "G10": "A0＋静态分支约束体内图",
    "G11": "A0＋相位条件分支约束体内图",
}
METRIC_NAMES = {"r2": "vector R²均值", "mae": "分量MAE m/s", "rmse": "分量RMSE m/s",
                "vector_rmse_m_s": "向量RMSE m/s", "direction_cosine": "方向余弦"}
HEADERS = ["实验臂", "方法", "记录性质", "状态"]
HEADERS += [f"{label} {METRIC_NAMES['r2']}" for label in STAGES.values()]
HEADERS += [f"{label} {METRIC_NAMES[k]}" for label in STAGES.values()
            for k in ("mae", "rmse", "vector_rmse_m_s", "direction_cosine")]
HEADERS += [f"Δ{label}R² vs U0" for label in STAGES.values()]
HEADERS += ["预注册主对照", "训练单元", "开发留出单元", "fold", "seed", "已完成epochs", "优化器更新数", "相位数", "checkpoint",
            "训练query", "评估query", "参数量", "训练秒", "缓存读入+80帧预测中位秒", "原提交作业", "实际执行作业", "运行目录", "配置路径", "原始metrics", "逐单元metrics", "数据split", "split SHA256", "stats SHA256", "备注"]
COL = {h: i + 1 for i, h in enumerate(HEADERS)}
GLOSSARY = [
    ("V5.2整周期专页", "本轮速度结果索引", f"{SHEET}：8臂D1/D2/A0/A1/G00/G01/G10/G11＋历史Iu/U0参照。206训练/55开发留出数据单元，患者分组fold0，seed1234，last150，80相位。", "新8次训练＋2条引用", "单seed开发筛选；55个数据单元不等同55名独立患者（含ILO before/after）。"),
    ("V5.2整周期专页", "vector R²与聚合", "每个数据单元内：体积权重归一化，相位等权，XYZ分量分别中心化计算vector R²；随后55个数据单元等权平均。", "无量纲，越高越好", "不是旧R²_cb，不是pooled，不是速度模长R²。所有格从既有per_case.jsonl复算均值并与metrics.json、report.json核对。"),
    ("V5.2整周期专页", "MAE/RMSE定义", "分量MAE=XYZ绝对误差均值；分量RMSE=sqrt(XYZ平方误差均值)；向量RMSE=sqrt(3)×分量RMSE。各项先逐单元计算，再等权平均。", "m/s，越低越好", "表中RMSE均值不等于sqrt(均值MSE)；不混入speed模长误差。"),
    ("V5.2整周期专页", "时间窗与方向", "整周期0–79；峰值窗17–26（非单峰帧）；谷底窗5–9及43–57；减速段27–42（均0基相位索引）。方向余弦仅真值速率≥0.01m/s，预测零向量计0。", "80个离散相位", "窗口按冻结训练波形定义；方向分母和有效覆盖范围沿原评估合同。"),
    ("V5.2整周期专页", "运行与引用边界", "Iu/U0引用已完成历史run，不计入本轮8次训练。原提交作业和实际执行作业分列，以execution.json记录恢复/重排后的实际来源。", "last150，无test选模", "推理秒数含缓存加载/传输/编码/80帧解码，不含原几何预处理/导出；不能直接称临床端到端时延。"),
]


def read(path):
    return json.loads(Path(path).read_text())


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def close(a, b):
    return isinstance(a, (int, float)) and isinstance(b, (int, float)) and math.isclose(a, b, abs_tol=1e-12, rel_tol=1e-10)


def cell_record(c):
    return (c.value, copy.copy(c._style), c.number_format, c.data_type,
            (c.comment.text, c.comment.author) if c.comment else None,
            (c.hyperlink.target, c.hyperlink.location, c.hyperlink.tooltip) if c.hyperlink else None)


def source_rows(report, matrix):
    contract = report["contract"]
    assert contract["heldout_count"] == 55 and len(contract["train_unit_ids"]) == 206
    assert contract["fold"] == 0 and contract["seed"] == 1234 and contract["phase_count"] == 80
    assert sha(contract["split_path"]) == contract["split_sha256"]
    matrix_arms = {x["arm"]: x for x in matrix["arms"]}
    submission = read(EXP / "submission.json")
    rows, audits = [], []
    u0 = report["models"]["U0"]["task_metrics"]["velocity"]
    for arm in ARMS:
        model = report["models"][arm]
        assert model["formal_complete"] and model["status"] == "completed" and not model["errors"], arm
        run = Path(model["out_dir"])
        metric_path, case_path = run / "metrics.json", run / "per_case.jsonl"
        raw = read(metric_path)
        cfg = read(run / "config.json")
        units = [json.loads(line) for line in case_path.read_text().splitlines() if line.strip()]
        assert len(units) == 55 and {x["unit_id"] for x in units} == set(contract["heldout_unit_ids"])
        assert raw["selection"] == "last" and raw["phase_count"] == 80 and raw["seed"] == 1234
        assert cfg["train"]["epochs"] == 150 and model["last_completed_epoch"] == 150
        assert cfg["tasks"] == ["velocity"]
        assert sha(cfg["data"]["stats_path"]) == model["stats_sha256"]
        history = [json.loads(line) for line in (run / "history.jsonl").read_text().splitlines() if line.strip()]
        complete = read(run / "training_complete.json")
        assert len(history) == 150 and history[-1]["epoch"] == 150 and history[-1]["step"] == 7800
        assert complete["epochs"] == 150 and complete["steps"] == 7800 and not complete.get("smoke")
        assert cfg["data"]["train_query_n"]["velocity"] == 1024 and cfg["data"]["eval_query_n"]["velocity"] == 16384
        record = {h: None for h in HEADERS}
        record.update({"实验臂": arm, "方法": DESCRIPTIONS[arm], "记录性质": "历史引用（本轮未重训）" if arm in ("Iu", "U0") else "本轮单seed训练",
                       "状态": "已完成150轮及55单元评估", "训练单元": 206, "开发留出单元": 55, "fold": 0, "seed": 1234,
                       "已完成epochs": 150, "优化器更新数": 7800, "相位数": 80, "checkpoint": "ckpt_last.pt（固定150轮）",
                       "训练query": 1024, "评估query": 16384, "参数量": model["parameter_count"], "训练秒": model["train_seconds"],
                       "缓存读入+80帧预测中位秒": model["sampled_query_median_s"], "运行目录": str(run), "配置路径": model["config_path"],
                       "原始metrics": str(metric_path), "逐单元metrics": str(case_path), "数据split": contract["split_path"],
                       "split SHA256": contract["split_sha256"], "stats SHA256": model["stats_sha256"],
                       "预注册主对照": "/".join(matrix_arms[arm]["primary_controls"]) if arm in matrix_arms else "历史参照",
                       "备注": "物理速度向量；逐单元体积加权指标的55单元等权均值。单seed开发筛选，非独立临床验证；峰值为窗口非单帧。"})
        execution = read(run / "execution.json")
        actual_job = str(execution["job_id"])
        if execution.get("array_task_id") is not None:
            actual_job += "_" + str(execution["array_task_id"])
        record["实际执行作业"] = actual_job
        if arm in matrix_arms:
            spec = matrix_arms[arm]
            record["原提交作业"] = f"{submission['batch1_job' if spec['batch'] == 1 else 'batch2_job']}_{spec['array_index']}"
        else:
            record["原提交作业"] = actual_job
        audit = {"arm": arm, "metrics_path": str(metric_path), "per_case_path": str(case_path),
                 "metrics_sha256": sha(metric_path), "per_case_sha256": sha(case_path), "n_units": len(units), "metric_checks": []}
        for stage, label in STAGES.items():
            for key, metric_label in METRIC_NAMES.items():
                vals = [x["tasks"]["velocity"][stage][key] for x in units]
                assert all(isinstance(v, (int, float)) and math.isfinite(v) for v in vals), (arm, stage, key)
                value = statistics.fmean(vals)
                raw_value = raw["tasks"]["velocity"][stage][key]
                report_value = model["task_metrics"]["velocity"][stage][key]
                assert close(value, raw_value) and close(value, report_value), (arm, stage, key, value, raw_value, report_value)
                header = f"{label} {metric_label}"
                record[header] = value
                audit["metric_checks"].append({"header": header, "field": f"tasks.velocity.{stage}.{key}", "value": value,
                                                "raw_metrics_value": raw_value, "report_value": report_value,
                                                "max_abs_difference": max(abs(value - raw_value), abs(value - report_value))})
            record[f"Δ{label}R² vs U0"] = record[f"{label} {METRIC_NAMES['r2']}"] - u0[stage]["r2"]
        rows.append(record)
        audits.append(audit)
    return rows, audits


def update(book, report_path, evidence_path):
    report = read(report_path)
    matrix = read(ROOT / "training_wss_min/configs/velocity_phase_v52_20260930/matrix.json")
    rows, audits = source_rows(report, matrix)
    before_sha = sha(book)
    wb = load_workbook(book)
    allowed_cells = {(INDEX, c) for c in ("C6", "D6", "E6")}
    originals = [s for s in wb.worksheets if s.title != SHEET]
    old_cells = {(s.title, c.coordinate): cell_record(c) for s in originals for row in s for c in row
                 if (s.title, c.coordinate) not in allowed_cells}
    old_dims = {s.title: ({k: dict(v) for k, v in s.column_dimensions.items()}, {k: dict(v) for k, v in s.row_dimensions.items()}) for s in originals}
    old_merges = {s.title: tuple(map(str, s.merged_cells.ranges)) for s in originals}
    original_rows = {s.title: s.max_row for s in originals}
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    source = wb["速度与压力实验矩阵"]
    ws = wb.create_sheet(SHEET, wb.sheetnames.index("速度与压力实验矩阵") + 1)
    ws["A1"] = "V5.2速度谷底与减速段优化｜8个单seed实验＋2条历史参照"
    ws["A2"] = "206训练/55开发留出数据单元；患者分组fold0；seed1234；150轮/7800更新/last；80相位；query1024→16384。"
    ws["A3"] = "主指标为逐数据单元vector R²的等权均值；体积加权；与旧R²_cb、speed模长指标分开。Iu/U0本轮未重训。"
    for r in (1, 2, 3):
        ws.cell(r, 1)._style = copy.copy(source.cell(r, 1)._style)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=28)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 34
    for col, label in ((1, "01 实验与主结果"), (9, "02 物理误差与方向"), (25, "03 相对U0变化"), (29, "04 协议、成本和可追溯来源")):
        ws.cell(4, col, label)._style = copy.copy(source.cell(4, 1)._style)
    for i, header in enumerate(HEADERS, 1):
        cell = ws.cell(5, i, header)
        cell._style = copy.copy(source.cell(5, 1)._style)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 60
    metric_cell_checks = []
    for r, (record, audit) in enumerate(zip(rows, audits), 6):
        for header, col in COL.items():
            cell = ws.cell(r, col, record[header])
            cell._style = copy.copy(source.cell(6 if r % 2 == 0 else 7, 1)._style)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if isinstance(record[header], float):
                cell.number_format = "0.0000"
            if header in ("运行目录", "配置路径", "原始metrics", "逐单元metrics", "数据split"):
                cell.hyperlink = record[header]
            if len(str(record[header] or "")) > 75:
                cell.comment = Comment(str(record[header]), "velocity workbook sync")
        if record["记录性质"].startswith("历史"):
            for cell in ws[r]:
                cell.fill = PatternFill("solid", fgColor="E9EEF3")
        ws.row_dimensions[r].height = 62
        audit["sheet_row"] = r
        for check in audit["metric_checks"]:
            coord = ws.cell(r, COL[check["header"]]).coordinate
            check["cell"] = coord
            metric_cell_checks.append((coord, check["value"]))
    last_col = get_column_letter(len(HEADERS))
    table_ref = f"A5:{last_col}15"
    ws.add_table(Table(displayName="VelocityPhaseV52Results", ref=table_ref,
                       tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)))
    ws.freeze_panes = "E6"
    for i in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    for h, width in (("实验臂", 11), ("方法", 44), ("记录性质", 23), ("状态", 29), ("运行目录", 55), ("备注", 65)):
        ws.column_dimensions[get_column_letter(COL[h])].width = width
    ws.column_dimensions.group("AC", last_col, hidden=False)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.print_title_rows = "1:5"
    ws.print_area = "A1:AB15"
    glossary_row = 18
    ws.cell(glossary_row, 1, "本页指标与边界").font = Font(bold=True, color="17365D")
    for idx, entry in enumerate(GLOSSARY, glossary_row + 1):
        ws.cell(idx, 1, entry[1])
        ws.cell(idx, 2, "；".join(entry[2:]))
        ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=14)
        ws.cell(idx, 2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[idx].height = 48
    index = wb[INDEX]
    index["C6"] = "WSS实验矩阵：直接WSS；速度与压力实验矩阵：历史体场；TAWSS_OSI周期量矩阵：周期积分量；V52速度周期优化：206/55单seed整周期速度vector指标。方法说明对照保留历史说明。"
    index["D6"] = "4个结果页＋2个说明页"
    index["E6"] = "各页采用各自冻结数据/评估协议；新增V52速度周期优化采用逐单元vector指标均值，禁止与旧R²_cb混读。"
    indexed = {index.cell(r, 2).value: r for r in range(1, index.max_row + 1)
               if index.cell(r, 1).value == "V5.2整周期专页"}
    index_rows = []
    for entry in GLOSSARY:
        r = indexed.get(entry[1], index.max_row + 1)
        for c, value in enumerate(entry, 1):
            cell = index.cell(r, c, value)
            cell._style = copy.copy(index.cell(14, c)._style)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        if r > original_rows[INDEX]:
            index.row_dimensions[r].height = 58
        index_rows.append(r)
    # The appended owned glossary is allowed to refresh idempotently.
    for r in index_rows:
        for c in range(1, 6):
            old_cells.pop((INDEX, f"{get_column_letter(c)}{r}"), None)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    evidence_path.parent.mkdir(parents=True, exist_ok=True)
    backup = evidence_path.parent / f"workbook_before_velocity_phase_{stamp}.xlsx"
    backup.write_bytes(book.read_bytes())
    fd, name = tempfile.mkstemp(prefix=book.stem + ".", suffix=".xlsx", dir=book.parent)
    os.close(fd)
    candidate = Path(name)
    try:
        wb.save(candidate)
        checked = load_workbook(candidate)
        for (sheet, coord), value in old_cells.items():
            assert cell_record(checked[sheet][coord]) == value, f"historical cell changed: {sheet}!{coord}"
        for sheet, merges in old_merges.items():
            assert tuple(map(str, checked[sheet].merged_cells.ranges)) == merges
            cols, dims = old_dims[sheet]
            assert {k: dict(v) for k, v in checked[sheet].column_dimensions.items()} == cols
            for k, old in dims.items():
                assert dict(checked[sheet].row_dimensions[k]) == old, (sheet, k, "row dimension changed")
        for coord, value in metric_cell_checks:
            assert close(checked[SHEET][coord].value, value), (coord, value, checked[SHEET][coord].value)
        for row, record in enumerate(rows, 6):
            for header, value in record.items():
                got = checked[SHEET].cell(row, COL[header]).value
                assert close(got, value) if isinstance(value, float) else got == value, (row, header, got, value)
        assert [checked[SHEET].cell(r, 1).value for r in range(6, 16)] == list(ARMS)
        assert sha(book) == before_sha, "workbook changed concurrently; refusing overwrite"
        os.replace(candidate, book)
    finally:
        candidate.unlink(missing_ok=True)
    evidence = {"schema": "velocity_phase_workbook_update_v1", "passed": True,
                "updated_at": datetime.now(timezone.utc).isoformat(), "workbook": str(book), "sheet": SHEET,
                "table_ref": table_ref, "rows": {arm: i + 6 for i, arm in enumerate(ARMS)},
                "record_count": 10, "new_training_runs": 8, "historical_reference_runs": 2,
                "index_sheet": INDEX, "index_changed_cells": ["C6", "D6", "E6"], "index_glossary_rows": index_rows,
                "historical_cells_verified": len(old_cells), "historical_formulas_verified": sum(v[3] == "f" for v in old_cells.values()),
                "historical_dimensions_and_merges_verified": True, "reopen_numeric_cells_verified": len(metric_cell_checks),
                "agreement_tolerance": {"absolute": 1e-12, "relative": 1e-10},
                "maximum_raw_report_per_case_abs_difference": max(c["max_abs_difference"] for a in audits for c in a["metric_checks"]),
                "report": str(report_path), "report_sha256": sha(report_path), "report_generated_at": report["generated_at"],
                "backup": str(backup), "before_sha256": before_sha, "after_sha256": sha(book), "source_audits": audits}
    evidence_path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    return {k: v for k, v in evidence.items() if k != "source_audits"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=BOOK)
    parser.add_argument("--report", type=Path, default=EXP / "report.json")
    parser.add_argument("--evidence", type=Path, default=EXP / "analysis_20260930/workbook_update.json")
    args = parser.parse_args()
    print(json.dumps(update(args.workbook, args.report, args.evidence), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
