"""Idempotent workbook backfill for volume_time_20260919 (pressure / velocity full-cycle time arms).

Owns exactly the rows whose column A equals GROUP on sheet ``速度与压力实验矩阵``
(135 columns, table ``VolumeResults``): 6 peak-frame fold bases (best only) and 18 time arms
(best + last) = 42 rows.  Column semantics follow the existing pressure/velocity rows.

本批与 §17 / 波 2 体场行的关键口径差异，全部写进备注列，不改列语义：

* 数据是 v5.1 母库 + `cv3_v51` 三折，每臂只读自己的留出折（46 / 46 / 46 例），`test34` 本轮未用；
  波 2 那批是旧 v5 的 train138/test34，两者的 R²_cb **不可直接比**。
* 6–125 列仍是**峰值帧**（步 1162）读数，和峰值帧臂同尺；周期指标（全周期 / 谷底 / 最差帧 /
  时均场 / 峰时误差 / 对 T-null 的差）是本批独有，写在备注与 ΔR² 参照列里。
* 时间基臂（PTB*/VTB*）的峰值帧预测由系数重建而来，不是直接回归峰值帧。

    python -m training_wss_min.tools.update_volume_time_xlsx [--workbook PATH]
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import tempfile
import time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

from training_wss_min import config as C
from training_wss_min.tools.report_wss_recovery import cell_record, get_metric
from training_wss_min.tools.update_wss_local_wave2_volume_xlsx import COLUMNS, sha

ROOT = C.PROJECT_ROOT
NAME = "volume_time_20260919"
CONFIGS = ROOT / "training_wss_min/configs" / NAME
EXP = ROOT / "training_wss_min/experiments" / NAME
RUNS = ROOT / "training_wss_min/runs" / NAME
BOOK = ROOT / "docs/03-汇报材料/WSS_PointNet实验矩阵与结果汇总last.xlsx"
SHEET = "速度与压力实验矩阵"
GROUP = "体场全周期时间阶段2｜2026-09-19（PT0/PTB8/PTB16、VT0/VTB4/VTB8×cv3，含峰值帧折底座）"
MISSING = "—"
CYCLE_KEYS = ("cycle_r2cb", "trough_r2cb", "min_frame_r2cb", "tmean_r2cb", "ts_corr",
              "peak_time_err_med_frames", "amp_err_med")


def read_json(path, default=None):
    path = Path(path)
    return json.loads(path.read_text()) if path.is_file() else default


def metrics_of(run_dir: Path, checkpoint: str):
    path = run_dir / f"eval/ckpt_{checkpoint}/metrics.json"
    payload = read_json(path)
    return (payload["test"] if payload else None), str(path)


def kind_of(aid: str) -> str:
    return "P" if aid[0] == "P" else "V"


def describe(aid: str, arm: dict, gate: dict | None) -> tuple[str, str]:
    """(描述, 结构) — 折底座 / 相位查询 / 时间基头三类。"""
    fold, seed = arm["fold"], arm.get("seed", 1234)
    kind = "压力" if kind_of(aid) == "P" else "速度"
    if "F6_v51" in aid:
        return (f"{kind}峰值帧折底座：老 v5 的 {'PF6' if kind_of(aid) == 'P' else 'VF6'} 配方"
                f"（18D 体场特征 + Murray log_q/log_tau0）搬到 v5.1 数据与 cv3_v51 fold{fold}，seed {seed}；"
                f"本批时间臂的配对父臂与峰值帧读数基准",
                f"{'PF6' if kind_of(aid) == 'P' else 'VF6'} 原配方（峰值帧，fold{fold}）")
    if "T0" in aid:
        return (f"{kind}相位查询：折底座 + 4 相位列（q_norm/dq_norm/t_sin/t_cos，共 24D，新增列零初始化），"
                f"random_frame + 峰值保底 1/9，逐帧线性 z，EMA(train_loss, α0.2) 选 best；"
                f"fold{fold} 留出折 81 帧，seed {seed}",
                f"折底座 + 相位条件（输入 +4 维）")
    k = int(aid.split("TB")[1].split("_")[0])
    channels = 1 if kind_of(aid) == "P" else 3
    return (f"{kind}时间基头 K={k}（out_dim {channels * (k + 1)} = {channels} 分量 × [b0, a_1..a_K]，"
            f"{'三分量共用一套基向量，' if channels == 3 else ''}训练折 PCA 基，系数 MSE，pinball 关），"
            f"无时间输入，一次推理重建 81 帧；fold{fold} 留出折，seed {seed}",
            f"折底座 + 时间基输出头 K={k}（out_dim {channels * (k + 1)}）")


def arm_rows():
    matrix = read_json(CONFIGS / "matrix.json")
    queue = read_json(EXP / "queue_status.json", {"arms": {}})
    gates = read_json(EXP / "gate_report.json", {"families": {}})
    rows = []
    for arm in matrix["arms"]:
        aid = arm["id"]
        config = read_json(CONFIGS / arm["config"])
        record = queue["arms"].get(aid, {})
        run = RUNS / aid
        base_id = arm["parent"] if arm["phase"] == 2 else None
        family = "pressure" if kind_of(aid) == "P" else "velocity"
        block = gates.get("families", {}).get(family, {})
        stem = aid.split("_f")[0]
        for checkpoint in arm.get("evaluate", ["best"]):
            stage = record.get("stages", {}).get(f"eval_{checkpoint}", {})
            metrics, path = metrics_of(run, checkpoint)
            accepted = record.get("status") == "complete" and stage.get("returncode") == 0 and metrics is not None
            refs = []
            if base_id:
                ref_metrics, ref_path = metrics_of(RUNS / base_id, "best")
                if ref_metrics is not None:
                    refs.append((f"{base_id}（同折峰值帧底座 best）", ref_metrics, ref_path))
            rows.append(dict(id=aid, arm=arm, kind=kind_of(aid), checkpoint=checkpoint, run_name=config["name"],
                             run_dir=str(run), config_path=str(CONFIGS / arm["config"]), metrics=metrics,
                             metrics_path=path, refs=refs, accepted=accepted, family=family,
                             diagnostics=read_json(run / "training_diagnostics.json", {}),
                             epochs=config["train"]["epochs"], fold=arm["fold"],
                             input_dim=len(config["data"]["input_features"]),
                             gate=(block.get("arms", {}).get(stem, {}).get(checkpoint) if base_id else None),
                             t_null=block.get("t_null_cycle"), base_peak=block.get("fold_base", {}).get("peak_r2cb"),
                             status=record.get("status", "pending"),
                             split=config["data"]["split_path"]))
    return matrix, queue, gates, rows


def row_values(item: dict, job_id, matrix: dict) -> tuple[list, str, list[str]]:
    kind, aid, gate = item["kind"], item["id"], item["gate"]
    values = [MISSING] * 135
    values[0] = GROUP
    description, structure = describe(aid, item["arm"], gate)
    values[1] = f"{aid}｜{description}｜{item['status']}"
    values[2] = "压力 · Pa" if kind == "P" else "速度 · m/s"
    values[3] = item["checkpoint"]
    values[4] = (f"v5.1 cv3_v51 fold{item['fold']} 留出折 · "
                 + ("壁面∪内部" if kind == "P" else "内部"))
    sources = [item["config_path"], str(Path(item["run_dir"]) / "history.jsonl"),
               str(EXP / "results.json"), str(EXP / "gate_report.json")]
    seed = item["arm"].get("seed", 1234)
    note = [f"Job {job_id}；{item['status']}；单 seed {seed}、只作筛选。",
            f"数据 = v5.1 母库 + cv3_v51 fold{item['fold']}（训练两折、读数留出折），test34 本轮未用；"
            f"与波 2 体场行（旧 v5 train138/test34）不可直接比。",
            f"6–125 列是峰值帧（步 1162）读数"
            + ("；时间基臂的峰值帧预测由系数重建而来，不是直接回归峰值帧。" if "TB" in aid and "F6" not in aid
               else "。")]
    if item["accepted"]:
        m = item["metrics"]
        for col, dotted in COLUMNS.items():
            value = get_metric(m, dotted)
            if value is not None:
                values[col - 1] = value
        values[11] = f"{int(m['aggregate']['r2_negative_cases'])} / {int(m['aggregate']['n_cases'])}"
        values[89] = item["diagnostics"].get("elapsed_seconds", MISSING)
        values[90] = item["diagnostics"].get("peak_cuda_allocated_mb", MISSING)
        peak = get_metric(m, "efficiency.peak_cuda_memory_bytes")
        values[91] = peak / 2 ** 20 if peak is not None else MISSING
        own_r2 = get_metric(m, "field_casebalanced.r2")
        for j, (label, ref_metrics, ref_path) in enumerate(item["refs"][:2]):
            base = 94 + j * 6
            ref_r2 = get_metric(ref_metrics, "field_casebalanced.r2")
            values[base] = f"{label}｜{'压力' if kind == 'P' else '速度'}｜峰值帧"
            values[base + 1] = own_r2 - ref_r2 if (own_r2 is not None and ref_r2 is not None) else MISSING
            values[base + 5] = "单seed筛选，无门槛判定"
            sources.append(ref_path)
            if j == 0:
                values[12] = values[base + 1]
                values[13] = label
        values[92] = ("单seed筛选（时间臂对同折峰值帧底座；周期指标对 A3 T-null）" if item["refs"]
                      else "单seed筛选（折底座本身，无同期对照）")
        values[93] = ("峰值帧折底座" if "F6_v51" in aid else "相位查询（T0）" if "T0" in aid else "时间基头（TB）")
        cyc = m.get("cycle")
        if cyc:
            tn = item["t_null"] or {}
            note.append(
                "周期指标（该折 81 帧、固定子采样：压力 20000 内部单元 + 4000 壁面节点，速度 20000 内部单元）："
                f"全周期 R²_cb {cyc['cycle_r2cb']:.4f}、谷底 {cyc['trough_r2cb']:.4f}、"
                f"最差帧 {cyc['min_frame_r2cb']:.4f}@{cyc['argmin_frame']}、时均场 {cyc['tmean_r2cb']:.4f}"
                f"（MAE {cyc['tmean_mae']:.3f} {'Pa' if kind == 'P' else 'm/s'}）、"
                f"时序相关 {cyc['ts_corr']:.3f}、峰时误差中位 {cyc['peak_time_err_med_frames']:.2f} 帧。")
            if tn:
                note.append(f"对照 A3 T-null（真值峰值场 × 训练折时间形状，三折均值）："
                            f"全周期 {tn['cycle_r2cb']:.4f}、谷底 {tn['trough_r2cb']:.4f}、"
                            f"时均 {tn['tmean_r2cb']:.4f}。")
        if gate:
            note.append(
                f"三折均值门控（脚本 offline/report.py）：峰值 Δ vs 折底座 {gate['d_peak_vs_base']:+.4f}"
                f"（{gate['peak_sign_agreement']}/3 折为正）、全周期 Δ vs T-null {gate['d_cycle_vs_tnull']:+.4f}"
                f"（{gate['cycle_sign_agreement']}/3 折胜出）→ "
                + ("过门" if gate["gate"] else "不过门"))
    else:
        note.append("本行尚无可接受的评估产物。")
    values[126] = "Pa" if kind == "P" else "m/s"
    values[127] = ("PointNeXt-R + LocalGeoPE + L-SA2 + QAD-lite；"
                   + ("PF6：原 SA3 后全局 BT ＋ Murray 先验" if kind == "P"
                      else "VF6：新 SA3 单半径 0.20 加 BT ＋ Murray 先验"))
    values[128] = structure
    values[129] = (f"{item['epochs']} epoch；support 随机5000；独立 query；{item['input_dim']}D 输入；"
                   + ("相对压力 p−逐帧体积均值" if kind == "P" else "速度三分量")
                   + ("、峰值帧全局线性 z、MSE" if "F6_v51" in aid else "、逐帧线性 z")
                   + ("、系数 MSE" if "TB" in aid and "F6" not in aid else "、MSE" if "F6_v51" not in aid else ""))
    values[130] = 5000
    values[131] = item["run_name"]
    sources.append(item["metrics_path"])
    values[132] = " ".join(note)[:32000]
    values[133] = "\n".join(sources)
    values[134] = "[]"
    return values, " ".join(note), sources


def update_workbook(book: Path):
    matrix, queue, gates, rows = arm_rows()
    before = sha(book)
    wb = load_workbook(book)
    if wb.sheetnames[:3] != ["WSS实验矩阵", "速度与压力实验矩阵", "指标说明"]:
        raise ValueError("unexpected workbook sheet schema")
    ws = wb[SHEET]
    if ws.max_column != 135 or ws["EB5"].value != "原始实验 ID" or "VolumeResults" not in ws.tables:
        raise ValueError("unexpected volume sheet 135-column schema")
    owned = [r for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value == GROUP]
    if owned and len(owned) != len(rows):
        raise ValueError(f"owned rows must number exactly {len(rows)}")
    keys = [(item["run_name"], item["checkpoint"]) for item in rows]
    locations = ({(ws.cell(r, 132).value, ws.cell(r, 4).value): r for r in owned} if owned else
                 {key: ws.max_row + i + 1 for i, key in enumerate(keys)})
    if set(locations) != set(keys):
        raise ValueError("run identities differ from the workbook's owned rows")
    preserved = {(s.title, c.coordinate): cell_record(c) for s in wb for row in s for c in row
                 if not (s.title == ws.title and c.row in owned)}
    dimensions = {s.title: (copy.deepcopy(s.column_dimensions), copy.deepcopy(s.row_dimensions)) for s in wb}
    merges = {s.title: tuple(map(str, s.merged_cells.ranges)) for s in wb}
    job_id = queue.get("job_id", "尚未提交")
    for item in rows:
        row = locations[(item["run_name"], item["checkpoint"])]
        values, note, sources = row_values(item, job_id, matrix)
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col)
            cell._style = copy.copy(ws.cell(6 if row % 2 == 0 else 7, col)._style)
            cell.value = value
            cell.comment = None
            cell.hyperlink = None
        ws.cell(row, 133).comment = Comment(note[:32000], "volume time report")
        ws.cell(row, 134).comment = Comment("\n".join(sources), "volume time report")
        ws.row_dimensions[row].height = 72
    ws.tables["VolumeResults"].ref = f"A5:EE{ws.max_row}"
    ws.tables["VolumeResults"].autoFilter.ref = f"A5:EE{ws.max_row}"
    ws.print_area = f"A1:N{ws.max_row}"
    fd, temp = tempfile.mkstemp(dir=book.parent, prefix=f".{NAME}_", suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(temp)
        checked = load_workbook(temp)
        for (sheet, coordinate), record in preserved.items():
            if cell_record(checked[sheet][coordinate]) != record:
                raise AssertionError(f"historical cell changed: {sheet}!{coordinate}")
        assert {s.title: tuple(map(str, s.merged_cells.ranges)) for s in checked} == merges
        for sheet in checked:
            old_cols, old_rows = dimensions[sheet.title]
            assert {k: dict(v) for k, v in sheet.column_dimensions.items()} == {k: dict(v) for k, v in old_cols.items()}
            for key, value in old_rows.items():
                if sheet.title != ws.title or key not in owned:
                    assert dict(sheet.row_dimensions[key]) == dict(value)
        for item in rows:
            r = locations[(item["run_name"], item["checkpoint"])]
            assert checked[SHEET].cell(r, 1).value == GROUP and checked[SHEET].cell(r, 132).value == item["run_name"]
            if not item["accepted"]:
                for column in [*range(6, 14), *range(15, 127)]:
                    assert checked[SHEET].cell(r, column).value == MISSING
        checked.close()
        if sha(book) != before:
            raise RuntimeError("workbook changed concurrently; refusing replacement")
        backup = (book.parent / f"{book.stem}_backup_before_{NAME}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
        backup.write_bytes(book.read_bytes())
        os.replace(temp, book)
    finally:
        if Path(temp).exists():
            Path(temp).unlink()
    evidence = {"passed": True, "before_sha256": before, "after_sha256": sha(book), "backup": str(backup),
                "rows": {f"{k[0]}@{k[1]}": v for k, v in locations.items()}, "records": len(rows),
                "accepted_result_rows": sum(r["accepted"] for r in rows),
                "historical_cells_verified": len(preserved), "table_ref": ws.tables["VolumeResults"].ref,
                "row_span": [min(locations.values()), max(locations.values())]}
    (EXP / "xlsx_acceptance.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
    return evidence


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=BOOK)
    args = parser.parse_args(argv)
    print(json.dumps(update_workbook(args.workbook), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
