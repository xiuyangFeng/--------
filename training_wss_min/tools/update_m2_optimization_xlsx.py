"""Update only the marked M2 sections; preserve historical cells and formulas.

Twenty base rows and four reserved combination rows use fixed positions. Existing
rows are never deleted or inserted, even when other sections follow this one.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import tempfile

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from training_wss_min.tools import report_m2_optimization as R
from training_wss_min.tools.update_v5_rerun_xlsx import BOOK, EXPECTED_SHEETS, LSA2
from training_wss_min.tools.update_v6_followup_xlsx import followup_overview_values, teacher_metric_values

PREFIX = "M2优化20260909"
TITLE = PREFIX + "｜20基础臂及最多4组合；单seed1234；best训练损失选模，last交叉检查"
END = PREFIX + "｜本节结束"
CAPACITY = 25  # historical M2, 20 base arms, at most 4 combinations
RECT_HEIGHT = CAPACITY + 4  # title, headers, 25 data rows, note, end
WIDTHS = {"实验矩阵总览": 49, "教师汇报视图": 19, "汇总对比": 23}
NOTE = ("仅单seed1234、已暴露test34探索。总览/教师填写best；汇总同时列best与last。"
        "损失改变也改变best选择准则；不从best/last择优。Pa R²_cb病例等权；MAE顶点pool。"
        "教师O为log_z p99比；汇总高值比均为Pa。"
        "同期MO0成绩偏离历史M2，需结合配套contemporary_control_diagnosis.md审计解释；不能作为稳定提升证据。"
        "新几何仍待人工验收。")


def tag(kind):
    return f"{R.NAME}|{kind}"


def marker(cell, kind):
    cell.comment = Comment(tag(kind), "M2 optimization report")


def snapshot_cells(workbook, excluded=None):
    excluded = excluded or {}
    output = {}
    for ws in workbook.worksheets:
        rect = excluded.get(ws.title)
        for row in ws:
            for cell in row:
                if rect and rect[0] <= cell.row <= rect[1] and cell.column <= rect[2]:
                    continue
                if cell.value is not None:
                    output[(ws.title, cell.coordinate)] = cell.value
    return output


def locate_section(ws):
    starts = [c.row for c in ws["A"] if c.value == TITLE]
    ends = [c.row for c in ws["A"] if c.value == END]
    if not starts and not ends:
        return None
    if len(starts) != 1 or len(ends) != 1 or ends[0] - starts[0] != RECT_HEIGHT - 1:
        raise ValueError(f"Ambiguous or resized M2 section in {ws.title}")
    first = starts[0]
    expected = {0: "title", 1: "header", CAPACITY + 2: "note", CAPACITY + 3: "end"}
    expected.update({2 + i: f"slot:{i}" for i in range(CAPACITY)})
    for offset, kind in expected.items():
        c = ws.cell(first + offset, 1)
        if c.comment is None or c.comment.text != tag(kind):
            raise ValueError(f"Unowned or changed M2 row {ws.title}!{c.row}")
    return first


def write_row(ws, row, values, width, header=False):
    if len(values) > width:
        raise ValueError("Row exceeds owned rectangle")
    for column in range(1, width + 1):
        cell = ws.cell(row, column)
        cell.value = values[column - 1] if column <= len(values) else None
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if header:
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.font = Font(bold=True)
        elif isinstance(cell.value, float):
            cell.number_format = "0.0000"
    ws.row_dimensions[row].height = 48 if header else 60


def make_headers(sheet):
    if sheet == "汇总对比":
        return ["ID", "族/组成", "状态", "best Pa R²_cb", "Δbest vs历史M2", "Δbest vs MO0",
                "last Pa R²_cb", "Δlast vs历史M2", "Δlast vs MO0", "best log_z R²_cb", "best病例均值R²",
                "best P10", "best MAE Pa", "best high-WSS R²", "best top10比", "best p99比", "best IoU",
                "best负R²数", "门槛", "变化", "结果目录", "best参数量", "训练秒数"]
    if sheet == "教师汇报视图":
        return ["阶段", "实验/状态", "模型", "训练/选模", "Pa R²_cb", "病例均值R²", "pooled RMSE Pa",
                "pooled NMAE(range)", "case mean NMAE", "high-WSS R²", "log_z R²_cb", "log_z NMAE(range)",
                "Spearman", "top10 IoU", "log_z p99比", "Pa fit R²_cb", "Pa fit R² pooled", "Pa fit slope", "Pa fit intercept"]
    return ["实验", "数据", "模型", "变化", "训练/选模", "support点数", "Pa R²_cb", "Pa pooled R²",
            "病例均值R²", "病例中位R²", "病例P10", "负R²数", "pooled RMSE Pa", "pooled MAE Pa",
            "pooled range-NRMSE", "case mean NRMSE", "pooled NMAE", "case mean NMAE", "high-WSS R²",
            "top10幅值比", "top10 IoU", "log_z R²_cb", "log_z pooled R²", "log_z病例均值R²", "log_z病例中位R²",
            "log_z病例P10", "log_z pooled MAE", "log_z pooled RMSE", "log_z pooled NRMSE", "log_z case NRMSE",
            "log_z pooled NMAE", "log_z case NMAE", "Spearman", "high-WSS Spearman", "log_z top10幅值比", "log_z p99比",
            "结果目录", "Pa fit R²_cb", "Pa fit slope cb", "Pa fit intercept cb", "Pa fit R² pooled",
            "Pa fit slope pooled", "Pa fit intercept pooled", "log_z fit R²_cb", "log_z fit slope cb",
            "log_z fit intercept cb", "log_z fit R² pooled", "log_z fit slope pooled", "log_z fit intercept pooled"]


def make_rows(arms, reports, runs_dir=R.RUNS):
    rows = {sheet: [] for sheet in WIDTHS}
    synthetic = {"id": "historical_m2", "family": "reference", "hypothesis": "历史M2；原25D、原损失、独立query/3NN",
                 "run_name": R.ANCHOR, "config": "historical_m2.json", "resolved_config": {"model": {}}}
    for arm in [synthetic, *arms]:
        aid = arm["id"]
        historic = aid == "historical_m2"
        best = reports["best"]["reference"] if historic else reports["best"]["arms"][aid]
        last = reports["last"]["reference"] if historic else reports["last"]["arms"][aid]
        valid = historic or best["evidence"]["valid_scientific_result"]
        last_valid = historic or last["evidence"]["valid_scientific_result"]
        metric = R.load_metrics(arm["run_name"], "best", runs_dir) if valid else None
        status = "历史参照" if historic else best["evidence"]["status"]
        model = LSA2 + " + 多尺度局部壁面分支；" + arm["hypothesis"]
        protocol = "random5000 support/独立query5000；s1234；400轮；本行best"
        if metric is not None:
            # Reuse only the established 49 numeric columns; the legacy helper
            # expects a different change-manifest schema. Our text is set below.
            adjusted = {**arm, "changes": {}}
            overview = followup_overview_values(adjusted, metric)
            overview[0], overview[2], overview[3], overview[4], overview[36] = f"{PREFIX}·{aid}", model, arm["hypothesis"], protocol, arm["run_name"]
        else:
            overview = [None] * 49
            overview[:6] = [f"{PREFIX}·{aid}", "V5 train138/test34；峰值单帧", model, arm["hypothesis"] + "｜" + status, protocol, 5000]
            overview[36] = arm["run_name"]
        rows["实验矩阵总览"].append(overview)
        rows["教师汇报视图"].append(["M2优化", f"{aid} {arm['hypothesis']}｜{status}", model, protocol] + teacher_metric_values(best if valid else {}, metric))
        get = lambda key: best.get(key) if valid else None
        delta = lambda entry, ref: entry.get(f"deltas_vs_{ref}", {}).get("physical_r2_cb")
        gate = reports["selection"].get("candidates", {}).get(aid)
        gate_text = "参照" if historic or aid == "MO0" else "待整批完成" if gate is None else "通过探索门槛" if gate["qualified"] else "未通过：" + "；".join(gate["issues"])
        members = arm.get("members", arm.get("components", []))
        rows["汇总对比"].append([aid, "+".join(members) if members else arm["family"], status, get("physical_r2_cb"),
            delta(best, "historical_m2"), delta(best, "mo0"), last.get("physical_r2_cb") if last_valid else None,
            delta(last, "historical_m2"), delta(last, "mo0"), get("normalized_r2_cb"), get("case_mean"), get("case_p10"),
            get("mae"), get("high_wss_r2"), get("top10_ratio"), get("p99_ratio"), get("top10_iou"), get("negative_cases"),
            gate_text, arm["hypothesis"], arm["run_name"], best.get("efficiency", {}).get("parameters"), best.get("evidence", {}).get("training_seconds")])
    for sheet in rows:
        while len(rows[sheet]) < CAPACITY:
            rows[sheet].append([f"{PREFIX}·组合预留{len(rows[sheet]) - 20}", "未登记组合；不是已完成实验"])
        if len(rows[sheet]) != CAPACITY:
            raise ValueError("M2 workbook capacity exceeded")
    return rows


def update_in_memory(workbook, rows):
    if workbook.sheetnames != EXPECTED_SHEETS:
        raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    old_starts = {name: locate_section(workbook[name]) for name in WIDTHS}
    excluded = {name: (start, start + RECT_HEIGHT - 1, WIDTHS[name]) for name, start in old_starts.items() if start is not None}
    preserved = snapshot_cells(workbook, excluded)
    merges = {ws.title: [str(r) for r in ws.merged_cells.ranges] for ws in workbook.worksheets}
    result = {}
    for name, width in WIDTHS.items():
        ws = workbook[name]
        first = old_starts[name] or ws.max_row + 1
        result[name] = {"first_row": first, "last_row": first + RECT_HEIGHT - 1, "data_rows": list(range(first + 2, first + CAPACITY + 2))}
        title = ws.cell(first, 1)
        title.value = TITLE
        title.fill = PatternFill("solid", fgColor="1F4E78")
        title.font = Font(bold=True, color="FFFFFF")
        title.alignment = Alignment(wrap_text=True, vertical="center")
        marker(title, "title")
        ws.row_dimensions[first].height = 50
        if old_starts[name] is None:
            ws.merge_cells(start_row=first, start_column=1, end_row=first, end_column=width)
        write_row(ws, first + 1, make_headers(name), width, header=True)
        marker(ws.cell(first + 1, 1), "header")
        for index, row in enumerate(rows[name]):
            write_row(ws, first + 2 + index, row, width)
            marker(ws.cell(first + 2 + index, 1), f"slot:{index}")
        write_row(ws, first + CAPACITY + 2, [PREFIX + "｜说明", NOTE], width)
        marker(ws.cell(first + CAPACITY + 2, 1), "note")
        write_row(ws, first + CAPACITY + 3, [END], width)
        marker(ws.cell(first + CAPACITY + 3, 1), "end")
    return result, preserved, merges


def verify_preserved(workbook, preserved, merges):
    if workbook.sheetnames != EXPECTED_SHEETS:
        raise AssertionError("Workbook sheet names/order changed")
    for (name, coordinate), value in preserved.items():
        if workbook[name][coordinate].value != value:
            raise AssertionError(f"Historical cell changed: {name}!{coordinate}")
    for name, ranges in merges.items():
        current = {str(r) for r in workbook[name].merged_cells.ranges}
        if not set(ranges) <= current:
            raise AssertionError(f"Historical merged cells changed in {name}")


def update_workbook(book=BOOK, matrix_path=R.MATRIX, combinations_path=R.COMBINATIONS, experiment_dir=R.EXP,
                    runs_dir=R.RUNS, require_complete=True):
    book, experiment_dir = Path(book), Path(experiment_dir)
    reports = R.report_all(matrix_path, combinations_path, experiment_dir, runs_dir)
    if require_complete and not all(reports[ck]["matrix_finalized"] for ck in ("best", "last")):
        raise RuntimeError("Refusing final workbook update until all registered runs and evaluations are verified")
    arms = R.matrix_arms(matrix_path, combinations_path)
    original = book.read_bytes()
    original_hash = hashlib.sha256(original).hexdigest()
    backup = experiment_dir / "WSS_PointNet实验矩阵与结果汇总last_backup_before_m2_optimization_20260909.xlsx"
    if not backup.exists():
        backup.write_bytes(original)
    rows = make_rows(arms, reports, runs_dir)
    descriptor, candidate_name = tempfile.mkstemp(prefix=".m2_optimization_", suffix=".xlsx", dir=book.parent)
    os.close(descriptor)
    candidate = Path(candidate_name)
    try:
        candidate.write_bytes(original)
        workbook = load_workbook(candidate, data_only=False)
        positions, preserved, merges = update_in_memory(workbook, rows)
        workbook.save(candidate)
        workbook.close()
        check = load_workbook(candidate, data_only=False)
        verify_preserved(check, preserved, merges)
        checked = 0
        for sheet, position in positions.items():
            if locate_section(check[sheet]) != position["first_row"]:
                raise AssertionError("M2 section ownership did not survive workbook serialization")
            for row_index, expected in zip(position["data_rows"], rows[sheet]):
                for column, value in enumerate(expected, 1):
                    actual = check[sheet].cell(row_index, column).value
                    same = abs(actual - value) < 1e-12 if isinstance(value, float) and isinstance(actual, (float, int)) else actual == value
                    if not same:
                        raise AssertionError(f"Result transcription mismatch {sheet}!{row_index},{column}")
                    checked += 1
        check.close()
        if hashlib.sha256(book.read_bytes()).hexdigest() != original_hash:
            raise RuntimeError("Workbook changed concurrently; refusing to overwrite")
        candidate.replace(book)
    finally:
        candidate.unlink(missing_ok=True)
    evidence = {"book": str(book), "backup": str(backup), "before_sha256": original_hash, "after_sha256": R.sha256(book),
                "positions": positions, "preserved_nonempty_cells": len(preserved),
                "preserved_formulas": sum(isinstance(v, str) and v.startswith("=") for v in preserved.values()),
                "transcription_checks": checked, "all_historical_cells_and_formulas_preserved": True,
                "completed_best": reports["best"]["completed_count"], "completed_last": reports["last"]["completed_count"],
                "matrix_finalized": all(reports[ck]["matrix_finalized"] for ck in ("best", "last")), "row_deletion": False}
    (experiment_dir / "workbook_update.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n")
    return evidence


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--book", type=Path, default=BOOK)
    parser.add_argument("--matrix", type=Path, default=R.MATRIX)
    parser.add_argument("--combinations", type=Path, default=R.COMBINATIONS)
    parser.add_argument("--experiment-dir", type=Path, default=R.EXP)
    parser.add_argument("--runs-dir", type=Path, default=R.RUNS)
    parser.add_argument("--allow-pending", action="store_true", help="Explicitly register pending rows with blank metrics")
    args = parser.parse_args(argv)
    result = update_workbook(args.book, args.matrix, args.combinations, args.experiment_dir, args.runs_dir, not args.allow_pending)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
