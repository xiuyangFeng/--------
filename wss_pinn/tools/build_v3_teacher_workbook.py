"""Build a teacher-facing XLSX for the completed V3 six-arm experiment.

The workbook is generated from frozen summary CSV/JSON artifacts and embeds
the already-audited metric and loss-convergence figures.  It does not recompute
test metrics or select checkpoints.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage


ROOT = Path(__file__).resolve().parents[2]
SUMMARY = ROOT / "outputs/wss_pinn/volume_uvwp_peak_qs_smooth_v3/summary"
LOSS_DIR = SUMMARY / "loss_convergence"
DEFAULT_OUTPUT = SUMMARY / "WSS_PINN_V3_六臂核心实验与指标汇报.xlsx"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
GREEN = "548235"
LIGHT_GREEN = "E2F0D9"
ORANGE = "C65911"
LIGHT_ORANGE = "FCE4D6"
RED = "C00000"
LIGHT_RED = "F4CCCC"
PURPLE = "7030A0"
LIGHT_PURPLE = "E4DFEC"
GRAY = "666666"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
YELLOW = "FFF2CC"

THIN_GRAY = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _number(value: Any) -> float | int | str | None:
    if value in (None, "", "n/a", "disabled"):
        return value
    try:
        result = float(value)
    except (TypeError, ValueError):
        return value
    if not math.isfinite(result):
        return None
    if result.is_integer() and abs(result) < 1e12:
        return int(result)
    return result


def _style_title(ws, text: str, end_column: int, *, row: int = 1) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=end_column)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name="Microsoft YaHei", size=18, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 26
    ws.row_dimensions[row + 1].height = 12
    for r in range(row, row + 2):
        for c in range(1, end_column + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=NAVY)


def _section(ws, row: int, text: str, end_column: int, color: str = BLUE) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name="Microsoft YaHei", size=12, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 22
    return row + 1


def _header_row(ws, row: int, headers: list[str], *, fill: str = NAVY) -> None:
    for column, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=column, value=header)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 32


def _body_cell(cell, *, center: bool = True, wrap: bool = False) -> None:
    cell.font = Font(name="Microsoft YaHei", size=10, color="222222")
    cell.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center",
        wrap_text=wrap,
    )
    cell.border = BORDER


def _set_widths(ws, widths: dict[int, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width


def _add_image(ws, path: Path, anchor: str, *, width: int) -> None:
    with PILImage.open(path) as image:
        ratio = width / image.width
        height = int(image.height * ratio)
    xl_image = XLImage(str(path))
    xl_image.width = width
    xl_image.height = height
    ws.add_image(xl_image, anchor)


def _format_numeric_columns(ws, start_row: int, end_row: int, columns: list[int]) -> None:
    for row in range(start_row, end_row + 1):
        for column in columns:
            ws.cell(row, column).number_format = "0.0000"


def _apply_r2_scale(ws, cell_range: str) -> None:
    ws.conditional_formatting.add(
        cell_range,
        ColorScaleRule(
            start_type="min",
            start_color=LIGHT_RED,
            mid_type="percentile",
            mid_value=50,
            mid_color=YELLOW,
            end_type="max",
            end_color=LIGHT_GREEN,
        ),
    )


def _apply_lower_better_scale(ws, cell_range: str) -> None:
    ws.conditional_formatting.add(
        cell_range,
        ColorScaleRule(
            start_type="min",
            start_color=LIGHT_GREEN,
            mid_type="percentile",
            mid_value=50,
            mid_color=YELLOW,
            end_type="max",
            end_color=LIGHT_RED,
        ),
    )


def _judgment(arm: str) -> str:
    return {
        "E1": "xyz DATA基线",
        "E2": "BC：总体速度近中性，边界指标改善",
        "E3": "PDE No-Go：残差下降但速度精度变差",
        "E4": "geom稳定正增量；pooled speed最佳",
        "E5": "当前推荐：case-balanced speed最佳",
        "E6": "PDE No-Go：物理一致性提升但速度精度下降",
    }[arm]


def _increment_judgment(key: str) -> str:
    return {
        "BC_xyz": "近中性：case-balanced微升，pooled微降",
        "PDE_on_BC_xyz": "No-Go：速度下降，PDE residual明显改善",
        "full_physics_xyz": "No-Go：完整physics未提升速度",
        "BC_xyz_geom": "近中性：边界指标改善，总体速度基本不变",
        "PDE_on_BC_xyz_geom": "No-Go：速度下降，PDE residual明显改善",
        "full_physics_xyz_geom": "No-Go：完整physics未提升速度",
        "geom_DATA": "Go：speed R²与MAE同步改善",
        "geom_BC": "Go：speed R²与MAE同步改善",
        "geom_BCPDE": "Go：speed R²与MAE同步改善",
    }[key]


def _teacher_sheet(wb: Workbook, primary: list[dict[str, str]]) -> None:
    ws = wb.active
    ws.title = "教师汇报"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    _style_title(ws, "WSS-PINN V3 六臂核心实验与指标汇报", 14)

    metadata = [
        ("路线", "volume_uvwp_peak_qs_smooth_v3", "科学定位", "Fluent瞬态peak标签 + Carreau–Yasuda准稳态PINN正则"),
        ("数据划分", "train123 / val15 / test35", "主checkpoint", "best_validation_data（test35不参与选模）"),
        ("完成状态", "6/6训练完成；18/18 checkpoint评估完成", "训练预算", "2500 epoch / 155000 step / 单seed1234"),
    ]
    row = 4
    for left_key, left_value, right_key, right_value in metadata:
        ws.cell(row, 1, left_key)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        ws.cell(row, 2, left_value)
        ws.cell(row, 8, right_key)
        ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=14)
        ws.cell(row, 9, right_value)
        for column in (1, 8):
            cell = ws.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            cell.font = Font(name="Microsoft YaHei", bold=True, color=NAVY)
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER
        for start, end in ((2, 6), (9, 14)):
            for column in range(start, end + 1):
                cell = ws.cell(row, column)
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = BORDER
        ws.row_dimensions[row].height = 25
        row += 1

    row = _section(ws, 8, "一、核心结论", 14)
    conclusions = [
        (LIGHT_GREEN, GREEN, "1. geom是唯一稳定正增量：E4−E1、E5−E2、E6−E3的speed R²_cb均提高约+0.19。"),
        (YELLOW, ORANGE, "2. BC对总体速度近中性，但能改善入口流量和出口压力；不能写成稳定提升speed R²。"),
        (LIGHT_RED, RED, "3. PDE首轮No-Go：continuity/momentum residual下降约89–90%/60–62%，但speed R²_cb下降约0.05。"),
        (LIGHT_PURPLE, PURPLE, "4. 训练loss已平台，但validation data/physics loss早期最优后持续变差，属于优化收敛后的泛化过拟合，不应继续加epoch。"),
        (LIGHT_ORANGE, ORANGE, "5. near-wall speed R²六臂全部为负；本实验不预测WSS，也不能形成WSS可靠性结论。"),
    ]
    for fill, color, text in conclusions:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=14)
        cell = ws.cell(row, 1, text)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color=color)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.row_dimensions[row].height = 30
        row += 1

    row = _section(ws, 15, "二、六臂主结果（best_validation_data）", 14)
    headers = [
        "臂", "输入", "模式", "speed R²_cb", "speed R²_pooled", "speed MAE", "speed RMSE",
        "near-wall R²_cb", "core R²_cb", "p R²_cb", "p R²_pooled", "continuity RMS",
        "momentum RMS(Pa/m)", "核心判定",
    ]
    _header_row(ws, row, headers)
    table_start = row + 1
    for item in primary:
        values = [
            item["arm"], item["input"], item["mode"], item["speed_case_balanced_r2"],
            item["speed_pooled_r2"], item["speed_case_balanced_mae_m_s"],
            item["speed_case_balanced_rmse_m_s"], item["near_wall_speed_case_balanced_r2"],
            item["core_speed_case_balanced_r2"], item["pressure_case_balanced_r2"],
            item["pressure_pooled_r2"], item["continuity_rms_dimensionless"],
            item["momentum_rms_pa_per_m"], _judgment(item["arm"]),
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=table_start, column=column, value=_number(value))
            _body_cell(cell, center=column != 14, wrap=column == 14)
        verdict_cell = ws.cell(table_start, 14)
        if "No-Go" in verdict_cell.value:
            verdict_cell.fill = PatternFill("solid", fgColor=LIGHT_RED)
            verdict_cell.font = Font(name="Microsoft YaHei", color=RED, bold=True)
        elif item["arm"] in {"E4", "E5"}:
            verdict_cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
            verdict_cell.font = Font(name="Microsoft YaHei", color=GREEN, bold=True)
        else:
            verdict_cell.fill = PatternFill("solid", fgColor=YELLOW)
        table_start += 1
    _format_numeric_columns(ws, row + 1, table_start - 1, list(range(4, 14)))
    _apply_r2_scale(ws, f"D{row + 1}:E{table_start - 1}")
    _apply_r2_scale(ws, f"H{row + 1}:K{table_start - 1}")
    _apply_lower_better_scale(ws, f"F{row + 1}:G{table_start - 1}")
    _apply_lower_better_scale(ws, f"L{row + 1}:M{table_start - 1}")

    row = table_start + 1
    row = _section(ws, row, "三、核心指标图", 14)
    _add_image(ws, SUMMARY / "six_arm_primary_metrics.png", f"A{row}", width=1050)
    row += 31
    row = _section(ws, row, "四、Data / Physics Loss 收敛总览", 14)
    _add_image(ws, LOSS_DIR / "teacher_loss_overview.png", f"A{row}", width=1120)

    _set_widths(
        ws,
        {1: 8, 2: 13, 3: 15, 4: 12, 5: 13, 6: 11, 7: 11, 8: 14, 9: 12,
         10: 11, 11: 12, 12: 14, 13: 17, 14: 38},
    )
    ws.freeze_panes = "A17"
    ws.auto_filter.ref = "A16:N22"
    ws.sheet_view.zoomScale = 80
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def _main_metrics_sheet(wb: Workbook, primary: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("六臂主指标")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BLUE
    headers = [
        "臂", "实验ID", "输入", "模式", "u R²_cb", "v R²_cb", "w R²_cb",
        "speed R²_cb", "speed R²_pooled", "speed MAE(m/s)", "speed RMSE(m/s)",
        "p R²_cb", "p R²_pooled", "near-wall speed R²_cb", "core speed R²_cb",
        "wall RMS(m/s)", "inlet flow ARE", "outlet p MAE(Pa)", "continuity RMS",
        "momentum RMS(Pa/m)", "正R²病例数", "最差病例", "最差speed R²",
        "speed方差比中位数", "病例间均值方差比", "判定",
    ]
    _style_title(ws, "六臂主指标明细（主checkpoint：best_validation_data）", len(headers))
    _header_row(ws, 4, headers)
    for row_index, item in enumerate(primary, 5):
        values = [
            item["arm"], item["run"], item["input"], item["mode"],
            item["u_case_balanced_r2"], item["v_case_balanced_r2"],
            item["w_case_balanced_r2"], item["speed_case_balanced_r2"],
            item["speed_pooled_r2"], item["speed_case_balanced_mae_m_s"],
            item["speed_case_balanced_rmse_m_s"], item["pressure_case_balanced_r2"],
            item["pressure_pooled_r2"], item["near_wall_speed_case_balanced_r2"],
            item["core_speed_case_balanced_r2"], item["wall_speed_rms_m_s"],
            item["inlet_flow_absolute_relative_error"], item["outlet_pressure_mae_pa"],
            item["continuity_rms_dimensionless"], item["momentum_rms_pa_per_m"],
            item["speed_positive_r2_cases"], item["speed_worst_case"],
            item["speed_worst_case_r2"], item["speed_prediction_to_truth_variance_ratio_median"],
            item["between_case_speed_mean_variance_ratio"], _judgment(item["arm"]),
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, _number(value))
            _body_cell(cell, center=column not in {2, 22, 26}, wrap=column in {22, 26})
        if row_index % 2 == 0:
            for column in range(1, len(headers) + 1):
                ws.cell(row_index, column).fill = PatternFill("solid", fgColor="F7F9FB")
    _format_numeric_columns(ws, 5, 10, list(range(5, 21)) + [23, 24, 25])
    _apply_r2_scale(ws, "E5:I10")
    _apply_r2_scale(ws, "L5:O10")
    _apply_lower_better_scale(ws, "J5:K10")
    _apply_lower_better_scale(ws, "P5:T10")
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}10"
    ws.freeze_panes = "E5"
    _set_widths(ws, {1: 7, 2: 29, 3: 13, 4: 16, 22: 32, 26: 42})
    for column in range(5, len(headers)):
        ws.column_dimensions[get_column_letter(column)].width = max(
            ws.column_dimensions[get_column_letter(column)].width or 0, 13
        )
    ws.sheet_view.zoomScale = 75
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _increment_sheet(wb: Workbook, increments: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("增量分析")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN
    _style_title(ws, "BC / PDE / geom 预注册增量分析", 18)
    metric_names = [
        "speed_case_balanced_r2", "speed_pooled_r2", "speed_case_balanced_mae_m_s",
        "speed_case_balanced_rmse_m_s", "near_wall_speed_case_balanced_r2",
        "core_speed_case_balanced_r2", "pressure_case_balanced_r2", "wall_speed_rms_m_s",
        "inlet_flow_absolute_relative_error", "outlet_pressure_mae_pa",
        "continuity_rms_dimensionless", "momentum_rms_pa_per_m",
    ]
    headers = [
        "比较", "说明", "实验臂", "对照臂", "Δspeed R²_cb", "Δspeed R²_pooled",
        "Δspeed MAE", "Δspeed RMSE", "Δnear-wall R²", "Δcore R²", "Δp R²_cb",
        "Δwall RMS", "Δinlet flow ARE", "Δoutlet p MAE", "Δcontinuity RMS",
        "Δmomentum RMS", "判定", "阅读规则",
    ]
    _header_row(ws, 4, headers)
    for row_index, item in enumerate(increments, 5):
        metrics = item["metrics"]
        deltas = [metrics[name]["delta"] for name in metric_names]
        values = [
            item["comparison"], item["label"], item["treatment"], item["control"],
            *deltas, _increment_judgment(item["comparison"]),
            "R²正值更好；MAE/RMSE/残差负值更好",
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, _number(value))
            _body_cell(cell, center=column not in {2, 17, 18}, wrap=column in {2, 17, 18})
        judgment = ws.cell(row_index, 17)
        if "No-Go" in judgment.value:
            judgment.fill = PatternFill("solid", fgColor=LIGHT_RED)
            judgment.font = Font(name="Microsoft YaHei", color=RED, bold=True)
        elif "Go" in judgment.value:
            judgment.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
            judgment.font = Font(name="Microsoft YaHei", color=GREEN, bold=True)
        else:
            judgment.fill = PatternFill("solid", fgColor=YELLOW)
        for column in range(5, 17):
            value = ws.cell(row_index, column).value
            if not isinstance(value, (int, float)):
                continue
            higher_better = column in {5, 6, 9, 10, 11}
            improved = value > 0 if higher_better else value < 0
            ws.cell(row_index, column).fill = PatternFill(
                "solid", fgColor=LIGHT_GREEN if improved else LIGHT_RED
            )
    _format_numeric_columns(ws, 5, 13, list(range(5, 17)))
    ws.auto_filter.ref = "A4:R13"
    ws.freeze_panes = "E5"
    _set_widths(ws, {1: 23, 2: 30, 3: 10, 4: 10, 17: 42, 18: 32})
    for column in range(5, 17):
        ws.column_dimensions[get_column_letter(column)].width = 15
    ws.sheet_view.zoomScale = 75
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _checkpoint_sheet(wb: Workbook, all_checkpoints: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("Checkpoint敏感性")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ORANGE
    headers = [
        "臂", "输入", "模式", "checkpoint", "speed R²_cb", "speed R²_pooled",
        "speed MAE", "speed RMSE", "p R²_cb", "p R²_pooled", "near-wall R²_cb",
        "core R²_cb", "continuity RMS", "momentum RMS(Pa/m)", "wall RMS",
        "inlet flow ARE", "outlet p MAE(Pa)", "说明",
    ]
    _style_title(ws, "三类Checkpoint敏感性（主比较固定best_validation_data）", len(headers))
    _header_row(ws, 4, headers)
    for row_index, item in enumerate(all_checkpoints, 5):
        values = [
            item["arm"], item["input"], item["mode"], item["checkpoint"],
            item["speed_case_balanced_r2"], item["speed_pooled_r2"],
            item["speed_case_balanced_mae_m_s"], item["speed_case_balanced_rmse_m_s"],
            item["pressure_case_balanced_r2"], item["pressure_pooled_r2"],
            item["near_wall_speed_case_balanced_r2"], item["core_speed_case_balanced_r2"],
            item["continuity_rms_dimensionless"], item["momentum_rms_pa_per_m"],
            item["wall_speed_rms_m_s"], item["inlet_flow_absolute_relative_error"],
            item["outlet_pressure_mae_pa"],
            "主checkpoint" if item["checkpoint"] == "best_validation_data" else (
                "敏感性：total最优" if item["checkpoint"] == "best_validation_total" else "敏感性：训练末尾"
            ),
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, _number(value))
            _body_cell(cell, center=column != 18, wrap=column == 18)
        fill = LIGHT_BLUE if item["checkpoint"] == "best_validation_data" else (
            LIGHT_PURPLE if item["checkpoint"] == "best_validation_total" else LIGHT_GRAY
        )
        for column in range(1, len(headers) + 1):
            ws.cell(row_index, column).fill = PatternFill("solid", fgColor=fill)
    _format_numeric_columns(ws, 5, 22, list(range(5, 18)))
    _apply_r2_scale(ws, "E5:F22")
    _apply_r2_scale(ws, "I5:L22")
    _apply_lower_better_scale(ws, "G5:H22")
    _apply_lower_better_scale(ws, "M5:Q22")
    ws.auto_filter.ref = "A4:R22"
    ws.freeze_panes = "E5"
    _set_widths(ws, {1: 7, 2: 13, 3: 16, 4: 24, 18: 18})
    for column in range(5, 18):
        ws.column_dimensions[get_column_letter(column)].width = 14
    ws.sheet_view.zoomScale = 75
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _convergence_sheet(wb: Workbook, convergence: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("收敛与损失")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PURPLE
    _style_title(ws, "Data Loss / Physics Loss 收敛监控", 15)
    row = _section(ws, 4, "损失定义", 15, color=PURPLE)
    formulas = [
        "L_data = (L_u + L_v + L_w + L_p) / 4",
        "L_BC = (L_wall + L_inlet + L_outlet) / 3",
        "L_PDE = (L_cont + (L_mx + L_my + L_mz) / 3) / 2",
        "L_phy = lambda_BC * L_BC + lambda_PDE * L_PDE；本轮lambda_BC=lambda_PDE=1",
    ]
    for formula in formulas:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=15)
        cell = ws.cell(row, 1, formula)
        cell.font = Font(name="Consolas", size=11, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=LIGHT_PURPLE)
        cell.alignment = Alignment(vertical="center")
        cell.border = BORDER
        row += 1
    row += 1
    headers = [
        "臂", "输入", "模式", "best val data epoch", "best val data loss",
        "last val data loss", "last/best data", "最后10% data变化(%)",
        "best val phy epoch", "best val phy loss", "last val phy loss",
        "last/best phy", "最后10% phy变化(%)", "late clip均值", "收敛判读",
    ]
    _header_row(ws, row, headers, fill=PURPLE)
    start = row + 1
    for row_index, item in enumerate(convergence, start):
        values = [
            item["arm"], item["input"], item["mode"], item["best_validation_data_epoch"],
            item["best_validation_data_loss"], item["last_validation_data_loss"],
            item["data_last_to_best_ratio"], item["last_10pct_validation_data_change_percent"],
            item["best_validation_physics_epoch"], item["best_validation_physics_loss"],
            item["last_validation_physics_loss"], item["physics_last_to_best_ratio"],
            item["last_10pct_validation_physics_change_percent"],
            item["late_gradient_clipping_fraction_mean"],
            "训练已平台；validation早期最优后变差，属于过拟合，不应继续加epoch",
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, _number(value))
            _body_cell(cell, center=column != 15, wrap=column == 15)
        for column in (7, 12):
            value = ws.cell(row_index, column).value
            if isinstance(value, (int, float)) and value > 1.05:
                ws.cell(row_index, column).fill = PatternFill("solid", fgColor=LIGHT_RED)
                ws.cell(row_index, column).font = Font(name="Microsoft YaHei", color=RED, bold=True)
        ws.cell(row_index, 15).fill = PatternFill("solid", fgColor=YELLOW)
    end = start + len(convergence) - 1
    _format_numeric_columns(ws, start, end, [5, 6, 7, 8, 10, 11, 12, 13, 14])
    ws.auto_filter.ref = f"A{row}:O{end}"
    ws.freeze_panes = f"D{start}"
    row = end + 2
    row = _section(ws, row, "六臂Data / Physics Loss总览", 15, color=PURPLE)
    _add_image(ws, LOSS_DIR / "teacher_loss_overview.png", f"A{row}", width=1200)
    _set_widths(ws, {1: 7, 2: 13, 3: 16, 4: 18, 5: 17, 6: 17, 7: 14, 8: 20,
                     9: 18, 10: 17, 11: 17, 12: 14, 13: 20, 14: 14, 15: 48})
    ws.sheet_view.zoomScale = 75
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _case_sheet(wb: Workbook, cases: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("逐病例")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GRAY
    fields = [
        ("arm", "臂"), ("input", "输入"), ("mode", "模式"), ("case_id", "病例"),
        ("evaluated_points", "评估点数"), ("speed_r2", "speed R²"),
        ("speed_mae_m_s", "speed MAE"), ("speed_rmse_m_s", "speed RMSE"),
        ("near_wall_speed_r2", "near-wall R²"), ("core_speed_r2", "core R²"),
        ("pressure_r2", "p R²"), ("pressure_mae_pa", "p MAE(Pa)"),
        ("pressure_gauge_error_pa", "pressure gauge误差(Pa)"),
        ("wall_speed_rms_m_s", "wall RMS"),
        ("inlet_flow_absolute_relative_error", "inlet flow ARE"),
        ("outlet_pressure_mae_pa", "outlet p MAE(Pa)"),
        ("continuity_rms_dimensionless", "continuity RMS"),
        ("momentum_rms_pa_per_m", "momentum RMS(Pa/m)"),
        ("speed_prediction_variance_ratio", "speed预测/真值方差比"),
    ]
    _style_title(ws, "六臂 × test35 逐病例指标（best_validation_data）", len(fields))
    _header_row(ws, 4, [label for _, label in fields])
    for row_index, item in enumerate(cases, 5):
        for column, (field, _) in enumerate(fields, 1):
            cell = ws.cell(row_index, column, _number(item[field]))
            _body_cell(cell, center=column != 4, wrap=column == 4)
        if row_index % 2 == 0:
            for column in range(1, len(fields) + 1):
                ws.cell(row_index, column).fill = PatternFill("solid", fgColor="F7F9FB")
    end = 4 + len(cases)
    _format_numeric_columns(ws, 5, end, list(range(6, len(fields) + 1)))
    _apply_r2_scale(ws, f"F5:F{end}")
    _apply_r2_scale(ws, f"I5:K{end}")
    _apply_lower_better_scale(ws, f"G5:H{end}")
    _apply_lower_better_scale(ws, f"L5:R{end}")
    ws.auto_filter.ref = f"A4:{get_column_letter(len(fields))}{end}"
    ws.freeze_panes = "F5"
    _set_widths(ws, {1: 7, 2: 13, 3: 16, 4: 34, 5: 13})
    for column in range(6, len(fields) + 1):
        ws.column_dimensions[get_column_letter(column)].width = 16
    ws.sheet_view.zoomScale = 70


def _provenance_sheet(wb: Workbook, output_path: Path) -> None:
    ws = wb.create_sheet("口径与文件")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GRAY
    _style_title(ws, "评估口径、限制与文件索引", 6)
    entries = [
        ("科学定位", "Fluent瞬态peak标签 + Carreau–Yasuda准稳态PINN正则；不是稳态CFD surrogate"),
        ("输出", "u/v/w/p；本路线不预测WSS"),
        ("split", "train123 / val15 / test35；SHA256 c80cb65ad95f7d76fadaff82ea02c35c27ad24ab2474e1050bf981eff97493f9"),
        ("主checkpoint", "best_validation_data；best_validation_total与last仅作敏感性分析"),
        ("test35规则", "不参与训练、选模、续训判断或checkpoint反选"),
        ("case-balanced", "每例先计算指标，再对35例等权平均"),
        ("pooled", "合并test35全部评估点后计算；易受大病例点数影响，必须与case-balanced并列报告"),
        ("出口流量限制", "冻结资产有exact outlet pressure，但没有exact outlet flow target/面法向/面积权重；PCA质量守恒仅为诊断"),
        ("收敛判读", "训练loss平台；validation早期最优后变差，属于泛化过拟合，不支持续训到5000"),
        ("单seed限制", "本轮只有seed1234；后续确认实验需在修正物理合同后再做多seed"),
        ("工作簿", str(output_path)),
        ("主结果README", str(SUMMARY / "README.md")),
        ("损失收敛README", str(LOSS_DIR / "README.md")),
        ("9页损失PDF", str(LOSS_DIR / "teacher_loss_convergence.pdf")),
        ("六臂主表", str(SUMMARY / "six_arm_metrics_best_validation_data.csv")),
        ("增量JSON", str(SUMMARY / "incremental_effects_best_validation_data.json")),
        ("逐病例CSV", str(SUMMARY / "case_metrics_best_validation_data.csv")),
    ]
    _header_row(ws, 4, ["项目", "内容", "说明/备注", "文件是否存在", "大小(bytes)", "更新时间"])
    for row_index, (key, value) in enumerate(entries, 5):
        path = Path(value) if value.startswith("/") else None
        exists = path.is_file() if path else "—"
        size = path.stat().st_size if path and path.is_file() else "—"
        mtime = (
            __import__("datetime").datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
            if path and path.is_file()
            else "—"
        )
        values = [key, value, "", exists, size, mtime]
        for column, item in enumerate(values, 1):
            cell = ws.cell(row_index, column, item)
            _body_cell(cell, center=column in {1, 4, 5, 6}, wrap=column in {2, 3})
        ws.row_dimensions[row_index].height = 30
    _set_widths(ws, {1: 20, 2: 95, 3: 35, 4: 14, 5: 16, 6: 22})
    ws.freeze_panes = "A5"
    ws.sheet_view.zoomScale = 85


def build(output_path: Path) -> None:
    primary = _read_csv(SUMMARY / "six_arm_metrics_best_validation_data.csv")
    all_checkpoints = _read_csv(SUMMARY / "six_arm_metrics_all_checkpoints.csv")
    increments = _read_json(SUMMARY / "incremental_effects_best_validation_data.json")
    cases = _read_csv(SUMMARY / "case_metrics_best_validation_data.csv")
    convergence = _read_csv(LOSS_DIR / "teacher_loss_convergence_summary.csv")
    if not (len(primary) == 6 and len(all_checkpoints) == 18 and len(increments) == 9 and len(cases) == 210 and len(convergence) == 6):
        raise RuntimeError("summary input counts do not match the frozen six-arm contract")

    workbook = Workbook()
    workbook.properties.title = "WSS-PINN V3 六臂核心实验与指标汇报"
    workbook.properties.subject = "Fluent瞬态peak标签 + Carreau–Yasuda准稳态PINN正则"
    workbook.properties.creator = "Codex"
    workbook.properties.description = "六臂主指标、增量、checkpoint敏感性、收敛与逐病例结果"

    _teacher_sheet(workbook, primary)
    _main_metrics_sheet(workbook, primary)
    _increment_sheet(workbook, increments)
    _checkpoint_sheet(workbook, all_checkpoints)
    _convergence_sheet(workbook, convergence)
    _case_sheet(workbook, cases)
    _provenance_sheet(workbook, output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix="wss_pinn_v3_", suffix=".xlsx", dir=output_path.parent, delete=False
    ) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        os.replace(temporary, output_path)
        output_path.chmod(0o664)
    finally:
        temporary.unlink(missing_ok=True)

    checked = load_workbook(output_path, read_only=False, data_only=False)
    expected = ["教师汇报", "六臂主指标", "增量分析", "Checkpoint敏感性", "收敛与损失", "逐病例", "口径与文件"]
    if checked.sheetnames != expected:
        raise RuntimeError(f"workbook sheet contract drift: {checked.sheetnames}")
    if checked["六臂主指标"].max_row != 10 or checked["逐病例"].max_row != 214:
        raise RuntimeError("workbook row-count validation failed")
    if len(checked["教师汇报"]._images) != 2 or len(checked["收敛与损失"]._images) != 1:
        raise RuntimeError("workbook embedded-image validation failed")
    checked.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output = args.output.resolve()
    build(output)
    print(json.dumps({"status": "completed", "output": str(output)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
